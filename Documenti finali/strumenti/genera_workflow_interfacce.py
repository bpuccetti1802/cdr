# -*- coding: utf-8 -*-
"""Genera «WORKFLOW_Interfacce-SIPO-ANSC_v0.1» — le interfacce del workflow, pagina per pagina.

Richiesta dell'utente (21/10/2026): un documento dedicato alla gestione delle interfacce del
workflow, con un capitolo per ciascun punto del workflow di img/workflow_ansc_sipo.pptx; per
ogni fase gli stati in SIPO e in ANSC, le pagine, i servizi ANSC chiamati e le tabelle SIPO
coinvolte; ripresa delle informazioni del capitolo 13 dell'analisi; un diagramma a stati.

Fonti, lette a ogni esecuzione (una sola fonte per ciascun dato):
  · img/workflow_ansc_sipo.pptx — i punti del workflow e il servizio di ciascuno;
  · ANALISI_Integrazione-ANSC_v3.44.docx — la tabella del cap. 13 (testo e immagini delle
    celle), le schermate di Side del 13.1, la tabella degli stati di SIPO;
  · ansc/docs/Decodifiche/11_*.csv — gli stati ANSC.
Le tabelle SIPO di ciascuna pagina vengono dal codice (front-end statocivile-web e nascita-web,
back-end statocivile-be e nascita-be) e dalla MAPPA delle nascite: le schermate del cap. 13 sono
di una nascita, che fa da atto di esempio.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/genera_workflow_interfacce.py
"""
import csv
import glob
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D                 # noqa: E402
import diagrammi_workflow_stati         # noqa: E402
import v3_43_evento as V43              # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
STR = os.path.join(FIN, 'strumenti')
IMG = os.path.join(STR, 'img')
TEMPLATE = os.path.join(BASE, 'Template documentale', 'template_DAD_roma-capitale.docx')
ANALISI = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.44.docx')
PPTX = os.path.join(IMG, 'workflow_ansc_sipo.pptx')
VERSIONE = '0.1'
NOME = f'WORKFLOW_Interfacce-SIPO-ANSC_v{VERSIONE}'
OUT = os.path.join(FIN, NOME + '.docx')
DATA = '21/10/2026'
TAGLIA_DA = 29
TMP = tempfile.mkdtemp(prefix='wf_')


# ------------------------------------------------------------------------------ le fonti
def punti_pptx():
    """I box del PowerPoint, per riga: pagina (colonna sinistra) e servizio (colonna 2)."""
    s = Presentation(PPTX).slides[0]
    box = [(int(x.top / 12700), int(x.left / 12700),
            re.sub(r'[\x00-\x08\x0b-\x1f]', ' ', x.text_frame.text).strip())
           for x in s.shapes if x.has_text_frame and x.text_frame.text.strip()]
    pagine = sorted([b for b in box if -1 < b[1] < 40 and b[0] > 80 and b[0] < 600],
                    key=lambda b: b[0])
    servizi = [b for b in box if 200 <= b[1] <= 215 and 80 < b[0] < 600]
    out = []
    for top, _, testo in pagine:
        serv = [t for tp, _, t in servizi if abs(tp - top) <= 12]
        out.append((testo.split('\n')[0].strip(), ' / '.join(x.replace('\n', ' · ')
                                                               for x in serv) or '—'))
    return out


def figura_pptx():
    """La diapositiva in alta risoluzione (Quick Look rende la prima diapositiva)."""
    shutil.copy(PPTX, os.path.join(TMP, 'wf.pptx'))
    subprocess.run(['qlmanage', '-t', '-s', '2400', '-o', TMP, os.path.join(TMP, 'wf.pptx')],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
    png = os.path.join(TMP, 'wf.pptx.png')
    if not os.path.exists(png):                      # ripiego: l'esportazione dell'utente
        png = os.path.join(IMG, 'workflow_ansc_sipo', 'Diapositiva1.png')
    return png


def analisi():
    """Dalla v3.44: righe del cap. 13 (testo e immagini) e schermate del 13.1."""
    d = docx.Document(ANALISI)
    t = [x for x in d.tables if x.rows[0].cells[0].text.strip() == 'Passaggio'][0]
    righe = []
    for ri, r in enumerate(t.rows[1:], start=1):
        celle = [c.text for c in r.cells]
        imm = []
        for k, b in enumerate(r.cells[1]._tc.findall('.//' + qn('a:blip'))):
            f = os.path.join(TMP, f'cap13_r{ri}_{k}.png')
            open(f, 'wb').write(d.part.related_parts[b.get(qn('r:embed'))].blob)
            imm.append(f)
        righe.append(dict(passaggio=pulisci(celle[0]), testo=pulisci(celle[1]),
                          ansc=pulisci(celle[2]), sipo=pulisci(celle[3]), immagini=imm))
    ps = d.paragraphs
    i0 = [k for k, p in enumerate(ps) if p.style.name == 'Heading 2'
          and p.text.startswith('Processo ANSC di validazione')][0]
    schermate = []
    for k in range(i0, i0 + 25):
        for j, b in enumerate(ps[k]._p.findall('.//' + qn('a:blip'))):
            blob = d.part.related_parts[b.get(qn('r:embed'))].blob
            if Image.open(io.BytesIO(blob)).size[0] < 400:
                continue                              # il ritaglio dell'identificativo
            f = os.path.join(TMP, f'side_{k}_{j}.png')
            open(f, 'wb').write(blob)
            schermate.append(f)
    stati_sipo = [x for x in d.tables if [c.text.strip() for c in x.rows[0].cells][:2]
                  == ['ID', 'Significato']][0]
    sipo = [[c.text.strip() for c in r.cells] for r in stati_sipo.rows]
    return righe, schermate, sipo


def pulisci(s):
    s = re.sub(r'\n\s*\n+', '\n', s.strip())
    return s.replace(' ', ' ')


def stati_ansc():
    f = glob.glob(os.path.join(BASE, 'ansc', 'docs', 'Decodifiche', '11_*.csv'))[0]
    return [(r['ID'], r['DESCRIZIONE']) for r in csv.DictReader(open(f, encoding='utf-8'))]


# ------------------------------------------------------------------------------ i contenuti
NEL_WF = {
    '1': 'Solo dalla web app (bozza) o per la prenotazione del parto plurimo; nel nostro '
         'percorso compare come 0, ad uso applicativo.',
    '2': 'Dopo la validazione R009 riuscita.',
    '3': 'Dopo R011 su un atto non formato, o dopo la formazione con una rettifica di tipo '
         'annullamento.',
    '4': 'Dopo le firme dei dichiaranti, R006 o R012.',
    '5': 'Dopo la firma dell’USC, R007 /firma_usc: l’atto è formato.',
    '6': 'Non prodotto dal percorso.',
    '7': 'Non prodotto dal percorso (annotazioni).',
    '8': 'Solo con il provvedimento di rifiuto, R016.',
    '9': 'Cancellazione della firma dell’ufficiale, R011.',
    '10': 'Annullamento per inefficacia, solo dalla web app.',
    '11': 'Dopo l’attestazione di conformità, R007 /attestazione/firma.',
}

LOCALI = [
    ['Colonna di ANSC_STATO_ATTO', 'Valori', 'Che cosa dice'],
    ['STATO', 'i valori di ANSC_11, più IN_PREPARAZIONE e BOZZA_APPLICATIVA (0)',
     'Lo stato dell’atto in ANSC. IN_PREPARAZIONE: in ANSC non c’è ancora nulla. 0: in ANSC '
     'c’è un atto confermato ma i dati sono cambiati dopo la validazione e vanno validati di '
     'nuovo; per ANSC coincide con 1 BOZZA.'],
    ['COD_FASE', 'IN_LAVORAZIONE, UC_DETERMINATO, DOCUMENTI_ACQUISITI, PREVALIDATO, VALIDATO, '
     'FIRMATO, ANNULLATO', 'A che punto è la nostra lavorazione. Non si confonde con STATO: una '
     'sola delle due colonne ci appartiene.'],
    ['COD_PAGINA_RIPRESA', 'DETTAGLIO, RIEPILOGO, ANTEPRIMA, VALIDAZIONE, FIRME_DICHIARANTI, '
     'CONFORMITA, FIRMA_USC', 'L’ultima pagina raggiunta: è quella che si riapre quando la '
     'redazione sospesa si riprende.'],
    ['FLG_SOSPESO', 'S, N', 'Se la redazione è sospesa.'],
]

# Le fasi. «riga» = righe della tabella del cap. 13 che vi corrispondono; «pptx» = indice del
# punto del PowerPoint; «schermate» = indici delle schermate di Side del 13.1.
FASI = [
    dict(
        titolo='Accesso e scelte iniziali', pptx=0, righe=[0, 1, 2], schermate=[],
        sintesi='Login, scelta della tipologia di gestione e del tipo di atto: le pagine sono '
                'quelle di oggi e restano invariate. È qui che si fissano il tipo atto e il '
                'Modello, da cui la configurazione ricava gli UC candidati; l’UC si determinerà '
                'più avanti, nel riepilogo.',
        pagine=[['Menu di SIPO', 'statocivile-web, menu dello stato civile', 'invariata'],
                ['Scelta della tipologia di atto', 'statocivile-web, /sceltaTipologiaAtto '
                 '(SceltaTipologiaAttoController)', 'invariata']],
        servizi=[['—', 'Nessuna chiamata ad ANSC.']],
        sipo=[['V_CONF_TIPO_ATTI', 'lettura', '/cercaTipoAttoForNasciteFromVistaRefactored, '
               '/findTipoAttoByAreaAndSottoareaAndIdUtenteAndTipoRito (statocivile-be, '
               'ConfTipoAttiController)'],
              ['CONF_TIPO_ATTI', 'lettura', '/getConfTipoAttoById: il tipo atto con Modello '
               '(ID_MODELLO_ATTO) e maschera (MASCHERA_UI)'],
              ['ATTO, ATTO_NASCITA, CONF_PARTE_SERIE, ESPONENTE', 'scrittura', '/saveAtto '
               '(nascita-be, SalvataggioTabAttoController): la pratica nasce con un numero '
               'provvisorio (PKG_PRATICA.get_temp_numatto)']],
        ansc_usr=[['ANSC_CFG_UC', 'lettura', 'Gli UC candidati per il Modello; nessuna '
                   'scrittura.']],
        stati=('—', '1 IN LAVORAZIONE, alla creazione della pratica', 'nessun atto',
               'nessun atto', '— (la riga dell’atto nasce nel dettaglio)'),
        transizioni=['Avanti: alla ricerca dei dati anagrafici.'],
        aperti=['La figura pone lo stato SIPO a «N.A.» per questo punto e fa nascere lo stato 1 '
                'alla ricerca dei dati anagrafici; la tabella del capitolo 13 lo pone già alla '
                'scelta del tipo di atto, ed è lì che il codice crea la pratica. Va allineata '
                'una delle due.']),
    dict(
        titolo='Ricerca dati anagrafici', pptx=1, righe=[3], schermate=[],
        sintesi='Si cercano nell’anagrafe locale di SIPO i soggetti dell’atto, fra cui '
                'l’intestatario. Nessuna chiamata ad ANSC: l’identificazione nazionale avviene '
                'nella pagina successiva.',
        pagine=[['Ricerca del soggetto', 'statocivile-web, /ricercaSoggetto '
                 '(RicercaSoggettoController); per la nascita i frammenti dei genitori di '
                 'nascita-web', 'invariata, salvo l’indicazione dell’intestatario']],
        servizi=[['—', 'Nessuna chiamata ad ANSC: «Anagrafe locale SIPO».']],
        sipo=[['SOGGETTO (MATR_USR)', 'lettura e scrittura', '/getRicercheAnagraficheStatoCivile, '
               '/insertSoggetto (statocivile-be, SoggettoController e '
               'GestioneSoggettoController)'],
              ['ANAG_USR.SOGGETTO', 'lettura', '/findIdSoggettoFromAnagrafe: il legame con '
               'l’anagrafe per CODICE_INDIVIDUALE']],
        ansc_usr=[['SOGGETTI_USECASE', 'lettura', 'Chi è l’intestatario per l’UC, secondo il '
                   'capitolo 13.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', 'nessun atto', 'nessun atto', '—'),
        transizioni=['Avanti: al dettaglio.', 'Indietro: alla scelta del tipo di atto.'],
        aperti=['Il capitolo 13 indica SOGGETTI_USECASE per sapere chi sia l’intestatario, ma '
                'quella tabella è per UC e l’UC si determina solo nel riepilogo: qui sono noti '
                'gli UC candidati del Modello, non l’UC. Va stabilito se l’intestatario si '
                'ricava dai candidati (lo stesso per tutti?) o da un dato del Modello.']),
    dict(
        titolo='Dettaglio (creazione atto)', pptx=2, righe=[4], schermate=[],
        sintesi='La maschera dell’atto, con i dati dei soggetti. Si aggiungono, per ciascun '
                'soggetto, la ricerca in ANSC e i tre codici che ne risultano: identificativo '
                'della scheda ANPR, identificativo unico nazionale, identificativo del soggetto '
                'in ANSC. Non si prosegue se mancano dati obbligatori.',
        pagine=[['Dettaglio dell’atto', 'nascita-web, maschere AttoNascitaTipoNN (estendono '
                 'AttoNascitaTipo00Controller)', 'modificata: pulsante di ricerca ANSC per '
                 'soggetto e i tre codici nella scheda']],
        servizi=[['R005', 'Consultazione dei soggetti, per ciascun soggetto; richiede l’OTP di '
                  'sessione. Per l’intestatario la ricerca è obbligatoria, non il risultato.']],
        sipo=[['ATTO, ATTO_NASCITA, ATTO_NASCITA_SOGGETTO, ATTO_NASCITA_EXTEND, SOGGETTO',
               'scrittura', '/saveAttoNascita (nascita-be, SalvataggioController): il '
               'salvataggio dell’atto di nascita'],
              ['ANAG_USR.SOGGETTO', 'scrittura, solo se vuote', 'ID_SCHEDA_ANPR e '
               'ID_UNICO_NAZIONALE dalla risposta di R005']],
        ansc_usr=[['ANSC_SOGGETTO_IDENTIFICATIVI', 'scrittura', 'L’identificativo ANSC del '
                   'soggetto e l’esito della ricerca.'],
                  ['ANSC_STATO_ATTO', 'scrittura', 'Nasce la riga dell’atto: COD_FASE '
                   'IN_LAVORAZIONE, COD_PAGINA_RIPRESA DETTAGLIO.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase SOGGETTO, con richiesta e risposta.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', 'nessun atto', 'nessun atto',
               'DETTAGLIO'),
        transizioni=['Avanti: al riepilogo, solo se le ricerche obbligatorie sono state fatte e '
                     'non mancano dati obbligatori.', 'Indietro: alla ricerca dei dati '
                     'anagrafici.'],
        aperti=['La figura rende obbligatoria la ricerca per ciascun soggetto; la tabella del '
                'capitolo 13 per il solo intestatario (OP-76 e).',
                'La ricerca richiede l’OTP: se è obbligatoria, il token non può più essere '
                'facoltativo, mentre il percorso in otto passi dell’analisi lo dichiara tale '
                '(OP-76 d).',
                'ANSC può restituire più soggetti fra cui scegliere, e la scelta può '
                'invalidare dati già scritti in SIPO.']),
    dict(
        titolo='Riepilogo atto', pptx=3, righe=[5], schermate=[],
        sintesi='Si determina l’UC e si mostrano codice e descrizione (ANSC_03), i firmatari con '
                'il loro contrassegno, i documenti previsti e i dati ANSC mancanti. Non si '
                'prosegue finché la pre-verifica non è soddisfatta; nella prima fase il '
                'controllo può essere semplificato. Se l’atto ha già un identificativo nazionale '
                '— si torna indietro dopo la validazione — lo si mostra.',
        pagine=[['Riepilogo dell’atto', 'statocivile-web, /dettaglioAttoStatoCivile '
                 '(DettaglioAttoStatoCivileController)', 'modificata, o una pagina nuova '
                 'richiamata da un pulsante: la scelta è dei realizzatori']],
        servizi=[['—', 'Nessuna chiamata obbligatoria: determinazione dell’UC e pre-verifica '
                  'sono locali.'],
                 ['R001', 'Invio dei documenti, se avviene qui (OP-76 a).']],
        sipo=[['V_RIEPILOGO_ATTO_GEN, GTT_RIEPILOGO_ATTO', 'lettura', '/getRiepilogoAttoGen '
               '(statocivile-be, DettaglioAttiController)'],
              ['V_DETTAGLIO_ATTO_GEN, GTT_DETTAGLIO_ATTO', 'lettura', '/getDataDettaglioAtto'],
              ['ATTO', 'lettura', '/getStatoByIdAtto: lo stato della pratica'],
              ['ALLEGATO (MATR_USR)', 'scrittura', 'I documenti caricati, nello stato 0 '
               '«Caricato in SIPO»']],
        ansc_usr=[['ANSC_CFG_UC, LOGICHE_DI_SCELTA', 'lettura', 'La determinazione dell’UC.'],
                  ['ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO, VALORE_DOMINIO', 'lettura',
                   'La pre-verifica dei dati.'],
                  ['ALLEGATI_USECASE, SOGGETTI_USECASE', 'lettura', 'I documenti e i '
                   'firmatari previsti dall’UC.'],
                  ['firmatari', 'scrittura', 'Chi firma e in quale modo.'],
                  ['ANSC_STATO_ATTO', 'scrittura', 'ID_UC_ANSC, COD_LOGICA, COD_ORIGINE_UC, '
                   'ID_VERSIONE; COD_FASE UC_DETERMINATO; COD_PAGINA_RIPRESA RIEPILOGO.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fasi DETERMINAZIONE e PREVERIFICA.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', 'nessun atto (o 2 se si torna dopo la '
               'validazione)', 'nessun atto (o 0 se i dati sono cambiati)', 'RIEPILOGO'),
        transizioni=['Avanti: all’anteprima, con la pre-verifica soddisfatta.',
                     'Indietro: al dettaglio.'],
        aperti=['In quale pagina parta l’invio dei documenti (OP-76 a).',
                'Il capitolo 13 chiede di prevedere le operazioni del concentratore per salvare '
                'i dati e restituire il codice (commento dell’utente): non sono ancora nel '
                'contratto ansc-v1 (OP-75 b).']),
    dict(
        titolo='Anteprima SIPO', pptx=4, righe=[6], schermate=[],
        sintesi='Si costruisce il testo dell’atto e si acquisisce il numero comunale definitivo, '
                'quello della numerazione digitale (lo stesso della web app): senza di esso '
                'l’atto non si salva in ANSC, e va mostrato. Si aggiunge il pulsante «Valida». '
                'Per un atto digitale la posizione dell’«occhio» non serve più, salvo per il '
                'registro di emergenza.',
        pagine=[['Anteprima dell’atto', 'statocivile-web, /anteprimaAttoStatoCivile '
                 '(AnteprimaAttoController)', 'modificata: numero comunale e pulsante «Valida»']],
        servizi=[['—', 'Nessuna chiamata ad ANSC (R010, l’anteprima di ANSC, resta facoltativa).'],
                 ['num-ansc-sipo', 'Lo stacco del numero comunale: è un servizio del Comune, '
                  'non di ANSC.']],
        sipo=[['ATTO', 'lettura e scrittura', '/generaAnteprimaAtto, '
               '/salvaPosizioneOcchio e /recuperaPosizioneOcchio (statocivile-be, '
               'AnteprimaAttoController)'],
              ['ATTO (numero)', 'scrittura', 'Il numero comunale va salvato anche sull’atto di '
               'SIPO: la colonna è da decidere (OP-60)']],
        ansc_usr=[['PRENOTAZIONE_ATTI, DATI_ATTO_ESTERNO', 'scrittura', 'Il numero staccato, '
                   'nello stato 1 Disponibile.'],
                  ['ANSC_STATO_ATTO', 'scrittura', 'NUM_COMUNALE; TXT_EVENTO con il modello '
                   'evento costruito; COD_PAGINA_RIPRESA ANTEPRIMA.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase PAYLOAD.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', 'nessun atto', 'nessun atto',
               'ANTEPRIMA'),
        transizioni=['Valida: alla pagina di validazione.', 'Indietro: al riepilogo.'],
        aperti=['La figura colloca qui utenza, PIN e OTP del servizio di firma remota, che '
                'servono invece alle firme R007 (OP-76 b).',
                'Il pulsante «Definisci» di questa pagina: la tabella del capitolo 13 lo '
                'rinomina «stampe» (OP-76 c).']),
    dict(
        titolo='Anteprima · Validazione', pptx=5, righe=[], schermate=[],
        sintesi='Pagina nuova, simile a quella di Side: si prepara il modello evento e lo si '
                'invia a R009, che valida metadati e allegati. Finché la validazione non riesce '
                'si torna a qualunque pagina precedente a correggere. Se riesce, R009 '
                'restituisce l’identificativo nazionale dell’atto e quelli dei soggetti che '
                'ANSC conosce, che si salvano.',
        pagine=[['Validazione', 'nuova', 'da realizzare, sul modello di Side']],
        servizi=[['R009', '/validazione/evento: deposita l’evento e restituisce idAnsc. Una '
                  'risposta 400 non crea nulla in ANSC.']],
        sipo=[['ANAG_USR.ATTO (ID_ATTO_ANSC, ANNO_ANSC, DATA_ANSC)', 'da decidere',
               'Le colonne ANSC che SIPO ha già sull’atto di anagrafe (OP-29)']],
        ansc_usr=[['ANSC_STATO_ATTO', 'scrittura', 'ID_ANSC, ID_EVENTO_ANSC, STATO CONFERMATO, '
                   'TXT_HASH_VALIDATO, DATA_VALIDAZIONE; COD_FASE VALIDATO; COD_PAGINA_RIPRESA '
                   'VALIDAZIONE.'],
                  ['ANSC_SOGGETTO_IDENTIFICATIVI', 'scrittura', 'Gli identificativi dei '
                   'soggetti restituiti.'],
                  ['forzatura_validazione', 'scrittura', 'Il codice di forzatura, se ANSC '
                   'lo concede.'],
                  ['DATI_ATTO_ESTERNO', 'scrittura', 'Il numero comunale passa a 2 Utilizzato, '
                   'con l’identificativo nazionale.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase DEPOSITO.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', 'nessun atto (o 0)', '2 CONFERMATO',
               'VALIDAZIONE'),
        transizioni=['R009 accettata: alle firme dei dichiaranti, o alla conformità se non ci '
                     'sono firmatari.', 'R009 respinta: si torna alle pagine dei dati, si '
                     'corregge e si ripete.', 'Timeout: esito indeterminato, si accerta con '
                     'R005; mai un nuovo invio alla cieca.'],
        aperti=['Con quale campo, id o idAnsc, una nuova validazione riconosca l’atto già '
                'confermato, dopo un ritorno sui dati (OP-75 a, OP-46).']),
    dict(
        titolo='Anteprima · Firme dei dichiaranti', pptx=6, righe=[], schermate=[0, 1],
        sintesi='Pagina nuova: si fanno firmare i dichiaranti, in modo digitale o caricando il '
                'cartaceo firmato e scansionato, che si associa al firmatario. Se non ci sono '
                'firmatari la pagina si salta.',
        pagine=[['Firme dei dichiaranti', 'nuova', 'da realizzare, sul modello di Side']],
        servizi=[['R006', 'Firma cartacea: caricamento del verbale firmato.'],
                 ['R012', 'Firma elettronica del dichiarante.']],
        sipo=[['—', '—', 'Nessuna tabella di SIPO: le firme si registrano nelle tabelle '
               'dell’integrazione']],
        ansc_usr=[['firmatari', 'scrittura', 'Lo stato della firma di ciascun firmatario.'],
                  ['ANSC_STATO_ATTO', 'scrittura', 'STATO FIRMATO_DICHIARANTE; '
                   'COD_PAGINA_RIPRESA FIRME_DICHIARANTI.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase FIRMA_DICH.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', '2 CONFERMATO',
               '4 FIRMATO DAL DICHIARANTE', 'FIRME_DICHIARANTI'),
        transizioni=['Avanti: alla conformità.', 'Dalla firma del dichiarante i dati sono '
                     'vincolati: correggerli significa tornare alla validazione e far firmare '
                     'di nuovo [DA VERIFICARE].'],
        aperti=['Il DDL di firmatari, ereditato da Side, va corretto prima dell’adozione '
                '(OP-75 d).']),
    dict(
        titolo='Anteprima · Conformità', pptx=7, righe=[], schermate=[2],
        sintesi='Si genera l’attestazione di conformità degli allegati e l’USC la firma '
                'digitalmente.',
        pagine=[['Attestazione di conformità', 'nuova', 'da realizzare, sul modello di Side']],
        servizi=[['R007', '/attestazione/firma, con le credenziali di firma remota '
                  '(utenza, PIN, OTP del servizio di firma).']],
        sipo=[['ALLEGATO (MATR_USR)', 'lettura e scrittura', 'Quali documenti includere '
               'nell’attestazione (fg_incluso_att_conformita)']],
        ansc_usr=[['ANSC_STATO_ATTO', 'scrittura', 'STATO FIRMATO_CONFORMITA; '
                   'COD_PAGINA_RIPRESA CONFORMITA.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase FIRMA_CONFORMITA.']],
        stati=('1 IN LAVORAZIONE', '1 IN LAVORAZIONE', '4 (o 2 senza firmatari)',
               '11 FIRMATO PER CONFORMITÀ', 'CONFORMITA'),
        transizioni=['Avanti: alla firma dell’USC.'],
        aperti=[]),
    dict(
        titolo='Anteprima · Firma dell’USC', pptx=8, righe=[], schermate=[3],
        sintesi='L’USC firma l’atto finito. La firma lo salva in ANSC e determina il '
                'salvataggio in SIPO: da qui l’atto non si modifica più. Il contrassegno '
                '«Comunicazione anagrafica» comanda la chiamata a R014.',
        pagine=[['Firma dell’USC', 'nuova', 'da realizzare, sul modello di Side; pulsante '
                 '«Chiudi» secondo la figura']],
        servizi=[['R007', '/firma_usc, con le credenziali di firma remota.'],
                 ['R014', 'L’aggiornamento anagrafico (data di decorrenza), se richiesto.']],
        sipo=[['ATTO, CONF_STATO_PRATICA', 'scrittura', '/finalizzaAtto (statocivile-be, '
               'AnteprimaAttoController → InitFlow): stato 2 DEFINITO e DATA_DEFINIZIONE_ATTO, '
               'oppure 3 se l’aggiornamento anagrafico non riesce'],
              ['CONF_ATTO_FLUSSI_ELABORAZIONE, ATTO_FLUSSI_ELABORAZIONE, '
               'ATTO_FLUSSI_ELABORAZIONE_LOG', 'lettura e scrittura', 'I flussi '
               'dell’aggiornamento anagrafico (InitFlow)'],
              ['PKG_PRATICA', 'procedure', 'IntegrazioneATTO assegna il numero definitivo; '
               'CompletaATTO porta l’atto a 2']],
        ansc_usr=[['ANSC_STATO_ATTO', 'scrittura', 'STATO FIRMATO_USC; COD_FASE FIRMATO; '
                   'COD_PAGINA_RIPRESA FIRMA_USC; TXT_EVENTO non cambia più.'],
                  ['ANSC_LOG_AUDIT', 'scrittura', 'Fase FIRMA_USC.']],
        stati=('1 IN LAVORAZIONE', '2 DEFINITO (3 se l’aggiornamento anagrafico non riesce)',
               '11 FIRMATO PER CONFORMITÀ', '5 FIRMATO DA USC', 'FIRMA_USC'),
        transizioni=['Chiudi: l’atto è formato; punto di non ritorno.'],
        aperti=['PKG_PRATICA.IntegrazioneATTO assegna oggi un numero definitivo della serie di '
                'SIPO, mentre l’atto digitale porta il numero comunale di num-ansc-sipo: va '
                'deciso quale numero resti sull’atto di SIPO (OP-60).',
                'Il nome e il ruolo del pulsante: la figura lo chiama ora «Chiudi» e vi pone il '
                'salvataggio in SIPO; l’analisi lo descrive come «Definisci» (OP-76 c).']),
    dict(
        titolo='Chiudi', pptx=9, righe=[7], schermate=[],
        sintesi='Il salvataggio in SIPO è avvenuto con la firma. L’atto non si modifica più, '
                'nemmeno fino alle 24 del giorno di redazione come avviene oggi: dopo la firma '
                'la correzione è un procedimento di ANSC (annotazione, rettifica, '
                'annullamento).',
        pagine=[['Uscita', 'statocivile-web', 'invariata']],
        servizi=[['—', 'Nessuna chiamata.']],
        sipo=[['ATTO', 'lettura', 'La pratica definita']],
        ansc_usr=[['ANSC_STATO_ATTO', 'lettura', 'Lo stato finale.']],
        stati=('2 DEFINITO', '2 DEFINITO', '5 FIRMATO DA USC', '5 FIRMATO DA USC', '—'),
        transizioni=['Nessuna nel percorso ordinario; i passaggi successivi sono nel capitolo '
                     '«Dopo la formazione».'],
        aperti=['La rettifica di SIPO (/updLavorazioneRettifica) riporta oggi la pratica da 2 a '
                '1: per un atto formato in ANSC va inibita o ricondotta a R013.']),
]


# ------------------------------------------------------------------------------ primitive
def testo(par, s):
    if par.runs:
        par.runs[0].text = s
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(s)


def h(d, livello, t):
    return d.add_paragraph(t, style=f'Heading {livello}')


def par(d, t, stile='Normal', corsivo=False):
    p = d.add_paragraph(style=stile)
    for pezzo, grassetto in D.segmenta('', t):
        r = p.add_run(pezzo)
        r.bold = grassetto
        if corsivo:
            r.italic = True
            r.font.size = Pt(9)
    return p


def voce(d, t):
    return par(d, t, 'List Paragraph')


def tabella(d, righe, larghezze):
    t = d.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(righe):
        for j, v in enumerate(r):
            run = t.cell(i, j).paragraphs[0].add_run(str(v))
            run.bold = i == 0
            run.font.size = Pt(9)
    for j, w in enumerate(larghezze):
        for r in t.rows:
            r.cells[j].width = Inches(w)
    griglia = t._tbl.find(qn('w:tblGrid'))
    for col, w in zip(griglia.findall(qn('w:gridCol')), larghezze):
        col.set(qn('w:w'), str(int(round(Inches(w).twips))))
    d.add_paragraph()
    return t


def figura(d, png, larghezza, didascalia):
    d.add_picture(png, width=Inches(larghezza))
    d.paragraphs[-1].alignment = 1
    par(d, didascalia, corsivo=True)


def costruisci():
    shutil.copyfile(TEMPLATE, OUT)
    d = docx.Document(OUT)
    body = d.element.body
    for ch in list(body.iterchildren())[TAGLIA_DA:]:
        if ch.tag != qn('w:sectPr'):
            body.remove(ch)
    for p in d.paragraphs:
        if '“Lorem Ipsum”' in p.text:
            testo(p, '“Integrazione SIPO – ANSC”')
        elif p.style.name == 'Subtitle':
            testo(p, 'La gestione delle interfacce del workflow: pagine, servizi, tabelle e stati')
    return d


# ------------------------------------------------------------------------------ il documento
def scrivi(d):
    punti = punti_pptx()
    assert len(punti) == len(FASI), (len(punti), [p[0] for p in punti])
    righe13, schermate, sipo_stati = analisi()
    ansc = stati_ansc()

    h(d, 1, 'Scopo del documento')
    par(d, 'Il documento descrive, pagina per pagina, il workflow con cui SIPO forma un atto '
           'insieme ad ANSC: per ciascun punto, le pagine di SIPO coinvolte, i servizi di ANSC '
           'chiamati, le tabelle di SIPO e dell’integrazione su cui si lavora, e gli stati che '
           'l’atto assume in SIPO e in ANSC prima e dopo il passaggio.')
    par(d, 'Riprende il processo del capitolo «Il processo di gestione dei dati ANSC in SIPO» '
           'dell’analisi di integrazione [R1] e la sua figura [R2], e vi aggiunge ciò che il '
           'codice di SIPO dice di ciascuna pagina. Le schermate di SIPO sono quelle di un atto '
           'di nascita, che fa da esempio: per le altre famiglie cambiano le maschere del '
           'dettaglio e le tabelle specifiche dell’atto, non il workflow.')
    h(d, 2, 'Riferimenti')
    tabella(d, [['#', 'Documento', 'Che cosa se ne prende'],
                ['[R1]', 'ANALISI_Integrazione-ANSC_v3.44.docx', 'Il capitolo 13, il modello '
                 'dati (ANSC_STATO_ATTO e le tabelle dell’integrazione), gli Open Point.'],
                ['[R2]', 'strumenti/img/workflow_ansc_sipo.pptx', 'I punti del workflow, i '
                 'servizi e gli stati. È la fonte della figura del capitolo successivo.'],
                ['[R3]', 'MAPPA_Nascite_Schermate-Servizi-Tabelle_v0.1.docx', 'Le tabelle '
                 'toccate dagli endpoint dell’area nascite.'],
                ['[R4]', 'ansc/docs/Decodifiche (ANSC_11)', 'Gli stati dell’atto in ANSC.'],
                ['[R5]', 'Codice di SIPO: statocivile-web, nascita-web, statocivile-be, '
                 'nascita-be', 'Gli endpoint chiamati da ciascuna pagina e le tabelle che '
                 'raggiungono.']], [0.6, 2.6, 3.3])
    h(d, 2, 'Glossario')
    tabella(d, [['Termine', 'Significato'],
                ['Stato SIPO', 'MATR_USR.ATTO.ID_STATO_PRATICA, decodificato da '
                 'CONF_STATO_PRATICA.'],
                ['Stato ANSC', 'Lo stato dell’atto in ANSC (decodifica ANSC_11), replicato in '
                 'ANSC_STATO_ATTO.STATO.'],
                ['Stato 0', 'Stato ad uso applicativo: per ANSC coincide con 1 BOZZA; dice che '
                 'l’atto ha già l’identificativo nazionale ma i dati sono cambiati dopo la '
                 'validazione.'],
                ['Pagina di ripresa', 'ANSC_STATO_ATTO.COD_PAGINA_RIPRESA: l’ultima pagina '
                 'raggiunta, che si riapre quando una redazione sospesa si riprende.'],
                ['UC', 'Il caso d’uso di ANSC (decodifica ANSC_03), determinato nel riepilogo.'],
                ['OTP di sessione', 'Il codice di quattro ore generato sulla web app di ANSC, '
                 'necessario a ogni chiamata.'],
                ['Credenziali di firma remota', 'Utenza, PIN e OTP del servizio di firma '
                 'qualificata, richiesti da R007.']], [1.8, 4.7])

    h(d, 1, 'Il workflow in sintesi')
    par(d, 'La figura è la diapositiva [R2]: in colonna le pagine di SIPO, il servizio che '
           'ciascuna chiama, lo stato dell’atto in ANSC e quello in SIPO. La stella segnala le '
           'pagine che richiedono l’OTP di sessione.')
    figura(d, figura_pptx(), 6.2, 'Il processo di gestione dei dati ANSC in SIPO [R2].')
    righe = [['#', 'Punto del workflow', 'Servizio', 'Stato SIPO in uscita',
              'Stato ANSC in uscita']]
    for i, (f, (titolo_pptx, serv)) in enumerate(zip(FASI, punti), start=1):
        righe.append([str(i), titolo_pptx, serv, f['stati'][1], f['stati'][3]])
    tabella(d, righe, [0.3, 1.8, 1.7, 1.4, 1.3])

    h(d, 1, 'Gli stati')
    h(d, 2, 'In ANSC')
    par(d, 'Gli undici stati della decodifica ANSC_11 [R4], con il punto del workflow che li '
           'produce.')
    tabella(d, [['ID', 'Stato', 'Nel workflow']] + [[i, s, NEL_WF.get(i, '')] for i, s in ansc],
            [0.4, 2.0, 4.1])
    h(d, 2, 'In SIPO')
    par(d, 'In SIPO lo stato è racchiuso in una sola colonna, MATR_USR.ATTO.ID_STATO_PRATICA, '
           'decodificata da CONF_STATO_PRATICA. Le righe di quella tabella non sono nel '
           'repository: alcune descrizioni sono ricavate dall’uso [R1].')
    tabella(d, sipo_stati, [0.6, 2.6, 3.3])
    par(d, 'Il workflow usa 1 e 2. Il 3 compare se l’aggiornamento anagrafico che segue la '
           'definizione non riesce o resta in attesa, e torna a 2 con /annullaSospensioneAtto.')
    h(d, 2, 'Gli stati locali dell’integrazione')
    par(d, 'ANSC_STATO_ATTO porta, accanto alla replica dello stato ANSC, tre informazioni '
           'nostre [R1].')
    tabella(d, LOCALI, [1.6, 2.4, 2.5])

    h(d, 1, 'Il diagramma a stati')
    par(d, 'Ogni stato del diagramma è la coppia stato SIPO · stato ANSC. Sulle frecce, la '
           'pagina o il servizio che provoca il passaggio. Il riquadro viola raccoglie gli '
           'stati in cui la redazione si può sospendere e riprendere; la linea rossa è il punto '
           'di non ritorno.')
    figura(d, diagrammi_workflow_stati.disegna(), 6.4, 'Il workflow come macchina a stati.')
    tabella(d, [['Da', 'A', 'Che cosa lo provoca'],
                ['Nessun atto', 'In lavorazione', 'Scelta del tipo di atto: SIPO crea la '
                 'pratica (/saveAtto).'],
                ['In lavorazione', 'In lavorazione', 'Ricerca, dettaglio (R005), riepilogo, '
                 'anteprima; una validazione respinta.'],
                ['In lavorazione', 'Validato', 'R009 accettata.'],
                ['Validato', 'Da rivalidare', 'Modifica dei dati dopo la validazione.'],
                ['Da rivalidare', 'Validato', 'Nuova validazione R009.'],
                ['Validato', 'Abbandonato', 'R011 annulla.'],
                ['Validato', 'Firmato dai dichiaranti', 'R006 o R012.'],
                ['Validato', 'Firmato per conformità', 'Senza firmatari: R007 '
                 '/attestazione/firma.'],
                ['Firmato dai dichiaranti', 'Firmato per conformità', 'R007 '
                 '/attestazione/firma.'],
                ['Firmato per conformità', 'Atto formato', 'R007 /firma_usc, R014 se '
                 'richiesto, salvataggio in SIPO.'],
                ['Atto formato', 'Sospeso in SIPO', 'Aggiornamento anagrafico non riuscito o '
                 'in attesa.'],
                ['Sospeso in SIPO', 'Atto formato', '/annullaSospensioneAtto.'],
                ['Atto formato', 'Dopo la formazione', 'R011, rettifica di tipo annullamento, '
                 'annullamento per inefficacia dalla web app.']], [1.7, 1.7, 3.1])

    for n, (f, (titolo_pptx, serv)) in enumerate(zip(FASI, punti), start=1):
        h(d, 1, f['titolo'])
        par(d, f['sintesi'])
        par(d, f'Nella figura [R2]: «{titolo_pptx}»; servizio: {serv}.', corsivo=True)
        h(d, 2, 'Pagine')
        tabella(d, [['Pagina', 'Dove sta in SIPO', 'Intervento']] + f['pagine'],
                [1.6, 3.0, 1.9])
        h(d, 2, 'Servizi ANSC chiamati')
        tabella(d, [['Servizio', 'Che cosa fa qui']] + f['servizi'], [1.2, 5.3])
        h(d, 2, 'Tabelle coinvolte')
        par(d, 'Le tabelle di SIPO, con l’endpoint che le raggiunge [R5]:')
        tabella(d, [['Tabella SIPO', 'Uso', 'Endpoint e note']] + f['sipo'], [2.3, 1.1, 3.1])
        par(d, 'Le tabelle dell’integrazione [R1]:')
        tabella(d, [['Tabella', 'Uso', 'Che cosa']] + f['ansc_usr'], [2.3, 1.1, 3.1])
        h(d, 2, 'Stati')
        s_in, s_out, a_in, a_out, pag = f['stati']
        tabella(d, [['', 'In ingresso', 'In uscita'],
                    ['Stato SIPO', s_in, s_out],
                    ['Stato ANSC', a_in, a_out],
                    ['Pagina di ripresa', '', pag]], [1.6, 2.4, 2.5])
        h(d, 2, 'Transizioni')
        for t in f['transizioni']:
            voce(d, t)
        if f['righe']:
            h(d, 2, 'Dal capitolo 13 dell’analisi')
            for ri in f['righe']:
                r = righe13[ri]
                par(d, f'**{r["passaggio"]}**')
                for riga in r['testo'].split('\n'):
                    if riga.strip():
                        par(d, riga.strip())
                for img in r['immagini']:
                    figura(d, img, 5.6, f'Schermata di SIPO: {r["passaggio"]}.')
        if f['schermate']:
            h(d, 2, 'Schermate di riferimento')
            par(d, 'Schermate del sistema Side del Comune di Milano, riprese dal capitolo 13.1 '
                   'dell’analisi [R1]: la pagina di SIPO andrà realizzata in modo simile.')
            for k in f['schermate']:
                figura(d, schermate[k], 6.0, f'Side: {f["titolo"].split("·")[-1].strip()}.')
        if f['aperti']:
            h(d, 2, 'Punti aperti')
            for a in f['aperti']:
                voce(d, a)

    h(d, 1, 'Dopo la formazione')
    par(d, 'Formato l’atto, gli stati ANSC cambiano solo per gli istituti previsti, e alcuni '
           'soltanto dalla web app [R1]:')
    tabella(d, [['Momento', 'Stato ANSC', 'Da dove']] +
            [[r[0], r[1], r[2]] for r in [
                ['Cancellazione della firma dell’ufficiale', '9 CANCELLATO', 'R011'],
                ['Rettifica di tipo annullamento', '3 ANNULLATO', 'R013'],
                ['Annullamento per inefficacia', '10 INEFFICACE', 'solo la web app']]],
            [2.6, 1.6, 2.3])
    par(d, 'Chi usa la web app per queste operazioni riconcilia subito dopo dal back-office: '
           'è l’unico modo perché lo stato registrato in SIPO resti quello vero.')

    h(d, 1, 'Sospensione e ripresa')
    par(d, 'Dalla pagina di dettaglio fino alla firma dell’USC la redazione si può sospendere e '
           'riprendere. ANSC_STATO_ATTO conserva il modello evento intero, la pagina raggiunta '
           'e la revisione [R1]. La ripresa si regge su sette regole.')
    for testa, corpo in V43.RIPRESA:
        par(d, f'**{testa}**' + corpo.replace('i firmatari dalla tabella firmatari',
                                              'i firmatari dalla tabella firmatari'),
            'List Paragraph')
    par(d, 'Lo stato 0. ' + V43.STATO_ZERO.replace('Il problema. ', ''))

    h(d, 1, 'Punti aperti')
    par(d, 'I punti aperti delle singole pagine, riuniti. Dove esiste, il rimando all’Open Point '
           'dell’analisi [R1].')
    righe = [['#', 'Punto del workflow', 'Questione']]
    k = 0
    for f in FASI:
        for a in f['aperti']:
            k += 1
            righe.append([f'WF-{k}', f['titolo'], a])
    tabella(d, righe, [0.6, 1.6, 4.3])
    return k


def main():
    d = costruisci()
    n_aperti = scrivi(d)
    D.frontespizio(d, VERSIONE, NOME, data=DATA)
    st = D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')
    r = st.rows[1]
    for cella, v in zip(r.cells, [DATA, VERSIONE, 'Tutti', 'Prima emissione: il workflow '
                                  'SIPO-ANSC pagina per pagina, con servizi ANSC, tabelle e '
                                  'stati; diagramma a stati.']):
        D.riscrivi_cella(cella, v)
    for riga in list(st.rows[2:]):               # le righe vuote del template
        if not ''.join(c.text for c in riga.cells).strip():
            st._tbl.remove(riga._tr)
    d.save(OUT)
    cap, tab, fig = D.riepilogo(OUT)
    print(f'  capitoli {cap}, tabelle {tab}, immagini {fig}, punti aperti {n_aperti}')
    print('  scritto', OUT)


if __name__ == '__main__':
    main()
