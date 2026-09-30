# -*- coding: utf-8 -*-
"""Un foglio per caso d'uso, morte e nascita, in un unico file.

Struttura del file: i fogli di governo davanti (indice, lavoro residuo, quadro dei percorsi,
blocchi assenti) e poi un foglio per UC, nella forma del foglio di lavoro compilato a mano.
L'indice è navigabile: ogni riga porta il collegamento al proprio foglio.

⚠️ DUE FAMIGLIE, DUE SITUAZIONI DIVERSE, e il file non le confonde:
  · NASCITA — esiste il foglio compilato a mano di `Dic_Nasc_001`: la colonna SIPO si propaga
    a tutti gli UC in cui lo stesso percorso ANSC ricorre.
  · MORTE — non esiste alcun foglio compilato: c'è solo ciò che il codice dell'area decessi
    sa dire. Un percorso compilato per la nascita NON si propaga alla morte se nomina una
    tabella della nascita (`ATTO_NASCITA…`): stessa forma, altra sorgente. Dove invece la
    tabella è comune — `SOGGETTO`, `ATTO` — si riporta come suggerimento di altra famiglia,
    marcato come tale, perché un suggerimento non è un rilievo.

    /usr/bin/python3 "Documenti finali/strumenti/uc_workbook.py"
"""
import collections
import csv
import os
import sys

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import uc_famiglia as UF   # noqa: E402

BASE = UF.BASE
CARTELLA = UF.CARTELLA
USCITA = os.path.join(CARTELLA, 'MAPPATURA_UC_Nascite-Morti_v0.1.xlsx')
ALTRA = 'suggerita da altra famiglia'
COLORI = dict(UF.COLORI)
COLORI[ALTRA] = PatternFill('solid', fgColor='EDE4F5')

FAMIGLIE = (('morte', 'decessi', None), ('nascita', 'nascita', 'Dic_Nasc_001'))

COL_UC = [
    ('UseCase', 12, 'id_uc'), ('Sezione FE ANSC', 22, 'sezione'), ('Campo', 30, 'campo'),
    ('Binding Object', 30, 'bo'), ('Binding Field', 24, 'bf'),
    ('Dato obbligatorio SI/NO', 10, 'obbligatorio'),
    ('Tabella/Campo SIPO', 34, 'sipo'),
    ('Regole di transcodifica', 44, 'transcodifica'),
    ('Condizioni particolari/Logiche di business', 44, 'note'),
    ("Condizioni obbligatorieta' ANSC", 40, 'condizione'),
    ("Obbligatorio per l'UC", 12, 'esito'), ('Come è stato deciso', 30, 'come'),
    ('Origine della colonna SIPO', 22, 'origine'), ('Evidenza', 46, 'evidenza'),
]
COL_INDICE = [
    ('Caso d’uso', 20, 'uc'), ('ID usecase', 12, 'id_uc'), ('Famiglia', 12, 'famiglia'),
    ('Righe', 8, 'righe'), ('Compilate a mano', 12, 'a_mano'),
    ('Proposte', 10, 'proposta'), ('Da altra famiglia', 12, 'altra'),
    ('Da compilare', 12, 'scoperto'), ('Soggetto assente', 12, 'assente'),
    ('Obbligatori', 10, 'obbligatori'), ('Vai al foglio', 16, 'link'),
]


def id_usecase():
    f = os.path.join(BASE, 'ansc', 'docs', 'Mapping_casi_uso', '3_dec_use_case.csv')
    with open(f, encoding='utf-8-sig') as fh:
        return {r['COD ANSC'].strip(): r['ID USECASE'].strip()
                for r in csv.DictReader(fh) if r.get('COD ANSC')}


def specifica_di_nascita(valore):
    """Una colonna che nomina una tabella della nascita non vale per la morte."""
    t = (valore or '').split('.')[0].upper()
    return 'NASCITA' in t or 'NASC' in t


def costruisci():
    ids = id_usecase()
    tutte, percorsi_glob = [], {}
    mano_nascita = UF.lavoro_a_mano('Dic_Nasc_001')
    for famiglia, area, sorgente in FAMIGLIE:
        if sorgente:
            righe, percorsi = UF.raccogli(famiglia, area, sorgente)
        else:
            # nessun foglio a mano per questa famiglia: si passa un indice vuoto e si
            # completa dopo con i suggerimenti dell'altra, filtrati
            righe, percorsi = UF.raccogli(famiglia, area, sorgente_vuota())
            for r in righe:
                k = f"{r['bo']}.{r['bf']}".strip('.')
                if r['origine'] not in (UF.SCOPERTO, UF.PROPOSTA) or k not in mano_nascita:
                    continue
                sipo, transc, note = mano_nascita[k]
                if not sipo:
                    continue
                if not UF.ben_formata(sipo):
                    # annotazione, non colonna: si riporta fra le note e la cella resta da fare
                    r['note'] = (f'Dal foglio di Dic_Nasc_001: {sipo} {note}').strip()
                elif not specifica_di_nascita(sipo):
                    # ⚠️ il lavoro umano dell'altra famiglia batte la proposta automatica:
                    # su `numeroatto` il compilatore ha scelto NUM_COMUNALE_ANSC dove il
                    # codice propone NUMERO_ATTO, e sono due colonne diverse.
                    r.update(sipo=UF.normalizza(sipo), transcodifica=transc, note=note,
                             origine=ALTRA, evidenza='Dic_Nasc_001.xlsx (altra famiglia)')
            rifai_percorsi(righe, percorsi)
        for r in righe:
            r['famiglia'] = famiglia
            r['id_uc'] = ids.get(r['uc'], '')
        for p in percorsi:
            p['famiglia'] = famiglia
        tutte.extend(righe)
        percorsi_glob[famiglia] = percorsi
    return tutte, percorsi_glob


def sorgente_vuota():
    """Un foglio di lavoro inesistente: `raccogli` accetta il nome, qui non deve trovare nulla."""
    return '__nessuno__'


def rifai_percorsi(righe, percorsi):
    """Riallinea l'origine dei percorsi dopo l'innesto dei suggerimenti di altra famiglia."""
    per_k = {}
    for r in righe:
        k = f"{r['bo']}.{r['bf']}".strip('.')
        if k:
            per_k.setdefault(k, r)
    for p in percorsi:
        r = per_k.get(p['percorso'])
        if r:
            p['origine'], p['sipo'], p['evidenza'] = r['origine'], r['sipo'], r['evidenza']


def _foglio(wb, nome, colonne, righe, evidenzia=None, link=None):
    ws = wb.create_sheet(nome)
    for j, (_, largh, _) in enumerate(colonne, 1):
        ws.column_dimensions[get_column_letter(j)].width = largh
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    testa = []
    for titolo, _, _ in colonne:
        c = WriteOnlyCell(ws, value=titolo)
        c.font = Font(bold=True, color='FFFFFF', size=9)
        c.fill = UF.BLU
        c.alignment = Alignment(vertical='center', wrap_text=True)
        testa.append(c)
    ws.append(testa)
    n = 1
    for r in righe:
        fill = evidenzia(r) if evidenzia else None
        celle = []
        for _, _, k in colonne:
            v = r.get(k, '')
            c = WriteOnlyCell(ws, value=v)
            if fill:
                c.fill = fill
            if link and k == 'link' and v:
                c.value = 'apri'
                c.hyperlink = f"#'{v}'!A1"
                c.font = Font(color='0563C1', underline='single')
            celle.append(c)
        ws.append(celle)
        n += 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(colonne))}{n}'
    return ws


def _legenda(wb, righe, percorsi):
    ws = wb.create_sheet('Legenda', 0)
    ws.column_dimensions['A'].width = 44
    ws.column_dimensions['B'].width = 104
    cp = collections.Counter(p['origine'] for p in percorsi)
    cr = collections.Counter(r['origine'] for r in righe if r['origine'])
    ce = collections.Counter(r['esito'] for r in righe)
    n_uc = len({(r['famiglia'], r['uc']) for r in righe})
    voci = [
        ('Mappatura dei casi d’uso — morte e nascita', 'SIPO ↔ ANSC · Roma Capitale', True),
        ('Che cos’è',
         f'Un foglio per ciascuno dei {n_uc} casi d’uso delle due famiglie, nella forma del '
         f'foglio di lavoro compilato a mano. Davanti stanno i fogli di governo: l’indice '
         f'navigabile, il lavoro residuo, il quadro dei percorsi, i blocchi che SIPO non ha.',
         False),
        ('Come si aggiorna',
         'Si rigenera con «Documenti finali/strumenti/uc_workbook.py». Il lavoro umano si '
         'scrive nel foglio di lavoro dell’UC sorgente (Dic_Nasc_001.xlsx): alla '
         'rigenerazione si propaga da solo a tutti gli UC che usano lo stesso percorso.',
         False),
        ('⚠️ Il lavoro si fa PER PERCORSO, non per UC',
         f'I casi d’uso sono {n_uc}, i percorsi ANSC distinti {len(percorsi)}. Lo stesso '
         f'percorso ricorre in decine di UC con la stessa sorgente SIPO: compilarlo una '
         f'volta lo risolve ovunque. Il foglio «Da compilare» è ordinato per quante righe '
         f'ciascun percorso sblocca.', False),
        ('⚠️ Le due famiglie non sono nella stessa condizione', None, True),
        ('  nascita',
         'Esiste il foglio compilato a mano di Dic_Nasc_001: la colonna SIPO si propaga a '
         'tutti gli UC in cui lo stesso percorso ricorre.', False),
        ('  morte',
         'Non esiste alcun foglio compilato: c’è ciò che il codice dell’area decessi sa '
         'dire, più i percorsi comuni suggeriti dalla nascita. ⚠️ Un valore che nomina una '
         'tabella della nascita NON è stato propagato: stessa forma, altra sorgente.', False),
        ('Le origini della colonna SIPO', None, True),
        (f'  {UF.A_MANO}', 'Dal foglio di Dic_Nasc_001, stesso percorso ANSC. È una '
                           'propagazione, non un accertamento.', False),
        (f'  {UF.PROPOSTA}', 'Dal codice: maschera → DTO → colonna, con evidenza file:riga. '
                             'Sul banco di prova concorda con il compilatore 36 volte su 41 '
                             'e le 5 divergenze danno ragione a lui.', False),
        (f'  {ALTRA}', 'Percorso compilato per l’altra famiglia, su tabella non specifica '
                       '(SOGGETTO, ATTO). Da verificare prima dell’uso.', False),
        (f'  {UF.SCOPERTO}', 'Il lavoro che resta.', False),
        (f'  {UF.ASSENTE}', 'Manca il blocco intero: parte di maschera da progettare.', False),
        ('I conti (percorsi · righe UC×campo)', None, True),
    ]
    for k in (UF.A_MANO, UF.PROPOSTA, ALTRA, UF.SCOPERTO, UF.ASSENTE):
        voci.append((f'  · {k}', f'{cp.get(k, 0)} percorsi · {cr.get(k, 0)} righe', False))
    voci += [
        ('  · totale', f'{len(percorsi)} percorsi · {len(righe)} righe', False),
        ('L’obbligatorietà', None, True),
        ('  la regola',
         'Governa «Condizioni obbligatorieta’»; «Obbligatorio» interviene dove quella tace. '
         '⚠️ Le due colonne del mapping si contraddicono in un terzo delle righe: su '
         'Dic_Nasc_001 «Obbligatorio» dichiara facoltativo il cognome della madre. La '
         'semantica ufficiale non è pubblicata: va chiesta a Sogei.', False),
        ('  esito', ' · '.join(f'{k or "(non dichiarato)"}: {v}' for k, v in ce.most_common()),
         False),
    ]
    for a, b, testa in voci:
        c1, c2 = WriteOnlyCell(ws, value=a), WriteOnlyCell(ws, value=b)
        c1.font = Font(bold=True, size=11 if testa else 9,
                       color='FFFFFF' if b is None else '000000')
        if b is None:
            c1.fill = UF.BLU
        c2.alignment = Alignment(wrap_text=True, vertical='top')
        ws.append([c1, c2])


if __name__ == '__main__':
    righe, percorsi_fam = costruisci()
    percorsi = [p for v in percorsi_fam.values() for p in v]
    wb = openpyxl.Workbook(write_only=True)
    _legenda(wb, righe, percorsi)

    per_uc = collections.OrderedDict()
    for r in righe:
        per_uc.setdefault((r['famiglia'], r['uc']), []).append(r)
    indice = []
    for (fam, uc), rr in per_uc.items():
        c = collections.Counter(x['origine'] for x in rr)
        indice.append({
            'uc': uc, 'id_uc': rr[0]['id_uc'], 'famiglia': fam, 'righe': len(rr),
            'a_mano': c.get(UF.A_MANO, 0), 'proposta': c.get(UF.PROPOSTA, 0),
            'altra': c.get(ALTRA, 0), 'scoperto': c.get(UF.SCOPERTO, 0),
            'assente': c.get(UF.ASSENTE, 0),
            'obbligatori': sum(1 for x in rr if x['esito'] == 'SI'), 'link': uc,
        })
    _foglio(wb, 'Indice', COL_INDICE, indice, link=True,
            evidenzia=lambda r: UF.ROSSO if r['scoperto'] > r['a_mano'] else None)
    da_fare = sorted((p for p in percorsi if p['origine'] == UF.SCOPERTO),
                     key=lambda p: -p['n_uc'])
    COL_P = [('Famiglia', 12, 'famiglia')] + UF.COL_PERC
    _foglio(wb, 'Da compilare', COL_P, da_fare, evidenzia=lambda p: UF.ROSSO)
    _foglio(wb, 'Percorsi (tutti)', COL_P, sorted(percorsi, key=lambda p: -p['n_uc']),
            evidenzia=lambda p: COLORI.get(p['origine']))
    _foglio(wb, 'Soggetti assenti in SIPO', COL_P,
            sorted((p for p in percorsi if p['origine'] == UF.ASSENTE),
                   key=lambda p: -p['n_uc']), evidenzia=lambda p: UF.GRIGIO)
    for (fam, uc), rr in per_uc.items():
        _foglio(wb, uc[:31], COL_UC, rr, evidenzia=lambda r: COLORI.get(r['origine']))
    wb.save(USCITA)
    print('scritto :', os.path.relpath(USCITA, BASE))
    print('fogli   :', 5 + len(per_uc), '· UC:', len(per_uc), '· righe:', len(righe))
    for k, v in collections.Counter(p['origine'] for p in percorsi).most_common():
        print(f'   {k or "(senza binding)":28} {v:5} percorsi')
