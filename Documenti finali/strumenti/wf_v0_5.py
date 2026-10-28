# -*- coding: utf-8 -*-
"""v0.5 del workflow delle interfacce SIPO-ANSC, dalla v0.4 dell'utente (28/10/2026).

Allineamento alle decisioni recepite nell'analisi (v3.53-v3.54):
  - il servizio della numerazione è be-numeratore-comunale, con l'operazione POST /richiesta;
  - TIPO_LOGICHE_DI_SCELTA/LOGICHE_DI_SCELTA deprecate: la logica di scelta è un codice libero;
  - ANSC_STATO_ATTO senza FLG_SOSPESO, TXT_HASH_VALIDATO, ID_EVENTO_ANSC, ID_VERSIONE:
    la «sospensione» diventa la redazione INTERROTTA (nessun contrassegno; pulsante INTERROMPI),
    lo stato 0 si imposta quando i dati cambiano dopo la validazione; NUM_COMUNALE → NUMERO_ATTO,
    DATA_VALIDAZIONE → TS_VALIDAZIONE, STATO → ID_STATO / STATO.
Figure rigenerate: schermate_sipo_wf.disegna_v3(), diagrammi_pagine_wf.comune(),
diagrammi_workflow_stati.disegna().

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/wf_v0_5.py
"""
import os
import sys

import docx

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.4.docx')
OUT = os.path.join(BASE, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.5.docx')
IMG = os.path.join(QUI, 'img')


def _paragrafi(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs


def sost(doc, vecchio, nuovo, attese=1):
    n, visti = 0, set()
    for p in _paragrafi(doc):
        if id(p._p) in visti or vecchio not in p.text:
            continue
        visti.add(id(p._p))
        for r in p.runs:
            if vecchio in r.text:
                r.text = r.text.replace(vecchio, nuovo)
                break
        else:
            D.testo_di(p, p.text.replace(vecchio, nuovo))
        n += 1
    if n != attese:
        raise SystemExit(f'«{vecchio[:70]}»: attese {attese}, trovate {n}')


def tabella(doc, *intestazioni):
    for t in doc.tables:
        testa = [c.text.strip() for c in t.rows[0].cells]
        if all(i in testa for i in intestazioni):
            return t
    raise SystemExit('tabella non trovata: %s' % (intestazioni,))


def riga(t, primo):
    for r in t.rows:
        if r.cells[0].text.strip().startswith(primo):
            return r
    raise SystemExit('riga non trovata: ' + primo)


def stati_locali(doc):
    sost(doc, 'che si riapre quando una redazione sospesa si riprende.',
         'che si riapre quando una redazione interrotta si riprende.')
    sost(doc, 'ANSC_STATO_ATTO porta, accanto alla replica dello stato ANSC, tre informazioni '
              'nostre [R1].',
         'ANSC_STATO_ATTO porta, accanto alla replica dello stato ANSC, due informazioni nostre '
         '[R9].')
    t = tabella(doc, 'Colonna di ANSC_STATO_ATTO', 'Valori')
    r = riga(t, 'STATO')
    D.riscrivi_cella(r.cells[0], 'ID_STATO / STATO')
    D.riscrivi_cella(r.cells[1], 'i codici di ANSC_11 con la loro descrizione, più '
                     'IN_PREPARAZIONE e BOZZA_APPLICATIVA (0)')
    sost(doc, 'L’ultima pagina raggiunta: è quella che si riapre quando la redazione sospesa si '
              'riprende.',
         'L’ultima pagina raggiunta: è quella che si riapre quando la redazione interrotta si '
         'riprende.')
    r = riga(t, 'FLG_SOSPESO')
    if D.ha_commenti(r._tr):
        raise SystemExit('commento sulla riga FLG_SOSPESO')
    t._tbl.remove(r._tr)
    sost(doc, 'Il riquadro viola raccoglie gli stati in cui la redazione si può sospendere e '
              'riprendere;',
         'Il riquadro viola raccoglie gli stati in cui la redazione si può interrompere e '
         'riprendere;')


def pagine_sipo(doc):
    sost(doc, 'riaprendo un atto in stato di sospensione/bozza',
         'riaprendo un atto la cui redazione è stata interrotta, o in bozza,')
    r = _riga_ovunque(doc, 'ANSC_CFG_UC, LOGICHE_DI_SCELTA')
    D.riscrivi_cella(r.cells[0], 'ANSC_CFG_UC')
    D.riscrivi_cella(r.cells[2], 'La determinazione dell’UC. La logica di scelta '
                     '(LOGICA_DI_SCELTA) è un codice libero, senza catalogo: la valuta il codice '
                     'applicativo.')
    sost(doc, 'ID_UC_ANSC, COD_LOGICA, COD_ORIGINE_UC, ID_VERSIONE; COD_FASE UC_DETERMINATO;',
         'ID_UC_ANSC, COD_LOGICA, COD_ORIGINE_UC; COD_FASE UC_DETERMINATO;')
    r = _riga_ovunque(doc, 'num-ansc-sipo')
    D.riscrivi_cella(r.cells[0], 'be-numeratore-comunale (POST /richiesta)')
    sost(doc, 'NUM_COMUNALE; TXT_EVENTO con il modello evento costruito;',
         'NUMERO_ATTO; TXT_EVENTO con il modello evento costruito;')
    sost(doc, 'il numero comunale di num-ansc-sipo', 'il numero comunale di be-numeratore-comunale',
         attese=2)


def _riga_ovunque(doc, primo):
    trovate = [r for t in doc.tables for r in t.rows
               if r.cells[0].text.strip().startswith(primo)]
    if len(trovate) != 1:
        raise SystemExit(f'«{primo}»: {len(trovate)} righe')
    return trovate[0]


def pagine_nuove(doc):
    sost(doc, 'SOSPENDI chiude la redazione e la lascia in modalità riprendibile.',
         'INTERROMPI lascia la redazione: la riga resta com’è all’ultimo salvataggio, senza '
         'contrassegno, e la redazione si riprende dalla ricerca degli atti.')
    sost(doc, 'la sospensione e la ripresa, l’eliminazione dell’evento.',
         'l’interruzione e la ripresa, l’eliminazione dell’evento.')
    r = [r for t in doc.tables for r in t.rows if r.cells[0].text.strip() == 'Pagina attiva'
         and r.cells[1].text.strip().startswith('SOSPENDI')]
    assert len(r) == 1, len(r)
    r = r[0]
    D.riscrivi_cella(r.cells[1], 'INTERROMPI, fine della sessione di SIPO, OTP non rinnovato')
    D.riscrivi_cella(r.cells[2], 'Interrotta')
    D.riscrivi_cella(r.cells[3], 'Nessun contrassegno: modello evento, pagina e revisione sono '
                     'già salvati dall’ultimo cambio di pagina o dall’ultima chiamata.')
    r = _riga_ovunque(doc, 'Sospesa')
    D.riscrivi_cella(r.cells[0], 'Interrotta')

    validati = ('ID_ANSC, ID_EVENTO_ANSC, STATO CONFERMATO, TXT_HASH_VALIDATO, DATA_VALIDAZIONE;',
                'ID_ANSC, stato 2 CONFERMATO (ID_STATO e STATO), TS_VALIDAZIONE;')
    sost(doc, *validati, attese=2)
    sost(doc, 'numero comunale a 2 Utilizzato.',
         'numero comunale a 2 Utilizzato (PUT /utilizza di be-numeratore-comunale).')
    sost(doc, 'Stato 0 BOZZA_APPLICATIVA: l’impronta del modello non è più TXT_HASH_VALIDATO.',
         'Stato 0 BOZZA_APPLICATIVA: i dati non sono più quelli validati.')


def ripresa(doc):
    h = [p for p in doc.paragraphs if p.style.name.startswith('Heading')
         and p.text.strip() == 'Sospensione e ripresa']
    assert len(h) == 1
    D.testo_di(h[0], 'Interruzione e ripresa')
    sost(doc, 'la redazione si può sospendere e riprendere. ANSC_STATO_ATTO conserva il modello '
              'evento intero, la pagina raggiunta e la revisione [R1].',
         'la redazione si può interrompere e riprendere. ANSC_STATO_ATTO conserva il modello '
         'evento intero, la pagina raggiunta e la revisione [R9].')
    sost(doc, 'I valori che ANSC restituisce (id, idAnsc, stato, contrassegno di secretazione)',
         'I valori che ANSC restituisce (idAnsc, stato, contrassegno di secretazione)')
    sost(doc, 'La sospensione.', 'L’interruzione.')
    sost(doc, 'Esplicita, con un comando «Sospendi» nelle pagine del percorso, o implicita, quando '
              'la sessione di SIPO si chiude o l’OTP scade a metà lavoro. In entrambi i casi la '
              'riga resta com’è all’ultimo salvataggio e porta il contrassegno di sospensione.',
         'Esplicita, con il comando «Interrompi» nelle pagine del percorso, o implicita, quando '
         'la sessione di SIPO si chiude o l’OTP scade a metà lavoro. In entrambi i casi la riga '
         'resta com’è all’ultimo salvataggio: non serve un contrassegno, perché una redazione è '
         'interrotta finché la sua fase non è conclusiva.')
    sost(doc, 'che mostra le redazioni sospese.', 'che mostra le redazioni non concluse.')
    sost(doc, 'Una redazione sospesa non porta con sé l’autorità di chi l’ha sospesa.',
         'Una redazione interrotta non porta con sé l’autorità di chi l’ha lasciata.')
    sost(doc, 'riprendere la redazione sospesa da un collega,',
         'riprendere la redazione lasciata da un collega,')
    sost(doc, 'Lo stato 0 si imposta quando l’impronta del modello evento salvato in '
              'ANSC_STATO_ATTO differisce da quella dell’ultima validazione riuscita '
              '(TXT_HASH_VALIDATO), e si chiude',
         'Lo stato 0 si imposta quando i dati cambiano dopo l’ultima validazione riuscita, e si '
         'chiude')


def punti_aperti(doc):
    sost(doc, 'sta in ANSC_STATO_ATTO.NUM_COMUNALE, insieme all’idAnsc;',
         'sta in ANSC_STATO_ATTO.NUMERO_ATTO, insieme all’idAnsc;')


def riferimenti(doc):
    t = tabella(doc, '#', 'Documento', 'Che cosa se ne prende')
    D.clona_riga(t, ('[R9]', 'ANALISI_Integrazione-ANSC_v3.54.docx',
                     'Il modello dati corrente — ANSC_STATO_ATTO, ANSC_CFG_UC, la numerazione '
                     'comunale con il servizio be-numeratore-comunale — e la ripresa di una '
                     'redazione interrotta. Dove diverge da [R1], prevale.'))


FIGURE = [
    ('Il workflow come macchina a stati.', 'workflow_stati.png'),
    ('Il percorso fra le schermate.', 'wf3_percorso.png'),
    ('Il comportamento comune delle pagine nuove.', 'wf_sd_comune.png'),
    ('Ipotesi di schermata: Validazione.', 'wf3_sch_validazione.png'),
    ('Ipotesi di schermata: Firme dei dichiaranti.', 'wf3_sch_firme.png'),
    ('Ipotesi di schermata: Conformità.', 'wf3_sch_conformita.png'),
    ('Ipotesi di schermata: Firma dell’USC.', 'wf3_sch_firma_usc.png'),
]


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    riferimenti(doc)
    stati_locali(doc)
    pagine_sipo(doc)
    pagine_nuove(doc)
    ripresa(doc)
    punti_aperti(doc)
    for didascalia, png in FIGURE:
        D.sostituisci_immagine(doc, didascalia, os.path.join(IMG, png))
    D.storia(doc, '28/10/2026', '0.5',
             'Riferimenti · Gli stati locali · Riepilogo · Anteprima · Pagine nuove · Validazione '
             '· Interruzione e ripresa · Punti aperti',
             'Allineato al modello dati corrente dell’analisi [R9]: il servizio della numerazione '
             'è be-numeratore-comunale (POST /richiesta, PUT /utilizza); la logica di scelta è un '
             'codice libero, senza le tabelle del catalogo; ANSC_STATO_ATTO senza contrassegno '
             'di sospensione, impronta del modello validato, ID_EVENTO_ANSC e baseline. La '
             'sospensione diventa la redazione interrotta, con il pulsante INTERROMPI; lo stato 0 '
             'si imposta quando i dati cambiano dopo la validazione. Rigenerate le figure del '
             'percorso, degli stati e delle schermate.')
    D.frontespizio(doc, '0.5', 'WORKFLOW_Interfacce-SIPO-ANSC_v0.5', '28/10/2026')
    doc.save(OUT)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
