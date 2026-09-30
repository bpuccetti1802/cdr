# -*- coding: utf-8 -*-
"""Genera «DISEGNO_Back-Office_ANSC» — il disegno dell'applicazione di back-office.

Il documento parte dal menu e scende alle singole pagine. Ogni pagina ha un wireframe e una
scheda applicativa che dichiara le tabelle sottese e il comportamento di campi e bottoni.

⚠️ I wireframe sono prodotti da `pagine_bo.py` e vanno rigenerati quando cambia il modello
dati: una schermata che mostra una colonna inesistente induce in errore chi costruisce.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/genera_disegno_bo.py
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = os.path.join(BASE, 'Template documentale', 'template_DAD_roma-capitale.docx')
OUT = os.path.join(BASE, 'Documenti finali', 'DISEGNO_Back-Office_ANSC_v0.2.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')
TAGLIA_DA = 29


def testo(par, s):
    if par.runs:
        par.runs[0].text = s
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(s)


def costruisci():
    shutil.copyfile(TEMPLATE, OUT)
    d = docx.Document(OUT)
    body = d.element.body
    for ch in list(body.iterchildren())[TAGLIA_DA:]:
        if ch.tag != qn('w:sectPr'):
            body.remove(ch)
    for p in d.paragraphs:
        if '“Lorem Ipsum”' in p.text:
            testo(p, '“Integrazione SIPO – ANSC”')
        elif p.style.name == 'Subtitle':
            testo(p, 'Il back-office: disegno dell’applicazione e delle sue pagine')
    return d


def h(d, liv, t):
    p = d.add_paragraph(style=f'Heading {liv}')
    p.add_run(t)
    return p


def par(d, t, stile='Normal'):
    p = d.add_paragraph(style=stile)
    for pezzo, gr in D.segmenta('', t):
        r = p.add_run(pezzo)
        r.bold = gr
    return p


def voce(d, testa, corpo):
    p = d.add_paragraph(style='List Paragraph')
    r = p.add_run(testa + ' ')
    r.bold = True
    for pezzo, gr in D.segmenta('', corpo):
        rr = p.add_run(pezzo)
        rr.bold = gr
    return p


def tabella(d, righe, larghezze=None, corpo=8.5):
    t = d.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(righe):
        for j, v in enumerate(r):
            cel = t.cell(i, j)
            cel.text = ''
            p = cel.paragraphs[0]
            for pezzo, gr in D.segmenta('', str(v)):
                run = p.add_run(pezzo)
                run.bold = gr or (i == 0)
                run.font.size = Pt(corpo)
    if larghezze:
        for j, w in enumerate(larghezze):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    return t


def ddl(d, testo_ddl):
    for riga in testo_ddl.split('\n'):
        p = d.add_paragraph()
        r = p.add_run(riga)
        r.font.name = 'Courier New'
        r.font.size = Pt(8)
        p.paragraph_format.space_after = Pt(0)


def figura(d, png, didascalia, larghezza=6.4):
    d.add_paragraph().add_run().add_picture(os.path.join(IMG, png), width=Inches(larghezza))
    cap = d.add_paragraph()
    r = cap.add_run(didascalia)
    r.italic = True
    r.font.size = Pt(9)


# ═══════════════════════════════════════════════════ le schede di pagina
# (titolo, png, [paragrafi di scopo], [righe «tabelle sottese»],
#  [righe «campi, colonne e azioni»], nota conclusiva | None)
PAGINE = [
 ('La home: il menu delle aree', 'bo_home.png',
  ['La home non è un cruscotto: è un **indice**. Riprende la forma osservata in S.I.De. — '
   'mattonelle disposte su tre colonne, una riga di sottotitolo che dice che cosa si trova '
   'dietro ciascuna — perché è la forma che regge quando le aree crescono e perché non '
   'promette numeri che il back-office non è tenuto a calcolare.',
   'La scelta di non mettere contatori in home è deliberata. Un contatore invita a leggere '
   'la home come una misura dello stato del sistema; ma il back-office è una superficie di '
   '**eccezione e configurazione**, e il numero che conta — quanti atti si sono fermati — '
   'vive nella pagina che li tratta, dove accanto al numero c’è l’azione.',
   '⚠️ Il **distintivo OTP** accanto al titolo è l’unico elemento di stato presente in ogni '
   'pagina. Segnala se la sessione verso ANSC è attiva e quanto le resta; quando non lo è, '
   'le azioni che dialogano con ANSC sono disabilitate con la ragione visibile, non '
   'nascoste.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['Nessuna', 'Nessun accesso ai dati', 'La home è statica: le mattonelle sono rotte di '
    'navigazione filtrate dalle abilitazioni del profilo.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Mattonella d’area', 'Rotta', 'Porta alla pagina d’elenco dell’area. È **nascosta** se '
    'il profilo non possiede alcuna abilitazione dell’area: non disabilitata, nascosta.'],
   ['Distintivo OTP', 'Stato', 'Verde con tempo residuo se la sessione è attiva, rosso se '
    'assente. Al clic apre la finestra di acquisizione del codice.'],
   ['Barra applicativa', 'Cornice', 'Fornita dalla shell: municipio e sede di lavoro, '
    'ambiente, utente. Il back-office non la disegna.']],
  None),

 ('Supervisione atti', 'bo_atti.png',
  ['È la pagina più frequentata e la ragione per cui il back-office esiste: raccoglie gli '
   'atti che **non hanno concluso il percorso ordinario**. Non è una superficie di lavoro — '
   'l’atto si compila nelle maschere di SIPO — ma di ripresa e di diagnosi.',
   'La struttura è quella di ogni pagina d’elenco dell’applicazione: pannello di filtri '
   'richiudibile, tabella dei risultati, paginazione in basso a destra. I filtri sono '
   'disposti su due righe perché le due domande che ci si pone sono diverse: la prima riga '
   'chiede **a che punto** si è fermato un atto, la seconda **quale** atto.',
   '⚠️ La colonna «Stato» porta il valore di ANSC, non un nostro giudizio. È una replica di '
   'un dato altrui e il vincolo di tabella ne ammette dodici valori: gli undici pubblicati '
   'da ANSC più quello iniziale, che è nostro. La colonna «Fase» invece è nostra e dice a '
   'quale dei quattro passi del percorso l’atto è arrivato.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_STATO_ATTO**', 'Lettura', 'È la tabella dell’elenco. Le colonne mostrate sono '
    'ID_ATTO_SIPO, ID_TIPO_EVENTO, ID_UC_ANSC, COD_FASE, STATO, ID_ANSC, NUM_COMUNALE, '
    'OPERATORE; i filtri agiscono su queste più COD_MUNICIPIO, FLG_EMERGENZA e DATA_UPD.'],
   ['ANSC_CFG_UC', 'Lettura per raccordo', 'La descrizione leggibile dell’UC accanto al '
    'codice, presa per valore da COD_UC_ANSC.'],
   ['ANSC_LOG_AUDIT', 'Lettura in conteggio', 'L’ultimo errore mostrato in riga proviene da '
    'ANSC_STATO_ATTO.ULTIMO_ERRORE; il dettaglio completo sta nell’audit.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Fase', 'Filtro', 'Valori da COD_FASE: Lavorazione, Documenti, Validazione, Firma. '
    'È il filtro che risponde alla domanda «dove si è fermato».'],
   ['Stato ANSC', 'Filtro', 'I dodici valori ammessi dal vincolo di ANSC_STATO_ATTO.'],
   ['Categoria di eccezione', 'Filtro', 'Non è una colonna: è una **vista derivata** da '
    'fase, stato ed esito. Undici categorie, fra cui «Rifiutati in validazione», '
    '«Indeterminati», «UC ambiguo», «In emergenza da rientrare».'],
   ['ID ANSC / Numero comunale', 'Filtro', 'Ricerca puntuale. ⚠️ In SIPO non esiste oggi '
    'alcuna API che ricerchi un atto per identificativo nazionale: qui la ricerca è sullo '
    'store dell’integrazione, non sull’archivio di SIPO.'],
   ['Occhio', 'Azione di riga', 'Apre il dettaglio dell’atto. Sempre abilitata.'],
   ['Matita', 'Azione di riga', 'Riporta l’operatore alla maschera di SIPO che ha in carico '
    'l’atto, usando MASCHERA_UI della configurazione. Disabilitata dopo la firma.'],
   ['RICONCILIA SELEZIONATI', 'Azione di pagina', 'Interroga ANSC sullo stato reale degli '
    'atti scelti e riallinea ANSC_STATO_ATTO. ⚠️ Richiede sessione OTP attiva; non è un '
    'ritentativo, è una lettura.'],
   ['ESPORTA', 'Azione di pagina', 'Scarica l’elenco filtrato. Non tocca alcuna tabella.']],
  '⚠️ **Una scelta da confermare**: la selezione multipla esiste per la riconciliazione, che '
  'è una lettura, ma non per azioni dispositive. È il criterio adottato per le notifiche e '
  'qui se ne tiene conto per coerenza.'),

 ('Dettaglio dell’atto', 'bo_atto.png',
  ['Raccoglie in una sola pagina ciò che di un atto il componente sa: lo scenario, lo stato '
   'di completezza sezione per sezione, i documenti, la cronologia delle chiamate e il '
   'payload che è stato depositato.',
   'La riga di sezioni con spunta e croce è ripresa da S.I.De., dove assolve esattamente a '
   'questa funzione nella maschera di stesura dell’atto. Qui però **non si compila**: si '
   'legge perché una sezione è incompleta. Per questo accanto alla sezione in errore compare '
   'il riquadro che nomina il campo mancante e il codice con cui ANSC rifiuterebbe l’atto.',
   '⚠️ La scheda «Payload inviato» merita una nota. Il payload è conservato sull’atto e non '
   'soltanto nell’audit, perché **non è sempre ricostruibile**: rigenerarlo da atto e '
   'configurazione darebbe un risultato diverso dopo l’attivazione di una nuova versione, '
   'mentre l’atto depositato è quello che è. Letto insieme a ID_VERSIONE dice che cosa è '
   'stato inviato e con quale configurazione.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_STATO_ATTO**', 'Lettura e scrittura', 'Lo scenario in testa (ID_MODELLO_ATTO, '
    'ID_CONF_TIPO_ATTO, ID_UC_ANSC, COD_ORIGINE_UC, ID_VERSIONE) e il payload '
    '(TXT_PAYLOAD). Le azioni di ripresa aggiornano COD_FASE, STATO e TIPO_ESITO.'],
   ['**ANSC_LOG_AUDIT**', 'Lettura', 'La cronologia: FASE, servizio, ESITO, DURATA_MS, '
    'OPERATORE, DATA. Richiesta e risposta si aprono dal singolo rigo.'],
   ['ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO', 'Lettura', 'L’elenco delle sezioni dell’UC e, per '
    'ciascuna, i campi obbligatori: è ciò che produce la spunta o la croce.'],
   ['ALLEGATO', 'Lettura', 'Il conteggio dei documenti e il loro stato, nella scheda '
    'dedicata.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Scenario', 'Pannello', 'Sola lettura. Mostra anche **come** l’UC è stato determinato: '
    'COD_ORIGINE_UC distingue la scelta automatica dalla scelta dell’operatore in caso di '
    'ambiguità.'],
   ['Sezioni', 'Accordion', 'Una riga per ANSC_CFG_SEZIONE, ordinata per NUM_ORDINE. Spunta '
    'se tutti i campi obbligatori sono valorizzati, croce altrimenti. All’apertura elenca i '
    'campi con la corrispondenza SIPO e il valore corrente.'],
   ['RIPRENDI DA VALIDAZIONE', 'Azione', 'Rimanda l’atto alla fase indicata dopo che la '
    'causa è stata rimossa. Scrive COD_FASE. ⚠️ Non è un ritentativo cieco: è abilitata solo '
    'se la preverifica locale ora passa.'],
   ['RICONCILIA CON ANSC', 'Azione', 'Interroga ANSC e riallinea lo stato. È la risposta '
    'agli esiti indeterminati, dove un ritentativo produrrebbe un duplicato.'],
   ['SCARICA PAYLOAD', 'Azione', 'Restituisce TXT_PAYLOAD come file. Utile al fornitore in '
    'diagnosi.'],
   ['ANNULLA ATTO', 'Azione', 'Cancellazione logica in ANSC. ⚠️ L’identificativo nazionale '
    'resta consumato: l’azione chiede conferma esplicita e ne dà conto nel testo.'],
   ['REGISTRO DI EMERGENZA', 'Azione', 'Porta l’atto in modalità di emergenza. Scrive '
    'FLG_EMERGENZA, ID_ATTO_EMERGENZA, COD_MOTIVO_RECUPERO e '
    'DATA_INGRESSO_EMERGENZA. In emergenza l’atto si prepara ma non si forma.']],
  None),

 ('Allegati dell’atto', 'bo_allegati.png',
  ['SIPO oggi non gestisce alcun documento allegato agli atti di stato civile: la funzione è '
   'nuova per intero. La pagina distingue due elenchi, e la distinzione non è grafica ma di '
   'sostanza.',
   'Il primo elenco è **governato dalla configurazione**: sono i documenti che l’UC '
   'determinato richiede, con la loro obbligatorietà e l’eventuale condizione. Non si '
   'aggiungono e non si tolgono da qui; si caricano. Il secondo elenco raccoglie gli **atti '
   'a testo libero**, che l’operatore descrive lui e aggiunge quando servono.',
   '⚠️ La colonna «Stato ANSC» decide se si può proseguire. Il dizionario degli stati di '
   'allegato ne dichiara cinque e **solo «Inserito» consente di andare avanti**: un documento '
   'in scansione non è un documento caricato, e la pagina non deve lasciar credere il '
   'contrario.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ALLEGATI_USECASE**', 'Lettura', 'L’elenco richiesto dall’UC: ds_allegato, '
    'ty_presenza (obbligatorietà), cd_logica (la condizione), fg_incluso_att_conformita, '
    'nr_ordinamento. Filtrata per cd_usecase e id_versione.'],
   ['**ALLEGATO**', 'Lettura e scrittura', 'I documenti effettivi: nm_file, ty_file, '
    'cd_stato, id_ansc_allegato, fg_testo_libero, fg_incluso_att_conformita, cd_hash. Il '
    'contenuto sta in oj_allegato.'],
   ['VALORE_DOMINIO', 'Lettura', 'La descrizione leggibile del tipo di allegato e dello '
    'stato, dai dizionari ANSC_09 e ANSC_08.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Riga di documento richiesto', 'Elenco', 'Non modificabile: viene dalla configurazione. '
    'Il segno più carica il file e crea la riga in ALLEGATO con fg_testo_libero = N.'],
   ['Condizione', 'Colonna', 'Testo leggibile della logica che rende il documento '
    'obbligatorio. Quando la condizione non è soddisfatta il documento resta facoltativo e '
    'la riga è attenuata.'],
   ['Descrizione + segno più', 'Campo e azione', 'Aggiunge un **atto a testo libero**: crea '
    'la riga con fg_testo_libero = S. È la via per i documenti che l’UC non prevede.'],
   ['Conformità', 'Colonna', 'Corrisponde a fg_incluso_att_conformita. Il valore proposto '
    'viene dalla configurazione; l’operatore può cambiarlo sul singolo documento.'],
   ['Cestino', 'Azione di riga', 'Rimuove il documento. Disabilitata se lo stato ANSC è '
    '«Inserito» e l’atto è già depositato.'],
   ['INVIA AD ANSC', 'Azione', 'Trasmette i documenti non ancora inviati e attende l’esito '
    'della scansione antivirus. Scrive id_ansc_allegato e cd_stato. Richiede sessione OTP.']],
  '⚠️ I vincoli sul file — dimensione massima, nome senza spazi né caratteri speciali, '
  'formati ammessi — sono dichiarati in testa alla pagina e non solo nel messaggio di '
  'errore: è la convenzione di S.I.De. e riduce i tentativi a vuoto.'),

 ('Gestione delle notifiche', 'bo_notifiche.png',
  ['Le notifiche sono il flusso in ingresso: ANSC comunica al Comune che un atto formato '
   'altrove richiede una proposta di annotazione, una presa visione o una trascrizione. È '
   'stato deciso che si lavorino **nel nostro back-office** e non sulla web app di ANSC, '
   'perché la conferma è un atto di volontà dell’ufficiale e il presidio deve stare dove '
   'stanno le pratiche.',
   '⚠️ La conferma ha una conseguenza che la pagina dichiara nel riquadro in calce: finché '
   'tutti i comuni coinvolti non hanno confermato, **la comunicazione anagrafica non parte**, '
   'e un solo rifiuto la annulla. Una notifica pendente blocca un effetto giuridico su '
   'un’altra persona, in un altro comune: non è una coda di cortesia.',
   '⚠️ La conferma è **per singola notifica**. S.I.De. offre anche la conferma massiva della '
   'pagina; qui è stata esclusa perché darebbe l’apparenza di un atto unico dove ce ne sono '
   'molti. È una scelta che si paga in tempo di lavorazione ed è fra i punti aperti.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_NOTIFICA**', 'Lettura e scrittura', 'Tutto l’elenco: DESC_INTESTATARIO, '
    'COD_GENERE, ID_COMUNE_FORMAZIONE, TXT_COMPOSIZIONE, DATA_NOTIFICA, COD_STATO, '
    'COD_STATO_SOLLECITO. La conferma o il rifiuto scrivono COD_STATO, DATA_ESITO, '
    'UTENTE_ESITO e TXT_MOTIVO_RIFIUTO.'],
   ['ANSC_STATO_ATTO', 'Lettura', 'Il raccordo con l’atto locale quando esiste '
    '(ID_ATTO_SIPO), per aprire la pratica dal dettaglio della notifica.'],
   ['VALORE_DOMINIO', 'Lettura', 'Le descrizioni di genere, stato, canale e stato di '
    'sollecito dai dizionari ANSC_101, ANSC_102, ANSC_105 e ANSC_124.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Genere', 'Filtro', 'Annotazione, trascrizione, presa visione. Il flusso di conferma usa '
    'i primi due.'],
   ['Sollecito', 'Filtro e colonna', 'Da visionare, gestita, **scaduto**. ⚠️ Il termine che '
    'determina «scaduto» non è pubblicato da ANSC: la colonna riporta ciò che ANSC dichiara, '
    'non un conteggio nostro.'],
   ['Spunta', 'Azione di riga', 'Conferma la notifica. ⚠️ Richiede sessione OTP: è un’azione '
    'dispositiva senza firma, e la sessione è l’unica prova di presenza dell’ufficiale.'],
   ['Croce', 'Azione di riga', 'Rifiuta. ⚠️ Il contratto di ANSC **non prevede un campo per '
    'la motivazione**: la si registra in TXT_MOTIVO_RIFIUTO per memoria interna, e la pagina '
    'dice che non viaggia.'],
   ['Occhio', 'Azione di riga', 'Apre il dettaglio con il testo integrale della proposta e, '
    'se l’atto è nostro, il collegamento alla pratica.'],
   ['SCARICA NOTIFICHE RECENTI', 'Azione di pagina', 'Interroga ANSC e inserisce le notifiche '
    'nuove. È un comando manuale, non un processo pianificato.']],
  None),

 ('Configurazione dei casi d’uso — elenco', 'bo_uc_elenco.png',
  ['Da qui si governa **quali casi d’uso il Comune adotta** e come li aggancia alle proprie '
   'strutture: il modello di atto, il tipo atto, la maschera che li lavora. È la pagina in '
   'cui il lato ANSC e il lato Comune si incontrano.',
   'Il sottotitolo dichiara sempre su quale versione si sta lavorando, e la dichiara perché '
   'è l’informazione che rende innocua la pagina: **le modifiche non hanno effetto finché la '
   'versione non è attivata**. Senza quel promemoria, la stessa schermata sembrerebbe '
   'toccare l’esercizio.',
   '⚠️ La colonna «Priorità» non è un ornamento: quando più logiche rispondono sullo stesso '
   'atto, vince il caso d’uso con priorità minore. È la conseguenza della scelta di '
   'determinare l’UC con espressioni sui dati anziché con regole dichiarative: la copertura '
   'e la mutua esclusione non sono più dimostrabili a priori, e al loro posto ci sono la '
   'priorità e la simulazione.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_CFG_UC**', 'Lettura e scrittura', 'L’elenco: COD_UC_ANSC, DESCRIZIONE, '
    'ID_MODELLO_ATTO, ID_CONF_TIPO_ATTO, MASCHERA_UI, SERIE, NUM_PRIORITA, STATO, '
    'DATA_INIZIO_VALIDITA e DATA_FINE_VALIDITA. Filtrata per ID_VERSIONE.'],
   ['LOGICHE_DI_SCELTA', 'Lettura', 'La logica associata, mostrata come dominio più codice '
    '(ID_DOMINIO, COD_LOGICA_DI_SCELTA).'],
   ['ANSC_ANA_UC', 'Lettura per controllo', 'Serve a segnalare un UC adottato che ANSC ha '
    'nel frattempo ritirato.'],
   ['ANSC_CFG_VERSIONE', 'Lettura', 'La versione in lavorazione e il suo stato.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Versione', 'Filtro', 'Consente di consultare una versione storica. ⚠️ In sola lettura: '
    'una versione storica non si modifica.'],
   ['Stato', 'Colonna', 'Attivo, bozza, sospeso. Un UC sospeso resta configurato ma non '
    'partecipa alla determinazione.'],
   ['Matita', 'Azione di riga', 'Apre il dettaglio dell’UC con le sue quattro schede.'],
   ['IMPORTA DA MAPPING', 'Azione di pagina', 'Prepara nella bozza sezioni, campi, allegati '
    'e formule a partire dal mapping ufficiale. ⚠️ Riscrive solo le righe che nessuno ha '
    'ancora esaminato; per le altre produce un elenco di scostamenti.'],
   ['NUOVO UC', 'Azione di pagina', 'Adotta un caso d’uso presente nel catalogo di ANSC ma '
    'non ancora configurato. La scelta avviene sul catalogo, non a mano libera.']],
  None),

 ('Configurazione di un caso d’uso — dettaglio', 'bo_uc_dettaglio.png',
  ['È la pagina dove si svolge il lavoro di configurazione vero e proprio, ed è organizzata '
   'in quattro schede perché le quattro cose che un UC dichiara hanno volumi molto diversi: '
   'poche sezioni, decine di campi, una manciata di allegati, alcune formule.',
   'La scheda «Campi» è la più densa e la si è disegnata come **superficie di revisione**, '
   'non di immissione: l’importazione porta le righe dal mapping ufficiale e il funzionario '
   'completa il lato SIPO. Il filtro «Senza corrispondenza SIPO» è quello che ordina il '
   'lavoro, perché è esattamente l’elenco di ciò che manca.',
   '⚠️ La colonna «Campo SIPO» è l’unico lavoro umano accumulato e l’unica cosa che una '
   'reimportazione può distruggere. Il documento di analisi pone la regola che si riporti '
   'dalla versione attiva a ogni rigenerazione; la pagina la rende visibile segnalando in '
   'rosso i campi ancora scoperti.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_CFG_CAMPO**', 'Lettura e scrittura', 'Scheda «Campi»: OGGETTO_ANSC, CAMPO_ANSC, '
    'FLG_OBBLIGATORIO, SCHEMA_SIPO, TABELLA_SIPO, CAMPO_SIPO, ID_DECODIFICA_ANSC, '
    'ID_BUSINESS_LOGIC, VALORE_DEFAULT, MESSAGGIO, NOTE, ORDINAMENTO, OPERATIVO.'],
   ['**ANSC_CFG_SEZIONE**', 'Lettura e scrittura', 'Scheda «Sezioni»: SEZIONE_FE_ANSC, '
    'OGGETTO_ANSC, DESCRIZIONE, NUM_ORDINE, COD_ORIGINE e la logica che ne condiziona la '
    'presenza (ID_DOMINIO, COD_LOGICA).'],
   ['**ALLEGATI_USECASE**', 'Lettura e scrittura', 'Scheda «Allegati».'],
   ['**ANSC_CFG_FORMULA**', 'Lettura e scrittura', 'Scheda «Formule»: COD_FORMULA, '
    'TXT_FORMULA, FLG_OBBLIGATORIA (lo dice ANSC) e FLG_ADOTTATA (lo sceglie il Comune).'],
   ['LOGICHE_DI_SCELTA', 'Lettura', 'Le espressioni selezionabili su sezioni e campi.'],
   ['DOMINIO_DECODIFICA', 'Lettura', 'L’elenco delle decodifiche associabili a un campo.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Sezione', 'Filtro', 'Restringe i campi a una sezione. Con oltre ottanta campi per UC è '
    'il filtro che rende la pagina praticabile.'],
   ['Senza corrispondenza SIPO', 'Filtro', 'Isola le righe con CAMPO_SIPO vuoto. È la coda '
    'di lavoro del funzionario.'],
   ['Obbl.', 'Colonna', 'Da FLG_OBBLIGATORIO. ⚠️ Nel mapping di ANSC l’obbligatorietà è '
    'dichiarata in due colonne che si contraddicono: qui si mostra il valore adottato, e la '
    'nota di riga dice da quale delle due proviene.'],
   ['Logica', 'Colonna', 'La logica di campo che trasforma il valore SIPO in valore ANSC. '
    'Quando manca e la decodifica è valorizzata, la traduzione passa dalla riconciliazione.'],
   ['SIMULA SU ATTO REALE', 'Azione di pagina', 'Applica la configurazione a un atto '
    'esistente e mostra il payload che ne risulterebbe, senza depositarlo. ⚠️ È il controllo '
    'principale, in assenza della verifica di copertura.'],
   ['REIMPORTA SEZIONE', 'Azione di pagina', 'Riporta dal mapping le sole righe della sezione '
    'corrente, con l’elenco degli scostamenti.'],
   ['SALVA', 'Azione di pagina', 'Scrive nella bozza. Disabilitata se la versione '
    'selezionata non è in stato BOZZA.']],
  None),

 ('Catalogo dei casi d’uso di ANSC', 'bo_catalogo_uc.png',
  ['È una pagina di sola consultazione, e lo dichiara nel sottotitolo. Mostra **ciò che ANSC '
   'pubblica**: non si modifica, si aggiorna. Esiste perché la configurazione deve poter '
   'nominare un caso d’uso prima di adottarlo, e perché serve un luogo dove verificare che '
   'un UC adottato sia ancora valido.',
   '⚠️ La colonna «Fine validità» è la più importante della pagina. **ANSC non cancella un '
   'caso d’uso: gli chiude la validità.** Un UC con fine validità valorizzata e ancora '
   'adottato nella configurazione è una discordanza che il report d’impatto segnala, e che '
   'qui si vede a colpo d’occhio.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_ANA_UC**', 'Lettura', 'COD_UC_ANSC, COD_MOTORE, DESCRIZIONE, COD_FAMIGLIA, '
    'ID_TIPO_EVENTO, ID_TIPO_DOCUMENTO, COD_VERSIONE, DATA_INIZIO_VALIDITA, '
    'DATA_FINE_VALIDITA.'],
   ['ANSC_CFG_UC', 'Lettura per raccordo', 'La colonna «Adottato» dice se esiste una riga di '
    'configurazione per quel codice nella versione corrente.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Validi al', 'Filtro', 'Applica la validità temporale a una data. Predefinito: oggi.'],
   ['Codice motore', 'Colonna', 'Il nome con cui il mapping ufficiale nomina il caso d’uso '
    '(per esempio Morte_001). Serve a ritrovare il file di mapping.'],
   ['Adottato', 'Colonna', 'Sì o no. Un «no» non è un difetto: il perimetro prevede di '
    'configurare prima i casi d’uso più frequenti.'],
   ['AGGIORNA DA ANSC', 'Azione di pagina', 'Ricarica il catalogo. ⚠️ Non tocca la '
    'configurazione: aggiorna soltanto la replica di ciò che ANSC dichiara.']],
  None),

 ('Logiche di scelta', 'bo_logiche.png',
  ['Le logiche sono le espressioni che il concentratore valuta sui dati dell’atto **già '
   'salvato in SIPO**. Si dividono per dominio, e i tre domini fanno cose diverse: scegliere '
   'il caso d’uso, decidere se una sezione è presente, trasformare il valore di un campo.',
   'La pagina le tratta come un catalogo unico perché condividono la stessa natura e lo '
   'stesso motore; il dominio è una colonna, non un’applicazione a parte. ⚠️ Le espressioni '
   'sono **codice eseguibile**: il linguaggio è volutamente non deciso in questa fase, e la '
   'pagina non lo presuppone.',
   'Il pannello di simulazione in calce è il presidio che sostituisce i controlli di '
   'copertura: si sceglie un atto reale e si vede che cosa la logica risponde, con l’esito e '
   'la durata. È anche il modo per accorgersi che una logica è diventata costosa.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**LOGICHE_DI_SCELTA**', 'Lettura e scrittura', 'COD_LOGICA_DI_SCELTA, '
    'DESC_LOGICA_DI_SCELTA, TXT_LOGICA e il dominio di appartenenza.'],
   ['**TIPO_LOGICHE_DI_SCELTA**', 'Lettura', 'I domini: 1 scelta dell’UC, 2 presenza della '
    'sezione, 3 logica di campo. Alimenta il filtro e la colonna.'],
   ['ANSC_CFG_UC, ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO', 'Lettura in conteggio', 'La colonna '
    '«Usata da» conta i riferimenti: è ciò che impedisce di cancellare una logica in uso.'],
   ['ANSC_STATO_ATTO', 'Lettura', 'L’atto su cui la simulazione viene eseguita.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Dominio', 'Filtro e colonna', 'Da TIPO_LOGICHE_DI_SCELTA. Cambiando dominio cambia il '
    'significato della colonna «Usata da».'],
   ['Espressione', 'Colonna', 'TXT_LOGICA, mostrata a carattere fisso. In modifica si apre '
    'in un’area ampia con la descrizione obbligatoria accanto.'],
   ['Usata da', 'Colonna', 'Numero di UC, sezioni o campi che la richiamano. Al clic ne '
    'elenca i riferimenti.'],
   ['Cestino', 'Azione di riga', '⚠️ Disabilitato se la logica è usata: si toglie il '
    'riferimento prima, non dopo.'],
   ['ESEGUI SIMULAZIONE', 'Azione', 'Valuta la logica sull’atto indicato e mostra esito, '
    'eventuale UC determinato e durata. Non scrive nulla.']],
  '⚠️ La descrizione della logica è **obbligatoria**. È il regime che risponde alla critica '
  'rivolta al motore di regole già presente in SIPO, dove le espressioni sono script senza '
  'documentazione: qui una logica senza descrizione non si salva.'),

 ('Dizionari ANSC', 'bo_dizionari.png',
  ['I dizionari sono la replica locale delle decodifiche pubblicate da ANSC. La pagina li '
   'mostra in due elenchi affiancati — i domini a sinistra, i valori del dominio scelto a '
   'destra — perché è il modo in cui si consultano: si cerca un dominio e se ne leggono i '
   'valori.',
   '⚠️ La pagina serve il back-office e la diagnosi, **non le maschere di SIPO**. Le maschere '
   'leggono i valori validi da una vista, per sinonimo e in sola lettura: è una scelta '
   'dichiarata, che tiene separati i domini di guasto.',
   '⚠️ Nella colonna «ID dominio» lo stesso codice può comparire due volte, e non è un '
   'errore: due identificativi su centoquarantatré coprono due tabelle diverse. Per questo '
   'la chiave comprende il nome del dominio e non il solo codice — senza quella correzione '
   'il carico completo violerebbe il vincolo.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**DOMINIO_DECODIFICA**', 'Lettura', 'id_dominio e nm_dominio: l’elenco di sinistra.'],
   ['**VALORE_DOMINIO**', 'Lettura', 'cd_valore, ds_valore, nr_ordinamento, '
    'dt_inizio_validita, dt_fine_validita: l’elenco di destra.'],
   ['ANSC_CFG_CAMPO, ALLEGATI_USECASE', 'Lettura in conteggio', 'La colonna «Usato da» conta '
    'i campi e gli allegati che richiamano il dominio.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Validi al', 'Filtro', 'Applica la validità temporale. Un valore scaduto resta visibile '
    'ma attenuato: serve a leggere gli atti vecchi.'],
   ['SCARICA DA ANSC (R901)', 'Azione di pagina', 'Ricarica domini e valori. ⚠️ È un comando '
    'manuale e non pianificato; il servizio restituisce il contenuto compresso.'],
   ['STORICO DEGLI SCARICHI', 'Azione di pagina', '⚠️ Le strutture ereditate non conservano '
    'lo storico dei caricamenti: la rinuncia è dichiarata nell’analisi. Il pulsante rimanda '
    'al registro dei comandi, dove l’esecuzione è tracciata.'],
   ['Riga di dominio', 'Selezione', 'Seleziona il dominio e aggiorna l’elenco di destra. '
    'Nessuna modifica: i dizionari si aggiornano da ANSC, non a mano.']],
  None),

 ('Riconciliazione delle decodifiche', 'bo_riconciliazione.png',
  ['È la pagina che colma un buco che il modello aveva: sapere **quale** dizionario ANSC vale '
   'per un campo non dice **come tradurre** il valore che SIPO ha in tabella. La '
   'riconciliazione è quella corrispondenza, valore per valore.',
   'Il raccordo si fa contro le tabelle di configurazione di SIPO — le CONF_* — e la pagina '
   'lo espone con le descrizioni di entrambi i lati, perché un confronto fra codici nudi non '
   'è verificabile da un funzionario.',
   '⚠️ Un valore SIPO senza corrispondenza **blocca in preverifica ogni atto che lo usi**. '
   'Per questo il filtro «Solo non riconciliati» è predefinito a sì: la pagina si apre sul '
   'lavoro da fare, non sull’elenco completo.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**RICONCILIAZ_DIZIONARI**', 'Lettura e scrittura', 'DECODIFICA, VALORE_SIPO, '
    'DESCRIZIONE_SIPO, VALORE_ANSC, DESCRIZIONE_ANSC, CONDIZIONE, DATA_INIZIO_VALIDITA e '
    'DATA_FINE_VALIDITA. Filtrata per ID_VERSIONE.'],
   ['DOMINIO_DECODIFICA, VALORE_DOMINIO', 'Lettura', 'Il lato ANSC: alimenta la tendina dei '
    'valori ammessi e la descrizione a fianco.'],
   ['Tabelle CONF_* di SIPO', 'Lettura per sinonimo', 'Il lato SIPO: l’elenco dei valori da '
    'riconciliare e la loro descrizione.'],
   ['ANSC_CFG_CAMPO', 'Lettura', 'Individua quali campi useranno la corrispondenza.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Decodifica ANSC', 'Filtro', 'Il dominio su cui si sta lavorando.'],
   ['Tabella SIPO', 'Filtro', 'La tabella di configurazione corrispondente. ⚠️ Il raccordo '
    'fra le due è materia di analisi e non è deducibile dai nomi.'],
   ['Valore ANSC', 'Campo', 'Tendina dei valori validi del dominio. Un valore non ancora '
    'scelto è evidenziato come «da mappare».'],
   ['Condizione', 'Campo', 'Per i casi in cui la traduzione dipende dal contesto e non dal '
    'solo valore. Vuota nella maggior parte delle righe.'],
   ['IMPORTA DA EXCEL', 'Azione di pagina', 'Carica le corrispondenze dal foglio di lavoro '
    'usato in analisi, in stato da confermare.'],
   ['Validità', 'Colonna', 'Una corrispondenza può cambiare nel tempo senza perdere il '
    'passato: gli atti già depositati restano leggibili con la corrispondenza in vigore '
    'allora.']],
  None),

 ('Versioni della configurazione', 'bo_versioni.png',
  ['La configurazione non è un carico iniziale: è un flusso. Il mapping di ANSC è rivisto in '
   'media ogni diciassette giorni, e senza un governo delle versioni ogni revisione '
   'arriverebbe in esercizio senza che nessuno l’abbia guardata.',
   'Il meccanismo è per copia integrale: una nuova bozza **copia** l’attiva — portandosi '
   'dietro il lavoro umano sui campi SIPO — e vi applica il solo scostamento pubblicato. Una '
   'sola versione può essere attiva; il ritorno indietro è un cambio di stato, non un '
   'ripristino.',
   'Il report d’impatto è la parte utile della pagina, perché dice **che cosa si romperà**: '
   'campi nuovi senza corrispondenza, campi spariti dal mapping, logiche che perdono '
   'copertura, casi d’uso ritirati e ancora adottati, valori non riconciliati. Il funzionario '
   'lavora i buchi, non l’intera configurazione.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_CFG_VERSIONE**', 'Lettura e scrittura', 'COD_VERSIONE, COD_STATO, '
    'COD_VERSIONE_ANSC, ID_VERSIONE_ANSC, DESCRIZIONE, NUM_UC_TOCCATI, DATA_ATTIVAZIONE, '
    'UTENTE_ATTIVAZIONE.'],
   ['Tutte le tabelle versionate', 'Lettura e copia', 'ANSC_CFG_UC, ANSC_CFG_SEZIONE, '
    'ANSC_CFG_CAMPO, ANSC_CFG_FORMULA, ALLEGATI_USECASE e RICONCILIAZ_DIZIONARI sono '
    'duplicate alla creazione della bozza.'],
   ['ANSC_ANA_UC', 'Lettura per controllo', 'Alimenta la voce «UC ritirati e ancora '
    'adottati» del report.'],
   ['ANSC_STATO_ATTO', 'Lettura', 'ID_VERSIONE sull’atto è un timbro storico: sopravvive '
    'all’archiviazione delle versioni vecchie.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Stato', 'Colonna', 'Bozza, attiva, storica. Un indice unico su funzione garantisce che '
    'l’attiva sia una sola.'],
   ['NUOVA BOZZA DA ATTIVA', 'Azione di pagina', 'Crea la bozza per copia. Disabilitata se '
    'una bozza esiste già: le bozze non si accumulano.'],
   ['Report d’impatto', 'Pannello', 'Calcolato sulla bozza. Ogni voce è cliccabile e porta '
    'alla pagina che consente di risolverla.'],
   ['ATTIVA LA VERSIONE', 'Azione', '⚠️ Consentita anche con scostamenti aperti, ma con '
    'conferma esplicita che ne elenca le conseguenze. Scrive COD_STATO, DATA_ATTIVAZIONE e '
    'UTENTE_ATTIVAZIONE, e porta la precedente a storica.'],
   ['CONFRONTA CON L’ATTIVA', 'Azione', 'Mostra lo scostamento riga per riga fra bozza e '
    'attiva.'],
   ['Cestino', 'Azione di riga', 'Elimina la bozza. Non disponibile su attiva e storiche.']],
  '⚠️ **Chi possa attivare una versione è materia dell’organizzazione, non del progetto.** '
  'La pagina offre la funzione e registra chi l’ha esercitata; la delega è un punto aperto '
  'dell’analisi.'),

 ('Comandi e operazioni di servizio', 'bo_comandi.png',
  ['Alcune operazioni sono lunghe, rare e non interattive: lo scarico dei dizionari, '
   'l’importazione del mapping, le verifiche di copertura, la riconciliazione massiva. È '
   'stato dichiarato che siano eseguibili **anche da riga di comando**; questa pagina ne '
   'governa l’avvio dall’interfaccia e ne conserva l’esito.',
   'La doppia via non è una ridondanza: la riga di comando serve alle esecuzioni pianificate '
   'e all’assistenza, l’interfaccia serve al funzionario che deve poter avviare '
   'un’importazione senza chiedere a nessuno. Entrambe scrivono nello stesso registro, ed è '
   'il registro a rendere la pagina utile.',
   '⚠️ L’opzione «Solo simulazione» è predefinita a sì sui comandi che scrivono. Un comando '
   'che tocca decine di migliaia di righe deve poter essere provato prima di essere eseguito.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_LOG_AUDIT**', 'Lettura e scrittura', 'Il registro delle esecuzioni: FASE come '
    'nome del comando, ESITO, DURATA_MS, OPERATORE, DATA; RICHIESTA e RISPOSTA conservano '
    'parametri e riepilogo.'],
   ['DOMINIO_DECODIFICA, VALORE_DOMINIO', 'Scrittura', 'Destinazione dello scarico dei '
    'dizionari.'],
   ['ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO, ALLEGATI_USECASE, ANSC_CFG_FORMULA', 'Scrittura',
    'Destinazione dell’importazione del mapping, nella sola versione in bozza.'],
   ['ANSC_STATO_ATTO', 'Lettura e scrittura', 'Destinazione della riconciliazione massiva.'],
   ['ANSC_CFG_VERSIONE', 'Lettura', 'La versione di destinazione dei comandi che scrivono '
    'configurazione.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Esito', 'Colonna', 'Ok, con rilievi, in errore, e gli esiti propri del comando (per '
    'esempio «nuova revisione» per il controllo delle revisioni di ANSC).'],
   ['Occhio', 'Azione di riga', 'Apre il verbale dell’esecuzione: parametri, righe toccate, '
    'scostamenti, errori.'],
   ['Versione di destinazione', 'Campo', '⚠️ Limitata alle versioni in stato BOZZA: nessun '
    'comando scrive sull’attiva.'],
   ['Solo simulazione', 'Campo', 'Esegue senza scrivere e produce il verbale degli effetti '
    'che avrebbe avuto.'],
   ['AVVIA', 'Azione', 'Avvia il comando in secondo piano. La pagina non attende: l’esito '
    'compare nel registro.']],
  None),

 ('Numerazione comunale', 'bo_numerazione.png',
  ['Il numero comunale dell’atto è a carico del Comune: **nessun servizio di ANSC lo '
   'assegna**. Alla scala di Roma più operatori lavorano insieme, e un’assegnazione presa al '
   'momento del deposito produrrebbe collisioni o salti.',
   'La pagina riprende quasi alla lettera la funzione osservata in S.I.De., dove i numeri si '
   'richiedono in blocco per registro e anno e restano disponibili finché non vengono '
   'consumati o restituiti. È anche la funzione che il committente ha chiesto in modo '
   'esplicito: poter staccare il numero da SIPO prima di usarlo altrove.',
   '⚠️ Per il parto plurimo i numeri vanno staccati **in blocco e in anticipo**, perché la '
   'prenotazione in ANSC li richiede tutti insieme. È il caso che rende la funzione '
   'necessaria e non comoda.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_STATO_ATTO**', 'Lettura e scrittura', 'NUM_COMUNALE è la colonna che la pagina '
    'governa; ID_ANSC dice se il numero è stato consumato da un atto depositato.'],
   ['Registro dei numeri', '⚠️ Da definire', 'Il modello attuale **non ha** una tabella per i '
    'numeri staccati e non ancora consumati: NUM_COMUNALE vive sull’atto. Serve una '
    'struttura propria, o una riga di atto in stato iniziale. È un punto aperto.'],
   ['ANSC_CFG_UC', 'Lettura', 'Registro e serie a cui il numero appartiene (SERIE).']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['Numero di elementi', 'Campo', 'Quanti numeri staccare in un colpo. Per il parto plurimo '
    'si indica il numero di nati.'],
   ['RICHIEDI', 'Azione', 'Stacca il blocco e lo registra come disponibile, con operatore e '
    'momento.'],
   ['Stato', 'Colonna', 'Disponibile, consumato, restituito. Un numero consumato porta '
    'l’identificativo ANSC dell’atto che lo ha usato.'],
   ['Spunta', 'Azione di riga', 'Assegna il numero a un atto in lavorazione.'],
   ['Croce', 'Azione di riga', 'Restituisce il numero. ⚠️ La restituzione è esplicita: la '
    'numerazione non tollera salti e un numero abbandonato va dichiarato tale.']],
  None),

 ('Pagine e menu', 'bo_amministrazione.png',
  ['È la pagina che governa il registro descritto nel capitolo «Il registro delle pagine e '
   'del menu»: decide che cosa compare nella home, in quale ordine, con quale aspetto e per '
   'chi. È l’unica pagina dell’applicazione che non tratta dati di stato civile ma '
   'l’applicazione stessa.',
   'L’elenco mostra tutte le pagine dichiarate, comprese quelle che **non compaiono nel '
   'menu** — il dettaglio di un atto si raggiunge da un elenco, non da una mattonella — e '
   'comprese quelle di aree **non ancora rilasciate**, che restano dichiarate con stato '
   'apposito. È così che il registro documenta anche ciò che manca.',
   '⚠️ Una scelta di disegno da dichiarare: il registro **non** contiene l’indirizzo da cui '
   'si scarica il micro-frontend. Quell’indirizzo cambia da un ambiente all’altro e '
   'appartiene alla configurazione di esercizio, non al dato applicativo. Il registro dice '
   'quale applicazione ospita una pagina; dove quell’applicazione si trovi lo dice la '
   'configurazione della shell, come spiegato nel capitolo su Kubernetes.'],
  [['Tabella', 'Uso', 'Che cosa se ne mostra'],
   ['**ANSC_CFG_PAGINA**', 'Lettura e scrittura', 'TITOLO, SOTTOTITOLO, COD_ICONA, ROTTA, '
    'URL_GUIDA, COD_STATO, FLG_IN_MENU, NUM_ORDINE, DATA_INIZIO_VALIDITA e '
    'DATA_FINE_VALIDITA.'],
   ['**ANSC_CFG_APPLICAZIONE**', 'Lettura e scrittura', 'COD_APPLICAZIONE e DESCRIZIONE '
    'alimentano la colonna e il filtro; lo stato dell’applicazione prevale su quello della '
    'pagina.'],
   ['**ANSC_CFG_PAGINA_ABILITAZ**', 'Lettura e scrittura', 'COD_ABILITAZIONE e COD_TIPO: '
    'l’abilitazione di accesso e quelle delle singole azioni.'],
   ['**ANSC_CFG_TEMA**', 'Lettura', 'Il tema associato all’applicazione, mostrato in sola '
    'lettura nel dettaglio.']],
  [['Elemento', 'Genere', 'Comportamento'],
   ['In menu', 'Colonna e campo', 'Da FLG_IN_MENU. Una pagina raggiungibile ma non esposta '
    'come mattonella ha «no»: è il caso dei dettagli.'],
   ['Ordine', 'Campo', 'NUM_ORDINE. Determina la posizione della mattonella. Le pagine '
    'fuori menu non lo usano.'],
   ['Stato', 'Campo', 'Attiva, sospesa, non rilasciata. ⚠️ «Non rilasciata» dichiara una '
    'pagina prevista il cui codice non esiste ancora: compare qui ma non nel menu, e serve a '
    'non perdere il disegno.'],
   ['Abilitazione di accesso', 'Campo', 'Una sola per pagina. Senza di essa la mattonella '
    'non compare e la rotta non si carica.'],
   ['Abilitazioni di azione', 'Campo', 'Più di una: governano i singoli bottoni della '
    'pagina. ⚠️ Nascondono il bottone, non autorizzano: il controllo resta sul servizio.'],
   ['Guida in linea', 'Campo', 'Indirizzo della scheda di aiuto mostrata dal punto '
    'interrogativo della pagina.'],
   ['ANTEPRIMA DEL MENU', 'Azione di pagina', 'Mostra la home come apparirebbe a un profilo '
    'scelto. ⚠️ È il controllo che evita di accorgersi di un errore dal racconto di un '
    'operatore.'],
   ['SALVA', 'Azione', 'Scrive immediatamente. Il menu si rilegge al prossimo caricamento: '
    'non serve un rilascio.']],
  '⚠️ **L’accesso a questa pagina va trattato come amministrativo.** Chi la governa decide '
  'che cosa gli altri vedono, e per errore potrebbe nascondere un’area a tutti. Per questo '
  'l’anteprima esiste, e per questo ogni modifica è tracciata con utente e momento.'),
]

COPERTURA = [
 ['Tabella', 'Pagine che la governano', 'Copertura'],
 ['ANSC_STATO_ATTO', 'Supervisione atti · Dettaglio atto · Numerazione comunale',
  'Completa in lettura; in scrittura per fase, stato, emergenza e numero comunale.'],
 ['ANSC_LOG_AUDIT', 'Dettaglio atto (scheda Cronologia) · Comandi di servizio',
  'Completa in lettura. La scrittura è del concentratore, non dell’interfaccia.'],
 ['ALLEGATO', 'Allegati dell’atto', 'Completa.'],
 ['ALLEGATI_USECASE', 'Configurazione UC (scheda Allegati) · Allegati dell’atto',
  'Completa: configurata nella prima, applicata nella seconda.'],
 ['ANSC_NOTIFICA', 'Gestione delle notifiche', 'Completa.'],
 ['ANSC_CFG_UC', 'Configurazione UC — elenco e dettaglio', 'Completa.'],
 ['ANSC_ANA_UC', 'Catalogo dei casi d’uso', 'Sola lettura, come deve essere.'],
 ['ANSC_CFG_SEZIONE', 'Configurazione UC (scheda Sezioni) · Dettaglio atto',
  'Completa in configurazione; in sola lettura nel dettaglio dell’atto.'],
 ['ANSC_CFG_CAMPO', 'Configurazione UC (scheda Campi) · Dettaglio atto', 'Completa.'],
 ['ANSC_CFG_FORMULA', 'Configurazione UC (scheda Formule)', 'Completa.'],
 ['ANSC_CFG_VERSIONE', 'Versioni della configurazione', 'Completa.'],
 ['TIPO_LOGICHE_DI_SCELTA', 'Logiche di scelta',
  '⚠️ In sola lettura: i tre domini sono d’impianto. Se dovessero diventare configurabili '
  'servirebbe una pagina di amministrazione dedicata.'],
 ['LOGICHE_DI_SCELTA', 'Logiche di scelta', 'Completa, con simulazione.'],
 ['DOMINIO_DECODIFICA', 'Dizionari ANSC', 'Sola lettura: si aggiorna da ANSC.'],
 ['VALORE_DOMINIO', 'Dizionari ANSC', 'Sola lettura: si aggiorna da ANSC.'],
 ['RICONCILIAZ_DIZIONARI', 'Riconciliazione delle decodifiche', 'Completa.'],
 ['ANSC_CFG_APPLICAZIONE', 'Pagine e menu',
  'Completa. ⚠️ Nuova in questa versione: non appartiene al modello dell’integrazione ma a '
  'quello dell’applicazione.'],
 ['ANSC_CFG_PAGINA', 'Pagine e menu', 'Completa. Nuova in questa versione.'],
 ['ANSC_CFG_PAGINA_ABILITAZ', 'Pagine e menu', 'Completa. Nuova in questa versione.'],
 ['ANSC_CFG_TEMA', 'Pagine e menu',
  'In sola lettura dall’interfaccia: i temi sono pochi e d’impianto. Nuova in questa '
  'versione.'],
]

APERTI = [
 ['#', 'Questione', 'Perché è aperta'],
 ['BO-1', '**Chiuso in questa versione.** Granularità dei micro-frontend',
  'Deciso: quattro unità di rilascio, una per area funzionale, con le pagine come rotte a '
  'caricamento differito. La motivazione e l’alternativa scartata sono nel capitolo «Il '
  'contesto tecnico».'],
 ['BO-2', 'Se la conferma delle notifiche debba restare solo singola',
  'S.I.De. offre anche la conferma massiva della pagina. La scelta della conferma singola è '
  'motivata, ma il costo cresce con il volume, e il volume alla scala di Roma non è noto.'],
 ['BO-3', 'Dove vivano i numeri comunali staccati e non ancora consumati',
  'Il modello non ha una tabella per il blocco di numeri: oggi NUM_COMUNALE vive sull’atto. '
  'La pagina della numerazione presuppone una struttura che va aggiunta o surrogata.'],
 ['BO-4', 'Se la numerazione comunale appartenga al back-office o alle maschere di SIPO',
  'È una funzione operativa, non di supervisione. Qui è nel back-office perché è trasversale '
  'ai registri; in S.I.De. sta nel percorso dell’atto.'],
 ['BO-5', 'Quale sia il vocabolario delle abilitazioni',
  'Ogni pagina e ogni azione di questo documento presuppone un’abilitazione che le consenta. '
  'L’elenco non esiste ancora: è la stessa lacuna registrata nell’analisi del front-end.'],
 ['BO-6', 'Se il registro delle postazioni rientri nel perimetro',
  'La mattonella «Postazioni e certificati» compare nel menu perché un registro delle '
  'postazioni con certificati esiste già in SIPO, e alimenta il canale verso ANPR. Se debba '
  'servire anche l’integrazione ANSC, e con quali schermate, non è deciso.'],
 ['BO-7', 'Come si presenta una sezione condizionata che non si applica',
  'Nel dettaglio dell’atto una sezione la cui logica di presenza è falsa può essere nascosta '
  'o mostrata attenuata. Nasconderla è più pulito; mostrarla spiega perché un campo atteso '
  'non è richiesto.'],
 ['BO-8', 'Se serva una pagina di consultazione degli atti conclusi',
  'Il back-office tratta le eccezioni. Ritrovare un atto andato a buon fine oggi richiede le '
  'maschere di SIPO, che però non cercano per identificativo nazionale.'],
 ['BO-9', 'Da dove proviene il vocabolario delle icone',
  'Il registro conserva un codice di icona (COD_ICONA), ma l’insieme dei codici ammessi '
  'appartiene alla libreria condivisa. Se la libreria non ne pubblica un catalogo, il campo '
  'diventa testo libero e una svista si vede solo a video.'],
 ['BO-10', 'Se il registro debba valere anche per le altre applicazioni dell’ente',
  'Le quattro tabelle stanno nel nostro schema e governano il nostro back-office. Se le '
  'nuove applicazioni di Roma Capitale saranno molte, un registro d’ente sarebbe più '
  'corretto: la forma qui adottata è compatibile con quella promozione, ma la decisione non '
  'compete a questo progetto.'],
 ['BO-11', 'Chi amministra il registro',
  'La pagina «Pagine e menu» decide che cosa gli altri vedono e per errore può nascondere '
  'un’area a tutti. Va stabilito se l’abilitazione competa al funzionario configuratore o a '
  'un profilo amministrativo distinto.'],
 ['BO-12', 'Se la testata debba consentire il cambio di ruolo attivo',
  'Il disegno prevede che il ruolo attivo sia mostrato e commutabile, perché da esso '
  'dipendono le azioni disponibili. ⚠️ In SIPO oggi il comportamento è deciso dal **primo** '
  'ruolo dell’utente, senza che nessuno lo scelga: adottare la commutazione esplicita è un '
  'miglioramento, ma va confermato che sia voluto.'],
]


def scrivi(d):
    # ─────────────────────────────────────────────────────────── 1. scopo
    h(d, 1, 'Scopo del documento')
    par(d, 'Questo documento disegna **l’applicazione di back-office** del componente che '
           'integra SIPO con ANSC. Parte dal menu e scende fino alla singola pagina; per '
           'ciascuna mostra un wireframe e ne dichiara il funzionamento: quali tabelle sono '
           'sottese, che cosa fa ogni campo, che cosa fa ogni bottone e quando è disabilitato.')
    par(d, 'Non è un progetto grafico. I wireframe dicono **che cosa** una pagina contiene e '
           'come sono disposte le informazioni; il rivestimento — colori, tipografia, '
           'spaziature — appartiene al design system della libreria condivisa e non è materia '
           'di questo documento. Dove il disegno assume qualcosa che non è stato deciso, lo '
           'dichiara e lo registra fra i punti aperti anziché darlo per stabilito.')
    par(d, '⚠️ **Il documento non introduce funzioni nuove.** Le pagine esistono per governare '
           'le sedici tabelle già progettate nell’analisi dell’integrazione; il criterio di '
           'completezza è che ciascuna di quelle tabelle abbia almeno una pagina che la '
           'governi, e il capitolo «La copertura del modello dati» lo verifica riga per riga.')

    h(d, 2, 'Perimetro e metodo')
    par(d, 'Il disegno si appoggia a tre basi di evidenza. La prima è il **modello dati** '
           'dell’analisi dell’integrazione, da cui vengono le sedici tabelle e i loro campi. '
           'La seconda sono gli **schermi di S.I.De.**, il sistema di stato civile del Comune '
           'di Milano, raccolti nella cartella delle sorgenti: da lì vengono le convenzioni '
           'di interazione, perché sono già in esercizio su un dominio identico. La terza '
           'sono gli schermi della **web app di ANSC**, che mostrano che cosa la piattaforma '
           'nazionale chiede e come lo chiede.')
    par(d, 'Il metodo è quello di non inventare dove si può osservare. Dove S.I.De. risolve '
           'un problema che abbiamo anche noi — la sessione a tempo, i filtri richiudibili, '
           'lo stato di completezza per sezione, lo stacco dei numeri — il disegno ne riprende '
           'la soluzione e lo dichiara. ⚠️ Dove invece la nostra analisi ha deciso '
           'diversamente da S.I.De., il documento lo segnala: non sono sviste, sono scelte, e '
           'due di esse sono fra i punti aperti.')

    h(d, 2, 'Glossario')
    tabella(d, [
        ['Termine', 'Significato in questo documento'],
        ['Back-office', 'L’applicazione qui disegnata: supervisione delle eccezioni e '
                        'governo della configurazione. **Non** è la superficie su cui si '
                        'compila un atto, che resta nelle maschere di SIPO.'],
        ['Pagina', 'Una rotta dell’applicazione, con il proprio indirizzo e le proprie '
                   'abilitazioni.'],
        ['Scheda', 'Una linguetta interna a una pagina di dettaglio. Non cambia rotta.'],
        ['UC', 'Il caso d’uso, struttura di ANSC, identificato da un codice numerico.'],
        ['Modello', 'La struttura del Comune di Roma che indirizza un atto a uno o più UC.'],
        ['Versione', 'Una baseline della configurazione: bozza, attiva o storica.'],
        ['Logica di scelta', 'Un’espressione valutata sui dati dell’atto già salvato in '
                             'SIPO, per scegliere l’UC, condizionare una sezione o '
                             'trasformare un campo.'],
        ['Sessione OTP', 'Il codice temporaneo, valido quattro ore, che abilita il dialogo '
                         'con ANSC. Non è un accesso all’applicazione.'],
        ['Abilitazione', 'La stringa che il profilo dell’operatore porta e che consente una '
                         'pagina o un’azione.'],
    ], larghezze=[1.4, 5.1])

    h(d, 2, 'Riferimenti')
    tabella(d, [
        ['#', 'Riferimento', 'Natura'],
        ['[R1]', 'ANALISI_Integrazione-ANSC_v3.30.docx — impianto, modello dati e decisioni '
                 'dell’integrazione SIPO→ANSC', 'Documento di progetto'],
        ['[R2]', 'ANALISI_Front-End-Angular_v0.4.docx — micro-frontend, libreria condivisa, '
                 'autorizzazione per abilitazioni', 'Documento di progetto'],
        ['[R3]', 'ASIS_Autenticazione-Profilazione_SIPO_v0.2.docx — che cosa SIPO offre oggi '
                 'in fatto di identità e permessi', 'Documento di progetto'],
        ['[R4]', 'SIDE/Grafica — schermate del sistema di stato civile del Comune di Milano',
         'Evidenza osservata'],
        ['[R5]', 'SIDE/webAppANSC — schermate della web app di ANSC in preproduzione',
         'Evidenza osservata'],
        ['[R6]', 'SPEC_API_Integrazione-ANSC_v0.1.docx e i contratti OpenAPI dei tre pod',
         'Documento di progetto'],
    ], larghezze=[0.55, 4.45, 1.5])

    # ────────────────────────────────────────────── 2. impianto documentale
    h(d, 1, 'L’impianto documentale')
    par(d, 'Il documento è costruito su una forma ripetuta. Dopo il contesto e la mappa '
           'dell’applicazione, **ogni pagina ha la propria sezione**, e ogni sezione ha '
           'sempre le stesse quattro parti, nello stesso ordine.')
    voce(d, 'Il wireframe.',
         'Che cosa la pagina mostra e come è disposto. È uno schizzo funzionale: le colonne, '
         'i filtri e i bottoni disegnati sono quelli che il modello dati consente, non '
         'quelli che starebbero bene.')
    voce(d, 'Lo scopo.',
         'Perché la pagina esiste, che problema risolve e quali scelte di disegno non sono '
         'ovvie. È la parte che spiega, e dove il disegno diverge da un’alternativa '
         'ragionevole ne dà la ragione.')
    voce(d, 'Le tabelle sottese.',
         'Quali tabelle la pagina legge e quali scrive, e **quali colonne** di ciascuna. È '
         'la parte che rende il documento verificabile: una pagina che mostra una colonna '
         'inesistente si riconosce da qui.')
    voce(d, 'Campi, colonne e azioni.',
         'Il comportamento di ciascun elemento interattivo: che cosa fa, che cosa scrive, '
         'quando è disabilitato. ⚠️ È la parte che serve a chi costruisce, e che di solito '
         'manca nei documenti di interfaccia.')
    par(d, 'Chiudono il documento due capitoli di controllo: la **copertura del modello dati**, '
           'che verifica che nessuna tabella sia rimasta senza pagina, e i '
           '**punti aperti**, che raccolgono le decisioni che il disegno ha dovuto assumere '
           'e che vanno confermate.')

    # ──────────────────────────────────────────────────── 3. contesto
    h(d, 1, 'Il contesto tecnico')
    h(d, 2, 'Il back-office come insieme di micro-frontend')
    par(d, 'Le nuove applicazioni del Comune sono impostate come micro-frontend Angular '
           'ospitati da una shell che ne condivide la libreria [R2]. Il back-office non è '
           'però **un** micro-frontend: è **quattro**, uno per area funzionale.')
    figura(d, 'bo_mfe.png',
           'Figura 1 — La composizione. La shell ospita quattro remote; le pagine sono rotte '
           'a caricamento differito interne a ciascuno. Il registro governa il menu e non '
           'dipende da come il codice è impacchettato.')
    par(d, 'La ripartizione segue le aree già dichiarate nella mappa dell’applicazione, e la '
           'ragione è che **le aree hanno ritmi di rilascio davvero diversi**. L’area '
           'operativa cambia quando cambia il percorso di formazione dell’atto; quella di '
           'configurazione quando ANSC rivede il mapping, in media ogni diciassette giorni; '
           'quella di servizio quasi mai. Rilasciarle insieme significherebbe far passare per '
           'il collaudo dell’area operativa una correzione a una schermata di configurazione.')
    tabella(d, [
        ['Unità di rilascio', 'Pagine', 'Che cosa la fa cambiare'],
        ['mfe-operativa', 'Supervisione atti · Dettaglio atto · Allegati · Numerazione',
         'Il percorso di formazione dell’atto e le categorie di eccezione.'],
        ['mfe-notifiche', 'Gestione delle notifiche',
         'Il contratto del servizio di adempimenti di ANSC.'],
        ['mfe-configurazione', 'UC elenco e dettaglio · Catalogo · Logiche · Dizionari · '
                               'Riconciliazione · Versioni',
         'Le revisioni del mapping di ANSC e l’evoluzione del modello di configurazione.'],
        ['mfe-servizio', 'Comandi e operazioni · Pagine e menu',
         'Quasi nulla: è l’area più stabile.'],
    ], larghezze=[1.4, 2.8, 2.3])
    par(d, 'Ciò che la shell fornisce, e che nessuno dei quattro deve disegnare: la testata, '
           'il piè di pagina, il profilo con le sue abilitazioni, la sessione e le rotte di '
           'primo livello. **Il micro-frontend disegna dalle briciole di pane in giù.** I '
           'wireframe mostrano comunque la testata, perché una pagina si valuta per come '
           'appare all’operatore e non per come è ripartita fra componenti.')
    par(d, '⚠️ La libreria condivisa va caricata **una volta sola**. Se i quattro remote ne '
           'portassero ciascuno una copia, oltre al peso si avrebbero quattro servizi di '
           'sessione distinti, e il distintivo OTP di un’area non saprebbe ciò che un’altra '
           'ha appena acquisito. È il motivo per cui la libreria è dichiarata come dipendenza '
           'condivisa in singola istanza, con la versione allineata fra i quattro.')

    h(d, 2, 'Perché non un micro-frontend per pagina')
    par(d, 'L’ipotesi di far corrispondere **una pagina a un’applicazione** è stata valutata '
           'e scartata. Vale la pena dire perché, perché è un’idea ragionevole che non regge '
           'alla prova dei conti.')
    voce(d, 'L’indipendenza comprata non si usa.',
         'Il rilascio indipendente serve quando due parti cambiano per ragioni diverse. Qui '
         '«Dizionari» e «Riconciliazione» cambiano per la stessa ragione — una revisione del '
         'mapping — e nessuno rilascerebbe l’una senza l’altra. L’indipendenza si pagherebbe '
         'senza esercitarla.')
    voce(d, 'Il costo è quattordici volte, non una.',
         'Quattordici catene di costruzione, quattordici artefatti da versionare, '
         'quattordici manifesti da tenere allineati sulla stessa versione della libreria. '
         '⚠️ Il disallineamento di versione fra remote è il difetto tipico di questa '
         'architettura, e cresce con il numero di remote.')
    voce(d, 'La navigazione interna diventa un confine.',
         'Passare dall’elenco degli atti al dettaglio di un atto è la transizione più '
         'frequente dell’applicazione. Con due remote distinti attraversa un confine fra '
         'applicazioni: o si accetta un caricamento, o si costruisce un raccordo che fa '
         'sembrare unito ciò che si è appena diviso.')
    par(d, 'Il criterio adottato è quindi: **si divide dove cambia il motivo per cambiare**, '
           'non dove cambia la schermata. ⚠️ Resta vero, e va detto, che la divisione per area '
           'non è irreversibile: una pagina può essere estratta in un remote proprio se un '
           'giorno acquisisse un ciclo di vita suo, e il registro descritto nel capitolo '
           'seguente rende quell’operazione invisibile all’operatore, perché il menu non '
           'dipende da come il codice è impacchettato.')
    par(d, '⚠️ Il **caricamento differito** resta necessario a qualunque granularità. Un '
           'back-office di quindici pagine, molte delle quali visitate di rado, è il caso in '
           'cui caricare tutto all’avvio è più costoso che utile: la guardia che protegge la '
           'rotta deve impedire il **caricamento**, non solo l’accesso. È una raccomandazione '
           'già registrata nel documento del front-end.')

    h(d, 2, 'Autorizzazione e sessione')
    par(d, 'Due meccanismi distinti governano che cosa un operatore vede e che cosa può fare, '
           'e non vanno confusi perché falliscono in modi diversi.')
    voce(d, 'Le abilitazioni.',
         'Decidono **che cosa si vede**. Una mattonella di area la cui abilitazione manca non '
         'è disabilitata: è assente. Una colonna di azioni mostra solo le icone consentite. '
         '⚠️ Una guardia lato client non è un controllo di sicurezza — nasconde, non impedisce '
         '— e il controllo vero resta sui servizi.')
    voce(d, 'La sessione OTP.',
         'Decide **che cosa si può fare ora**. Vale quattro ore ed è personale. Le azioni che '
         'dialogano con ANSC sono disabilitate quando manca, con la ragione visibile; il '
         'distintivo in testa alla pagina ne mostra lo stato e il tempo residuo.')
    par(d, '⚠️ La distinzione ha una ricaduta di disegno che conviene enunciare: **un’azione '
           'non consentita dal profilo non si mostra, un’azione non eseguibile ora si mostra '
           'disabilitata.** Nascondere la seconda farebbe credere che la funzione non esista; '
           'disabilitare la prima rivelerebbe l’esistenza di funzioni che non competono.')
    par(d, 'Sotto la soglia di guardia — una proposta ragionevole è quindici minuti — conviene '
           'che il distintivo inviti a rinnovare la sessione **prima** di iniziare '
           'un’operazione lunga, invece di lasciarla interrompere a metà.')

    h(d, 2, 'Le convenzioni ereditate da S.I.De.')
    par(d, 'Sette convenzioni sono riprese quasi alla lettera, perché risolvono problemi che '
           'abbiamo anche noi e perché sono già state provate su operatori di stato civile.')
    tabella(d, [
        ['Convenzione', 'Come è ripresa'],
        ['Home a mattonelle su tre colonne',
         'Indice delle aree, con una riga che dice che cosa c’è dietro. Nessun contatore.'],
        ['Briciole di pane sempre presenti',
         'Home › Integrazione ANSC › area › elemento. Sono la sola via di risalita.'],
        ['Titolo con distintivo di sessione e azioni a destra',
         'Il titolo porta il distintivo OTP; le azioni di pagina stanno a destra, non fra i '
         'filtri.'],
        ['Pannelli di filtro richiudibili',
         'Aperti al primo accesso, richiusi dopo la ricerca per lasciare spazio ai risultati.'],
        ['Tabella con colonna «Azioni» a icone',
         'Spunta, croce, occhio, matita, cestino. Sempre nell’ultima colonna, sempre nello '
         'stesso ordine.'],
        ['Paginazione in basso a destra',
         'Righe per pagina, intervallo mostrato e totale. Il totale è informazione di lavoro: '
         'dice quanto resta da fare.'],
        ['Stato di completezza per sezione',
         'Spunta o croce accanto al nome della sezione nel dettaglio dell’atto.'],
    ], larghezze=[2.0, 4.5])
    par(d, '⚠️ Due convenzioni di S.I.De. sono invece state **scartate**, e la ragione non è '
           'estetica. La prima è la **conferma massiva delle notifiche**: darebbe l’apparenza '
           'di un atto unico dove ce ne sono molti, ciascuno con effetti su persone diverse. '
           'La seconda è il **numero comunale digitato a mano** nella conferma dell’evento, '
           'che sulla web app di ANSC è un campo con incremento manuale: alla scala di Roma '
           'la concorrenza fra operatori lo rende impraticabile, e da qui la pagina di '
           'numerazione.')

    # ──────────────────────────────────── 3bis. testata e piè di pagina
    h(d, 1, 'La testata e il piè di pagina')
    par(d, 'Testata e piè di pagina **sono della shell** e valgono uguali per ogni '
           'applicazione dell’ente: è la scelta che garantisce all’operatore di trovare le '
           'stesse informazioni nello stesso posto quando passa da un applicativo all’altro. '
           'Il back-office vi contribuisce soltanto la parte contestuale. Questo capitolo ne '
           'specifica il contenuto in forma di requisito verso la shell.')
    figura(d, 'bo_chrome.png',
           'Figura 2 — Anatomia della testata e del piè di pagina, con la ripartizione fra '
           'ciò che fornisce la shell e ciò che contribuisce il micro-frontend.')
    h(d, 2, 'Che cosa deve contenere la testata')
    tabella(d, [
        ['Elemento', 'Chi lo fornisce', 'Perché serve'],
        ['Stemma ed ente', 'Shell', 'Identifica l’amministrazione. Obbligatorio per la PA.'],
        ['Nome dell’applicazione', 'Shell',
         'Distingue il back-office dalle altre applicazioni ospitate dalla stessa shell.'],
        ['**Ambiente**', 'Shell',
         '⚠️ DEV, COLLAUDO o ESERCIZIO, in evidenza. È l’elemento che impedisce l’errore più '
         'costoso: operare in esercizio credendo di essere in collaudo. In esercizio '
         'l’etichetta può essere attenuata, ma negli altri ambienti deve essere '
         '**impossibile da non vedere**.'],
        ['Sede di lavoro', 'Shell',
         'Municipio e sportello dell’operatore. Determina l’ambito organizzativo e quindi '
         'che cosa vede negli elenchi.'],
        ['**Tipo di utente**', 'Shell',
         'Il ruolo con cui si sta operando: ufficiale di stato civile, funzionario '
         'configuratore, supporto. ⚠️ Se un operatore possiede più ruoli, la testata deve '
         'dire **quale è attivo** e consentirne il cambio, perché da quello dipendono le '
         'azioni disponibili.'],
        ['Utente', 'Shell', 'Nome e cognome, con accesso al profilo e alla disconnessione.'],
        ['Briciole di pane', 'Micro-frontend', 'La sola via di risalita nella gerarchia.'],
        ['Titolo e sottotitolo della pagina', 'Micro-frontend',
         'Provengono dal registro, non dal codice: si cambiano senza rilascio.'],
        ['**Distintivo di sessione con tempo residuo**', 'Micro-frontend',
         '⚠️ Non basta dire se la sessione è attiva: deve dire **quanto le resta**. La '
         'sessione dura quattro ore e la web app di ANSC lo dichiara espressamente; un '
         'operatore che sta per iniziare una sequenza lunga deve poter decidere di rinnovarla '
         'prima. Sotto la soglia di guardia il distintivo cambia aspetto e invita al rinnovo.'],
        ['Azioni della pagina', 'Micro-frontend',
         'A destra del titolo, mai fra i filtri.'],
    ], larghezze=[1.7, 1.2, 3.6])
    h(d, 2, 'Che cosa deve contenere il piè di pagina')
    tabella(d, [
        ['Elemento', 'Perché serve'],
        ['Ente e struttura di riferimento', 'Obbligo istituzionale.'],
        ['**Versione e identificativo di build**',
         '⚠️ Non è un ornamento. Durante un aggiornamento progressivo in Kubernetes '
         'convivono due versioni dell’applicazione, e un difetto segnalato da un operatore è '
         'irriproducibile se non si sa quale stava usando. È il primo dato che l’assistenza '
         'chiede.'],
        ['Riferimenti per l’assistenza', 'Dove segnalare un malfunzionamento.'],
        ['Note legali, privacy, accessibilità',
         'Obbligatori per un servizio della pubblica amministrazione.'],
    ], larghezze=[2.0, 4.5])
    par(d, '⚠️ Una cosa che testata e piè di pagina **non devono contenere**: contatori di '
           'lavorazioni pendenti. Un numero in testata è letto come un compito, viene '
           'ricalcolato a ogni cambio di pagina e diventa una richiesta continua al server '
           'per un dato che nessuno ha chiesto. Il numero di notifiche da confermare sta '
           'nella pagina delle notifiche, dove accanto c’è l’azione.')

    # ──────────────────────────────── 3ter. il registro delle pagine
    h(d, 1, 'Il registro delle pagine e del menu')
    par(d, 'Il menu non è scritto nel codice: è **un dato**. Quattro tabelle dello schema '
           'ANSC_USR dichiarano quali applicazioni esistono, quali pagine contengono, come si '
           'presentano e chi le può usare; la home si costruisce leggendole.')
    par(d, 'La ragione di questa scelta è che le cose che si vogliono cambiare più spesso — '
           'il titolo di una voce, il suo posto nel menu, l’icona, chi può vederla, se sia '
           'attiva — sono le stesse che nel codice richiederebbero un rilascio. Portandole '
           'nel dato si ottiene che **una riorganizzazione del menu non sia un rilascio**, e '
           'che la mappa dell’applicazione sia ispezionabile da una tabella invece che '
           'leggendo i sorgenti.')
    par(d, 'C’è un secondo effetto, meno evidente e più importante: **il registro disaccoppia '
           'il menu dall’impacchettamento**. Se un giorno una pagina venisse estratta in un '
           'micro-frontend proprio, o due aree venissero fuse, il registro cambierebbe una '
           'riga e l’operatore non se ne accorgerebbe. È ciò che rende reversibile la scelta '
           'di granularità discussa nel capitolo precedente.')

    h(d, 2, 'Le quattro tabelle')
    tabella(d, [
        ['Tabella', 'Che cosa dichiara'],
        ['**ANSC_CFG_APPLICAZIONE**',
         'Il micro-frontend: codice, descrizione, nome del remote, stato, ordine e tema '
         'associato. Una riga per unità di rilascio — oggi quattro.'],
        ['**ANSC_CFG_PAGINA**',
         'La pagina: titolo e sottotitolo del menu, icona, rotta, guida in linea, stato, se '
         'compaia nel menu, ordine e validità temporale.'],
        ['**ANSC_CFG_PAGINA_ABILITAZ**',
         'Chi può usarla: l’abilitazione di accesso alla pagina e quelle delle singole '
         'azioni. Più righe per pagina.'],
        ['**ANSC_CFG_TEMA**',
         'L’aspetto: il tema della libreria condivisa da applicare e i pochi valori che '
         'possono variare per applicazione.'],
    ], larghezze=[1.9, 4.6])
    par(d, '⚠️ I nomi seguono la grammatica dello standard aziendale già adottata dal resto '
           'del modello: prefisso di modulo, marcatore di tipo «CFG», nome dell’oggetto. '
           'Qualificate dallo schema si leggono ANSC_USR.ANSC_CFG_PAGINA, come le altre '
           'tabelle di configurazione.')

    h(d, 2, 'Le strutture')
    ddl(d, """CREATE TABLE ANSC_USR.ANSC_CFG_APPLICAZIONE (
  ID_APPLICAZIONE   NUMBER GENERATED ALWAYS AS IDENTITY,
  COD_APPLICAZIONE  VARCHAR2(30 CHAR)   NOT NULL,   -- OPERATIVA, NOTIFICHE, ...
  DESCRIZIONE       VARCHAR2(100 CHAR)  NOT NULL,
  NOME_REMOTE       VARCHAR2(60 CHAR)   NOT NULL,   -- il nome esposto dal remote
  COD_STATO         VARCHAR2(20 CHAR)   NOT NULL,
  NUM_ORDINE        NUMBER(4)           NOT NULL,
  ID_TEMA           NUMBER,
  DATA_INS          DATE  DEFAULT SYSDATE NOT NULL,
  DATA_UPD          DATE,
  UTENTE_INS        VARCHAR2(30 CHAR)   NOT NULL,
  UTENTE_UPD        VARCHAR2(30 CHAR),
  CONSTRAINT PK_ansc_cfg_applicazione PRIMARY KEY (ID_APPLICAZIONE),
  CONSTRAINT UK_ansc_cfg_applicazione UNIQUE (COD_APPLICAZIONE),
  CONSTRAINT FK_ansc_cfg_applicaz_tema FOREIGN KEY (ID_TEMA)
      REFERENCES ANSC_USR.ANSC_CFG_TEMA (ID_TEMA),
  CONSTRAINT CK_ansc_cfg_applicaz_stato
      CHECK (COD_STATO IN ('ATTIVA','SOSPESA','NON_RILASCIATA','RITIRATA'))
) TABLESPACE ANSC_USR;""")
    d.add_paragraph()
    ddl(d, """CREATE TABLE ANSC_USR.ANSC_CFG_PAGINA (
  ID_PAGINA             NUMBER GENERATED ALWAYS AS IDENTITY,
  ID_APPLICAZIONE       NUMBER              NOT NULL,
  COD_PAGINA            VARCHAR2(40 CHAR)   NOT NULL,
  TITOLO                VARCHAR2(100 CHAR)  NOT NULL,
  SOTTOTITOLO           VARCHAR2(200 CHAR),
  COD_ICONA             VARCHAR2(40 CHAR),
  ROTTA                 VARCHAR2(200 CHAR)  NOT NULL,
  URL_GUIDA             VARCHAR2(300 CHAR),
  COD_STATO             VARCHAR2(20 CHAR)   NOT NULL,
  FLG_IN_MENU           CHAR(1 CHAR)        DEFAULT 'S' NOT NULL,
  NUM_ORDINE            NUMBER(4),
  DATA_INIZIO_VALIDITA  DATE                DEFAULT SYSDATE NOT NULL,
  DATA_FINE_VALIDITA    DATE,
  DATA_INS              DATE  DEFAULT SYSDATE NOT NULL,
  DATA_UPD              DATE,
  UTENTE_INS            VARCHAR2(30 CHAR)   NOT NULL,
  UTENTE_UPD            VARCHAR2(30 CHAR),
  CONSTRAINT PK_ansc_cfg_pagina PRIMARY KEY (ID_PAGINA),
  CONSTRAINT UK_ansc_cfg_pagina UNIQUE (ID_APPLICAZIONE, COD_PAGINA),
  CONSTRAINT FK_ansc_cfg_pagina_applic FOREIGN KEY (ID_APPLICAZIONE)
      REFERENCES ANSC_USR.ANSC_CFG_APPLICAZIONE (ID_APPLICAZIONE),
  CONSTRAINT CK_ansc_cfg_pagina_stato
      CHECK (COD_STATO IN ('ATTIVA','SOSPESA','NON_RILASCIATA','RITIRATA')),
  CONSTRAINT CK_ansc_cfg_pagina_menu   CHECK (FLG_IN_MENU IN ('S','N'))
) TABLESPACE ANSC_USR;""")
    d.add_paragraph()
    ddl(d, """CREATE TABLE ANSC_USR.ANSC_CFG_PAGINA_ABILITAZ (
  ID_PAGINA_ABILITAZ  NUMBER GENERATED ALWAYS AS IDENTITY,
  ID_PAGINA           NUMBER              NOT NULL,
  COD_ABILITAZIONE    VARCHAR2(60 CHAR)   NOT NULL,
  COD_TIPO            VARCHAR2(10 CHAR)   NOT NULL,   -- ACCESSO | AZIONE
  COD_AZIONE          VARCHAR2(40 CHAR),              -- valorizzato se COD_TIPO = AZIONE
  DATA_INS            DATE  DEFAULT SYSDATE NOT NULL,
  DATA_UPD            DATE,
  UTENTE_INS          VARCHAR2(30 CHAR)   NOT NULL,
  UTENTE_UPD          VARCHAR2(30 CHAR),
  CONSTRAINT PK_ansc_cfg_pagina_abilitaz PRIMARY KEY (ID_PAGINA_ABILITAZ),
  CONSTRAINT UK_ansc_cfg_pagina_abilitaz
      UNIQUE (ID_PAGINA, COD_ABILITAZIONE, COD_AZIONE),
  CONSTRAINT FK_ansc_cfg_abilitaz_pagina FOREIGN KEY (ID_PAGINA)
      REFERENCES ANSC_USR.ANSC_CFG_PAGINA (ID_PAGINA),
  CONSTRAINT CK_ansc_cfg_abilitaz_tipo  CHECK (COD_TIPO IN ('ACCESSO','AZIONE')),
  CONSTRAINT CK_ansc_cfg_abilitaz_coer
      CHECK ((COD_TIPO = 'ACCESSO' AND COD_AZIONE IS NULL)
          OR (COD_TIPO = 'AZIONE'  AND COD_AZIONE IS NOT NULL))
) TABLESPACE ANSC_USR;""")
    d.add_paragraph()
    ddl(d, """CREATE TABLE ANSC_USR.ANSC_CFG_TEMA (
  ID_TEMA          NUMBER GENERATED ALWAYS AS IDENTITY,
  COD_TEMA         VARCHAR2(30 CHAR)   NOT NULL,   -- il tema della libreria condivisa
  DESCRIZIONE      VARCHAR2(100 CHAR)  NOT NULL,
  COD_COLORE_BARRA VARCHAR2(7 CHAR),               -- deroga per ambiente o applicazione
  TXT_TOKEN        CLOB,                           -- scostamenti dal tema, in JSON
  FLG_PREDEFINITO  CHAR(1 CHAR)        DEFAULT 'N' NOT NULL,
  DATA_INS         DATE  DEFAULT SYSDATE NOT NULL,
  DATA_UPD         DATE,
  UTENTE_INS       VARCHAR2(30 CHAR)   NOT NULL,
  UTENTE_UPD       VARCHAR2(30 CHAR),
  CONSTRAINT PK_ansc_cfg_tema PRIMARY KEY (ID_TEMA),
  CONSTRAINT UK_ansc_cfg_tema UNIQUE (COD_TEMA),
  CONSTRAINT CK_ansc_cfg_tema_predef CHECK (FLG_PREDEFINITO IN ('S','N')),
  CONSTRAINT CK_ansc_cfg_tema_json   CHECK (TXT_TOKEN IS JSON)
) TABLESPACE ANSC_USR;""")
    par(d, '⚠️ Sul tema conviene essere sobri. **La grafica appartiene alla libreria '
           'condivisa**, e consentire di ridefinirla da tabella riaprirebbe la strada a '
           'quattro applicazioni che si somigliano senza essere uguali. Qui si configura '
           '**quale** tema applicare e pochi scostamenti dichiarati — tipicamente il colore '
           'della barra per distinguere l’ambiente a colpo d’occhio. Il campo degli '
           'scostamenti è vincolato a contenere JSON valido, così un errore di compilazione '
           'si ferma alla scrittura e non alla resa.')

    h(d, 2, 'Come si costruisce il menu')
    par(d, 'La home chiede al concentratore le pagine visibili per il profilo corrente; il '
           'servizio interroga le quattro tabelle e restituisce **solo ciò che l’operatore '
           'può vedere**, già ordinato. La selezione applica, in questo ordine, quattro '
           'condizioni.')
    voce(d, 'Lo stato dell’applicazione.',
         'Se l’applicazione è sospesa o non rilasciata, nessuna delle sue pagine compare, '
         'qualunque sia lo stato della singola pagina.')
    voce(d, 'Lo stato e la validità della pagina.',
         'Attiva e dentro l’intervallo di validità. La validità temporale consente di '
         'programmare l’apparizione di una funzione senza presidiare il momento.')
    voce(d, 'Il contrassegno di presenza nel menu.',
         'Le pagine di dettaglio esistono nel registro — perché hanno una rotta, '
         'un’abilitazione e una guida — ma non sono mattonelle.')
    voce(d, 'L’abilitazione di accesso.',
         'Il profilo deve possederla. ⚠️ **Non si restituisce una voce disabilitata**: non si '
         'restituisce affatto, perché l’elenco di ciò che non si può fare è a sua volta '
         'un’informazione.')
    par(d, 'Il risultato è memorizzato per la durata della sessione del browser e riletto a '
           'ogni accesso, così una modifica dal back-office si vede subito senza costringere '
           'a una richiesta per ogni cambio di pagina.')

    h(d, 2, 'Che cosa il registro non deve fare')
    par(d, 'Tre limiti vanno dichiarati, perché un registro che governa il menu è facilmente '
           'scambiato per un registro che governa la sicurezza.')
    voce(d, '⚠️ Non è un controllo di accesso.',
         'Le abilitazioni del registro decidono che cosa **si vede**. Il controllo vero sta '
         'sui servizi, che verificano l’autorizzazione a ogni chiamata. Togliere una voce dal '
         'menu non impedisce a nessuno di raggiungere la rotta: è la guardia lato client a '
         'non caricarla, e il servizio a rifiutare.')
    voce(d, '⚠️ Non contiene l’indirizzo dei remote.',
         'Dove si scarichi un micro-frontend cambia da un ambiente all’altro e appartiene '
         'alla configurazione di esercizio. Metterlo in tabella significherebbe avere una '
         'base dati diversa per ambiente, che è il difetto che si voleva evitare.')
    voce(d, '⚠️ Non sostituisce il rilascio.',
         'Si può sospendere una pagina, rinominarla, spostarla, nasconderla. **Non si può '
         'aggiungere una pagina che non esiste**: la riga «non rilasciata» dichiara '
         'l’intenzione, il codice deve comunque essere scritto e distribuito.')

    # ──────────────────────────────────────────── 4. mappa dell'applicazione
    h(d, 1, 'La mappa dell’applicazione')
    par(d, 'L’applicazione ha **quindici pagine** distribuite su quattro aree funzionali. '
           'La ripartizione non è tematica ma per **ritmo di lavoro**: ciò che si guarda ogni '
           'giorno, ciò che si guarda quando qualcosa arriva, ciò che si tocca poche volte '
           'l’anno, ciò che si esegue.')
    tabella(d, [
        ['Area', 'Pagine', 'Chi la usa e quando'],
        ['Operativa', 'Supervisione atti · Dettaglio atto · Allegati dell’atto · '
                      'Numerazione comunale',
         'L’ufficiale e il supporto, ogni giorno. È l’area che giustifica l’esistenza '
         'dell’applicazione.'],
        ['In ingresso', 'Gestione delle notifiche',
         'L’ufficiale, quando ANSC notifica. Il ritmo lo detta ANSC, non il Comune.'],
        ['Configurazione', 'Configurazione UC (elenco e dettaglio) · Catalogo UC · '
                           'Logiche di scelta · Dizionari · Riconciliazione · Versioni',
         'Il funzionario configuratore. Poche volte l’anno, con revisioni del mapping ogni '
         'diciassette giorni in media.'],
        ['Servizio', 'Comandi e operazioni · Pagine e menu · Audit e tracciato',
         'Il supporto e il fornitore, in diagnosi o su richiesta.'],
    ], larghezze=[1.2, 2.9, 2.4])
    par(d, '⚠️ **Audit e tracciato** non ha una sezione propria in questo documento: la '
           'cronologia di un atto è una scheda del suo dettaglio e il registro dei comandi è '
           'parte della pagina dei comandi. Una pagina di audit trasversale — «tutte le '
           'chiamate in un intervallo» — è utile alla diagnosi ma non aggiunge disegno: '
           'sarebbe l’ennesimo elenco con filtri sulla stessa tabella.')

    h(d, 2, 'La home')
    figura(d, 'bo_home.png',
           'Figura 3 — La home. Le mattonelle sono rotte, non riepiloghi: nessun contatore '
           'promette una misura che il back-office non è tenuto a calcolare.')

    # ──────────────────────────────────────────────────── 5. le pagine
    h(d, 1, 'Le pagine')
    par(d, 'Seguono le quindici pagine nella forma dichiarata nell’impianto documentale. '
           'La home, già mostrata, apre l’elenco per completezza.')
    for i, (titolo, png, scopo, tab, elem, nota) in enumerate(PAGINE, start=1):
        h(d, 2, f'{i}. {titolo}')
        if png != 'bo_home.png':
            # le prime due figure sono l'architettura e la cornice; la home è la terza.
            # Da qui la numerazione prosegue con l'indice della pagina più due.
            figura(d, png, f'Figura {i + 2} — {titolo}.')
        for p in scopo:
            par(d, p)
        par(d, 'Tabelle sottese', 'Heading 3')
        tabella(d, tab, larghezze=[1.7, 1.3, 3.5])
        par(d, 'Campi, colonne e azioni', 'Heading 3')
        tabella(d, elem, larghezze=[1.6, 1.1, 3.8])
        if nota:
            par(d, nota)

    # ─────────────────────────────────────────────────── 6. copertura
    h(d, 1, 'La copertura del modello dati')
    par(d, 'Il criterio di completezza è che ciascuna tabella progettata abbia almeno una pagina '
           'che la governi: le sedici del modello dell’integrazione più le quattro del '
           'registro introdotto da questa versione. La verifica segue.')
    tabella(d, COPERTURA, larghezze=[1.5, 2.6, 2.4])
    par(d, '**Venti tabelle su venti sono coperte.** Due osservazioni sulla qualità della '
           'copertura, che il conteggio da solo nasconde.')
    voce(d, 'La copertura in scrittura è più stretta di quella in lettura, ed è voluto.',
         'I dizionari e il catalogo degli UC si leggono soltanto: sono repliche di ciò che '
         'ANSC dichiara, e consentirne la modifica locale creerebbe divergenze invisibili.')
    voce(d, '⚠️ Una struttura manca al modello, e la pagina della numerazione la rende '
            'visibile.',
         'Non esiste una tabella dei numeri comunali staccati e non ancora consumati: oggi il '
         'numero vive sull’atto. Finché non c’è, la pagina della numerazione presuppone '
         'qualcosa che il modello non offre — ed è il punto aperto BO-3.')

    # ───────────────────────────────────────────── 7. pattern trasversali
    h(d, 1, 'I comportamenti trasversali')
    par(d, 'Alcune decisioni valgono per tutte le pagine e si enunciano una volta sola, '
           'perché ripeterle in quindici sezioni le renderebbe invisibili.')
    tabella(d, [
        ['Situazione', 'Comportamento adottato'],
        ['Elenco vuoto',
         'Non una tabella con zero righe, ma un testo che dice perché è vuoto e che cosa '
         'fare: «Nessun atto si è fermato nel periodo scelto» è diverso da «Nessun '
         'risultato».'],
        ['Azione non consentita dal profilo', 'Non si mostra.'],
        ['Azione non eseguibile ora',
         'Si mostra disabilitata, con la ragione leggibile senza doverla indovinare '
         '(sessione assente, versione non in bozza, stato dell’atto incompatibile).'],
        ['Azione irreversibile',
         'Conferma esplicita che ne enuncia la conseguenza. ⚠️ Per l’annullamento di un atto '
         'la conferma dice che l’identificativo nazionale resta consumato.'],
        ['Errore restituito da ANSC',
         'Si mostra il codice e il testo di ANSC, non una riscrittura: il codice è ciò che '
         'consente al fornitore e all’assistenza di ritrovare il caso.'],
        ['Operazione lunga',
         'Non si attende: si avvia, e l’esito compare nel registro dei comandi. La pagina '
         'resta utilizzabile.'],
        ['Dato che viene da ANSC',
         'Si mostra come ANSC lo dichiara, anche quando la nostra nomenclatura sarebbe più '
         'chiara. Lo stato dell’atto ne è l’esempio: è una replica, non un giudizio.'],
        ['Modifica su versione non attiva',
         'La pagina dichiara in testa su quale versione si sta lavorando. È il promemoria che '
         'rende innocue le schermate di configurazione.'],
    ], larghezze=[1.7, 4.8])

    # ────────────────────────────────────────────────── 9bis. kubernetes
    h(d, 1, 'Indicazioni per l’esercizio in Kubernetes')
    par(d, 'L’applicazione è destinata a un impianto Kubernetes, e quattro micro-frontend '
           'più una shell pongono qualche questione che non si presenterebbe con '
           'un’applicazione unica. Si riportano le sole indicazioni che **incidono sul '
           'disegno**; il dimensionamento e la topologia appartengono al documento di '
           'analisi dell’integrazione.')

    h(d, 2, 'Che cosa si distribuisce')
    par(d, 'Un micro-frontend, una volta costruito, è **un insieme di file statici**: non è '
           'un processo applicativo. Ciascuno dei quattro è un’unità di distribuzione '
           'autonoma con un servizio che ne espone i file, e la shell li compone nel browser '
           'dell’operatore.')
    tabella(d, [
        ['Oggetto', 'Quante istanze', 'Nota'],
        ['Shell', 'Una distribuzione',
         'Espone la rotta radice e carica i remote. È l’unico che deve conoscere gli '
         'indirizzi degli altri.'],
        ['I quattro micro-frontend', 'Una distribuzione ciascuno',
         'File statici serviti da un servente leggero. ⚠️ Nessuno stato: si possono '
         'replicare e ricreare liberamente.'],
        ['Concentratore', 'Già previsto dall’analisi',
         'Non è oggetto di questo documento: il back-office ne è un consumatore.'],
    ], larghezze=[1.8, 1.8, 2.9])
    par(d, '⚠️ **Le risorse vanno chieste per quello che sono**: servire file statici '
           'consuma poca memoria e quasi nessun processore. Dimensionare questi contenitori '
           'come un’applicazione Java significherebbe sprecare il quadruplo di quanto serve, '
           'moltiplicato per quattro applicazioni.')

    h(d, 2, 'La configurazione che cambia per ambiente')
    par(d, 'È il punto che decide se si dovrà ricostruire l’applicazione per ogni ambiente o '
           'no, e la risposta deve essere no.')
    voce(d, 'Gli indirizzi dei remote.',
         'La shell deve sapere dove trovare i quattro micro-frontend, e l’indirizzo cambia '
         'fra sviluppo, collaudo ed esercizio. ⚠️ Va letto **all’avvio da configurazione '
         'montata nel contenitore**, non incorporato nella costruzione: altrimenti lo stesso '
         'artefatto non può essere promosso da un ambiente al successivo, che è il modo in '
         'cui un artefatto dovrebbe viaggiare.')
    voce(d, 'L’etichetta dell’ambiente.',
         'DEV, COLLAUDO o ESERCIZIO in testata proviene dalla stessa configurazione. È '
         'l’esempio più semplice e più dimostrativo: se fosse compilata dentro '
         'l’applicazione, un artefatto collaudato non sarebbe lo stesso che va in esercizio.')
    voce(d, 'Gli indirizzi dei servizi.',
         'Il back-office chiama il concentratore. Se shell e micro-frontend sono esposti '
         'dallo stesso ingresso di rete del servizio, le chiamate restano di stessa origine e '
         'non serve alcuna concessione per richieste incrociate. ⚠️ È la disposizione da '
         'preferire, perché una concessione aperta fra origini diverse è una superficie in '
         'più da presidiare.')

    h(d, 2, 'Il versionamento e l’aggiornamento progressivo')
    par(d, '⚠️ È la questione specifica di questa architettura, e va affrontata in fase di '
           'disegno perché si manifesta come un difetto incomprensibile.')
    par(d, 'Durante un aggiornamento progressivo convivono per qualche minuto la versione '
           'vecchia e la nuova. Un operatore che ha caricato la shell prima '
           'dell’aggiornamento e apre un’area dopo può ricevere un remote di versione '
           'diversa da quella che la shell si aspetta. Il sintomo è una schermata bianca o un '
           'errore che non si riproduce: la peggiore specie di difetto.')
    tabella(d, [
        ['Accorgimento', 'Perché'],
        ['Il manifesto del remote non si mette in cache',
         'Il file che elenca i moduli va servito con divieto di conservazione, mentre i file '
         'con impronta nel nome possono essere conservati a lungo. ⚠️ Invertire le due '
         'politiche è l’errore che fa vedere agli operatori una versione vecchia per ore.'],
        ['La shell rileva il cambio di versione',
         'Se la versione dichiarata dal servente non è quella caricata, la shell lo segnala e '
         'propone di ricaricare. Meglio un avviso che un errore.'],
        ['La versione è nel piè di pagina',
         'È il dato che rende diagnosticabile una segnalazione: senza, la stessa schermata '
         'può venire da due versioni diverse.'],
        ['Compatibilità all’indietro per un rilascio',
         'Un remote nuovo deve continuare a funzionare con la shell precedente per il tempo '
         'dell’aggiornamento. È un vincolo di disciplina, non di strumento.'],
    ], larghezze=[2.0, 4.5])

    h(d, 2, 'Sonde, riservatezza e rete')
    voce(d, 'Sonde di pronto e di vivo.',
         'Un servente di file statici è pronto quando risponde: una richiesta alla radice '
         'basta. ⚠️ Non si usi come sonda una pagina che a sua volta chiami il concentratore, '
         'altrimenti un disservizio del back-end farebbe riavviare i front-end, che non '
         'c’entrano.')
    voce(d, 'Nessun segreto nel front-end.',
         'Il codice servito al browser è leggibile da chiunque. ⚠️ Il certificato del Comune, '
         'il codice di sessione e le credenziali restano nel concentratore; il front-end non '
         'ne detiene alcuno. È già l’impianto stabilito dall’analisi e qui si ribadisce '
         'perché la tentazione di «passare la configurazione al front-end» si presenta '
         'proprio quando si introduce un registro.')
    voce(d, 'Il registro si raggiunge per servizio, non per base dati.',
         'La home chiede il menu al concentratore, che legge le quattro tabelle. ⚠️ Nessun '
         'micro-frontend parla con Oracle: sarebbe impossibile e non va nemmeno progettato.')
    voce(d, 'Regole di rete.',
         'I front-end hanno bisogno di ricevere richieste dall’ingresso e di nient’altro; le '
         'chiamate ai servizi partono dal browser dell’operatore, non dai contenitori che '
         'servono i file. È un perimetro molto stretto, e conviene dichiararlo tale.')

    # ─────────────────────────────────────────────────── 8. punti aperti
    h(d, 1, 'Punti aperti di disegno')
    par(d, 'Le decisioni che il disegno ha dovuto assumere per poter procedere, e che vanno '
           'confermate prima di realizzare. Sono numerate con il prefisso BO per non '
           'confondersi con i registri degli altri documenti.')
    tabella(d, APERTI, larghezze=[0.5, 1.9, 4.1])


if __name__ == '__main__':
    d = costruisci()
    for t in d.tables:
        if t.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in t.rows:
                v = {'Progetto': 'Integrazione SIPO – ANSC',
                     'Data consegna': '26/09/2026', 'Versione': '0.2',
                     'Documento': 'DISEGNO_Back-Office_ANSC_v0.2'}.get(r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    st = D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')
    for r in list(st.rows)[1:]:
        if not any(c.text.strip() for c in r.cells):
            r._tr.getparent().remove(r._tr)
    D.storia(d, '26/09/2026', '0.2',
             'Contesto tecnico · Testata e piè (nuovo) · Registro (nuovo) · '
             'Pagine · Kubernetes (nuovo) · Punti aperti',
             'Definita la granularità dei micro-frontend: quattro unità di rilascio, una '
             'per area, con motivazione dell’alternativa scartata (BO-1 chiuso). Nuovo '
             'capitolo sulla testata e sul piè di pagina, con il contenuto obbligatorio '
             'posto come requisito verso la shell. Nuovo capitolo sul registro delle pagine '
             'e del menu, con le quattro tabelle e le relative DDL, e nuova pagina «Pagine e '
             'menu» che le governa. Nuovo capitolo di indicazioni per l’esercizio in '
             'Kubernetes. Le pagine passano da quattordici a quindici, le tabelle coperte da '
             'sedici a venti, i punti aperti da otto a dodici.')
    D.storia(d, '26/09/2026', '0.1', 'Tutti',
             'Prima stesura. Disegno dell’applicazione di back-office a partire dal menu: '
             'quattordici pagine, ciascuna con wireframe, tabelle sottese e comportamento di '
             'campi e azioni. Le convenzioni di interazione sono riprese da S.I.De.; la '
             'copertura delle sedici tabelle del modello è verificata in un capitolo '
             'dedicato. Otto punti aperti di disegno.')
    scrivi(d)
    d.save(OUT)
    print('scritto:', os.path.relpath(OUT, BASE))
    print('capitoli/tabelle/immagini:', D.riepilogo(OUT))
