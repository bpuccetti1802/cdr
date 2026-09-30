# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.26 -> v3.27 (24/09/2026).

Il back-office non era stato rivisto dopo le v3.23-3.26 e diceva cose che il modello non
dice più. Questa versione lo riallinea, rifà i wireframe e aggiunge due cose:

  · i COMANDI: ogni comando è disponibile dal back-office e da riga di comando, con le
    stesse regole e la stessa traccia;
  · il RACCORDO della riconciliazione con le tabelle di SIPO, verificato sul codice.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_27_backoffice_e_raccordo.py
"""
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.26.docx')
DST = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.27.docx')
IMG = os.path.join(BASE, 'strumenti', 'img')

if os.path.exists(DST):
    os.remove(DST)
shutil.copy(SRC, DST)
doc = docx.Document(DST)
fatti = []


def par(inizio):
    t = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(t) != 1:
        raise SystemExit(f'attesa 1 occorrenza di «{inizio[:60]}», trovate {len(t)}')
    return t[0]


def riga_con(t, prima_cella):
    for r in t.rows:
        if r.cells[0].text.strip() == prima_cella:
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


def inserisci_riga(t, dopo, valori):
    nuova = D.clona_riga(t, valori)
    riga_con(t, dopo)._tr.addnext(nuova._tr)
    return nuova


MOD = D.trova_tabella(doc, 'colonna', 'tipo', 'note')


# ═════════════════════════ 1. scheda e storia
for t in doc.tables:
    if t.rows[0].cells[0].text.strip().lower().startswith(('area organizzativa', 'progetto')):
        for nome, val in (('Data consegna', '24/09/2026'), ('Versione', '3.27')):
            try:
                D.riscrivi_cella(riga_con(t, nome).cells[1], val)
            except SystemExit:
                pass

D.storia(doc, '24/09/2026', '3.27',
         'Back-office · Gestione dei dizionari · Open Point · Wireframe',
         'Il back-office è riallineato al modello: la scelta dell’UC da parte dell’operatore '
         'in caso di ambiguità, la fase della lavorazione fra i filtri della worklist, il '
         'numero comunale e la logica applicata nel dettaglio dell’atto, le schermate delle '
         'sezioni e della riconciliazione dei valori, le fasi di audit aggiornate e la '
         'nomenclatura per caso d’uso al posto di quella per operazioni. Dichiarato che ogni '
         'comando è disponibile tanto dal back-office quanto da riga di comando, con le stesse '
         'regole di identificazione e di tracciamento. Aggiunto il raccordo fra le decodifiche '
         'da riconciliare e le tabelle di configurazione di SIPO, verificato sul codice. '
         'Rifatti i wireframe e aggiunti quelli della riconciliazione e dei comandi.')


# ═════════════════════════ 2. l'UC ambiguo: la regola nuova
t = D.trova_tabella(doc, 'categoria', 'stato reale tipico', 'causa')
r = riga_con(t, 'UC ambiguo')
D.riscrivi_cella(r.cells[3], 'Il sistema presenta all’operatore gli UC candidati e gli '
                             'consente di sceglierne uno: la scelta è registrata come tale '
                             '(COD_ORIGINE_UC = OPERATORE) e l’atto prosegue. ⚠️ È una via '
                             'd’uscita, non la soluzione: la configurazione resta ambigua e va '
                             'corretta, e la presenza di scelte manuali è la misura del '
                             'lavoro arretrato.')
inserisci_riga(t, 'UC ambiguo', (
    'Valore non traducibile', 'IN_PREPARAZIONE (in ANSC non esiste nulla)',
    'Un campo codificato porta un valore di SIPO per il quale la riconciliazione non dichiara '
    'alcun corrispondente ANSC, oppure lo dichiara con una validità scaduta.',
    'Si completa la riconciliazione dei valori per quella decodifica. ⚠️ Non si forza il '
    'deposito: un valore inventato supera il pre-filtro e viene respinto da ANSC, o peggio '
    'accettato con un significato diverso.', 'Nessuno: l’atto non raggiunge ANSC'))

D.sostituisci(doc, 'Quando si attivano più regole con esiti diversi, l’UC è ambiguo. Qui la '
                   'causa è sempre nella configurazione, mai nel dato',
              'Quando si attivano più regole con esiti diversi, l’UC è ambiguo. La causa è '
              'sempre nella configurazione, mai nel dato')
D.para(doc, par('Entrambe le condizioni si prevengono con la simulazione')._p,
       '⚠️ Su questa seconda condizione il disegno ammette una via d’uscita presidiata, che la '
       'prima non ha: poiché gli UC candidati sono noti e la loro descrizione è pubblicata da '
       'ANSC, il sistema li presenta all’operatore e gli consente di sceglierne uno, '
       'registrando la scelta come manuale. È ammessa soltanto qui, perché soltanto qui il '
       'sistema sa che la risposta sta fra due alternative dichiarate; e resta il segnale di '
       'una configurazione da correggere, non una modalità ordinaria di lavoro.')
fatti.append('UC ambiguo: regola nuova; aggiunta la categoria «Valore non traducibile»')


# ═════════════════════════ 3. la worklist e il dettaglio
D.sostituisci(doc, 'filtrabile per municipio, ufficiale, categoria di eccezione e stato reale '
                   'in ANSC',
              'filtrabile per municipio, ufficiale, categoria di eccezione, fase della '
              'lavorazione (COD_FASE) e stato reale in ANSC', attese=1,
              etichetta='worklist: filtro per fase', fatti=fatti)
D.sostituisci(doc, 'ogni riga porta l’atto, l’evento, lo stato reale in ANSC',
              'ogni riga porta l’atto, l’evento, la fase raggiunta, lo stato reale in ANSC',
              attese=1, etichetta='worklist: colonna della fase', fatti=fatti)
D.sostituisci(doc, 'mostra lo stato reale in ANSC accertato con R005 (con l’idAnsc assegnato)',
              'mostra lo stato reale in ANSC accertato con R005 — con l’identificativo '
              'nazionale, il numero comunale, l’UC determinato e la logica che l’ha prodotto, '
              'e la baseline con cui il payload è stato costruito —', attese=1,
              etichetta='dettaglio atto: i dati conservati', fatti=fatti)


# ═════════════════════════ 4. le schermate
t = D.trova_tabella(doc, 'schermata', 'scopo', 'azioni principali')
r = riga_con(t, 'Configurazione — Catalogo delle logiche')
D.riscrivi_cella(r.cells[1], 'TIPO_LOGICHE_DI_SCELTA e LOGICHE_DI_SCELTA: le logiche dei tre '
                             'domini — scelta dell’UC, attivazione di una sezione, '
                             'trasformazione di un campo — ciascuna con la propria descrizione '
                             'e l’espressione che la realizza.')
r = riga_con(t, 'Configurazione — Campi, allegati e formule per UC')
D.riscrivi_cella(r.cells[0], 'Configurazione — Sezioni, campi, allegati e formule per UC')
D.riscrivi_cella(r.cells[1], 'ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO, ALLEGATI_USECASE e '
                             'ANSC_CFG_FORMULA per singolo UC: quali blocchi del modello '
                             'evento si popolano e a quale condizione, da dove viene ogni '
                             'dato, quali documenti servono e quali diciture sono previste.')
D.riscrivi_cella(r.cells[2], 'Precompila dal foglio (sezioni, campi, allegati e formule '
                             'insieme), esamina e conferma riga per riga, completa la '
                             'corrispondenza con SIPO, richiama la logica del campo dal '
                             'catalogo, apre la riconciliazione dei valori dove il campo è '
                             'codificato.')
nuova = D.clona_riga(t, (
    'Configurazione — Riconciliazione dei valori',
    'RICONCILIAZ_DIZIONARI: per ogni decodifica ANSC, che cosa diventa un valore di SIPO, con '
    'la condizione e la validità. Affianca la decodifica replicata, che dice quali valori ANSC '
    'ammette, alla tabella di SIPO che ospita i valori locali.',
    'Consulta per decodifica, dichiara una corrispondenza, chiude la validità di una '
    'corrispondenza superata, apre la decodifica ANSC e la tabella SIPO a confronto.',
    'Admin'))
riga_con(t, 'Configurazione — Sezioni, campi, allegati e formule per UC')._tr.addnext(nuova._tr)
D.clona_riga(t, (
    'Comandi ed esecuzioni',
    'I comandi del componente con il loro storico: allineamento dei dizionari nelle tre '
    'modalità, importazione della configurazione, attivazione di una baseline, scarico delle '
    'notifiche, riconciliazione di un atto. Per ciascuno è dichiarato l’equivalente da riga di '
    'comando.',
    'Esegui, consulta lo storico con esito e conteggi, scarica il registro delle esecuzioni.',
    'Amministratore'))
fatti.append('schermate: sezioni, riconciliazione, comandi, catalogo a tre domini')

# la nomenclatura rimasta indietro
D.sostituisci(doc, 'catalogo degli UC, configurazione dei Modelli e dei campi per UC con '
                   'versionamento, condizioni di applicabilità degli UC, regole di controllo, '
                   'dizionari ANSC',
              'catalogo degli UC, configurazione dei casi d’uso e delle sezioni, campi, '
              'allegati e formule con versionamento, catalogo delle logiche, riconciliazione '
              'dei valori, dizionari ANSC, comandi ed esecuzioni', attese=1,
              etichetta='mappa dei menu', fatti=fatti)
t = D.trova_tabella(doc, 'ruolo', 'responsabilità')
D.riscrivi_cella(riga_con(t, 'Amministratore').cells[1],
                 'Configurazione (casi d’uso, sezioni, campi, allegati, formule, logiche, '
                 'riconciliazione dei valori, versioni), importazione dai fogli, aggiornamento '
                 'dei dizionari ANSC su comando nelle tre modalità e consultazione dello '
                 'storico, esecuzione dei comandi dal back-office o da riga di comando, '
                 'ambienti e connettività, registro delle postazioni autorizzate, politiche di '
                 'conservazione.')
t = D.trova_tabella(doc, 'attività', 'schermate')
D.riscrivi_cella(riga_con(t, 'Verifica').cells[1],
                 'Pre-filtro locale (RF-9) prima del deposito; Dettaglio atto (esito R009/R007); '
                 'Audit, con le fasi del flusso operativo: DETERMINAZIONE, PREVERIFICA, '
                 'ALLEGATI, SOGGETTO, PAYLOAD, ANTEPRIMA, DEPOSITO, FIRMA_DICH, FIRMA_USC, '
                 'ANNULLAMENTO, RICONCILIAZIONE.')
D.riscrivi_cella(riga_con(t, 'Amministrazione').cells[1],
                 'Configurazione (casi d’uso, sezioni, campi, allegati, formule, versioni), '
                 'Catalogo delle logiche, Riconciliazione dei valori, Importazione della '
                 'configurazione, Dizionari ANSC (aggiornamento su comando e storico), Comandi '
                 'ed esecuzioni, Sistema & Connettività.')
fatti.append('menu, ruoli e attività: nomenclatura allineata')


# ═════════════════════════ 5. i comandi, anche da riga di comando
ancora = D.h(doc, 2, 'Il governo della configurazione alla scala del dominio')._p
D.para(doc, ancora, 'I comandi: dal back-office e da riga di comando', stile='Heading 2')
for t in [
    'Alcune attività del componente non sono operazioni su un singolo atto ma comandi che '
    'agiscono su un insieme: allineare i dizionari, importare una configurazione, attivare una '
    'baseline, scaricare le notifiche, accertare lo stato di un atto. Il back-office li offre '
    'come pulsanti; ⚠️ gli stessi comandi devono essere eseguibili da riga di comando, e non '
    'come funzione di ripiego ma come seconda via dichiarata.',
    'La ragione è operativa. Un aggiornamento dei dizionari si fa volentieri fuori orario; una '
    'importazione di sessantamila righe non si presidia davanti a un browser; e il giorno in '
    'cui l’interfaccia non è raggiungibile, il componente deve restare governabile. Chi '
    'conduce l’esercizio lavora con strumenti di automazione, e un comando che esiste solo '
    'dentro una pagina web non vi si può collocare.',
    '⚠️ La regola che rende la cosa sostenibile è una sola: le due vie sono la stessa '
    'funzione, non due realizzazioni. Stesso codice, stessi controlli, stessa traccia. Un '
    'comando che dalla riga salti la verifica che il back-office esegue è un modo per '
    'introdurre di nascosto uno stato che nessuna schermata sa spiegare.',
]:
    D.para(doc, ancora, t)
D.tabella(doc, ancora, [
    ('Comando', 'Che cosa fa', 'Chi lo esegue', 'Nota'),
    ('Allineamento dei dizionari', 'R901 nelle tre modalità: completa, selettiva, '
     'simulazione.', 'Amministratore',
     'La simulazione non scrive nulla: è il modo di valutare l’impatto prima di aggiornare.'),
    ('Importazione della configurazione', 'Legge il foglio e prepara le righe in stato '
     'proposto.', 'Amministratore',
     'Riscrive soltanto le righe che nessuno ha ancora esaminato; per le altre produce un '
     'elenco di scostamenti.'),
    ('Attivazione di una baseline', 'Porta una bozza in ATTIVA, dopo i controlli.',
     'Amministratore (OP-37: chi la autorizzi è decisione organizzativa)',
     'Un indice unico garantisce che l’attiva sia una sola: il comando non può violarlo.'),
    ('Scarico delle notifiche', 'Interroga R008 e aggiorna lo store locale.',
     'Amministratore o pianificazione',
     '⚠️ Eseguibile senza presidio solo se OP-23 si chiude in senso favorevole.'),
    ('Riconciliazione di un atto', 'R005: accerta lo stato reale e allinea quello locale.',
     'Supporto', 'È una lettura: non modifica nulla in ANSC.'),
    ('⚠️ Ciò che non è un comando', 'Deposito, firme e annullamento.', 'USC in «Finalizza»',
     'Sono atti dell’ufficiale in sessione OTP: non esistono né come pulsante di back-office '
     'né come comando di riga.'),
], MOD)
D.para(doc, ancora, 'I comandi del componente, con chi li esegue.', corsivo=True)
for t in [
    'Tre condizioni valgono per entrambe le vie. La prima: l’esecuzione è sempre attribuita a '
    'una persona, anche dalla riga di comando, dove l’identità va fornita esplicitamente — un '
    'comando eseguito da nessuno non è ricostruibile. La seconda: ogni esecuzione lascia una '
    'riga nell’audit con il proprio esito e i propri conteggi, e la schermata dei comandi non '
    'fa che leggere quella traccia. La terza: due esecuzioni dello stesso comando non devono '
    'sovrapporsi, ed è la questione già segnalata per l’aggiornamento dei dizionari — serve un '
    'blocco, non una raccomandazione.',
]:
    D.para(doc, ancora, t)
D.immagine(doc, ancora, os.path.join(IMG, 'bo_comandi.png'), 6.3,
           'Wireframe — Comandi ed esecuzioni: per ciascun comando, che cosa fa, l’equivalente '
           'da riga di comando e l’esito dell’ultima esecuzione.')
fatti.append('nuova sezione sui comandi e relativo wireframe')


# ═════════════════════════ 6. il raccordo con le tabelle di SIPO
ancora = par('⚠️ Resta fuori ciò che nessuna tabella può dire')._p.getnext()
D.para(doc, ancora, 'Il raccordo con le tabelle di SIPO', stile='Heading 3')
for t in [
    'La riconciliazione si costruisce per coppie: da una parte la decodifica che ANSC pubblica, '
    'dall’altra la tabella di SIPO che ospita i valori locali. La seconda non è dichiarata da '
    'nessuna fonte di ANSC e va individuata nel patrimonio esistente, dove le tabelle di '
    'codifica sono molte: ⚠️ le entità di SIPO ne mappano trecentododici con prefisso CONF_, '
    'oltre alla famiglia CFG_ del motore di regole.',
    'Il foglio di riconciliazione del Comune elenca trenta decodifiche da raccordare e per '
    'quattro di esse indica già la tabella di SIPO. La verifica sul codice — condotta sui nomi '
    'delle entità — dà questo esito, e conviene tenerlo presente perché due nomi su quattro non '
    'corrispondono a una tabella esistente.',
]:
    D.para(doc, ancora, t)
D.tabella(doc, ancora, [
    ('Decodifica ANSC', 'Tabella SIPO indicata nel foglio', 'Esito della verifica sul codice'),
    ('ANSC_61 — stato civile', 'CONF_STATO_CIVILE', '✔ esiste con questo nome.'),
    ('ANSC_52 — richieste ricevute', 'CONF_RICHIESTE_ALTRO_RICEVUTE', '✔ esiste con questo '
     'nome. ⚠️ Esiste anche CONF_RICHIESTE_RICEVUTE, che è un’altra tabella: da confermare '
     'quale delle due.'),
    ('ANSC_32 — difficoltà degli sposi', 'conf_diff_sposi',
     '⚠️ Non esiste: il nome della tabella è CONF_DIFFICOLTA_SPOSI. Il foglio usa '
     'un’abbreviazione.'),
    ('ANSC_48 — parentela', 'CONF_PARENTELA_AUTORIZZ',
     '⚠️ Non esiste: la tabella è CONF_PARENTELE_AUTORIZZ, al plurale.'),
    ('Le altre ventisei', '— (non indicata)', 'Da individuare. ⚠️ La configurazione dei campi '
     'contiene già l’indicazione: la colonna SIPO di un campo codificato dice da quale tabella '
     'proviene il valore, ed è da lì che il raccordo si ricava senza cercarlo a mano.'),
], MOD)
D.para(doc, ancora, 'Il raccordo fra le decodifiche da riconciliare e le tabelle di SIPO, '
                    'verificato sui nomi delle entità.', corsivo=True)
for t in [
    '⚠️ Una avvertenza sull’elenco: contiene anche ANSC_109, che nel corpus pubblicato da ANSC '
    'non esiste — gli identificativi si fermano a 108 e riprendono dai 170. Va chiarito se sia '
    'un refuso o una decodifica attesa e non ancora pubblicata, perché una riconciliazione '
    'verso una decodifica inesistente non è verificabile.',
    'Il metodo che ne discende è dichiarato qui perché eviterà un lavoro inutile: non si '
    'costruisce il raccordo decodifica per decodifica interrogando chi conosce SIPO, ma si '
    'parte dai campi configurati. Un campo che porta una decodifica ANSC porta anche la propria '
    'colonna di SIPO; la tabella di quella colonna è la candidata naturale, e l’unica verifica '
    'che resta da fare a mano è che i suoi valori siano quelli attesi.',
]:
    D.para(doc, ancora, t)
D.immagine(doc, ancora, os.path.join(IMG, 'bo_riconciliazione.png'), 6.3,
           'Wireframe — Riconciliazione dei valori: la decodifica ANSC da una parte, la tabella '
           'di SIPO dall’altra, e per ciascun valore locale il corrispondente nazionale con la '
           'sua validità.')
fatti.append('nuova sezione sul raccordo con le tabelle di SIPO')

t = D.trova_tabella(doc, '#', 'tema', 'questione')
D.clona_riga(t, ('OP-64', 'Raccordo fra decodifiche ANSC e tabelle di SIPO',
                 'Delle trenta decodifiche da riconciliare, quattro indicano la tabella di '
                 'SIPO e due di queste con un nome che non esiste; per le altre ventisei la '
                 'tabella va individuata. Va inoltre chiarito ANSC_109, che nel corpus '
                 'pubblicato non compare.', 'Aperto', 'Analisi / Cliente', 'Media'))
fatti.append('OP-64 aperto')


# ═════════════════════════ 7. i wireframe rifatti
for inizio, png, nuova_didascalia in (
    ('Wireframe — Vista di supervisione', 'bo_supervisione.png',
     'Wireframe — Supervisione atti: worklist delle eccezioni per categoria e per fase, con il '
     'dettaglio dell’atto — stato reale, dati conservati, timeline dell’audit e azioni '
     'contestuali.'),
    ('Wireframe — Configurazione (Casi d’uso)', 'bo_config_uc.png',
     'Wireframe — Configurazione (Casi d’uso): una riga per UC adottato, con Modello, priorità, '
     'logica di scelta richiamata dal catalogo e origine della riga.'),
    ('Wireframe — Ricerca / Tracciabilità', 'bo_ricerca.png',
     'Wireframe — Ricerca e tracciabilità: un atto si ritrova anche per identificativo '
     'nazionale e per numero comunale; la scheda mostra il percorso nelle sue fasi.'),
    ('Wireframe — Configurazione (Campi, allegati e formule per UC)', 'bo_config_campi.png',
     'Wireframe — Configurazione (Sezioni, campi, allegati e formule per UC): quattro schede '
     'per un solo UC, con la logica del campo richiamata dal catalogo e il rimando alla '
     'riconciliazione dei valori.'),
    ('Wireframe — Audit & Log', 'bo_audit.png',
     'Wireframe — Audit e log: le fasi del flusso operativo, comprese quelle che non chiamano '
     'ANSC, con esito, durata, operatore e messaggio.'),
):
    D.sostituisci_immagine(doc, inizio, os.path.join(IMG, png))
    D.testo_di(par(inizio), nuova_didascalia)
fatti.append('cinque wireframe rifatti e didascalie aggiornate')

doc.save(DST)

# ───────────────────────────────── controlli
import zipfile   # noqa: E402
z = zipfile.ZipFile(DST)
ncom = z.read('word/comments.xml').decode().count('<w:comment ')
assert ncom == 23, f'commenti persi: {ncom}'
d2 = docx.Document(DST)
for p in d2.paragraphs:
    if p.style.name.startswith('Heading') and len(p.text) > 95:
        raise SystemExit('heading anomalo: ' + p.text[:70])
print('\n'.join(' · ' + f for f in fatti))
print('commenti:', ncom, '· capitoli/tabelle/immagini:', D.riepilogo(DST))
print('scritto:', os.path.relpath(DST, BASE))
