# -*- coding: utf-8 -*-
"""WORKFLOW_Interfacce-SIPO-ANSC v0.2 → v0.3 (21/10/2026).

Indicazione dell'utente: le pagine nuove gestite a linguette non vanno, «si perde la
transizione». La v0.3 le propone come SCHERMATE SEPARATE di SIPO:
  · le quattro ipotesi di schermata sono ridisegnate senza linguette; ciascuna si apre con un
    banner verde di arrivo che dice da quale schermata si arriva e che cosa è cambiato
    (schermate_sipo_wf.disegna_v3);
  · nel capitolo comune, la voce sulle linguette diventa quella sulle schermate separate, e
    una sezione nuova «Il percorso fra le schermate» mostra la sequenza con il comando, il
    servizio e lo stato su ogni passaggio, con la tabella dei passaggi;
  · una schermata nuova, «Esito della chiusura», dopo CHIUDI nella Firma dell'USC.
Diagrammi a stati, testi e punti aperti della v0.2 restano com'erano.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/wf_v0_3.py
"""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D            # noqa: E402
import schermate_sipo_wf as SCH    # noqa: E402

QUI = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(QUI)
SORGENTE = os.path.join(DOCS, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.2.docx')
DEST = os.path.join(DOCS, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.3.docx')
IMG = os.path.join(QUI, 'img')
LARGO = 6.5

SOSTITUZIONI = (
    ("Ipotesi di schermata: Validazione.", 'wf3_sch_validazione.png'),
    ("Ipotesi di schermata: Firme dei dichiaranti.", 'wf3_sch_firme.png'),
    ("Ipotesi di schermata: Conformità.", 'wf3_sch_conformita.png'),
    ("Ipotesi di schermata: Firma dell’USC.", 'wf3_sch_firma_usc.png'),
)

VOCE_LINGUETTE = "Le linguette delle quattro pagine."
VOCE_SCHERMATE = (
    "Una schermata per pagina, con il banner di arrivo. Le pagine nuove sono schermate "
    "separate di SIPO, ciascuna con il proprio titolo, come le maschere esistenti; si passa "
    "dall’una all’altra con i pulsanti. Ogni schermata si apre con un banner verde che dice da "
    "quale schermata si arriva e che cosa è cambiato — il numero comunale staccato, "
    "l’identificativo nazionale assegnato, lo stato ANSC raggiunto: è il punto in cui "
    "l’operatore vede la transizione. Tornando INDIETRO il banner non compare, perché i dati "
    "non cambiano.")

PERCORSO_TESTO = (
    "La figura mostra le schermate in sequenza, dall’anteprima di SIPO all’esito della "
    "chiusura. Su ogni passaggio ci sono il pulsante che lo provoca, il servizio chiamato e lo "
    "stato che ne risulta; il tratteggio indica i passaggi che saltano una schermata, quando "
    "l’atto non ha firmatari o non ha allegati da attestare. Le schermate sono quelle dei "
    "capitoli seguenti, ridotte.")

PASSAGGI = [
    ("Schermata di partenza", "Comando [condizione]", "Schermata di arrivo",
     "Che cosa dice il banner di arrivo"),
    ("Anteprima SIPO", "VALIDA", "Validazione",
     "Numero comunale staccato, modello evento costruito e salvato."),
    ("Validazione", "AVANTI [ci sono firmatari]", "Firme dei dichiaranti",
     "Validazione riuscita (R009): identificativo nazionale, stato ANSC 2 CONFERMATO."),
    ("Validazione", "AVANTI [nessun firmatario]", "Conformità",
     "Lo stesso esito della validazione."),
    ("Firme dei dichiaranti", "AVANTI", "Conformità",
     "Firme acquisite: stato ANSC 4 FIRMATO DAL DICHIARANTE, dati vincolati."),
    ("Validazione o Firme dei dichiaranti", "AVANTI [nessun allegato da attestare]",
     "Firma dell’USC", "L’esito della schermata di partenza e che la conformità non serve."),
    ("Conformità", "AVANTI o CONTINUA", "Firma dell’USC",
     "Attestazione firmata (R007): stato ANSC 11 FIRMATO PER CONFORMITÀ; con CONTINUA, nessun "
     "allegato attestato."),
    ("Firma dell’USC", "CHIUDI", "Esito della chiusura",
     "Atto firmato e chiuso; la schermata riporta l’esito di firma, data di validità, "
     "salvataggio in SIPO e aggiornamento anagrafico."),
    ("qualunque", "INDIETRO", "schermata precedente", "Nessun banner: i dati non cambiano."),
    ("Firme dei dichiaranti, Conformità", "TORNA ALLA COMPILAZIONE [regola del ritorno dopo le "
     "firme]", "pagine dei dati", "Dati riaperti: stato 0, firme decadute, nuova validazione da "
     "fare."),
]

ESITO_TESTO = (
    "Dopo CHIUDI si apre una schermata di esito, che chiude il percorso: per ciascuna delle "
    "operazioni della chiusura — firma dell’atto, data di validità, salvataggio in SIPO, "
    "aggiornamento anagrafico — il servizio e l’esito. Corrisponde agli stati finali del "
    "diagramma della pagina: se il salvataggio in SIPO non riesce, la stessa schermata lo "
    "mostra in rosso con RIPROVA, che ripete il salvataggio e mai la firma. Da qui si torna "
    "alla ricerca degli atti, alle stampe o a un atto nuovo.")


def sostituisci_figura(doc, didascalia, png):
    """Riscrive il PNG e ne corregge le proporzioni, che con il blob non cambiano da sole."""
    ps = doc.paragraphs
    for i, p in enumerate(ps):
        if p.text.strip() == didascalia:
            el = ps[i - 1]._p
            blip = el.findall('.//' + qn('a:blip'))
            assert len(blip) == 1, didascalia
            doc.part.related_parts[blip[0].get(qn('r:embed'))]._blob = open(png, 'rb').read()
            from PIL import Image
            w, h = Image.open(png).size
            for tag in ('wp:extent', 'a:ext'):
                for e in el.iter(qn(tag)):
                    if e.get('cx'):
                        e.set('cy', str(int(int(e.get('cx')) * h / w)))
            return
    raise SystemExit('didascalia non trovata: ' + didascalia)


def heading_dopo(doc, h1, livello, titolo):
    dentro = False
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            if dentro:
                break
            dentro = p.text.strip().startswith(h1)
        elif dentro and p.style.name == f'Heading {livello}' and p.text.strip() == titolo:
            return p._p
    raise SystemExit(f'{titolo} non trovato in {h1}')


def tabella_sinistra(doc, prima, righe, modello, larghezze):
    t = D.tabella(doc, prima, righe, modello, larghezze)
    for r in t.rows:
        for c in r.cells:
            for par in c.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    D.para(doc, prima, '')


def main():
    figure = SCH.disegna_v3()
    print('figure:', len(figure))
    doc = docx.Document(SORGENTE)
    fatti = []

    for didascalia, png in SOSTITUZIONI:
        sostituisci_figura(doc, didascalia, os.path.join(IMG, png))
    fatti.append('quattro schermate ridisegnate senza linguette, con il banner di arrivo')

    voce = [p for p in doc.paragraphs if p.text.startswith(VOCE_LINGUETTE)]
    assert len(voce) == 1
    D.testo_di(voce[0], VOCE_SCHERMATE)
    fatti.append('voce sulle linguette → schermate separate')

    modello = D.trova_tabella(doc, 'Da', 'Che cosa lo provoca')
    prima = heading_dopo(doc, "Le pagine nuove", 2, "Il comportamento comune")
    D.para(doc, prima, "Il percorso fra le schermate", stile='Heading 2')
    D.para(doc, prima, PERCORSO_TESTO)
    D.immagine(doc, prima, os.path.join(IMG, 'wf3_percorso.png'), LARGO,
               "Il percorso fra le schermate.")
    tabella_sinistra(doc, prima, PASSAGGI, modello, [1.45, 1.6, 1.3, 2.15])
    fatti.append('sezione «Il percorso fra le schermate»')

    prima = heading_dopo(doc, "Anteprima · Firma dell’USC", 2, "Il diagramma a stati della pagina")
    D.para(doc, prima, ESITO_TESTO)
    D.immagine(doc, prima, os.path.join(IMG, 'wf3_sch_esito.png'), LARGO,
               "Ipotesi di schermata: Esito della chiusura.")
    fatti.append('schermata «Esito della chiusura»')

    fatti += D.frontespizio(doc, '0.3', 'WORKFLOW_Interfacce-SIPO-ANSC_v0.3', '21/10/2026')
    D.storia(doc, '21/10/2026', '0.3', 'Pagine nuove; Firma dell’USC',
             "Pagine nuove come schermate separate di SIPO invece che a linguette: ogni "
             "schermata si apre con il banner di arrivo, che mostra la transizione; sezione «Il "
             "percorso fra le schermate» con la sequenza e la tabella dei passaggi; schermata "
             "di esito della chiusura dopo CHIUDI.")
    doc.save(DEST)
    for f in fatti:
        print(' ·', f)
    print('scritto:', DEST, D.riepilogo(DEST))


if __name__ == '__main__':
    main()
