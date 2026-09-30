# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.25 -> v3.26 (24/09/2026).

Allinea il modello alla sorgente `Integrazione_SIPO_ANSC.xlsx` (16:49) e chiude le
divergenze fra schede e DDL, secondo le decisioni prese con l'utente:

  · il catalogo delle logiche prende i nomi della sorgente — TIPO_LOGICHE_DI_SCELTA e
    LOGICHE_DI_SCELTA — e con essi cade la somiglianza con i dizionari ANSC (OP-56);
  · il dominio 3, «logica dei campi», entra nel catalogo: ANSC_CFG_CAMPO.ID_BUSINESS_LOGIC
    diventa un rimando e non più uno script scritto nella riga;
  · ANSC_CFG_UC: forma mista — ID_TIPO_EVENTO e ID_TIPO_DOCUMENTO, più DESCRIZIONE e STATO;
  · ANSC_STATO_ATTO: valgono i nomi della scheda e CHIAVE_ANTI_DUPLICATO è eliminata (OP-63);
  · ANSC_CFG_SEZIONE resta: dal DDL dei campi escono ID_UC, SEZIONE_FE_ANSC e la condizione;
  · nuova RICONCILIAZ_DIZIONARI, la traduzione dei valori fra SIPO e ANSC (OP-59).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_26_logiche_e_riconciliazione.py
"""
import os
import shutil
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.25.docx')
DST = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.26.docx')
IMG = os.path.join(BASE, 'strumenti', 'img')

if os.path.exists(DST):
    os.remove(DST)
shutil.copy(SRC, DST)
doc = docx.Document(DST)
fatti = []


# ─────────────────────────────────────────────────────────── utilità
def par(inizio):
    t = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(t) != 1:
        raise SystemExit(f'attesa 1 occorrenza di «{inizio[:60]}», trovate {len(t)}')
    return t[0]


def riga_con(t, prima_cella):
    for r in t.rows:
        if r.cells[0].text.strip() == prima_cella:
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


def inserisci_riga(t, dopo, valori):
    nuova = D.clona_riga(t, valori)
    riga_con(t, dopo)._tr.addnext(nuova._tr)
    return nuova


def mono(testo):
    t = [p for p in doc.paragraphs if p.text.rstrip() == testo.rstrip()]
    if len(t) != 1:
        raise SystemExit(f'riga DDL non univoca: «{testo[:70]}» ({len(t)})')
    return t[0]


def elimina(elementi, cosa):
    for e in elementi:
        if D.ha_commenti(e):
            raise SystemExit(f'{cosa}: commento ancorato qui, non si cancella')
    for e in elementi:
        e.getparent().remove(e)
    fatti.append('eliminato: ' + cosa)


def riga_ddl_via(testo):
    elimina([mono(testo)._p], 'riga DDL ' + testo.strip()[:40])


def blocco(create_line):
    """I paragrafi di un CREATE TABLE: serve quando la stessa riga ricorre in più tabelle."""
    p = mono(create_line)
    out, e = [p._p], p._p.getnext()
    while e is not None:
        out.append(e)
        if e.tag.endswith('}p') and Paragraph(e, doc).text.strip().startswith(') TABLESPACE'):
            return out
        e = e.getnext()
    raise SystemExit('blocco non chiuso: ' + create_line)


def riga_del_blocco_via(create_line, testo):
    dentro = [e for e in blocco(create_line)
              if e.tag.endswith('}p') and Paragraph(e, doc).text.rstrip() == testo.rstrip()]
    if len(dentro) != 1:
        raise SystemExit(f'in {create_line[:40]}: «{testo.strip()[:40]}» trovata {len(dentro)}')
    elimina(dentro, 'riga ' + testo.strip()[:38] + ' da ' + create_line.split('.')[-1][:-2])


def ddl_dopo(ancora, righe):
    D.ddl(doc, mono(ancora)._p.getnext(), righe)


MOD = D.trova_tabella(doc, 'colonna', 'tipo', 'note')


# ═════════════════════════ 1. scheda e storia
for t in doc.tables:
    if t.rows[0].cells[0].text.strip().lower().startswith(('area organizzativa', 'progetto')):
        for nome, val in (('Data consegna', '24/09/2026'), ('Versione', '3.26')):
            try:
                D.riscrivi_cella(riga_con(t, nome).cells[1], val)
            except SystemExit:
                pass

D.storia(doc, '24/09/2026', '3.26',
         'Adeguamenti database · Gestione dei dizionari · Flusso operativo · Open Point · '
         'Appendice A',
         'Il modello è allineato alla sorgente del Comune. Il catalogo delle logiche prende i '
         'nomi della sorgente — TIPO_LOGICHE_DI_SCELTA e LOGICHE_DI_SCELTA — e con essi cade '
         'la somiglianza con i dizionari ANSC che rendeva le due coppie confondibili. Il '
         'catalogo guadagna un terzo dominio, la logica dei campi: la trasformazione di un '
         'valore non è più uno script scritto nella riga del campo ma un codice che la '
         'richiama. ANSC_CFG_CAMPO perde le tre colonne che duplicavano ANSC_CFG_SEZIONE; '
         'ANSC_STATO_ATTO adotta i nomi della scheda e perde CHIAVE_ANTI_DUPLICATO, con il '
         'controllo anti-duplicato registrato come punto aperto. Introdotta '
         'RICONCILIAZ_DIZIONARI, che dichiara come un valore di SIPO si traduce nel '
         'corrispondente valore di ANSC.')


# ═════════════════════════ 2. i nomi del catalogo delle logiche
n = D.sostituisci(doc, 'VALORI_DOMINIO', 'LOGICHE_DI_SCELTA')
fatti.append(f'VALORI_DOMINIO → LOGICHE_DI_SCELTA in {n} punti')
n = D.sostituisci(doc, 'TIPO_DOMINIO', 'TIPO_LOGICHE_DI_SCELTA')
fatti.append(f'TIPO_DOMINIO → TIPO_LOGICHE_DI_SCELTA in {n} punti')

# la scheda del catalogo porta già i nomi della sorgente: si allinea il solo DDL
D.sostituisci(doc, 'DESCRIZIONE_DI_SCELTA', 'DESC_LOGICA_DI_SCELTA', attese=2,
              etichetta='DDL del catalogo', fatti=fatti)

# il terzo dominio
t = D.tabella_colonne(doc, 'DESCRIZIONE_DOMINIO')
D.riscrivi_cella(riga_con(t, 'DESCRIZIONE_DOMINIO').cells[2],
                 'A che cosa serve il dominio. Sono tre: 1 = logica di scelta dell’UC, '
                 '2 = logica di attivazione di una sezione, 3 = logica dei campi, cioè la '
                 'trasformazione da applicare al valore letto da SIPO.')
D.para(doc, par('TIPO_LOGICHE_DI_SCELTA — i domini di logica')._p,
       '⚠️ Il terzo dominio è la novità di questa versione: la trasformazione di un valore — '
       'estrarre la data da una stringa, ricavarne l’ora, tradurre un codice — smette di '
       'essere un testo scritto dentro la riga del campo e diventa una voce del catalogo che '
       'i campi richiamano. Le prime tre sono DATA_COMPLETA, ORA e MINUTI, e la stessa '
       'trasformazione vale per tutti i campi che la citano invece di essere riscritta su '
       'ciascuno.')
fatti.append('catalogo: colonne allineate e terzo dominio dichiarato')


# ═════════════════════════ 3. ANSC_CFG_UC — forma mista
t = D.tabella_colonne(doc, 'ID_UC_CFG')
D.riscrivi_cella(riga_con(t, 'ID_DOMINIO / COD_LOGICA').cells[0],
                 'ID_DOMINIO / COD_LOGICA_DI_SCELTA')
inserisci_riga(t, 'COD_UC_ANSC', (
    'DESCRIZIONE', 'VARCHAR2(200)',
    'Descrizione del Modello di atto come la nomina il Comune: «MORTE IN ABITAZIONE A ROMA». '
    'Serve a chi configura, che ragiona per Modelli e non per codici.'))
inserisci_riga(t, 'MASCHERA_UI', (
    'STATO', 'VARCHAR2(1)',
    '[DA VERIFICARE: i valori ammessi. Un solo carattere; la sorgente dichiara la colonna ma '
    'non la valorizza. '
    '⚠️ Da non confondere con lo stato dell’atto: qui si tratta della riga di configurazione.]'))
D.sostituisci(doc, '  COD_TIPO_EVENTO      VARCHAR2(20 CHAR)  NOT NULL,',
              '  ID_TIPO_EVENTO       VARCHAR2(10 CHAR)  NOT NULL,', attese=1,
              etichetta='CFG_UC: tipo evento', fatti=fatti)
D.sostituisci(doc, '  TIPO_OPERAZIONE  \tVARCHAR2(20 CHAR)  NOT NULL,',
              '  ID_TIPO_DOCUMENTO    VARCHAR2(10 CHAR)  NOT NULL,', attese=1,
              etichetta='CFG_UC: tipo documento', fatti=fatti)
# ⚠️ due righe scritte a mano: tabulazioni al posto degli spazi e, su STATO, la virgola
# mancante — il DDL così com'è non compila.
D.sostituisci(doc, '  DESCRIZIONE\t\tVARCHAR2 (200 CHAR),',
              '  DESCRIZIONE          VARCHAR2(200 CHAR),', attese=1,
              etichetta='CFG_UC: DESCRIZIONE normalizzata', fatti=fatti)
D.sostituisci(doc, '  STATO\t\t\tVARCHAR2 (1 CHAR)',
              '  STATO                VARCHAR2(1 CHAR),', attese=1,
              etichetta='CFG_UC: STATO, virgola mancante', fatti=fatti)


# ═════════════════════════ 4. ANSC_STATO_ATTO — valgono i nomi della scheda
D.sostituisci(doc, '  COD_TIPO_EVENTO         VARCHAR2(20 CHAR)  NOT NULL,',
              '  ID_TIPO_EVENTO          VARCHAR2(10 CHAR)  NOT NULL,', attese=1,
              etichetta='STATO_ATTO: tipo evento', fatti=fatti)
ddl_dopo('  ID_TIPO_EVENTO          VARCHAR2(10 CHAR)  NOT NULL,',
         ['  ID_TIPO_CONTENUTO       VARCHAR2(10 CHAR),'])
D.sostituisci(doc, '  COD_UC_ANSC             VARCHAR2(20 CHAR),',
              '  ID_UC_ANSC              VARCHAR2(20 CHAR),', attese=1,
              etichetta='STATO_ATTO: UC determinato', fatti=fatti)
for riga in ('  CHIAVE_ANTI_DUPLICATO   VARCHAR2(64 CHAR)  NOT NULL,',
             '  CONSTRAINT UK_ANSC_STATO_ATTO_ANTIDUP UNIQUE (CHIAVE_ANTI_DUPLICATO),'):
    riga_ddl_via(riga)
intestazione = mono('COMMENT ON COLUMN ANSC_USR.ANSC_STATO_ATTO.CHIAVE_ANTI_DUPLICATO IS')._p
elimina([intestazione, intestazione.getnext()], 'COMMENT di CHIAVE_ANTI_DUPLICATO')
D.sostituisci(doc, 'l\'\'UC ANSC e\'\' in COD_UC_ANSC', 'l\'\'UC ANSC e\'\' in ID_UC_ANSC')
D.sostituisci(doc, 'ANSC_STATO_ATTO.COD_UC_ANSC', 'ANSC_STATO_ATTO.ID_UC_ANSC')
D.sostituisci(doc, 'con COD_UC_ANSC, COD_LOGICA', 'con ID_UC_ANSC, COD_LOGICA')
D.sostituisci(doc, 'UC determinato (COD_UC_ANSC)', 'UC determinato (ID_UC_ANSC)')

# il controllo che cade insieme alla colonna
D.para(doc, D.h(doc, 2, 'La traccia di ciò che accade')._p,
       '⚠️ Una avvertenza che discende dal modello dati. Lo store di stato non porta più una '
       'chiave anti-duplicato: fino alla v3.25 un vincolo di unicità impediva che lo stesso '
       'atto fosse depositato due volte per la stessa operazione, ed era un presidio locale '
       'perché ANSC non offre alcuna idempotenza. Venuto meno il vincolo, il controllo va '
       'realizzato altrove — nella lettura dello stato prima del deposito, o in un indice '
       'equivalente — e fino ad allora un secondo deposito dello stesso atto consuma un '
       'secondo identificativo nazionale senza che nulla se ne accorga (OP-63).')
fatti.append('STATO_ATTO allineato alla scheda; avvertenza sull’anti-duplicato')


# ═════════════════════════ 5. ANSC_CFG_CAMPO — la sezione resta, la logica si cita
CAMPI = 'CREATE TABLE ANSC_USR.ANSC_CFG_CAMPO ('
# ⚠️ le righe qui sotto sono state riscritte a mano con tabulazioni: si citano come stanno.
for riga in ('  ID_UC                \t\tNUMBER             NOT NULL,',
             '  SEZIONE_FE_ANSC      \t\tVARCHAR2(200 CHAR) NOT NULL,',
             '  DESC_COND_OBBL_EVENTO\t\tVARCHAR2(1000 CHAR),'):
    riga_del_blocco_via(CAMPI, riga)

# la logica del campo diventa un rimando al catalogo: serve il dominio accanto al codice
D.sostituisci(doc, '  ID_BUSINESS_LOGIC\t\t\tVARCHAR2 (10),',
              '  ID_DOMINIO           NUMBER DEFAULT 3,\n'
              '  ID_BUSINESS_LOGIC    VARCHAR2(30 CHAR),', attese=1,
              etichetta='CFG_CAMPO: dominio della logica', fatti=fatti)
# ⚠️ difetto trovato nel DDL scritto a mano: tipo senza lunghezza, non compila
D.sostituisci(doc, '  FLG_OBBLIGATORIO\t\t\tCHAR () NOT NULL,',
              '  FLG_OBBLIGATORIO     CHAR(1 CHAR)       NOT NULL,', attese=1,
              etichetta='CFG_CAMPO: CHAR senza lunghezza', fatti=fatti)
D.sostituisci(doc, '  CONSTRAINT UK_ANSC_CFG_CAMPO UNIQUE (ID_SEZIONE, OGGETTO_ANSC, CAMPO_ANSC),',
              '  CONSTRAINT UK_ANSC_CFG_CAMPO UNIQUE (ID_SEZIONE, OGGETTO_ANSC, CAMPO_ANSC),\n'
              '  CONSTRAINT FK_ANSC_CFG_CAMPO_LOGICA\n'
              '    FOREIGN KEY (ID_DOMINIO, ID_BUSINESS_LOGIC)\n'
              '    REFERENCES ANSC_USR.LOGICHE_DI_SCELTA (ID_DOMINIO, COD_LOGICA_DI_SCELTA),',
              attese=1, etichetta='CFG_CAMPO: chiave esterna al catalogo', fatti=fatti)

t = D.tabella_colonne(doc, 'ID_CAMPO')
r = riga_con(t, 'ID_BUSINESS_LOGIC')
D.riscrivi_cella(r.cells[0], 'ID_DOMINIO / ID_BUSINESS_LOGIC')
D.riscrivi_cella(r.cells[1], 'NUMBER / VARCHAR2(30) (FK)')
D.riscrivi_cella(r.cells[2], 'La trasformazione da applicare al valore letto da SIPO, presa '
                             'dal catalogo delle logiche (dominio 3). ⚠️ Non è uno script '
                             'scritto nella riga: è un codice, e la stessa trasformazione vale '
                             'per tutti i campi che la citano. La descrizione in lingua '
                             'corrente resta nella colonna precedente.')

# le due chiavi esterne che cercavano la colonna con il nome vecchio
n = D.sostituisci(doc, 'REFERENCES ANSC_USR.LOGICHE_DI_SCELTA (ID_DOMINIO, COD_LOGICA),',
                  'REFERENCES ANSC_USR.LOGICHE_DI_SCELTA (ID_DOMINIO, COD_LOGICA_DI_SCELTA),')
fatti.append(f'chiavi esterne verso il catalogo corrette: {n}')


# ═════════════════════════ 6. RICONCILIAZ_DIZIONARI
ancora = D.h(doc, 2, 'Fruizione da parte di SIPO')._p
D.para(doc, ancora, 'La riconciliazione dei valori fra SIPO e ANSC', stile='Heading 2')
for t in [
    'I dizionari dicono quali valori ANSC ammette; non dicono quale valore di SIPO '
    'corrisponda a quale valore di ANSC. È una distinzione che si vede bene su un esempio: lo '
    'stato civile di una persona in SIPO è un codice delle tabelle di configurazione locali, '
    'in ANSC è un valore della decodifica ANSC_61, e i due non coincidono. Il payload deve '
    'portare il secondo.',
    'Fino alla v3.25 il modello dichiarava quale decodifica governa un campo e si fermava lì: '
    'la traduzione restava nella descrizione della logica del campo, cioè in una frase in '
    'lingua corrente rivolta a chi sviluppa. ⚠️ Una traduzione scritta in prosa non è '
    'verificabile e non è riusabile: la stessa corrispondenza va riscritta su ogni campo che '
    'la richiede, e nessuno può sapere, guardando la configurazione, se due campi traducano '
    'lo stesso valore in due modi diversi.',
    'La sorgente del Comune la mette in tabella, ed è la scelta che questa versione recepisce. '
    'La riga dichiara, per una decodifica, che cosa diventa un dato valore di SIPO, con la '
    'validità temporale: una corrispondenza che cambia non si riscrive, si chiude e se ne apre '
    'una nuova, perché gli atti già formati vanno riletti con la corrispondenza in vigore '
    'allora.',
]:
    D.para(doc, ancora, t)
D.para(doc, ancora, 'RICONCILIAZ_DIZIONARI — la traduzione dei valori.')
D.tabella(doc, ancora, [
    ('Colonna', 'Tipo', 'Note'),
    ('ID_RICONCILIAZIONE', 'NUMBER (PK)', 'Chiave tecnica.'),
    ('ID_VERSIONE', 'NUMBER (FK)', 'Baseline di appartenenza (ANSC_CFG_VERSIONE). ⚠️ Aggiunta '
     'rispetto alla sorgente: la traduzione è una decisione del Comune, e come le altre si '
     'attiva e si storicizza in blocco.'),
    ('DECODIFICA', 'VARCHAR2(100)', 'La decodifica ANSC interessata. ⚠️ Nella sorgente ha due '
     'forme: il solo identificativo (ANSC_01) oppure la colonna di SIPO che lo precede '
     '(conf_diff_sposi.ANSC_32), quando la stessa decodifica si traduce diversamente secondo '
     'la colonna di partenza. Le due forme convivono e vanno normalizzate prima del carico.'),
    ('VALORE_SIPO / DESCRIZIONE_SIPO', 'VARCHAR2(30) / VARCHAR2(400)', 'Il valore come SIPO lo '
     'registra, con la sua descrizione: è la descrizione che consente a chi compila di '
     'riconoscere la riga senza risalire al codice.'),
    ('VALORE_ANSC / DESCRIZIONE_ANSC', 'VARCHAR2(30) / VARCHAR2(400)', 'Il valore da '
     'trasmettere e la descrizione pubblicata da ANSC. ⚠️ Il valore deve esistere nella '
     'decodifica replicata: è la verifica che il pre-filtro esegue leggendo per chiave.'),
    ('CONDIZIONE', 'VARCHAR2(1000)', 'Quando la corrispondenza si applica, per i casi in cui '
     'non è unica. Nella sorgente è testo libero; ha la stessa grammatica delle logiche e '
     'potrà citarne il catalogo.'),
    ('DATA_INIZIO_VALIDITA / DATA_FINE_VALIDITA', 'DATE', 'Validità della corrispondenza. Una '
     'traduzione cessata resta, perché serve a rileggere gli atti formati quando era in '
     'vigore.'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici di audit previsti dallo standard di nomenclatura [R5].'),
], MOD)
D.para(doc, ancora, 'Vincolo di unicità su (ID_VERSIONE, DECODIFICA, VALORE_SIPO, '
                    'DATA_INIZIO_VALIDITA): una sola traduzione per valore, dentro una '
                    'baseline e per un periodo.')
D.para(doc, ancora, '⚠️ Resta fuori ciò che nessuna tabella può dire: quali valori di SIPO non '
                    'abbiano alcun corrispondente in ANSC. Sono i casi in cui la '
                    'configurazione va completata a mano e, finché non lo è, il pre-filtro li '
                    'segnala come valore non traducibile invece di lasciarli passare.')
fatti.append('nuova sezione e scheda di RICONCILIAZ_DIZIONARI')

# il DDL, in coda alle tabelle di configurazione
D.ddl(doc, mono('CREATE TABLE ANSC_USR.DOMINIO_DECODIFICA (')._p, [
    'CREATE TABLE ANSC_USR.RICONCILIAZ_DIZIONARI (',
    '  ID_RICONCILIAZIONE   NUMBER GENERATED ALWAYS AS IDENTITY,',
    '  ID_VERSIONE          NUMBER             NOT NULL,',
    '  DECODIFICA           VARCHAR2(100 CHAR) NOT NULL,',
    '  VALORE_SIPO          VARCHAR2(30 CHAR)  NOT NULL,',
    '  DESCRIZIONE_SIPO     VARCHAR2(400 CHAR),',
    '  VALORE_ANSC          VARCHAR2(30 CHAR),',
    '  DESCRIZIONE_ANSC     VARCHAR2(400 CHAR),',
    '  CONDIZIONE           VARCHAR2(1000 CHAR),',
    '  DATA_INIZIO_VALIDITA DATE DEFAULT SYSDATE            NOT NULL,',
    '  DATA_FINE_VALIDITA   DATE DEFAULT DATE \'9999-12-31\'  NOT NULL,',
    '  DATA_INS             TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,',
    '  DATA_UPD             TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,',
    '  UTENTE_INS           VARCHAR2(40 CHAR),',
    '  UTENTE_UPD           VARCHAR2(40 CHAR),',
    '  CONSTRAINT PK_RICONCILIAZ_DIZIONARI PRIMARY KEY (ID_RICONCILIAZIONE),',
    '  CONSTRAINT FK_RICONCILIAZ_DIZ_VERSIONE',
    '    FOREIGN KEY (ID_VERSIONE) REFERENCES ANSC_USR.ANSC_CFG_VERSIONE (ID_VERSIONE),',
    '  CONSTRAINT UK_RICONCILIAZ_DIZIONARI',
    '    UNIQUE (ID_VERSIONE, DECODIFICA, VALORE_SIPO, DATA_INIZIO_VALIDITA)',
    ') TABLESPACE ANSC_USR;',
    'CREATE INDEX ANSC_USR.IX_RICONCILIAZ_DIZ_SIPO',
    '  ON ANSC_USR.RICONCILIAZ_DIZIONARI (DECODIFICA, VALORE_SIPO) TABLESPACE ANSC_USR;',
    'COMMENT ON TABLE ANSC_USR.RICONCILIAZ_DIZIONARI IS',
    "  'Come un valore di SIPO si traduce nel corrispondente valore di ANSC. I dizionari",
    "   dicono quali valori ANSC ammette, non quale valore locale vi corrisponde.';",
    'COMMENT ON COLUMN ANSC_USR.RICONCILIAZ_DIZIONARI.DECODIFICA IS',
    "  'ATTENZIONE: nella sorgente ha due forme, il solo identificativo (ANSC_01) oppure la",
    "   colonna di SIPO che lo precede (conf_diff_sposi.ANSC_32). Da normalizzare al carico.';",
    ' ',
])
fatti.append('Appendice A: DDL di RICONCILIAZ_DIZIONARI')


# ═════════════════════════ 7. open point
t = D.trova_tabella(doc, '#', 'tema', 'questione')
r = riga_con(t, 'OP-56')
D.riscrivi_cella(r.cells[2], 'Chiuso nella v3.26: il catalogo delle logiche prende i nomi '
                             'della sorgente del Comune — TIPO_LOGICHE_DI_SCELTA e '
                             'LOGICHE_DI_SCELTA — che non si confondono più con '
                             'DOMINIO_DECODIFICA e VALORE_DOMINIO dei dizionari ANSC.')
D.riscrivi_cella(r.cells[3], 'Chiuso')
D.riscrivi_cella(r.cells[5], '—')
r = riga_con(t, 'OP-59')
D.riscrivi_cella(r.cells[2], 'Chiuso nella v3.26 con la tabella RICONCILIAZ_DIZIONARI, che '
                             'dichiara la traduzione di un valore di SIPO nel corrispondente '
                             'valore di ANSC, con validità temporale. ⚠️ Resta da compilarla: '
                             'la sorgente elenca le decodifiche interessate ma non le '
                             'corrispondenze, che sono lavoro dei funzionari.')
D.riscrivi_cella(r.cells[3], 'Chiuso')
D.riscrivi_cella(r.cells[5], '—')
D.clona_riga(t, ('OP-63', 'Controllo anti-duplicato del deposito',
                 'La colonna CHIAVE_ANTI_DUPLICATO è stata eliminata dallo store di stato e '
                 'con essa il vincolo di unicità che impediva di depositare due volte lo '
                 'stesso atto per la stessa operazione. ANSC non offre idempotenza: il '
                 'controllo va ricostruito altrove, e va deciso dove. Finché non c’è, un '
                 'secondo deposito consuma un secondo identificativo nazionale.',
                 'Aperto', 'Analisi', 'Alta'))
fatti.append('OP-56 e OP-59 chiusi, OP-63 aperto')


# ═════════════════════════ 8. figure
for didascalia, png in (('Schema ANSC_USR.', 'erd_ansc_usr.png'),
                        ('Dalla maschera SIPO al payload ANSC.', 'catena_configurazione.png'),
                        ('Il flusso operativo nelle sue quattro fasi.', 'flusso_operativo.png')):
    D.sostituisci_immagine(doc, didascalia, os.path.join(IMG, png))
fatti.append('tre figure rigenerate')

doc.save(DST)

# ───────────────────────────────── controlli
import zipfile   # noqa: E402
z = zipfile.ZipFile(DST)
xml = z.read('word/document.xml').decode()
ncom = z.read('word/comments.xml').decode().count('<w:comment ')
assert ncom == 23, f'commenti persi: {ncom}'
for vietato in ('VALORI_DOMINIO', 'CHIAVE_ANTI_DUPLICATO'):
    resti = xml.count(vietato)
    print(f'  residui di {vietato}: {resti}')
print('\n'.join(' · ' + f for f in fatti))
print('commenti:', ncom, '· capitoli/tabelle/immagini:', D.riepilogo(DST))
print('scritto:', os.path.relpath(DST, BASE))
