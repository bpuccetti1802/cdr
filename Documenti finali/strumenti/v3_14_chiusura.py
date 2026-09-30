# -*- coding: utf-8 -*-
"""v3.14 — la chiusura: l'ERD aggiornato, il capitolo di scopo, la storia del documento."""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.14.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img', 'erd_ansc_usr.png')

MODIFICHE = [
    ('La configurazione è un dato del Comune',
     'Le versioni precedenti descrivevano l’importazione dal mapping ufficiale come il '
     'meccanismo principale di popolamento della configurazione. Non è più così: la '
     'configurazione la scrivono i funzionari, e l’importazione prepara le righe in stato '
     '«proposto». Ogni riga dichiara la propria origine, e la reimportazione riscrive le sole '
     'righe che nessuno ha ancora esaminato.'),
    ('L’UC sostituisce il Modello come oggetto della configurazione',
     'ANSC_CFG_OPERAZIONE, una riga per Modello, diventa ANSC_CFG_UC, una riga per caso d’uso '
     'adottato. Il Modello, il tipo atto e la maschera restano come attributi. Sparisce il '
     'doppio luogo in cui l’UC poteva essere dichiarato — la colonna della riga di operazione '
     'e le regole — che le versioni precedenti lasciavano senza una precedenza scritta.'),
    ('La determinazione dell’UC è una condizione sui dati di SIPO',
     'La condizione non è più una terna campo, operatore, valore ma un’interrogazione in sola '
     'lettura sulle tabelle di SIPO, associata all’UC configurato. Il fondamento è che SIPO '
     'registra l’atto sulle proprie tabelle prima di depositarlo in ANSC: il dato su cui '
     'decidere è già scritto, ed è quello vero. Il prezzo — la rinuncia alla verifica '
     'automatica di copertura e mutua esclusione — è dichiarato in PC-9, insieme a ciò che '
     'lo sostituisce: la priorità esplicita fra gli UC e la simulazione su atti reali.'),
    ('Restano dichiarative le regole di controllo e di generazione',
     'La revisione riguarda la sola specie di determinazione. Le altre due continuano a essere '
     'configurate dai funzionari nella forma prevista dalla v3.0.'),
    ('Gli allegati diventano una funzione da costruire, non solo da configurare',
     'La verifica sul codice conferma che l’area di stato civile di SIPO non gestisce alcun '
     'documento allegato agli atti: mancano il file, il tipo e lo stato. Servono maschere di '
     'caricamento in SIPO, un luogo dove conservare i documenti fino al deposito e il '
     'raccordo con R001 e la scansione antivirus.'),
]


def applica(doc):
    fatti = []

    D.sostituisci_immagine(doc, 'Schema ANSC_USR', IMG)
    fatti.append('ERD di §8.1 rigenerato')

    # ------------------------------------------------- il capitolo di scopo
    scopo = D.h(doc, 1, 'Scopo del documento')
    i = D.indice_di(doc, scopo)
    fine = next(j for j in range(i + 1, len(doc.paragraphs))
                if doc.paragraphs[j].style.name == 'Heading 1')
    dopo = doc.paragraphs[fine]._p
    D.para(doc, dopo, 'Modifiche rispetto alla versione 3.13', stile='Heading 2')
    D.para(doc, dopo,
           'Questa versione rivede il modo in cui la configurazione mette in relazione i dati '
           'di SIPO con i casi d’uso di ANSC. La revisione nasce da due riscontri: '
           'l’osservazione diretta della web app di ANSC, che mostra come il caso d’uso venga '
           'composto rispondendo a domande in sequenza, e la decisione del Comune di affidare '
           'ai funzionari la configurazione, con la precompilazione in funzione di aiuto.')
    for testa, corpo in MODIFICHE:
        D.voce(doc, dopo, testa, corpo)
    fatti.append(f'capitolo Scopo: «Modifiche rispetto alla 3.13» ({len(MODIFICHE)} voci)')

    D.storia(doc, '06/09/2026', '3.14',
             'Cap. Modello dati · Cap. Le regole di determinazione e di controllo · '
             'Cap. Back-office · PC-9 · App. A',
             'La configurazione diventa un dato del Comune: ANSC_CFG_UC sostituisce '
             'ANSC_CFG_OPERAZIONE con una riga per caso d’uso, la determinazione dell’UC '
             'diventa una condizione in forma di interrogazione sui dati di SIPO '
             '(ANSC_CFG_UC_CONDIZIONE, PC-9 rivista), ogni riga di configurazione dichiara la '
             'propria origine perché la reimportazione non cancelli il lavoro dei funzionari. '
             'Schermate del back-office allineate; ERD rigenerato.')
    fatti.append('storia del documento')
    return fatti


if __name__ == '__main__':
    doc = docx.Document(DOC)
    for f in applica(doc):
        print('  ·', f)
    doc.save(DOC)
    print('salvato')
