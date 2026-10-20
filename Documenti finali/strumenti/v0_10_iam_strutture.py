# -*- coding: utf-8 -*-
"""ANALISI_Identita-Profilazione-IAM v0.9 → v0.10 (20/10/2026).

Richiesta dell'utente: ricavare le strutture dati di L4 (roma-capitale-management: utenti,
profili, abilitazioni), dichiarando che — in attesa delle DDL — sono estratte dalla mappatura
JPA; produrne una DDL per ORACLE (non PostgreSQL) e una serie di raccomandazioni perché queste
tabelle stiano in un dominio a parte per motivi di sicurezza.

Una sola specifica (TABELLE) alimenta sia le tabelle descrittive del documento sia la DDL:
così le due cose non possono divergere.

Fonti (tutte in serv_roma-capitale-management-main/, [F13]): le 17 classi di
src/main/java/.../management/entity/, i file di interrogazione in src/main/resources/.../repository/
(JPQL, nessuna interrogazione nativa), pom.xml:113 (driver PostgreSQL), application.properties:9
(currentSchema=management — ⚠️ il file contiene credenziali in chiaro: non riportarle).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_10_iam_strutture.py
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402
import diagrammi_iam      # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
IMG = os.path.join(FIN, 'strumenti', 'img')
SRC = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.9.docx')
DST = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.10.docx')

OWN = 'PROFILI_OWN'          # proprietario dello schema (proposta, PI-40)
APP = 'PROFILI_APP'          # utenza applicativa di roma-capitale-management (proposta)
TS = 'TS_PROFILI'            # tablespace dedicato (proposta)
RL_W = 'RL_PROFILI_SCRITTURA'
RL_R = 'RL_PROFILI_LETTURA'
TITOLO = 'Le strutture dati del servizio di profilazione'

# ------------------------------------------------------------------------------ la specifica
# colonna: (nome, tipo JPA, tipo Oracle, obbligatoria, nota)
# chiavi: pk, uk [(nome, colonne, stato)], fk [(nome, colonne, tabella, colonne_rif, stato)],
#         ck [(nome, espressione, evidenza)], seq
DICH, PROP, DED = 'dichiarata', 'proposta', 'dedotta dal nome'


def c(nome, jpa, ora, obbl=False, nota=''):
    return (nome, jpa, ora, obbl, nota)


TABELLE = [
    dict(nome='ambito', classe='AmbitoEntity', gruppo='Catalogo',
         desc='Gli ambiti della piattaforma (per esempio un’area dell’ente).',
         col=[c('cd_ambito', 'String(2)', 'VARCHAR2(2 CHAR)', True, 'Chiave.'),
              c('ds_ambito', 'String(50)', 'VARCHAR2(50 CHAR)', True)],
         pk=['cd_ambito']),
    dict(nome='applicazione', classe='ApplicazioneEntity', gruppo='Catalogo',
         desc='Le applicazioni, con indirizzo, contesto, cloud e modalità di accesso.',
         col=[c('cd_ambito', 'String(2)', 'VARCHAR2(2 CHAR)', True, 'Chiave, con cd_applicazione.'),
              c('cd_applicazione', 'String(10)', 'VARCHAR2(10 CHAR)', True, 'Chiave.'),
              c('ds_applicazione', 'String(60)', 'VARCHAR2(60 CHAR)', True),
              c('ds_endpoint', 'String(125)', 'VARCHAR2(125 CHAR)', True),
              c('ds_context_root', 'String(125)', 'VARCHAR2(125 CHAR)', True,
                'È ciò che il service mesh risolve a ogni instradamento.'),
              c('cd_doc_manuale', 'String(36)', 'VARCHAR2(36 CHAR)'),
              c('cd_cloud', 'String(1)', 'CHAR(1)', True),
              c('ds_service', 'String(50)', 'VARCHAR2(50 CHAR)', True),
              c('ds_label', 'String(30)', 'VARCHAR2(30 CHAR)', True),
              c('nm_ordine', 'Long', 'NUMBER(19)'),
              c('id_applicazione_sa', 'Long', 'NUMBER(19)', True,
                'Univoca: è l’unico vincolo di unicità che la mappatura dichiara.'),
              c('cd_ambito_padre', 'String(2)', 'VARCHAR2(2 CHAR)', False,
                'Con cd_applicazione_padre: l’applicazione padre.'),
              c('cd_applicazione_padre', 'String(10)', 'VARCHAR2(10 CHAR)'),
              c('ty_login', 'String(1)', 'CHAR(1)', False, 'Valori ammessi [DA VERIFICARE].'),
              c('ty_auth', 'String(1)', 'CHAR(1)', False,
                'G = di gruppo, unisce le abilitazioni delle figlie; altrimenti P.'),
              c('fg_anonymous', 'String(1)', 'CHAR(1)'),
              c('fg_lista_strutture', 'String(1)', 'CHAR(1)')],
         pk=['cd_ambito', 'cd_applicazione'],
         uk=[('UK_APPLICAZIONE_SA', ['id_applicazione_sa'], DICH)],
         fk=[('FK_APPLICAZIONE_AMBITO', ['cd_ambito'], 'ambito', ['cd_ambito'], DICH),
             ('FK_APPLICAZIONE_PADRE', ['cd_ambito_padre', 'cd_applicazione_padre'],
              'applicazione', ['cd_ambito', 'cd_applicazione'], DICH)],
         ck=[('CK_APPLICAZIONE_TY_AUTH', "ty_auth IN ('G','P')", 'UtenteDomain.java:304')]),
    dict(nome='menu', classe='MenuEntity', gruppo='Catalogo',
         desc='Le voci di menu a cui un’abilitazione può corrispondere.',
         col=[c('id_menu', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('cd_menu', 'String(50)', 'VARCHAR2(50 CHAR)', True),
              c('ds_menu', 'String(100)', 'VARCHAR2(100 CHAR)'),
              c('id_applicazione_sa', 'Long', 'NUMBER(19)', False,
                'Riferimento all’applicazione dedotto dal nome: la mappatura non lo dichiara.')],
         pk=['id_menu'], seq='menu_id_menu_seq',
         fk=[('FK_MENU_APPLICAZIONE', ['id_applicazione_sa'], 'applicazione',
              ['id_applicazione_sa'], DED)]),
    dict(nome='abilitazione', classe='AbilitazioneEntity', gruppo='Catalogo',
         desc='Il diritto elementare di un’applicazione, citato per codice nel token.',
         col=[c('id_abilitazione', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('cd_ambito', 'String', 'VARCHAR2(2 CHAR)', True,
                'Lunghezza non dichiarata: allineata ad applicazione.'),
              c('cd_applicazione', 'String', 'VARCHAR2(10 CHAR)', True,
                'Lunghezza non dichiarata: allineata ad applicazione.'),
              c('cd_abilitazione', 'String', 'VARCHAR2(255 CHAR)', True,
                'Lunghezza non dichiarata: 255 è il valore predefinito della JPA [DA VERIFICARE].'),
              c('ds_abilitazione', 'String(10)', 'VARCHAR2(10 CHAR)', True,
                'Dieci caratteri per una descrizione (MA-10) [DA VERIFICARE].'),
              c('cd_esterno', 'String(100)', 'VARCHAR2(100 CHAR)'),
              c('id_menu', 'Long', 'NUMBER(19)')],
         pk=['id_abilitazione'], seq='abilitazione_id_abilitazione_seq',
         uk=[('UK_ABILITAZIONE_CODICE', ['cd_ambito', 'cd_applicazione', 'cd_abilitazione'], PROP)],
         fk=[('FK_ABILITAZIONE_APPLICAZ', ['cd_ambito', 'cd_applicazione'], 'applicazione',
              ['cd_ambito', 'cd_applicazione'], DICH),
             ('FK_ABILITAZIONE_MENU', ['id_menu'], 'menu', ['id_menu'], DICH)]),
    dict(nome='profilo', classe='ProfiloEntity', gruppo='Profili',
         desc='Un insieme di abilitazioni di un’applicazione, con un codice di ruolo.',
         col=[c('id_profilo', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('ds_profilo', 'String(50)', 'VARCHAR2(50 CHAR)', True),
              c('fg_profilo_default', 'String(1)', 'CHAR(1)', True),
              c('cd_ruolo_profilo', 'String(50)', 'VARCHAR2(50 CHAR)', False,
                'Il posto del ruolo di SIPO.'),
              c('nm_valore', 'Long', 'NUMBER(19)'),
              c('cd_ambito', 'String(2)', 'VARCHAR2(2 CHAR)', False,
                'Facoltativa, come cd_applicazione: un profilo può non appartenere ad alcuna '
                'applicazione.'),
              c('cd_applicazione', 'String(10)', 'VARCHAR2(10 CHAR)')],
         pk=['id_profilo'], seq='profilo_id_profilo_seq',
         fk=[('FK_PROFILO_APPLICAZIONE', ['cd_ambito', 'cd_applicazione'], 'applicazione',
              ['cd_ambito', 'cd_applicazione'], PROP)]),
    dict(nome='profilo_per_abilitazione', classe='ProfiloPerAbilitazioneEntity', gruppo='Profili',
         desc='Quali abilitazioni compongono un profilo.',
         col=[c('id_profilo', 'Long', 'NUMBER(19)', True, 'Chiave, con id_abilitazione.'),
              c('id_abilitazione', 'Long', 'NUMBER(19)', True, 'Chiave.')],
         pk=['id_profilo', 'id_abilitazione'],
         fk=[('FK_PPA_PROFILO', ['id_profilo'], 'profilo', ['id_profilo'], DICH),
             ('FK_PPA_ABILITAZIONE', ['id_abilitazione'], 'abilitazione', ['id_abilitazione'],
              DICH)]),
    dict(nome='ente', classe='EnteEntity', gruppo='Utenti',
         desc='L’ente di appartenenza dell’utente.',
         col=[c('id_ente', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('cd_ente', 'Long', 'NUMBER(19)', True),
              c('ds_ente', 'String(250)', 'VARCHAR2(250 CHAR)')],
         pk=['id_ente'], seq='ente_id_ente_seq'),
    dict(nome='utente', classe='UtenteEntity', gruppo='Utenti',
         desc='Gli utenti della piattaforma, con i dati di contatto e lo stato dell’utenza.',
         col=[c('id_utente', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('ds_username', 'String(50)', 'VARCHAR2(50 CHAR)', True,
                'Univoca per proposta: l’inserimento ne controlla l’esistenza.'),
              c('ds_password', 'String(260)', 'VARCHAR2(260 CHAR)', True,
                'Cifrata in modo reversibile (UtenteDomain.java:185); vedi DS-08.'),
              c('ds_nome', 'String(60)', 'VARCHAR2(60 CHAR)', True),
              c('ds_cognome', 'String(60)', 'VARCHAR2(60 CHAR)', True),
              c('cd_fiscale', 'String(16)', 'VARCHAR2(16 CHAR)', True,
                'Univoca per proposta: l’inserimento ne controlla l’esistenza, e il login dei '
                'dipendenti cerca l’utente per codice fiscale.'),
              c('nm_telefono', 'String(20)', 'VARCHAR2(20 CHAR)', True),
              c('ds_email', 'String(255)', 'VARCHAR2(255 CHAR)', True),
              c('ds_ufficio', 'String(255)', 'VARCHAR2(255 CHAR)', True),
              c('dt_ultimo_login', 'Date', 'TIMESTAMP', False,
                'Regge la disattivazione dopo sei mesi di inattività.'),
              c('fg_attivo_sn', 'Character', 'CHAR(1)', True, 'S o N.'),
              c('nm_tentativi_pin', 'Long', 'NUMBER(19)'),
              c('nm_tentativi_login', 'Long', 'NUMBER(19)'),
              c('fg_scad_account', 'String(1)', 'CHAR(1)'),
              c('fg_scad_password', 'String(1)', 'CHAR(1)'),
              c('fg_cambio_password', 'String(1)', 'CHAR(1)'),
              c('ts_agg_password', 'Date', 'TIMESTAMP'),
              c('cd_cnt', 'Long', 'NUMBER(19)'),
              c('cd_stato', 'String(1)', 'CHAR(1)', False, 'Valori ammessi [DA VERIFICARE].'),
              c('ty_utente', 'String(1)', 'CHAR(1)', False, 'Valori ammessi [DA VERIFICARE].'),
              c('id_ente', 'Long', 'NUMBER(19)')],
         pk=['id_utente'], seq='utente_id_utente_seq',
         uk=[('UK_UTENTE_USERNAME', ['ds_username'], PROP),
             ('UK_UTENTE_CF', ['cd_fiscale'], PROP)],
         fk=[('FK_UTENTE_ENTE', ['id_ente'], 'ente', ['id_ente'], DICH)],
         ck=[('CK_UTENTE_ATTIVO', "fg_attivo_sn IN ('S','N')", 'LoginDomain.java:140 (auth)')]),
    dict(nome='struttura', classe='StrutturaEntity', gruppo='Organizzazione',
         desc='Gli alberi organizzativi.',
         col=[c('id_struttura', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('nm_struttura', 'String(200)', 'VARCHAR2(200 CHAR)', True),
              c('ds_struttura', 'String(1000)', 'VARCHAR2(1000 CHAR)', True)],
         pk=['id_struttura'], seq='struttura_id_struttura_seq'),
    dict(nome='nodo_struttura', classe='NodoStrutturaEntity', gruppo='Organizzazione',
         desc='I nodi di un albero organizzativo.',
         col=[c('id_nodo', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('id_struttura', 'Long', 'NUMBER(19)', True),
              c('id_nodo_padre', 'Long', 'NUMBER(19)', False,
                'Riferimento al padre dedotto dal nome: la mappatura non lo dichiara.'),
              c('cd_nodo', 'String(50)', 'VARCHAR2(50 CHAR)', True),
              c('nm_nodo', 'String(200)', 'VARCHAR2(200 CHAR)', True),
              c('ds_nodo', 'String(1000)', 'VARCHAR2(1000 CHAR)')],
         pk=['id_nodo'], seq='nodoStruttura_id_nodo_seq',
         fk=[('FK_NODO_STRUTTURA', ['id_struttura'], 'struttura', ['id_struttura'], DICH),
             ('FK_NODO_PADRE', ['id_nodo_padre'], 'nodo_struttura', ['id_nodo'], DED)]),
    dict(nome='visibilita', classe='VisibilitaEntity', gruppo='Organizzazione',
         desc='Le regole di visibilità sui dati: tabella, colonna, valore.',
         col=[c('id_visibilita', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('nm_tabella', 'String(250)', 'VARCHAR2(250 CHAR)'),
              c('nm_colonna', 'String(100)', 'VARCHAR2(100 CHAR)', True),
              c('ds_valore', 'String(1000)', 'VARCHAR2(1000 CHAR)', True),
              c('ds_visibilita', 'String(1000)', 'VARCHAR2(1000 CHAR)', True),
              c('ds_breve', 'String(50)', 'VARCHAR2(50 CHAR)', True),
              c('dipendenza', 'Long', 'NUMBER(19)', False, 'Rimanda a un’altra regola.')],
         pk=['id_visibilita'], seq='visibilita_id_visibilita_seq',
         fk=[('FK_VISIBILITA_DIPENDENZA', ['dipendenza'], 'visibilita', ['id_visibilita'], DICH)]),
    dict(nome='nodo_per_visibilita', classe='NodoPerVisibilitaEntity', gruppo='Organizzazione',
         desc='Quali regole di visibilità porta un nodo.',
         col=[c('id_nodo', 'Long', 'NUMBER(19)', True, 'Chiave, con id_visibilita.'),
              c('id_visibilita', 'Long', 'NUMBER(19)', True, 'Chiave.')],
         pk=['id_nodo', 'id_visibilita'],
         fk=[('FK_NPV_NODO', ['id_nodo'], 'nodo_struttura', ['id_nodo'], PROP),
             ('FK_NPV_VISIBILITA', ['id_visibilita'], 'visibilita', ['id_visibilita'], DICH)]),
    dict(nome='utente_per_profilo', classe='UtentePerProfiloEntity', gruppo='Utenti',
         desc='L’assegnazione di un profilo a un utente, eventualmente su un nodo organizzativo.',
         col=[c('id_utente', 'Long', 'NUMBER(19)', True, 'Chiave, con id_profilo.'),
              c('id_profilo', 'Long', 'NUMBER(19)', True, 'Chiave.'),
              c('id_nodo', 'Long', 'NUMBER(19)', False,
                'Nessuna operazione la scrive (MA-03).')],
         pk=['id_utente', 'id_profilo'],
         fk=[('FK_UPP_UTENTE', ['id_utente'], 'utente', ['id_utente'], PROP),
             ('FK_UPP_PROFILO', ['id_profilo'], 'profilo', ['id_profilo'], DICH),
             ('FK_UPP_NODO', ['id_nodo'], 'nodo_struttura', ['id_nodo'], DICH)]),
    dict(nome='utente_per_profilo_get', classe='UtentePerProfiloGetEntity',
         gruppo='Di una sola applicazione',
         desc='Le assegnazioni legate alle strutture del protocollo, per lo sportello unico.',
         col=[c('id_utente', 'Long', 'NUMBER(19)', True, 'Chiave, con le due seguenti.'),
              c('id_profilo', 'Long', 'NUMBER(19)', True),
              c('codice_struttura_ged', 'String', 'VARCHAR2(25 CHAR)', True,
                'Lunghezza allineata a v_str_organizzativa.codice; una vista non si può '
                'vincolare con una chiave esterna.')],
         pk=['id_utente', 'id_profilo', 'codice_struttura_ged'],
         fk=[('FK_UPPG_UTENTE', ['id_utente'], 'utente', ['id_utente'], PROP),
             ('FK_UPPG_PROFILO', ['id_profilo'], 'profilo', ['id_profilo'], DICH)]),
    dict(nome='configurazione', classe='ConfigurazioneEntity', gruppo='Registro della piattaforma',
         desc='La configurazione remota dei servizi, per applicazione.',
         col=[c('cd_ambito', 'String(2)', 'VARCHAR2(2 CHAR)', True, 'Chiave, con le due seguenti.'),
              c('cd_applicazione', 'String(10)', 'VARCHAR2(10 CHAR)', True),
              c('des_nome_parametro', 'String(100)', 'VARCHAR2(100 CHAR)', True),
              c('des_val_parametro', 'text', 'CLOB', True),
              c('ty_parametro', 'String(1)', 'CHAR(1)')],
         pk=['cd_ambito', 'cd_applicazione', 'des_nome_parametro'],
         fk=[('FK_CONFIGURAZIONE_APPLICAZ', ['cd_ambito', 'cd_applicazione'], 'applicazione',
              ['cd_ambito', 'cd_applicazione'], DICH)]),
    dict(nome='avviso', classe='AvvisoEntity', gruppo='Registro della piattaforma',
         desc='Gli avvisi agli utenti, per applicazione.',
         col=[c('id_avviso', 'Long', 'NUMBER(19)', True, 'Chiave, da sequenza.'),
              c('cd_ambito', 'String(2)', 'VARCHAR2(2 CHAR)', True),
              c('cd_applicazione', 'String(10)', 'VARCHAR2(10 CHAR)', True),
              c('ds_titolo_avviso', 'String(100)', 'VARCHAR2(100 CHAR)', True),
              c('ds_avviso', 'text', 'CLOB', True),
              c('dt_inizio', 'Timestamp', 'TIMESTAMP', True),
              c('dt_fine', 'Timestamp', 'TIMESTAMP', True),
              c('nm_priorita', 'Short', 'NUMBER(5)', True),
              c('fl_attivo', 'Boolean', 'NUMBER(1)', True, '0 o 1.')],
         pk=['id_avviso'], seq='avvisi_id_avviso_seq',
         fk=[('FK_AVVISO_APPLICAZIONE', ['cd_ambito', 'cd_applicazione'], 'applicazione',
              ['cd_ambito', 'cd_applicazione'], DICH)],
         ck=[('CK_AVVISO_ATTIVO', 'fl_attivo IN (0,1)', 'AvvisoDomain.java:206, 225')]),
]
VISTA = dict(nome='v_str_organizzativa', classe='StrutturaOrganizzativaGetEntity',
             col=[('codice', 'VARCHAR2(25 CHAR)'), ('descrizione', 'VARCHAR2(255 CHAR)'),
                  ('descrizione_breve', 'VARCHAR2(255 CHAR)'),
                  ('codice_str_padre', 'VARCHAR2(25 CHAR)'),
                  ('ordine_organigramma', 'VARCHAR2(50 CHAR)'),
                  ('codice_responsabile', 'NUMBER(19)')])
NUCLEO = ['utente', 'utente_per_profilo', 'profilo', 'profilo_per_abilitazione',
          'abilitazione', 'applicazione', 'ambito']
T = {t['nome']: t for t in TABELLE}


def controlla():
    """I nomi di vincolo stanno nei 30 byte; ogni FK punta a colonne che esistono."""
    for t in TABELLE:
        cols = {x[0] for x in t['col']}
        assert set(t['pk']) <= cols, t['nome']
        nomi = [f'PK_{t["nome"]}'.upper()] + [u[0] for u in t.get('uk', [])] + \
               [f[0] for f in t.get('fk', [])] + [k[0] for k in t.get('ck', [])]
        for n in nomi:
            assert len(n) <= 30, n
        for _, cc, rif, rc, _ in t.get('fk', []):
            assert set(cc) <= cols and set(rc) <= {x[0] for x in T[rif]['col']}, (t['nome'], rif)
    assert len(TABELLE) == 16


# ------------------------------------------------------------------------------ la DDL
def ddl():
    L = []
    a = L.append
    a('-- ============================================================================')
    a('-- Dominio dei profili: le tabelle di roma-capitale-management su Oracle')
    a('-- Ricostruita dalla mappatura JPA del servizio (versione 1.0.20), in attesa delle')
    a('-- DDL dello schema PostgreSQL «management» (PI-39). Vincoli marcati «proposta» o')
    a('-- «dedotta» non sono dichiarati nel codice. Nomi di proprietario, utenze, ruoli e')
    a('-- tablespace sono proposte da concordare (PI-40). Richiede Oracle 18c o successivo.')
    a('-- ============================================================================')
    a('')
    a('-- 1. Tablespace cifrato e proprietario senza accesso (eseguito dai DBA)')
    a(f'CREATE TABLESPACE {TS} DATAFILE SIZE 1G AUTOEXTEND ON NEXT 256M')
    a("  ENCRYPTION USING 'AES256' DEFAULT STORAGE (ENCRYPT);   -- richiede TDE [DA VERIFICARE]")
    a(f'CREATE USER {OWN} NO AUTHENTICATION')
    a(f'  DEFAULT TABLESPACE {TS} QUOTA UNLIMITED ON {TS};')
    a('')
    a('-- 2. Sequenze: nomi e passo come nella mappatura (allocationSize = 1).')
    a('--    Nella migrazione START WITH = massimo identificativo esistente + 1.')
    for t in TABELLE:
        if t.get('seq'):
            if len(t['seq']) > 30:
                a(f'-- {len(t["seq"])} caratteri: serve Oracle 12.2 o successivo')
            a(f'CREATE SEQUENCE {OWN}.{t["seq"]}')
            a('  START WITH 1 INCREMENT BY 1 NOCYCLE;')
    a('')
    a('-- 3. Tabelle, nell’ordine imposto dalle chiavi esterne')
    for t in TABELLE:
        a('')
        a(f'CREATE TABLE {OWN}.{t["nome"]} (')
        # ogni voce è una lista di righe; la virgola va in coda all'ultima riga della voce,
        # prima dell'eventuale commento. Righe entro i ~90 caratteri della pagina.
        voci = []
        w = max(len(x[0]) for x in t['col'])
        for nome, _, ora, obbl, _ in t['col']:
            voci.append(([f'  {nome.ljust(w)}  {ora}{" NOT NULL" if obbl else ""}'], ''))
        voci.append(([f'  CONSTRAINT PK_{t["nome"].upper()} PRIMARY KEY ({", ".join(t["pk"])})'],
                     ''))
        for n, cc, stato in t.get('uk', []):
            voci.append(([f'  CONSTRAINT {n} UNIQUE ({", ".join(cc)})'],
                         '' if stato == DICH else f'   /* {stato} */'))
        for n, cc, rif, rc, stato in t.get('fk', []):
            voci.append(([f'  CONSTRAINT {n} FOREIGN KEY ({", ".join(cc)})',
                          f'    REFERENCES {OWN}.{rif} ({", ".join(rc)})'],
                         '' if stato == DICH else f'   /* {stato} */'))
        for n, espr, _ in t.get('ck', []):
            voci.append(([f'  CONSTRAINT {n} CHECK ({espr})'], ''))
        for i, (righe, nota) in enumerate(voci):
            for r in righe[:-1]:
                a(r)
            a(righe[-1] + ('' if i == len(voci) - 1 else ',') + nota)
        a(f') TABLESPACE {TS};')
        a(f'COMMENT ON TABLE {OWN}.{t["nome"]} IS')
        a(f"  '{t['desc'].replace(chr(8217), chr(39) * 2)}';")
    a('')
    a('-- 4. Indici sulle colonne delle chiavi esterne non coperte dalla chiave primaria')
    for t in TABELLE:
        for n, cc, _, _, _ in t.get('fk', []):
            if cc == t['pk'][:len(cc)]:
                continue
            ix = 'IX_' + n[3:]
            assert len(ix) <= 30, ix
            a(f'CREATE INDEX {OWN}.{ix}')
            a(f'  ON {OWN}.{t["nome"]} ({", ".join(cc)}) TABLESPACE {TS};')
    a('')
    a('-- 5. La vista dell’organigramma: la definizione non è fra le sorgenti [DA VERIFICARE]')
    a(f'-- CREATE OR REPLACE VIEW {OWN}.{VISTA["nome"]} AS')
    a('--   SELECT ' + ', '.join(x[0] for x in VISTA['col'][:3]) + ',')
    a('--          ' + ', '.join(x[0] for x in VISTA['col'][3:]))
    a('--     FROM <origine dell’organigramma del protocollo>;')
    a('--   colonne attese:')
    for n, tipo in VISTA['col']:
        a(f'--     {n.ljust(20)} {tipo}')
    a('')
    a('-- 6. Ruoli: scrittura per il servizio, sola lettura per la revisione degli accessi')
    a(f'CREATE ROLE {RL_W};')
    a(f'CREATE ROLE {RL_R};')
    for t in TABELLE:
        a(f'GRANT SELECT, INSERT, UPDATE, DELETE ON {OWN}.{t["nome"]}')
        a(f'  TO {RL_W};')
        a(f'GRANT SELECT ON {OWN}.{t["nome"]} TO {RL_R};')
    for t in TABELLE:
        if t.get('seq'):
            a(f'GRANT SELECT ON {OWN}.{t["seq"]} TO {RL_W};')
    a('')
    a('-- 7. Utenza applicativa: solo sessione e ruolo, nessun privilegio di sistema.')
    a('--    La parola chiave sta nel gestore dei segreti, mai nella configurazione versionata.')
    a(f'CREATE USER {APP} IDENTIFIED BY "<dal gestore dei segreti>"')
    a('  PROFILE <profilo con scadenza e blocco dei tentativi> ACCOUNT UNLOCK;')
    a(f'GRANT CREATE SESSION TO {APP};')
    a(f'GRANT {RL_W} TO {APP};')
    a(f'ALTER USER {APP} DEFAULT ROLE ALL;')
    a('-- Le classi non nominano lo schema: con i sinonimi privati il codice non cambia.')
    for t in TABELLE:
        a(f'CREATE SYNONYM {APP}.{t["nome"]}')
        a(f'  FOR {OWN}.{t["nome"]};')
    for t in TABELLE:
        if t.get('seq'):
            a(f'CREATE SYNONYM {APP}.{t["seq"]}')
            a(f'  FOR {OWN}.{t["seq"]};')
    a('')
    a('-- 8. Tracciamento delle modifiche (Unified Auditing), conservato fuori dal dominio')
    a('CREATE AUDIT POLICY POL_PROFILI_MODIFICHE')
    traccia = ['utente', 'utente_per_profilo', 'profilo', 'profilo_per_abilitazione',
               'abilitazione', 'applicazione']
    voci = [f'{azione} ON {OWN}.{n}' for n in traccia for azione in ('INSERT', 'UPDATE', 'DELETE')]
    for i, v in enumerate(voci):
        a(('  ACTIONS ' if i == 0 else '          ') + v + (';' if i == len(voci) - 1 else ','))
    a('AUDIT POLICY POL_PROFILI_MODIFICHE;')
    a('-- Le modifiche di struttura e di permesso sono rare: si tracciano per tutte le utenze.')
    a('CREATE AUDIT POLICY POL_PROFILI_STRUTTURA')
    a('  ACTIONS CREATE TABLE, ALTER TABLE, DROP TABLE, CREATE INDEX, DROP INDEX,')
    a('          CREATE SEQUENCE, ALTER SEQUENCE, DROP SEQUENCE, CREATE SYNONYM,')
    a('          DROP SYNONYM, GRANT, REVOKE, ALTER USER;')
    a('AUDIT POLICY POL_PROFILI_STRUTTURA;')
    return L


# ------------------------------------------------------------------------------ i testi
TIPI = [
    ['Nella mappatura JPA', 'In PostgreSQL, presumibilmente', 'In Oracle'],
    ['Long', 'bigint', 'NUMBER(19)'],
    ['Short', 'smallint', 'NUMBER(5)'],
    ['String con lunghezza n', 'varchar(n)', 'VARCHAR2(n CHAR)'],
    ['String senza lunghezza', 'varchar(255), il valore predefinito della JPA',
     'VARCHAR2(255 CHAR), o la lunghezza della colonna a cui si lega'],
    ['String di lunghezza 1, Character', 'char(1) o varchar(1)', 'CHAR(1)'],
    ['String dichiarata «text»', 'text', 'CLOB'],
    ['Date, Timestamp', 'timestamp', 'TIMESTAMP'],
    ['Boolean', 'boolean', 'NUMBER(1) con vincolo 0/1; BOOLEAN dalla versione 23ai'],
]

LACUNE = [
    ['Aspetto', 'Che cosa dice la mappatura JPA', 'Che cosa fa la DDL per Oracle', 'Stato'],
    ['Tabelle e colonne', 'Nomi, tipi Java, lunghezze, obbligatorietà dove dichiarata.',
     'Le riporta tutte, con i tipi della tabella precedente.', 'Dichiarato'],
    ['Chiavi primarie', 'Dichiarate su ogni classe, semplici o composte.', 'Le riporta.',
     'Dichiarato'],
    ['Sequenze', 'Nome e passo (uno) per ciascuna chiave generata.', 'Le riporta con lo stesso '
     'nome, che il servizio usa a ogni inserimento.', 'Dichiarato'],
    ['Obbligatorietà', 'Solo dove la colonna lo dice; le altre risultano facoltative.',
     'La riporta così; il database vero può essere più severo.', 'Da confermare'],
    ['Chiavi esterne', 'Non dichiarate: la mappatura dice soltanto come le classi si legano.',
     '{nfk} chiavi esterne: {ndich} dai legami fra le classi, {nprop} proposte dove il legame è '
     'implicito, {nded} dedotte dal solo nome della colonna.', 'Proposto'],
    ['Unicità', 'Una sola colonna univoca: l’identificativo dell’applicazione per il servizio '
     'esterno.', 'La riporta e ne propone tre: nome utente, codice fiscale, codice '
     'dell’abilitazione nella sua applicazione.', 'Proposto'],
    ['Valori ammessi', 'Nessuno.', 'Vincoli solo dove il codice confronta valori precisi: '
     'utente attivo S/N, autorizzazione G/P, avviso attivo 0/1. Gli altri contrassegni restano '
     'liberi.', 'Da confermare'],
    ['Indici', 'Nessuno.', 'Uno per ogni chiave esterna non coperta dalla chiave primaria.',
     'Proposto'],
    ['La vista dell’organigramma', 'Le sole colonne.', 'Lasciata come schema da completare.',
     'Da verificare'],
    ['Permessi, spazio, tracciamento', 'Nulla: non sono materia della mappatura.',
     'Dalle raccomandazioni sul dominio dei profili.', 'Proposto'],
]

PASSAGGIO = [
    ['Punto', 'Che cosa tocca', 'Effetto sul servizio'],
    ['Driver e indirizzo', 'Il progetto dichiara il driver PostgreSQL (pom.xml:113) e '
     'l’indirizzo della base dati porta lo schema nel parametro di connessione '
     '(application.properties:9).', 'Driver Oracle al posto di quello PostgreSQL e nuovo '
     'indirizzo: è un rilascio del servizio.'],
    ['Dialetto', 'Non dichiarato: lo strato di persistenza lo ricava dalla connessione.',
     'Nessuno.'],
    ['Interrogazioni', 'Tutte in linguaggio di interrogazione degli oggetti, in file separati; '
     'nessuna interrogazione nativa.', 'Portabili. Alcune applicano la conversione in minuscolo '
     'anche a colonne numeriche (Applicazione.xml:42, 47): Oracle le converte da sé, ma va '
     'provato in collaudo [DA VERIFICARE].'],
    ['Schema', 'Le classi non nominano lo schema; oggi lo sceglie l’indirizzo di connessione.',
     'Sinonimi privati dell’utenza applicativa, già nella DDL, oppure lo schema predefinito '
     'dello strato di persistenza: il codice non cambia.'],
    ['Sequenze', 'Il nome della sequenza delle abilitazioni è lungo 32 caratteri.',
     'Serve Oracle 12.2 o successivo; per l’utenza senza accesso diretto serve la 18c.'],
    ['Testi lunghi e contrassegni', 'Due colonne dichiarate «text» e un booleano.',
     'CLOB e NUMBER(1). La dichiarazione «text» serve solo a generare lo schema, che il servizio '
     'non fa; la lettura del CLOB come stringa va provata [DA VERIFICARE].'],
    ['Chiavi esterne', 'Oggi, se esistono, non sono note. Per utenti e profili la gestione '
     'cancella prima le righe collegate (UtenteDomain.java:268, ProfiloDomain.java:164-169); per '
     'abilitazioni, applicazioni, menu, ambiti, strutture e regole di visibilità cancella la sola '
     'riga (AbilitazioneDomain.java:178, ApplicazioneDomain.java:256, MenuDomain.java:92, '
     'AmbitoDomain.java:104, StrutturaDomain.java:92, VisibilitaDomain.java:50).',
     'Con le chiavi esterne la cancellazione di un oggetto ancora usato viene rifiutata, invece '
     'di lasciare righe senza oggetto: è un cambiamento di comportamento voluto, ma il servizio '
     'restituirà un errore dove oggi riesce, e va provato (MA-05).'],
    ['Dati esistenti', 'Lo schema PostgreSQL in esercizio.', 'Migrazione dei dati e '
     'riallineamento delle sequenze al massimo identificativo più uno.'],
    ['Titolarità', 'Il servizio e la sua base dati sono della piattaforma, che serve anche '
     'altre aree dell’ente.', 'La scelta di Oracle va presa con il titolare (PI-41).'],
]

RACCOMANDAZIONI = [
    ['#', 'Raccomandazione', 'Perché', 'Come, in Oracle'],
    ['DS-01', '**Un dominio proprio** per le tabelle dei profili: una base dati collegabile '
     'dedicata o, come minimo, uno schema dedicato in un tablespace proprio.',
     'Chi scrive queste tabelle decide chi entra in tutte le applicazioni della piattaforma. Un '
     'difetto o una compromissione di un’altra applicazione non deve dare accesso ai profili.',
     'PDB dedicato; in subordine schema e tablespace dedicati nell’istanza condivisa.'],
    ['DS-02', 'Il proprietario dello schema **non si collega**.', 'Le strutture cambiano solo '
     'con un rilascio controllato, non dall’applicazione e non da una sessione qualunque.',
     'Utenza senza autenticazione; i rilasci si eseguono per delega da un’utenza amministrativa '
     'nominativa.'],
    ['DS-03', '**Una sola utenza applicativa**, quella di roma-capitale-management, con i soli '
     'privilegi sulle righe.', 'Oggi il servizio non dichiara alcun permesso per operazione: il '
     'perimetro della base dati è l’ultimo limite a ciò che può fare.', 'Sessione e ruolo di '
     'scrittura sulle sole tabelle; nessun privilegio di sistema; sinonimi privati.'],
    ['DS-04', 'Gli altri servizi passano **dalle API, non dalla base dati**.', 'È già così: '
     'roma-capitale-auth e il service mesh non hanno tabelle e leggono management con i '
     'percorsi interni. Va mantenuto, e quei percorsi vanno chiusi (MA-12).', 'Nessuna utenza '
     'per auth e per il service mesh.'],
    ['DS-05', '**Nessun collegamento diretto con gli schemi di SIPO**.', 'Domini di guasto e di '
     'compromissione separati; coerente con la scelta di lasciare il vecchio ambiente com’è. '
     'Se SIPO deve alimentare la piattaforma, lo fa per API (modalità A).', 'Niente database '
     'link, sinonimi pubblici o chiavi esterne fra MATR_USR, ANAG_USR, ANSC_USR e il dominio.'],
    ['DS-06', 'Il dominio è **raggiungibile solo dai pod di management** e dall’accesso '
     'amministrativo.', 'Una credenziale sottratta non basta, se non si è nel posto giusto.',
     'Servizio di base dati dedicato, regola di rete in Kubernetes, liste di accesso del '
     'listener; connessione cifrata.'],
    ['DS-07', 'Le **credenziali di connessione** fuori dalla configurazione versionata, e '
     'cambiate al passaggio.', 'Oggi sono in chiaro nel file del repository (MA-14): vanno '
     'considerate note.', 'Gestore dei segreti; profilo dell’utenza con scadenza e blocco dei '
     'tentativi.'],
    ['DS-08', 'Le **password come impronta**, e facoltative.', 'Oggi sono cifrate in modo '
     'reversibile con una chiave unica. Per i dipendenti l’accesso passa da IAM e la password '
     'locale non serve.', 'Colonna facoltativa; impronta con sale calcolata dal servizio, non '
     'dal database.'],
    ['DS-09', '**Cifratura a riposo** dei dati personali.', 'Codice fiscale, posta, telefono e '
     'password: le copie e i file di dati non devono essere leggibili fuori dal database.',
     'Tablespace cifrato; copie di sicurezza cifrate [DA VERIFICARE: licenza].'],
    ['DS-10', 'Il **tracciamento** di ogni modifica a utenti, profili e assegnazioni, fuori '
     'dalla portata dell’utenza applicativa.', 'Oggi le assegnazioni si cancellano e si '
     'reinseriscono senza traccia (MA-06). Il tracciamento del database non sostituisce la '
     'storia applicativa (MI-01), ma la rende verificabile.', 'Politiche di Unified Auditing; '
     'traccia esportata verso il sistema di raccolta dei log.'],
    ['DS-11', '**Separazione dei compiti** per chi amministra la base dati.', 'Chi amministra '
     'l’istanza non deve potersi concedere un profilo senza che resti traccia.', 'Tracciamento '
     'degli utenti privilegiati; se disponibile, un recinto di Database Vault sullo schema '
     '[DA VERIFICARE: licenza].'],
    ['DS-12', '**Integrità nel database**, non solo nel codice.', 'La gestione è fatta di '
     'operazioni elementari senza controlli di autorità: i vincoli impediscono almeno '
     'assegnazioni a profili inesistenti e codici duplicati.', 'Chiavi esterne, unicità e '
     'vincoli della DDL.'],
    ['DS-13', '**Ambienti di prova senza dati reali**.', 'Una copia della produzione porta '
     'fuori codici fiscali e assegnazioni di tutto l’ente.', 'Dati sintetici o mascherati; '
     'nessun aggiornamento da produzione a collaudo senza mascheramento.'],
]

PI = [
    ('PI-39', 'Le DDL dello schema «management» in esercizio, per esempio l’estrazione della sola '
     'struttura dal catalogo PostgreSQL, per confermare tipi, obbligatorietà, vincoli, indici e '
     'valori ammessi che la mappatura JPA non dichiara e che la DDL per Oracle propone.',
     'Titolare della piattaforma / Sistemi'),
    ('PI-40', 'La versione e le opzioni dell’Oracle di destinazione: identificatori lunghi (12.2), '
     'utenze senza autenticazione (18c), tipo booleano (23ai), base dati collegabile dedicata, '
     'cifratura a riposo e separazione dei compiti, con le relative licenze; i nomi di '
     'proprietario, utenze, ruoli e tablespace proposti.', 'Sistemi / DBA'),
    ('PI-41', 'Che cosa significa portare la base dati dei profili su Oracle: un’istanza del '
     'servizio dedicata a SIPO con la propria base dati, oppure il passaggio dell’intera '
     'piattaforma, che serve anche altre aree dell’ente; e chi migra i dati esistenti.',
     'Titolare della piattaforma / Architettura'),
]

GLOSSARIO = [
    ('JPA', 'Java Persistence API: le annotazioni con cui le classi di un servizio dichiarano '
     'tabelle e colonne. In assenza delle DDL è la fonte delle strutture dati del servizio dei '
     'profili.'),
    ('DDL', 'Le istruzioni che creano tabelle, vincoli, sequenze, indici e permessi.'),
    ('Dominio dei profili', 'La collocazione separata proposta per le tabelle del servizio dei '
     'profili: una base dati collegabile dedicata, o almeno uno schema proprio.'),
]


def main():
    controlla()
    if os.path.exists(DST):
        os.remove(DST)                         # solo un'esecuzione precedente di questo script
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    modello = d.tables[3]
    fatti = []
    fk = [f for t in TABELLE for f in t.get('fk', [])]
    nfk = len(fk)
    conta = {s: sum(1 for f in fk if f[4] == s) for s in (DICH, PROP, DED)}

    # ------------------------------------------------ figure: la prossima numerazione
    nfig = max(int(m.group(1)) for p in d.paragraphs
               for m in [re.match(r'Figura (\d+)', p.text.strip())] if m) + 1
    png = os.path.join(IMG, diagrammi_iam.dominio(IMG))

    # ------------------------------------------------ Scopo, Executive summary, cap. 8
    scopo = [p for p in d.paragraphs if p.text.startswith('Il documento valuta anche il servizio')]
    assert len(scopo) == 1
    dopo = scopo[0]._p.getnext()
    D.para(d, dopo, 'Ne ricostruisce infine le strutture dati, che in attesa delle DDL sono '
           'estratte dalla mappatura JPA del codice, ne propone la DDL per Oracle e indica come '
           'collocarle in un dominio a parte per ragioni di sicurezza.')

    parte1 = D.h(d, 1, 'Parte I')
    D.para(d, parte1._p, pezzi=D.segmenta('', (
        'Le strutture dati del servizio dei profili, in attesa delle DDL dello schema, sono state '
        'ricostruite dalla mappatura JPA del codice. Ne discendono una **DDL per Oracle** e la '
        'raccomandazione di collocare queste tabelle in un **dominio a parte** — una base dati o '
        'almeno uno schema propri, raggiunti da una sola utenza applicativa e senza collegamenti '
        'diretti con gli schemi di SIPO —, perché chi le scrive decide chi entra in tutte le '
        'applicazioni della piattaforma.')))
    fatti.append('Scopo ed Executive summary: le strutture dati e il dominio a parte')

    accesso = D.h(d, 2, 'Come partecipa all’accesso')
    D.para(d, accesso._p, f'Il dettaglio colonna per colonna, i vincoli e la DDL per Oracle sono '
           f'nel capitolo «{TITOLO}».')

    # ------------------------------------------------ glossario e punti aperti
    gl = D.trova_tabella(d, 'Termine', 'Significato in questo documento')
    for r in GLOSSARIO:
        D.clona_riga(gl, r)
    pa = D.trova_tabella(d, '#', 'Questione', 'A chi compete')
    assert pa.rows[-1].cells[0].text.strip() == 'PI-38', pa.rows[-1].cells[0].text
    for r in PI:
        D.clona_riga(pa, r)
    fatti.append('glossario +3, PI-39…PI-41')

    # ------------------------------------------------ il capitolo nuovo
    fine = d.element.body[-1]
    assert fine.tag == qn('w:sectPr')

    def p(testo):
        D.para(d, fine, pezzi=D.segmenta('', testo)) if '**' in testo else D.para(d, fine, testo)

    def h2(testo):
        D.para(d, fine, testo, stile='Heading 2')

    D.para(d, fine, TITOLO, stile='Heading 1')
    p('**In attesa delle DDL dello schema, che non sono fra le sorgenti, le informazioni di questo '
      'capitolo sono state estratte dalla mappatura JPA del servizio** [F13]: le diciassette '
      'classi con cui roma-capitale-management dichiara tabelle e colonne, e i file delle sue '
      'interrogazioni. La mappatura dice con certezza nomi, tipi, lunghezze e chiavi primarie; '
      'non dice vincoli, indici né valori ammessi. Dove la DDL che segue va oltre ciò che la '
      'mappatura dichiara, lo dice, e le DDL vere dovranno confermarla (PI-39).')
    p('Il servizio non genera lo schema dal codice e non ha script di creazione né di '
      'migrazione: la base dati in esercizio è stata creata a parte. roma-capitale-auth e il '
      'service mesh non hanno tabelle proprie e leggono management attraverso i suoi percorsi '
      'interni: le strutture di questo capitolo sono quindi tutte quelle del servizio dei '
      'profili.')

    h2('Il metodo: che cosa dice la mappatura, e che cosa no')
    righe = [r[:] for r in LACUNE]
    righe[5][2] = righe[5][2].format(nfk=nfk, ndich=conta[DICH], nprop=conta[PROP],
                                     nded=conta[DED])
    D.tabella(d, fine, righe, modello, [1.3, 1.9, 2.3, 0.9])
    p('I tipi della mappatura si traducono così. La colonna di PostgreSQL è presumibile, perché '
      'la mappatura dichiara il tipo Java e non quello della base dati.')
    D.tabella(d, fine, TIPI, modello, [2.0, 2.0, 2.4])

    h2('Utenti, profili e abilitazioni, colonna per colonna')
    p('Le sette tabelle del nucleo sono quelle che decidono che cosa un operatore può fare: '
      'l’utente riceve profili, il profilo raccoglie abilitazioni, l’abilitazione appartiene a '
      'un’applicazione di un ambito. Per ciascuna, la classe da cui è ricavata.')
    for nome in NUCLEO:
        t = T[nome]
        D.titolo_ddl(d, fine, f'{nome} — {t["desc"]} Classe {t["classe"]}.')
        righe = [['Colonna', 'Nella JPA', 'In Oracle', 'Obbl.', 'Note']]
        for col, jpa, ora, obbl, nota in t['col']:
            righe.append([col, jpa, ora, 'sì' if obbl else '', nota])
        D.tabella(d, fine, righe, modello, [1.6, 1.0, 1.4, 0.5, 2.0])

    h2('Le altre tabelle')
    p('Le altre nove tabelle e la vista completano il servizio: l’ente dell’utente, '
      'l’organizzazione e la visibilità sui dati, il registro della piattaforma e la parte '
      'specifica dello sportello unico.')
    righe = [['Tabella', 'Gruppo', 'Chiave', 'Riferimenti', 'Classe']]
    for t in TABELLE:
        if t['nome'] in NUCLEO:
            continue
        rif = '; '.join(f'{", ".join(cc)} → {r}' for _, cc, r, _, _ in t.get('fk', [])) or '—'
        righe.append([t['nome'], t['gruppo'], ', '.join(t['pk']), rif, t['classe']])
    righe.append([VISTA['nome'], 'Di una sola applicazione', 'codice',
                  'Vista dell’organigramma del protocollo; definizione non fra le sorgenti',
                  VISTA['classe']])
    tz = D.tabella(d, fine, righe, modello, [1.6, 1.1, 1.3, 1.9, 1.5])
    # nomi di classe lunghi: Word allargherebbe la colonna
    tz._tbl.tblPr.append(docx.oxml.parse_xml(
        '<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'w:type="fixed"/>'))
    for r in tz.rows[1:]:
        for cc in (r.cells[0], r.cells[4]):
            for run in cc.paragraphs[0].runs:
                run.text = run.text.replace('_', '_​').replace('Entity', '​Entity')

    h2('Il passaggio a Oracle: che cosa tocca il servizio')
    p('Portare le tabelle su Oracle non richiede di riscrivere il servizio, perché le '
      'interrogazioni non dipendono dalla base dati; richiede però un rilascio, e alcune prove.')
    D.tabella(d, fine, PASSAGGIO, modello, [1.3, 2.4, 2.7])

    h2('La DDL per Oracle')
    p(f'La DDL crea il dominio dei profili nell’ordine imposto dalle chiavi esterne: tablespace e '
      f'proprietario, le {sum(1 for t in TABELLE if t.get("seq"))} sequenze, le '
      f'{len(TABELLE)} tabelle con le loro chiavi, gli indici, la vista da completare, i ruoli, '
      f'l’utenza applicativa con i sinonimi e il tracciamento. Nomi di tabelle e colonne sono '
      f'quelli della mappatura, perché il servizio li usa così; i nomi dei vincoli seguono i '
      f'prefissi PK_, UK_, FK_, CK_, IX_ e stanno nei trenta caratteri. I nomi di proprietario '
      f'({OWN}), utenza ({APP}), ruoli e tablespace sono proposte da concordare (PI-40). I '
      f'vincoli che la mappatura non dichiara sono marcati come proposta o come deduzione.')
    D.ddl(d, fine, ddl())

    h2('Un dominio a parte: le raccomandazioni di sicurezza')
    p('Queste tabelle non sono dati di un’applicazione fra le altre: dicono chi può entrare in '
      'ciascuna applicazione della piattaforma, e con quali diritti. Per questo **vanno '
      'collocate in un dominio a parte**, con un perimetro proprio. Le raccomandazioni valgono '
      'su Oracle e varrebbero anche se la base dati restasse PostgreSQL.')
    D.immagine(d, fine, png, 6.3,
               f'Figura {nfig} — Il dominio dei profili: le tabelle del servizio in un perimetro '
               f'proprio, raggiunto da una sola utenza applicativa, senza collegamenti diretti '
               f'con gli schemi di SIPO.')
    D.tabella(d, fine, RACCOMANDAZIONI, modello, [0.6, 1.8, 2.2, 1.8])
    p('Il minimo è lo schema dedicato in un tablespace cifrato, con il proprietario senza accesso, '
      'una sola utenza applicativa e nessun collegamento con gli schemi di SIPO (DS-01…DS-07). '
      'Il raccomandato è la base dati collegabile dedicata, che aggiunge un confine di istanza, '
      'copie di sicurezza e amministrazione separate. Quale delle due sia praticabile dipende '
      'dalla versione e dalle opzioni di Oracle dell’ente (PI-40).')
    fatti.append(f'nuovo capitolo «{TITOLO}»: {len(TABELLE)} tabelle + 1 vista, '
                 f'{nfk} chiavi esterne, DDL di {len(ddl())} righe, DS-01…DS-13, Figura {nfig}')

    D.storia(d, '20/10/2026', '0.10',
             'Scopo · Executive summary · Il servizio di profilazione della piattaforma · '
             'Glossario · Punti aperti · nuovo capitolo sulle strutture dati',
             'Ricostruite le strutture dati del servizio dei profili: in attesa delle DDL, le '
             'informazioni sono estratte dalla mappatura JPA del codice. Aggiunti il dettaglio '
             'colonna per colonna del nucleo utenti, profili e abilitazioni, la DDL per Oracle '
             'con vincoli proposti, ruoli, utenza applicativa e tracciamento, ciò che il '
             'passaggio a Oracle tocca nel servizio e le raccomandazioni per collocare le '
             'tabelle in un dominio a parte. Nuovi PI-39…PI-41.')
    for f in D.frontespizio(d, '0.10', 'ANALISI_Identita-Profilazione-IAM_v0.10'):
        fatti.append(f)

    d.save(DST)
    with open(os.path.join(FIN, 'DDL_Profili-Oracle_v0.1.sql'), 'w',
              encoding='utf-8') as fh:
        fh.write('\n'.join(ddl()) + '\n')
    print('\n'.join('  · ' + f for f in fatti))
    cap, tab, fig = D.riepilogo(DST)
    print(f'  capitoli {cap}, tabelle {tab}, immagini {fig}')
    print('  scritto', DST)


if __name__ == '__main__':
    main()
