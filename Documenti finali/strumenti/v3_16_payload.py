# -*- coding: utf-8 -*-
"""v3.16 — il payload inviato ad ANSC si conserva sull'atto, non solo nell'audit.

⚠️ Il payload c'era già, ma non dove serve: `ANSC_LOG_AUDIT` registra richiesta e risposta di
OGNI chiamata (verifica, allegati, soggetto, deposito, firma, riconciliazione). Per rispondere
alla domanda che si pone davvero in diagnosi — «che cosa abbiamo mandato per QUESTO atto?» —
bisogna scorrere l'audit, filtrare la fase di deposito e prendere l'ultima riga.

Tre ragioni per tenerne una copia sull'atto, e la terza è quella che decide:
  · **trovabilità** — una colonna invece di una ricerca sull'audit;
  · **conservazione** — l'audit conserva payload integrali che sono dati personali e in parte
    particolari, e il documento stesso prevede di minimizzarli: se l'audit si purga, sparisce
    anche la prova di ciò che è stato depositato;
  · ⚠️ **ricostruibilità** — il payload si potrebbe rifare da atto SIPO più configurazione, ma
    SOLO finché la configurazione non cambia. Dopo l'attivazione di una nuova baseline non è
    più ricostruibile, e l'atto è già formato. `ID_VERSIONE` è già sull'atto: insieme al
    payload dice non solo che cosa è stato inviato, ma con quale configurazione è stato
    prodotto.

Costo dichiarato: 144.285 atti l'anno secondo la ricognizione dei tipi atto (52.142 decessi,
42.239 nascite, 26.226 cittadinanze, 23.111 matrimoni, 567 unioni) — fra 1 e 4 GB l'anno
secondo la dimensione media del payload.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.16.docx')

COLONNA = ('TXT_PAYLOAD', 'CLOB (IS JSON)',
           'Il modello evento come è stato costruito e inviato al deposito: la copia esatta di '
           'ciò che ANSC ha ricevuto per questo atto. ⚠️ Si conserva qui e non solo nell’audit '
           'perché, dopo l’attivazione di una nuova baseline della configurazione, il payload '
           'non è più ricostruibile dall’atto: letto insieme a ID_VERSIONE dice che cosa è '
           'stato inviato e con quale configurazione è stato prodotto. Su Oracle 12.2 è un '
           'CLOB con vincolo «IS JSON»; il tipo JSON nativo non è disponibile.')

PROSA = [
    'La conservazione del payload sull’atto risponde a due esigenze pratiche che l’audit da '
    'solo non copre. La prima è di diagnosi: davanti a un atto rifiutato la domanda è sempre '
    '«che cosa abbiamo mandato», e con il payload sulla riga dell’atto la risposta è immediata '
    'invece di richiedere una ricerca fra le chiamate registrate. La seconda è di prova: '
    'l’audit conserva i payload integrali di ogni chiamata, che sono dati personali e in parte '
    'particolari, e per questo il presente documento ne prevede la minimizzazione e una '
    'conservazione limitata nel tempo; la copia sull’atto sopravvive a quella pulizia perché '
    'ne è il documento essenziale, uno per atto e non uno per chiamata.',

    '⚠️ La ragione che rende la copia necessaria e non comoda è però un’altra: il payload non è '
    'sempre ricostruibile. Finché la configurazione resta la stessa, si potrebbe rigenerare '
    'dall’atto di SIPO applicando le regole; ma la configurazione cambia — il mapping di ANSC '
    'è rivisto una ventina di volte l’anno — e dopo l’attivazione di una nuova baseline la '
    'rigenerazione produrrebbe un payload diverso da quello effettivamente depositato. L’atto, '
    'nel frattempo, è già formato e non si tocca. Il payload conservato, letto insieme alla '
    'baseline registrata sulla stessa riga, è quindi l’unico modo per sapere a posteriori che '
    'cosa il Comune abbia dichiarato e su quali regole.',

    'Il costo è misurabile e va dichiarato. La ricognizione dei tipi atto del Comune conta '
    '144.285 atti l’anno per le aree censite — 52.142 decessi, 42.239 nascite, 26.226 '
    'cittadinanze, 23.111 matrimoni, 567 unioni — che a una dimensione media fra otto e trenta '
    'kilobyte per payload valgono fra uno e quattro gigabyte l’anno. È un volume gestibile con '
    'una politica di conservazione dichiarata, ed è la ragione per cui si conserva il payload '
    'del deposito e non quello di ogni tentativo.',
]


def applica(doc):
    fatti = []
    t = D.tabella_colonne(doc, 'ID_STATO_ATTO')
    esistenti = [r.cells[0].text.strip() for r in t.rows]
    if 'TXT_PAYLOAD' not in esistenti:
        righe = [[c.text for c in r.cells] for r in t.rows[1:]]
        fuori = []
        for v in righe:
            fuori.append(v)
            if v[0].startswith('ID_ANSC'):
                fuori.append(list(COLONNA))
        D.riscrivi_sicura(t, fuori)
        fatti.append('ANSC_STATO_ATTO: aggiunta TXT_PAYLOAD dopo ID_ANSC')

    # il DDL
    riga = next((p for p in doc.paragraphs
                 if p.text.strip().startswith('ID_ANSC') and 'VARCHAR2(50 CHAR)' in p.text), None)
    if riga is not None:
        D.riga_dopo(doc, riga, '  TXT_PAYLOAD          CLOB,')
        fine = next((p for p in doc.paragraphs
                     if p.text.strip().startswith('COMMENT ON COLUMN ANSC_USR.ANSC_STATO_ATTO')),
                    None)
        if fine is not None:
            D.riga_dopo(doc, fine, "  CONSTRAINT CK_ANSC_STATO_ATTO_PAYLOAD CHECK (TXT_PAYLOAD IS JSON),")
        fatti.append('DDL di ANSC_STATO_ATTO: aggiunta TXT_PAYLOAD')

    # la prosa, dopo il paragrafo sugli indici della tabella
    ancora = next((p for p in doc.paragraphs
                   if p.text.strip().startswith('Indici: chiave primaria ID_STATO_ATTO')), None)
    if ancora is not None:
        i = D.indice_di(doc, ancora)
        dopo = doc.paragraphs[i + 1]._p
        for testo in PROSA:
            D.para(doc, dopo, testo)
        fatti.append(f'prosa sulla conservazione del payload ({len(PROSA)} paragrafi)')
    return fatti


if __name__ == '__main__':
    doc = docx.Document(DOC)
    for f in applica(doc):
        print('  ·', f)
    doc.save(DOC)
    print('salvato')
