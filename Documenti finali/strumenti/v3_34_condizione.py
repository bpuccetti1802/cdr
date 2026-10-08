# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.33 → v3.34 (07/10/2026).

Due interventi chiesti dal committente:
  1. RICONCILIAZ_DIZIONARI perde la colonna CONDIZIONE: si è deciso di non usarla;
  2. ANSC_CFG_DOMINIO_SIPO si popola a mano e non avrà una pagina di back-office: va
     detto nel documento, perché il criterio di completezza del disegno del back-office
     è che ogni tabella abbia una pagina che la governi.

⚠️ Nel disegno del back-office esistono TRE «Condizione» e solo due vanno tolte: quella
della pagina degli allegati appartiene ad ALLEGATI_USECASE («testo leggibile della logica
che rende il documento obbligatorio») e resta. Qui si tocca la sola analisi; il back-office
si allinea nella sua prossima versione, insieme alle schermate dell'OTP già attese.

⚠️ Corretta anche una svista della v3.33: nell'ERD la scheda della riconciliazione non
aveva ricevuto ID_DOMINIO_SIPO, aggiunta al DDL ma non alla figura.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_34_condizione.py
"""
import os
import shutil
import subprocess
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.33.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.34.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')


def main():
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)

    def commenti(p):
        return int(subprocess.run(
            ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                           "grep -o '<w:comment ' | wc -l" % repr(p)],
            capture_output=True, text=True).stdout.strip() or 0)
    prima = commenti(DST)

    d = docx.Document(DST)
    fatti = []

    # ------------------------------------------------ 1. via CONDIZIONE
    ps = d.paragraphs
    riga = next(p for p in ps
                if p.text.strip() == 'CONDIZIONE           VARCHAR2(1000 CHAR),')
    if D.ha_commenti(riga._p):
        raise SystemExit('la riga CONDIZIONE porta un commento: non si elimina')
    riga._p.getparent().remove(riga._p)
    fatti.append('colonna CONDIZIONE rimossa dal DDL di RICONCILIAZ_DIZIONARI')

    # ⚠️ le tabelle «Colonna | Tipo | Note» sono molte: trova_tabella prenderebbe la
    # prima. Si individua quella giusta dalla colonna che solo lei contiene.
    t37 = D.tabella_colonne(d, 'ID_RICONCILIAZIONE')
    for r in list(t37.rows):
        if r.cells[0].text.strip() == 'CONDIZIONE':
            if any(D.ha_commenti(c._tc) for c in r.cells):
                raise SystemExit('la riga CONDIZIONE della scheda porta un commento')
            r._tr.getparent().remove(r._tr)
            fatti.append('riga CONDIZIONE rimossa dalla scheda delle colonne')
            break
    else:
        raise SystemExit('riga CONDIZIONE non trovata nella scheda')

    D.para(d, D.h(d, 3, 'RICONCILIAZ_DIZIONARI')._p.getnext(),
           '⚠️ **Dalla v3.34 la colonna CONDIZIONE non c’è più**: era prevista per i casi '
           'in cui la traduzione dipende dal contesto e non dal solo valore, e si è deciso '
           'di non usarla. Quei casi, se emergeranno, si rappresentano con una riga per '
           'ciascun contesto — il che è già possibile, perché il campo di partenza entra '
           'nella chiave.')
    fatti.append('dichiarato nel documento perché la colonna è stata tolta')

    # ------------------------------------------------ 2. la nota sul popolamento
    anc = D.h(d, 3, 'ANSC_CFG_DOMINIO_SIPO')
    nxt = anc._p.getnext()
    D.para(d, nxt,
           '⚠️ **Questa tabella si popola a mano e non ha una pagina di back-office.** È '
           'una scelta, e conviene dichiararne la ragione: le righe sono poche — una per '
           'ciascuna corrispondenza fra un dominio di ANSC e la colonna di SIPO che ne '
           'porta il valore — cambiano di rado, e si scrivono una volta sola quando si '
           'configura una famiglia di atti. Costruire una pagina per governarle costerebbe '
           'più di quanto farebbe risparmiare.')
    D.para(d, nxt,
           '⚠️ **Ne discende un’eccezione da ricordare quando si legge il disegno del '
           'back-office**, il cui criterio di completezza è che ciascuna tabella abbia '
           'almeno una pagina che la governi: questa non ce l’ha, e non per dimenticanza. '
           'Il prezzo è che un refuso nel nome di una tabella o di una colonna non viene '
           'intercettato da alcuna maschera: si manifesta come una riconciliazione che non '
           'trova il proprio catalogo. Conviene quindi che il carico sia accompagnato da un '
           'controllo che verifichi l’esistenza reale degli oggetti nominati — una query '
           'sul dizionario dei dati, non una maschera.')
    fatti.append('nota sul popolamento manuale e sull’assenza di pagina di back-office')

    # ------------------------------------------------ 3. l'ERD
    subprocess.run([sys.executable, os.path.join(os.path.dirname(
        os.path.abspath(__file__)), 'diagrammi_erd.py')], check=True,
        cwd=os.path.dirname(os.path.abspath(__file__)))
    D.sostituisci_immagine(d, 'Schema ANSC_USR', os.path.join(IMG, 'erd_ansc_usr.png'))
    fatti.append('ERD rigenerato: via CONDIZIONE, e ID_DOMINIO_SIPO che la v3.33 '
                 'aveva messo nel DDL ma non nella figura')

    # ------------------------------------------------ 4. testata e storia
    for tb in d.tables:
        if tb.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in tb.rows:
                v = {'Versione': '3.34',
                     'Documento': 'ANALISI_Integrazione-ANSC_v3.34'}.get(
                        r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    D.storia(d, '07/10/2026', '3.34',
             'Modello dati (RICONCILIAZ_DIZIONARI, ANSC_CFG_DOMINIO_SIPO, ERD) · '
             'Appendice DDL',
             'Rimossa da RICONCILIAZ_DIZIONARI la colonna CONDIZIONE, che si è deciso di '
             'non usare: i casi in cui la traduzione dipende dal contesto si '
             'rappresenteranno con una riga per contesto, poiché il campo di partenza è '
             'già parte della chiave. Dichiarato inoltre che il catalogo delle '
             'corrispondenze si popola a mano e non avrà una pagina di back-office, con la '
             'ragione della scelta e il prezzo che comporta: nessuna maschera intercetta un '
             'refuso nel nome di una tabella, quindi il carico va accompagnato da un '
             'controllo sul dizionario dei dati. Allineata la figura dello schema, cui la '
             'versione precedente non aveva aggiunto la colonna di raccordo.')
    fatti.append('testata e storia aggiornate')

    # ------------------------------------------------ 5. grassetti
    def grassetti(p):
        if '**' not in p.text or D.ha_commenti(p._p):
            return 0
        pezzi = D.segmenta('', p.text)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        for testo, gr in pezzi:
            run = p.add_run(testo)
            run.bold = gr
        return 1
    n = sum(grassetti(p) for p in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    n += grassetti(p)
    fatti.append('%d paragrafi con grassetto applicato' % n)

    d.save(DST)
    dopo = commenti(DST)
    if dopo != prima:
        raise SystemExit('COMMENTI PERSI: erano %d, sono %d' % (prima, dopo))
    fatti.append('commenti di Word conservati: %d → %d' % (prima, dopo))

    print('\n'.join(' · ' + f for f in fatti))
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))
    print('scritto:', os.path.relpath(DST, BASE))


if __name__ == '__main__':
    main()
