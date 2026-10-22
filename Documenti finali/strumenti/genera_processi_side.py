# -*- coding: utf-8 -*-
"""Due documenti di processo di Side (Comune di Milano) adattati a Roma, 22/10/2026.

  PROCESSO_Numerazione-Comunale_SIPO-ANSC_v0.1.docx
      da «CdMI SIDe MI_ANSC - Processo di gestione della numerazione comunale v1.0»; le pagine
      stanno nel back-office dell'integrazione; serie unica per anno, sei cifre (analisi v3.41+).
  PROCESSO_Attestato-Conformita_SIPO-ANSC_v0.1.docx
      da «CdMI SIDe MI_ANSC - Processo di gestione e firma dell'attestato di conformità v1.1»;
      la pagina sta nel percorso di SIPO (documento del workflow v0.3).

Entrambi conservano la struttura del documento di Milano e dichiarano gli scostamenti in un
capitolo apposito. Si costruiscono dal template DAD (svuotamento dopo l'elemento 29, come
genera_workflow_interfacce.py). Le tabelle dei servizi R007, R010 e del modello evento si
leggono dai contratti di ANSC, quelle del servizio della numerazione da ansc-numerazione-v1.yaml.
Figure: schermate_processi_side.py.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/genera_processi_side.py
"""
import os
import shutil
import sys

import docx
import yaml
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D                 # noqa: E402
import schermate_processi_side as F     # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
IMG = os.path.join(FIN, 'strumenti', 'img')
OPENAPI = os.path.join(BASE, 'ansc', 'docs', 'openapi')
TEMPLATE = os.path.join(BASE, 'Template documentale', 'template_DAD_roma-capitale.docx')
VERSIONE = '0.1'
DATA = '22/10/2026'
TAGLIA_DA = 29
SIDE_NUM = 'CdMI SIDe MI_ANSC - Processo di gestione della numerazione comunale v1.0.docx'
SIDE_ATT = ('CdMI SIDe MI_ANSC - Processo di gestione e firma dell’attestato di conformità '
            'v1.1.docx')


# ------------------------------------------------------------------------------ primitive
def testo(p, s):
    if p.runs:
        p.runs[0].text = s
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(s)


def h(d, livello, t):
    return d.add_paragraph(t, style=f'Heading {livello}')


def par(d, t, stile='Normal', corsivo=False):
    p = d.add_paragraph(style=stile)
    for pezzo, grassetto in D.segmenta('', t):
        r = p.add_run(pezzo)
        r.bold = grassetto
        if corsivo:
            r.italic = True
            r.font.size = Pt(9)
    return p


def voce(d, t):
    return par(d, t, 'List Paragraph')


def tabella(d, righe, larghezze):
    t = d.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(righe):
        for j, v in enumerate(r):
            run = t.cell(i, j).paragraphs[0].add_run(str(v))
            run.bold = i == 0
            run.font.size = Pt(9)
    for j, w in enumerate(larghezze):
        for r in t.rows:
            r.cells[j].width = Inches(w)
    griglia = t._tbl.find(qn('w:tblGrid'))
    for col, w in zip(griglia.findall(qn('w:gridCol')), larghezze):
        col.set(qn('w:w'), str(int(round(Inches(w).twips))))
    # larghezze fisse: senza, Word allarga le colonne con i percorsi lunghi
    pr = t._tbl.tblPr
    layout = pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        pr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    d.add_paragraph()
    return t


def figura(d, nome, larghezza, didascalia):
    d.add_picture(os.path.join(IMG, nome), width=Inches(larghezza))
    d.paragraphs[-1].alignment = 1
    par(d, didascalia, corsivo=True)


def costruisci(out, sottotitolo):
    shutil.copyfile(TEMPLATE, out)
    d = docx.Document(out)
    body = d.element.body
    for ch in list(body.iterchildren())[TAGLIA_DA:]:
        if ch.tag != qn('w:sectPr'):
            body.remove(ch)
    for p in d.paragraphs:
        if '“Lorem Ipsum”' in p.text:
            testo(p, '“Integrazione SIPO – ANSC”')
        elif p.style.name == 'Subtitle':
            testo(p, sottotitolo)
    return d


def chiudi(d, out, nome, sintesi):
    D.frontespizio(d, VERSIONE, nome, data=DATA)
    st = D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')
    for cella, v in zip(st.rows[1].cells, [DATA, VERSIONE, 'Tutti', sintesi]):
        D.riscrivi_cella(cella, v)
    for riga in list(st.rows[2:]):
        if not ''.join(c.text for c in riga.cells).strip():
            st._tbl.remove(riga._tr)
    d.save(out)
    print('  scritto', out, D.riepilogo(out))


def schema(file_yaml, nome):
    """Righe «Campo | Tipo | Esempio | Descrizione» di uno schema di un contratto ANSC."""
    y = yaml.safe_load(open(os.path.join(OPENAPI, file_yaml)))
    s = y['components']['schemas'][nome]
    props = dict(s.get('properties', {}))
    for a in s.get('allOf', []):
        props.update(a.get('properties', {}))
    righe = [['Campo', 'Tipo', 'Esempio', 'Descrizione']]
    for n, v in props.items():
        tipo = v.get('type') or v.get('$ref', '').split('/')[-1]
        if tipo == 'array':
            tipo = 'array di ' + v.get('items', {}).get('$ref', '').split('/')[-1]
        righe.append([n, tipo, str(v.get('example', '')),
                      ' '.join((v.get('description') or '').split())])
    return righe, y['info'].get('version')


# =========================================================================== NUMERAZIONE
NUM_ADATTAMENTI = [
    ['Elemento', 'In Side', 'A Roma', 'Perché'],
    ['Dove stanno le pagine', 'Nel sottomenu di ciascun registro di stato civile.',
     'Nel back-office dell’integrazione ANSC: una mattonella «Numerazione comunale» con le '
     'due voci.', 'La serie è una sola per tutti i registri: due pagine bastano. La gestione '
     'della numerazione è governo della serie, non redazione dell’atto.'],
    ['La serie', 'Per registro e anno, da un numeratore del gestionale.',
     'Una sequence Oracle per anno, SEQ_NUM_COMUNALE_<anno>, unica per tutti i registri, sei '
     'cifre al massimo.', 'Assunto del Comune (analisi, «La numerazione comunale»); le fonti '
     'ANSC dichiarano cinque caratteri (NC-1).'],
    ['Quali numeri passano di qui', 'Solo quelli degli atti da redigere sulla web app.',
     'Tutti: anche quelli degli atti formati in SIPO, che il concentratore stacca '
     'all’anteprima.', 'La serie è verificabile in un luogo solo e un buco è spiegabile.'],
    ['Registro della richiesta (ty_evento)', 'Campo della richiesta e del filtro.',
     'Tolto da tabelle, pagine e filtri.', 'Con la serie unica non ha più funzione; il '
     'registro dell’atto resta sull’atto di SIPO.'],
    ['Id della richiesta', 'Gestore del numeratore da creare.', 'Colonna IDENTITY.',
     'Lo offre il database.'],
    ['Anno sui numeri', 'Solo sulla richiesta.', 'Ripetuto su DATI_ATTO_ESTERNO.',
     'Permette il vincolo di unicità su anno e numero.'],
    ['Legame con l’atto', 'id_evento, verso la tabella degli eventi di Side.', 'id_atto_sipo.',
     'SIPO non ha la tabella degli eventi.'],
    ['Segna come utilizzato', 'Cambio di stato «ad uso e consumo dell’operatore», senza vincoli.',
     'Chiede l’identificativo nazionale dell’atto; Scartato ed Evento acquisito sono finali.',
     'Un numero Utilizzato deve poter essere ricondotto all’atto che lo porta.'],
    ['SIDe non disponibile (scenari 4 e 5)', 'Al ritorno si richiede un numero, che '
     '«coinciderà» con quello indicato dalla web app.', 'Al ritorno si registra il numero '
     'effettivamente usato, con l’identificativo nazionale (registrazione a posteriori).',
     'La coincidenza dipende dall’ordine delle richieste; la registrazione no.'],
    ['Anno della richiesta', 'Modificabile liberamente.', 'Il corrente o il precedente.',
     'Il precedente serve a cavallo d’anno; altri anni non hanno uno scenario.'],
    ['Numeri consecutivi', 'Richiesti per i gemelli, senza dire come garantirli.',
     'Stacchi serializzati con un blocco esclusivo sulla tabella dei numeri.',
     'Una sequence da sola non garantisce la contiguità di un blocco.'],
    ['Buchi della serie', 'Non trattati.', 'Mostrati dalla pagina di ricerca, serie per serie.',
     'Un valore preso dalla sequence in una transazione fallita non torna indietro.'],
    ['Cornice grafica', 'Comune di Milano.', 'Roma Capitale, come il resto del back-office.',
     'È l’unico intervento solo grafico.'],
]

NUM_PUNTI = [
    ['#', 'Questione', 'Rimando'],
    ['NC-1', 'ANSC dichiara per il numero dell’atto cinque caratteri al massimo; l’assunto del '
     'Comune è di sei cifre. Va confermato che ANSC accetti un numero di sei cifre.', 'OP-73 a'],
    ['NC-2', 'Che ANSC accetti un progressivo comunale unico fra i registri: le fonti non lo '
     'dichiarano né lo escludono.', 'OP-73'],
    ['NC-3', 'A quale anno appartiene il numero di un atto staccato il 31 dicembre e formato il '
     '1° gennaio: la regola proposta è scartarlo e staccarne uno della serie nuova.', 'OP-73'],
    ['NC-4', 'Il privilegio CREATE SEQUENCE all’utenza che esegue il job e l’utenza del '
     'servizio della numerazione.', 'OP-73 d'],
    ['NC-5', 'Un tetto al numero di elementi di una richiesta: la colonna ne ammette fino a '
     '999, un parto plurimo ne chiede pochi. [DA VERIFICARE: il tetto pratico.]', '—'],
    ['NC-6', 'Chi può segnare come Scartato un numero staccato da un altro operatore: il '
     'contratto riserva i cambi di stato al ruolo USC, senza distinguere l’autore della '
     'richiesta.', '—'],
    ['NC-7', 'Il disegno del back-office v0.10 descrive ancora le due pagine con il registro: '
     'va riallineato a questo documento.', '—'],
    ['NC-8', 'Se la registrazione a posteriori richieda un profilo distinto da quello '
     'dell’ufficiale: può far avanzare la serie e lasciarvi buchi.', 'BO-30'],
]


def numerazione():
    nome = f'PROCESSO_Numerazione-Comunale_SIPO-ANSC_v{VERSIONE}'
    out = os.path.join(FIN, nome + '.docx')
    d = costruisci(out, 'Processo di gestione della numerazione comunale')

    h(d, 1, 'Introduzione')
    par(d, 'Il documento adatta a Roma il processo di gestione della numerazione comunale del '
           'sistema Side del Comune di Milano [R1]. In Side il processo serve agli atti da '
           'redigere sulla web app di ANSC, il cui numero comunale si chiede al gestionale '
           'perché nessuno lo assegni due volte. A Roma lo stesso processo governa tutti i '
           'numeri comunali — anche quelli degli atti formati in SIPO, che il concentratore '
           'stacca all’anteprima — e le sue pagine stanno nel back-office dell’integrazione '
           'ANSC [R3].')
    par(d, 'La struttura è quella del documento di origine: attori, processo e scenari, le due '
           'pagine, lo sviluppo propedeutico, i quesiti. Dove Roma si discosta, il capitolo '
           '«Che cosa cambia rispetto a Side» lo dice con la ragione. Le schermate sono '
           'ridisegnate nella cornice istituzionale di Roma Capitale del back-office; i valori '
           'sono di esempio.')
    h(d, 2, 'Riferimenti')
    tabella(d, [['#', 'Documento', 'Che cosa se ne prende'],
                ['[R1]', 'Sorgenti Documentali/' + SIDE_NUM, 'Il processo di origine: attori, '
                 'scenari, pagine, tabelle.'],
                ['[R2]', 'ANALISI_Integrazione-ANSC_v3.44.docx, «La numerazione comunale»',
                 'La serie per anno, le garanzie, le strutture ereditate, i percorsi, il '
                 'servizio num-ansc-sipo, lo stacco passo per passo.'],
                ['[R3]', 'DISEGNO_Back-Office_ANSC_v0.10.docx', 'La cornice, il menu e i ruoli '
                 'del back-office.'],
                ['[R4]', 'ansc-numerazione-v1.yaml (1.1.0) e SPEC_API_Integrazione-ANSC_v0.3',
                 'Le dieci operazioni del servizio della numerazione e i loro ruoli.'],
                ['[R5]', 'WORKFLOW_Interfacce-SIPO-ANSC_v0.3.docx', 'Il punto del percorso '
                 'di SIPO in cui il numero si stacca: l’anteprima.']], [0.6, 2.7, 3.2])
    h(d, 2, 'Glossario')
    tabella(d, [['Termine', 'Significato'],
                ['Numero comunale', 'Il progressivo che il Comune assegna all’atto; viaggia nel '
                 'modello evento (evento.numeroatto) e compone l’identificativo nazionale.'],
                ['Serie', 'I numeri di un anno: una sequence Oracle, SEQ_NUM_COMUNALE_<anno>.'],
                ['Stacco', 'Il prelievo di uno o più numeri dalla serie, registrati come '
                 'Disponibili insieme alla richiesta che li ha prodotti.'],
                ['Buco', 'Un valore preso dalla sequence e mai registrato: non ha una riga.'],
                ['Identificativo nazionale (idAnsc)', 'Anno, numero nazionale, numero comunale, '
                 'codice ISTAT, assegnato da ANSC alla validazione.'],
                ['num-ansc-sipo', 'Il servizio della numerazione: l’unico componente che stacca '
                 'numeri. Non parla con ANSC.']], [1.9, 4.6])

    h(d, 1, 'Attori')
    tabella(d, [['Nome', 'Descrizione', 'Attore di sistema'],
                ['Ufficiale di stato civile', 'Abilitato al back-office e ad ANSC. Chiede i '
                 'numeri per gli atti da formare sulla web app, ne aggiorna lo stato, registra '
                 'a posteriori quelli usati mentre SIPO non era disponibile. Ruolo USC.', 'No'],
                ['Amministratore del back-office', 'Crea la serie di un anno quando il job non '
                 'l’ha creata; consulta serie e buchi.', 'No'],
                ['Supporto, Auditor', 'Consultano richieste, numeri e serie, senza cambiarli.',
                 'No'],
                ['Concentratore operativo (all-ansc-sipo)', 'Stacca il numero di un atto formato '
                 'in SIPO all’anteprima, lo porta a Utilizzato al deposito riuscito o a '
                 'Scartato se l’atto è abbandonato. Agisce con il token dell’operatore.', 'Sì'],
                ['Servizio della numerazione (num-ansc-sipo)', 'Governa la serie: stacchi '
                 'serializzati, stati, registrazioni a posteriori, lettura dei buchi.', 'Sì'],
                ['Job JOB_CREA_SEQ_NUM_COMUNALE', 'Il 1° dicembre crea la sequence dell’anno '
                 'successivo.', 'Sì']], [2.0, 3.6, 0.9])

    h(d, 1, 'Processo')
    par(d, 'In Side l’operatore che forma un atto sulla web app, al passo della validazione, '
           'deve indicare il numero comunale. La web app suggerisce l’ultimo numero registrato in '
           'ANSC per l’anno, ma perché lo stesso numero non vada a due atti la numerazione la '
           'controlla il gestionale. Roma adotta lo stesso principio e lo estende: la serie è una '
           'sola, ogni numero passa dal servizio della numerazione e si registra nella stessa '
           'tabella, qualunque sia la superficie su cui l’atto si forma [R2].')
    par(d, 'Le pagine sono due, come in Side: la richiesta e la ricerca e consultazione. Stanno '
           'nel back-office e si raggiungono dalla mattonella «Numerazione comunale» del menu '
           'dell’integrazione. Side valuta di sostituire le voci ripetute per registro con una '
           'voce sola e una pagina di scelta; con una serie unica la soluzione si semplifica '
           'ancora: una mattonella, due voci.')
    figura(d, 'pr_num_menu.png', 6.3, 'Il menu del back-office: la mattonella «Numerazione '
           'comunale» con le due pagine.')
    h(d, 2, 'Scenari')
    par(d, 'I primi sei sono quelli di Side, adattati; gli ultimi tre sono propri di Roma, '
           'perché qui passano dalla serie anche gli atti formati in SIPO.')
    scenari = [
        ('Scenario 1 – Atto formato sulla web app: uso ordinario', [
            'Completo l’atto sulla web app e mi fermo alla validazione, che chiede il numero '
            'comunale.',
            'Nella pagina di richiesta del back-office stacco un numero dell’anno corrente.',
            'Il numero compare nella tabella, nello stato Disponibile.',
            'Inserisco il numero nella web app e valido l’atto.',
            'Segno il numero come Utilizzato, indicando l’identificativo nazionale che ANSC ha '
            'assegnato: lo stato passa da 1 Disponibile a 2 Utilizzato.']),
        ('Scenario 2 – Numeri consecutivi per gli atti dei gemelli', [
            'Prima di cominciare stacco, con una sola richiesta, tanti numeri quanti i nati.',
            'I numeri sono consecutivi: gli stacchi sono serializzati e nessuna richiesta '
            'contemporanea può intercalarsi.',
            'Li uso sulla web app, oppure in SIPO, dove la prenotazione degli atti dei gemelli '
            '(R009 /prenota/evento) li richiede tutti insieme.',
            'Ciascun numero segue poi la sorte del proprio atto.']),
        ('Scenario 3 – Numero non utilizzato', [
            'Stacco un numero.',
            'Non mi serve più: l’atto non si conclude o viene annullato.',
            'Segno il numero come Scartato, dopo una conferma: lo stato passa da 1 Disponibile '
            'a 3 Scartato ed è finale.']),
        ('Scenario 4 – SIPO non disponibile: il numero suggerito da ANSC', [
            'Formo l’atto sulla web app con il numero suggerito da ANSC.',
            'Al ritorno di SIPO lo registro dal pannello «Registrazione a posteriori» della '
            'pagina di richiesta, con l’identificativo nazionale dell’atto.',
            'Il numero entra nella serie come Utilizzato. Se è uguale o superiore al prossimo '
            'valore della sequence, la serie avanza fino a lì; se è inferiore e libero, colma '
            'un buco; se è già registrato, la registrazione è respinta.']),
        ('Scenario 5 – SIPO non disponibile a cavallo d’anno', [
            'Come lo scenario 4, indicando l’anno precedente: l’anno ammesso è il corrente o il '
            'precedente.']),
        ('Scenario 6 – Consultazione delle richieste', [
            'Apro la pagina di ricerca e consultazione.',
            'Imposto i filtri: le due date sono obbligatorie.',
            'Leggo le richieste e i numeri in tabella, con le serie in testa.',
            'Sui numeri Disponibili posso cambiare lo stato; su un Utilizzato posso legare '
            'l’atto di SIPO.']),
        ('Scenario 7 – Atto formato in SIPO', [
            'All’anteprima il concentratore stacca un numero legato all’atto, senza che '
            'l’ufficiale apra il back-office.',
            'Se la validazione R009 riesce, il numero passa a Utilizzato con l’identificativo '
            'nazionale.',
            'Se la validazione è respinta, il numero resta all’atto e si riusa al tentativo '
            'successivo: la serie non ne consuma un altro.',
            'Se l’atto è abbandonato prima del deposito, il numero passa a Scartato.',
            'Le righe sono visibili nella pagina di ricerca, con l’atto di SIPO valorizzato.']),
        ('Scenario 8 – Aggancio di un atto formato sulla web app', [
            'Un atto formato sulla web app viene agganciato in SIPO.',
            'Il numero, già Utilizzato, si lega all’atto di SIPO e passa a 4 Evento acquisito, '
            'che è finale. È ciò che Side annuncia come «futuro allineamento».']),
        ('Scenario 9 – Governo delle serie', [
            'Il 1° dicembre il job crea la sequence dell’anno successivo.',
            'Se manca, lo stacco fallisce in modo visibile; l’amministratore la crea con CREA '
            'SERIE.',
            'Nella pagina di ricerca la tabella delle serie mostra, anno per anno, il prossimo '
            'numero, i conteggi per stato e i buchi.']),
    ]
    for titolo, passi in scenari:
        h(d, 3, titolo)
        for p in passi:
            voce(d, p)
    h(d, 2, 'I percorsi e le operazioni del servizio')
    tabella(d, [['Percorso', 'Chi chiama', 'Operazioni [R4]'],
                ['Atto formato in SIPO', 'Il concentratore, con il token dell’operatore',
                 'creaPrenotazione con l’atto; aggiornaNumero a Utilizzato al deposito riuscito, '
                 'a Scartato se l’atto è abbandonato.'],
                ['Atto formato sulla web app', 'L’ufficiale, dalla pagina di richiesta',
                 'creaPrenotazione; aggiornaNumero a Utilizzato o a Scartato; a Evento acquisito '
                 'all’aggancio.'],
                ['Parto plurimo', 'L’ufficiale, prima della prenotazione R009',
                 'creaPrenotazione con tanti numeri quanti i nati.'],
                ['SIPO non disponibile', 'L’ufficiale, al ritorno di SIPO', 'creaNumero.'],
                ['Governo delle serie', 'Amministratore, Supporto, Auditor', 'elencaSerie, '
                 'leggiSerie, creaSerie; elencaPrenotazioni, leggiPrenotazione, elencaNumeri, '
                 'leggiNumero.']], [1.6, 1.9, 3.0])

    h(d, 1, 'Pagina della richiesta')
    par(d, 'Quando l’operatore apre la pagina, l’anno è il corrente e il numero di elementi è '
           'uno. Anno e numero di elementi sono obbligatori; l’atto di SIPO è facoltativo e '
           'ammesso solo con un numero, per legare subito il numero a un atto. Non c’è il '
           'registro.')
    figura(d, 'pr_num_richiesta.png', 6.3, 'Numerazione comunale — richiesta. Tre numeri '
           'consecutivi per un parto trigemino; sotto, la registrazione a posteriori.')
    voce(d, '**RICHIEDI.** Chiama creaPrenotazione: il servizio inserisce una riga in '
            'PRENOTAZIONE_ATTI per la richiesta e una in DATI_ATTO_ESTERNO per ogni numero '
            'staccato, nello stato 1 Disponibile e senza identificativo nazionale.')
    voce(d, '**Anno.** Il corrente o il precedente, per gli scenari a cavallo d’anno.')
    voce(d, '**Più numeri.** Una richiesta può staccare più numeri consecutivi, senza buchi, '
            'per gli atti dei gemelli.')
    voce(d, '**Tabella.** I numeri staccati compaiono sotto, ordinati per richiesta e numero.')
    voce(d, '**Spunta — Segna come utilizzato.** Solo su un numero Disponibile. Apre una '
            'finestra che chiede l’identificativo nazionale, obbligatorio, e l’atto di SIPO, '
            'facoltativo; chiama aggiornaNumero.')
    voce(d, '**Croce — Segna come scartato.** Solo su un numero Disponibile, dopo una conferma; '
            'chiama aggiornaNumero. Scartato è finale.')
    voce(d, '**Registrazione a posteriori.** Anno, numero usato e identificativo nazionale; '
            'REGISTRA chiama creaNumero.')
    figura(d, 'pr_num_utilizzato.png', 4.6, 'La finestra «Segna come utilizzato».')

    h(d, 1, 'Pagina di ricerca e consultazione')
    par(d, 'Quando l’operatore apre la pagina, le date sono quella del giorno e l’operatore è '
           'quello collegato; solo le due date sono obbligatorie, come in Side. Al posto del '
           'registro ci sono lo stato, il numero e l’atto di SIPO. In testa, la tabella delle '
           'serie dice per ciascun anno il prossimo numero, i conteggi per stato e i buchi.')
    figura(d, 'pr_num_ricerca.png', 6.3, 'Numerazione comunale — ricerca e consultazione.')
    tabella(d, [['Elemento', 'Che cosa fa', 'Operazione [R4]'],
                ['Serie della numerazione', 'Una riga per anno; «vedi» sui buchi li elenca.',
                 'elencaSerie, leggiSerie'],
                ['CERCA', 'Richieste e numeri secondo i filtri, paginati, ordinati per '
                 'richiesta e numero.', 'elencaPrenotazioni, elencaNumeri'],
                ['Occhio', 'La richiesta che ha prodotto il numero e i suoi legami.',
                 'leggiNumero, leggiPrenotazione'],
                ['Spunta e croce', 'Come nella pagina di richiesta, sui soli Disponibili.',
                 'aggiornaNumero'],
                ['Matita', 'Su un Utilizzato senza atto di SIPO: lega l’atto e porta a Evento '
                 'acquisito.', 'aggiornaNumero'],
                ['CREA SERIE', 'Crea la sequence dell’anno corrente o del successivo. Solo '
                 'amministratore.', 'creaSerie']], [1.7, 3.3, 1.5])

    h(d, 1, 'Sviluppo propedeutico')
    par(d, 'Side chiede un gestore del numeratore dell’id di richiesta e due tabelle. A Roma il '
           'numeratore è una colonna IDENTITY, le tabelle conservano i nomi di Side in ANSC_USR '
           'e i numeri si prendono da una sequence per anno [R2].')
    h(d, 2, 'La tabella PRENOTAZIONE_ATTI')
    tabella(d, [['Campo', 'Tipo', 'Descrizione', 'Alimentazione'],
                ['id_prenotazione_atti', 'NUMBER (PK)', 'Id della richiesta', 'IDENTITY'],
                ['nr_anno_atto', 'NUMBER(4)', 'Anno della serie', 'Corrente o precedente'],
                ['nr_id_ansc_da_prenotare', 'NUMBER(3)', 'Quanti numeri si chiedono',
                 'Dalla richiesta'],
                ['cd_operatore_prenotazione', 'VARCHAR2(40)', 'Chi ha fatto la richiesta',
                 'Utenza collegata'],
                ['ts_prenotazione', 'TIMESTAMP', 'Quando', 'Istante della richiesta']],
            [1.9, 1.2, 1.8, 1.6])
    h(d, 2, 'La tabella DATI_ATTO_ESTERNO')
    tabella(d, [['Campo', 'Tipo', 'Descrizione', 'Alimentazione'],
                ['id_prenotazione_atti', 'NUMBER (PK, FK)', 'La richiesta', 'Chiave esterna'],
                ['nr_atto_comunale', 'NUMBER(6) (PK)', 'Il numero comunale',
                 'Dalla sequence dell’anno'],
                ['nr_anno_atto', 'NUMBER(4)', 'Anno, ripetuto dalla richiesta',
                 'Vincolo di unicità con il numero'],
                ['cd_stato_numero_prenotato', 'CHAR(1)', '1 Disponibile · 2 Utilizzato · '
                 '3 Scartato · 4 Evento acquisito', 'Alla creazione 1'],
                ['id_ansc_collegato', 'VARCHAR2(50)', 'Identificativo nazionale dell’atto',
                 'Segna come utilizzato, registrazione a posteriori, deposito riuscito'],
                ['id_atto_sipo', 'NUMBER', 'L’atto di SIPO (in Side id_evento)',
                 'Stacco del concentratore, matita, finestra'],
                ['ts_allineamento_ansc_side / cd_operatore_allineamento',
                 'TIMESTAMP / VARCHAR2(40)', 'Quando e da chi il numero è stato ricongiunto '
                 'all’atto', 'Con il legame all’atto']], [1.9, 1.3, 1.8, 1.5])
    h(d, 2, 'La serie')
    par(d, 'Una sequence per anno, creata il 1° dicembre per l’anno successivo da '
           'JOB_CREA_SEQ_NUM_COMUNALE, che esegue PRC_CREA_SEQ_NUM_COMUNALE; la stessa procedura '
           'si esegue all’installazione per l’anno in corso. Lo stacco non crea mai una '
           'sequence: se manca, fallisce.')
    for r in ('CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_2027',
              '  START WITH 1 INCREMENT BY 1 MAXVALUE 999999',
              '  NOCACHE ORDER NOCYCLE;'):
        p = par(d, r)
        p.runs[0].font.name = 'Courier New'
        p.runs[0].font.size = Pt(8)
        p.paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    h(d, 2, 'Il servizio num-ansc-sipo')
    par(d, 'Le pagine non scrivono sulle tabelle: chiamano il servizio della numerazione, che è '
           'l’unico punto in cui gli stacchi si serializzano. Lo stacco prende un blocco '
           'esclusivo su DATI_ATTO_ESTERNO, inserisce la richiesta, prende i valori dalla '
           'sequence, registra i numeri e conferma: tutto in una transazione [R2].')
    y = yaml.safe_load(open(os.path.join(FIN, 'ansc-numerazione-v1.yaml')))
    righe = [['Operazione', 'Metodo e percorso', 'Che cosa fa', 'Ruoli']]
    for p_, ops in y['paths'].items():
        for m, o in ops.items():
            if m in ('get', 'post', 'patch', 'put', 'delete'):
                righe.append([o['operationId'], f'{m.upper()} {p_}', o.get('summary', ''),
                              ', '.join(o.get('x-ruoli', []))])
    tabella(d, righe, [1.3, 2.4, 1.6, 1.2])

    h(d, 1, 'Che cosa cambia rispetto a Side')
    par(d, 'Ciò che non compare nella tabella si adotta come Side lo descrive.')
    tabella(d, NUM_ADATTAMENTI, [1.3, 1.6, 1.8, 1.8])

    h(d, 1, 'Quesiti e considerazioni')
    h(d, 2, 'L’allineamento con gli atti formati sulla web app')
    par(d, 'Side annuncia un processo futuro che aggiorni il gestionale con gli atti inseriti o '
           'conclusi sulla web app: cercare l’atto in ANSC, aggiornare stato, identificativo '
           'nazionale, evento, istante e utenza di allineamento, salvare l’evento ed '
           'eventualmente il modello intero, anche con un batch schedulato.')
    par(d, 'A Roma i pezzi ci sono già, e cambia il modo. Lo stato Evento acquisito e i legami '
           'del numero si scrivono con aggiornaNumero all’aggancio dell’atto; il modello evento '
           'intero sta su ANSC_STATO_ATTO.TXT_EVENTO; la lettura dell’atto in ANSC è R005. Il '
           'batch non si adotta: la riconciliazione con ANSC si fa su richiesta, non a '
           'calendario, come stabilito nell’analisi [R2].')
    h(d, 2, 'Punti aperti')
    tabella(d, NUM_PUNTI, [0.6, 4.9, 1.0])
    chiudi(d, out, nome, 'Prima emissione: il processo di Side adattato a Roma, con le pagine '
           'nel back-office e la serie unica per anno.')
    return out


# ============================================================================ CONFORMITÀ
ATT_ADATTAMENTI = [
    ['Elemento', 'In Side', 'A Roma', 'Perché'],
    ['La pagina', 'Pagina dinamica di Side, raggiunta nel percorso di firma.',
     'Schermata separata di SIPO, con il banner di arrivo dalla pagina precedente [R3].',
     'Le pagine nuove di SIPO si susseguono come schermate distinte.'],
    ['Chi chiama ANSC', 'Il gestionale, con le proprie API (firma-attestazione).',
     'Il concentratore all-ansc-sipo, che firma il token e la richiesta; la pagina non parla '
     'mai con ANSC.', 'Il concentratore è l’unico luogo del certificato e dell’OTP.'],
    ['Codice di sessione', 'Bottone OTP nella testata.', 'Tempo residuo dell’OTP nella testata '
     'dell’atto.', 'Lo stesso elemento delle altre pagine nuove.'],
    ['REGISTRO DI EMERGENZA', 'Bottone nella pagina.', 'Non presente.', 'In modalità di '
     'emergenza non si firma: la pagina si presenta solo al rientro.'],
    ['TORNA ALLA COMPILAZIONE', 'Possibile solo prima delle firme dei dichiaranti.',
     'Anche dopo, se tutti i firmatari sono in condizione di firmare di nuovo.',
     'Decisione del Comune, documento del workflow [R3].'],
    ['Dove si salva la scelta', 'allegato.fg_attestazione e fg_incluso_att_conformita.',
     'Le stesse colonne, su ALLEGATO in MATR_USR, legato a id_atto_sipo.',
     'SIPO non ha la tabella degli eventi di Side.'],
    ['Sospensione e ripresa', 'Non descritte.', 'SOSPENDI; alla ripresa le scelte salvate si '
     'ripropongono e l’anteprima va generata di nuovo.', 'Comportamento comune delle pagine '
     'nuove [R3].'],
    ['Cornice grafica', 'Comune di Milano.', 'SIPO.', 'Solo grafico.'],
]

ATT_PROVE = [
    ['#', 'Caso', 'Che cosa si fa', 'Esito atteso'],
    ['1', 'Sequenza ordinaria', 'Si apre la pagina, si genera, si firma.',
     'Al primo accesso tutti i documenti selezionati; dopo la firma ANSC 11 e la schermata '
     'della firma dell’USC.'],
    ['2', 'Nessun documento da attestare', 'Si deselezionano tutti e si preme CONTINUA.',
     'Si passa alla firma dell’USC; la scelta è salvata.'],
    ['3', 'Atto con documenti e dichiaranti', 'Si firmano i dichiaranti.',
     'Dopo le firme si presenta la pagina della conformità.'],
    ['4', 'Atto con documenti e senza dichiaranti', 'Si valida.',
     'Dopo la validazione si presenta la pagina della conformità.'],
    ['5', 'Atto senza documenti', 'Si valida, con o senza dichiaranti.',
     'La pagina si salta: si passa alla firma dell’USC.'],
    ['6', 'Annotazione art. 98', 'Un’annotazione con documento.',
     '[DA VERIFICARE: il percorso delle annotazioni in SIPO non è ancora disegnato.]'],
    ['7', 'Modalità di emergenza', 'Si recupera un atto preparato in emergenza.',
     'La pagina si presenta al rientro come nel percorso ordinario.'],
    ['8', 'Ripresa dopo una sospensione', 'Si sospende dopo GENERA e si riprende.',
     'Le scelte salvate si ripropongono; l’anteprima va generata di nuovo.'],
    ['9', 'Contrassegno assente', 'Un documento con fg_incluso_att_conformita nullo.',
     'È proposto come selezionato.'],
    ['10', 'Cambio di selezione', 'Dopo GENERA si toglie un documento.',
     'L’anteprima scompare, FIRMA si disabilita.'],
    ['11', 'Credenziali errate', 'FIRMA con PIN o OTP del servizio di firma sbagliati.',
     'Messaggio di ANSC (406001, 406004); PIN e OTP da reinserire; stato invariato.'],
    ['12', 'OTP di sessione scaduto', 'GENERA con la sessione scaduta (406002).',
     'Richiesta di un nuovo OTP; poi si ripete GENERA.'],
    ['13', 'Esito indeterminato', 'FIRMA senza risposta.',
     'Si legge lo stato con R005: se è 11 si prosegue senza firmare di nuovo.'],
    ['14', 'Ritorno ai dati dopo le firme', 'TORNA ALLA COMPILAZIONE con un firmatario che non '
     'può firmare di nuovo.', 'Il ritorno è inibito; resta ELIMINA EVENTO.'],
]

ATT_PUNTI = [
    ['#', 'Questione', 'Rimando'],
    ['AC-1', 'Il contratto del concentratore (ansc-v1.yaml) non ha ancora le due operazioni '
     'della pagina: generare l’anteprima con la scelta dei documenti (R010) e firmare '
     'l’attestazione (R007). Vanno aggiunte e la specifica delle API rigenerata.', 'OP-75 b'],
    ['AC-2', 'Da dove si prealimenta l’utenza di firma remota: Side usa quella dell’operatore '
     'collegato, SIPO non la conserva. [DA VERIFICARE]', '—'],
    ['AC-3', 'La preselezione: l’analisi segue Side (tutti selezionati); la guida ANSC dice che '
     'i documenti già firmati digitalmente non richiedono l’attestazione. Va deciso se '
     'lasciare deselezionati i P7M.', 'WF-18'],
    ['AC-4', 'Il percorso del servizio di anteprima è /anteprima/attestazione nel contratto e '
     '/attestazione/anteprima nel Changelog; l’algoritmo dell’impronta non è dichiarato. Fa '
     'fede il contratto.', 'WF-18'],
    ['AC-5', 'Le annotazioni con documenti (art. 98): il percorso in SIPO non è disegnato.',
     '—'],
    ['AC-6', 'Il ritorno ai dati dopo le firme richiede che ANSC accetti una nuova validazione '
     'di un atto già firmato dal dichiarante.', 'WF-16'],
]


def conformita():
    nome = f'PROCESSO_Attestato-Conformita_SIPO-ANSC_v{VERSIONE}'
    out = os.path.join(FIN, nome + '.docx')
    d = costruisci(out, 'Processo di gestione e firma dell’attestato di conformità degli '
                        'allegati')

    h(d, 1, 'Introduzione')
    par(d, 'Il documento adatta a Roma il processo di gestione e firma dell’attestato di '
           'conformità degli allegati del sistema Side del Comune di Milano [R1], scritto per il '
           'rilascio 1.40 di ANSC, che consente all’ufficiale di scegliere quali documenti '
           'attestare. In SIPO il processo diventa la pagina «Conformità» del percorso di '
           'formazione dell’atto [R3]; l’analisi lo adotta così com’è, con gli aggiustamenti '
           'dichiarati [R2].')
    par(d, 'L’attestato non lo scrive il Comune: lo genera ANSC, come PDF, a partire dai '
           'documenti scelti. La formula è fissa — «NOME COGNOME, ufficiale dello stato civile, '
           'attesta, che le copie allegate ed elencate nella sezione Allegati, sono conformi ai '
           'sensi di legge» — seguita dagli estremi dell’evento e dall’elenco dei documenti con '
           'la loro impronta [R4]. Serve per i documenti cartacei digitalizzati; quelli già '
           'firmati digitalmente sono originali elettronici.')
    par(d, 'La struttura è quella del documento di origine: attori, servizi, modifica del '
           'processo, pagina, casi di prova. Le schermate sono ridisegnate nello stile di SIPO; '
           'i valori sono di esempio.')
    h(d, 2, 'Riferimenti')
    tabella(d, [['#', 'Documento', 'Che cosa se ne prende'],
                ['[R1]', 'Sorgenti Documentali/' + SIDE_ATT, 'Il processo di origine: servizi, '
                 'regole della pagina, casi di prova.'],
                ['[R2]', 'ANALISI_Integrazione-ANSC_v3.44.docx, «L’attestazione di conformità»',
                 'L’adozione del processo di Side e i due contrassegni.'],
                ['[R3]', 'WORKFLOW_Interfacce-SIPO-ANSC_v0.3.docx, «Anteprima · Conformità»',
                 'La posizione nel percorso, il diagramma a stati, la regola del ritorno ai '
                 'dati, il comportamento comune delle pagine nuove.'],
                ['[R4]', 'ansc/docs/Guida_operativa_ANSC.pdf, §6.8', 'La formula e il contenuto '
                 'dell’attestato.'],
                ['[R5]', 'ansc/docs/openapi: R007_firma_usc, R010_anteprima, model_evento',
                 'Le strutture dei servizi, lette dai contratti.']], [0.6, 2.7, 3.2])
    h(d, 2, 'Glossario')
    tabella(d, [['Termine', 'Significato'],
                ['Attestato di conformità', 'Il documento, generato da ANSC, con cui '
                 'l’ufficiale attesta che le copie allegate sono conformi.'],
                ['Documento', 'Un allegato dell’atto, registrato in MATR_USR.ALLEGATO e '
                 'trasmesso ad ANSC con R001.'],
                ['Credenziali di firma remota', 'Utenza, PIN e OTP del servizio di firma '
                 'qualificata, richiesti da R007.'],
                ['OTP di sessione', 'Il codice di quattro ore generato sulla web app di ANSC, '
                 'necessario a ogni chiamata.']], [1.9, 4.6])

    h(d, 1, 'Attori')
    tabella(d, [['Nome', 'Descrizione', 'Attore di sistema'],
                ['Ufficiale di stato civile', 'Abilitato a SIPO e ad ANSC: sceglie i documenti '
                 'da attestare e firma l’attestato con la firma remota.', 'No'],
                ['Concentratore operativo (all-ansc-sipo)', 'Chiama R010 e R007 per conto della '
                 'pagina, con il token firmato e la sessione OTP dell’ufficiale; salva le '
                 'scelte su ALLEGATO.', 'Sì'],
                ['ANSC', 'Genera il PDF dell’attestato e ne registra la firma; porta l’atto a '
                 '11 FIRMATO PER CONFORMITÀ.', 'Sì']], [2.0, 3.6, 0.9])

    h(d, 1, 'Servizi')
    par(d, 'Il rilascio 1.40 di ANSC ha introdotto tre modifiche [R1]: il campo '
           'flagInclusoAttConformita nel modello del documento allegato; l’elenco dei documenti, '
           'con il contrassegno, nella richiesta di firma dell’attestato (R007 '
           '/attestazione/firma) e in quella della sua anteprima (R010 /anteprima/attestazione). '
           'Le tabelle che seguono sono lette dai contratti [R5].')
    h(d, 2, 'Il modello del documento: flagInclusoAttConformita e flagAttestazione')
    righe, ver = schema('model_evento.yaml', 'ModelAllegatoRif')
    par(d, f'ModelAllegatoRif, nel modello evento (contratto {ver}).')
    tabella(d, righe, [1.6, 0.8, 1.6, 2.5])
    voce(d, '**Due tipi per la stessa informazione.** Il contrassegno è stringa nel modello '
            'evento (flagInclusoAttConformita) e booleano nei servizi R007 e R010 '
            '(flagAttConformita).')
    voce(d, '**flagAttestazione.** Sogei ha risposto nel repository pubblico di ANSC che si può '
            'ignorare (discussione 787, citata da Side). In SIPO i due campi si valorizzano '
            'sempre con lo stesso valore, come stabilito nell’analisi [R2].')
    h(d, 2, 'R010 — anteprima dell’attestato')
    righe, ver = schema('R010_anteprima.yaml', 'AnteprimaAttestazioneRequest')
    par(d, f'POST /anteprima/attestazione/{{version}} (contratto {ver}). La richiesta:')
    tabella(d, righe, [1.6, 1.6, 0.9, 2.4])
    righe, _ = schema('R010_anteprima.yaml', 'AllegatoInput')
    par(d, 'Ciascun documento dell’elenco (AllegatoInput):')
    tabella(d, righe, [1.4, 0.8, 1.2, 3.1])
    righe, _ = schema('R010_anteprima.yaml', 'AnteprimaAttestazioneResponse')
    par(d, 'La risposta: il PDF in base64 e un contrassegno che dice se c’è qualcosa da '
           'attestare.')
    tabella(d, righe, [1.8, 2.0, 0.8, 1.9])
    h(d, 2, 'R007 — firma dell’attestato')
    righe, ver = schema('R007_firma_usc.yaml', 'FirmaAttestazioneRequest')
    par(d, f'POST /attestazione/firma/{{version}} (contratto {ver}). La richiesta porta lo '
           'stesso elenco di documenti e i parametri della firma remota:')
    tabella(d, righe, [1.4, 1.9, 0.8, 2.4])
    righe, _ = schema('R007_firma_usc.yaml', 'ParametriFirma')
    tabella(d, righe, [1.3, 0.8, 1.0, 3.4])
    par(d, 'Le credenziali si inviano in base64 (idModalitaInput 2): la modalità in chiaro non '
           'è accettata in produzione. Il documento di Milano scambia R007 e R010 nei titoli di '
           'due paragrafi: qui i nomi seguono i contratti.')

    h(d, 1, 'Modifica del processo')
    par(d, 'La pagina si presenta dopo la firma dei dichiaranti, se l’atto ne ha, oppure subito '
           'dopo la validazione. Se l’atto non ha documenti si salta e si passa alla firma '
           'dell’ufficiale [R3]. Come in Side:')
    voce(d, 'L’ufficiale non trova subito l’attestato ma l’elenco di tutti i documenti '
            'dell’atto, compreso il processo verbale delle firme dei dichiaranti.')
    voce(d, 'Sceglie quali documenti attestare.')
    voce(d, 'Solo se almeno un documento è selezionato si può generare l’attestato.')
    voce(d, 'Solo se l’attestato è stato generato si può firmare.')
    voce(d, 'Se tutti i documenti sono deselezionati, si passa alla firma finale con CONTINUA.')
    par(d, 'A Roma la pagina è una schermata di SIPO a sé, che si apre con il banner di arrivo '
           'dalla pagina precedente, porta in testa i dati dell’atto e il tempo residuo della '
           'sessione con ANSC, e ha i pulsanti comuni alle pagine nuove: INDIETRO, TORNA ALLA '
           'COMPILAZIONE, ELIMINA EVENTO, SOSPENDI [R3].')

    h(d, 1, 'Pagina della firma per conformità')
    h(d, 2, 'Scelta dei documenti da attestare')
    par(d, 'La pagina elenca i documenti dell’atto, con tipo, formato, stato e impronta. Al '
           'primo accesso sono tutti selezionati; agli accessi successivi ripropone la scelta '
           'salvata. Un contrassegno nullo vale come selezionato [R2]. Le credenziali di firma '
           'sono disabilitate finché l’attestato non è generato.')
    figura(d, 'pr_conf_iniziale.png', 6.3, 'Primo accesso: documenti tutti selezionati, '
           'GENERA abilitato, FIRMA e CONTINUA disabilitati.')
    par(d, 'Se l’atto è firmato dai dichiaranti, i dati sono vincolati: TORNA ALLA COMPILAZIONE '
           'è disabilitato, salvo che tutti i firmatari siano in condizione di firmare di nuovo. '
           'Senza firme dei dichiaranti si torna ai dati liberamente [R3].')
    h(d, 2, 'Generazione dell’attestato')
    par(d, 'GENERA ATTESTAZIONE DI CONFORMITÀ chiama R010 con l’elenco dei documenti e il '
           'contrassegno di ciascuno, mostra l’anteprima del PDF restituito da ANSC e abilita le '
           'credenziali di firma, con l’utenza prealimentata (AC-2). Salva la scelta su ALLEGATO '
           '— fg_incluso_att_conformita e fg_attestazione, con lo stesso valore — perché la '
           'pagina possa riproporla. L’ufficiale può ancora cambiare la selezione: l’anteprima '
           'scompare, le credenziali si disabilitano e l’attestato va generato di nuovo.')
    figura(d, 'pr_conf_anteprima.png', 6.3, 'Attestato generato su due documenti: anteprima '
           'del PDF di ANSC e firma abilitata.')
    h(d, 2, 'Firma')
    par(d, 'FIRMA è abilitato solo con l’anteprima visibile e chiama R007 con lo stesso elenco. '
           'A firma riuscita ANSC porta l’atto a 11 FIRMATO PER CONFORMITÀ, la pagina di ripresa '
           'diventa quella della firma dell’ufficiale e si apre la schermata della firma '
           'dell’USC. Se le credenziali sono rifiutate (406001, 406004) la pagina lo dice, PIN e '
           'OTP del servizio di firma vanno reinseriti e lo stato non cambia. Se ANSC non '
           'risponde, lo stato si legge con R005 prima di qualunque nuovo tentativo.')
    h(d, 2, 'Nessun documento da attestare')
    par(d, 'Se l’ufficiale deseleziona tutti i documenti, GENERA si disabilita e CONTINUA porta '
           'direttamente alla firma dell’ufficiale. Anche CONTINUA salva la scelta.')
    figura(d, 'pr_conf_nessuno.png', 6.3, 'Nessun documento selezionato: CONTINUA abilitato.')
    h(d, 2, 'I pulsanti')
    tabella(d, [['Pulsante', 'Abilitato quando', 'Che cosa fa'],
                ['GENERA ATTESTAZIONE DI CONFORMITÀ', 'Almeno un documento selezionato.',
                 'R010; anteprima; salva la scelta.'],
                ['FIRMA', 'Anteprima visibile.', 'R007; poi la firma dell’USC.'],
                ['CONTINUA', 'Nessun documento selezionato.', 'Salva la scelta; poi la firma '
                 'dell’USC.'],
                ['INDIETRO', 'Sempre.', 'Torna alla pagina precedente, senza cambiare i dati.'],
                ['TORNA ALLA COMPILAZIONE', 'Nessuna firma dei dichiaranti, o tutti i firmatari '
                 'in condizione di firmare di nuovo.', 'Riapre i dati; l’atto torna nello stato '
                 '0.'],
                ['ELIMINA EVENTO', 'Sempre, con conferma.', 'R011: l’atto è annullato in ANSC.'],
                ['SOSPENDI', 'Sempre.', 'Chiude la redazione, riprendibile da questa pagina.']],
            [2.1, 2.2, 2.2])

    h(d, 1, 'Casi di prova')
    par(d, 'I primi sette sono quelli di Side; gli altri sono propri di Roma.')
    tabella(d, ATT_PROVE, [0.4, 1.6, 2.1, 2.4])

    h(d, 1, 'Che cosa cambia rispetto a Side')
    par(d, 'Ciò che non compare nella tabella si adotta come Side lo descrive.')
    tabella(d, ATT_ADATTAMENTI, [1.3, 1.6, 1.9, 1.7])

    h(d, 1, 'Punti aperti')
    tabella(d, ATT_PUNTI, [0.6, 4.9, 1.0])
    chiudi(d, out, nome, 'Prima emissione: il processo di Side adattato alla pagina di '
           'conformità di SIPO.')
    return out


def main():
    F.disegna_tutte()
    numerazione()
    conformita()


if __name__ == '__main__':
    main()
