# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.43 → v3.44 (21/10/2026).

Decisione dell'utente: una sola tabella per l'atto. La v3.43 aveva introdotto ANSC_EVENTO
accanto ad ANSC_STATO_ATTO, in rapporto 1 a 1: la stessa entità divisa in due. ANSC_EVENTO
sparisce e il suo contenuto entra in ANSC_STATO_ATTO: le otto colonne della testata del modello
evento che lo store non aveva, il modello evento intero (TXT_EVENTO), l'impronta dell'ultima
validazione e i dati della ripresa. Tutti i rimandi ad ANSC_STATO_ATTO restano validi.

Il testo della v3.43 (tracciato, regole della ripresa, stato 0) si riusa da v3_43_evento.py:
qui cambia solo dove sta.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_44_tabella_unica.py
"""
import copy
import os
import shutil
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D          # noqa: E402
import diagrammi_erd             # noqa: E402
import diagrammi_flusso          # noqa: E402
import v3_43_evento as V         # noqa: E402

FIN, IMG = V.FIN, V.IMG
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.43.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.44.docx')
SA = 'ANSC_STATO_ATTO'
TITOLO_RIPRESA = 'La ripresa di una redazione sospesa'

# ------------------------------------------------------------------------------ 5.2.1
TITOLO_521_VECCHIO = 'ANSC_STATO_ATTO — store di stato ANSC per atto.'
TITOLO_521 = 'ANSC_STATO_ATTO — l’atto in ANSC: stato, modello evento e ripresa.'

# le righe di colonna che si aggiungono alla tabella descrittiva, prima dei campi tecnici
NUOVE_COLONNE = [r for r in V.COLONNE_EVENTO[1:]
                 if r[0] not in ('ID_EVENTO', 'ID_STATO_ATTO', 'ID_ATTO_SIPO')
                 and not r[0].startswith('DATA_INS')]
NUOVE_COLONNE[0] = ['ID_EVENTO_ANSC', 'VARCHAR2(50)', 'Il campo id del modello evento, '
                    'assegnato da ANSC «per uso interno».']
assert [r[0] for r in NUOVE_COLONNE][-1] == 'NUM_REVISIONE'

RIPRESA = [list(x) for x in V.RIPRESA]
for r in RIPRESA:
    r[1] = r[1].replace('gli identificativi dei soggetti da ANSC_SOGGETTO_IDENTIFICATIVI; il '
                        'numero comunale e l’UC da ANSC_STATO_ATTO.',
                        'gli identificativi dei soggetti da ANSC_SOGGETTO_IDENTIFICATIVI; il '
                        'numero comunale e l’UC dalla riga stessa.')
assert 'dalla riga stessa' in RIPRESA[2][1]

# ------------------------------------------------------------------------------ DDL
COLONNE_DDL = [
    '  ID_EVENTO_ANSC          VARCHAR2(50 CHAR),',
    '  FLG_SECRETATO           CHAR(1 CHAR),',
    '  COD_VERSIONE_ANSC       VARCHAR2(10 CHAR),',
    '  DATA_FORMAZIONE         TIMESTAMP,',
    '  COD_RUOLO_OPERATORE     VARCHAR2(10 CHAR),',
    '  FLG_CASA_COMUNALE       CHAR(1 CHAR),',
    '  ID_COMUNE_REGISTRAZIONE VARCHAR2(10 CHAR),',
    '  DATA_DECORRENZA         TIMESTAMP,',
    '  TXT_EVENTO              CLOB,',
    '  TXT_HASH_VALIDATO       VARCHAR2(64 CHAR),',
    '  DATA_VALIDAZIONE        TIMESTAMP,',
    "  COD_PAGINA_RIPRESA      VARCHAR2(30 CHAR) DEFAULT 'DETTAGLIO' NOT NULL,",
    "  FLG_SOSPESO             CHAR(1 CHAR) DEFAULT 'N' NOT NULL,",
    '  DATA_SOSPENSIONE        TIMESTAMP,',
    '  UTENTE_SOSPENSIONE      VARCHAR2(40 CHAR),',
    '  NUM_REVISIONE           NUMBER DEFAULT 0 NOT NULL,',
]
VINCOLI_DDL = [
    '  CONSTRAINT CK_ANSC_STATO_ATTO_EVENTO CHECK (TXT_EVENTO IS JSON),',
    "  CONSTRAINT CK_ANSC_STATO_ATTO_SECRETATO CHECK (FLG_SECRETATO IN ('0','1')),",
    "  CONSTRAINT CK_ANSC_STATO_ATTO_CASA CHECK (FLG_CASA_COMUNALE IN ('S','N')),",
    "  CONSTRAINT CK_ANSC_STATO_ATTO_SOSPESO CHECK (FLG_SOSPESO IN ('S','N')),",
    '  CONSTRAINT CK_ANSC_STATO_ATTO_PAGINA',
    "    CHECK (COD_PAGINA_RIPRESA IN ('DETTAGLIO','RIEPILOGO','ANTEPRIMA','VALIDAZIONE',",
    "                                  'FIRME_DICHIARANTI','CONFORMITA','FIRMA_USC'))",
]
INDICE_DDL = [
    'CREATE INDEX ANSC_USR.IX_ANSC_STATO_ATTO_SOSPESO',
    '  ON ANSC_USR.ANSC_STATO_ATTO (FLG_SOSPESO, UTENTE_SOSPENSIONE) TABLESPACE ANSC_USR;',
]
COMMENTI_DDL = [
    'COMMENT ON COLUMN ANSC_USR.ANSC_STATO_ATTO.TXT_EVENTO IS',
    "  'Modello evento in redazione; formato l''atto, la copia depositata e firmata.';",
    'COMMENT ON COLUMN ANSC_USR.ANSC_STATO_ATTO.TXT_HASH_VALIDATO IS',
    "  'SHA-256 del modello all''ultima validazione riuscita: se diverso, si rivalida (stato 0).';",
]
for n in ('CK_ANSC_STATO_ATTO_SECRETATO', 'IX_ANSC_STATO_ATTO_SOSPESO'):
    assert len(n) <= 30, n

# ------------------------------------------------------------------------------ rimandi
SOSTITUZIONI = [
    ('la riga dell’atto in ANSC_STATO_ATTO, in IN_LAVORAZIONE, e quella in ANSC_EVENTO con la '
     'pagina di ripresa.',
     'la riga dell’atto in ANSC_STATO_ATTO, in IN_LAVORAZIONE, con la pagina di ripresa.'),
    ('ANSC_EVENTO.TXT_EVENTO: il modello evento in redazione',
     'ANSC_STATO_ATTO.TXT_EVENTO: il modello evento in redazione'),
    ('Il modello evento in ANSC_EVENTO.', 'Il modello evento, in ANSC_STATO_ATTO.TXT_EVENTO.'),
    ('ID_OPERAZIONE_ANSC e stato CONFERMATO; ANSC_EVENTO: l’impronta del modello validato.',
     'ID_OPERAZIONE_ANSC, stato CONFERMATO e l’impronta del modello validato.'),
    ('ogni salvataggio aggiorna anche ANSC_EVENTO con il modello evento e la pagina raggiunta',
     'ogni salvataggio aggiorna anche, in ANSC_STATO_ATTO, il modello evento e la pagina '
     'raggiunta'),
    ('(si veda «ANSC_EVENTO — l’atto in redazione in ANSC e il punto di ripresa»)',
     f'(si veda «{TITOLO_RIPRESA}»)'),
    ('ANSC_STATO_ATTO conserva l’esito e ANSC_EVENTO il modello evento depositato,',
     'ANSC_STATO_ATTO conserva l’esito e il modello evento depositato,'),
    ('del modello evento salvato in ANSC_EVENTO differisce',
     'del modello evento salvato in ANSC_STATO_ATTO differisce'),
    (' Il modello evento non sta più qui: è in ANSC_EVENTO.', ''),
    ('REFERENCES ANSC_USR.ANSC_EVENTO(ID_EVENTO);',
     'REFERENCES ANSC_USR.ANSC_STATO_ATTO(ID_STATO_ATTO);'),
    ('La tabella dell’evento in redazione', 'L’atto in redazione e la ripresa'),
]

STORIA = ('21/10/2026', '3.44',
          'Adeguamenti database · Il flusso operativo · Il processo di gestione dei dati ANSC in '
          'SIPO · Appendice A',
          'Una sola tabella per l’atto: ANSC_EVENTO, introdotta nella versione precedente '
          'accanto ad ANSC_STATO_ATTO in rapporto uno a uno, è eliminata e il suo contenuto '
          'entra in ANSC_STATO_ATTO — le otto colonne della testata del modello evento che lo '
          'store non aveva, il modello evento intero, l’impronta dell’ultima validazione e i '
          'dati per sospendere e riprendere la redazione. forzatura_validazione richiama '
          'ANSC_STATO_ATTO. Rigenerate la figura del flusso operativo e lo schema ANSC_USR.')


# ------------------------------------------------------------------------------ primitive
def paragrafo(d, testo):
    ps = [p for p in d.paragraphs if p.text == testo]
    assert len(ps) == 1, (testo, len(ps))
    return ps[0]


def clona_dopo(p, righe):
    """Inserisce dopo p tanti paragrafi quante le righe, con la stessa forma di p."""
    ultimo = p._p
    for testo in righe:
        nuovo = copy.deepcopy(p._p)
        ultimo.addnext(nuovo)
        np = Paragraph(nuovo, p._parent)
        for r in np.runs[1:]:
            r._r.getparent().remove(r._r)
        np.runs[0].text = testo
        ultimo = nuovo
    return Paragraph(ultimo, p._parent)


def elimina(elementi):
    for el in elementi:
        assert not D.ha_commenti(el)
        el.getparent().remove(el)


# ------------------------------------------------------------------------------ main
def main():
    if os.path.exists(DST):
        os.remove(DST)                       # solo un'esecuzione precedente di questo script
    shutil.copy(SRC, DST)
    prima = V.commenti(DST)
    d = docx.Document(DST)
    fatti = []
    modello = D.trova_tabella(d, 'Colonna', 'Tipo', 'Note')

    # ------------------------------------------------ 5.2: via la sezione ANSC_EVENTO
    elimina(D.sezione(d, 'ANSC_EVENTO — l’atto in redazione', livello=3))
    fatti.append('tolta la sezione ANSC_EVENTO del cap. 5.2')

    # ------------------------------------------------ 5.2.1: titolo, colonne, testata, ripresa
    D.testo_di(D.h(d, 3, TITOLO_521_VECCHIO, esatto=True), TITOLO_521)
    t = V.tabella_con(d, 'ULTIMO_ERRORE')
    i = [k for k, r in enumerate(t.rows) if r.cells[0].text.strip() == 'ULTIMO_ERRORE'][0]
    for k, valori in enumerate(NUOVE_COLONNE):
        V.riga_dopo(t, i + k, valori)
    fatti.append(f'5.2.1: +{len(NUOVE_COLONNE)} righe nella tabella delle colonne')

    fine = D.h(d, 3, 'ANSC_LOG_AUDIT — traccia')._p
    righe_tr = [(i_, c, tp, dove.replace('ANSC_EVENTO.', f'{SA}.'), nota)
                for i_, c, tp, dove, nota in V.tracciato()]
    n_sa = sum(1 for x in righe_tr if x[3].startswith(SA))
    n_der = sum(1 for x in righe_tr if x[3] == '— derivato')
    n_json = sum(1 for x in righe_tr if x[3] == '— nel JSON')
    n_dep = sum(1 for x in righe_tr if x[3] == '— deprecato')
    assert n_sa + n_der + n_json + n_dep == 32

    D.para(d, fine, 'La testata del modello evento', stile='Heading 4')
    D.para(d, fine, 'La tabella lega l’atto redatto in SIPO all’atto in redazione in ANSC, ne '
           'conserva il modello evento intero e consente di sospendere la redazione e di '
           'riprenderla dalla pagina in cui era, con tutti i dati già inseriti per ANSC: '
           'soggetti e loro identificativi, documenti, firmatari. È il corrispondente della '
           'tabella evento del sistema Side, a cui le tabelle ereditate firmatari e '
           'forzatura_validazione si legano.')
    D.para(d, fine, pezzi=D.segmenta('', (
        f'Le colonne della testata sono ricavate dalle righe 2-33 del foglio ModelEvento di '
        f'{V.FOGLIO_TRACCIATO}: trentadue campi, prima che il modello cominci a descrivere '
        f'dichiarante, soggetti e sezioni. Il criterio è uno: **una colonna per ciascun campo su '
        f'cui si cerca, si filtra, si decide o si riprende; tutto il resto sta nel modello '
        f'evento, che è la copia integrale.**')))
    D.para(d, fine, f'Dei trentadue campi {n_sa} hanno una colonna in questa tabella, {n_der} si '
           f'ricavano da altre colonne, {n_json} restano nel modello evento e {n_dep} sono '
           f'deprecati da ANSC.')
    righe = [['Riga', 'Campo', 'Tipo', 'Dove si registra', 'Note']]
    for i_, campo, tipo, dove, nota in righe_tr:
        righe.append([str(i_), campo, tipo, dove, nota])
    tz = D.tabella(d, fine, righe, modello, [0.4, 1.6, 0.8, 1.8, 2.0])
    for r in tz.rows[1:]:
        for c in (r.cells[1], r.cells[3]):
            for run in c.paragraphs[0].runs:
                run.text = run.text.replace('_', '_​').replace('.', '.​')
    D.para(d, fine, 'Testata del modello evento e dove si registra ciascun campo.', corsivo=True)
    D.para(d, fine, pezzi=D.segmenta('', (
        '**Il modello evento è la copia di lavoro**: in redazione cambia a ogni salvataggio, '
        'formato l’atto coincide con ciò che è stato depositato e firmato, perché una modifica '
        'dopo la validazione porta allo stato 0 e impone una nuova validazione prima delle '
        'firme. Una sola colonna basta così per entrambi gli usi.')))

    D.para(d, fine, TITOLO_RIPRESA, stile='Heading 4')
    D.para(d, fine, 'La ripresa si regge su sette regole.')
    for testa, corpo in RIPRESA:
        D.voce(d, fine, testa, corpo)
    D.para(d, fine, 'Il volume è quello del modello evento: con i 144.285 atti l’anno della '
           'ricognizione dei tipi atto e una dimensione media fra dieci e trenta kilobyte, fra '
           'uno e quattro gigabyte l’anno. Si conserva un modello per atto, non uno per '
           'salvataggio: la storia dei passaggi è nell’audit.')
    fatti.append(f'5.2.1: testata del modello evento ({n_sa} campi in colonna) e ripresa')

    # ------------------------------------------------ tabella degli oggetti del cap. 5
    ob = D.trova_tabella(d, 'Oggetto', 'Ruolo', 'Note principali')
    r = [x for x in ob.rows if x.cells[0].text.strip() == 'ANSC_EVENTO'][0]
    elimina([r._tr])
    r = [x for x in ob.rows if x.cells[0].text.strip() == SA][0]
    D.riscrivi_cella(r.cells[2], r.cells[2].text.strip() + ' Porta anche la testata del '
                     'modello evento, il modello intero e la pagina da cui riprendere una '
                     'redazione sospesa.')

    # ------------------------------------------------ rimandi nel testo
    for vecchio, nuovo in SOSTITUZIONI:
        V.sost(d, vecchio, nuovo, fatti)

    # ------------------------------------------------ Appendice A
    titolo = [p for p in d.paragraphs if p.style.name == 'Heading 3' and p.text == 'ANSC_EVENTO']
    assert len(titolo) == 1
    elementi = D.sezione(d, 'ANSC_EVENTO', livello=3)
    assert Paragraph(elementi[0], d).text == 'ANSC_EVENTO'
    elimina(elementi)
    clona_dopo(paragrafo(d, '  ID_ANSC                 VARCHAR2(50 CHAR),'), COLONNE_DDL)
    ult = paragrafo(d, "    CHECK (COD_ORIGINE_UC IN ('AUTOMATICA','OPERATORE'))")
    ult.runs[0].text = ult.text + ','
    for r in ult.runs[1:]:
        r._r.getparent().remove(r._r)
    clona_dopo(ult, VINCOLI_DDL)
    clona_dopo(paragrafo(d, '  ON ANSC_USR.ANSC_STATO_ATTO (ID_ATTO_SIPO) TABLESPACE ANSC_USR;'),
               INDICE_DDL)
    clona_dopo(paragrafo(d, "  'Esito negativo qualificato: SCARTATO (400 in validazione) o "
                            "INDETERMINATO (timeout).';"), COMMENTI_DDL)
    fatti.append('Appendice A: tolta la DDL di ANSC_EVENTO, colonne e vincoli in ANSC_STATO_ATTO')

    # ------------------------------------------------ nessun residuo
    residui = [p.text[:80] for p in V.tutti_i_paragrafi(d)
               if 'ANSC_EVENTO' in p.text and not p.text.startswith(('Allineati al processo',
                                                                      'Una sola tabella'))]
    assert not residui, residui

    # ------------------------------------------------ figure
    V.immagine_proporzionata(d, 'Il flusso operativo nelle sue quattro fasi',
                             diagrammi_flusso.flusso(os.path.join(IMG, 'flusso_operativo.png')))
    V.immagine_proporzionata(d, 'Schema ANSC_USR, con ALLEGATO', diagrammi_erd.disegna())
    fatti.append('figure rigenerate')

    D.storia(d, *STORIA)
    fatti += D.frontespizio(d, '3.44', data='21/10/2026')
    d.save(DST)
    dopo = V.commenti(DST)
    print('\n'.join('  · ' + f for f in fatti))
    print(f'  commenti di Word conservati: {prima} → {dopo}')
    if dopo != prima:
        raise SystemExit('commenti persi')
    cap, tab, fig = D.riepilogo(DST)
    print(f'  capitoli {cap}, tabelle {tab}, immagini {fig}')
    print('  scritto', DST)


if __name__ == '__main__':
    main()
