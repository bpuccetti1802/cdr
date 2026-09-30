# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.24 -> v3.25 (24/09/2026).

Il valore di un dizionario deve potersi estrarre con la chiave «dominio + valore».

  · È già la chiave primaria della replica locale, ma mancava l'operazione che la usa:
    aggiunta `leggiValoreDizionario` al contratto (8 operazioni).
  · ⚠️ La coppia però NON è univoca per due identificativi su 143 (134 e 135 coprono due
    tabelle ciascuno): il nome della tabella entra nella chiave delle due tabelle di replica,
    e la configurazione dei campi lo porta con sé, così che la lettura non sia mai ambigua.
  · La stessa chiave vale sulla vista letta da SIPO: cambia il canale, non la chiave.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_25_chiave_dizionari.py
"""
import os
import shutil
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.24.docx')
DST = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.25.docx')
IMG = os.path.join(BASE, 'strumenti', 'img')

if os.path.exists(DST):
    os.remove(DST)
shutil.copy(SRC, DST)
doc = docx.Document(DST)
fatti = []


def par(inizio, stile=None):
    t = [p for p in doc.paragraphs
         if p.text.strip().startswith(inizio) and (stile is None or p.style.name == stile)]
    if len(t) != 1:
        raise SystemExit(f'attesa 1 occorrenza di «{inizio[:60]}», trovate {len(t)}')
    return t[0]


def riga_con(t, prima_cella):
    for r in t.rows:
        if r.cells[0].text.strip() == prima_cella:
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


def inserisci_riga(t, dopo, valori):
    nuova = D.clona_riga(t, valori)
    riga_con(t, dopo)._tr.addnext(nuova._tr)
    return nuova


def mono(testo):
    t = [p for p in doc.paragraphs if p.text.rstrip() == testo.rstrip()]
    if len(t) != 1:
        raise SystemExit(f'riga DDL non univoca: «{testo[:60]}» ({len(t)})')
    return t[0]


def elimina(elementi, cosa):
    for e in elementi:
        if D.ha_commenti(e):
            raise SystemExit(f'{cosa}: commento ancorato qui, non si cancella')
    for e in elementi:
        e.getparent().remove(e)
    fatti.append('eliminato: ' + cosa)


def blocco_ddl(create_line, fine=') TABLESPACE'):
    # `fine` è il prefisso della riga che chiude il blocco, dopo lo strip
    p = mono(create_line)
    out, e = [p._p], p._p.getnext()
    while e is not None:
        out.append(e)
        if e.tag.endswith('}p') and Paragraph(e, doc).text.strip().startswith(fine):
            return out
        e = e.getnext()
    raise SystemExit('blocco DDL non chiuso: ' + create_line)


def sostituisci_blocco(create_line, righe, fine=') TABLESPACE'):
    el = blocco_ddl(create_line, fine)
    dopo = el[-1].getnext()
    elimina(el, 'blocco DDL ' + create_line.split('.')[-1].split()[0])
    D.ddl(doc, dopo, righe)


MOD = D.trova_tabella(doc, 'colonna', 'tipo', 'note')


# ═════════════════════════════════════ 1. scheda e storia
for t in doc.tables:
    if t.rows[0].cells[0].text.strip().lower().startswith(('area organizzativa', 'progetto')):
        for nome, val in (('Data consegna', '24/09/2026'), ('Versione', '3.25')):
            try:
                D.riscrivi_cella(riga_con(t, nome).cells[1], val)
            except SystemExit:
                pass

D.storia(doc, '24/09/2026', '3.25',
         'Gestione dei dizionari · Costruzione del payload · Appendice A · Appendice B · '
         'Open Point',
         'Il valore di una decodifica si estrae per chiave, che è la coppia decodifica più '
         'valore: al contratto si aggiunge la lettura puntuale corrispondente, e la stessa '
         'chiave vale sulla vista che SIPO legge in base dati. Poiché nel corpus pubblicato '
         'due identificativi coprono due tabelle ciascuno, la coppia non basta in quei due '
         'casi: il nome della tabella entra nella chiave delle tabelle di replica e la '
         'configurazione dei campi lo porta con sé, così che la traduzione di un codice non '
         'sia mai ambigua. Aggiornati il DDL, lo schema entità-relazioni e OP-28.')


# ═════════════════════════════════════ 2. l'estrazione per chiave
ancora = par('Due precisazioni sul confronto dei testi')._p
D.para(doc, ancora, 'L’estrazione per chiave: decodifica più valore', stile='Heading 3')
for t in [
    'La lettura più frequente non è né una ricerca né un elenco: è la traduzione di un singolo '
    'codice. Dato il campo, la configurazione dice quale decodifica lo governa; dato il valore, '
    'serve la sua descrizione — per scriverla in un atto, per mostrarla in una maschera, per '
    'spiegare un rifiuto che cita un codice. ⚠️ La chiave di questa lettura è la coppia '
    'decodifica più valore, ed è la chiave primaria della replica locale: la lettura per chiave '
    'non richiede quindi alcuna struttura nuova, richiede che sia dichiarata.',
    'La coppia vale su entrambi i canali, e questo è il punto che rende la scelta sostenibile: '
    'il back-office e la diagnosi la usano attraverso l’operazione del componente, le maschere '
    'di SIPO la usano in base dati sulla vista di sola lettura. Cambia il canale, non la '
    'chiave: una descrizione letta dall’una e dall’altra parte è la stessa.',
    '⚠️ La validità temporale non filtra la lettura per chiave. Un valore cessato deve potersi '
    'leggere, perché serve a interpretare gli atti formati quando era in vigore; la risposta '
    'dichiara se sia ancora valido, e chi propone una scelta all’operatore esclude i cessati. '
    'Filtrare qui significherebbe non saper più leggere i propri atti.',
    '⚠️ Resta un caso in cui la coppia non basta, e va affrontato adesso perché tocca la chiave '
    'primaria. Nel corpus pubblicato due identificativi coprono due tabelle ciascuno — il 134 e '
    'il 135 — e per quelli la stessa coppia designa due valori diversi. Le conseguenze sono '
    'due: il carico completo violerebbe la chiave primaria della struttura ereditata, e la '
    'lettura per chiave restituirebbe un valore arbitrario. Per questo il nome della tabella '
    'entra nella chiave delle due tabelle di replica, e la configurazione del campo lo porta '
    'con sé quando serve: per le altre centoquarantuno decodifiche resta indicato soltanto '
    'l’identificativo, e la coppia è sufficiente.',
]:
    D.para(doc, ancora, t)
D.tabella(doc, ancora, [
    ('Chi legge', 'Come', 'Che cosa ottiene'),
    ('Costruzione del payload', 'Vista in base dati, per chiave.',
     'Verifica che il valore trasmesso esista nella decodifica dichiarata dalla '
     'configurazione. ⚠️ È la verifica di esistenza, non la traduzione: che cosa scrivere al '
     'posto di un valore di SIPO lo dice la transcodifica (OP-59).'),
    ('Maschere di SIPO', 'Vista in base dati, per sinonimo e in sola lettura.',
     'La descrizione da mostrare, e l’elenco dei valori validi quando la maschera propone una '
     'scelta.'),
    ('Back-office e diagnosi', 'Operazione leggiValoreDizionario del componente.',
     'Il valore con la sua validità; 404 se la decodifica non è replicata o non contiene il '
     'valore, e le due cause si distinguono nel messaggio.'),
    ('Presentazione di un atto già formato', 'Vista in base dati, per chiave, senza filtro di '
     'validità.', 'La descrizione in vigore all’epoca dell’atto, anche se il valore è stato '
     'poi cessato.'),
], MOD)
D.para(doc, ancora, 'Chi legge per chiave, con quale canale e che cosa ne ottiene.',
       corsivo=True)
fatti.append('nuova sezione «L’estrazione per chiave»')

# la tabella delle interfacce guadagna l'operazione
t = D.trova_tabella(doc, 'operazione', 'che cosa fa', 'chi la usa')
inserisci_riga(t, 'elencaValoriDizionario', (
    'leggiValoreDizionario',
    '⚠️ Nuova: legge un singolo valore per chiave, decodifica più valore. Non filtra per '
    'validità: la dichiara. Risponde 409, invece di indovinare, quando l’identificativo copre '
    'più tabelle e il nome non è indicato.',
    'Traduzione di un codice in descrizione; verifica di esistenza; diagnosi.'))


# ═════════════════════════════════════ 3. il modello dati
t = D.tabella_colonne(doc, 'nm_dominio')
D.riscrivi_cella(riga_con(t, 'id_dominio').cells[1], 'NUMBER (PK)')
D.riscrivi_cella(riga_con(t, 'nm_dominio').cells[1], 'VARCHAR2(50) (PK)')
D.riscrivi_cella(riga_con(t, 'nm_dominio').cells[2],
                 'Nome della tabella restituito dal servizio (es. dec_stato_evento). ⚠️ Entra '
                 'nella chiave: due identificativi su 143 coprono due tabelle ciascuno, e con '
                 'la sola chiave numerica il carico completo fallirebbe e la lettura per '
                 'chiave sarebbe ambigua. È la correzione minima alla struttura ereditata.')

t = D.tabella_colonne(doc, 'id_valore')
D.riscrivi_cella(riga_con(t, 'id_dominio').cells[1], 'NUMBER (PK, FK)')
inserisci_riga(t, 'id_dominio', (
    'nm_dominio', 'VARCHAR2(50) (PK, FK)',
    'Tabella di appartenenza. ⚠️ Insieme all’identificativo e al valore forma la chiave: è ciò '
    'che consente di replicare per intero le due decodifiche che ne condividono '
    'l’identificativo.'))
D.riscrivi_cella(riga_con(t, 'id_valore').cells[2],
                 'Il valore codificato. Insieme alla decodifica forma la chiave di lettura: è '
                 'la coppia con cui il payload verifica e le maschere traducono.')

t = D.tabella_colonne(doc, 'ID_DECODIFICA_ANSC')
inserisci_riga(t, 'ID_DECODIFICA_ANSC', (
    'NM_DECODIFICA_ANSC', 'VARCHAR2(50)',
    'Nome della tabella, da indicare soltanto per i due identificativi che ne coprono più di '
    'una. ⚠️ Vuoto negli altri casi: è il modo per non appesantire la configurazione con un '
    'dato che serve in due casi su centoquarantatré, senza per questo lasciare la lettura '
    'ambigua dove lo sarebbe.'))
fatti.append('chiave di dominio_decodifica e valore_dominio, NM_DECODIFICA_ANSC su CFG_CAMPO')


# ═════════════════════════════════════ 4. il DDL
sostituisci_blocco('CREATE TABLE ANSC_USR.DOMINIO_DECODIFICA (', [
    'CREATE TABLE ANSC_USR.DOMINIO_DECODIFICA (',
    '  id_dominio           NUMBER             NOT NULL,',
    '  nm_dominio           VARCHAR2(50 CHAR)  NOT NULL,',
    '  CONSTRAINT dominio_decodifica_pk PRIMARY KEY (id_dominio, nm_dominio)',
    ') TABLESPACE ANSC_USR;',
    'COMMENT ON TABLE ANSC_USR.DOMINIO_DECODIFICA IS',
    "  'Catalogo delle decodifiche ANSC replicate. Struttura ereditata da Side, con una",
    "   correzione: il nome entra nella chiave, perche nel corpus pubblicato gli identificativi",
    "   134 e 135 coprono due tabelle ciascuno e la sola chiave numerica non regge il carico.';",
])
sostituisci_blocco('CREATE TABLE ANSC_USR.VALORE_DOMINIO (', [
    'CREATE TABLE ANSC_USR.VALORE_DOMINIO (',
    '  id_dominio           NUMBER             NOT NULL,',
    '  nm_dominio           VARCHAR2(50 CHAR)  NOT NULL,',
    '  id_valore            VARCHAR2(10 CHAR)  NOT NULL,',
    '  ds_valore            CLOB,',
    '  dt_inizio_validita   DATE,',
    '  dt_fine_validita     DATE,',
    '  nr_ordinamento       NUMBER,',
    '  cd_valore            VARCHAR2(10 CHAR),',
    '  CONSTRAINT valore_dominio_pk PRIMARY KEY (id_dominio, nm_dominio, id_valore),',
    '  CONSTRAINT valore_dominio_tipo_dominio_fk',
    '    FOREIGN KEY (id_dominio, nm_dominio)',
    '    REFERENCES ANSC_USR.DOMINIO_DECODIFICA (id_dominio, nm_dominio)',
    ') TABLESPACE ANSC_USR;',
    '-- la lettura per chiave e servita dalla chiave primaria; l indice seguente serve la',
    '-- ricerca testuale del back-office, che non ha una colonna di appoggio.',
    'CREATE INDEX ANSC_USR.IX_VALORE_DOMINIO_RICERCA',
    '  ON ANSC_USR.VALORE_DOMINIO (id_dominio, cd_valore) TABLESPACE ANSC_USR;',
    'COMMENT ON COLUMN ANSC_USR.VALORE_DOMINIO.dt_fine_validita IS',
    "  'La validita va applicata in lettura quando si PROPONE un valore, non quando lo si",
    "   TRADUCE: i valori cessati restano perche servono a interpretare gli atti gia formati.';",
])
# la vista: la chiave completa, e il filtro di validità solo dove si propone
sostituisci_blocco('CREATE OR REPLACE VIEW ANSC_USR.V_ANSC_DIZ_VALIDO AS', [
    'CREATE OR REPLACE VIEW ANSC_USR.V_ANSC_DIZ_VALIDO AS',
    '  SELECT v.id_dominio, v.nm_dominio, v.id_valore, v.cd_valore,',
    '         v.ds_valore, v.nr_ordinamento,',
    "         CASE WHEN SYSDATE BETWEEN NVL(v.dt_inizio_validita, DATE '0001-01-01')",
    "                              AND NVL(v.dt_fine_validita,   DATE '9999-12-31')",
    "              THEN 'S' ELSE 'N' END AS fg_valido",
    '    FROM ANSC_USR.VALORE_DOMINIO     v',
    '    JOIN ANSC_USR.DOMINIO_DECODIFICA d',
    '      ON d.id_dominio = v.id_dominio AND d.nm_dominio = v.nm_dominio;',
    'COMMENT ON TABLE ANSC_USR.V_ANSC_DIZ_VALIDO IS',
    "  'Contratto di fruizione verso SIPO: si legge per chiave (id_dominio + id_valore, piu",
    "   nm_dominio per i due identificativi ambigui) e si raggiunge per sinonimo con grant di",
    "   sola lettura. La colonna fg_valido dice se il valore sia ancora proponibile: chi traduce",
    "   un codice legge comunque, chi propone una scelta filtra su fg_valido = S.';",
], fine='siano le tabelle sottostanti')
# la colonna nuova sulla configurazione dei campi
D.ddl(doc, mono('  ID_DECODIFICA_ANSC   NUMBER,')._p.getnext(),
      ['  NM_DECODIFICA_ANSC   VARCHAR2(50 CHAR),'])
fatti.append('Appendice A: chiavi dei dizionari, vista con fg_valido, NM_DECODIFICA_ANSC')


# 4b — le COMMENT superate, rimaste fuori dai blocchi sostituiti
def togli_commento(prima_riga_del_corpo, righe_corpo):
    """Elimina una COMMENT ereditata: la riga di intestazione più il corpo.

    ⚠️ Si parte dal corpo, che è univoco, e si risale all'intestazione: le intestazioni sono
    due, perché la COMMENT nuova ripete lo stesso oggetto.
    """
    corpo = [mono(prima_riga_del_corpo)._p]
    for _ in range(righe_corpo - 1):
        corpo.append(corpo[-1].getnext())
    assert Paragraph(corpo[-1], doc).text.rstrip().endswith("';"), \
        'la COMMENT da eliminare non finisce dove previsto: ' + prima_riga_del_corpo[:50]
    intestazione = corpo[0].getprevious()
    assert Paragraph(intestazione, doc).text.strip().startswith('COMMENT ON'), \
        'sopra il corpo non c’è l’intestazione attesa'
    elimina([intestazione] + corpo, 'COMMENT superata ' + prima_riga_del_corpo[:42])


togli_commento("  'Catalogo delle decodifiche ANSC replicate. Struttura ereditata da Side.", 3)
togli_commento("  'La validita va applicata in lettura: i valori cessati restano perche servono "
               "a interpretare", 2)


# ═════════════════════════════════════ 5. Appendice B e open point
t = [x for x in doc.tables
     if [c.text.strip() for c in x.rows[0].cells][:2] == ['Metodo e percorso', 'Identificativo']
     and any('dizionari' in r.cells[0].text for r in x.rows[1:])][0]
nuova = D.clona_riga(t, ('GET /ansc/v1/dizionari/{codDizionario}/valori/{codValore}',
                         'leggiValoreDizionario', 'Legge un valore per chiave, decodifica più '
                         'valore', 'AMMINISTRATORE, AUDITOR, SUPPORTO'))
riga_con(t, 'GET /ansc/v1/dizionari/valori')._tr.addnext(nuova._tr)

t = D.trova_tabella(doc, '#', 'tema', 'questione')
r = riga_con(t, 'OP-28')
D.riscrivi_cella(r.cells[2], 'Nel corpus pubblicato due identificativi corrispondono a più di '
                             'una tabella: il 134 a dec_dichiarante_trascr_nascita e a '
                             'dec_dichiarante_trascr_postuma, che sono tabelle distinte, e il '
                             '135 a due file che differiscono per un refuso nel nome. Sul '
                             'nostro modello la questione è risolta dalla v3.25: il nome della '
                             'tabella entra nella chiave della replica e la configurazione del '
                             'campo lo indica dove serve, così che il carico completo non '
                             'violi la chiave e la lettura per chiave non sia ambigua. ⚠️ Resta '
                             'da chiarire con il fornitore come il modello dell’evento '
                             'discrimini le due tabelle, che referenzia in modo condizionato al '
                             'caso d’uso: è la sola parte che non possiamo decidere noi.')
fatti.append('Appendice B e OP-28 aggiornati')


# ═════════════════════════════════════ 6. figura
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import diagrammi_erd as E   # noqa: E402
E.TABELLE = [
    (t if t[4] != 'DOMINIO_DECODIFICA' else
     (t[0], t[1], t[2], t[3], t[4], [('PK', 'id_dominio + nm_dominio')], t[6]))
    for t in E.TABELLE]
E.TABELLE = [
    (t if t[4] != 'VALORE_DOMINIO' else
     (t[0], t[1], t[2], t[3], t[4],
      [('PK', 'id_dominio + nm_dominio'), ('', '+ id_valore'), ('', 'ds_valore, validità'),
       ('', 'nr_ordinamento, cd_valore')], t[6]))
    for t in E.TABELLE]
E.disegna()
D.sostituisci_immagine(doc, 'Schema ANSC_USR.', os.path.join(IMG, 'erd_ansc_usr.png'))
fatti.append('ERD: chiave dei dizionari aggiornata')

doc.save(DST)

# ────────────────────────────────────────────── controlli
import zipfile   # noqa: E402
ncom = zipfile.ZipFile(DST).read('word/comments.xml').decode().count('<w:comment ')
assert ncom == 27, f'commenti persi: {ncom}'
print('\n'.join(' · ' + f for f in fatti))
print('commenti:', ncom, '· capitoli/tabelle/immagini:', D.riepilogo(DST))
print('scritto:', os.path.relpath(DST, BASE))
