# -*- coding: utf-8 -*-
"""Le tre voci di chiusura del capitolo notifiche: lead-in in grassetto separato dal corpo."""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D  # noqa: E402
from v3_20_notifiche import DOC, OP_NUOVI  # noqa: E402

doc = docx.Document(DOC)
atteso = {c: (tema, q) for c, tema, q, _, _, _ in OP_NUOVI}
fatti = 0
for p in doc.paragraphs:
    t = p.text.strip()
    for cod, (tema, q) in atteso.items():
        if t.startswith(cod + '. '):
            prima = q.split('. ')[0] + '.'
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            for txt, bold in D.segmenta(f'{cod} — {tema}. ', prima):
                p.add_run(txt).bold = bold
            fatti += 1
assert fatti == 3, fatti
doc.save(DOC)
print('voci sistemate:', fatti)
