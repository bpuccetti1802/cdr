# -*- coding: utf-8 -*-
"""ANALISI_Identita-Profilazione-IAM v0.5 → v0.6 (13/10/2026).

Revisione del capitolo «Scelte effettuate», chiesta dal committente su tre punti:

  1. la soluzione scelta deve valere **per entrambe le generazioni di front-end** — le
     trentacinque applicazioni Thymeleaf da ritoccare e le nuove in Angular — e ciò che
     non vale per una delle due va **escluso dichiarando il motivo**;
  2. il **passaggio del token nelle intestazioni** va spiegato meglio, tratta per tratta;
  3. va spiegato **come si armonizzano** i due mondi, anche attraverso il BFF già
     introdotto nel capitolo delle soluzioni, con il come si fa e le buone pratiche.

Tre sottocapitoli nuovi, uno riscritto, e due punti aperti.

⚠️ Il capitolo contiene prosa riscritta dall'utente in Word: si interviene per
sottocapitoli interi dove il ragionamento cambia e **non** si tocca il resto. Tre refusi
della stesura manuale si correggono (dovà, Andrebber, «tiene può tenere»).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_6_iam_armonia.py
"""
import os
import shutil
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.5.docx')
DST = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.6.docx')


# --------------------------------------------------------------------- contenuti

DUE_MONDI = [
    ['', 'I trentacinque front-end Thymeleaf', 'Le applicazioni Angular nuove'],
    ['Chi autentica la persona',
     'sipo-auth, raggiunto con un reindirizzamento dal front-end.',
     'sipo-auth, raggiunto con un reindirizzamento dalla shell.'],
    ['Dove sta il token',
     'Nella **sessione lato server** del front-end, accanto all’oggetto utente. Il '
     'browser non lo vede mai.',
     'Nel **BFF**, che tiene la sessione. Il browser non lo vede mai.'],
    ['Che cosa porta il browser',
     'Il solo identificativo di sessione.',
     'Il solo cookie di sessione del BFF.'],
    ['Come il token raggiunge i back-end',
     '`Authorization: Bearer`, messa dal componente di accesso ai servizi: **la stessa '
     'intestazione di oggi**.',
     '`Authorization: Bearer`, messa dal BFF.'],
    ['Che cosa va ritoccato',
     'La libreria di profilazione e quella di accesso ai servizi. **Nessuna delle 4.781 '
     'chiamate applicative.**',
     'La shell e la libreria Angular condivisa: il BFF sostituisce la conservazione del '
     'gettone nel browser.'],
    ['Che cosa NON funziona, e perché',
     'Che ciascun front-end diventi client OIDC per conto proprio: sarebbero '
     'trentacinque rilasci e trentacinque regressioni, e il punto da cambiare sono le '
     'librerie, non i punti di chiamata.',
     'Che il gettone resti nel browser: **è ciò che si fa oggi** (`localStorage`, chiave '
     '`auth`) ed è l’unica cosa che il BFF esiste per togliere.'],
]

ESCLUSIONI = [
    ['Che cosa si esclude', 'Motivo', 'Che cosa si fa invece'],
    ['Due emittenti di token, uno per mondo',
     'Il token si inoltra da back-end a back-end e si verifica con le chiavi pubbliche '
     'dell’emittente: due emittenti vorrebbero dire due catene di fiducia, due registri '
     'di accesso e due insiemi di claim da tenere allineati. E i due mondi devono '
     'decidere sugli **stessi** claim, perché le regole dei back-end e le abilitazioni '
     'delle pagine del back-office parlano degli stessi ruoli.',
     'Un solo emittente, sipo-auth, e un solo formato di token.'],
    ['Il gettone nel browser per le applicazioni a pagina singola',
     'Ogni luogo in cui il browser può conservarlo è leggibile da chiunque ottenga '
     'esecuzione di script nella pagina. Non è un rischio ipotetico: è la pratica '
     'corrente della libreria Angular.',
     'Il BFF tiene il gettone lato server; al browser resta un cookie di sessione.'],
    ['L’adozione di OIDC front-end per front-end',
     'Le 4.781 chiamate che dipendono dal profilo non si toccano una per una, e un '
     'rilascio per applicazione moltiplica le regressioni senza anticipare nessun '
     'beneficio.',
     'Il cambiamento avviene dentro le due librerie condivise, a parità di firma.'],
    ['La modalità header per le chiamate che il **browser** fa ai back-end',
     'Il browser di una applicazione a pagina singola chiama direttamente i servizi: su '
     'quelle chiamate l’Access Manager non inietta nulla se il percorso non è protetto, '
     'e non c’è modo di farsi dare l’identità dal browser senza fidarsi del browser.',
     'Le chiamate passano dal BFF, che è lato server e riceve le intestazioni su ogni '
     'richiesta. ⚠️ Esiste però un’alternativa prevista dal manuale e da verificare: il '
     'dominio dedicato ai servizi di back-end, che espone le stesse variabili di '
     'intestazione purché il client inoltri i cookie (PI-30).'],
]

TRATTE = [
    ['Tratta', 'Che cosa viaggia', 'In quale intestazione', 'Chi la verifica'],
    ['Browser → perimetro (portale)',
     'Nulla che riguardi l’identità applicativa.',
     'Il cookie di sessione del portale.',
     'L’Access Manager di IAM.'],
    ['Perimetro → front-end Thymeleaf',
     'L’identità della persona, in chiaro.',
     '`iv-user`, `iv-codfis`, `iv-nome`, `iv-cognome`, `sysgroup`, e le altre di '
     'Tabella 2 che la policy concede.',
     '⚠️ **Nessuno**: il front-end si fida. È il difetto che la soluzione chiude.'],
    ['sipo-auth → front-end Thymeleaf (sequenza B)',
     'Il token SIPO, oltre alle intestazioni di oggi.',
     '**`X-Sipo-Token`**: un nome proprio e riservato, non una `iv-*`.',
     'Il front-end, con le chiavi pubbliche di sipo-auth.'],
    ['Browser → BFF (mondo nuovo)',
     'Nulla che riguardi l’identità.',
     'Il cookie di sessione del BFF, `Secure`, `HttpOnly`, `SameSite`.',
     'Il BFF, sulla propria sessione.'],
    ['Front-end Thymeleaf → back-end',
     'Il token SIPO dell’operatore.',
     '**`Authorization: Bearer`** — la stessa intestazione di oggi.',
     'Il back-end, con le chiavi pubbliche.'],
    ['BFF → back-end',
     'Il token SIPO dell’operatore.',
     '**`Authorization: Bearer`**.',
     'Il back-end, con le chiavi pubbliche.'],
    ['Back-end → back-end',
     'Lo stesso token, inoltrato senza riemetterlo.',
     '**`Authorization: Bearer`**.',
     'Il back-end chiamato.'],
    ['Trattamenti pianificati',
     'Un’utenza tecnica dedicata, non una persona.',
     '**`Authorization: Bearer`**.',
     'Il back-end.'],
]

IGIENE = [
    ['Regola', 'Perché, e che cosa accade se salta'],
    ['Il perimetro **cancella** le intestazioni di identità che arrivano dal client',
     'Vale sia per le `iv-*` sia per `X-Sipo-Token`. Se il perimetro inoltra ciò che il '
     'browser gli manda, chiunque si dichiara chi vuole: la soluzione non migliora di '
     'nulla il punto di partenza.'],
    ['I front-end accettano traffico **solo** dal perimetro',
     'Altrimenti la regola precedente si scavalca raggiungendo il pod per un’altra '
     'strada. È l’accertamento da cui dipende la scelta fra le due sequenze.'],
    ['Il token non va mai nella stringa di interrogazione',
     'Gli indirizzi finiscono nei registri di accesso del proxy e nel riferimento di '
     'provenienza del browser. Un token in un registro è un token compromesso.'],
    ['Il token non va in un cookie leggibile da script',
     'Un cookie senza `HttpOnly` è esattamente l’archivio del browser da cui il BFF '
     'serve ad uscire.'],
    ['Il nome dell’intestazione è riservato e documentato',
     'Un nome generico come `X-Token` prima o poi collide con quello di un altro '
     'componente della catena, e la collisione si manifesta come un accesso negato '
     'intermittente.'],
]

DIMENSIONE = [
    ['Quanto pesa', 'Conseguenza'],
    ['Un token con identità, organizzazione, postazione e un ruolo: **1-2 KB**',
     'Nessun problema: resta ampiamente sotto i limiti consueti dei proxy.'],
    ['Lo stesso token con l’elenco completo delle abilitazioni',
     '⚠️ Cresce in proporzione alle righe che l’utente ha nella tabella dei permessi. '
     'Alcune decine sono innocue, alcune centinaia no.'],
    ['Oltre il limite dell’intestazione (tipicamente **8 KB**)',
     '⚠️ Il proxy risponde con un errore su **tutte** le chiamate di quell’utente, non '
     'solo su quelle grandi: il guasto è per persona e sembra inspiegabile.'],
]

ARMONIA = [
    ['Mondo', 'Chi tiene la sessione', 'Chi custodisce il token', 'Che cosa è, in sostanza'],
    ['Thymeleaf, oggi e domani',
     'Il front-end stesso, lato server.',
     'Il front-end, nella sessione.',
     '**È già un BFF**, per costruzione: la pagina la compone il server e il gettone non '
     'lascia la macchina.'],
    ['Angular, oggi',
     'Nessuno: lo stato sta nel browser.',
     '⚠️ Il browser (`localStorage`).',
     'Il caso che il BFF esiste per correggere.'],
    ['Angular, domani',
     'Il BFF, uno per shell.',
     'Il BFF.',
     'La stessa proprietà del mondo Thymeleaf, ottenuta con un componente in più perché '
     'una pagina che gira nel browser non ha un server proprio.'],
]

BUONE_PRATICHE = [
    ['Punto', 'Come si fa', 'Perché'],
    ['Granularità',
     'Un BFF **per shell**, non per micro-frontend.',
     'La sessione è della shell; quattro BFF per quattro remote darebbero quattro '
     'sessioni da tenere allineate e nessun vantaggio.'],
    ['Il cookie',
     '`Secure`, `HttpOnly`, `SameSite=Lax`, `Path` sul contesto della shell, nessun '
     'dominio più ampio del necessario.',
     '`HttpOnly` è ciò che tiene il gettone fuori dalla portata degli script; `Lax` '
     'serve perché il ritorno dal reindirizzamento di autenticazione è una navigazione '
     'proveniente da un altro sito.'],
    ['Protezione dalle richieste indotte',
     'Un gettone anti-CSRF (doppio invio o intestazione dedicata), oppure '
     '`SameSite=Strict` più la verifica dell’origine.',
     '⚠️ **È il difetto che il BFF introduce**, e viene dimenticato di regola: quando '
     'l’autorizzazione viaggia in un cookie, il browser la allega anche a una richiesta '
     'partita da un altro sito. Con il gettone in un’intestazione il problema non '
     'esisteva.'],
    ['Perimetro di inoltro',
     'Il BFF inoltra **solo i percorsi dichiarati**, non tutto ciò che riceve.',
     'Un proxy generico autenticato è una via di accesso a qualunque servizio '
     'raggiungibile dalla sua rete.'],
    ['Sessione fuori dalla memoria',
     'Un archivio condiviso fra le repliche.',
     'Con la sessione in memoria due repliche producono uscite impreviste. È il rilievo '
     'che il documento già segnala per i front-end: la costruzione del BFF è '
     'l’occasione per non ereditarlo.'],
    ['Il manifesto dei micro-frontend',
     'Servito dalla stessa origine del BFF, o con origine e cookie allineati.',
     'Se il manifesto e i remote arrivano da un’altra origine le chiamate diventano '
     'inter-origine, e un cookie `SameSite` non parte: la shell si carica e le chiamate '
     'tornano non autorizzate.'],
    ['Prontezza',
     'La sonda del BFF non dipende dalla raggiungibilità di IAM.',
     'Altrimenti un disservizio esterno fa riavviare un componente sano, e il ciclo di '
     'riavvii peggiora il disservizio.'],
    ['Interfaccia commutabile',
     'Separare «costruisci il reindirizzamento» da «gestisci il ritorno» dal primo '
     'giorno.',
     'È ciò che consente di partire in modalità header e passare a OIDC cambiando un '
     'parametro, invece di riscrivere il componente.'],
]

PUNTI_NUOVI = [
    ('PI-29', 'Quale protezione dalle richieste indotte adotta il BFF: un gettone '
              'anti-CSRF a doppio invio, un’intestazione dedicata, oppure cookie '
              'rigorosamente di primo sito con verifica dell’origine. La questione nasce '
              'con il BFF e non esisteva finché l’autorizzazione viaggiava in '
              'un’intestazione: va decisa prima di scrivere la shell, perché cambia il '
              'contratto fra shell e BFF.',
     'Architettura / Referenti SIPO'),
    ('PI-30', 'Se i back-end di SIPO possano essere esposti sul dominio che l’ente '
              'dedica ai servizi di back-end e alle applicazioni a pagina singola. Il '
              'manuale dichiara che su quel dominio sono disponibili le stesse variabili '
              'di intestazione, a condizione che le richieste del client inoltrino i '
              'cookie: se la risposta è positiva, il mondo nuovo può partire in modalità '
              'header come il vecchio, senza attendere le decisioni su OIDC.',
     'Presidio IAM / Sistemi'),
]


def main():
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    modello = d.tables[3]
    fatti = []

    def inserisci(prima, blocchi):
        """Innesta i blocchi prima dell'elemento dato. ('p'|'b'|'v'|'h2'|'tab')."""
        for tipo, contenuto in blocchi:
            if tipo == 'p':
                D.para(d, prima, contenuto)
            elif tipo == 'b':                      # paragrafo con grassetti
                D.para(d, prima, pezzi=D.segmenta('', contenuto))
            elif tipo == 'v':
                D.voce(d, prima, *contenuto)
            elif tipo == 'h2':
                D.para(d, prima, contenuto, stile='Heading 2')
            elif tipo == 'tab':
                righe, larghezze = contenuto
                D.tabella(d, prima, righe, modello, larghezze)

    def ancora(titolo, livello=2):
        """L'elemento XML dell'intestazione di una sezione: ci si innesta prima."""
        return D.h(d, livello, titolo)._p

    # ============================================================ 1. i due mondi
    inserisci(ancora('Perché non si comincia dalla libreria'), [
        ('h2', 'Un solo impianto per due generazioni di front-end'),
        ('b', 'Il sistema avrà per anni **due generazioni di front-end**: le trentacinque '
              'applicazioni Thymeleaf, che si ritoccano e restano, e le nuove applicazioni '
              'Angular, che nascono adesso. La soluzione scelta è valida **solo se regge su '
              'entrambe**, e questo sottocapitolo lo verifica voce per voce invece di darlo '
              'per scontato: un impianto che funzionasse solo per il mondo nuovo '
              'lascerebbe la parte più grande del sistema dov’è oggi, e uno che funzionasse '
              'solo per il vecchio obbligherebbe il back-office a costruirsi un accesso per '
              'conto proprio — cioè il secondo impianto che si vuole evitare.'),
        ('tab', (DUE_MONDI, [1.4, 2.6, 2.6])),
        ('b', 'La lettura della tabella è che **l’impianto è lo stesso e cambia soltanto chi '
              'custodisce il token**: nel mondo Thymeleaf la sessione del front-end, nel '
              'mondo Angular il BFF. Tutto il resto — l’emittente, il formato del token, '
              'l’intestazione con cui raggiunge i back-end, la verifica a valle — è '
              'identico. Non è una coincidenza fortunata: è la proprietà che si è scelto '
              'di preservare, e va difesa a ogni decisione successiva.'),
        ('p', 'Quattro cose, invece, non si possono fare: vanno escluse adesso e con il '
              'motivo scritto, perché sono le quattro strade che a prima vista sembrano '
              'più brevi.'),
        ('tab', (ESCLUSIONI, [1.6, 3.0, 2.0])),
    ])
    fatti.append('nuovo sottocapitolo «Un solo impianto per due generazioni di front-end»')

    # ============================================================ 2. il token nelle intestazioni
    inserisci(ancora('Scadenza, uscite, indisponibilità'), [
        ('h2', 'Il token nelle intestazioni, tratta per tratta'),
        ('b', 'Il fatto che rende l’intervento proporzionato è questo: **l’intestazione non '
              'cambia, cambiano chi firma e che cosa c’è dentro**. Oggi i back-end ricevono '
              'già `Authorization: Bearer` con un token firmato a chiave simmetrica il cui '
              'soggetto è un’utenza tecnica; domani ricevono la stessa intestazione, con un '
              'token firmato a chiave asimmetrica il cui soggetto è una persona. Il '
              'meccanismo di trasporto è identico, ed è il motivo per cui le 4.781 chiamate '
              'applicative non si toccano: cambia la provenienza del token dentro una sola '
              'libreria.'),
        ('p', 'Nel resto della catena le intestazioni sono invece diverse fra loro, e '
              'confonderle è l’errore che fa fallire le prove di integrazione. Ecco che '
              'cosa viaggia su ciascuna tratta.'),
        ('tab', (TRATTE, [1.7, 1.9, 2.0, 1.6])),
        ('b', 'Due osservazioni sulla tabella. La prima è che **le intestazioni `iv-*` e il '
              'token non sono alternativi ma sovrapposti** durante la transizione: in '
              'sequenza B il front-end riceve le prime, già valorizzate da dati verificati, '
              'e accanto il token da spendere a valle. La seconda è che **il token non va '
              'mai messo in una `iv-*`**: quelle appartengono all’Access Manager, e il '
              'perimetro le riscrive — un token messo lì verrebbe sovrascritto o, peggio, '
              'accettato quando arriva dal client.'),
        ('p', 'Da qui cinque regole di igiene del perimetro. Non sono raccomandazioni: sono '
              'le condizioni perché il trasporto in intestazione sia sicuro, e ciascuna ha '
              'un modo preciso di fallire.'),
        ('tab', (IGIENE, [2.6, 4.0])),
        ('b', 'Resta un vincolo di misura che si scopre tardi se non lo si mette in conto. '
              'Il manuale dell’ente chiede esplicitamente di ridurre le variabili di '
              'intestazione «a quelle strettamente indispensabili» **per limitarne la '
              'dimensione**, e un token è a tutti gli effetti una variabile di '
              'intestazione.'),
        ('tab', (DIMENSIONE, [3.0, 3.6])),
        ('b', 'La conseguenza di disegno è che **la decisione su dove stiano le abilitazioni '
              'va presa su una misura, non su un principio**. Se un utente ha poche decine '
              'di permessi, portarli nel token è semplice e va bene; se ne ha centinaia, '
              'nel token restano identità, organizzazione, postazione, ruoli e profilo, e '
              'le abilitazioni si leggono una volta per sessione da un servizio dedicato, '
              'tenendole poi nella sessione del front-end o del BFF. La misura è una sola '
              'interrogazione sulla tabella dei permessi — quante righe ha l’utente con più '
              'righe — e va fatta prima di fissare il contenuto del token.'),
    ])
    fatti.append('nuovo sottocapitolo «Il token nelle intestazioni, tratta per tratta»')

    # ============================================================ 3. armonizzare col BFF
    inserisci(ancora('Scadenza, uscite, indisponibilità'), [
        ('h2', 'Armonizzare il vecchio e il nuovo: che cosa fa il BFF'),
        ('b', 'Il capitolo delle soluzioni ha introdotto il Backend For Frontend come '
              'variante del gateway. Letto dal punto di vista dell’armonizzazione, il BFF '
              'si capisce meglio così: **è il nome di ciò che le applicazioni Thymeleaf già '
              'fanno**. Una pagina composta dal server ha una sessione lato server, e il '
              'gettone non lascia la macchina; il browser porta soltanto un identificativo '
              'di sessione. Un’applicazione che gira nel browser non ha un server proprio, '
              'e per avere la stessa proprietà deve prenderselo in prestito: quel server è '
              'il BFF.'),
        ('b', 'Armonizzare i due mondi non vuol dire quindi rendere il vecchio simile al '
              'nuovo. Vuol dire **dare al nuovo la proprietà che il vecchio ha per '
              'costruzione**, ed enunciarla come un vincolo valido per tutti: *il gettone '
              'non attraversa mai il browser; il browser porta solo un cookie o un '
              'identificativo di sessione*. È una frase che si può verificare con gli '
              'strumenti di sviluppo del browser, e per questo è un buon vincolo.'),
        ('tab', (ARMONIA, [1.4, 1.6, 1.5, 2.3])),
        ('b', 'Ne segue una risposta e una conseguenza. La risposta: **i trentacinque '
              'front-end non hanno bisogno di un BFF**, perché sono già server; l’unica cosa '
              'che manca loro è che il token nella sessione **ci sia**, ed è il lavoro di '
              'sipo-auth e della libreria di accesso ai servizi. La conseguenza, meno '
              'evidente: **la sequenza B è il BFF applicato al mondo vecchio**. Mettendo '
              'sipo-auth davanti ai front-end si fa per loro esattamente il mestiere del '
              'BFF — terminare l’autenticazione e consegnare a valle un’identità '
              'verificata. Le due sequenze non sono una per mondo: sono due modi di '
              'collocare lo stesso componente.'),
        ('p', 'Sul come si costruisce, otto punti. I primi due sono di impianto, il terzo è '
              'quello che si dimentica, gli altri si pagano in esercizio.'),
        ('tab', (BUONE_PRATICHE, [1.5, 2.5, 2.6])),
        ('b', 'Un’ultima avvertenza, perché è il modo tipico in cui questo disegno degenera: '
              '**il BFF non è il luogo dove mettere la logica applicativa.** Autentica, '
              'custodisce la sessione, allega il token e inoltra; non traduce payload, non '
              'aggrega risposte di comodo, non prende decisioni di dominio. Appena comincia '
              'a farlo diventa un secondo back-end non dichiarato, con una propria '
              'versione della verità e nessuna delle garanzie di quelli veri.'),
    ])
    fatti.append('nuovo sottocapitolo «Armonizzare il vecchio e il nuovo: che cosa fa il BFF»')

    # ============================================================ 4. chiamate ai back-end
    D.sostituisci_sezione(
        d, 'Le chiamate verso i back-end', 'Le chiamate verso i back-end',
        [('p', 'Dal front-end al back-end. Il codice applicativo chiama il componente di '
               'accesso ai servizi come oggi, con la stessa firma e lo stesso parametro: '
               'nessuna delle chiamate esistenti cambia. È dentro quel componente che il '
               'gettone dell’utenza tecnica viene sostituito dal token SIPO trovato nella '
               'sessione corrente, rinnovandolo prima se sta per scadere. Il back-end '
               'verifica firma, emittente, scadenza e destinatario, concede l’autorità '
               'corrispondente al profilo e legge chi opera dal token, non dal corpo della '
               'richiesta.'),
         ('p', 'Dalla shell Angular al back-end il percorso è lo stesso con un passaggio in '
               'più: la chiamata parte dal browser verso il BFF, che vi allega il token e '
               'la inoltra. Per il back-end le due provenienze sono indistinguibili, ed è '
               'esattamente l’effetto che si cerca: un solo modo di autorizzare, due modi '
               'di arrivarci.'),
         ('p', 'Da back-end a back-end lo stesso token viene inoltrato, così il servizio che '
               'dialoga con ANPR può dichiarare l’operatore e la postazione reali. I '
               'processi senza utente — i trattamenti pianificati — non hanno una richiesta '
               'in corso: usano un’utenza tecnica dedicata al trattamento, non il codice '
               'fiscale di una persona, con il gettone tenuto in memoria.'),
         ('p', 'Il sottocapitolo seguente riporta, tratta per tratta, quale intestazione '
               'porta che cosa, le regole di igiene del perimetro e il vincolo di '
               'dimensione: è la parte che serve a chi deve configurare il proxy e scrivere '
               'il client, e per questo sta separata dal racconto.')])
    fatti.append('riscritta la sezione «Le chiamate verso i back-end» (aggiunta la tratta del BFF)')

    # ============================================================ 5. sostenibilità: rimando
    D.sostituisci(
        d,
        'La soluzione è stata concepita a partire dal mondo Java, che è la parte più '
        'grande del sistema. Reggerà anche sul front-end Angular e la variante del '
        'Backend For Frontend è la strada da percorrere: le applicazioni a pagina singola '
        'non possono usare le intestazioni, e tenere il gettone nel browser è ciò che il '
        'BFF evita. Ci sono però alcuni punti da definire.',
        'La soluzione è stata concepita a partire dal mondo Java, che è la parte più '
        'grande del sistema. Che regga anche sul front-end Angular è verificato voce per '
        'voce nel sottocapitolo «Un solo impianto per due generazioni di front-end», e il '
        'modo di armonizzare i due mondi — con il Backend For Frontend — è descritto in '
        '«Armonizzare il vecchio e il nuovo». Restano i punti da definire qui sotto.',
        attese=1, etichetta='rimando ai sottocapitoli nuovi', fatti=fatti)

    # ============================================================ 6. refusi
    for vecchio, nuovo in [('che dovà dialogare', 'che dovrà dialogare'),
                           ('Andrebber quindi', 'Andrebbero quindi'),
                           ('il BFF tiene può tenere', 'il BFF può tenere')]:
        D.sostituisci(d, vecchio, nuovo, attese=1, etichetta='refuso', fatti=fatti)

    # ============================================================ 7. intestazioni vuote
    # ⚠️ Due intestazioni senza testo e senza run (una di primo e una di secondo
    # livello, fra il Glossario e l'Executive summary) producevano **due righe bianche
    # nell'indice**: l'indice è generato dagli stili, non dal testo, quindi un heading
    # vuoto è una voce vuota. Si rimuovono con il controllo sui commenti, per non
    # cancellare un'ancora.
    vuote = 0
    for p in list(d.paragraphs):
        if p.style.name.startswith('Heading') and not p.text.strip() and not p.runs:
            if D.ha_commenti(p._p):
                raise SystemExit('intestazione vuota con un commento: non si cancella')
            p._p.getparent().remove(p._p)
            vuote += 1
    if vuote:
        fatti.append('rimosse %d intestazioni vuote (righe bianche nell\'indice)' % vuote)

    # ============================================================ 8. punti aperti
    tabella_punti = D.trova_tabella(d, '#', 'questione', 'a chi compete')
    for codice, testo, compete in PUNTI_NUOVI:
        D.clona_riga(tabella_punti, (codice, testo, compete))
    fatti.append('aggiunti %d punti aperti' % len(PUNTI_NUOVI))

    # ============================================================ 9. frontespizio e storia
    for fatto in D.frontespizio(d, '0.6', 'ANALISI_Identita-Profilazione-IAM_v0.6'):
        fatti.append(fatto)

    D.storia(d, '13/10/2026', '0.6',
             'Scelte effettuate (tre sottocapitoli) · Punti aperti',
             'Rivisto il capitolo delle scelte su tre punti. Reso esplicito che '
             'l’impianto deve valere per entrambe le generazioni di front-end, con il '
             'confronto voce per voce e quattro esclusioni motivate. Aggiunta la '
             'descrizione del passaggio del token nelle intestazioni tratta per tratta, '
             'con le cinque regole di igiene del perimetro e il vincolo di dimensione, da '
             'cui segue che la collocazione delle abilitazioni va decisa su una misura. '
             'Aggiunto il sottocapitolo sull’armonizzazione attraverso il BFF: il mondo '
             'Thymeleaf è già un BFF per costruzione, la sequenza B è il BFF applicato al '
             'vecchio, e otto buone pratiche di costruzione. Nuovi PI-29 (protezione dalle '
             'richieste indotte) e PI-30 (dominio dei servizi di back-end).')

    d.save(DST)

    cap, tab, fig = D.riepilogo(DST)
    print('scritto %s' % os.path.basename(DST))
    for f in fatti:
        print('  ' + str(f))
    print('  capitoli %d | tabelle %d | figure %d' % (cap, tab, fig))
    print('  punti aperti: %d' % (len(tabella_punti.rows) - 1))


if __name__ == '__main__':
    main()
