# -*- coding: utf-8 -*-
"""v3.14 — quello che la verifica di congruenza sui capp. 8-10 ha trovato.

Quattro tabelle esistevano **solo come DDL**, senza che il corpo del documento ne descrivesse
le colonne: `ANSC_CFG_REGOLA`, `ANSC_CFG_REGOLA_CONDIZIONE`, `ANSC_ANA_UC` e
`ANSC_CFG_VERSIONE`. Le prime due sono un buco aperto da questa stessa versione — sostituendo
«La forma della regola» con «La forma della condizione» ho tolto la descrizione della regola
senza rimpiazzarla per le due specie che restano.

⚠️ E il DDL di `ANSC_CFG_REGOLA` ammetteva ancora `'DETERMINAZIONE'` fra i tipi, con un
vincolo che pretendeva `COD_UC_ANSC` per quel caso: contraddiceva PC-9 in modo eseguibile, che
è la specie peggiore di contraddizione, perché il database l'avrebbe imposta.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402
from v3_14_configurazione import riscrivi_tabella   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.14.docx')

REGOLA = [
    ['Colonna', 'Tipo', 'Note'],
    ['ID_REGOLA', 'NUMBER (PK)', 'Chiave tecnica.'],
    ['COD_REGOLA', 'VARCHAR2(40)', 'Nome parlante della regola, univoco nella baseline.'],
    ['COD_TIPO_REGOLA', 'VARCHAR2(20)',
     'CONTROLLO o GENERAZIONE. ⚠️ La terza specie, DETERMINAZIONE, non esiste più: dalla v3.14 '
     'l’UC è scelto dalle condizioni di applicabilità di ANSC_CFG_UC (PC-9).'],
    ['COD_AREA', 'VARCHAR2(50)', 'Area applicativa di SIPO a cui la regola si riferisce.'],
    ['ID_MODELLO_ATTO', 'VARCHAR2(5)', 'Modello di atto su cui la regola interviene.'],
    ['ID_CONF_TIPO_ATTO', 'NUMBER', 'Tipo atto, quando la regola distingue le varianti.'],
    ['COD_UC_ANSC', 'VARCHAR2(20)',
     'UC a cui la regola si limita, quando non vale per tutti quelli del Modello. Non è più '
     'l’esito della regola ma il suo ambito.'],
    ['COD_ERRORE', 'VARCHAR2(40)', 'Codice dell’esito negativo, per le regole di controllo.'],
    ['MESSAGGIO', 'VARCHAR2(400)', 'Testo mostrato all’operatore quando il controllo fallisce.'],
    ['COD_CAMPO_DESTINAZIONE', 'VARCHAR2(120)',
     'Campo che la regola valorizza, per le regole di generazione.'],
    ['TXT_ESPRESSIONE', 'VARCHAR2(400)', 'Come si calcola il valore, per le generazioni.'],
    ['NUM_PRIORITA', 'NUMBER', 'Ordine di valutazione fra le regole dello stesso ambito.'],
    ['FLG_MANDATORY / FLG_AGGREGATE', 'CHAR(1)',
     'Composizione delle condizioni, con la semantica del motore già presente in SIPO: '
     'congiunzione obbligatoria e aggregazione.'],
    ['FLG_ABILITATA', 'CHAR(1)', 'S/N: consente di sospendere una regola senza cancellarla.'],
    ['DESCRIZIONE', 'VARCHAR2(400)', 'Che cosa la regola stabilisce, in lingua corrente.'],
    ['ID_VERSIONE', 'NUMBER (FK)', 'Baseline di appartenenza (ANSC_CFG_VERSIONE).'],
    ['DATA_INIZIO_VALIDITA / DATA_FINE_VALIDITA', 'DATE', 'Validità temporale.'],
    ['DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici previsti dallo standard di nomenclatura [R5].'],
]

CONDIZIONE = [
    ['Colonna', 'Tipo', 'Note'],
    ['ID_REGOLA_CONDIZIONE', 'NUMBER (PK)', 'Chiave tecnica.'],
    ['ID_REGOLA', 'NUMBER (FK)', 'Regola a cui la condizione appartiene.'],
    ['NUM_ORDINE', 'NUMBER', 'Ordine di valutazione fra le condizioni della stessa regola.'],
    ['CAMPO_SIPO', 'VARCHAR2(120)', 'Campo osservato, nella forma TABELLA.COLONNA.'],
    ['COD_OPERATORE', 'VARCHAR2(20)',
     'UGUALE, DIVERSO, IN, NON_IN, VUOTO, NON_VUOTO, MAGGIORE, MINORE.'],
    ['VALORE', 'VARCHAR2(400)', 'Termine di confronto; assente per VUOTO e NON_VUOTO.'],
    ['COD_DECODIFICA_ANSC', 'VARCHAR2(20)',
     'Dizionario con cui validare il valore, quando il campo è codificato.'],
    ['DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici previsti dallo standard di nomenclatura [R5].'],
]

ANA_UC = [
    ['Colonna', 'Tipo', 'Note'],
    ['ID_UC', 'NUMBER (PK)', 'Chiave tecnica locale.'],
    ['COD_UC_ANSC', 'VARCHAR2(20)',
     'Codice numerico pubblicato da ANSC (es. 11111000). ⚠️ Non è opaco: è la sequenza delle '
     'scelte che compongono il caso d’uso, e la web app di ANSC lo costruisce cifra per cifra '
     'mentre l’operatore risponde alle domande.'],
    ['COD_MOTORE', 'VARCHAR2(40)', 'Codice parlante del motore (es. Dic_Nasc_001).'],
    ['DESCRIZIONE', 'VARCHAR2(400)', 'Denominazione pubblicata da ANSC.'],
    ['COD_FAMIGLIA', 'VARCHAR2(30)', 'Famiglia di evento: nascita, morte, matrimonio, …'],
    ['COD_CATEGORIA', 'VARCHAR2(30)', 'Sottoinsieme della famiglia, dove ANSC lo dichiara.'],
    ['COD_VERSIONE_MAPPING', 'VARCHAR2(20)',
     'Revisione del mapping da cui la riga è stata replicata.'],
    ['DATA_INIZIO_VALIDITA / DATA_FINE_VALIDITA', 'DATE',
     'Validità dichiarata da ANSC. ⚠️ ANSC non cancella un UC: gli chiude la validità — sono i '
     'tre casi ritirati, 374 validi su 377 righe del catalogo.'],
    ['DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici previsti dallo standard di nomenclatura [R5].'],
]


def applica(doc):
    fatti = []
    modello = D.trova_tabella(doc, 'Colonna', 'Tipo', 'Note')

    # ------------------------------------------- il DDL che contraddiceva PC-9
    D.sostituisci(doc,
                  "CHECK (COD_TIPO_REGOLA IN ('DETERMINAZIONE','CONTROLLO','GENERAZIONE')),",
                  "CHECK (COD_TIPO_REGOLA IN ('CONTROLLO','GENERAZIONE')),",
                  attese=1, etichetta='vincolo sui tipi di regola', fatti=fatti)
    D.sostituisci(doc,
                  "CHECK ( (COD_TIPO_REGOLA = 'DETERMINAZIONE' AND COD_UC_ANSC IS NOT NULL)",
                  "CHECK ( (COD_TIPO_REGOLA = 'CONTROLLO'     AND COD_ERRORE  IS NOT NULL)",
                  attese=1, etichetta='vincolo condizionale (primo ramo)', fatti=fatti)
    D.sostituisci(doc, "     OR (COD_TIPO_REGOLA = 'CONTROLLO'     AND COD_ERRORE  IS NOT NULL)\n",
                  '', attese=None, etichetta='ramo duplicato', fatti=fatti)
    riga = next((p for p in doc.paragraphs
                 if p.text.strip().startswith("OR (COD_TIPO_REGOLA = 'CONTROLLO'")), None)
    if riga is not None:
        riga._p.getparent().remove(riga._p)
        fatti.append('DDL: rimosso il ramo ora duplicato del vincolo')

    # ------------------------------------------- le tabelle descrittive mancanti
    ancora = D.h(doc, 2, 'Le regole già presenti in SIPO')
    dopo = ancora._p
    D.para(doc, dopo, 'ANSC_CFG_REGOLA — le regole di controllo e di generazione.')
    D.tabella(doc, dopo, REGOLA, modello=modello)
    D.para(doc, dopo, 'ANSC_CFG_REGOLA_CONDIZIONE — le condizioni di una regola.')
    D.tabella(doc, dopo, CONDIZIONE, modello=modello)
    D.para(doc, dopo,
           'Le due tabelle conservano la forma dichiarativa scelta nella v3.0: una condizione '
           'è una terna campo, operatore, valore, e le condizioni di una regola si compongono '
           'secondo i contrassegni di obbligatorietà e aggregazione. ⚠️ È la forma che la v3.14 '
           'ha abbandonato per la sola determinazione dell’UC, dove non bastava; per i '
           'controlli e per le generazioni resta adeguata e più verificabile di una query.')
    fatti.append('cap. 10: descritte ANSC_CFG_REGOLA e ANSC_CFG_REGOLA_CONDIZIONE')

    ancora = next(p for p in doc.paragraphs
                  if p.text.strip().startswith('ANSC_CFG_UC —'))
    D.para(doc, ancora._p, 'ANSC_ANA_UC — il catalogo degli UC replicato da ANSC.')
    D.tabella(doc, ancora._p, ANA_UC, modello=modello)
    fatti.append('cap. 8: descritta ANSC_ANA_UC')

    # ------------------------------------------- la colonna che mancava
    t = D.tabella_colonne(doc, 'ID_STATO_ATTO')
    if 'CHIAVE_ANTI_DUPLICATO' not in [r.cells[0].text.strip() for r in t.rows]:
        righe = [[c.text for c in r.cells] for r in t.rows[1:]]
        fuori = []
        for v in righe:
            fuori.append(v)
            if v[0].startswith('ID_ATTO_SIPO'):
                fuori.append(['CHIAVE_ANTI_DUPLICATO', 'VARCHAR2(64)',
                              'Chiave locale che impedisce di accodare due volte lo stesso '
                              'atto per la stessa operazione. ⚠️ È locale: ANSC non espone '
                              'alcuna chiave di idempotenza.'])
        riscrivi_tabella(t, fuori)
        fatti.append('cap. 8: aggiunta CHIAVE_ANTI_DUPLICATO ad ANSC_STATO_ATTO')
    return fatti


if __name__ == '__main__':
    doc = docx.Document(DOC)
    for f in applica(doc):
        print('  ·', f)
    doc.save(DOC)
    print('salvato')
