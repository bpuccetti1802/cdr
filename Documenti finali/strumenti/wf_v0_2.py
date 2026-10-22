# -*- coding: utf-8 -*-
"""WORKFLOW_Interfacce-SIPO-ANSC v0.1 → v0.2 (21/10/2026).

Parte dalla v0.1 RIVISTA A MANO dall'utente (salvata il 21/10 alle 21:36): si modifica il file
reale, non si rigenera da genera_workflow_interfacce.py, che perderebbe le modifiche.

Interviene solo sulle quattro pagine nuove (Validazione, Firme dei dichiaranti, Conformità,
Firma dell'USC), su indicazione dell'utente:
  · un capitolo nuovo con le parti e il comportamento comuni alle quattro pagine;
  · per ciascuna pagina un'ipotesi di schermata nello stile di SIPO (schermate_sipo_wf.py) e
    il diagramma a stati della pagina con la tabella dei passaggi (diagrammi_pagine_wf.py);
  · la regola decisa dall'utente per il ritorno ai dati dopo le firme dei dichiaranti: si
    segue Side (che lo vieta) salvo che tutti i firmatari possano firmare di nuovo;
  · WF-10: la nuova validazione riconosce l'atto dal numero comunale associato all'atto;
  · correzioni: la citazione dell'attestazione nel capitolo Conformità (troncata
    nell'incolla), ANSC_EVENTO → ANSC_STATO_ATTO nel paragrafo sullo stato 0;
  · i punti aperti nuovi WF-15…WF-21 e i riferimenti [R6]…[R8].

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/wf_v0_2.py
"""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D            # noqa: E402
import diagrammi_pagine_wf as SD   # noqa: E402
import schermate_sipo_wf as SCH    # noqa: E402

QUI = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(QUI)
SORGENTE = os.path.join(DOCS, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.1.docx')
DEST = os.path.join(DOCS, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.2.docx')
IMG = os.path.join(QUI, 'img')
LARGO = 6.5
COLONNE = (1.35, 1.75, 1.25, 2.15)
TESTA = ("Da", "Evento [condizione]", "A", "Effetti")


# ------------------------------------------------------------------------- utilità
def fine_capitolo(doc, titolo):
    """L'elemento da cui comincia il capitolo successivo (si inserisce prima di esso)."""
    dentro = False
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            if dentro:
                return p._p
            dentro = p.text.strip().startswith(titolo)
    raise SystemExit('capitolo non trovato o ultimo: ' + titolo)


def blocchi(doc, prima, modello, elementi):
    for tipo, *v in elementi:
        if tipo == 'h2':
            D.para(doc, prima, v[0], stile='Heading 2')
        elif tipo == 'p':
            D.para(doc, prima, v[0])
        elif tipo == 'v':
            D.para(doc, prima, v[0], stile='List Paragraph')
        elif tipo == 'img':
            D.immagine(doc, prima, os.path.join(IMG, v[0]), v[1], v[2])
        elif tipo == 'tab':
            t = D.tabella(doc, prima, [TESTA] + v[0], modello, list(COLONNE))
            for riga in t.rows:            # come le tabelle della v0.1: a sinistra
                for cella in riga.cells:
                    for par in cella.paragraphs:
                        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            D.para(doc, prima, '')


def schermata(nome, didascalia, larghezza=LARGO):
    return ('img', nome, larghezza, "Ipotesi di schermata: " + didascalia)


def diagramma(nome, didascalia):
    return ('img', nome, LARGO, didascalia)


# --------------------------------------------------------------------------- contenuti
COMUNE = [
    ('p', "Le quattro pagine nuove — Validazione, Firme dei dichiaranti, Conformità e Firma "
          "dell’USC — hanno una parte comune, che il capitolo descrive una volta sola; i "
          "capitoli delle singole pagine la danno per presupposta. Le ipotesi di schermata "
          "seguono lo stile delle pagine esistenti di SIPO [R8]: testata rossa, menu "
          "orizzontale, pannelli con barra grigia, etichette sopra i campi, pulsanti "
          "rettangolari in maiuscolo. Il contenuto viene dai contratti di ANSC, dalle "
          "schermate di Side del capitolo 13.1 dell’analisi [R1] e dal documento di Milano "
          "sull’attestato di conformità [R6]. I valori mostrati sono di esempio."),
    ('h2', "Le parti comuni"),
    ('v', "La testata dell’atto. Apre ogni pagina: caso d’uso ANSC con la descrizione "
          "(ANSC_03), numero comunale, identificativo nazionale, stato ANSC, stato SIPO, "
          "ufficiale dello stato civile e tempo residuo dell’OTP di sessione. È ciò che "
          "dice all’operatore quale atto ha davanti e a che punto è."),
    ('v', "Le linguette delle quattro pagine. Rossa quella attiva, grigia con la spunta "
          "quella conclusa, chiara quella non ancora raggiungibile. Riprendono le linguette "
          "delle maschere di SIPO e mostrano l’avanzamento che in Side e nella web app di "
          "ANSC mostra la barra dei passi."),
    ('v', "I pulsanti con lo stesso comportamento ovunque. INDIETRO torna alla pagina "
          "precedente senza toccare i dati. SOSPENDI chiude la redazione e la lascia "
          "riprendibile. ELIMINA EVENTO annulla l’atto in ANSC con R011, solo se l’atto ha "
          "l’identificativo nazionale e prima della firma dell’USC; chiede conferma. Il "
          "pulsante che porta avanti è a destra e resta disabilitato finché la pagina non ha "
          "raggiunto il proprio esito."),
    ('h2', "Il comportamento comune"),
    ('p', "Il diagramma riunisce i passaggi che possono avvenire in qualunque delle quattro "
          "pagine: la sessione con ANSC da rinnovare, l’esito da accertare quando ANSC non "
          "risponde, la sospensione e la ripresa, l’eliminazione dell’evento. Nelle figure "
          "delle singole pagine non sono ripetuti."),
    diagramma('wf_sd_comune.png', "Il comportamento comune delle pagine nuove."),
    ('tab', [
        ("Pagina attiva", "chiamata rifiutata per l’OTP di sessione (406001 non valido, "
         "406002 scaduto)", "Sessione da rinnovare",
         "Si chiede un nuovo OTP, generato sulla web app di ANSC. Codici d’errore di ANSC."),
        ("Sessione da rinnovare", "nuovo OTP consegnato al concentratore",
         "Pagina attiva", "Si ripete l’azione interrotta."),
        ("Pagina attiva", "nessuna risposta o errore 5xx su R009, R006, R012, R007, R014",
         "Esito da accertare", "Nessun nuovo invio alla cieca: si legge lo stato con R005."),
        ("Esito da accertare", "R005: l’operazione non è avvenuta", "Pagina attiva",
         "L’azione si può ripetere."),
        ("Esito da accertare", "R005: l’operazione è avvenuta",
         "stato successivo della pagina", "Si registrano i valori letti."),
        ("Pagina attiva", "SOSPENDI, fine della sessione di SIPO, OTP non rinnovato",
         "Sospesa", "FLG_SOSPESO = S. Modello evento, pagina e revisione sono già salvati "
         "dall’ultimo cambio di pagina o dall’ultima chiamata."),
        ("Sospesa", "riapertura dalla ricerca degli atti", "Ripresa",
         "OTP di sessione e credenziali di firma remota si chiedono di nuovo: non si "
         "conservano."),
        ("Ripresa", "R005: lo stato reale è quello registrato", "Pagina attiva",
         "Si riapre la pagina registrata in COD_PAGINA_RIPRESA."),
        ("Ripresa", "R005: atto annullato o cambiato sulla web app", "stato letto",
         "Lo stato letto prevale su quello registrato e lo sostituisce."),
        ("Pagina attiva", "ELIMINA EVENTO confermato [atto con idAnsc, prima della firma "
         "dell’USC]", "Annullato", "R011; ANSC 3 ANNULLATO. Effetti sulla pratica di SIPO e "
         "sul numero comunale: WF-21."),
    ]),
]

VALIDAZIONE = [
    ('h2', "Ipotesi di schermata"),
    ('p', "La pagina mostra, sotto la testata dell’atto, i dati della validazione (versione "
          "del modello ANSC, data di formazione, revisione del modello evento, ultima "
          "validazione riuscita), l’esito restituito da ANSC e, solo quando ANSC lo concede, "
          "il pannello della forzatura. Ogni errore porta la pagina in cui si corregge e il "
          "collegamento per raggiungerla. VALIDA CON FORZATURA si abilita solo con la "
          "conferma dell’ufficiale; AVANTI solo dopo una validazione riuscita. La figura "
          "mostra il caso di una validazione respinta con anomalie forzabili."),
    schermata('wf_sch_validazione.png', "Validazione."),
    ('h2', "Il diagramma a stati della pagina"),
    diagramma('wf_sd_validazione.png', "Validazione: gli stati della pagina."),
    ('tab', [
        ("ingresso", "[in ANSC nessun atto]", "Pronta da validare",
         "Modello evento costruito, numero comunale presente."),
        ("ingresso", "[stato 0: atto confermato, dati cambiati]", "Da rivalidare",
         "La testata mostra l’idAnsc; un avviso spiega che i dati vanno validati di nuovo."),
        ("Pronta da validare", "VALIDA", "Validazione in corso",
         "Si salva il modello evento (TXT_EVENTO, NUM_REVISIONE) e si chiama R009 "
         "/validazione/evento."),
        ("Da rivalidare", "VALIDA", "Validazione in corso",
         "R009 riferita all’atto esistente, riconosciuto dal numero comunale associato "
         "all’atto (WF-10)."),
        ("Validazione in corso", "200", "Validata",
         "ID_ANSC, ID_EVENTO_ANSC, STATO CONFERMATO, TXT_HASH_VALIDATO, DATA_VALIDAZIONE; "
         "identificativi dei soggetti; numero comunale a 2 Utilizzato."),
        ("Validazione in corso", "400, 404, 416 senza codice di forzatura", "Respinta",
         "Errori nella tabella dell’esito, ciascuno con la pagina in cui si corregge; in "
         "ANSC nulla."),
        ("Validazione in corso", "errori con forcingCode", "Forzabile",
         "forzatura_validazione: codice di forzatura e anomalie."),
        ("Respinta", "VAI ALLA PAGINA o INDIETRO", "pagine dei dati",
         "Si corregge e si torna con VALIDA."),
        ("Forzabile", "conferma dell’ufficiale e VALIDA CON FORZATURA",
         "Validazione in corso", "R009 con forcingCode (WF-15)."),
        ("Forzabile", "INDIETRO", "pagine dei dati", "Si corregge invece di forzare."),
        ("Validata", "INDIETRO e modifica dei dati", "pagine dei dati",
         "Stato 0 BOZZA_APPLICATIVA: l’impronta del modello non è più TXT_HASH_VALIDATO."),
        ("Validata", "AVANTI [ci sono firmatari]", "Firme dei dichiaranti",
         "COD_PAGINA_RIPRESA FIRME_DICHIARANTI."),
        ("Validata", "AVANTI [nessun firmatario]", "Conformità", "COD_PAGINA_RIPRESA CONFORMITA."),
    ]),
]

FIRME = [
    ('h2', "Ipotesi di schermata"),
    ('p', "La pagina elenca i firmatari scelti nel riepilogo, ciascuno con la modalità di firma "
          "(autografa con R006 o digitale con R012), l’email per il link, lo stato della "
          "firma e le azioni sul link. Le firme autografe si raccolgono in due passi: si "
          "stampa l’anteprima dell’atto, la si fa firmare ai presenti e si carica il "
          "documento, di cui si segue lo stato (ANSC_08). VERIFICA STATO EVENTO legge lo "
          "stato con R005; AVANTI si abilita quando ANSC porta l’atto a 4 FIRMATO DAL "
          "DICHIARANTE. La figura mostra una firma mista, autografa e digitale (WF-17)."),
    schermata('wf_sch_firme.png', "Firme dei dichiaranti."),
    ('p', "Dopo che i dichiaranti hanno firmato, TORNA ALLA COMPILAZIONE non riapre "
          "direttamente i dati: apre una finestra che chiede di confermare, firmatario per "
          "firmatario, che ciascuno è in condizione di firmare di nuovo. Finché manca una "
          "conferma, RIAPRI I DATI resta disabilitato e l’alternativa è ELIMINA EVENTO."),
    schermata('wf_sch_riapertura.png', "la conferma per riaprire i dati dopo le firme.", 5.2),
    ('h2', "Il diagramma a stati della pagina"),
    ('p', "La regola del ritorno ai dati. In Side, dopo le firme dei dichiaranti, TORNA ALLA "
          "COMPILAZIONE è disabilitato (schermate del capitolo 13.1) e il documento di Milano "
          "lo dice espressamente: firmato dai dichiaranti, l’evento non torna alla bozza [R6]. "
          "La pagina segue Side con una condizione: il ritorno è consentito se tutti i "
          "firmatari sono in condizione di firmare di nuovo, altrimenti è inibito. Riaprendo i "
          "dati le firme raccolte decadono, lo stato passa a 0 e, dopo la nuova validazione, "
          "si firma di nuovo. Se ANSC accetti la nuova validazione di un atto già firmato dal "
          "dichiarante va provato in preproduzione (WF-16)."),
    diagramma('wf_sd_firme.png', "Firme dei dichiaranti: gli stati della pagina."),
    ('tab', [
        ("ingresso", "—", "Firme da raccogliere",
         "Firmatari e modalità dalla tabella firmatari, scelti nel riepilogo."),
        ("Firme da raccogliere", "CARICA (R006) o INVIA LINK (R012 invio-email)",
         "Firme in corso", "firmatari.cd_stato_firma, id_presa_visione."),
        ("Firme in corso", "documento scartato, link annullato o scaduto (412001)",
         "Firme in corso", "Si ricarica il documento o si invia di nuovo il link."),
        ("Firme in corso", "VERIFICA STATO EVENTO: ANSC porta l’atto a 4",
         "Firmato dai dichiaranti", "STATO FIRMATO_DICHIARANTE; AVANTI abilitato."),
        ("Firme da raccogliere", "INDIETRO", "Validazione", "I dati restano validi."),
        ("Firme da raccogliere", "TORNA ALLA COMPILAZIONE [nessuna firma acquisita]",
         "pagine dei dati", "Stato 0."),
        ("Firme in corso, Firmato dai dichiaranti", "TORNA ALLA COMPILAZIONE [tutti i "
         "firmatari possono firmare di nuovo, confermato firmatario per firmatario]",
         "Dati riaperti", "Stato 0; firme azzerate; link ancora aperti annullati (R012 "
         "annulla-link); la conferma si registra nell’audit. Dopo la nuova validazione si "
         "firma di nuovo (WF-16)."),
        ("Firmato dai dichiaranti", "TORNA ALLA COMPILAZIONE [non tutti possono firmare di "
         "nuovo]", "—", "RIAPRI I DATI disabilitato; resta ELIMINA EVENTO (R011) e un atto "
         "nuovo."),
        ("Firmato dai dichiaranti", "AVANTI", "Conformità", "COD_PAGINA_RIPRESA CONFORMITA."),
        ("firmatario, autografa", "STAMPA ANTEPRIMA, CARICA, VERIFICA",
         "Da firmare → Documento in verifica → Inserito o Scartato",
         "R010 per l’anteprima, R006 per il caricamento; stato del documento ANSC_08. Uno "
         "scartato si ricarica."),
        ("firmatario, digitale", "INVIA LINK, AGGIORNA, ANNULLA LINK, INVIA DI NUOVO",
         "Link da inviare → Da visionare (0) → Visionato (1) o Annullato (2)",
         "R012 invio-email, pooling-statoatto, annulla-link; un link scaduto (412001) si "
         "invia di nuovo."),
    ]),
]

CONFORMITA = [
    ('h2', "Ipotesi di schermata"),
    ('p', "La pagina segue il processo di Milano [R6]: l’elenco degli allegati con la casella "
          "di inclusione, GENERA ATTESTAZIONE DI CONFORMITÀ, l’anteprima del PDF generato da "
          "ANSC con la formula dell’attestazione, le credenziali di firma remota. Un "
          "adattamento rispetto a Side: Side preseleziona tutti gli allegati, l’ipotesi "
          "preseleziona solo quelli che vanno attestati e lascia deselezionati i documenti già "
          "firmati digitalmente (formato P7M), che la guida ANSC dichiara non bisognosi di "
          "attestazione (WF-18). FIRMA si abilita solo con l’anteprima visibile; CONTINUA "
          "solo quando nessun allegato è selezionato."),
    schermata('wf_sch_conformita.png', "Conformità."),
    ('h2', "Il diagramma a stati della pagina"),
    diagramma('wf_sd_conformita.png', "Conformità: gli stati della pagina."),
    ('tab', [
        ("ingresso", "[l’atto non ha allegati da attestare]", "Firma dell’USC",
         "La pagina si salta (caso di prova 6.5 del documento di Milano)."),
        ("ingresso", "[ci sono allegati]", "Selezione degli allegati",
         "La prima volta preselezionati quelli da attestare; le volte successive le scelte "
         "salvate."),
        ("Selezione degli allegati", "deseleziona tutto", "Nessun allegato selezionato",
         "GENERA disabilitato, CONTINUA abilitato."),
        ("Nessun allegato selezionato", "seleziona un allegato", "Selezione degli allegati", "—"),
        ("Selezione degli allegati", "GENERA ATTESTAZIONE DI CONFORMITÀ", "Anteprima generata",
         "R010 /anteprima/attestazione con gli allegati e flagAttConformita; si salva "
         "ALLEGATO.fg_incluso_att_conformita; utenza di firma precompilata."),
        ("Anteprima generata", "cambia la selezione", "Selezione degli allegati",
         "L’anteprima scompare e le caselle della firma si disabilitano."),
        ("Anteprima generata", "FIRMA", "Firma in corso", "R007 /attestazione/firma."),
        ("Firma in corso", "esito positivo", "Firmato per conformità",
         "STATO FIRMATO_CONFORMITA; audit FIRMA_CONFORMITA."),
        ("Firma in corso", "406001, 406004", "Anteprima generata",
         "Messaggio; PIN e OTP del servizio di firma da reinserire."),
        ("Nessun allegato selezionato", "CONTINUA", "Firma dell’USC", "Si salvano le scelte."),
        ("Firmato per conformità", "AVANTI", "Firma dell’USC", "COD_PAGINA_RIPRESA FIRMA_USC."),
        ("Selezione degli allegati", "TORNA ALLA COMPILAZIONE [nessuna firma dei dichiaranti, "
         "o tutti possono firmare di nuovo]", "pagine dei dati",
         "La stessa regola della pagina delle firme dei dichiaranti."),
    ]),
]

FIRMA_USC = [
    ('h2', "Ipotesi di schermata"),
    ('p', "La pagina riprende quella di Side: il contrassegno della comunicazione anagrafica "
          "con la data di validità dell’evento, l’anteprima dell’atto, le credenziali di firma "
          "remota. Al posto di FIRMA c’è CHIUDI, come nella figura del workflow: firma l’atto "
          "in ANSC, invia la data di validità se è richiesta la comunicazione anagrafica e "
          "salva l’atto in SIPO. Un avviso lo dichiara prima del clic, perché dopo la firma "
          "l’atto non si modifica più."),
    schermata('wf_sch_firma_usc.png', "Firma dell’USC."),
    ('h2', "Il diagramma a stati della pagina"),
    diagramma('wf_sd_firma_usc.png', "Firma dell’USC: gli stati della pagina."),
    ('tab', [
        ("ingresso", "—", "Pronta alla firma",
         "ANSC 11, oppure 4 o 2 se non c’era nulla da attestare."),
        ("Pronta alla firma", "ANTEPRIMA DELL’ATTO", "Pronta alla firma",
         "R010, anteprima dell’evento."),
        ("Pronta alla firma", "CHIUDI", "Firma in corso",
         "R007 /firma_usc con le credenziali di firma remota, inviate in base64 "
         "(idModalitaInput 2: «in chiaro» non è accettato in produzione)."),
        ("Firma in corso", "406001, 406002, 406004", "Pronta alla firma",
         "Messaggio; PIN e OTP del servizio di firma da reinserire."),
        ("Firma in corso", "esito positivo", "Firmato dall’USC",
         "ANSC 5: l’atto è formato. Punto di non ritorno: INDIETRO ed ELIMINA EVENTO "
         "disabilitati."),
        ("Firmato dall’USC", "[comunicazione anagrafica]", "Data di validità",
         "R014 /evento/data-validita con la data indicata."),
        ("Firmato dall’USC", "[senza comunicazione anagrafica]", "Salvataggio in SIPO", "—"),
        ("Data di validità", "esito positivo o negativo", "Salvataggio in SIPO",
         "Un esito negativo non ferma la chiusura: si segnala (WF-19)."),
        ("Salvataggio in SIPO", "SIPO 2", "Chiuso",
         "/finalizzaAtto → InitFlow: ID_STATO_PRATICA 2, DATA_DEFINIZIONE_ATTO; "
         "ANSC_STATO_ATTO STATO FIRMATO_USC, COD_FASE FIRMATO."),
        ("Salvataggio in SIPO", "SIPO 3", "Chiuso, sospeso in SIPO",
         "Aggiornamento anagrafico non riuscito o in attesa."),
        ("Salvataggio in SIPO", "errore", "Formato in ANSC, non in SIPO",
         "Eccezione: si ripete il salvataggio, mai la firma (WF-20)."),
        ("Formato in ANSC, non in SIPO", "RIPROVA", "Salvataggio in SIPO", "—"),
        ("Chiuso, sospeso in SIPO", "/annullaSospensioneAtto", "Chiuso", "SIPO 2."),
    ]),
]

RIFERIMENTI = [
    ("[R6]", "Sorgenti Documentali/CdMI SIDe MI_ANSC - Processo di gestione e firma "
     "dell'attestato di conformità v1.1.docx",
     "Il processo di Milano per la conformità: pulsanti, abilitazioni, casi di prova; il "
     "divieto di tornare alla bozza dopo le firme dei dichiaranti."),
    ("[R7]", "ansc/docs/Guida_operativa_ANSC.pdf, §6.8",
     "La formula e il contenuto dell’attestazione di conformità generata da ANSC."),
    ("[R8]", "Sorgenti Documentali/Immagini sipo/ e il CSS di statocivile-web",
     "Lo stile delle ipotesi di schermata: testata, menu, pannelli, campi, pulsanti, colori."),
]

WF_NUOVI = [
    ("WF-15", "Anteprima · Validazione",
     "Chi può confermare la forzatura e con quale abilitazione. Un atto forzato entra in ANSC "
     "come anomalo e per gli atti anomali ANSC dichiara non ancora implementato il ciclo delle "
     "conferme (OP-21, OP-41)."),
    ("WF-16", "Anteprima · Firme dei dichiaranti",
     "Ritorno ai dati dopo le firme: si accerta con la conferma dell’operatore, firmatario per "
     "firmatario, registrata nell’audit (proposta). Va provato in preproduzione se ANSC accetti "
     "una nuova validazione R009 di un atto in stato 4 FIRMATO DAL DICHIARANTE; se non la "
     "accetta, il ritorno passa per R011 e un atto nuovo [DA VERIFICARE]."),
    ("WF-17", "Anteprima · Firme dei dichiaranti",
     "Firma mista, autografa e digitale, nello stesso atto: se ANSC la ammetta e quando porti "
     "l’atto a 4 con più firmatari. I valori di firmatari.cd_stato_firma (dominio 19 di Side) "
     "non sono documentati."),
    ("WF-18", "Anteprima · Conformità",
     "Preselezione degli allegati: Side li preseleziona tutti, l’ipotesi esclude quelli già "
     "firmati digitalmente (P7M), da confermare. Il percorso del servizio è "
     "/anteprima/attestazione nel contratto R010 e /attestazione/anteprima nel Changelog; "
     "l’algoritmo dell’impronta non è dichiarato."),
    ("WF-19", "Anteprima · Firma dell’USC",
     "Esito negativo di R014 dopo la firma: l’atto è formato ma la data di validità non è "
     "registrata. Come e da dove si ripete (back-office?) [DA VERIFICARE]."),
    ("WF-20", "Anteprima · Firma dell’USC",
     "Salvataggio in SIPO non riuscito dopo la firma in ANSC: l’eccezione «formato in ANSC, "
     "non in SIPO» va aggiunta alle categorie di eccezione del back-office."),
    ("WF-21", "Pagine nuove",
     "ELIMINA EVENTO (R011): effetto sulla pratica di SIPO e sul numero comunale già usato "
     "(DATI_ATTO_ESTERNO) [DA VERIFICARE]."),
]


# ------------------------------------------------------------------------------- main
def main():
    for p in (SORGENTE,):
        if not os.path.exists(p):
            raise SystemExit('manca ' + p)
    SCH.disegna_tutte()
    SD.disegna_tutti()
    doc = docx.Document(SORGENTE)
    fatti = []

    # correzioni puntuali
    D.sostituisci(doc, "le copie allzione Allegati, sono conformi",
                  "le copie allegate ed elencate nella sezione Allegati, sono conformi",
                  attese=1, fatti=fatti, etichetta='citazione dell’attestazione')
    D.sostituisci(doc, "(nell'esempio:  «Processo verbale»)",
                  "(nell’esempio: «Attestazione di nascita», «Processo verbale»)",
                  attese=1, fatti=fatti, etichetta='esempio degli allegati')
    D.sostituisci(doc, "salvato in ANSC_EVENTO", "salvato in ANSC_STATO_ATTO", attese=1,
                  fatti=fatti, etichetta='ANSC_EVENTO → ANSC_STATO_ATTO')
    D.sostituisci(doc, "Dalla firma del dichiarante i dati sono vincolati: correggerli "
                  "significa tornare alla validazione e far firmare di nuovo; in questo caso "
                  "lo stato passa a 0 BOZZA",
                  "Torna alla compilazione dopo le firme: consentito solo se tutti i firmatari "
                  "sono in condizione di firmare di nuovo; le firme decadono, lo stato passa a "
                  "0 e dopo la nuova validazione si firma di nuovo. Altrimenti è inibito e "
                  "resta ELIMINA EVENTO (R011)",
                  attese=1, fatti=fatti, etichetta='regola del ritorno ai dati')

    wf = D.trova_tabella(doc, 'Punto del workflow', 'Questione')
    for r in wf.rows:
        if r.cells[0].text.strip() == 'WF-10':
            D.riscrivi_cella(r.cells[2], "Con quale campo, id o idAnsc, una nuova validazione "
                             "riconosca l’atto già confermato, dopo un ritorno sui dati "
                             "(OP-75 a, OP-46). RISOLTO: l’atto si riconosce dal numero "
                             "comunale (evento.numeroatto), che è associato all’atto e sta in "
                             "ANSC_STATO_ATTO.NUM_COMUNALE, insieme all’idAnsc; l’ID_ATTO di "
                             "SIPO non è un campo di ANSC. Il comportamento di ANSC va "
                             "provato in preproduzione.")
            fatti.append('WF-10 riscritto')
    for riga in WF_NUOVI:
        D.clona_riga(wf, riga)
    fatti.append('punti aperti WF-15…WF-21')

    rif = D.trova_tabella(doc, 'Documento', 'Che cosa se ne prende')
    for riga in RIFERIMENTI:
        D.clona_riga(rif, riga)
    fatti.append('riferimenti [R6]…[R8]')

    modello = D.trova_tabella(doc, 'Da', 'Che cosa lo provoca')

    # il capitolo comune, prima della prima pagina nuova
    prima = D.h(doc, 1, "Anteprima · Validazione")._p
    D.para(doc, prima, "Le pagine nuove: parti e comportamento comuni", stile='Heading 1')
    blocchi(doc, prima, modello, COMUNE)
    fatti.append('capitolo «Le pagine nuove: parti e comportamento comuni»')

    for titolo, contenuto in (("Anteprima · Validazione", VALIDAZIONE),
                              ("Anteprima · Firme dei dichiaranti", FIRME),
                              ("Anteprima · Conformità", CONFORMITA),
                              ("Anteprima · Firma dell’USC", FIRMA_USC)):
        blocchi(doc, fine_capitolo(doc, titolo), modello, contenuto)
        fatti.append('schermata e diagramma: ' + titolo)

    fatti += D.frontespizio(doc, '0.2', 'WORKFLOW_Interfacce-SIPO-ANSC_v0.2', '21/10/2026')
    D.storia(doc, '21/10/2026', '0.2',
             'Pagine nuove; Firme dei dichiaranti; Conformità; Punti aperti; Riferimenti',
             "Capitolo sulle parti e il comportamento comuni alle pagine nuove; per Validazione, "
             "Firme dei dichiaranti, Conformità e Firma dell’USC l’ipotesi di schermata in stile "
             "SIPO e il diagramma a stati della pagina con la tabella dei passaggi. Ritorno ai "
             "dati dopo le firme solo se tutti i firmatari possono firmare di nuovo; WF-10 "
             "risolto con il numero comunale; corretta la citazione dell’attestazione; "
             "WF-15…WF-21.")
    doc.save(DEST)
    for f in fatti:
        print(' ·', f)
    print('scritto:', DEST, D.riepilogo(DEST))


if __name__ == '__main__':
    main()
