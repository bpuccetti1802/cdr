# -*- coding: utf-8 -*-
"""Primitive per modificare il documento di analisi con python-docx.

Regola del workspace: il .docx si MODIFICA, non si rigenera — contiene edit manuali
dell'utente (fra cui i titoli dei capitoli, che l'utente rinomina). Queste funzioni
inseriscono paragrafi, tabelle e immagini in un punto preciso del corpo, preservando tutto
il resto.

Trappole già pagate e qui evitate:
  · «doc.paragraphs» e «doc.tables» costruiscono oggetti NUOVI a ogni accesso: per
    confrontare o cercare un indice si usa l'elemento XML sottostante (p._p, t._tbl), mai
    l'identità dell'oggetto.
  · i marcatori markdown «**...**» restano LETTERALI: il grassetto si ottiene con più run,
    ed è ciò che fa «pezzi».
  · clonare una riga di tabella eredita la formattazione; costruirla da zero no.

Questo file sta in «Documenti finali/strumenti/» e non nello scratchpad, che è di sessione.
"""
import copy

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

GIUST = WD_ALIGN_PARAGRAPH.JUSTIFY


# ------------------------------------------------------------------ ricerca
def h(doc, livello, testo, esatto=False):
    """Primo Heading del livello dato che comincia con (o è uguale a) il testo."""
    for p in doc.paragraphs:
        if p.style.name != f'Heading {livello}':
            continue
        t = p.text.strip()
        if (t == testo) if esatto else t.startswith(testo):
            return p
    raise SystemExit(f'Heading {livello} non trovato: {testo}')


def indice_di(doc, paragrafo):
    """Indice di un paragrafo, confrontando l'elemento XML e non l'oggetto."""
    return next(i for i, p in enumerate(doc.paragraphs) if p._p is paragrafo._p)


def trova_tabella(doc, *intestazioni):
    for t in doc.tables:
        testa = ' | '.join(c.text.strip().lower() for c in t.rows[0].cells)
        if all(i.lower() in testa for i in intestazioni):
            return t
    raise SystemExit('tabella non trovata: ' + ' / '.join(intestazioni))


def tabella_colonne(doc, colonna_caratteristica):
    """Fra le molte tabelle «Colonna | Tipo | Note», quella che contiene la colonna data."""
    for t in doc.tables:
        testa = ' | '.join(c.text.strip() for c in t.rows[0].cells)
        if not testa.startswith('Colonna | Tipo'):
            continue
        if any(r.cells[0].text.strip() == colonna_caratteristica for r in t.rows):
            return t
    raise SystemExit('tabella colonne non trovata per ' + colonna_caratteristica)


# ------------------------------------------------------------------ scrittura
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def ha_commenti(elemento):
    """Vero se l'elemento contiene l'ancora di un commento di Word."""
    for tag in ('commentRangeStart', 'commentRangeEnd', 'commentReference'):
        if elemento.find(f'.//{W}{tag}') is not None:
            return True
    return False


def testo_di(p, nuovo):
    """Sostituisce il testo di un paragrafo conservando formattazione E COMMENTI.

    ⚠️ La versione ingenua — scrivere nel primo run e cancellare gli altri — cancella anche i
    run che portano l'ancora di un commento di Word, e il commento sparisce dal documento
    senza lasciare traccia. I commenti sono istruzioni di lavoro dell'utente: vanno protetti
    dalle modifiche automatiche, non riscritti insieme al testo.
    """
    runs = p.runs
    if not runs:
        p.add_run(nuovo)
        return p
    primo = None
    for r in runs:
        if primo is None and not ha_commenti(r._r):
            primo = r
            r.text = nuovo
        elif r is not primo and not ha_commenti(r._r):
            r.text = ''
    if primo is None:                 # tutti i run sono ancore: si aggiunge senza toccarli
        p.add_run(nuovo)
    return p


def _scollegato(doc):
    p = doc.add_paragraph()
    p._p.getparent().remove(p._p)
    return p


def segmenta(testa, corpo):
    """«testo **in grassetto** testo» -> lista di (testo, grassetto) per «pezzi»."""
    pezzi = [(testa, True)] if testa else []
    for i, parte in enumerate(corpo.split('**')):
        if parte:
            pezzi.append((parte, i % 2 == 1))
    return pezzi


def para(doc, prima, testo='', stile='Normal', pezzi=None, corsivo=False, mono=False):
    """Crea un paragrafo e lo innesta prima dell'elemento «prima» (un p._p)."""
    p = _scollegato(doc)
    p.style = doc.styles[stile]
    p.alignment = GIUST
    if pezzi:
        for txt, bold in pezzi:
            r = p.add_run(txt)
            r.bold = bold
    elif testo:
        r = p.add_run(testo)
        if corsivo:
            r.italic = True
            r.font.size = Pt(9)
        if mono:
            r.font.name = 'Courier New'
            r.font.size = Pt(8)
    if mono:
        p.paragraph_format.space_after = Pt(0)
    prima.addprevious(p._p)
    return p


def voce(doc, prima, testa, corpo):
    """Voce puntata con lead-in in grassetto; interpreta anche i ** dentro il corpo."""
    return para(doc, prima, stile='List Paragraph', pezzi=segmenta(testa, corpo))


def ddl(doc, prima, righe):
    for r in righe:
        para(doc, prima, r, mono=True)


def titolo_ddl(doc, prima, nome):
    p = para(doc, prima, nome)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(9)
    return p


def tabella(doc, prima, righe, modello, larghezze=None):
    """Tabella nuova, con bordi e larghezza clonati da una tabella esistente."""
    t = doc.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = modello.style
    vecchio = t._tbl.find(qn('w:tblPr'))
    if vecchio is not None:
        t._tbl.remove(vecchio)
    t._tbl.insert(0, copy.deepcopy(modello._tbl.find(qn('w:tblPr'))))
    for i, riga in enumerate(righe):
        for j, val in enumerate(riga):
            p = t.cell(i, j).paragraphs[0]
            p.alignment = GIUST
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
    if larghezze:
        for j, w in enumerate(larghezze):
            for riga in t.rows:
                riga.cells[j].width = Inches(w)
        # ⚠️ Le larghezze di cella (w:tcW) non aggiornano la griglia (w:tblGrid/w:gridCol),
        # che `add_table` crea a colonne uguali: finché il documento non passa da Word —
        # che la ricalcola — la tabella si impagina sulla griglia e non sulle celle.
        # Va scritta anche quella, altrimenti la tabella nuova è l'unica a colonne uguali.
        griglia = t._tbl.find(qn('w:tblGrid'))
        if griglia is not None:
            for col, w in zip(griglia.findall(qn('w:gridCol')), larghezze):
                col.set(qn('w:w'), str(int(round(Inches(w).twips))))
    t._tbl.getparent().remove(t._tbl)
    prima.addprevious(t._tbl)
    return t


def immagine(doc, prima, percorso_png, pollici, didascalia):
    p = _scollegato(doc)
    p.alignment = GIUST
    p.add_run().add_picture(percorso_png, width=Inches(pollici))
    prima.addprevious(p._p)
    para(doc, prima, didascalia, corsivo=True)


def clona_riga(t, valori):
    """Aggiunge una riga copiando la formattazione dell'ultima e sostituendone i testi."""
    tr = copy.deepcopy(t.rows[-1]._tr)
    t._tbl.append(tr)
    riga = t.rows[-1]
    for cella, val in zip(riga.cells, valori):
        for p in cella.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        p = cella.paragraphs[0]
        if not p.runs:
            p.add_run('')
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = str(val)
    return riga


def riga_dopo(doc, ancora, testo):
    """Duplica un paragrafo (per ereditarne il carattere) e ne cambia il contenuto."""
    nuovo = copy.deepcopy(ancora._p)
    ancora._p.addnext(nuovo)
    np = Paragraph(nuovo, ancora._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    np.runs[0].text = testo
    return np


def sostituisci(doc, vecchio, nuovo, attese=None, etichetta='', fatti=None):
    """Sostituzione in paragrafi e celle, con controllo del numero di occorrenze."""
    n = 0
    for p in doc.paragraphs:
        if vecchio in p.text:
            testo_di(p, p.text.replace(vecchio, nuovo))
            n += 1
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    if vecchio in p.text:
                        testo_di(p, p.text.replace(vecchio, nuovo))
                        n += 1
    if attese is not None and n != attese:
        raise SystemExit(f'«{vecchio[:60]}»: attese {attese} occorrenze, trovate {n}')
    if n and fatti is not None:
        fatti.append(f'{etichetta or vecchio[:46]} — {n} occorrenz{"a" if n == 1 else "e"}')
    return n


def sostituisci_immagine(doc, inizio_didascalia, png):
    """Sostituisce il PNG di una figura individuata dalla sua didascalia.

    Serve quando il modello dati cambia: figura e modello sono una coppia.
    """
    ps = doc.paragraphs
    for i, p in enumerate(ps):
        if p.text.strip().startswith(inizio_didascalia):
            blips = ps[i - 1]._p.findall('.//' + qn('a:blip'))
            if len(blips) != 1:
                raise SystemExit(f'attesa 1 immagine prima di «{inizio_didascalia}»')
            rid = blips[0].get(qn('r:embed'))
            doc.part.related_parts[rid]._blob = open(png, 'rb').read()
            return True
    raise SystemExit('didascalia non trovata: ' + inizio_didascalia)


def storia(doc, data, versione, sezioni, sintesi):
    clona_riga(trova_tabella(doc, 'versione', 'sintesi dei cambiamenti'),
               (data, versione, sezioni, sintesi))


def riepilogo(percorso_docx):
    import docx as _d
    d = _d.Document(percorso_docx)
    return (sum(1 for p in d.paragraphs if p.style.name == 'Heading 1'),
            len(d.tables), len(d.inline_shapes))


def sezione(doc, titolo, livello=2):
    """Gli elementi XML di una sezione: dal suo heading fino al successivo di pari livello.

    Serve per sostituire un ragionamento intero quando la decisione che lo reggeva è
    cambiata: riscrivere i paragrafi uno per uno lascerebbe in mezzo le frasi della versione
    precedente, che è il modo più sicuro per pubblicare un capitolo che si contraddice.
    """
    corpo = list(doc.element.body.iterchildren())
    inizio = None
    for i, ch in enumerate(corpo):
        if not ch.tag.endswith('}p'):
            continue
        par = Paragraph(ch, doc)
        st, t = par.style.name, par.text.strip()
        if inizio is None:
            if st == f'Heading {livello}' and t.startswith(titolo):
                inizio = i
        elif st.startswith('Heading') and int(st.split()[-1]) <= livello:
            return corpo[inizio:i]
    return corpo[inizio:] if inizio is not None else []


def sostituisci_sezione(doc, titolo, nuovo_titolo, blocchi, livello=2):
    """Svuota una sezione e la riscrive. `blocchi` = lista di (tipo, contenuto).

    tipo: 'p' paragrafo · 'v' voce (testa, corpo) · 'h3' sottotitolo · 'tab' (righe, modello).
    """
    elementi = sezione(doc, titolo, livello)
    assert elementi, f'sezione non trovata: {titolo}'
    testa = elementi[0]
    testo_di(Paragraph(testa, doc), nuovo_titolo)
    for ch in elementi[1:]:
        ch.getparent().remove(ch)
    dopo = testa.getnext()
    for tipo, contenuto in blocchi:
        if tipo == 'p':
            para(doc, dopo, contenuto)
        elif tipo == 'h3':
            para(doc, dopo, contenuto, stile=f'Heading {livello + 1}')
        elif tipo == 'v':
            voce(doc, dopo, *contenuto)
        elif tipo == 'tab':
            righe, modello = contenuto
            tabella(doc, dopo, righe, modello=modello)
    return len(blocchi)


def riscrivi_cella(cella, testo):
    """Riscrive una cella conservando i commenti che vi sono ancorati."""
    if not cella.paragraphs:
        cella.text = testo
        return
    testo_di(cella.paragraphs[0], testo)
    for p in cella.paragraphs[1:]:
        if not ha_commenti(p._p):
            p._p.getparent().remove(p._p)


def riscrivi_sicura(t, righe):
    """Riscrive una tabella «Colonna | Tipo | Note» SENZA perdere i commenti.

    ⚠️ Le righe con un'ancora di commento non si eliminano mai: se la nuova configurazione
    ne prevede di meno, quelle righe restano e vanno svuotate a mano dall'autore. Cancellare
    una riga commentata significherebbe cancellare l'istruzione che qualcuno ci ha scritto.
    """
    while len(t.rows) - 1 < len(righe):
        clona_riga(t, [c.text for c in t.rows[-1].cells])
    for i, valori in enumerate(righe, start=1):
        for j, v in enumerate(valori):
            riscrivi_cella(t.rows[i].cells[j], v)
    for i in range(len(t.rows) - 1, len(righe), -1):
        if ha_commenti(t.rows[i]._tr):
            continue
        t._tbl.remove(t.rows[i]._tr)


def sposta_commenti(sorgente, destinazione):
    """Sposta le ancore dei commenti da un paragrafo a un altro, conservando i commenti.

    ⚠️ Serve quando si elimina un blocco duplicato su cui qualcuno ha commentato: il commento
    vive in `word/comments.xml`, ma senza l'ancora nel testo Word non lo mostra più e alla
    prima riapertura sparisce. Poiché il paragrafo gemello esiste nel blocco che resta, le
    ancore si spostano lì: il commento resta attaccato alla stessa frase, nella copia giusta.

    Si spostano `commentRangeStart`, `commentRangeEnd` e il run che contiene
    `commentReference`, nell'ordine in cui Word li vuole: l'apertura dopo le proprietà, la
    chiusura e il riferimento in coda.
    """
    ids = []
    for tag in ('commentRangeStart', 'commentRangeEnd'):
        for e in list(sorgente.findall(f'{W}{tag}')):
            ids.append(e.get(f'{W}id'))
            sorgente.remove(e)
            if tag == 'commentRangeStart':
                pPr = destinazione.find(f'{W}pPr')
                destinazione.insert(1 if pPr is not None else 0, e)
            else:
                destinazione.append(e)
    for r in list(sorgente.findall(f'{W}r')):
        if r.find(f'{W}commentReference') is not None:
            sorgente.remove(r)
            destinazione.append(r)
    return ids
