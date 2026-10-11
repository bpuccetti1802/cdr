# -*- coding: utf-8 -*-
"""DISEGNO_Back-Office_ANSC v0.8 → v0.9 (11/10/2026).

La numerazione comunale si riallinea all'analisi v3.38-v3.39 e alle due pagine del sistema
Side del Comune di Milano («Gestione numerazione comunale — Richiesta» e «— Ricerca e
consultazione», documento di processo v1.0):
  * la pagina unica «Numerazione comunale» diventa due pagine, richiesta e ricerca, che
    chiamano il servizio num-ansc-sipo (contratto ansc-numerazione-v1.yaml) invece di scrivere
    su ANSC_STATO_ATTO.NUM_COMUNALE;
  * la ricerca porta in testa le serie dell'anno con i BUCHI, che non hanno una riga e si
    vedono solo lì; la richiesta porta la registrazione a posteriori (scenario 4-5 di Side);
  * due adattamenti a Side dichiarati: stati finali e identificativo nazionale obbligatorio per
    «Utilizzato»; registrazione del numero effettivamente usato invece di una richiesta che
    «coinciderà» con il numero suggerito da ANSC;
  * BO-3 e BO-4 chiusi, BO-30 nuovo; copertura a venti tabelle del modello.
Verifica su SOGGETTI_USECASE: pagina, scheda del dettaglio UC, caricamento da file e copertura
c'erano già; mancavano nella tabella delle unità di rilascio (capitolo sul contesto tecnico),
dove mfe-configurazione elencava ancora le pagine della v0.7.

Le figure da 26 in poi scalano di uno (una pagina in più dopo la Figura 25).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v4.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v2.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/bo_v0_9.py
"""
import os
import re
import shutil
import subprocess
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.8.docx')
DST = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.9.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


USO = ('Tabella', 'Uso', 'Che cosa se ne mostra')
EL = ('Elemento', 'Genere', 'Comportamento')

PAGINE = [
    ('h2', 'Numerazione comunale — richiesta'),
    ('img', ('bo_num_richiesta.png', 25, 'Numerazione comunale — richiesta. Tre numeri '
             'consecutivi per un parto trigemino e, sotto, la registrazione a posteriori.')),
    ('p', 'Il numero comunale dell’atto è a carico del Comune: nessun servizio di ANSC lo '
          'assegna, e ANSC lo richiede per formare l’atto. Le due pagine della numerazione '
          'riprendono quelle del sistema Side del Comune di Milano [R11] — la richiesta e la '
          'ricerca con consultazione — e poggiano sul servizio della numerazione '
          'num-ansc-sipo che l’analisi introduce [R1]: le pagine non scrivono le tabelle, '
          'chiamano le operazioni del servizio [R6].'),
    ('p', 'La pagina si apre con il registro corrente, l’anno corrente e un numero, come in '
          'Side, e i tre campi sono modificabili. L’anno si cambia a cavallo d’anno, quando '
          'serve un numero della serie precedente; il numero di elementi si alza per il parto '
          'plurimo, che richiede numeri consecutivi. RICHIEDI stacca i numeri e li mostra in '
          'tabella nello stato Disponibile.'),
    ('p', '⚠️ Due adattamenti rispetto a Side, dichiarati. Il primo: in Side le azioni di riga '
          'sono «solo ad uso e consumo dell’operatore» e non hanno vincoli; qui gli stati '
          'Scartato ed Evento acquisito sono finali e Utilizzato chiede l’identificativo '
          'nazionale, perché in SIPO tutti i numeri passano dalla serie — anche quelli del '
          'percorso ordinario — e la serie è il luogo in cui un buco deve potersi spiegare. Il '
          'secondo: quando il gestionale non era disponibile e l’atto è stato formato sulla '
          'web app con il numero suggerito da ANSC, Side fa una richiesta ordinaria, nella '
          'fiducia che il numero ottenuto coincida; qui si registra il numero effettivamente '
          'usato, e la serie si adegua.'),
    ('h3', 'Tabelle sottese'),
    ('uso', [
        USO,
        ('PRENOTAZIONE_ATTI', 'Scrittura, tramite il servizio',
         'La richiesta: registro, anno, quantità, operatore e istante.'),
        ('DATI_ATTO_ESTERNO', 'Scrittura, tramite il servizio',
         'I numeri della richiesta, il loro stato, l’identificativo nazionale e l’atto di SIPO.'),
        ('SEQ_NUM_COMUNALE_<registro>_<anno>', 'Consumo, tramite il servizio',
         'La serie da cui i numeri si staccano. La pagina non la nomina: il servizio la sceglie '
         'da registro e anno.'),
        ('VALORE_DOMINIO', 'Lettura', 'I cinque registri della decodifica ANSC_01, per la '
                                      'tendina.'),
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('el', [
        EL,
        ('Registro', 'Campo', 'Tendina sui registri della decodifica ANSC_01. Predefinito il '
                              'registro corrente.'),
        ('Anno', 'Campo', 'Predefinito l’anno corrente; ammessi il corrente e il precedente.'),
        ('Numero di elementi', 'Campo', 'Predefinito 1. Per il parto plurimo, il numero dei '
                                        'nati: i numeri escono consecutivi.'),
        ('RICHIEDI', 'Azione', 'Stacca i numeri (creaPrenotazione) e li aggiunge alla tabella. '
                               'Se la serie dell’anno non esiste lo dice: la crea '
                               'l’amministratore dalla pagina di ricerca.'),
        ('Spunta', 'Azione di riga', 'Segna il numero come Utilizzato e chiede l’identificativo '
                                     'nazionale dell’atto (aggiornaNumero).'),
        ('Croce', 'Azione di riga', 'Segna il numero come Scartato (aggiornaNumero). Lo stato '
                                    'è finale.'),
        ('Registrazione a posteriori', 'Pannello',
         'Registro, anno, numero usato e identificativo nazionale; REGISTRA esegue creaNumero. '
         'Un numero già registrato è respinto, con l’indicazione dell’atto che lo porta.'),
    ]),

    ('h2', 'Numerazione comunale — ricerca e consultazione'),
    ('img', ('bo_num_ricerca.png', 26, 'Numerazione comunale — ricerca e consultazione. Le '
             'serie dell’anno in testa, con i buchi; la ricerca delle richieste sotto.')),
    ('p', 'La pagina riprende la ricerca di Side: registro, intervallo di date e operatore già '
          'impostati — il registro corrente, il giorno corrente, l’operatore corrente — e '
          'obbligatorie le sole due date. Alla ricerca di Side si aggiunge il filtro sull’atto '
          'di SIPO, che serve a ritrovare il numero di un atto in lavorazione.'),
    ('p', 'In testa sta ciò che Side non ha, perché non gli serve: lo stato delle serie '
          'dell’anno, registro per registro, con il prossimo numero, i conteggi per stato e i '
          'buchi. In Side passano dalla richiesta i soli numeri degli atti redatti sulla web '
          'app; in SIPO passano dalla serie tutti, e un buco nella serie di un registro è un '
          'fatto da spiegare. Il buco non ha una riga che lo descriva: questa tabella è l’unico '
          'posto in cui si vede, e il suo dettaglio ne elenca i numeri.'),
    ('p', 'CREA SERIE crea le sequence dei cinque registri per un anno. Le crea il job di '
          'dicembre per l’anno successivo; il comando serve all’installazione e quando il job '
          'non è riuscito, ed è riservato all’amministratore.'),
    ('h3', 'Tabelle sottese'),
    ('uso', [
        USO,
        ('PRENOTAZIONE_ATTI', 'Lettura', 'Le richieste che soddisfano i filtri.'),
        ('DATI_ATTO_ESTERNO', 'Lettura e scrittura, tramite il servizio',
         'I numeri delle richieste e i loro cambi di stato.'),
        ('SEQ_NUM_COMUNALE_<registro>_<anno>', 'Lettura; creazione dall’amministratore',
         'Il prossimo valore di ciascuna serie, da cui si calcolano i buchi. CREA SERIE esegue '
         'la procedura che crea le sequence.'),
        ('ANSC_STATO_ATTO', 'Lettura', 'Dal numero all’atto depositato: NUM_COMUNALE e ID_ANSC, '
                                       'per aprirne il dettaglio.'),
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('el', [
        EL,
        ('Serie dell’anno', 'Pannello', 'Una riga per registro con prossimo numero, conteggi '
                                        'per stato e buchi (elencaSerie); «vedi» elenca i '
                                        'numeri mancanti (leggiSerie).'),
        ('Registro · Richieste da · Richieste a · Operatore · Stato · Atto SIPO', 'Filtri',
         'Predefiniti il registro, il giorno e l’operatore correnti; obbligatorie le due date '
         '(elencaPrenotazioni).'),
        ('Spunta · Croce', 'Azioni di riga', 'Sui numeri Disponibili, come nella pagina di '
                                             'richiesta.'),
        ('Matita', 'Azione di riga', 'Su un numero Utilizzato: lega l’atto di SIPO e porta il '
                                     'numero a Evento acquisito, quando l’atto formato sulla '
                                     'web app viene agganciato in SIPO (aggiornaNumero).'),
        ('Occhio', 'Azione di riga', 'Il numero, la richiesta che lo ha prodotto e chi lo ha '
                                     'ricongiunto all’atto (leggiNumero); dal numero di un atto '
                                     'depositato si apre il dettaglio dell’atto.'),
        ('CREA SERIE', 'Azione di pagina', 'Crea le serie di un anno (creaSerie). Riservata '
                                           'all’amministratore.'),
    ]),
]

COPERTURA = [
    ('PRENOTAZIONE_ATTI', 'Numerazione — richiesta · Numerazione — ricerca',
     'Completa, tramite il servizio num-ansc-sipo. Nuova in questa versione.'),
    ('DATI_ATTO_ESTERNO', 'Numerazione — richiesta · Numerazione — ricerca',
     'Completa, compresi i cambi di stato e la registrazione a posteriori, tramite il servizio. '
     'Nuova in questa versione.'),
]

BO_NUOVI = [
    ('BO-30', 'Se la registrazione a posteriori di un numero richieda un profilo distinto',
     'Il contratto la consente all’ufficiale, come lo stacco. Ma un numero registrato a '
     'posteriori può far avanzare la serie e lasciarvi buchi: va stabilito se basti '
     'l’abilitazione dell’ufficiale o serva un profilo distinto, come per il caricamento da '
     'file (BO-28).'),
]

RIF = ('[R11]', 'Documento di processo «CdMI SIDe MI_ANSC — Processo di gestione della '
                'numerazione comunale» v1.0: gli scenari, le pagine di richiesta e di ricerca e '
                'consultazione, le tabelle PRENOTAZIONE_ATTI e DATI_ATTO_ESTERNO',
       'Evidenza fornita')

STORIA = ('11/10/2026', '0.9',
          'Le pagine (Numerazione comunale: richiesta, ricerca e consultazione) · Mappa · '
          'Contesto tecnico · Copertura · Riferimenti · Punti aperti',
          'Riallineata la numerazione comunale alla v3.39 dell’analisi, che la affida al '
          'servizio num-ansc-sipo, e alle due pagine del sistema Side del Comune di Milano. La '
          'pagina unica diventa due: la richiesta, con la registrazione a posteriori di un '
          'numero usato sulla web app, e la ricerca, che porta in testa le serie dell’anno con '
          'i buchi. Due adattamenti a Side dichiarati: stati finali e identificativo nazionale '
          'per «Utilizzato»; registrazione del numero effettivamente usato. PRENOTAZIONE_ATTI e '
          'DATI_ATTO_ESTERNO entrano nella copertura, che sale a venti tabelle. BO-3 e BO-4 '
          'chiusi, BO-30 nuovo. Completata la tabella delle unità di rilascio, che non elencava '
          'le pagine degli allegati e dei soggetti per caso d’uso e del caricamento da file.')


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    MOD_USO = D.trova_tabella(d, *USO)
    MOD_EL = D.trova_tabella(d, *EL)

    # ─────────────────────────── 0. le figure da 26 in poi scalano di uno
    for p in d.paragraphs:
        t = p.text.strip()
        m = re.match(r'^Figura (\d+) —', t)
        if m and int(m.group(1)) >= 26:
            D.testo_di(p, re.sub(r'^Figura \d+ —', 'Figura %d —' % (int(m.group(1)) + 1), t))
        elif re.search(r'Figura 2[67]\b', t):
            raise SystemExit('rimando a una figura da rinumerare: ' + t[:80])
    fatti.append('figure 26-27 rinumerate 27-28')

    # ─────────────────────────── 1. la pagina della numerazione → due pagine
    elementi = D.sezione(d, 'Numerazione comunale', 2)
    assert elementi and elementi[0].xpath('string(.)').strip() == 'Numerazione comunale'
    for el in elementi:
        if D.ha_commenti(el):
            raise SystemExit('commento nella sezione da sostituire')
    ancora = elementi[-1].getnext()
    for el in elementi:
        el.getparent().remove(el)
    for tipo, c in PAGINE:
        if tipo == 'p':
            D.para(d, ancora, c)
        elif tipo in ('h2', 'h3'):
            D.para(d, ancora, c, stile='Heading ' + tipo[1])
        elif tipo == 'uso':
            D.tabella(d, ancora, c, MOD_USO, [1.6, 1.2, 3.5])
            D.para(d, ancora, '')
        elif tipo == 'el':
            D.tabella(d, ancora, c, MOD_EL, [1.6, 1.2, 3.5])
            D.para(d, ancora, '')
        elif tipo == 'img':
            png, n, did = c
            D.immagine(d, ancora, os.path.join(IMG, png), 6.2, 'Figura %d — %s' % (n, did))
    fatti.append('numerazione: due pagine (richiesta; ricerca e consultazione)')

    # ─────────────────────────── 2. mappa, contesto tecnico, conteggi
    t = D.trova_tabella(d, 'Area', 'Pagine', 'Chi la usa e quando')
    r = [r for r in t.rows if r.cells[0].text.strip() == 'Operativa'][0]
    D.riscrivi_cella(r.cells[1], r.cells[1].text.strip().replace(
        'Numerazione comunale', 'Numerazione comunale — richiesta · Numerazione comunale — '
                                'ricerca e consultazione'))
    t = D.trova_tabella(d, 'Unità di rilascio', 'Pagine', 'Che cosa la fa cambiare')
    rr = {r.cells[0].text.strip(): r for r in t.rows}
    assert rr['mfe-operativa'].cells[1].text.strip().endswith('· Numerazione')
    D.riscrivi_cella(rr['mfe-operativa'].cells[1],
                     rr['mfe-operativa'].cells[1].text.strip() + ' (richiesta, ricerca)')
    assert 'Soggetti' not in rr['mfe-configurazione'].cells[1].text
    D.riscrivi_cella(rr['mfe-configurazione'].cells[1],
                     rr['mfe-configurazione'].cells[1].text.strip()
                     + ' · Allegati per UC · Soggetti per UC · Caricamento da file')
    D.sostituisci_immagine(d, 'Figura 1 —', os.path.join(IMG, 'bo_mfe.png'))
    fatti.append('unità di rilascio: numerazione e pagine di allegati, soggetti, caricamento')
    for a, b in [
        ('L’applicazione è stata pensata con diciotto pagine distribuite su quattro aree '
         'funzionali, cui si aggiunge una diciannovesima pagina',
         'L’applicazione è stata pensata con diciannove pagine distribuite su quattro aree '
         'funzionali, cui si aggiunge una ventesima pagina'),
        ('Seguono le diciotto pagine', 'Seguono le diciannove pagine'),
        ('Un back-office di diciotto pagine', 'Un back-office di diciannove pagine'),
        ('ripeterle in diciotto sezioni', 'ripeterle in diciannove sezioni'),
        ('le diciotto tabelle già progettate', 'le venti tabelle già progettate'),
        ('da cui vengono le diciotto tabelle', 'da cui vengono le venti tabelle'),
        ('le diciotto del modello dell’integrazione più le quattro del registro',
         'le venti del modello dell’integrazione più le quattro del registro'),
    ]:
        D.sostituisci(d, a, b, attese=1, fatti=fatti)

    # ─────────────────────────── 3. copertura
    t = D.trova_tabella(d, 'Tabella', 'Pagine che la governano', 'Copertura')
    rr = {r.cells[0].text.strip(): r for r in t.rows}
    D.riscrivi_cella(rr['ANSC_STATO_ATTO'].cells[1],
                     'Supervisione atti · Dettaglio atto · Numerazione — ricerca')
    D.riscrivi_cella(rr['ANSC_STATO_ATTO'].cells[2],
                     'Completa in lettura; in scrittura per fase, stato ed emergenza. Il numero '
                     'comunale vi si legge: lo scrive il concentratore al deposito, e la serie '
                     'sta in DATI_ATTO_ESTERNO.')
    D.riscrivi_cella(rr['SOGGETTI_USECASE'].cells[2], 'Completa.')
    dopo = rr['SOGGETTI_USECASE']._tr
    for c in COPERTURA:
        nuova = D.clona_riga(t, c)
        dopo.addnext(nuova._tr)
        dopo = nuova._tr
    fatti.append('copertura: PRENOTAZIONE_ATTI e DATI_ATTO_ESTERNO')

    # ─────────────────────────── 4. riferimenti, punti aperti, scheda, storia
    D.sostituisci(d, 'ANALISI_Integrazione-ANSC_v3.36.docx', 'ANALISI_Integrazione-ANSC_v3.39.docx',
                  attese=1, fatti=fatti, etichetta='[R1]')
    D.sostituisci(d, 'SPEC_API_Integrazione-ANSC_v0.1.docx e i contratti OpenAPI dei tre pod',
                  'SPEC_API_Integrazione-ANSC_v0.2.docx e i contratti OpenAPI dei quattro pod',
                  attese=1, fatti=fatti, etichetta='[R6]')
    t = D.trova_tabella(d, '#', 'Riferimento', 'Natura')
    assert t.rows[-1].cells[0].text.strip() == '[R10]'
    D.clona_riga(t, RIF)
    t = D.trova_tabella(d, '#', 'Questione', 'Perché è aperta')
    rr = {r.cells[0].text.strip(): r for r in t.rows}
    D.riscrivi_cella(rr['BO-3'].cells[1], 'Chiuso in questa versione. Dove vivano i numeri '
                                          'comunali staccati e non ancora consumati')
    D.riscrivi_cella(rr['BO-3'].cells[2],
                     'Deciso nell’analisi [R1]: nelle tabelle di Side PRENOTAZIONE_ATTI e '
                     'DATI_ATTO_ESTERNO, con una sequence per registro e anno, governate dal '
                     'servizio num-ansc-sipo. Le pagine della numerazione sono descritte su '
                     'quella base.')
    D.riscrivi_cella(rr['BO-4'].cells[1], 'Chiuso in questa versione. Se la numerazione '
                                          'comunale appartenga al back-office o alle maschere '
                                          'di SIPO')
    D.riscrivi_cella(rr['BO-4'].cells[2],
                     'Deciso nell’analisi [R1]: le pagine stanno nel back-office. Il percorso '
                     'ordinario non ne ha bisogno, perché il numero lo chiede il concentratore '
                     'al deposito; le pagine servono agli atti formati sulla web app, al parto '
                     'plurimo, alla registrazione a posteriori e al governo delle serie. In Side '
                     'sono voci del sottomenu dei registri.')
    assert t.rows[-1].cells[0].text.strip() == 'BO-29'
    for b in BO_NUOVI:
        D.clona_riga(t, b)
    scheda = d.tables[0]
    for r in scheda.rows:
        k = r.cells[0].text.strip()
        if k == 'Data consegna':
            D.riscrivi_cella(r.cells[1], '11/10/2026')
        elif k == 'Versione':
            D.riscrivi_cella(r.cells[1], '0.9')
        elif k == 'Documento':
            D.riscrivi_cella(r.cells[1], 'DISEGNO_Back-Office_ANSC_v0.9')
    D.storia(d, *STORIA)
    fatti.append('[R11], BO-3 e BO-4 chiusi, BO-30, scheda, storia')

    d.save(DST)
    dopo_c = commenti(DST)
    if dopo_c != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo_c} dopo')
    print('\n'.join('· ' + f for f in fatti))
    print(f'commenti preservati: {dopo_c}')
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))


if __name__ == '__main__':
    main()
