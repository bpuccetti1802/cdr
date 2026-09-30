# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.28 -> v3.29 (25/09/2026).

Tre cose, tutte nate da una revisione del committente:

  1. TERMINOLOGIA — ciò che il documento chiamava «documenti extra» sono gli **atti a testo
     libero**: l'atto che l'ufficiale redige liberamente e allega. Cambia anche il nome della
     colonna (fg_extra → fg_testo_libero), perché un contrassegno che non si chiama come la
     cosa che contrassegna si legge male.
  2. L'AGGANCIO di un atto formato sulla web app: l'USC lo cerca in ANSC per identificativo
     (R005), conferma, e SIPO ne registra i riferimenti; poi prosegue la compilazione come su
     un atto già concluso. È la proposta ricevuta, e il documento la adotta.
  3. LO STACCO DEL NUMERO COMUNALE prima che l'atto sia steso altrove: senza, due atti
     possono ricevere lo stesso numero.

Più due lacune del back-office trovate verificando le tabelle: i documenti di un atto e i
numeri comunali allocati non avevano una superficie.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_29_atti_testo_libero.py
"""
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.28.docx')
DST = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.29.docx')
IMG = os.path.join(BASE, 'strumenti', 'img')

if os.path.exists(DST):
    os.remove(DST)
shutil.copy(SRC, DST)
doc = docx.Document(DST)
fatti = []


def par(inizio):
    t = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(t) != 1:
        raise SystemExit(f'attesa 1 occorrenza di «{inizio[:60]}», trovate {len(t)}')
    return t[0]


def riga_con(t, prima_cella):
    for r in t.rows:
        if r.cells[0].text.strip() == prima_cella:
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


MOD = D.trova_tabella(doc, 'colonna', 'tipo', 'note')

for t in doc.tables:
    if t.rows[0].cells[0].text.strip().lower().startswith(('area organizzativa', 'progetto')):
        for nome, val in (('Data consegna', '25/09/2026'), ('Versione', '3.29')):
            try:
                D.riscrivi_cella(riga_con(t, nome).cells[1], val)
            except SystemExit:
                pass

D.storia(doc, '25/09/2026', '3.29',
         'Flusso operativo · SIPO e la web app · Back-office · Adeguamenti database · '
         'Open Point',
         'Recepita la revisione del committente. Ciò che il documento chiamava «documenti '
         'extra» sono gli atti a testo libero, e il contrassegno sulla tabella dei documenti '
         'prende quel nome. Aggiunto l’aggancio di un atto formato sulla web app: l’ufficiale '
         'lo cerca in ANSC per identificativo con R005, conferma, e SIPO ne registra i '
         'riferimenti, proseguendo poi la compilazione come su un atto già concluso; con esso '
         'lo stacco del numero comunale, che deve uscire da SIPO anche quando l’atto si forma '
         'altrove. Registrata come aperta la questione di perimetro sollevata: quali atti '
         'formati fuori SIPO si recuperano e con quale priorità. Aggiunte al back-office le '
         'due superfici che mancavano, i documenti di un atto e i numeri comunali allocati.')


# ═════════════════════════ 1. la terminologia
D.sostituisci(doc, 'e gli eventuali documenti extra, che la configurazione non elenca ma che '
                   'l’ufficiale ritiene di allegare',
              'e gli eventuali atti a testo libero, cioè gli atti che l’ufficiale redige '
              'liberamente e allega, che la configurazione non elenca')
D.sostituisci(doc, 'gli extra sono ammessi e contrassegnati come tali',
              'gli atti a testo libero sono ammessi e contrassegnati come tali')
D.sostituisci(doc, 'I primi sono verificabili dal pre-filtro, i secondi no: sono ammessi, '
                   'contati e distinti, non ignorati.',
              'I primi sono verificabili dal pre-filtro, i secondi no: sono ammessi, contati e '
              'distinti, non ignorati.')
n = D.sostituisci(doc, 'fg_extra', 'fg_testo_libero')
fatti.append(f'fg_extra → fg_testo_libero in {n} punti')
t = D.tabella_colonne(doc, 'id_allegato')
r = riga_con(t, 'fg_testo_libero')
D.riscrivi_cella(r.cells[2], '⚠️ Aggiunta: 0/1 secondo che il documento sia fra quelli che '
                             'l’UC prevede oppure sia un atto a testo libero, redatto '
                             'dall’ufficiale e allegato. I secondi sono ammessi ma non '
                             'verificabili dal pre-filtro, e vanno distinti perché la loro '
                             'assenza non è un errore e la loro presenza non soddisfa un '
                             'requisito.')
fatti.append('terminologia: atti a testo libero')


# ═════════════════════════ 2. l'aggancio di un atto formato sulla web app
ancora = D.h(doc, 2, 'La continuità operativa quando l’indisponibilità è nostra')._p
D.para(doc, ancora, 'L’aggancio di un atto formato sulla web app', stile='Heading 2')
for t in [
    'Un atto può nascere sulla web app: perché è un atto a testo libero, perché SIPO era fermo, '
    'o perché il caso non è fra quelli che il Comune ha configurato. In tutti questi casi l’atto '
    'esiste in ANSC, è formato e firmato, e non esiste in SIPO. La domanda posta in sede di '
    'revisione è se debba essere riportato in SIPO, e perché questo valga per gli atti a testo '
    'libero e non per gli altri atti formati sulla web app.',
    '⚠️ Il criterio non è il tipo di atto: è che cosa SIPO deve poter fare dopo. Finché SIPO '
    'resta l’oracolo del Comune — certifica, cerca, annota, conta, risponde allo sportello — un '
    'atto che in SIPO non c’è è un atto che il Comune non sa di avere. Se invece per una '
    'categoria di atti si accetta che l’oracolo sia ANSC, allora per quella categoria il '
    'recupero non serve. La distinzione fra atti a testo libero e altri atti non regge da sola: '
    'quella che regge è fra ciò che SIPO deve continuare a sapere e ciò che può ignorare. È una '
    'decisione di perimetro e di priorità, non tecnica, ed è registrata come punto aperto '
    '(OP-66).',
    'Dove il recupero si fa, il modo è quello proposto in revisione, e il presente documento lo '
    'adotta perché è il più economico dei modi possibili: si aggancia l’atto prima di compilarlo, '
    'invece di ricostruirlo e poi cercare di farlo coincidere.',
]:
    D.para(doc, ancora, t)
D.tabella(doc, ancora, [
    ('Passo', 'Che cosa accade', 'Nota'),
    ('1 — Avvio della stesura in SIPO', 'L’ufficiale apre la maschera dell’atto e dichiara che '
     'l’atto esiste già in ANSC.', 'È una scelta esplicita all’inizio, non una correzione a '
     'posteriori.'),
    ('2 — Ricerca in ANSC (R005)', 'Si cerca l’atto per identificativo nazionale, o per gli '
     'estremi che la consultazione accetta.',
     '⚠️ La ricerca per identificativo nazionale in SIPO oggi non esiste: va costruita (OP-29 e '
     'OP-30). È la stessa interfaccia che serve all’atto collegato.'),
    ('3 — Conferma della selezione', 'SIPO registra i riferimenti dell’atto: identificativo '
     'nazionale, numero comunale, data e ora di formazione, caso d’uso.',
     'Sono esattamente i quattro dati che il capitolo «Che cosa resta sull’atto» dichiara '
     'debbano restare: l’aggancio li scrive all’origine invece che alla fine.'),
    ('4 — Compilazione', 'L’ufficiale compila l’atto in SIPO come se stesse integrando un atto '
     'già concluso.', 'Regge perché SIPO già consente di modificare e integrare un atto dopo la '
     'conclusione: la prealimentazione non introduce un caso nuovo.'),
    ('⚠️ Che cosa NON accade', 'Nessun deposito, nessuna firma, nessun consumo di un '
     'identificativo.', 'L’atto è già formato in ANSC. L’aggancio ne registra l’esistenza: non '
     'lo riforma, e il percorso di finalizzazione non si esegue.'),
], MOD)
D.para(doc, ancora, 'L’aggancio di un atto formato sulla web app, passo per passo.',
       corsivo=True)
for t in [
    'Sullo store di stato l’atto agganciato nasce già concluso: la fase è FIRMATO e lo stato '
    'replica quello dichiarato da ANSC, mentre l’origine dell’UC resta quella manuale perché '
    'nessuna logica lo ha determinato. ⚠️ È una riga che non ha attraversato le quattro fasi, e '
    'la vista di supervisione deve poterlo distinguere: un atto agganciato non è un atto rimasto '
    'indietro.',
]:
    D.para(doc, ancora, t)

D.para(doc, ancora, 'Lo stacco del numero comunale', stile='Heading 3')
for t in [
    'Il numero comunale lo assegna il Comune: nessun servizio di ANSC lo attribuisce, e '
    'l’identificativo nazionale non lo contiene. Ne discende una conseguenza che vale per '
    'entrambi i percorsi e che la revisione ha richiamato: se un atto si forma sulla web app, il '
    'numero deve comunque uscire da SIPO, perché SIPO è l’unico luogo in cui la serie è '
    'governata. Un numero scelto a mano sulla web app è un numero che nessuno ha riservato.',
    'Serve quindi una funzione di stacco: l’ufficiale ottiene da SIPO il numero prima di stendere '
    'l’atto altrove, il numero risulta impegnato da quel momento, e l’aggancio lo ricongiunge '
    'all’atto quando questo torna in SIPO. ⚠️ Un numero staccato e non utilizzato va chiuso '
    'esplicitamente, altrimenti la serie presenta buchi che nessuno sa spiegare a distanza di '
    'tempo — è lo stesso problema della prenotazione degli identificativi per il parto plurimo, '
    'dove un atto chiuso lascia consumato il proprio identificativo.',
    'La funzione si lega a tre questioni già aperte e le rende più stringenti: l’allocazione in '
    'concorrenza alla scala di Roma (OP-44), la numerazione degli atti lavorati in emergenza '
    '(OP-47) e la colonna in cui il numero risiede, che sull’atto di stato civile oggi non esiste '
    '(OP-60).',
]:
    D.para(doc, ancora, t)
fatti.append('nuova sezione: aggancio e stacco del numero comunale')


# ═════════════════════════ 3. le due superfici mancanti nel back-office
t = D.trova_tabella(doc, 'schermata', 'scopo', 'azioni principali')
nuova = D.clona_riga(t, (
    'Documenti dell’atto',
    'ALLEGATO: i documenti di un atto con il tipo, l’impronta, lo stato restituito da ANSC e il '
    'contrassegno di atto a testo libero. È la superficie che mancava alla tabella dei '
    'documenti.',
    'Consulta ed apri un documento, verifica lo stato della scansione, ricarica un file '
    'respinto, distingue i documenti previsti dagli atti a testo libero.', 'USC / Supporto'))
riga_con(t, 'Dettaglio atto')._tr.addnext(nuova._tr)
D.clona_riga(t, (
    'Numeri comunali',
    'I numeri staccati, quelli utilizzati e quelli chiusi, con l’atto a cui ciascuno è legato: '
    'serve a governare la serie quando l’atto si forma sulla web app o in emergenza.',
    'Stacca un numero, lega un numero a un atto agganciato, chiudi un numero non utilizzato, '
    'consulta i buchi della serie.', 'USC / Amministratore'))
fatti.append('back-office: schermate dei documenti e dei numeri comunali')

t = D.trova_tabella(doc, 'schermata / componente', 'tipo intervento')
D.clona_riga(t, ('Avvio della stesura — atto già formato in ANSC', 'Maschera nuova',
                 'Ricerca dell’atto in ANSC per identificativo (R005) e conferma: SIPO registra '
                 'identificativo nazionale, numero comunale, data e ora di formazione e caso '
                 'd’uso, poi la compilazione prosegue come su un atto concluso.', 'Alta'))
D.clona_riga(t, ('Stacco del numero comunale', 'Funzione nuova',
                 'Ottiene da SIPO il numero prima che l’atto sia steso sulla web app; il numero '
                 'risulta impegnato e si ricongiunge all’atto con l’aggancio.', 'Alta'))
fatti.append('impatti sul front-end: aggancio e stacco del numero')


# ═════════════════════════ 4. la citazione normativa riapparsa
D.sostituisci(doc, 'il registro di emergenza previsto dall’articolo 10 del decreto ministeriale',
              'il registro di emergenza previsto dall’ordinamento, la cui base normativa resta '
              'da accertare (OP-49)', attese=1, etichetta='citazione non verificabile',
              fatti=fatti)


# ═════════════════════════ 5. open point
t = D.trova_tabella(doc, '#', 'tema', 'questione')
D.clona_riga(t, ('OP-66', 'Perimetro e priorità del recupero in SIPO degli atti formati fuori',
                 'Quali atti formati sulla web app — a testo libero, durante un fermo di SIPO, '
                 'o per casi non configurati — debbano essere riportati in SIPO, e con quale '
                 'priorità. Il criterio proposto è che vi rientri ciò su cui SIPO deve '
                 'continuare a rispondere: certificazione, ricerca, annotazioni, conteggi. È '
                 'una decisione di perimetro, non tecnica.', 'Aperto', 'Cliente', 'Alta'))
fatti.append('OP-66 aperto')


# ═════════════════════════ 6. figure
D.sostituisci_immagine(doc, 'Il flusso operativo nelle sue quattro fasi.',
                       os.path.join(IMG, 'flusso_operativo.png'))
D.sostituisci_immagine(doc, 'Wireframe — Supervisione atti',
                       os.path.join(IMG, 'bo_supervisione.png'))
fatti.append('figure aggiornate alla nuova terminologia')

doc.save(DST)

import zipfile   # noqa: E402
z = zipfile.ZipFile(DST)
xml = z.read('word/document.xml').decode()
ncom = z.read('word/comments.xml').decode().count('<w:comment ')
assert ncom == 23, f'commenti persi: {ncom}'
assert 'fg_extra' not in xml, 'resta fg_extra'
print('\n'.join(' · ' + f for f in fatti))
print('commenti:', ncom, '· capitoli/tabelle/immagini:', D.riepilogo(DST))
print('scritto:', os.path.relpath(DST, BASE))
