# -*- coding: utf-8 -*-
"""v3.54 dell'analisi ANSC, dalla v3.53 salvata dall'utente (28/10/2026).

Decisioni dell'utente sulle incongruenze segnalate nella v3.53:

  1. il servizio della numerazione si chiama be-numeratore-comunale (non più num-ansc-sipo né
     be-numerazione-comunale);
  2. le API della numerazione sono le quattro descritte nel capitolo del servizio (POST /richiesta,
     GET /ricerca, PUT /utilizza, PUT /scarta); ciò che non coprono va in OP-79;
  3. prefisso DATA_ per le colonne DATE, TS_ per le TIMESTAMP;
  4. ANSC_CFG_UC senza vincolo di unicità;
  5. le righe commentate delle DDL sono colonne e vincoli tolti di proposito, lasciati per memoria;
  6. riallineare tutto (gli altri documenti hanno i loro script).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_54_numeratore_ts.py
"""
import os
import re
import sys

import docx
from docx.text.paragraph import Paragraph

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402
from v3_53_ddl_sviluppatori import sost, par, via, riga_tabella, dopo  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.53.docx')
OUT = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.54.docx')
IMG = os.path.join(QUI, 'img')

SERVIZIO = 'be-numeratore-comunale'

# operationId delle quattro operazioni, come li fissa ansc-numerazione-v1.yaml
OPERAZIONI = [
    ('POST /api/v1/numerazione-comunale/richiesta', 'creaRichiesta',
     'Stacca uno o più numeri consecutivi e li registra come disponibili'),
    ('GET /api/v1/numerazione-comunale/ricerca', 'cercaNumeri',
     'Ricerca paginata dei numeri staccati'),
    ('PUT /api/v1/numerazione-comunale/utilizza', 'utilizzaNumero',
     'Segna un numero come utilizzato e registra l’atto SIPO'),
    ('PUT /api/v1/numerazione-comunale/scarta', 'scartaNumero',
     'Segna un numero come scartato'),
]


# ───────────────────────────────────────────────────────────── primitive locali
def storia_tbl(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith('Data')
                and 'Versione' in t.rows[0].cells[1].text)


def paragrafi_fuori_storia(doc):
    """Paragrafi del corpo e delle celle, esclusa la Storia del Documento."""
    st = storia_tbl(doc)._tbl
    yield from doc.paragraphs
    for t in doc.tables:
        if t._tbl is st:
            continue
        visti = set()
        for r in t.rows:
            for c in r.cells:
                if c._tc in visti:
                    continue
                visti.add(c._tc)
                yield from c.paragraphs


def sostituisci_parola(doc, mappa):
    """Sostituisce parole intere (per regex) fuori dalla Storia; restituisce i conteggi."""
    conti = {k: 0 for k in mappa}
    regole = [(k, re.compile(r'(?<![A-Za-z0-9_])' + re.escape(k) + r'(?![A-Za-z0-9_])'), v)
              for k, v in mappa.items()]
    for p in paragrafi_fuori_storia(doc):
        testo = p.text
        if not any(rx.search(testo) for _, rx, _ in regole):
            continue
        for k, rx, v in regole:
            conti[k] += len(rx.findall(testo))
        for r in p.runs:
            t = r.text
            for _, rx, v in regole:
                t = rx.sub(v, t)
            if t != r.text:
                r.text = t
        resto = p.text
        if any(rx.search(resto) for _, rx, _ in regole):   # parola spezzata fra run
            for _, rx, v in regole:
                resto = rx.sub(v, resto)
            D.testo_di(p, resto)
    return conti


def cella(t, primo, col):
    return riga_tabella(t, primo).cells[col]


def tabella_con(doc, intestazione_prima_cella, riga_contiene):
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip() == intestazione_prima_cella and \
                any(riga_contiene in r.cells[0].text for r in t.rows):
            return t
    raise SystemExit('tabella non trovata: ' + riga_contiene)


# ───────────────────────────────────────────────────────────── 1. nome del servizio
def nome_servizio(doc):
    conti = sostituisci_parola(doc, {'num-ansc-sipo': SERVIZIO,
                                     'be-numerazione-comunale': SERVIZIO})
    print('nome del servizio:', conti)
    h = D.h(doc, 2, 'Il servizio ' + SERVIZIO, esatto=True)
    D.testo_di(h, 'Le API del servizio ' + SERVIZIO)


# ───────────────────────────────────────────────────────────── 2. le quattro operazioni
def operazioni(doc):
    p = par(doc, 'Le operazioni sono dieci e stanno nel contratto')
    D.testo_di(p, 'Le operazioni sono quattro — POST /richiesta, GET /ricerca, PUT /utilizza e '
                  'PUT /scarta — e sono descritte nel paragrafo «Le API del servizio '
                  f'{SERVIZIO}»; stanno nel contratto ansc-numerazione-v1.yaml [R7] e '
                  'l’inventario è nell’Appendice B. La tabella che segue dice quale operazione '
                  'serve a ciascun percorso e che cosa le quattro operazioni non coprono ancora '
                  '(OP-79).')

    t = tabella_con(doc, 'Percorso', 'Governo delle serie')
    nuovi = {
        'Atto formato in SIPO':
            'POST /richiesta per un numero, all’Anteprima e prima del deposito: il numero resta '
            'sull’atto (ANSC_STATO_ATTO.NUMERO_ATTO) e si riusa se il deposito è respinto. '
            'PUT /utilizza con l’atto di SIPO, a deposito riuscito; PUT /scarta, se l’atto è '
            'abbandonato. L’identificativo nazionale non si registra sul numero, perché '
            'PUT /utilizza non lo porta (OP-79).',
        'Atto formato sulla web app':
            'POST /richiesta; poi PUT /utilizza o PUT /scarta. Il passaggio a Evento acquisito, '
            'all’aggancio dell’atto in SIPO, non ha un’operazione (OP-79).',
        'Parto plurimo':
            'POST /richiesta con numeroElementi pari ai nati; ciascun numero segue poi la sorte '
            'del proprio atto.',
        'SIPO non disponibile':
            'Nessuna operazione: la registrazione a posteriori del numero già usato, con '
            'l’allineamento del numeratore, è da realizzare (OP-79).',
        'Governo delle serie':
            'GET /ricerca, per anno, stato, richiesta, atto, numero, intervallo di date e '
            'operatore. La lettura della serie con i suoi buchi e la creazione anticipata della '
            'riga dell’anno non hanno un’operazione (OP-79); la riga si crea comunque alla prima '
            'richiesta.',
    }
    for primo, testo in nuovi.items():
        D.riscrivi_cella(cella(t, primo, 2), testo)

    # i percorsi descritti nel capitolo: la registrazione a posteriori non è esposta
    t = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Percorso'
             and t.rows[0].cells[1].text.strip() == 'Quando si stacca')
    if True:
        c = cella(t, 'SIPO non disponibile', 2)
        D.riscrivi_cella(c, c.text.rstrip() + ' L’operazione non è ancora esposta dal servizio '
                         '(OP-79).')

    # lo stacco passo per passo
    sost(doc, 'Se la richiesta indica un atto di SIPO che ha già un numero Disponibile nella stessa '
              'serie — un deposito respinto e ritentato — il servizio restituisce quel numero, '
              'conferma e si ferma: il numeratore non avanza.',
         'La richiesta non indica l’atto di SIPO. Se un deposito respinto si ritenta, il riuso '
         'del numero spetta al chiamante, che lo conserva sull’atto (ANSC_STATO_ATTO.NUMERO_ATTO) '
         'e non ne chiede un altro: il numeratore non avanza.')
    sost(doc, 'nello stato Disponibile e legata all’atto se indicato:',
         'nello stato Disponibile e senza atto, che si lega poi con PUT /utilizza:')
    sost(doc, '— il percorso «SIPO non disponibile» — si svolge con lo stesso blocco',
         '— il percorso «SIPO non disponibile» — non ha ancora un’operazione nel servizio '
         '(OP-79). Quando sarà realizzata, dovrà svolgersi con lo stesso blocco')
    t = next(t for t in doc.tables if [c.text.strip() for c in t.rows[0].cells] ==
             ['Codice', 'Stato', 'Come ci si arriva'])
    c = cella(t, '4', 2)
    D.riscrivi_cella(c, 'Non viene impostato da questi servizi (OP-79).')


# ───────────────────────────────────────────────────────────── 3. DATA_ e TS_
RINOMINE = {
    'DATA_INS': 'TS_INS', 'DATA_UPD': 'TS_UPD',
    'DATA_RICERCA': 'TS_RICERCA',
    'DATA_NOTIFICA': 'TS_NOTIFICA', 'DATA_ESITO': 'TS_ESITO', 'DATA_SOLLECITO': 'TS_SOLLECITO',
    'DATA_ATTIVAZIONE': 'TS_ATTIVAZIONE',
    'dt_inserimento': 'ts_inserimento', 'dt_aggiornamento': 'ts_aggiornamento',
}


def colonne_temporali(doc):
    conti = sostituisci_parola(doc, RINOMINE)
    print('rinomine:', conti)
    # ANSC_LOG_AUDIT.DATA: la sola colonna senza prefisso
    sost(doc, '  DATA       TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,',
         '  TS_REGISTRAZIONE TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,')
    t = tabella_con(doc, 'Colonna', 'OPERATORE / DATA')
    D.riscrivi_cella(cella(t, 'OPERATORE / DATA', 0), 'OPERATORE / TS_REGISTRAZIONE')

    # la convenzione, dichiarata dove si descrive lo standard
    p = par(doc, 'La denominazione degli oggetti segue lo standard aziendale di nomenclatura [R5]')
    D.testo_di(p, p.text.rstrip() + ' Per le colonne temporali il prefisso dipende dal tipo: '
                  'DATA_ per le colonne DATE, TS_ per le colonne TIMESTAMP, compresi i campi '
                  'tecnici TS_INS e TS_UPD. Lo standard prevede il solo DATA_: lo scostamento '
                  'è una scelta del Comune, perché il nome dica se la colonna porta anche l’ora.')
    sost(doc, 'Gli identificativi sono NUMBER con sequenza; le date sono TIMESTAMP.',
         'Gli identificativi sono NUMBER con sequenza; le colonne temporali sono DATE, con '
         'prefisso DATA_, o TIMESTAMP, con prefisso TS_.')


# ───────────────────────────────────────────────────────────── 4. ANSC_CFG_UC senza UK
def cfg_uc_senza_uk(doc):
    p = par(doc, 'CONSTRAINT UK_ANSC_CFG_UC')
    succ = Paragraph(p._p.getnext(), p._parent)
    assert succ.text.strip() == 'UNIQUE (ID_VERSIONE, COD_UC_ANSC, ID_TIPO_ATTO),', succ.text
    via(succ._p)
    via(p._p)

    t = tabella_con(doc, 'Colonna', 'ID_UC_CFG')
    c = cella(t, 'ID_CONF_TIPO_ATTO', 2)
    D.riscrivi_cella(c, c.text.replace(' Non entra nella chiave di unicità.', ''))
    c = cella(t, 'ID_TIPO_ATTO', 2)
    D.riscrivi_cella(c, 'Tipo atto di SIPO a cui la riga si applica: è la colonna con cui il tipo '
                        'atto si lega all’UC. Vuota per un UC del catalogo che nessun tipo atto '
                        'raggiunge ancora.')
    p = par(doc, 'La riga è per UC e per tipo atto che lo raggiunge, dentro una baseline')
    D.testo_di(p, p.text.rstrip() + ' La tabella non ha vincolo di unicità, per scelta: che uno '
                  'stesso tipo atto non sia legato due volte allo stesso UC nella stessa baseline '
                  'lo controlla l’applicazione al salvataggio.')
    sost(doc, 'È la ragione per cui ID_VERSIONE è nella chiave di unicità della configurazione.',
         'È la ragione per cui ogni lettura della configurazione filtra su ID_VERSIONE.')


# ───────────────────────────────────────────────────────────── 5. le righe commentate
def righe_commentate(doc):
    p = par(doc, 'Il DDL è indicativo per Oracle 12.2')
    dopo(p, 'Alcune righe delle DDL sono commentate, cioè precedute da «--» in testa a una '
            'colonna o a un vincolo: sono colonne e vincoli tolti di proposito, lasciati per '
            'memoria della scelta. Non sono dimenticanze e non vanno ripristinati. Le righe di '
            'commento scritte in prosa, invece, spiegano l’istruzione che le segue.')
    sost(doc, 'Il vincolo IS JSON è sospeso nel DDL.',
         'Il vincolo IS JSON è stato tolto di proposito: la riga commentata nel DDL ne conserva '
         'la memoria.')
    t = tabella_con(doc, 'Colonna', 'ALIAS_CAMPO_SIPO')
    c = cella(t, 'ID_VERSIONE', 2)
    D.riscrivi_cella(c, c.text.rstrip() + ' Le colonne tolte — sezione, schema, tabella e colonna '
                     'di SIPO, business logic, dominio, valore predefinito — restano nel DDL come '
                     'righe commentate, per memoria.')


# ───────────────────────────────────────────────────────────── 6. Appendice B, riferimenti, OP
def appendice_b(doc):
    t = tabella_con(doc, 'Aspetto', 'Contesto e versione')
    c = cella(t, 'Contesto e versione', 1)
    D.riscrivi_cella(c, c.text.replace(
        '/ansc/v1 per tutte le operazioni.',
        '/ansc/v1 per le operazioni dei concentratori e del processo di automazione. Il servizio '
        'della numerazione fa eccezione: espone le proprie operazioni sotto '
        '/api/v1/numerazione-comunale, come descritte nel capitolo «La numerazione comunale».'))
    c = cella(t, 'Forma delle risposte', 1)
    D.riscrivi_cella(c, c.text.rstrip() + ' Il servizio della numerazione fa eccezione anche qui: '
                     'restituisce i dati senza busta, con 200 e il corpo o 204 senza corpo.')

    h = D.h(doc, 3, 'Servizio della numerazione', esatto=False)
    D.testo_di(h, 'Servizio della numerazione — ' + SERVIZIO)
    t = next(t for t in doc.tables if any('creaPrenotazione' in r.cells[1].text for r in t.rows))
    righe = [r for r in t.rows[1:]]
    for r in righe[len(OPERAZIONI):]:
        via(r._tr)
    for r, (percorso, ident, titolo) in zip(righe, OPERAZIONI):
        for cc, testo in zip(r.cells, (percorso, ident, titolo,
                                       'Non dichiarati dal servizio')):
            D.riscrivi_cella(cc, testo)


def riferimenti(doc):
    for t in doc.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == 'R8' and 'SPEC_API_Integrazione-ANSC_v0.3' in \
                    r.cells[1].text:
                D.riscrivi_cella(r.cells[1], 'SPEC_API_Integrazione-ANSC_v0.4.docx')
                return
    raise SystemExit('R8 non trovato')


def open_point(doc):
    t = next(t for t in doc.tables if t.rows[1].cells[0].text.startswith('OP-'))
    ultima = t.rows[-1]
    assert ultima.cells[0].text.strip() == 'OP-78', ultima.cells[0].text
    assert not D.ha_commenti(ultima._tr), 'OP-78 ha un commento'
    valori = [
        'OP-79', 'Operazioni mancanti nel servizio della numerazione',
        f'Le quattro operazioni di {SERVIZIO} coprono lo stacco, la ricerca e i passaggi a '
        'Utilizzato e a Scartato. Restano senza operazione: (a) la registrazione a posteriori '
        'di un numero già usato fuori da SIPO, con l’allineamento del numeratore (percorso '
        '«SIPO non disponibile»); (b) il passaggio allo stato 4 Evento acquisito, all’aggancio '
        'di un atto formato sulla web app; (c) la registrazione dell’identificativo nazionale '
        'sul numero (id_ansc_collegato), che PUT /utilizza non porta; (d) la lettura della '
        'serie con i suoi buchi e la creazione anticipata della riga dell’anno; (e) l’operatore '
        'reale della richiesta, oggi sempre «admin».',
        'Aperto', 'Analisi / Sviluppo', 'Alta']
    D.clona_riga(t, valori[:len(ultima.cells)])

    # OP-62: i contratti sono stati riallineati all'UC e al tipo atto; resta il resto
    c = riga_tabella(t, 'OP-62').cells[2]
    D.riscrivi_cella(c, 'I contratti del concentratore espongono ora la configurazione per caso '
                     'd’uso (casi-uso, ANSC_CFG_UC legata al tipo atto) e i campi come ANSC_CFG_CAMPO; '
                     'la supervisione è per municipio. Restano da allineare: le operazioni per la '
                     'configurazione di allegati, soggetti, formule e interrogazioni per UC, che il '
                     'back-office prevede; le categorie di eccezione, cinque nel contratto contro '
                     'le undici dell’analisi; l’esito qualificato del deposito (codTipoEsito), che '
                     'il contratto conserva e lo store di stato non registra più. I contratti sono la '
                     'fonte autoritativa: l’allineamento si fa su di essi e l’appendice si rigenera, '
                     'non si riscrive a mano.')
    c = riga_tabella(t, 'OP-62').cells[1]
    D.riscrivi_cella(c, 'Allineamento dei contratti al modello della configurazione')


FIGURE = [
    ('Schema ANSC_USR, con ALLEGATO', 'erd_ansc_usr.png'),
    ('Il flusso operativo nelle sue quattro fasi', 'flusso_operativo.png'),
]
WORKFLOW_PNG = os.path.join(IMG, 'workflow_ansc_sipo.png')


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    nome_servizio(doc)
    operazioni(doc)
    colonne_temporali(doc)
    cfg_uc_senza_uk(doc)
    righe_commentate(doc)
    appendice_b(doc)
    riferimenti(doc)
    open_point(doc)
    for didascalia, png in FIGURE:
        D.sostituisci_immagine(doc, didascalia, os.path.join(IMG, png))
    if '--workflow' in sys.argv:
        D.sostituisci_immagine(doc, 'Il processo di gestione dei dati ANSC', WORKFLOW_PNG)
    D.storia(doc, '28/10/2026', '3.54',
             'Modello dati · Adeguamenti database · La numerazione comunale · Open Point · '
             'Riferimenti · Appendice A · Appendice B',
             f'Il servizio della numerazione si chiama {SERVIZIO} e le sue operazioni sono le '
             'quattro del servizio (POST /richiesta, GET /ricerca, PUT /utilizza, PUT /scarta): '
             'percorsi, stacco e Appendice B riallineati, e quanto non coperto raccolto in OP-79. '
             'Colonne temporali con prefisso DATA_ se DATE e TS_ se TIMESTAMP, campi tecnici '
             'compresi. Tolto il vincolo di unicità di ANSC_CFG_UC. Dichiarato che le righe '
             'commentate delle DDL sono colonne e vincoli tolti di proposito. OP-62 riformulato sui '
             'contratti riallineati.')
    D.frontespizio(doc, '3.54', 'ANALISI_Integrazione-ANSC_v3.54')
    doc.save(OUT)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
