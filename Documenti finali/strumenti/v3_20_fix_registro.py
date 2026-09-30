# -*- coding: utf-8 -*-
"""Le tre righe OP erano finite nella tabella sbagliata.

⚠️ Lezione: `doc.tables[n]` è un indice POSIZIONALE. Inserendo tabelle nuove prima di quella
cercata, l'indice slitta — e una riga «OP-53» si può ritrovare in mezzo a un elenco di
percorsi API senza che nulla protesti. Le tabelle si cercano per intestazione
(`trova_tabella`), mai per numero.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D  # noqa: E402
from v3_20_notifiche import DOC, OP_NUOVI  # noqa: E402

doc = docx.Document(DOC)
codici = {r[0] for r in OP_NUOVI}

# 1. via le righe finite dove non dovevano
tolte = 0
for t in doc.tables:
    for r in list(t.rows):
        if r.cells[0].text.strip() in codici and len(t.columns) != 6:
            r._tr.getparent().remove(r._tr)
            tolte += 1
assert tolte == 3, tolte

# 2. il registro vero, cercato per intestazione
reg = D.trova_tabella(doc, 'tema', 'questione')
assert reg is not None and len(reg.columns) == 6, 'registro degli Open Point non trovato'
assert reg.rows[-1].cells[0].text.strip() == 'OP-52', reg.rows[-1].cells[0].text
for riga in OP_NUOVI:
    D.clona_riga(reg, riga)

doc.save(DOC)
print('righe tolte:', tolte, '· registro ora:', len(reg.rows) - 1, 'voci')
