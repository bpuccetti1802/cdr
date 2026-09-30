# -*- coding: utf-8 -*-
"""v3.14 — allineamento del resto del documento alla configurazione per UC.

⚠️ La rinomina non è meccanica. `ANSC_CFG_OPERAZIONE` → `ANSC_CFG_UC` cambia anche la
CARDINALITÀ (da una riga per Modello a una per UC), quindi le frasi che dicono «per ogni
tipo di evento e operazione» vanno riscritte, non solo rinominate.

⚠️ `ANSC_CFG_REGOLA` **resta**: la v3.14 tocca la sola specie di DETERMINAZIONE, che diventa
una condizione dell'UC. Le regole di controllo e di generazione continuano a esistere nella
forma dichiarativa, e sono quelle che l'utente ha chiesto restino configurate dall'operatore.
Sostituire ovunque il nome avrebbe cancellato due specie su tre.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.14.docx')

PC9 = [
    'La corrispondenza fra il Modello del Comune e l’UC di ANSC non è biunivoca: lo stesso '
    'Modello alimenta più UC secondo i valori dei dati. Occorre quindi un luogo in cui '
    'esprimere la condizione che sceglie, e le alternative erano tre: un motore di regole di '
    'mercato, il riuso delle tabelle di regole già presenti in SIPO, oppure una condizione '
    'scritta nella configurazione del componente.',

    'Il motore di mercato resta escluso: le condizioni osservate sono congiunzioni di '
    'uguaglianze e appartenenze su campi noti, non richiedono inferenza, e un prodotto in più '
    'da governare avrebbe costo certo e beneficio nullo.',

    '⚠️ Sulla forma della condizione la presente versione corregge la scelta delle precedenti. '
    'Fino alla v3.13 la condizione era una terna campo, operatore, valore, scelta per non '
    'avere codice in configurazione. All’uso la forma si è rivelata insufficiente: i dati che '
    'discriminano l’UC non stanno tutti sull’atto — il discriminante delle nascite sta sul '
    'soggetto, la scelta dei documenti da allegare dipende da informazioni sparse su più '
    'tabelle — e per esprimerli con terne bisognerebbe aggiungere alla configurazione un modo '
    'per dichiarare da dove viene ciascun campo e come si lega all’atto: cioè reinventare le '
    'giunzioni, con una grammatica peggiore di quella che il sistema già possiede.',

    'Si adotta quindi l’interrogazione sul database di SIPO. Il fondamento è che, quando si '
    'deve stabilire quale caso d’uso si stia formando, SIPO ha già registrato l’atto sulle '
    'proprie tabelle: il dato su cui decidere esiste ed è quello vero. La condizione è '
    'associata all’UC configurato, non a una regola separata, e questo elimina anche il doppio '
    'luogo in cui l’UC poteva essere dichiarato.',

    'Il prezzo è dichiarato e non è piccolo: su una query non si dimostrano copertura e mutua '
    'esclusione, che erano il vantaggio della forma dichiarativa. Al loro posto stanno un '
    'ordine di valutazione esplicito fra gli UC dello stesso Modello e la simulazione su atti '
    'reali, che da ausilio diventa il controllo principale prima di attivare una baseline. '
    'Il rilievo mosso alle regole di SIPO — script non manutenibili in uno schema in esercizio '
    '— resta valido e viene affrontato dal regime imposto alla query: sola lettura, '
    'descrizione obbligatoria in lingua corrente, versionamento nella baseline, residenza in '
    'ANSC_USR e non in MATR_USR.',

    'Restano dichiarative, e configurate dall’operatore, le regole di controllo e di '
    'generazione: la revisione riguarda la sola specie di determinazione.',
]


def applica(doc):
    fatti = []

    # --------------------------------------------------------------- PC-9
    D.sostituisci(doc, 'PC-9 — Le regole sono dati dichiarativi',
                  'PC-9 — La determinazione dell’UC è una condizione sui dati di SIPO',
                  attese=1, etichetta='titolo PC-9', fatti=fatti)
    pc9 = D.h(doc, 2, 'PC-9 —')
    i = D.indice_di(doc, pc9)
    vecchi = [p for p in doc.paragraphs[i + 1:i + 4]]
    for p in vecchi:
        p._p.getparent().remove(p._p)
    dopo = doc.paragraphs[i + 1]._p
    for t in PC9:
        D.para(doc, dopo, t)
    fatti.append(f'PC-9 riscritta ({len(PC9)} paragrafi)')

    # ------------------------------------------------- le tre specie di regola
    t = D.trova_tabella(doc, 'Specie', 'Che cosa produce', 'Quando interviene')
    t.rows[1].cells[0].paragraphs[0].runs[0].text = 'Determinazione'
    D.testo_di(t.rows[1].cells[1].paragraphs[0],
               'L’UC di ANSC applicabile all’atto. ⚠️ Dalla v3.14 non è più una regola '
               'dichiarativa ma una condizione dell’UC configurato, in forma di '
               'interrogazione sui dati che SIPO ha già registrato (PC-9).')
    fatti.append('tabella delle tre specie: determinazione aggiornata')

    # --------------------------------------------------------- rinomine
    D.sostituisci(doc, 'ANSC_CFG_OPERAZIONE', 'ANSC_CFG_UC',
                  etichetta='ANSC_CFG_OPERAZIONE → ANSC_CFG_UC', fatti=fatti)
    D.sostituisci(doc, 'ID_CFG_OPERAZIONE', 'ID_UC_CFG', etichetta='chiave', fatti=fatti)
    # la cardinalità: le frasi che descrivevano la riga per Modello
    D.sostituisci(doc,
                  'È una riga per Modello, e riprende i valori che SIPO già possiede in '
                  'CONF_TIPO_ATTI.',
                  'È una riga per UC: il Modello, il tipo atto e la maschera sono suoi '
                  'attributi, e riprendono i valori che SIPO già possiede in CONF_TIPO_ATTI.',
                  etichetta='cardinalità della riga', fatti=fatti)
    return fatti


if __name__ == '__main__':
    doc = docx.Document(DOC)
    for f in applica(doc):
        print('  ·', f)
    doc.save(DOC)
    print('salvato')
