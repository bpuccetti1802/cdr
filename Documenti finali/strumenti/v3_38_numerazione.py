# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.37 → v3.38 (10/10/2026).

La numerazione comunale si EREDITA dal sistema Side del Comune di Milano («Processo di gestione
della numerazione comunale» v1.0 e la parte sui numeri del processo degli atti gemellari
v1.4), con un aggiustamento deciso dal committente: la serie è una SEQUENCE Oracle PER ANNO,
così che la numerazione riparta ogni anno.

⚠️ Tre fatti verificati che il disegno deve rispettare:
  * il numero comunale ha AL MASSIMO 5 CIFRE: `numeroatto` maxLength 5 in model_evento.yaml e
    in R005; «fino a 5 caratteri numerici» nella nota ANSC sulla nuova numerazione. Una serie
    sola per anno a Roma (144.285 atti/anno) supererebbe 99.999 → la serie è per REGISTRO e
    anno, che è anche la chiave di Side (ty_evento + nr_anno_atto) e la formulazione di OP-47;
  * una sequence NON garantisce l'assenza di buchi (un NEXTVAL in una transazione annullata è
    perso) né la contiguità di un blocco (le richieste concorrenti si intercalano), mentre Side
    pretende «in sequenza, non ci devono essere buchi» per i gemelli → assegnazioni
    serializzate, registrazione nella stessa transazione, buchi residui resi visibili;
  * Oracle 12.2 non ha ALTER SEQUENCE … RESTART (arriva con la 18c): «ripartire ogni anno» con
    una sequence sola non si può fare in modo pulito; con una sequence per anno sì, e la serie
    dell'anno precedente resta utilizzabile (scenario 5 di Side, a cavallo d'anno).
Le tabelle di Side non hanno tipi dichiarati (il documento elenca campi e descrizioni): i tipi
sono scelti qui e il documento lo dice.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/diagrammi_erd.py
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_38_numerazione.py
"""
import os
import shutil
import subprocess
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.37.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.38.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


PRENOTAZIONE = [
    ('Colonna', 'Tipo', 'Note'),
    ('id_prenotazione_atti', 'NUMBER (PK)',
     'Identificativo della richiesta, autogenerato. In Side lo produce un numeratore proprio; '
     'qui è una colonna IDENTITY.'),
    ('ty_evento', 'VARCHAR2(10)', 'Registro, secondo la decodifica ANSC_01: 1 nascita, 2 morte, '
                                  '3 matrimonio, 4 unione civile, 5 cittadinanza.'),
    ('nr_anno_atto', 'NUMBER(4)', 'Anno della serie da cui i numeri si staccano. Di norma '
                                  'l’anno corrente; può essere il precedente a cavallo d’anno.'),
    ('nr_id_ansc_da_prenotare', 'NUMBER(3)',
     'Quanti numeri si chiedono: uno, oppure tanti quanti i gemelli. Il nome è di Side.'),
    ('cd_operatore_prenotazione', 'VARCHAR2(40)', 'L’operatore che ha fatto la richiesta.'),
    ('ts_prenotazione', 'TIMESTAMP', 'Quando.'),
]

DATI_ESTERNO = [
    ('Colonna', 'Tipo', 'Note'),
    ('id_prenotazione_atti', 'NUMBER (PK, FK)', 'La richiesta da cui il numero proviene.'),
    ('nr_atto_comunale', 'NUMBER(5) (PK)',
     'Il numero. Cinque cifre al massimo, come dichiara ANSC: la sequence non può superarle.'),
    ('ty_evento / nr_anno_atto', 'VARCHAR2(10) / NUMBER(4)',
     'Aggiunte rispetto a Side: registro e anno ripetuti dalla richiesta, perché il database '
     'possa garantire che un numero non compaia due volte nella stessa serie (vincolo di '
     'unicità su registro, anno e numero). In Side l’unicità discende solo dal numeratore.'),
    ('cd_stato_numero_prenotato', 'CHAR(1)',
     'Dominio di Side (SIDe_17): 1 Disponibile, alla creazione; 2 Utilizzato; 3 Scartato; 4 '
     'Evento acquisito, quando l’atto formato altrove è stato agganciato in SIPO.'),
    ('id_ansc_collegato', 'VARCHAR2(50)', 'L’identificativo nazionale dell’atto che ha usato il '
                                          'numero.'),
    ('id_atto_sipo', 'NUMBER',
     'L’atto di SIPO. In Side la colonna si chiama id_evento e punta alla tabella degli eventi, '
     'che SIPO non ha: è lo stesso adattamento fatto per ALLEGATO.'),
    ('ts_allineamento_ansc_side / cd_operatore_allineamento', 'TIMESTAMP / VARCHAR2(40)',
     'Quando e da chi il numero è stato ricongiunto all’atto. Nomi di Side, conservati.'),
]

STATI_PERCORSI = [
    ('Percorso', 'Quando si stacca', 'Che cosa accade al numero'),
    ('Atto formato in SIPO', 'Un numero, dopo il pre-filtro e prima del deposito (R009), perché '
                             'il modello evento lo porta.',
     'Legato all’atto da subito. Deposito riuscito → 2 Utilizzato, con l’identificativo '
     'nazionale. Deposito respinto → resta all’atto e si riusa al tentativo successivo. Atto '
     'abbandonato prima del deposito → 3 Scartato.'),
    ('Atto formato sulla web app', 'Su richiesta dell’ufficiale, dalla pagina di richiesta, '
                                   'prima della validazione sulla web app.',
     '1 Disponibile → 2 Utilizzato o 3 Scartato, per mano dell’ufficiale. Con l’aggancio '
     'dell’atto in SIPO → 4 Evento acquisito.'),
    ('Parto plurimo', 'In blocco, tanti numeri quanti i gemelli, prima della prenotazione R009 '
                      'che li richiede tutti.',
     'Numeri consecutivi, garantiti dalla serializzazione. Ciascuno segue poi la sorte del '
     'proprio atto.'),
    ('SIPO non disponibile', 'Dopo, al ritorno di SIPO, indicando registro e anno: è lo '
                             'scenario di Side in cui l’atto è stato formato sulla web app con '
                             'il numero suggerito da ANSC.',
     'Si registra il numero già usato. ⚠️ Coincide con quello della serie solo se nel frattempo '
     'nessuno ha staccato numeri dello stesso registro: alla scala di Roma va verificato, non '
     'presunto.'),
    ('Modalità di emergenza', 'Al rientro, quando l’atto si deposita e quindi si forma.',
     'Come per l’atto formato in SIPO. L’ordine dei numeri segue l’ordine di formazione.'),
]

BLOCCHI = [
    ('h2', 'La numerazione comunale'),
    ('p', 'Il numero comunale è il progressivo che il Comune assegna all’atto e che viaggia nel '
          'modello evento. Nessun servizio di ANSC lo attribuisce: la piattaforma lo riceve e lo '
          'compone nell’identificativo nazionale, nella forma anno, numero nazionale, numero '
          'comunale, codice ISTAT. Governare la serie è quindi compito del Comune, e il luogo in '
          'cui la si governa è SIPO, anche quando l’atto si forma sulla web app.'),
    ('p', '⚠️ La gestione si eredita dal sistema Side del Comune di Milano [R12], con le sue '
          'due tabelle, i suoi quattro stati e i suoi casi; l’aggiustamento deciso per Roma è '
          'uno solo e riguarda la serie, che è una sequence di Oracle per anno, così che la '
          'numerazione riparta ogni anno.'),

    ('h3', 'La serie: per registro e per anno'),
    ('p', 'ANSC fissa il numero comunale ad al massimo cinque cifre: lo dichiarano il modello '
          'evento e il servizio di consultazione, e la nota ANSC sulla nuova numerazione degli '
          'atti digitali. La serie arriva quindi a 99.999 atti l’anno. Roma ne forma 144.285 '
          'secondo la ricognizione dei tipi atto: una serie sola per anno non basterebbe. Per '
          'registro basta con ampio margine — il più numeroso, quello dei decessi, conta 52.142 '
          'atti — ed è anche la chiave con cui Side organizza le richieste, per registro e anno.'),
    ('p', 'Ne discende una sequence per registro e per anno: cinque l’anno, una per ciascun '
          'valore della decodifica ANSC_01. Il nome segue lo standard di nomenclatura [R5]: '
          'SEQ_NUM_COMUNALE_<registro>_<anno>, per esempio SEQ_NUM_COMUNALE_2_2027 per le morti '
          'del 2027. [DA VERIFICARE: che ANSC si attenda il progressivo comunale per registro; le '
          'fonti lo implicano ma non lo dichiarano (OP-73).]'),
    ('v', ('Riparte ogni anno per costruzione. ', 'La sequence dell’anno nuovo comincia da '
           'uno. Su Oracle 12.2 una sequence sola non si potrebbe azzerare in modo pulito — il '
           'riavvio esplicito arriva con la versione 18 — e un azzeramento manuale a mezzanotte '
           'del 31 dicembre sarebbe l’operazione più fragile dell’anno.')),
    ('v', ('La serie dell’anno prima resta disponibile. ', 'A cavallo d’anno capita di dover '
           'staccare un numero della serie precedente: Side lo prevede esplicitamente, quando il '
           'gestionale torna disponibile dopo la fine dell’anno. Con una sequence per anno basta '
           'indicare l’anno.')),
    ('v', ('Cinque cifre come limite del database. ', 'MAXVALUE 99999: superarlo fa fallire '
           'lo stacco invece di produrre un numero che ANSC rifiuterebbe.')),
    ('v', ('Nessun numero perso per la cache. ', 'NOCACHE e ORDER: la cache delle sequence '
           'scarta i valori non consumati a ogni riavvio dell’istanza, e ogni valore scartato '
           'sarebbe un buco nella serie.')),
    ('v', ('Si creano in anticipo. ', 'Un job crea il 1° dicembre le sequence dell’anno '
           'successivo, e la procedura che le crea si esegue anche all’installazione per l’anno in '
           'corso. Lo stacco non crea mai una sequence: se manca, fallisce, ed è un guasto '
           'visibile invece di una struttura nata a metà giornata.')),

    ('h3', 'Due garanzie che una sequence non dà da sola'),
    ('p', '⚠️ Una sequence assegna numeri unici, ma non garantisce due cose che la numerazione '
          'degli atti richiede, e che Side pretende in modo esplicito per i gemelli — numeri in '
          'sequenza, senza buchi.'),
    ('v', ('La contiguità di un blocco. ', 'Due richieste contemporanee sullo stesso registro si '
           'intercalano: i numeri di una richiesta di tre gemelli possono uscire come 140, 142, '
           '143. Per questo ogni stacco, singolo o in blocco, si serializza: il concentratore '
           'acquisisce un blocco esclusivo su DATI_ATTO_ESTERNO, prende i valori, registra le '
           'righe e conferma, e solo allora il blocco si rilascia. Le letture non sono bloccate. '
           'Alla scala di Roma si tratta di qualche centinaio di stacchi al giorno, ciascuno di '
           'pochi millisecondi: il costo della serializzazione è trascurabile.')),
    ('v', ('L’assenza di buchi. ', 'Un valore preso dalla sequence in una transazione che poi '
           'fallisce è perduto: la sequence non torna indietro. Il disegno riduce la finestra al '
           'minimo — il valore si registra nella stessa transazione in cui lo si prende, prima di '
           'qualunque chiamata ad ANSC — e rende visibile ciò che resta: un numero mancante fra '
           'uno e l’ultimo registrato della serie è un buco, e la pagina del back-office lo '
           'mostra come tale. Un numero staccato e non usato, invece, non è un buco: è un numero '
           'Scartato, e lo dice.')),
    ('p', 'Il nome della sequence lo compone il concentratore da due valori controllati — il '
          'registro, che deve appartenere alla decodifica ANSC_01, e l’anno, di quattro cifre — '
          'e mai da un testo ricevuto: è un nome che finisce in un’istruzione SQL.'),

    ('h3', 'Le strutture ereditate da Side'),
    ('p', 'Due tabelle, con i nomi di origine. PRENOTAZIONE_ATTI registra ogni richiesta; '
          'DATI_ATTO_ESTERNO registra ogni numero staccato e ciò che ne è stato. Side dichiara '
          'campi e significato ma non i tipi, che sono scelti qui.'),
    ('p', 'PRENOTAZIONE_ATTI — la richiesta di numeri.'),
    ('tab', (PRENOTAZIONE, [1.9, 1.3, 3.4])),
    ('p', 'DATI_ATTO_ESTERNO — ogni numero e la sua sorte.'),
    ('tab', (DATI_ESTERNO, [1.9, 1.3, 3.4])),
    ('p', '⚠️ Un adattamento rispetto all’uso di Side. In Side DATI_ATTO_ESTERNO raccoglie i '
          'soli numeri chiesti per atti da redigere sulla web app, mentre gli atti formati nel '
          'gestionale prendono il numero da un numeratore interno. In SIPO tutti i numeri '
          'passano da qui, anche quelli del percorso ordinario: è ciò che rende la serie '
          'verificabile in un luogo solo, e un buco spiegabile. Il nome della tabella resta '
          'quello di origine, benché non descriva più soltanto atti esterni.'),
    ('p', 'Il numero dell’atto depositato resta anche sullo store di stato '
          '(ANSC_STATO_ATTO.NUM_COMUNALE), dove lo leggono la supervisione e la costruzione del '
          'payload. La sede in DATI_ATTO_ESTERNO è quella della serie; quella sullo store è '
          'quella dell’atto.'),

    ('h3', 'I cinque percorsi'),
    ('tab', (STATI_PERCORSI, [1.4, 2.2, 3.0])),
    ('p', ''),
    ('p', 'Le pagine sono quelle di Side: una pagina di richiesta, che propone registro e anno '
          'correnti e un numero, e una di ricerca e consultazione delle richieste, con i cambi '
          'di stato. In SIPO stanno nel back-office, nella pagina della numerazione comunale.'),
    ('p', 'Resta da stabilire a quale anno appartenga il numero di un atto che si deposita il 31 '
          'dicembre e si firma il 1° gennaio. Nel percorso ordinario il numero si stacca al '
          'deposito, e l’anno è quello del deposito; se l’atto si forma nell’anno successivo, il '
          'numero andrebbe scartato e staccato di nuovo dalla serie nuova. La regola va '
          'confermata (OP-73).'),
]

DDL = [
    'CREATE TABLE ANSC_USR.PRENOTAZIONE_ATTI (',
    '  id_prenotazione_atti      NUMBER GENERATED ALWAYS AS IDENTITY,',
    '  ty_evento                 VARCHAR2(10 CHAR)  NOT NULL,',
    '  nr_anno_atto              NUMBER(4)          NOT NULL,',
    '  nr_id_ansc_da_prenotare   NUMBER(3)          NOT NULL,',
    '  cd_operatore_prenotazione VARCHAR2(40 CHAR)  NOT NULL,',
    '  ts_prenotazione           TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,',
    '  CONSTRAINT prenotazione_atti_pk PRIMARY KEY (id_prenotazione_atti),',
    '  CONSTRAINT prenotazione_atti_qta_ck CHECK (nr_id_ansc_da_prenotare >= 1)',
    ') TABLESPACE ANSC_USR;',
    '',
    'CREATE TABLE ANSC_USR.DATI_ATTO_ESTERNO (',
    '  id_prenotazione_atti      NUMBER             NOT NULL,',
    '  nr_atto_comunale          NUMBER(5)          NOT NULL,',
    '  ty_evento                 VARCHAR2(10 CHAR)  NOT NULL,',
    '  nr_anno_atto              NUMBER(4)          NOT NULL,',
    "  cd_stato_numero_prenotato CHAR(1 CHAR) DEFAULT '1' NOT NULL,",
    '  id_ansc_collegato         VARCHAR2(50 CHAR),',
    '  id_atto_sipo              NUMBER,',
    '  ts_allineamento_ansc_side TIMESTAMP,',
    '  cd_operatore_allineamento VARCHAR2(40 CHAR),',
    '  CONSTRAINT dati_atto_esterno_pk',
    '    PRIMARY KEY (id_prenotazione_atti, nr_atto_comunale),',
    '  CONSTRAINT dati_atto_esterno_serie_uk',
    '    UNIQUE (ty_evento, nr_anno_atto, nr_atto_comunale),',
    '  CONSTRAINT dati_atto_esterno_pren_fk FOREIGN KEY (id_prenotazione_atti)',
    '    REFERENCES ANSC_USR.PRENOTAZIONE_ATTI (id_prenotazione_atti),',
    '  CONSTRAINT dati_atto_esterno_stato_ck',
    "    CHECK (cd_stato_numero_prenotato IN ('1','2','3','4'))",
    ') TABLESPACE ANSC_USR;',
    'CREATE INDEX ANSC_USR.IX_DATI_ATTO_ESTERNO_ATTO',
    '  ON ANSC_USR.DATI_ATTO_ESTERNO (id_atto_sipo) TABLESPACE ANSC_USR;',
    "COMMENT ON COLUMN ANSC_USR.DATI_ATTO_ESTERNO.cd_stato_numero_prenotato IS",
    "  '1 Disponibile | 2 Utilizzato | 3 Scartato | 4 Evento acquisito (dominio SIDe_17).';",
    '',
    '-- Una sequence per registro (ANSC_01) e per anno. Modello, per le morti del 2027:',
    'CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_2_2027',
    '  START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 99999 NOCYCLE NOCACHE ORDER;',
    '',
    'CREATE OR REPLACE PROCEDURE ANSC_USR.PRC_CREA_SEQ_NUM_COMUNALE (p_anno IN NUMBER) AS',
    '  n NUMBER;',
    'BEGIN',
    '  FOR r IN 1 .. 5 LOOP   -- i cinque registri di ANSC_01',
    '    SELECT COUNT(*) INTO n FROM USER_SEQUENCES',
    "     WHERE SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_' || r || '_' || p_anno;",
    '    IF n = 0 THEN',
    "      EXECUTE IMMEDIATE 'CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_' || r || '_' ||",
    "        p_anno || ' START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 99999'",
    "        || ' NOCYCLE NOCACHE ORDER';",
    '    END IF;',
    '  END LOOP;',
    'END;',
    '/',
    '',
    'BEGIN',
    '  DBMS_SCHEDULER.CREATE_JOB(',
    "    job_name        => 'ANSC_USR.JOB_CREA_SEQ_NUM_COMUNALE',",
    "    job_type        => 'PLSQL_BLOCK',",
    "    job_action      => 'BEGIN ANSC_USR.PRC_CREA_SEQ_NUM_COMUNALE(",
    "                          EXTRACT(YEAR FROM SYSDATE) + 1); END;',",
    "    repeat_interval => 'FREQ=YEARLY; BYMONTH=12; BYMONTHDAY=1; BYHOUR=2',",
    '    enabled         => TRUE);',
    'END;',
    '/',
    '-- All’installazione: EXEC ANSC_USR.PRC_CREA_SEQ_NUM_COMUNALE(<anno corrente>);',
    '',
]

OGGETTI = [
    ('PRENOTAZIONE_ATTI', 'Le richieste di numeri comunali.',
     'Ereditata da Side con i nomi di origine: registro, anno, quanti numeri, chi e quando. Una '
     'riga per richiesta, anche quando i numeri sono più d’uno.'),
    ('DATI_ATTO_ESTERNO', 'Ogni numero comunale staccato e la sua sorte.',
     'Ereditata da Side. Stato Disponibile, Utilizzato, Scartato o Evento acquisito; '
     'identificativo nazionale e atto SIPO collegati. In SIPO vi passano tutti i numeri, anche '
     'del percorso ordinario.'),
    ('SEQ_NUM_COMUNALE_<registro>_<anno>', 'La serie dei numeri comunali.',
     'Una sequence per registro e per anno, creata in anticipo da un job: la numerazione '
     'riparte ogni anno. Non è una tabella; è il solo aggiustamento rispetto a Side, che usa '
     'un numeratore proprio.'),
]

NUOVI_OP = [
    ('OP-73', 'Serie della numerazione comunale',
     'Due conferme e un permesso. (a) Che ANSC si attenda il progressivo comunale per registro: '
     'le fonti lo implicano — cinque cifre al massimo, chiave di Side per registro e anno — ma '
     'non lo dichiarano. (b) A quale anno appartenga il numero di un atto depositato in un anno '
     'e formato nel successivo. (c) Il privilegio CREATE SEQUENCE concesso direttamente allo '
     'schema ANSC_USR, senza il quale la procedura eseguita dal job non crea le sequence.',
     'Aperto', 'Analisi / Fornitore ANSC / DBA', 'Alta'),
]

STORIA = (
    '10/10/2026', '3.38',
    'SIPO e la web app (nuovo «La numerazione comunale») · Adeguamenti database · Modello dati '
    '(ERD) · Appendice A · Open Point',
    'La numerazione comunale si eredita dal sistema Side del Comune di Milano: le tabelle '
    'PRENOTAZIONE_ATTI e DATI_ATTO_ESTERNO, i quattro stati di un numero e i casi del processo. '
    'L’aggiustamento deciso per Roma è la serie, che è una sequence Oracle per anno e quindi '
    'riparte ogni anno; è anche per registro, perché ANSC limita il numero a cinque cifre e Roma '
    'forma più di 99.999 atti l’anno. Documentate le due garanzie che una sequence non dà da '
    'sola — contiguità dei blocchi e assenza di buchi — e come il disegno le ottiene o le rende '
    'visibili. Tutti i numeri, anche del percorso ordinario, passano dal registro di Side. '
    'Chiusi OP-44 e OP-47, ridotto OP-60, nuovo OP-73. ERD rigenerato con la seconda riga '
    'della fascia operativa.')


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    modello = D.trova_tabella(d, 'Oggetto', 'Ruolo')

    # ─────────────────────────── 1. il paragrafo nuovo
    ancora = D.h(d, 2, 'La continuità operativa quando l’indisponibilità è nostra')._p
    for tipo, c in BLOCCHI:
        if tipo == 'p':
            D.para(d, ancora, c)
        elif tipo in ('h2', 'h3'):
            D.para(d, ancora, c, stile='Heading ' + tipo[1])
        elif tipo == 'v':
            D.voce(d, ancora, *c)
        elif tipo == 'tab':
            D.tabella(d, ancora, c[0], modello, c[1])
    fatti.append('nuovo paragrafo «La numerazione comunale»')
    D.sostituisci(d, 'La funzione si lega a tre questioni già aperte e le rende più stringenti: '
                     'l’allocazione in concorrenza alla scala di Roma (OP-44), la numerazione '
                     'degli atti lavorati in emergenza (OP-47) e la colonna in cui il numero '
                     'risiede, che sull’atto di stato civile oggi non esiste (OP-60).',
                  'Come la funzione si realizzi — la serie, le strutture ereditate da Side e i '
                  'percorsi — è descritto nel paragrafo «La numerazione comunale».',
                  attese=1, fatti=fatti, etichetta='stacco: rimando al paragrafo nuovo')

    # ─────────────────────────── 2. modello dati
    t = D.trova_tabella(d, 'Oggetto', 'Ruolo')
    assert t.rows[-1].cells[0].text.strip() == 'SOGGETTI_USECASE'
    for o in OGGETTI:
        D.clona_riga(t, o)
    D.sostituisci(d, 'Schema ANSC_USR, con ALLEGATO che vi è adiacente perché sta in MATR_USR.',
                  'Schema ANSC_USR, con ALLEGATO che vi è adiacente perché sta in MATR_USR, e con '
                  'le sequence della numerazione comunale, che non sono tabelle.',
                  attese=1, fatti=fatti, etichetta='didascalia dell’ERD')
    D.sostituisci_immagine(d, 'Schema ANSC_USR, con ALLEGATO',
                           os.path.join(IMG, 'erd_ansc_usr.png'))
    fatti.append('elenco degli oggetti e ERD')

    # ─────────────────────────── 3. Appendice A, prima di ANSC_NOTIFICA (quella dell'appendice)
    notif = [p for p in d.paragraphs
             if p.style.name == 'Heading 3' and p.text.strip() == 'ANSC_NOTIFICA']
    assert len(notif) >= 1
    a = notif[-1]._p
    D.para(d, a, 'PRENOTAZIONE_ATTI, DATI_ATTO_ESTERNO e le sequence della numerazione',
           stile='Heading 3')
    D.ddl(d, a, DDL)
    fatti.append('Appendice A: due tabelle, sequence, procedura e job')

    # ─────────────────────────── 4. Open Point e storia
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == '#'
             and any(r.cells[0].text.strip() == 'OP-72' for r in t.rows))
    righe = {r.cells[0].text.strip(): r for r in t.rows}
    D.riscrivi_cella(righe['OP-44'].cells[2],
                     'Chiuso nella v3.38. La serie è una sequence per registro e per anno, che '
                     'riparte ogni anno; gli stacchi sono serializzati, così che i blocchi per il '
                     'parto plurimo siano consecutivi; ogni numero è registrato nella tabella di '
                     'Side DATI_ATTO_ESTERNO nella stessa transazione in cui è preso. Si veda il '
                     'paragrafo «La numerazione comunale». Restano le conferme di OP-73.')
    D.riscrivi_cella(righe['OP-44'].cells[3], 'Chiuso')
    D.riscrivi_cella(righe['OP-47'].cells[2],
                     'Chiuso nella v3.38. In modalità di emergenza l’atto si prepara e non si '
                     'forma: il numero si stacca al rientro, quando l’atto si deposita, e l’ordine '
                     'dei numeri segue quindi l’ordine di formazione. Diverso è il caso in cui a '
                     'mancare sia SIPO: lì si applica lo scenario di Side descritto nel paragrafo '
                     '«La numerazione comunale», con la cautela che vi è dichiarata.')
    D.riscrivi_cella(righe['OP-47'].cells[3], 'Chiuso')
    D.riscrivi_cella(righe['OP-60'].cells[2],
                     'Ridotto nella v3.38. Per l’integrazione il numero ha due sedi: '
                     'DATI_ATTO_ESTERNO per la serie e ANSC_STATO_ATTO.NUM_COMUNALE per l’atto. '
                     'Resta da decidere se debba comparire anche sull’atto di stato civile di '
                     'SIPO, dove oggi non esiste una colonna dedicata, perché le maschere e le '
                     'stampe di SIPO lo mostrino. Da chiudere insieme a OP-29.')
    assert t.rows[-1].cells[0].text.strip() == 'OP-72'
    for op in NUOVI_OP:
        D.clona_riga(t, op)
    D.storia(d, *STORIA)
    fatti.append('OP-44 e OP-47 chiusi, OP-60 ridotto, OP-73, storia')

    d.save(DST)
    dopo = commenti(DST)
    if dopo != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo} dopo')
    print('\n'.join('· ' + f for f in fatti))
    print(f'commenti preservati: {dopo}')


if __name__ == '__main__':
    main()
