# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.23 -> v3.24 (24/09/2026).

Verifica di come le tre tabelle ricevute sono trattate, e ciò che mancava:

  1. ALLEGATI_USECASE rispetto alla scelta dell'UC — l'aggancio c'è (la chiave porta
     cd_usecase) ma ⚠️ è un riferimento PER VALORE, non una chiave esterna: l'ERD dichiarava
     una FK che il DDL non ha. Corretto, e reso esplicito che cosa ne discende.
  2. DOMINIO_DECODIFICA e VALORE_DOMINIO — il caricamento era già coperto (allineaDizionari,
     tre modalità); ⚠️ mancava la RICERCA. Aggiunte al contratto `cercaValori` (trasversale
     alle decodifiche) e il parametro `ricerca` su elencaValoriDizionario; il capitolo dei
     dizionari dichiara ora le interfacce e la schermata che le usa.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_24_dizionari_e_allegati.py
"""
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.23.docx')
DST = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.24.docx')
IMG = os.path.join(BASE, 'strumenti', 'img')

if os.path.exists(DST):
    os.remove(DST)
shutil.copy(SRC, DST)
doc = docx.Document(DST)
fatti = []


def par(inizio, stile=None):
    t = [p for p in doc.paragraphs
         if p.text.strip().startswith(inizio) and (stile is None or p.style.name == stile)]
    if len(t) != 1:
        raise SystemExit(f'attesa 1 occorrenza di «{inizio[:60]}», trovate {len(t)}')
    return t[0]


def riga_con(t, prima_cella):
    for r in t.rows:
        if r.cells[0].text.strip() == prima_cella:
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


MOD = D.trova_tabella(doc, 'colonna', 'tipo', 'note')


# ═════════════════════════════════════ 1. scheda e storia
for t in doc.tables:
    if t.rows[0].cells[0].text.strip().lower().startswith(('area organizzativa', 'progetto')):
        for nome, val in (('Data consegna', '24/09/2026'), ('Versione', '3.24')):
            try:
                D.riscrivi_cella(riga_con(t, nome).cells[1], val)
            except SystemExit:
                pass

D.storia(doc, '24/09/2026', '3.24',
         'Gestione dei dizionari · Costruzione del payload · Back-office · Appendice B · '
         'Open Point',
         'Verificato il trattamento delle tre strutture ricevute. Per la configurazione degli '
         'allegati è reso esplicito che l’elenco dei documenti discende dall’UC determinato e '
         'che il legame è un riferimento per valore, non una chiave esterna: lo schema '
         'entità-relazioni, che dichiarava una chiave esterna assente dal DDL, è stato '
         'corretto. Per i dizionari, al caricamento su comando — già previsto nelle tre '
         'modalità — si aggiunge la ricerca, che mancava: il contratto del componente espone '
         'ora la ricerca di un valore in tutte le decodifiche e la ricerca testuale dentro una '
         'decodifica, e il capitolo dichiara quali interfacce servono la schermata del '
         'back-office e quali la diagnosi. Registrato l’allineamento dei contratti alla '
         'nomenclatura per caso d’uso (OP-62).')


# ═════════════════════════════════════ 2. gli allegati e l'UC determinato
ancora = par('Gli allegati provengono dallo stesso file dei campi')._p
D.para(doc, ancora,
       'Il legame con l’UC va letto con precisione, perché ne discende il momento in cui '
       'l’elenco dei documenti diventa noto. La chiave della configurazione porta il codice '
       'dell’UC come ANSC lo pubblica: finché la fase A non ha determinato l’UC non esiste '
       'alcun elenco di documenti da chiedere, e chiederli prima — sulla base del Modello — '
       'significa chiedere quelli sbagliati a chi si presenta allo sportello.')
D.para(doc, ancora,
       '⚠️ Il legame è un riferimento per valore e non una chiave esterna verso il catalogo '
       'degli UC, per la stessa ragione per cui lo è in ANSC_CFG_UC: si nomina il codice '
       'pubblicato da ANSC, non la riga locale, così che la configurazione possa nominare un '
       'UC prima che il catalogo sia stato ricaricato. Ne discende che l’esistenza dell’UC va '
       'verificata dall’importazione e dai controlli di attivazione della baseline, non da un '
       'vincolo del database: è una scelta di disegno e come tale va presidiata.')
fatti.append('allegati: esplicitato il legame con l’UC determinato (per valore)')


# ═════════════════════════════════════ 3. le interfacce dei dizionari
ancora = D.h(doc, 2, 'Fruizione da parte di SIPO')._p
D.para(doc, ancora, 'Ricerca e caricamento: le interfacce dei dizionari', stile='Heading 2')
for t in [
    'Le due tabelle replicate servono tre usi diversi, e conviene dire quale interfaccia serve '
    'quale uso, perché la risposta non è la stessa. Le maschere di SIPO leggono i valori in '
    'base dati, sulla vista di sola lettura: è la scelta dichiarata in PC-8 e non cambia. Il '
    'back-office e la diagnosi usano invece le interfacce del componente, che sono di due '
    'specie — quelle che caricano e quelle che cercano.',
    'Il caricamento era già previsto: è il comando di allineamento, nelle tre modalità '
    'completa, selettiva e di simulazione, con lo storico degli aggiornamenti eseguiti. La '
    'ricerca mancava, ed è ciò che questa versione aggiunge al contratto: ⚠️ senza di essa, '
    'per sapere in quale decodifica vive un valore di cui si conosce la descrizione, occorre '
    'aprire le decodifiche una per una — centoquarantatré — oppure interrogare il database a '
    'mano. È il lavoro che serve ogni volta che si riconcilia un valore di SIPO con quello di '
    'ANSC e ogni volta che un rifiuto cita un codice senza dire da quale tabella provenga.',
]:
    D.para(doc, ancora, t)
D.tabella(doc, ancora, [
    ('Operazione', 'Che cosa fa', 'Chi la usa'),
    ('allineaDizionari', 'Comanda l’allineamento nelle tre modalità: completa, selettiva su '
     'una o più decodifiche, simulazione che mostra il differenziale senza scrivere nulla.',
     'Schermata «Dizionari ANSC», ruolo Amministratore.'),
    ('elencaAggiornamenti, leggiAggiornamento', 'Lo storico dei caricamenti, con operatore, '
     'istante, esito e conteggi, e il dettaglio di un singolo aggiornamento.',
     'Stessa schermata, sezione dello storico; Auditor in lettura.'),
    ('elencaDizionari, leggiDizionario', 'Il catalogo locale: quali decodifiche sono '
     'replicate, con la versione dichiarata da ANSC e il numero di valori.',
     'Stessa schermata, elenco principale.'),
    ('elencaValoriDizionario', 'I valori di una decodifica, per difetto i soli in corso di '
     'validità, con ricerca testuale su codice e descrizione e paginazione.',
     'Apertura di una decodifica dalla schermata; diagnosi.'),
    ('cercaValori', '⚠️ Nuova: cerca un testo nel codice e nella descrizione di tutte le '
     'decodifiche replicate e restituisce, per ciascuna corrispondenza, la decodifica di '
     'appartenenza.', 'Riconciliazione dei valori fra SIPO e ANSC (OP-59); diagnosi di un '
     'rifiuto che cita un codice.'),
], MOD)
D.para(doc, ancora, 'Le interfacce del componente dei dizionari e la schermata che le usa.',
       corsivo=True)
for t in [
    'Due precisazioni sul confronto dei testi, che non sono dettagli implementativi ma '
    'condizioni perché la ricerca funzioni. Il corpus pubblicato non è uniforme quanto ad '
    'accenti e apostrofi tipografici — è la stessa ragione per cui sei descrizioni di allegato '
    'non trovano il proprio codice (OP-50) — e una ricerca che non li normalizzi non trova '
    'ciò che esiste. E la validità temporale resta applicata in lettura: i valori cessati si '
    'cercano, perché servono a interpretare gli atti già formati, ma non si propongono in '
    'scelta.',
]:
    D.para(doc, ancora, t)
fatti.append('nuova sezione «Ricerca e caricamento: le interfacce dei dizionari»')


# ═════════════════════════════════════ 4. il back-office
t = D.trova_tabella(doc, 'schermata', 'scopo', 'azioni principali')
r = riga_con(t, 'Dizionari ANSC')
D.riscrivi_cella(r.cells[1], 'Catalogo delle decodifiche replicate in locale con versione, '
                             'numero di valori, esito e data dell’ultimo aggiornamento; '
                             'apertura di una decodifica sui suoi valori con la validità '
                             'temporale; ricerca di un valore in tutte le decodifiche; storico '
                             'dei caricamenti.')
D.riscrivi_cella(r.cells[2], 'Aggiorna (completo / selettivo / simulazione), consulta i valori '
                             'di una decodifica, cerca un valore per codice o descrizione '
                             'senza sapere in quale decodifica stia, apri lo storico dei '
                             'caricamenti.')
fatti.append('schermata «Dizionari ANSC»: ricerca e apertura dei valori')


# ═════════════════════════════════════ 5. Appendice B
t = [x for x in doc.tables
     if [c.text.strip() for c in x.rows[0].cells][:2] == ['Metodo e percorso', 'Identificativo']
     and any('dizionari' in r.cells[0].text for r in x.rows[1:])][0]
nuova = D.clona_riga(t, ('GET /ansc/v1/dizionari/valori', 'cercaValori',
                         'Cerca un valore in tutte le decodifiche replicate',
                         'AMMINISTRATORE, AUDITOR, SUPPORTO'))
riga_con(t, 'GET /ansc/v1/dizionari/{codDizionario}')._tr.addnext(nuova._tr)
D.sostituisci(doc, 'GET /ansc/v1/dizionari/{codDizionario}/valori',
              'GET /ansc/v1/dizionari/{codDizionario}/valori', attese=1)
r = riga_con(t, 'GET /ansc/v1/dizionari/{codDizionario}/valori')
D.riscrivi_cella(r.cells[2], 'Elenca i valori di una decodifica, con ricerca testuale')

t2 = D.trova_tabella(doc, 'operazione del componente', 'servizio ansc invocato')
r = [x for x in t2.rows if x.cells[0].text.strip().startswith('allineaDizionari')][0]
D.riscrivi_cella(r.cells[0], 'allineaDizionari, elencaDizionari, elencaValoriDizionario, '
                             'cercaValori')
D.riscrivi_cella(r.cells[1], 'R901 (elenco e dettaglio delle decodifiche), invocato per il '
                             'tramite del concentratore operativo, che costruisce e firma il '
                             'token. ⚠️ Le sole operazioni di ricerca non contattano ANSC: '
                             'leggono la replica locale.')
fatti.append('Appendice B: inventario dei dizionari a sette operazioni')


# ═════════════════════════════════════ 6. open point
t = D.trova_tabella(doc, '#', 'tema', 'questione')
D.clona_riga(t, ('OP-62', 'Allineamento dei contratti alla nomenclatura per caso d’uso',
                 'I contratti del concentratore espongono la configurazione come «operazioni» '
                 'e ne descrivono la chiave come terna evento-operazione-casistica: è la '
                 'nomenclatura anteriore alla revisione che ha reso l’UC il soggetto della '
                 'configurazione. Mancano inoltre le operazioni per la configurazione delle '
                 'sezioni, degli allegati, delle formule e del catalogo delle logiche, che il '
                 'back-office prevede. I contratti sono la fonte autoritativa: l’allineamento '
                 'si fa su di essi e l’appendice si rigenera, non si riscrive a mano.',
                 'Aperto', 'Analisi', 'Alta'))
fatti.append('OP-62 sui contratti da allineare')


# ═════════════════════════════════════ 7. figura
D.sostituisci_immagine(doc, 'Schema ANSC_USR.', os.path.join(IMG, 'erd_ansc_usr.png'))
D.testo_di(par('Schema ANSC_USR. Le frecce piene'),
           'Schema ANSC_USR. Le frecce piene sono le chiavi esterne dichiarate nel DDL '
           'dell’Appendice A; le tratteggiate sono riferimenti per valore, senza vincolo di '
           'integrità: il codice dell’UC su ANSC_CFG_UC e su ALLEGATI_USECASE, la decodifica '
           'su ANSC_CFG_CAMPO e la baseline sullo store di stato.')
fatti.append('ERD corretto: ALLEGATI_USECASE richiama l’UC per valore')

doc.save(DST)

# ────────────────────────────────────────────── controlli
import zipfile   # noqa: E402
z = zipfile.ZipFile(DST)
ncom = z.read('word/comments.xml').decode().count('<w:comment ')
assert ncom == 27, f'commenti persi: {ncom}'
d2 = docx.Document(DST)
for p in d2.paragraphs:
    if p.style.name.startswith('Heading') and len(p.text) > 95:
        raise SystemExit('heading anomalo: ' + p.text[:70])
print('\n'.join(' · ' + f for f in fatti))
print('commenti:', ncom, '· capitoli/tabelle/immagini:', D.riepilogo(DST))
print('scritto:', os.path.relpath(DST, BASE))
