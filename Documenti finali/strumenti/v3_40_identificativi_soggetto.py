# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.39 → v3.40 (11/10/2026).

Nuova tabella ANSC_SOGGETTO_IDENTIFICATIVI: la corrispondenza fra il soggetto di SIPO e gli
identificativi che la consultazione R005 restituisce (richiesta dell'utente, sezione «Il
processo di gestione dei dati ANSC in SIPO», passaggio «Riepilogo dei soggetti»).

⚠️ Decisioni dell'utente (11/10/2026), per la regola dell'ATOMICITÀ DEL DATO:
  * i due codici ANPR NON si duplicano: esistono già su ANAG_USR.SOGGETTO —
    ID_SCHEDA_ANPR (Soggetto.java:122) e ID_UNICO_NAZIONALE (Soggetto.java:482), scritti dai
    canali ANPR (SoggettoBL.java:430 e 2131, all-anpr-sipo AllineamentoNascitaSipo.java:133).
    La tabella nuova tiene solo l'identificativo ANSC del soggetto e la traccia della ricerca;
    il risultato di R005 si scrive quindi su DUE tabelle;
  * su ANAG_USR.SOGGETTO R005 scrive SOLO SE VUOTO; un valore diverso non si sovrascrive, si
    registra come discordanza e si segnala.
⚠️ Nomi ANSC che si prestano all'equivoco (model_evento.yaml ModelSoggetto, R005
CriteriRicercaSoggetto): `idANPR` = identificativo UNICO nazionale (→ ID_UNICO_NAZIONALE);
`idsoggettoanpr` / `idAnpr` = identificativo della scheda (→ ID_SCHEDA_ANPR); `id` /
`idAnscSoggetto` = soggetto in ANSC (→ tabella nuova).
Legame stato civile → anagrafe: CODICE_INDIVIDUALE (SoggettoMatrRepository.java:78).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/diagrammi_erd.py
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_40_identificativi_soggetto.py
"""
import os
import shutil
import subprocess
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.39.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.40.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


CODICI = [
    ('Codice', 'Nome in ANSC', 'Dove si registra'),
    ('Identificativo della scheda ANPR', 'idsoggettoanpr nel modello evento; idAnpr fra i '
                                         'criteri di ricerca di R005',
     'ANAG_USR.SOGGETTO.ID_SCHEDA_ANPR — colonna esistente'),
    ('Identificativo unico nazionale', 'idANPR nel modello evento; idAnprUnico fra i criteri '
                                       'di ricerca di R005',
     'ANAG_USR.SOGGETTO.ID_UNICO_NAZIONALE — colonna esistente'),
    ('Identificativo del soggetto in ANSC', 'id nel modello evento; idAnscSoggetto fra i '
                                            'criteri di ricerca di R005',
     'ANSC_SOGGETTO_IDENTIFICATIVI.ID_ANSC_SOGGETTO — tabella nuova'),
]

COLONNE = [
    ('Colonna', 'Tipo', 'Note'),
    ('ID_SOGGETTO_IDENTIFICATIVI', 'NUMBER (PK)', 'Chiave tecnica.'),
    ('ID_SOGGETTO_SIPO', 'NUMBER (UK)',
     'Il soggetto di stato civile, MATR_USR.SOGGETTO.ID_SOGGETTO. Riferimento per valore: '
     'nessuna chiave esterna attraversa gli schemi. Una riga per soggetto.'),
    ('ID_ANSC_SOGGETTO', 'VARCHAR2(50)',
     'L’identificativo del soggetto in ANSC. Il contratto lo dichiara stringa senza lunghezza; '
     'cinquanta caratteri, come per gli altri identificativi ANSC del modello. Vuoto se la '
     'ricerca non ha trovato il soggetto.'),
    ('ID_ATTO_SIPO', 'NUMBER', 'L’atto durante la cui lavorazione si è fatta la consultazione.'),
    ('COD_ESITO_RICERCA', 'VARCHAR2(20)',
     'TROVATO, un solo risultato; SCELTO, scelto dall’operatore fra più risultati; '
     'NON_TROVATO.'),
    ('COD_ESITO_ANPR', 'VARCHAR2(20)',
     'Che cosa è accaduto ai due codici ANPR su ANAG_USR.SOGGETTO: SCRITTO, perché le colonne '
     'erano vuote; COINCIDENTE; DISCORDANTE, non sovrascritto e segnalato; NON_NUMERICO, '
     'l’identificativo della scheda non entra nella colonna numerica; ASSENTE_ANAGRAFE, il '
     'soggetto non ha una riga anagrafica (OP-74).'),
    ('DATA_RICERCA / UTENTE_RICERCA', 'TIMESTAMP / VARCHAR2(40)',
     'Quando e da chi è stata fatta l’ultima consultazione.'),
    ('ID_AUDIT', 'NUMBER',
     'La chiamata R005 in ANSC_LOG_AUDIT, dove resta la risposta integrale con i valori '
     'restituiti. Per valore e non per chiave esterna: l’audit si sfoltisce per '
     'minimizzazione, la riga del soggetto no.'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici dello standard [R5].'),
]

SCHEDA = [
    ('h3', 'ANSC_SOGGETTO_IDENTIFICATIVI — i codici del soggetto ottenuti da R005.'),
    ('p', 'La consultazione R005 restituisce, per ciascun soggetto, tre identificativi che il '
          'modello evento chiede per gli intestatari (RF-16) e che senza la consultazione non '
          'si conoscono. ANSC li nomina in un modo che si presta all’equivoco, e conviene '
          'fissarli una volta.'),
    ('tab', (CODICI, [1.7, 2.4, 2.5])),
    ('p', '⚠️ In ANSC il nome idANPR designa l’identificativo unico nazionale, non quello della '
          'scheda. È l’errore da non commettere scrivendo la transcodifica.'),
    ('p', 'Per la regola dell’atomicità del dato i due codici ANPR non si duplicano: SIPO li '
          'conserva già nell’anagrafe, dove li scrivono i canali di allineamento con ANPR. Il '
          'risultato di una consultazione si registra quindi su due tabelle, nella stessa '
          'transazione: i due codici ANPR sul soggetto anagrafico, l’identificativo ANSC in '
          'questa tabella. Il soggetto di stato civile raggiunge quello anagrafico per '
          'CODICE_INDIVIDUALE, il legame che SIPO già usa fra i due schemi.'),
    ('p', '⚠️ Due fonti per lo stesso dato. Sulle colonne anagrafiche scrivono già i canali ANPR, '
          'e la regola è che la consultazione non li contraddica in silenzio: scrive un codice '
          'solo se la colonna è vuota; se la trova valorizzata con lo stesso valore non fa '
          'nulla; se la trova valorizzata con un valore diverso non sovrascrive, registra la '
          'discordanza e la segnala al back-office. I valori restituiti da R005 restano nella '
          'risposta conservata in ANSC_LOG_AUDIT, senza essere copiati qui.'),
    ('mono', [
        'UPDATE ANAG_USR.SOGGETTO',
        '   SET ID_SCHEDA_ANPR = :id_scheda_anpr',
        ' WHERE CODICE_INDIVIDUALE = :codice_individuale',
        '   AND ID_SCHEDA_ANPR IS NULL;',
        '-- 0 righe: la colonna era già valorizzata, si confronta e si registra',
        '-- COINCIDENTE o DISCORDANTE. Lo stesso per ID_UNICO_NAZIONALE.',
    ]),
    ('p', '⚠️ È la prima volta che il componente scrive su ANAG_USR. Non ne modifica la '
          'struttura (RNF-7), ma richiede il permesso di aggiornare due colonne del soggetto '
          'anagrafico. ID_SCHEDA_ANPR è inoltre una colonna numerica, mentre il contratto '
          'dichiara l’identificativo come stringa: un valore non numerico non si scrive e si '
          'registra come tale.'),
    ('tab', (COLONNE, [1.9, 1.3, 3.4])),
    ('p', 'Una consultazione successiva aggiorna la riga. ⚠️ La riga è per soggetto di SIPO: se '
          'una riga di MATR_USR.SOGGETTO rappresenta la persona dentro un singolo atto e non la '
          'persona in generale, l’identificativo ANSC non si riusa da un atto all’altro e ogni '
          'atto richiede la propria consultazione. Il disegno regge comunque, ma il beneficio '
          'del riuso si perde. Da accertare (OP-74).'),
]

DDL = [
    'CREATE TABLE ANSC_USR.ANSC_SOGGETTO_IDENTIFICATIVI (',
    '  ID_SOGGETTO_IDENTIFICATIVI NUMBER GENERATED ALWAYS AS IDENTITY,',
    '  ID_SOGGETTO_SIPO           NUMBER             NOT NULL,',
    '  ID_ANSC_SOGGETTO           VARCHAR2(50 CHAR),',
    '  ID_ATTO_SIPO               NUMBER,',
    '  COD_ESITO_RICERCA          VARCHAR2(20 CHAR)  NOT NULL,',
    '  COD_ESITO_ANPR             VARCHAR2(20 CHAR),',
    '  DATA_RICERCA               TIMESTAMP          NOT NULL,',
    '  UTENTE_RICERCA             VARCHAR2(40 CHAR)  NOT NULL,',
    '  ID_AUDIT                   NUMBER,',
    '  DATA_INS                   TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,',
    '  DATA_UPD                   TIMESTAMP,',
    '  UTENTE_INS                 VARCHAR2(40 CHAR)  NOT NULL,',
    '  UTENTE_UPD                 VARCHAR2(40 CHAR),',
    '  CONSTRAINT PK_ansc_sogg_identif PRIMARY KEY (ID_SOGGETTO_IDENTIFICATIVI),',
    '  CONSTRAINT UK_ansc_sogg_identif UNIQUE (ID_SOGGETTO_SIPO),',
    '  CONSTRAINT CK_ansc_sogg_identif_ric',
    "    CHECK (COD_ESITO_RICERCA IN ('TROVATO','SCELTO','NON_TROVATO')),",
    '  CONSTRAINT CK_ansc_sogg_identif_id',
    "    CHECK (COD_ESITO_RICERCA = 'NON_TROVATO' OR ID_ANSC_SOGGETTO IS NOT NULL),",
    '  CONSTRAINT CK_ansc_sogg_identif_anpr',
    "    CHECK (COD_ESITO_ANPR IN ('SCRITTO','COINCIDENTE','DISCORDANTE',",
    "                              'NON_NUMERICO','ASSENTE_ANAGRAFE'))",
    ') TABLESPACE ANSC_USR;',
    'CREATE INDEX ANSC_USR.IX_ansc_sogg_identif_ansc',
    '  ON ANSC_USR.ANSC_SOGGETTO_IDENTIFICATIVI (ID_ANSC_SOGGETTO) TABLESPACE ANSC_USR;',
    'COMMENT ON TABLE ANSC_USR.ANSC_SOGGETTO_IDENTIFICATIVI IS',
    "  'Identificativo ANSC del soggetto di SIPO, da R005. I codici ANPR stanno su ANAG_USR.SOGGETTO.';",
    '',
    '-- Da eseguire dal proprietario di ANAG_USR (OP-74): lettura del soggetto e',
    '-- aggiornamento delle sole due colonne ANPR.',
    'GRANT SELECT ON ANAG_USR.SOGGETTO TO <UTENZA_CONCENTRATORE>;',
    'GRANT UPDATE (ID_SCHEDA_ANPR, ID_UNICO_NAZIONALE) ON ANAG_USR.SOGGETTO',
    '  TO <UTENZA_CONCENTRATORE>;',
    '',
]

OGGETTO = ('ANSC_SOGGETTO_IDENTIFICATIVI', 'Identificativo ANSC dei soggetti di SIPO.',
           'Una riga per soggetto, con l’identificativo ottenuto da R005 e la traccia della '
           'ricerca. I due codici ANPR della stessa risposta non stanno qui: si scrivono su '
           'ANAG_USR.SOGGETTO, dove esistono già, e solo se vuoti.')

OP74 = ('OP-74', 'Identificativi del soggetto',
        'Tre accertamenti. (a) Se una riga di MATR_USR.SOGGETTO rappresenti la persona o la '
        'persona dentro un singolo atto: nel secondo caso l’identificativo ANSC non si riusa da '
        'un atto all’altro. (b) Dove si registrino i codici ANPR dei soggetti che non hanno una '
        'riga in ANAG_USR.SOGGETTO, per esempio i non residenti: oggi la consultazione li '
        'ottiene e non li conserva. (c) Il permesso di aggiornare ID_SCHEDA_ANPR e '
        'ID_UNICO_NAZIONALE di ANAG_USR.SOGGETTO, l’utenza a cui concederlo — nel DDL il nome è '
        'un segnaposto — e chi tratta le discordanze segnalate.',
        'Aperto', 'Analisi / Settore tecnico / DBA', 'Alta')

STORIA = ('11/10/2026', '3.40',
          'Adeguamenti database (nuova ANSC_SOGGETTO_IDENTIFICATIVI) · Il processo di gestione '
          'dei dati ANSC in SIPO · Modello dati (ERD) · Appendice A · Open Point',
          'Nuova tabella ANSC_SOGGETTO_IDENTIFICATIVI, la corrispondenza fra il soggetto di SIPO '
          'e gli identificativi restituiti da R005. Per la regola dell’atomicità del dato i due '
          'codici ANPR non si duplicano: ID_SCHEDA_ANPR e ID_UNICO_NAZIONALE esistono già su '
          'ANAG_USR.SOGGETTO e vi si scrivono solo se vuoti, mentre una differenza si registra e '
          'si segnala. La tabella tiene l’identificativo ANSC e la traccia della ricerca. '
          'Fissati i nomi ANSC dei tre codici, fra cui idANPR, che designa l’identificativo '
          'unico nazionale. Prima scrittura del componente su ANAG_USR, senza modifiche di '
          'struttura. ERD rigenerato; nuovo OP-74.')


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    modello = D.trova_tabella(d, 'Oggetto', 'Ruolo')

    # ─────────────────────────── 1. adeguamenti database: la scrittura su ANAG_USR
    D.sostituisci(d, 'richiede permessi di scrittura su uno schema in esercizio (OP-68).',
                  'richiede permessi di scrittura su uno schema in esercizio (OP-68). ⚠️ E il '
                  'componente scrive anche su ANAG_USR, senza modificarne la struttura: i due '
                  'codici ANPR del soggetto ottenuti da R005 si registrano nelle colonne che il '
                  'soggetto anagrafico già possiede, e solo se vuote (si veda '
                  'ANSC_SOGGETTO_IDENTIFICATIVI).', attese=1, fatti=fatti,
                  etichetta='tutto additivo: la scrittura su ANAG_USR')
    t = D.trova_tabella(d, 'Oggetto', 'Ruolo')
    D.clona_riga(t, OGGETTO)
    D.sostituisci(d, 'Alle tre tabelle native si aggiunge ALLEGATO',
                  'Alle quattro tabelle native si aggiunge ALLEGATO', attese=1, fatti=fatti)
    D.sostituisci_immagine(d, 'Schema ANSC_USR, con ALLEGATO', os.path.join(IMG, 'erd_ansc_usr.png'))
    fatti.append('elenco degli oggetti ed ERD')

    # ─────────────────────────── 2. la scheda, in coda alle tabelle operative
    ancora = D.h(d, 2, 'Struttura dettagliata delle tabelle di configurazione')._p
    for tipo, c in SCHEDA:
        if tipo == 'p':
            D.para(d, ancora, c)
        elif tipo == 'h3':
            D.para(d, ancora, c, stile='Heading 3')
        elif tipo == 'tab':
            D.tabella(d, ancora, c[0], modello, c[1])
            D.para(d, ancora, '')
        elif tipo == 'mono':
            D.ddl(d, ancora, c)
    fatti.append('scheda di ANSC_SOGGETTO_IDENTIFICATIVI')

    # ─────────────────────────── 3. il passaggio «Riepilogo dei soggetti»
    celle = [r.cells[1] for t in d.tables for r in t.rows
             if len(r.cells) > 1 and r.cells[0].text.strip() == 'Riepilogo dei soggetti']
    assert len(celle) == 1 and 'IdSoggettoAnsc' in celle[0].text
    celle[0].add_paragraph(
        'I tre codici si registrano come descritto in «ANSC_SOGGETTO_IDENTIFICATIVI»: i due '
        'codici ANPR sul soggetto anagrafico, dove esistono già, e l’identificativo ANSC nella '
        'tabella nuova.')
    fatti.append('rimando nel passaggio «Riepilogo dei soggetti»')

    # ─────────────────────────── 4. Appendice A, prima di ANSC_NOTIFICA
    notif = [p for p in d.paragraphs
             if p.style.name == 'Heading 3' and p.text.strip() == 'ANSC_NOTIFICA']
    a = notif[-1]._p
    D.para(d, a, 'ANSC_SOGGETTO_IDENTIFICATIVI', stile='Heading 3')
    D.ddl(d, a, DDL)
    fatti.append('Appendice A: tabella, vincoli, indice, grant su ANAG_USR')

    # ─────────────────────────── 5. Open Point e storia
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == '#'
             and any(r.cells[0].text.strip() == 'OP-73' for r in t.rows))
    assert t.rows[-1].cells[0].text.strip() == 'OP-73'
    D.clona_riga(t, OP74)
    D.storia(d, *STORIA)
    fatti.append('OP-74, storia')

    d.save(DST)
    dopo = commenti(DST)
    if dopo != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo} dopo')
    print('\n'.join('· ' + f for f in fatti))
    print(f'commenti preservati: {dopo}')


if __name__ == '__main__':
    main()
