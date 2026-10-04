# -*- coding: utf-8 -*-
"""Genera «ANALISI_Identita-Profilazione-IAM» — il documento unico su identità,
profilazione e integrazione con l'IAM di Roma Capitale.

Accorpa tre basi di evidenza:
  · ASIS_Autenticazione-Profilazione_SIPO_v0.2.docx (ricognizione sui sorgenti, 25/09)
  · Sorgenti Documentali/Analisi_AS-IS_Profilazione.txt (30/09, campione Decessi)
  · Sorgenti Documentali/Analisi_Integrazione_IAM_Profilazione.txt (30/09, TO-BE OIDC)

⚠️ Sostituisce l'ASIS v0.2, che resta agli atti come versione storica e non va più
aggiornato: due documenti sullo stesso tema divergono sempre.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/diagrammi_iam.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/genera_identita_iam.py
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = os.path.join(BASE, 'Template documentale', 'template_DAD_roma-capitale.docx')
OUT = os.path.join(BASE, 'Documenti finali', 'ANALISI_Identita-Profilazione-IAM_v0.1.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')
TAGLIA_DA = 29


def testo(par, s):
    if par.runs:
        par.runs[0].text = s
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(s)


def costruisci():
    shutil.copyfile(TEMPLATE, OUT)
    d = docx.Document(OUT)
    body = d.element.body
    for ch in list(body.iterchildren())[TAGLIA_DA:]:
        if ch.tag != qn('w:sectPr'):
            body.remove(ch)
    for p in d.paragraphs:
        if '“Lorem Ipsum”' in p.text:
            testo(p, '“SIPO — Sistema Informativo dei Servizi alla Persona”')
        elif p.style.name == 'Subtitle':
            testo(p, 'Identità, profilazione e integrazione con IAM')
    return d


def h(d, liv, t):
    p = d.add_paragraph(style=f'Heading {liv}')
    p.add_run(t)
    return p


def par(d, t, stile='Normal'):
    p = d.add_paragraph(style=stile)
    for pezzo, gr in D.segmenta('', t):
        r = p.add_run(pezzo)
        r.bold = gr
    return p


def voce(d, testa, corpo):
    p = d.add_paragraph(style='List Paragraph')
    r = p.add_run(testa + ' ')
    r.bold = True
    for pezzo, gr in D.segmenta('', corpo):
        rr = p.add_run(pezzo)
        rr.bold = gr
    return p


def tabella(d, righe, larghezze=None, corpo=8.5):
    t = d.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(righe):
        for j, v in enumerate(r):
            cel = t.cell(i, j)
            cel.text = ''
            p = cel.paragraphs[0]
            for pezzo, gr in D.segmenta('', str(v)):
                run = p.add_run(pezzo)
                run.bold = gr or (i == 0)
                run.font.size = Pt(corpo)
    if larghezze:
        for j, w in enumerate(larghezze):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    return t


def figura(d, png, didascalia, larghezza=6.4):
    d.add_paragraph().add_run().add_picture(os.path.join(IMG, png), width=Inches(larghezza))
    cap = d.add_paragraph()
    r = cap.add_run(didascalia)
    r.italic = True
    r.font.size = Pt(9)


# ═══════════════════════════════════════════════════════════════ dati

RIFERIMENTI = [
    ['#', 'Riferimento', 'Natura'],
    ['[F1]', 'ASIS_Autenticazione-Profilazione_SIPO_v0.2.docx — ricognizione sui sorgenti '
             'dell’intero perimetro SIPO (25/09/2026)', 'Documento di progetto · sostituito'],
    ['[F2]', 'Analisi_AS-IS_Profilazione.txt — ricognizione sul campione DecessiWeb / '
             'DecessiBE con il dettaglio operativo (30/09/2026)', 'Analisi del fornitore'],
    ['[F3]', 'Analisi_Integrazione_IAM_Profilazione.txt — analisi di integrazione OIDC, '
             'decisioni, piano e rischi (30/09/2026)', 'Analisi del fornitore'],
    ['[F4]', 'RMCAP-IAM-MAN «Integrazione applicativa con il servizio di Identity and '
             'Access Management», Ed. 16 del 23/06/2025 (RTI Fastweb-Leonardo)',
     'Specifica normativa dell’ente'],
    ['[F5]', 'Prologic — «Comunicazione tra Identity Broker e servizi SSO»',
     'Documentazione di fornitore'],
    ['[F6]', 'Sorgenti applicativi di SIPO: common/profilazione-utente, common/rest-security, '
             'common/rest-client, common/anpr-client, back-end/anagrafe-be, '
             'sipo-root/sql/SQL/collaudo', 'Sorgente applicativo'],
    ['[F7]', 'fsha_mf-shared-library-main — libreria Angular condivisa di Roma Capitale',
     'Sorgente applicativo'],
    ['[F8]', 'ANALISI_Front-End-Angular_v0.4.docx — impostazione dei nuovi front-end',
     'Documento di progetto'],
]

CATENA = [
    ['CONF_RUOLI.RUOLO', 'Autorità nel front-end', 'PROFILO_BE', 'Utenza tecnica',
     'Autorità nel back-end'],
    ['ADMIN (1)', 'ROLE_ADMIN', 'ADMIN', 'admin.user', 'ROLE_ADMIN'],
    ['FUNZ_ANA (2)', 'ROLE_FUNZ_ANA', 'AUTH_USER', 'auth.user', 'ROLE_AUTH_USER'],
    ['DIR (3)', 'ROLE_DIR', 'AUTH_USER', 'auth.user', 'ROLE_AUTH_USER'],
    ['POLIZIA_LOC (8)', 'ROLE_POLIZIA_LOC', 'AUTH_USER', 'auth.user', 'ROLE_AUTH_USER'],
    ['CITTADINO (4)', 'ROLE_CITTADINO', 'CITTADINO', 'cittadino', 'ROLE_CITTADINO'],
    ['CITTADINO_PROF (5)', 'ROLE_CITTADINO_PROF', 'CITTADINO_PROF', 'cittadino.prof',
     'ROLE_CITTADINO_PROF'],
    ['UFF_CERT_CORR (6)', 'ROLE_UFF_CERT_CORR', 'UFF_CERT_CORR', 'uff.cert.corr',
     'ROLE_UFF_CERT_CORR'],
    ['UFF_CERT_ENTI_EST (7)', 'ROLE_UFF_CERT_ENTI_EST', 'UFF_CERT_ENTI_EST',
     'uff.cert.enti.est', 'ROLE_UFF_CERT_ENTI_EST'],
    ['(nessuno, ripiego)', 'ROLE_CITTADINO_CRI_ON', 'CITTADINO_CRI_ON (solo in Java)',
     'da clientConfig', 'ROLE_CITTADINO_CRI_ON'],
]

RILIEVI = [
    ['#', 'Rilievo', 'Gravità', 'Evidenza', 'Origine'],
    ['RI-01', 'I servizi di back-end rispondono senza gettone',
     'Critica',
     'ResourceServerConfig non dichiara anyRequest(). Ogni percorso non elencato e privo di '
     '@PreAuthorize risponde anche a chiamanti anonimi, comprese operazioni di scrittura: '
     'salvataggio e cancellazione di atti, finalizzazione, stampa di estratti con dati '
     'personali.', 'C1 · R-03'],
    ['RI-02', 'L’identità dell’operatore non è legata al gettone',
     'Critica',
     'Identificativo utente, organizzazione, codice fiscale e ruolo arrivano nel corpo della '
     'richiesta e il back-end li accetta senza verifica. Con un gettone qualsiasi — o senza, '
     'vedi RI-01 — si opera «come» un altro ufficiale o un’altra organizzazione. Il ruolo '
     'ricevuto decide perfino con quale utenza si chiama il canale ANPR.',
     'C4 · R-05 · R-07'],
    ['RI-03', 'Segreti in chiaro e una sola chiave di firma',
     'Critica',
     'Password di base dati di produzione nei file di proprietà; security.signing-key '
     'identica in 57 file; segreti dei client con prefisso {noop}; chiave TripleDes nel '
     'codice. Chi conosce la chiave di firma può costruire un gettone con qualunque autorità, '
     'valido su tutti i back-end.', 'C2 · R-04 · R-13'],
    ['RI-04', 'L’identità in ingresso è un header di cui ci si fida',
     'Critica',
     'CdRLoginFilter accoglie iv-user e sysgroup senza alcuna verifica di provenienza o di '
     'firma. Se l’host del front-end è raggiungibile senza passare dal portale, un header '
     'basta a impersonare chiunque.', 'C3 · R-01'],
    ['RI-05', 'Il nome della vista è scelto dal client',
     'Critica',
     'In GestioneDecessiController.riprendiLavorazione il parametro idSelezionato è spezzato '
     'e la seconda parte diventa il valore restituito dal metodo, cioè il nome della vista '
     'risolta. Verificato alla riga 696.', 'C5'],
    ['RI-06', 'Esiste un percorso alternativo di identificazione',
     'Alta',
     'Con il parametro di richiesta mode=local l’identità è letta dagli attributi di sessione '
     'anziché dagli header. Non è condizionato al profilo né all’ambiente.', 'R-02'],
    ['RI-07', 'Il front-end non autorizza per funzione',
     'Alta',
     'Le rotte dei moduli non anagrafici non compaiono negli elenchi di SecurityConfig, che '
     'non ha una clausola di chiusura. Qualunque utente con sessione valida, anche un '
     'cittadino, raggiunge le funzioni degli altri moduli. Il menu è solo visivo.',
     'A1 · R-03'],
    ['RI-08', 'L’autorizzazione nel back-end è solo per profilo tecnico',
     'Alta',
     'Le espressioni di autorizzazione ragionano su ADMIN, AUTH_USER e simili. Quasi tutte le '
     'operazioni di stato civile ammettono anche CITTADINO e CITTADINO_PROF. Nessun controllo '
     'per area, funzionalità, municipio o appartenenza del dato.', 'A2'],
    ['RI-09', 'Verso ANPR operatore e postazione non sono attendibili',
     'Alta',
     'La postazione è scritta nel codice (ConfigHandler.java:67, che sovrascrive il parametro '
     'ricevuto); l’operatore dei processi automatici è fisso e corrisponde a una persona '
     'reale; in finalizzazione il codice fiscale inviato è quello di chi ha creato l’atto, '
     'non di chi lo sta formando. Lo stato è statico e non protetto fra richieste.',
     'A3 · R-16'],
    ['RI-10', 'Sei moduli espongono le proprie rotte senza autenticazione',
     'Alta',
     'Le classi BasicConfiguration di allineamento nascite, matrimoni e decessi, del modulo '
     'fax e posta certificata, della carta d’identità e di una copia fuori albero dichiarano '
     'permitAll e disabilitano la protezione contro la falsificazione di richiesta; il '
     'gestore delle credenziali è vuoto.', 'R-09'],
    ['RI-11', 'Il filtro a chiave API non protegge ciò che emette gettoni',
     'Alta',
     'ApiKeyFilter confronta il percorso con «/generate» e «/validate», ma il controller è '
     'mappato su «/api/token/…» e il modulo non dichiara un context-path: la condizione non '
     'si verifica mai. Nessun modulo consuma quegli endpoint.', 'R-17'],
    ['RI-12', 'Dati personali in file di configurazione versionati',
     'Alta',
     'I file di proprietà di produzione e preproduzione del client ANPR contengono il codice '
     'fiscale di operatori reali nel campo dell’operatore, accanto alle password dei '
     'keystore.', 'R-19'],
    ['RI-13', 'Impostazioni predefinite permissive',
     'Alta',
     'In assenza di ruoli il servizio assegna d’ufficio ROLE_CITTADINO_CRI_ON, e va in errore '
     'se l’utente non esiste. Tutte le utenze tecniche di collaudo hanno la stessa impronta '
     'della parola d’ordine. Un solo identificativo di risorsa per tutti i back-end.',
     'A4 · R-13'],
    ['RI-14', 'Trasporto non cifrato fra i servizi',
     'Alta',
     'Le chiamate da front-end a back-end e fra back-end viaggiano in http, con gettoni al '
     'portatore validi dodici ore.', 'A5'],
    ['RI-15', 'Il permesso è attribuito per utente, non per ruolo',
     'Media',
     'La lettura che ricava le funzionalità non passa dal ruolo: unisce utente, area tematica '
     'e funzionalità. Non esiste alcun luogo in cui sia scritto che cosa un ruolo comporti, e '
     'lo stesso ruolo può avere funzionalità diverse per utenti diversi.', 'R-15 · [F2] 11.2'],
    ['RI-16', 'Solo il primo ruolo determina i privilegi',
     'Media',
     'Il profilo passato al back-end, la descrizione mostrata in testata e il filtro delle '
     'azioni nelle pagine usano tutti il primo ruolo, cioè quello con identificativo minore. '
     'Per un utente con più ruoli il comportamento dipende dall’ordinamento della lettura.',
     'M1 · R-06'],
    ['RI-17', 'Il catalogo delle funzionalità ha un solo effetto autorizzativo',
     'Media',
     'La verifica avviene all’ingresso, con un confronto per contenimento sulla sola '
     'destinazione iniziale, e una destinazione vuota la supera. Le navigazioni successive '
     'non la ripercorrono.', 'R-08'],
    ['RI-18', 'Un gettone nuovo a ogni chiamata, e quarantaquattro emittenti',
     'Media',
     'Il client richiede un gettone all’inizio di ogni metodo, senza alcuna conservazione: due '
     'viaggi per ogni chiamata. Ogni back-end è insieme server di autorizzazione e di risorsa: '
     'nessuna gestione centrale, e la revoca è impossibile perché i gettoni non hanno stato.',
     'M2 · M3'],
    ['RI-19', 'La sessione non è governata',
     'Media',
     'La durata è configurata da un solo modulo; nessun attributo di sicurezza sul cookie; '
     'nessun archivio condiviso fra istanze; nessun limite di sessioni concorrenti; '
     'l’interceptor per la sessione scaduta non è registrato.', 'M6 · R-14'],
    ['RI-20', 'La configurazione di profilazione non è versionata',
     'Media',
     'Negli script del repository non compare alcun inserimento di aree tematiche, '
     'funzionalità o ruoli per l’ambito Stato Civile: la configurazione reale esiste solo '
     'sulle basi dati degli ambienti.', 'M7'],
    ['RI-21', 'La configurazione del canale sicuro verso ANPR è orfana',
     'Media',
     'Il file cxf.xml non è caricato da alcun componente e nomina dieci volte una classe '
     'inesistente; saml.properties è caricato solo da lì. Non è dimostrabile dal codice che '
     'l’asserzione SAML venga mai costruita.', 'R-18'],
    ['RI-22', 'Il modello a permessi per aspetti è stato abbandonato a metà',
     'Media',
     'Le quattro classi CheckPermission sono commentate riga per riga e importano una classe '
     'assente dal workspace. Era il tentativo di sostituire l’autorizzazione per indirizzo con '
     'una per permessi: la direzione non è inesplorata, è interrotta.', 'R-20'],
    ['RI-23', 'La firma dell’identificativo di postazione esiste in tre copie',
     'Media',
     'La stessa logica è scritta in anpr-client, nel programma a riga di comando e nel modulo '
     'anagrafico. La cifratura simmetrica esiste in sedici copie con tre chiavi distinte.',
     'R-21'],
    ['RI-24', 'Il tracciamento è puntuale e la sua chiave arriva dal client',
     'Media',
     'Nessun meccanismo automatico di registrazione; poche colonne «chi» con sei nomi diversi, '
     'valorizzate dal parametro di richiesta. Il registro dei rilasci di estratti ha una lista '
     'di esclusione configurabile.', 'R-10 · R-11'],
    ['RI-25', 'Gli errori di autorizzazione non arrivano all’operatore',
     'Media',
     'Le risposte 401 e 403 dei back-end sono assorbite dal front-end: l’operatore vede un '
     'elenco vuoto, non un messaggio di accesso negato.', 'M5'],
    ['RI-26', 'La configurazione per ambiente è legata al ramo compilato',
     'Media',
     'Blocchi commentati nella persistenza, bersaglio nel file del client, ambiente ANPR '
     'scritto nel codice: quale ambiente si raggiunge dipende dal ramo da cui si è costruita '
     'l’immagine.', 'M4'],
    ['RI-27', 'Codice di sicurezza presente e non attivo',
     'Bassa',
     'Il vincolo di giunzione per la Polizia Locale non è mai invocato; l’interceptor di '
     'sessione invalida non è registrato; il gestore che disattiva la validazione dei '
     'certificati resta nel percorso di classe di trenta moduli pur essendo irraggiungibile.',
     'B1 · R-12'],
    ['RI-28', 'L’impianto esiste in due copie',
     'Bassa',
     'Profilazione, sicurezza REST, client ed entità sono duplicati sotto la radice delle '
     'elezioni, con lo stesso identificativo di artefatto: la copia può sovrascrivere '
     'l’originale nel deposito locale. Un modulo di produzione dipende dal duplicato.',
     'B1 · R-15'],
    ['RI-29', 'Pila tecnologica fuori supporto',
     'Bassa',
     'Spring Boot 2.1, il modulo OAuth2 di Spring Security è fuori supporto, il tipo di '
     'concessione in uso è deprecato, log4j2 2.15.0 porta una vulnerabilità nota.', 'B2'],
]

DECISIONI = [
    ['#', 'Decisione', 'Alternative', 'Raccomandazione'],
    ['D1', 'Chi è il componente che dialoga con IAM',
     '(A) la libreria di profilazione dentro ciascun front-end · (B) l’Identity Broker del '
     'fornitore · (C) un componente di autenticazione centrale di SIPO',
     '**(A) per la prima fase**, dietro un’interfaccia che renda la scelta reversibile. È la '
     'via più breve e non introduce componenti nuovi. ⚠️ (B) va considerata solo se il '
     'supporto all’IAM dell’ente è confermato e se il ripiego «claim senza verifica della '
     'firma» si può disattivare: altrimenti si introdurrebbe sul percorso critico del login un '
     'componente che in caso di errore accetta asserzioni non verificate. (C) resta '
     'l’approdo naturale della seconda fase.'],
    ['D2', 'Chi restituisce i profili a partire dal codice fiscale',
     '(a) un servizio dell’IAM · (b) un servizio di SIPO che espone ANAG_USR',
     '**Modello ibrido.** L’IAM decide l’accesso di massima — se la persona è abilitata al '
     'back-office o al front-office — e SIPO resta la fonte dei profili di dettaglio, perché '
     'ruoli, aree, funzionalità e organizzazione hanno una granularità che un IAM non '
     'gestisce. ⚠️ Vale comunque la pena esporli con un servizio dedicato: oggi quelle letture '
     'girano dentro ciascuno dei trentacinque front-end con le credenziali di base dati '
     'dentro il pacchetto.'],
    ['D3', 'Quale chiave lega l’utente IAM all’utente SIPO',
     'codice fiscale · nome utente',
     '**Codice fiscale**, con ripiego transitorio sul nome utente, registrato e segnalato, '
     'finché la bonifica non è completa. Il codice fiscale è stabile e comune a tutti i '
     'canali. ⚠️ Oggi la colonna è facoltativa, non univoca e non indicizzata: la bonifica è '
     'un prerequisito, non un’attività parallela.'],
    ['D4', 'Come l’identità dell’operatore raggiunge i back-end',
     'nessun cambiamento · un gettone SIPO per l’utente · inoltrare il gettone IAM',
     '**Nessun cambiamento nella prima fase, gettone SIPO per l’utente nella seconda.** '
     'Inoltrare il gettone dell’IAM è sconsigliato: il destinatario dichiarato è il client '
     'IAM e i back-end di SIPO non lo sono. ⚠️ È la seconda fase, non la prima, a chiudere '
     'RI-02 e a dare ad ANPR un operatore reale.'],
    ['D5', 'Come si realizza la parte OIDC sulla pila attuale',
     'callback esplicita con le librerie già presenti · aggiornare prima la piattaforma',
     '**Callback esplicita.** Legare l’integrazione all’aggiornamento della piattaforma '
     'significherebbe subordinarla a un intervento che tocca trentacinque front-end e '
     'quarantaquattro back-end. ⚠️ L’aggiornamento resta necessario (RI-29), ma per altre '
     'ragioni e con altri tempi.'],
    ['D6', 'Come convivono le due modalità durante la migrazione',
     'passaggio secco · proprietà per ambiente',
     '**Proprietà per ambiente**, per poter tornare indietro senza un rilascio. ⚠️ In modalità '
     'OIDC gli header in ingresso vanno ignorati **e rimossi dalla richiesta**: lasciarli '
     'leggibili riaprirebbe il vettore che l’integrazione serve a chiudere (RI-04).'],
]

RACCORDO = [
    ['Rilievo', 'L’integrazione lo chiude?', 'Perché'],
    ['RI-04 Fiducia negli header', '**Sì, prima fase**',
     'L’identità diventa un gettone firmato che si verifica, non un header che si accoglie.'],
    ['RI-06 Percorso mode=local', '**Sì, prima fase**',
     'Il parametro sparisce con la dismissione della modalità header.'],
    ['RI-19 Sessione non governata', '**Sì, prima fase**',
     'Rigenerazione dell’identificativo dopo il login, attributi del cookie e durata esplicita '
     'sono requisiti dell’integrazione.'],
    ['RI-17 Convalida della destinazione', '**Sì, prima fase**',
     'La destinazione si conserva lato server e si convalida per corrispondenza esatta.'],
    ['RI-02 Identità non propagata', 'Solo seconda fase',
     'Richiede un gettone che rappresenti la persona e non l’utenza tecnica.'],
    ['RI-16 Solo il primo ruolo', 'Solo seconda fase',
     'Il gettone può portare tutti i ruoli; oggi la scelta è forzata dal meccanismo.'],
    ['RI-18 Quarantaquattro emittenti', 'Solo seconda fase',
     'Un emittente unico rende possibili la revoca e una gestione centrale.'],
    ['RI-09 Tracciabilità verso ANPR', 'Solo seconda fase',
     'Operatore e postazione reali arrivano solo se l’identità arriva.'],
    ['RI-13 Impostazioni permissive', 'Solo seconda fase',
     'Le utenze tecniche condivise si dismettono quando non servono più.'],
    ['RI-01 Servizi senza gettone', '**No**',
     '⚠️ Non dipende dall’autenticazione. È una clausola di chiusura nella configurazione: '
     'intervento piccolo, indipendente, da anticipare.'],
    ['RI-03 Segreti in chiaro', '**No**',
     'La chiave di firma unica sparisce con la seconda fase; tutto il resto no.'],
    ['RI-05 Vista scelta dal client', '**No**', 'È un difetto applicativo, non di accesso.'],
    ['RI-14 Trasporto non cifrato', '**No**', 'È una scelta di pubblicazione dei servizi.'],
    ['RI-07 Autorizzazione per funzione', '**No**',
     'L’integrazione porta l’identità, non il modello dei permessi.'],
    ['RI-12 Dati personali versionati', '**No**', 'È igiene del repository.'],
]

FASI = [
    ['Fase', 'Contenuto', 'Giorni/uomo', 'Che cosa porta'],
    ['0', 'Decisioni D1-D6 con i presidi, prova tecnica del flusso su un front-end, scheda '
          'del servizio e richieste di apertura', '6-9',
     'La verifica che la pila attuale regga, prima di impegnarsi.'],
    ['1', 'Accesso OIDC per back-office e front-office: adattatore, punto di ingresso, '
          'callback, convalida, servizio dei profili, bonifica della base dati, rilascio sui '
          'trentacinque front-end', '45-60',
     'Chiude RI-04, RI-06, RI-17 e RI-19. È il percorso minimo per avere l’accesso su IAM.'],
    ['2', 'Identità fino ai back-end sul campione: emittente unico, gettone dell’utente, '
          'identità letta dal gettone, operatore e postazione reali verso ANPR', '30-40',
     'Chiude RI-02, RI-09, RI-16, RI-18 e parte di RI-13. ⚠️ L’estensione agli altri '
     'back-end è un programma a sé.'],
    ['3', 'Front office: seconda autorizzazione per persona giuridica, delega e assistito, '
          'con un contesto operativo esplicito', '15-20',
     'Rende dichiarato per conto di chi si opera. Non riguarda il back-office.'],
    ['4', 'Dismissione della modalità header e delle utenze tecniche di ruolo', '3-5',
     'Chiude il doppio binario e il rischio di regressione.'],
]

APERTI = [
    ['#', 'Questione', 'A chi compete'],
    ['PI-01', 'Algoritmo di firma dei gettoni dell’IAM e disponibilità del materiale di '
              'verifica. ⚠️ L’esempio della specifica mostra una firma simmetrica, che mal si '
              'concilia con un verificatore per chiave pubblica.', 'Presidio IAM'],
    ['PI-02', 'Se la dimostrazione di possesso del codice sia obbligatoria per i client '
              'riservati.', 'Presidio IAM'],
    ['PI-03', 'Quali ambiti di consenso sono abilitabili per un client di back-office e per '
              'uno di front-office.', 'Presidio IAM'],
    ['PI-04', 'Se si possano registrare trentacinque indirizzi di ritorno su un solo client, '
              'e durata dei gettoni.', 'Presidio IAM'],
    ['PI-05', 'Se il servizio che restituisce i profili sia dell’IAM, e con quale contratto.',
     'Presidio IAM'],
    ['PI-06', 'Se il codice fiscale sia sempre valorizzato per le utenze dei dipendenti.',
     'Presidio IAM'],
    ['PI-07', 'Se con la pubblicazione OIDC le giunzioni restino dietro il proxy e gli header '
              'di infrastruttura — indirizzo e numero di serie del certificato di postazione — '
              'continuino a essere iniettati. ⚠️ Da questa risposta dipende la tracciabilità '
              'verso ANPR.', 'Presidio Portale'],
    ['PI-08', 'Se il proxy rimuova gli header di identità eventualmente inviati dal client '
              'prima di iniettare i propri.', 'Presidio Portale'],
    ['PI-09', 'Se l’Identity Broker supporti l’IAM dell’ente, se possa portare claim '
              'applicativi e se il ripiego senza verifica della firma sia disattivabile.',
     'Fornitore del Broker'],
    ['PI-10', 'Se il professionista resti una profilazione di SIPO o diventi una '
              'rappresentanza gestita dall’IAM.', 'Referenti SIPO'],
    ['PI-11', 'Chi esegue e approva la bonifica dei codici fiscali in ANAG_USR.',
     'Referenti SIPO'],
    ['PI-12', 'Raggiungibilità diretta degli host dei front-end e dei back-end senza passare '
              'dal proxy: decide la gravità reale di RI-01 e RI-04.', 'Sistemi / Rete'],
    ['PI-13', 'Esistenza dei sinonimi o delle concessioni che rendono leggibili le utenze '
              'tecniche dallo schema dei moduli di stato civile.', 'Base dati'],
    ['PI-14', 'Contenuto reale in produzione della configurazione di profilazione per '
              'l’ambito Stato Civile, assente dagli script.', 'Base dati / Referenti'],
    ['PI-15', 'Sorgente del package che legge il registro delle postazioni e DDL delle '
              'tabelle che lo compongono, assenti dal repository.', 'Base dati / Sistemi'],
    ['PI-16', 'Se l’indicazione dell’ambiente e la versione di build possano entrare nella '
              'cornice istituzionale dei nuovi front-end.', 'Design system / Cliente'],
]


def scrivi(d):
    # ───────────────────────────────────────────────────────── 1. scopo
    h(d, 1, 'Scopo del documento')
    par(d, 'Questo documento descrive **come SIPO riconosce le persone e decide che cosa '
           'possano fare**, e **come questo cambia** con l’adozione del servizio di gestione '
           'delle identità di Roma Capitale. Tiene insieme le due cose di proposito: un '
           'accesso robusto davanti a servizi che rispondono senza gettone non migliora la '
           'sicurezza complessiva, e separare lo stato di fatto dal piano di intervento '
           'avrebbe nascosto proprio questo.')
    par(d, 'È diviso in due parti. La **prima** è una ricognizione dello stato di fatto, '
           'condotta sui sorgenti: non propone nulla e registra ciò che trova, compreso ciò '
           'che non c’è. La **seconda** descrive l’integrazione con l’IAM: il flusso target, '
           'le decisioni da prendere con la raccomandazione motivata di chi scrive, gli '
           'interventi per componente e il piano per fasi.')
    par(d, '⚠️ **Sostituisce [F1]**, che resta agli atti come versione storica e non va più '
           'aggiornato. Due documenti sullo stesso tema divergono sempre, e qui il tema è uno '
           'solo.')

    h(d, 2, 'Perimetro e metodo')
    par(d, 'Il documento accorpa tre ricognizioni indipendenti. La prima [F1] ha battuto in '
           'ampiezza l’intero perimetro di SIPO. La seconda [F2] ha campionato in profondità '
           'una coppia di moduli — il front-end e il back-end dei decessi — ricostruendo la '
           'catena fino ad ANPR. La terza [F3] ha studiato l’integrazione con l’IAM a partire '
           'dalla specifica dell’ente [F4].')
    par(d, '**Le tre ricognizioni concordano su tutti i punti d’impianto**, pur essendo state '
           'condotte separatamente e con perimetri diversi. Dove divergevano, la divergenza è '
           'stata verificata sui sorgenti e risolta; i tre casi sono riportati qui sotto '
           'perché la convergenza non sia data per scontata.')
    tabella(d, [
        ['Punto', 'Prima lettura', 'Seconda lettura', 'Verifica'],
        ['Valori dell’ambito', 'quattro', 'tre',
         '**Tre negli inserimenti** del repository; il commento nel codice ne dichiara un '
         'quarto, «statistica», che nei dati non esiste. Entrambe le letture erano fedeli '
         'alla propria fonte.'],
        ['File con la chiave di firma', '45', '44',
         '**57** sull’intero workspace: i due conteggi campionavano perimetri diversi.'],
        ['Protezione contro la falsificazione di richiesta', 'attiva per difetto',
         'non trattata',
         'Verificata: la catena del front-end non la disattiva mai, quindi vale il '
         'comportamento predefinito. ⚠️ Si segnala perché il silenzio si presterebbe '
         'all’inferenza opposta.'],
    ], larghezze=[1.3, 1.1, 1.1, 3.0])
    par(d, 'Vale per tutto il documento il limite dichiarato in [F1]: **leggendo il codice si '
           'accerta che cosa il sistema fa, non che cosa il livello di rete davanti a esso '
           'impedisce.** Alcuni dei rilievi sono attenuati, in esercizio, da una '
           'configurazione di rete che non è fra le sorgenti: il documento lo dice dove '
           'accade e lo registra come punto aperto (PI-12).')

    h(d, 2, 'Glossario')
    tabella(d, [
        ['Termine', 'Significato in questo documento'],
        ['Autenticazione', 'Il procedimento con cui si stabilisce chi è l’utente.'],
        ['Autorizzazione', 'La decisione se un utente già riconosciuto possa compiere una '
                           'certa azione.'],
        ['Profilazione', 'I dati che descrivono l’utente ai fini di quella decisione: ruoli, '
                         'funzionalità abilitate, ambito organizzativo.'],
        ['Modalità header', 'La modalità di integrazione con l’IAM in cui l’identità viaggia '
                            'in intestazioni HTTP aggiunte da un proxy. È quella che SIPO usa '
                            'oggi, ed è prevista dalla specifica dell’ente.'],
        ['Modalità OIDC', 'La modalità in cui l’identità è provata da un gettone firmato, '
                          'ottenuto con un reindirizzamento e uno scambio fuori banda.'],
        ['Utenza tecnica', 'Un’identità applicativa usata per chiamare i servizi di '
                           'back-end, distinta dalle persone.'],
        ['Profilo di back-end', 'Il valore che lega il ruolo della persona all’utenza tecnica '
                                'con cui si chiamano i servizi.'],
        ['Contesto operativo', 'Nella seconda fase del front office: per conto di chi una '
                               'persona sta operando, distinto da chi si è autenticato.'],
    ], larghezze=[1.5, 5.0])

    h(d, 2, 'Riferimenti')
    tabella(d, RIFERIMENTI, larghezze=[0.55, 4.25, 1.7])

    # ───────────────────────────────────────────────────────── 2. sintesi
    h(d, 1, 'Il quadro in una pagina')
    par(d, 'SIPO ha **due livelli di identità, distinti e scollegati**. Il primo è la persona, '
           'e vive soltanto nei front-end: arriva dal portale sotto forma di intestazioni '
           'HTTP, viene arricchita con i profili letti dalla base dati anagrafica e deposta in '
           'sessione. Il secondo è l’identità tecnica verso i back-end: a ogni chiamata il '
           'front-end ottiene un gettone presentando **un’utenza applicativa scelta in base al '
           'ruolo**, non la persona.')
    par(d, '⚠️ **Fra i due livelli l’identità dell’operatore si perde.** Nel gettone che '
           'raggiunge i servizi il soggetto è «amministratore» o «utente autorizzato», mai il '
           'codice fiscale di chi sta lavorando. Chi è davvero l’operatore viaggia nel corpo '
           'della richiesta, come un dato qualunque, e nessuno lo verifica. È la radice di '
           'buona parte dei rilievi della prima parte.')
    par(d, 'Sull’integrazione con l’IAM il punto di partenza è migliore di quanto sembri: '
           '**SIPO non è fuori dall’IAM, ne usa la modalità più debole**. Le intestazioni '
           '«iv-» non sono un arrangiamento proprietario ma la modalità header prevista dal '
           'capitolo 3.2 della specifica dell’ente [F4]. Passare al flusso con gettone firmato '
           'significa **cambiare soltanto lo strato di ingresso dell’identità**: ruoli, aree e '
           'funzionalità restano dove sono, e l’oggetto che i trentacinque front-end leggono '
           'dalla sessione non cambia. Si configurano, non si riscrivono.')
    par(d, '⚠️ Ciò che l’integrazione **non** porta va detto subito, perché è la metà meno '
           'comoda: i servizi che rispondono senza gettone, i segreti in chiaro, il trasporto '
           'non cifrato e l’assenza di un’autorizzazione per funzione non dipendono da come '
           'si entra. Il capitolo «Che cosa l’integrazione chiude» lo mette in tabella.')

    # ══════════════════════════ PARTE I ══════════════════════════
    h(d, 1, 'Parte I — Lo stato di fatto')
    par(d, 'La ricognizione che segue descrive il sistema come è oggi. Registra anche le '
           'assenze, perché un’assenza accertata evita che la stessa ricerca sia rifatta e '
           'impedisce di assumere come presente ciò che non c’è.')

    h(d, 2, 'L’autenticazione: SIPO non autentica')
    par(d, 'Non esiste un punto di accesso applicativo: nessun controller che riceva '
           'credenziali, nessuna pagina che le chieda, nessun servizio che le verifichi. '
           'L’identità arriva dal portale come intestazioni HTTP — identificativo utente, '
           'gruppo, codice fiscale, dati anagrafici — ed è accolta da un filtro registrato sul '
           'solo percorso di ingresso. Spring Security è presente, ma il fornitore di '
           'autenticazione è quello della pre-autenticazione e **le credenziali sono la '
           'stringa vuota**: serve solo ad autorizzare.')
    figura(d, 'aut_catena.png',
           'Figura 1 — La catena dell’accesso. Il pallino rosso segna i punti in cui un '
           'elemento è accolto senza essere verificato.')
    par(d, 'Il corollario pratico è che **in SIPO non esiste un luogo in cui intervenire '
           'sull’accesso**: qualunque cambiamento riguardi il modo in cui gli operatori si '
           'identificano si decide altrove e arriva come un cambiamento del formato delle '
           'intestazioni. È esattamente ciò che la seconda parte di questo documento '
           'descrive.')
    par(d, 'Verso i back-end avviene invece un’autenticazione vera: un flusso a concessione '
           'per password, con un gettone richiesto **a ogni chiamata** e mai conservato. '
           'L’identità che vi viaggia è però quella dell’utenza tecnica. La catena che porta '
           'dal ruolo della persona a quell’utenza è questa.')
    tabella(d, CATENA, larghezze=[1.5, 1.5, 1.3, 1.2, 1.5])
    par(d, '⚠️ **Si noti che cosa si perde nel passaggio.** Le differenze fra funzionario '
           'anagrafico, dirigente e polizia locale esistono nel front-end e scompaiono nel '
           'back-end: tutti e tre diventano «utente autorizzato». Aree tematiche e '
           'funzionalità non attraversano affatto il confine.')

    h(d, 2, 'La profilazione')
    par(d, 'Sei tabelle dello schema anagrafico descrivono chi può fare che cosa. Il perno è '
           'la tabella di assegnazione, che lega in una sola riga l’utente, il ruolo, l’area '
           'tematica e la funzionalità.')
    figura(d, 'prof_modello.png',
           'Figura 2 — Le tabelle della profilazione. Il permesso non è associato al ruolo: '
           'passa sempre per l’utente.')
    par(d, '⚠️ **Il ruolo non è collegato alle funzionalità.** La lettura che ricava le '
           'funzionalità di un utente non passa dal ruolo: unisce utente, area e funzionalità. '
           'Ne discende la proprietà che determina il costo di gestione dell’intero impianto: '
           '**non esiste un luogo in cui sia scritto che cosa un ruolo comporti**, e lo stesso '
           'ruolo può avere funzionalità diverse per utenti diversi. Per sapere che cosa '
           'significhi un ruolo bisogna esaminare le righe di chi lo possiede; per modificarlo, '
           'intervenire su tutte.')
    par(d, 'Il catalogo delle funzionalità ha un solo effetto autorizzativo, all’ingresso, con '
           'un confronto per contenimento sulla sola destinazione iniziale; per il resto '
           'costruisce il menu. Nel back-end le espressioni di autorizzazione ragionano '
           'esclusivamente su liste di profili tecnici scritte nel codice: fra tutte, **una '
           'sola** autorizza su una funzionalità.')
    par(d, 'L’ambito organizzativo esiste su tre dimensioni — struttura interna, sede di '
           'municipio, struttura convenzionata — ma ⚠️ **non è applicato dal server**: arriva '
           'come parametro dal chiamante, e nulla lo confronta con l’utente autenticato. Né '
           'potrebbe: il gettone porta l’utenza tecnica, che di organizzazione non ne ha.')

    h(d, 2, 'L’identità verso l’esterno')
    par(d, 'C’è una seconda sicurezza in SIPO, che convive con la prima senza toccarla ed è '
           'molto più matura: quella delle **macchine**. Il client ANPR porta cinquantuno file '
           'dedicati — caricamento di keystore, firma crittografica, asserzioni, firma dei '
           'messaggi — ed è dipendenza di trenta moduli di back-end. Esiste un **registro '
           'delle postazioni** con i certificati, e un’operazione di arruolamento che li '
           'riceve, ne estrae chiave e certificato, firma l’identificativo di postazione e lo '
           'registra.')
    par(d, '⚠️ **E tuttavia, verso ANPR, né l’operatore né la postazione sono attendibili.** '
           'L’identificativo di postazione passato alla procedura che seleziona il certificato '
           'è scritto nel codice e sovrascrive quello ricevuto: ogni chiamata proviene dalla '
           'stessa postazione, chiunque stia operando. L’operatore dei processi automatici è '
           'fisso. In finalizzazione il codice fiscale inviato è quello di chi ha **creato** '
           'l’atto, non di chi lo sta formando.')
    par(d, 'È lo stesso problema dell’identità che non arriva ai back-end, visto dal lato '
           'macchina — e ha la stessa soluzione: finché l’identità della persona non attraversa '
           'la catena, non può arrivare nemmeno all’ente terzo. ⚠️ Si aggiunga che la '
           'configurazione che collega il canale sicuro al componente che costruisce '
           'l’asserzione sta in **due file orfani**, uno dei quali nomina dieci volte una '
           'classe che non esiste: dal solo codice non è dimostrabile che quella parte venga '
           'eseguita (RI-21, PI-15).')

    h(d, 2, 'Il tracciamento')
    par(d, '**Non esiste alcuna forma di registrazione automatica.** Le annotazioni che il '
           'framework mette a disposizione per attribuire a ogni riga il suo autore non '
           'compaiono in nessun punto del perimetro. Le colonne «chi» sono una quindicina, con '
           'sei nomi diversi, e ⚠️ **il valore arriva dal client**: una traccia così formata '
           'registra ciò che il chiamante ha dichiarato di essere, non ciò che il sistema ha '
           'accertato.')
    par(d, 'L’unico registro sistematico degli accessi è una riga scritta a ogni connessione '
           'in un file di log. Nello stato civile non esiste alcuna tabella di audit: l’unica '
           'presente nel perimetro appartiene al dominio elettorale.')

    h(d, 2, 'I rilievi')
    par(d, 'Si raccolgono qui, in un registro unico, i rilievi delle tre ricognizioni. La '
           'colonna «origine» conserva la corrispondenza con i registri dei documenti di '
           'partenza, così che il raccordo resti verificabile. La gravità è quella attribuita '
           'in [F2], estesa ai rilievi che provengono da [F1].')
    par(d, '⚠️ **La valutazione del rischio reale richiede di sapere se gli host dei servizi '
           'siano raggiungibili senza passare dal proxy** (PI-12). Da quella risposta dipende '
           'se RI-01 e RI-04 siano difetti teorici o praticabili.')
    tabella(d, RILIEVI, larghezze=[0.5, 1.4, 0.6, 3.1, 0.9])

    # ══════════════════════════ PARTE II ══════════════════════════
    h(d, 1, 'Parte II — L’integrazione con IAM')

    h(d, 2, 'Il punto di partenza')
    par(d, 'La specifica dell’ente [F4] prevede due modalità di integrazione. Nella prima '
           'l’identità viaggia in intestazioni HTTP aggiunte da un proxy; nella seconda è '
           'provata da un gettone firmato ottenuto con un reindirizzamento. **SIPO usa oggi la '
           'prima**, e la usa come la specifica la descrive: le intestazioni che legge sono '
           'quelle della tabella del capitolo 3.2.')
    par(d, 'Due scostamenti vanno però annotati, perché sono anche due opportunità. Il primo: '
           'la specifica prevede che i gruppi dell’utente arrivino in un’intestazione '
           'dedicata, e **SIPO non la usa** — i ruoli vengono dalla base dati anagrafica. Il '
           'secondo: il tipo di utente arriva in un’intestazione non prevista dalla specifica, '
           'una variabile aggiunta dalla politica del proxy.')
    par(d, '⚠️ Ne discende l’osservazione che riorienta tutto il capitolo: **l’autenticazione '
           'è già delegata, l’autorizzazione è già in SIPO**. Passare al gettone firmato cambia '
           'il protocollo e il modello di fiducia, non la ripartizione delle responsabilità. '
           'Il passo «profili per codice fiscale» del flusso target corrisponde a ciò che oggi '
           'fanno il controller di autenticazione e il suo accesso ai dati, con la chiave '
           'cambiata.')
    par(d, 'Una conferma collaterale viene dai nuovi front-end. La libreria Angular condivisa '
           '[F7] rinvia, in assenza di gettone, a un indirizzo che contiene «loginIAM» [F8]: '
           'le applicazioni nuove **puntano già dove SIPO dovrà arrivare**. Il che rende la '
           'migrazione anche una convergenza, non solo un adeguamento.')

    h(d, 2, 'Il flusso target')
    figura(d, 'iam_flusso.png',
           'Figura 3 — Il flusso di autenticazione target. In verde ciò che va costruito.')
    par(d, 'Otto passi, dei quali **cinque sono nuovi e tre restano come sono**. Si costruisce '
           'il punto di ingresso che genera i parametri di sicurezza e reindirizza; la '
           'callback che riceve il codice e ne verifica la corrispondenza; lo scambio fuori '
           'banda con la convalida completa del gettone; la lettura dei profili per codice '
           'fiscale; la mappatura sui campi dell’oggetto di sessione.')
    par(d, '**Ciò che non cambia è il contratto verso le applicazioni.** L’oggetto depositato '
           'in sessione resta quello di oggi, con gli stessi campi: i trentacinque front-end '
           'vanno configurati e ricollaudati, non riscritti. È la ragione per cui questo '
           'intervento è proporzionato — e anche il motivo per cui conviene farlo dentro la '
           'libreria condivisa e non in ciascuna applicazione.')
    par(d, '⚠️ Due intestazioni non sono identità e vanno garantite comunque: l’indirizzo di '
           'provenienza e il **numero di serie del certificato di postazione**, che alimenta '
           'il canale ANPR. Se la pubblicazione cambia e quelle intestazioni spariscono, si '
           'perde l’unico legame fra la sessione e la postazione fisica (PI-07).')

    h(d, 2, 'Le decisioni da prendere')
    par(d, 'Sei decisioni precedono la scrittura di codice. Per ciascuna si riportano le '
           'alternative e la raccomandazione di chi scrive, con la ragione: la scelta resta '
           'del committente.')
    tabella(d, DECISIONI, larghezze=[0.4, 1.3, 1.8, 3.0])
    par(d, '⚠️ Una raccomandazione trasversale, che vale qualunque sia l’esito di D1: **il '
           'codice vada scritto dietro un’interfaccia** che separi «costruisci il '
           'reindirizzamento» e «gestisci il ritorno» dal resto. Così la scelta fra le tre '
           'alternative non costringe a riscrivere la libreria, e un eventuale cambio di '
           'fornitore non è un rifacimento.')

    h(d, 2, 'Che cosa l’integrazione chiude, e che cosa no')
    par(d, 'È la tabella che risponde alla domanda più utile: che cosa si compra con questo '
           'intervento.')
    figura(d, 'iam_raccordo.png',
           'Figura 4 — Il raccordo fra i rilievi dello stato di fatto e ciò che '
           'l’integrazione risolve.')
    tabella(d, RACCORDO, larghezze=[1.6, 1.3, 3.6])
    par(d, '⚠️ **La colonna dei «no» non si risolve da sé e non va rinviata.** Il caso più '
           'netto è RI-01: i servizi che rispondono senza gettone. Correggerlo è un intervento '
           'piccolo — una clausola di chiusura nella configurazione di sicurezza — '
           'indipendente dall’integrazione e realizzabile subito. Finché resta aperto, un '
           'accesso robusto davanti a servizi aperti non migliora la sicurezza complessiva: '
           'la sposta soltanto di un passo.')

    h(d, 2, 'Gli interventi per componente')
    tabella(d, [
        ['Componente', 'Intervento', 'Note'],
        ['Libreria di profilazione', 'Punto di ingresso, callback, convalida del gettone, '
                                     'mappatura dei claim, proprietà per la coesistenza, '
                                     'logout, governo della sessione, clausola di chiusura',
         '**È il punto di integrazione principale.** Il contratto verso le applicazioni resta '
         'invariato.'],
        ['Servizio dei profili (nuovo)', 'Espone per codice fiscale utente, ruoli, aree, '
                                         'funzionalità, organizzazione e dati di postazione',
         '⚠️ Sostituisce le letture che oggi girano dentro ciascun front-end **con le '
         'credenziali di base dati nel pacchetto**. Il guadagno non è solo di pulizia.'],
        ['Base dati anagrafica', 'Bonifica dei codici fiscali, indice univoco e obbligatorietà, '
                                 'registro degli accessi, versionamento della configurazione',
         '⚠️ **Prerequisito, non attività parallela**: la chiave di ricerca diventa il codice '
         'fiscale, che oggi è facoltativo, non univoco e non indicizzato.'],
        ['I trentacinque front-end', 'Configurazione, ricollaudo dell’accesso e della '
                                     'navigazione', 'Nessuna riscrittura.'],
        ['Client e sicurezza REST', 'Seconda fase: gettone dell’utente, conservazione, solo '
                                    'server di risorsa, identità letta dal gettone',
         'Chiude RI-02 e RI-18.'],
        ['Canale ANPR', 'Seconda fase: operatore e postazione reali',
         'Dipende dalla risposta a PI-07.'],
        ['Front office', 'Terza fase: contesto operativo per persona giuridica, delega e '
                         'assistito',
         'Non riguarda il back-office. ⚠️ Oggi il professionista è ricavato dalla base dati, '
         'senza rappresentanza né deleghe.'],
        ['Consumatori esterni delle intestazioni', 'Censimento e piano dedicato',
         '⚠️ Altri prodotti leggono le stesse intestazioni: cambiare la giunzione li tocca '
         '(PI-08).'],
    ], larghezze=[1.4, 2.3, 2.8])

    h(d, 2, 'Il piano per fasi')
    tabella(d, FASI, larghezze=[0.4, 2.4, 0.8, 2.9])
    par(d, 'Le stime sono indicative e vanno riviste dopo le decisioni e dopo la prova '
           'tecnica. Non comprendono i tempi di risposta dei presidi esterni né la bonifica '
           'dei dati, che è a carico del business. ⚠️ **Il percorso minimo per avere l’accesso '
           'su IAM è fase 0 più fase 1**; sono le fasi 2 e 3 a portare i benefici di sicurezza '
           'e di tracciabilità.')
    par(d, '⚠️ Una raccomandazione sull’ordine: **la clausola di chiusura sui servizi (RI-01) '
           'va anticipata alla fase 0.** È un intervento di poche righe, non dipende da nessuna '
           'delle sei decisioni e rimuove il rilievo più grave. Rinviarlo alla fase 2 '
           'significherebbe tenere aperto per mesi ciò che si può chiudere in giorni.')

    h(d, 2, 'I rischi dell’intervento')
    tabella(d, [
        ['Rischio', 'Probabilità', 'Impatto', 'Come si contiene'],
        ['Codici fiscali mancanti o duplicati: dipendenti che non entrano più', 'Alta', 'Alto',
         'Bonifica anticipata, ripiego transitorio sul nome utente, verifica giornaliera degli '
         'esiti.'],
        ['La pila attuale non regge la parte OIDC', 'Media', 'Medio',
         'Prova tecnica in fase 0; callback esplicita invece di affidarsi al supporto nativo.'],
        ['Il Broker non è conforme all’IAM dell’ente', 'Media', 'Alto',
         'Verifica preliminare dei requisiti; in alternativa la prima opzione di D1.'],
        ['Trentacinque indirizzi di ritorno e altrettanti segreti da gestire', 'Alta', 'Medio',
         'Un client unico con più indirizzi, oppure un componente centrale.'],
        ['Perdita delle intestazioni di infrastruttura: la postazione ANPR', 'Media', 'Alto',
         '⚠️ Verifica con il presidio del portale **prima** del passaggio, non dopo (PI-07).'],
        ['Navigazione fra front-end: accessi ripetuti', 'Media', 'Medio',
         'Verifica del comportamento silenzioso in presenza di sessione attiva, in fase 0.'],
        ['La copia «elezioni» sovrascrive la libreria nel deposito locale', 'Alta', 'Medio',
         'Allineare o rinominare l’identificativo di artefatto prima del rilascio (RI-28).'],
        ['I servizi restano aperti nonostante l’accesso nuovo', 'Alta', 'Critico',
         '⚠️ Anticipare la clausola di chiusura: è il rischio che vanifica l’intero '
         'intervento.'],
    ], larghezze=[2.0, 0.8, 0.7, 3.0])

    # ─────────────────────────────────────────────── punti aperti
    h(d, 1, 'Punti aperti')
    par(d, 'Le questioni che le fonti disponibili non consentono di chiudere, con '
           'l’indicazione di chi può rispondere. Sono numerate con il prefisso PI per non '
           'confondersi con i registri degli altri documenti di progetto.')
    tabella(d, APERTI, larghezze=[0.5, 4.3, 1.7])
    par(d, '⚠️ Tre di questi non sono domande ma **prerequisiti**: senza la risposta a PI-01 '
           'non si può scrivere il verificatore del gettone; senza PI-07 non si può garantire '
           'la tracciabilità verso ANPR dopo il passaggio; senza PI-12 non si può dire quanto '
           'siano gravi i due rilievi critici della prima parte.')


if __name__ == '__main__':
    d = costruisci()
    for t in d.tables:
        if t.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in t.rows:
                v = {'Progetto': 'SIPO — Identità e profilazione',
                     'Data consegna': '02/10/2026', 'Versione': '0.1',
                     'Documento': 'ANALISI_Identita-Profilazione-IAM_v0.1'
                     }.get(r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    st = D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')
    for r in list(st.rows)[1:]:
        if not any(c.text.strip() for c in r.cells):
            r._tr.getparent().remove(r._tr)
    D.storia(d, '02/10/2026', '0.1', 'Tutti',
             'Prima stesura. Accorpa la ricognizione sui sorgenti dell’intero perimetro, '
             'l’analisi sul campione Decessi e l’analisi di integrazione con l’IAM dell’ente. '
             'Le tre ricognizioni concordano su tutti i punti d’impianto; le tre divergenze '
             'sono state verificate sui sorgenti e risolte. Registro unico di ventinove '
             'rilievi con la corrispondenza alle fonti, sei decisioni istruite con '
             'raccomandazione, raccordo fra i rilievi e ciò che l’integrazione chiude, piano '
             'per fasi e sedici punti aperti. Sostituisce ASIS_Autenticazione-Profilazione_'
             'SIPO_v0.2, che resta agli atti come versione storica.')
    scrivi(d)
    d.save(OUT)
    print('scritto:', os.path.relpath(OUT, BASE))
    print('capitoli/tabelle/immagini:', D.riepilogo(OUT))
