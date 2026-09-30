# -*- coding: utf-8 -*-
"""Motore della mappatura campo per campo UC ANSC ↔ modello evento ↔ SIPO.

Non produce documenti: costruisce le strutture dati da cui i generatori (il workbook e i
capitoli del documento) leggono. Tenerlo separato serve a poter rigenerare tutto quando ANSC
pubblica una revisione del mapping, senza riscrivere le regole.

Le tre fonti, e che cosa danno ciascuna:
  · openapi/model_evento.yaml       → la STRUTTURA: percorso, tipo, formato, decodifica citata
  · Mapping_casi_uso/<fam>/<UC>.csv → il FABBISOGNO per UC: quali campi, obbligatori, condizioni
  · Decessi_ANSC.xlsx / Nascite_ANSC_07.08.2026.xlsx → il lato SIPO, l'unica parte umana

Nota di metodo: i percorsi del mapping si normalizzano sugli indici di lista (evento.x[0].y e
evento.x.y sono lo stesso percorso), altrimenti la risoluzione sul modello fallisce su tutte le
annotazioni contestuali.
"""
import collections
import csv
import glob
import os
import re

import yaml

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ANSC = os.path.join(BASE, 'ansc', 'docs')
SORG = os.path.join(BASE, 'Sorgenti Documentali')

FAMIGLIE = {
    'morte': 'Morte', 'nascita': 'Nascita', 'riconoscimenti': 'Riconoscimenti',
    'matrimoni': 'Matrimoni', 'unioni_civili': 'Unioni civili',
    'cittadinanza': 'Cittadinanza', 'trascrizioni': 'Trascrizioni',
}


# ------------------------------------------------------------------ modello evento
def _rif(v):
    """(schema, è_lista) del riferimento di una proprietà, comunque sia espresso."""
    if '$ref' in v:
        return v['$ref'].split('/')[-1], False
    for a in v.get('allOf') or []:
        if '$ref' in a:
            return a['$ref'].split('/')[-1], False
    if v.get('type') == 'array':
        n, _ = _rif(v.get('items') or {})
        if n:
            return n, True
    return None, False


class Modello:
    """Indice piatto dell'albero del modello evento: percorso → caratteristiche del campo."""

    def __init__(self, percorso_yaml=None):
        self.sch = yaml.safe_load(
            open(percorso_yaml or os.path.join(ANSC, 'openapi', 'model_evento.yaml'))
        )['components']['schemas']
        self.campi = {}          # percorso -> dict
        self.liste = set()       # percorsi che sono liste
        self._espandi('ModelEvento', '', {'ModelEvento'})

    def _props(self, nome):
        s = self.sch.get(nome) or {}
        p = dict(s.get('properties') or {})
        for a in s.get('allOf') or []:
            if '$ref' in a:
                p.update(self._props(a['$ref'].split('/')[-1]))
            else:
                p.update(a.get('properties') or {})
        return p

    def _espandi(self, nome, prefisso, visti):
        for k, v in self._props(nome).items():
            perc = f'{prefisso}.{k}' if prefisso else k
            rif, lista = _rif(v)
            desc = v.get('description', '') or ''
            dec = re.search(r'ANSC_\d+', desc)
            self.campi[perc] = {
                'percorso': perc,
                'tipo': v.get('type') or ('oggetto' if rif else ''),
                'formato': v.get('format', ''),
                'schema': rif or '',
                'lista': lista,
                'descrizione': desc,
                'decodifica': dec.group(0) if dec else '',
                'deprecato': bool(v.get('deprecated')),
                'esempio': '' if v.get('example') is None else str(v.get('example')),
                'multilingua': bool(rif and rif.endswith('ML')),
            }
            if lista:
                self.liste.add(perc)
            if rif and rif in self.sch and rif not in visti:
                self._espandi(rif, perc + '[]' if lista else perc, visti | {rif})

    # ------------------------------------------------------------- risoluzione
    @staticmethod
    def normalizza(percorso):
        """evento.x[0].y → x[].y : toglie la radice «evento.» e appiattisce gli indici."""
        p = re.sub(r'\[\d+\]', '[]', percorso.strip())
        return p[len('evento.'):] if p.startswith('evento.') else p

    def risolvi(self, binding_object, binding_field):
        """Cerca il campo nell'albero, provando anche a inserire l'indice di lista omesso."""
        grezzo = '.'.join(x for x in (binding_object.strip(), binding_field.strip()) if x)
        chiave = self.normalizza(grezzo)
        if chiave in self.campi:
            return grezzo, chiave, self.campi[chiave]
        # il mapping talvolta omette l'indice su una proprietà che nel modello è una lista
        pezzi = chiave.split('.')
        for i in range(len(pezzi) - 1):
            prova = '.'.join(pezzi[:i + 1]) + '[]' + ''.join('.' + x for x in pezzi[i + 1:])
            if prova in self.campi:
                return grezzo, prova, self.campi[prova]
        return grezzo, chiave, None


# ------------------------------------------------------------------ conversioni
# Classi di conversione fra il dato di SIPO e il campo di ANSC. L'ordine conta: si applica la
# prima regola che riconosce il campo.
CONVERSIONI = [
    ('Identificativo ANSC',
     lambda p, c: bool(re.search(r'(^|\.)idAnsc', p)) or p.endswith('.idAnsc'),
     'Identificativo nazionale assegnato da ANSC: non esiste in SIPO. Si ottiene dalla '
     'consultazione (R005) per i soggetti e dal deposito (R009) per gli atti; per l’atto '
     'collegato si veda OP-30.'),
    ('Territoriale (titolarità ANPR)',
     lambda p, c: bool(re.search(r'id(Stato|Comune|Provincia|Nazionalita|Cittadinanza)', p)),
     'Traduzione della codifica locale in quella nazionale: gli stati passano per '
     'CONF_STATO_ESTERO.CODICE_ANPR (archivio ANPR_02), i comuni e le province per le tabelle '
     'COMUNE e PROVINCIA di ANAG_USR. ANSC non pubblica questi cataloghi (DV-32).'),
    ('Decodifica ANSC',
     lambda p, c: bool(c and c['decodifica']),
     'Il valore è un codice di una decodifica pubblicata da ANSC: si traduce dalla tabella '
     'CONF_* di SIPO al codice ANSC leggendo il dizionario replicato in locale (RF-10).'),
    ('Booleano da flag',
     lambda p, c: (c and c['tipo'] == 'boolean') or bool(re.search(r'(^|\.)flag', p)),
     'In SIPO il contrassegno è tipicamente CHAR(1) con valori S/N: va convertito in booleano, '
     'e l’assenza del dato non equivale a «falso».'),
    ('Data',
     lambda p, c: bool(c and c['formato'] in ('date', 'date-time')),
     'Da DATE di Oracle al formato ISO 8601 richiesto dal contratto; attenzione alle date '
     'parziali, che ANSC tratta con i campi di formato dedicati.'),
    ('Numerico',
     lambda p, c: bool(c and c['tipo'] in ('integer', 'number')),
     'Conversione di tipo; verificare che il dato di SIPO non sia una stringa con zeri di '
     'riempimento.'),
    ('Testo diretto',
     lambda p, c: bool(c and c['tipo'] == 'string'),
     'Trasferimento diretto del valore; restano da verificare le lunghezze massime e la '
     'normalizzazione di maiuscole e apostrofi.'),
    ('Struttura',
     lambda p, c: bool(c and c['tipo'] == 'oggetto'),
     'Non è un campo foglia ma un blocco del modello: si valorizza riempiendo i campi che '
     'contiene.'),
    ('Da determinare',
     lambda p, c: True,
     'Il percorso non trova riscontro nel modello evento: prima di configurarlo va chiarito con '
     'il fornitore di ANSC (OP-52).'),
]


def classifica(percorso, campo):
    """Il campo non risolto NON si classifica: il suo trattamento va deciso, non dedotto."""
    if campo is None:
        return 'Non risolto sul modello', CONVERSIONI[-1][2]
    for nome, prova, regola in CONVERSIONI:
        if prova(percorso, campo):
            return nome, regola
    return 'Da determinare', ''


# ------------------------------------------------------------------ tipi di allegato
def _norm(s):
    """Normalizza una descrizione per il confronto: accenti, apostrofi tipografici, spazi."""
    import unicodedata
    s = unicodedata.normalize('NFKD', s or '')
    s = ''.join(c for c in s if not unicodedata.combining(c))
    for a, b in (('\u2019', "'"), ('\u2018', "'"), ('\u201c', '"'), ('\u201d', '"'),
                 ('\u2013', '-'), ('\u2014', '-')):
        s = s.replace(a, b)
    return re.sub(r'[\s]+', ' ', s).strip().lower()


def tipi_allegato():
    """descrizione normalizzata -> (codice ANSC_09, descrizione ufficiale).

    Il mapping nomina gli allegati per DESCRIZIONE, ANSC li codifica in ANSC_09: il raccordo
    e' per testo normalizzato, ed e' proprio cio' che lascia scoperte le poche descrizioni
    senza corrispondenza (OP-50).
    """
    fuori = {}
    p = os.path.join(ANSC, 'Decodifiche', '9_dec_tipo_allegato.csv')
    if not os.path.exists(p):
        return fuori
    with open(p, newline='', encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh):
            fuori[_norm(r.get('DESCRIZIONE', ''))] = (r.get('ID', '').strip(),
                                                      (r.get('DESCRIZIONE') or '').strip())
    return fuori


# ------------------------------------------------------------------ catalogo UC
def catalogo_uc():
    """codice motore → (codice numerico, descrizione) da 3_dec_use_case.csv (ANSC_03)."""
    per_motore = {}
    with open(os.path.join(ANSC, 'Mapping_casi_uso', '3_dec_use_case.csv'),
              newline='', encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh):
            per_motore[r['COD ANSC'].strip()] = r['ID USECASE'].strip()
    desc = {}
    dec = os.path.join(ANSC, 'Decodifiche', '3_dec_use_case.csv')
    if os.path.exists(dec):
        with open(dec, newline='', encoding='utf-8-sig') as fh:
            for r in csv.DictReader(fh):
                desc[r.get('ID', '').strip()] = r.get('DESCRIZIONE', '').strip()
    return {m: (n, desc.get(n, '')) for m, n in per_motore.items()}


# ------------------------------------------------------------------ lato SIPO
INTESTAZIONI_SIPO = ('binding object', 'binding field')


def _fogli_sipo(percorso_xlsx):
    """Estrae da un workbook di ricognizione le righe (binding, campo SIPO, note)."""
    import openpyxl
    w = openpyxl.load_workbook(percorso_xlsx, read_only=True, data_only=True)
    fuori = []
    for nome in w.sheetnames:
        righe = list(w[nome].iter_rows(values_only=True))
        testa = None
        for i, r in enumerate(righe):
            valori = [str(c).strip().lower() if c is not None else '' for c in r]
            if all(x in valori for x in INTESTAZIONI_SIPO):
                testa = (i, valori)
                break
        if not testa:
            continue
        i, valori = testa
        col = {v: j for j, v in enumerate(valori) if v}
        # gli UC citati nel foglio: colonna «Codice Motore»/«Codice Motore ANSC» in testa
        motori = set()
        testata = {}
        for k, r in enumerate(righe[:i]):
            valori_r = [str(c).strip() if c is not None else '' for c in (r or ())]
            for c in valori_r:
                if re.fullmatch(r'[A-Za-z_]+_\d{3}', c):
                    motori.add(c)
            # il blocco «DESCRIZIONE | ID_CONF_TIPO_ATTO | ID_MODELLO_ATTO | SERIE | TIPO_RITO |
            # MASCHERA_UI» dichiara la maschera: e' l'aggancio al template da cui si ricava il
            # campo del DB, quindi va letto e non ignorato
            if 'MASCHERA_UI' in valori_r and k + 1 < i:
                sotto = [str(c).strip() if c is not None else '' for c in (righe[k + 1] or ())]
                testata = {a: b for a, b in zip(valori_r, sotto) if a}
        for r in righe[i + 1:]:
            def v(nome_col):
                j = col.get(nome_col)
                return '' if j is None or j >= len(r) or r[j] is None else str(r[j]).strip()
            bo, bf = v('binding object'), v('binding field')
            if not bf:
                continue
            fuori.append({
                'foglio': nome, 'motori': motori,
                'maschera': testata.get('MASCHERA_UI', '').split(':')[-1],
                'id_conf_tipo_atto': testata.get('ID_CONF_TIPO_ATTO', ''),
                'id_modello_atto': testata.get('ID_MODELLO_ATTO', ''),
                'binding': '.'.join(x for x in (bo, bf) if x),
                'area': v('area'), 'gruppo': v('gruppo'),
                'campo': v('nome campo'),
                'obbligatorio_sipo': v('obbligatorio'),
                'tipo_sipo': v('tipo campo'),
                'note': ' · '.join(x for x in (v('note'), v('note campo')) if x),
            })
    w.close()
    return fuori


def lato_sipo(modello):
    """Percorso normalizzato → riga SIPO, con l'insieme degli UC per cui vale."""
    out = collections.defaultdict(list)
    for f in ('Decessi_ANSC.xlsx', 'Nascite_ANSC_07.08.2026.xlsx'):
        p = os.path.join(SORG, f)
        if not os.path.exists(p):
            continue
        for riga in _fogli_sipo(p):
            riga['fonte'] = f
            out[modello.normalizza(riga['binding'])].append(riga)
    return out


# ------------------------------------------------------------------ mappatura per UC
def mappatura(modello=None, sipo=None):
    """Elenco di righe: una per (UC, campo dichiarato dal mapping), già arricchite."""
    modello = modello or Modello()
    sipo = sipo if sipo is not None else lato_sipo(modello)
    uc = catalogo_uc()
    allegati = tipi_allegato()
    righe = []
    for f in sorted(glob.glob(os.path.join(ANSC, 'Mapping_casi_uso', '*', '*.csv'))):
        famiglia = FAMIGLIE.get(os.path.basename(os.path.dirname(f)), '?')
        motore = os.path.basename(f)[:-4]
        numerico, descrizione = uc.get(motore, ('', ''))
        with open(f, newline='', encoding='utf-8-sig') as fh:
            for r in csv.DictReader(fh):
                sez = (r.get('Sezione') or '').strip()
                campo = (r.get('Campo') or '').strip()
                bo = (r.get('Binding Object') or '').strip()
                bf = (r.get('Binding Field') or '').strip()
                if bf == 'Binding Field':
                    continue
                grezzo, chiave, c = (modello.risolvi(bo, bf) if bf else ('', '', None))
                classe, regola = (classifica(chiave, c) if bf else ('—', ''))
                # Il lato SIPO vale per gli UC che il foglio di ricognizione dichiara: non si
                # estende per analogia, perché lo stesso percorso ha significati diversi in
                # famiglie diverse (l'intestatario è il defunto o il neonato).
                noti = sipo.get(chiave, []) if bf else []
                cand = [s for s in noti if motore in s['motori']]
                altrove = sorted({m for s in noti for m in s['motori']} - {motore})
                stato = ('Rilevato' if cand
                         else f'Ereditabile da {altrove[0]}' if altrove else 'Da rilevare')
                # il rilievo fatto per un altro UC si MOSTRA ma non si adotta: lo stesso
                # percorso ha significati diversi in famiglie diverse, la scelta e' del funzionario
                altro = next((s for s in noti if not cand), None)
                cod_all = desc_all = ''
                if sez == 'Allegati':
                    cod_all, desc_all = allegati.get(_norm(campo), ('', ''))
                righe.append({
                    'allegato_codice': cod_all, 'allegato_descrizione': desc_all,
                    'motore': motore, 'uc': numerico, 'descrizione_uc': descrizione,
                    'famiglia': famiglia, 'sezione': sez, 'campo': campo,
                    'obbligatorio': (r.get('Obbligatorio') or '').strip(),
                    'condizione': (r.get("Condizioni obbligatorieta'") or '').strip(),
                    'formule': (r.get("Note obbligatorieta' formule") or '').strip(),
                    'binding': grezzo, 'percorso': chiave,
                    'risolto': c is not None,
                    'tipo': c['tipo'] if c else '', 'formato': c['formato'] if c else '',
                    'lista': bool(c and c['lista']),
                    'decodifica': c['decodifica'] if c else '',
                    'deprecato': bool(c and c['deprecato']),
                    'descrizione_campo': c['descrizione'] if c else '',
                    'classe': classe, 'regola': regola,
                    'sipo_area': cand[0]['area'] or cand[0]['gruppo'] if cand else '',
                    'sipo_campo': cand[0]['campo'] if cand else '',
                    'sipo_tipo': cand[0]['tipo_sipo'] if cand else '',
                    'sipo_note': cand[0]['note'] if cand else '',
                    'sipo_fonte': cand[0]['fonte'] if cand else '',
                    'sipo_altrove_campo': altro['campo'] if altro else '',
                    'sipo_altrove_uc': ', '.join(altrove)[:120],
                    'sipo_stato': stato if bf else '—',
                })
    return righe


if __name__ == '__main__':
    m = Modello()
    s = lato_sipo(m)
    print('campi del modello:', len(m.campi), '· liste:', len(m.liste))
    print('percorsi con lato SIPO noto:', len(s))
    r = mappatura(m, s)
    print('righe di mappatura:', len(r), '· UC:', len({x['motore'] for x in r}))
    c = collections.Counter(x['classe'] for x in r)
    for k, v in c.most_common():
        print(f'  {k:32} {v:6}')
    print('righe con lato SIPO:', sum(1 for x in r if x['sipo_campo']))
