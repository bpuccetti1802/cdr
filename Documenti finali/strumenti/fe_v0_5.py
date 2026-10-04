# -*- coding: utf-8 -*-
"""ANALISI_Front-End-Angular v0.4 → v0.5 (02/10/2026).

Tre aggiunte chieste dal committente:
  1. la gestione dei certificati di postazione, assegnata al remote mfOperation, con le
     schermate reali di «screen-app-certificati.docx»;
  2. la testata e il piè di pagina secondo «Footer cdr.jpeg», che toglie la colonna
     del menu e porta a sette i canali;
  3. un sottocapitolo sul menu configurato con un ConfigMap: perché è stato scelto,
     perché non è una buona pratica e come si rientra dopo il primo rilascio.

⚠️ Si modifica il file reale. La v0.4 non porta commenti di Word, ma porta riscritture
manuali del committente: gli ancoraggi usano il testo come è oggi, non come l'avevo
generato.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/fe_v0_5.py
"""
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Front-End-Angular_v0.4.docx')
DST = os.path.join(BASE, 'Documenti finali', 'ANALISI_Front-End-Angular_v0.5.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')

MENU_CONFRONTO = [
    ['Aspetto', 'Menu da ConfigMap (primo rilascio)', 'Menu da base dati (bersaglio)'],
    ['Chi lo cambia', 'Chi ha accesso al cluster: sistemisti o rilascio.',
     'Il funzionario abilitato, da una pagina di amministrazione.'],
    ['Quando ha effetto', 'Al riavvio dei pod, oppure al successivo ricaricamento se il file '
                          'è montato e riletto.',
     'Al successivo caricamento del menu, senza toccare i pod.'],
    ['Tracciabilità della modifica', '⚠️ Nessuna dentro l’applicazione: resta nella storia '
                                     'del repository di configurazione, se c’è.',
     'Utente e momento su ogni riga, come per il resto della configurazione.'],
    ['Abilitazioni per voce', 'Dichiarate nel file, senza vincolo di integrità verso il '
                              'vocabolario reale.',
     'In tabella, con il vincolo verso l’elenco delle abilitazioni esistenti.'],
    ['Differenze fra ambienti', '⚠️ Un file per ambiente: il menu di collaudo e quello di '
                                'esercizio divergono senza che nulla lo segnali.',
     'Un solo modello, con i dati propri di ciascun ambiente.'],
    ['Costo di realizzazione', 'Pressoché nullo: è un file.',
     'Quattro tabelle, un servizio di lettura e una pagina di amministrazione.'],
]

CERT_FLUSSO = [
    ['Passo', 'Che cosa fa l’operatore', 'Che cosa serve costruire'],
    ['Elenco', 'Cerca per nome del certificato e scorre l’elenco, con la sede accanto a '
               'ciascuno.',
     'Nulla di nuovo: tabella, campo di ricerca e impaginazione sono della libreria.'],
    ['Caricamento', 'Preme «Carica certificato»: la riga di ricerca lascia il posto al '
                    'modulo con nome, password e selezione del file.',
     'Il componente di scelta del file e la commutazione fra i due modi della stessa '
     'pagina. ⚠️ Il file è un contenitore PKCS#12 e la password serve ad aprirlo.'],
    ['Conferma', 'Riceve una notifica in alto a destra — «Certificato caricato con '
                 'successo» — e vede la riga comparire nell’elenco.',
     'Nulla: il componente di notifica è della libreria. La pagina non cambia.'],
    ['Eliminazione', 'Apre il menu della riga, sceglie «Elimina» e conferma in una finestra '
                     'che nomina il certificato.',
     'Nulla: menu di riga e finestra di conferma sono della libreria.'],
]


def main():
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    fatti = []

    # ───────────── 1. l'assegnazione al remote, dentro il modello a micro-frontend
    D.sostituisci(
        d, 'Per il perimetro dell’integrazione si prevedono tre micro-frontend',
        'Per il perimetro dell’integrazione si prevedono tre micro-frontend',
        attese=1, etichetta='ancoraggio al modello a micro-frontend', fatti=fatti)
    for p in d.paragraphs:
        if p.text.strip().startswith('Per il perimetro dell’integrazione si prevedono tre'):
            D.para(d, p._p.getnext() if p._p.getnext() is not None else p._p,
                   '**La gestione dei certificati di postazione appartiene al remote delle '
                   'operazioni, mfOperation.** È la scelta coerente con la natura di quelle '
                   'schermate: non configurano l’integrazione e non formano atti, '
                   'amministrano gli strumenti con cui una postazione si presenta agli enti '
                   'esterni. ⚠️ Vi si accede però con la stessa abilitazione amministrativa '
                   'delle altre funzioni di servizio, non con quella di chi forma gli atti: '
                   'appartenenza al remote e visibilità nel menu sono due cose distinte.')
            fatti.append('assegnazione dei certificati a mfOperation')
            break

    # ───────────── 2. il sottocapitolo sul menu da ConfigMap
    ancora = D.h(d, 2, 'La shell e la condivisione della libreria')

    def par(t, stile='Normal'):
        return D.para(d, ancora._p, t, stile=stile)

    par('La composizione del menu: una deroga per il primo rilascio', 'Heading 3')
    par('La home delle applicazioni è una pagina a mattonelle: ogni mattonella è una rotta, '
        'con il proprio titolo, sottotitolo, icona, ordine e abilitazione richiesta. Il '
        'disegno del back-office [F8] prevede che quelle voci siano **un dato**, non codice: '
        'quattro tabelle dichiarano quali applicazioni esistono, quali pagine contengono, '
        'come si presentano e chi le può usare, e la home si costruisce leggendole. Il '
        'pregio di quella scelta è che una riorganizzazione del menu non è un rilascio, e '
        'che la mappa dell’applicazione si ispeziona da una tabella invece che leggendo i '
        'sorgenti.')
    par('⚠️ **Per il primo rilascio si è deciso diversamente: il menu sarà configurato con '
        'un ConfigMap sul server.** La ragione è di tempi: il registro su base dati richiede '
        'le tabelle, il servizio che le legge e la pagina che le amministra, mentre un file '
        'di configurazione esiste il giorno in cui lo si scrive. In una prima messa in '
        'esercizio il menu cambia poco e cambia sotto il controllo di chi rilascia: il '
        'guadagno di flessibilità del registro non si esercita, e il suo costo si paga per '
        'intero.')
    par('La scelta è legittima e si dichiara per quello che è: **una soluzione tampone**, '
        'non il disegno. Vale la pena però essere espliciti su che cosa si rinuncia, perché '
        'è il modo di non dimenticarsene.')
    D.tabella(d, ancora._p, MENU_CONFRONTO, modello=d.tables[4], larghezze=[1.3, 2.6, 2.6])
    par('⚠️ **Le due righe che pesano di più sono la tracciabilità e la divergenza fra '
        'ambienti.** La prima perché una modifica al menu è una modifica a che cosa gli '
        'operatori vedono, e nel modello a file non resta traccia dentro l’applicazione di '
        'chi l’ha fatta e quando. La seconda perché un file per ambiente, nel tempo, diventa '
        'ambienti che non si somigliano più — ed è il difetto che si scopre quando una '
        'funzione collaudata non compare in esercizio.')
    par('**Come si rientra.** Il passaggio non è un rifacimento, a una condizione: che il '
        'front-end legga il menu **da un servizio**, e non dal file. Se la shell chiede le '
        'voci a un’interfaccia, al primo rilascio quell’interfaccia può restituire il '
        'contenuto del ConfigMap e in seguito quello delle tabelle, senza che il front-end se '
        'ne accorga. ⚠️ Se invece il file viene letto direttamente dal codice del front-end, '
        'il rientro costa quanto scrivere la funzione due volte: è la differenza fra una '
        'deroga e un debito.')
    par('La raccomandazione è quindi duplice: **si adotti il ConfigMap, ma dietro '
        'un’interfaccia**, e si fissi il momento del rientro — la prima revisione dopo la '
        'messa in esercizio — invece di lasciarlo a quando ci sarà tempo. Una soluzione '
        'tampone senza una data è una soluzione definitiva che non è stata dichiarata tale.')
    fatti.append('sottocapitolo sul menu da ConfigMap, con confronto e modo di rientro')

    # ───────────── 3. la testata e il piè di pagina
    ancora2 = D.h(d, 2, 'L’inventario dei componenti')

    def par2(t, stile='Normal'):
        return D.para(d, ancora2._p, t, stile=stile)

    par2('La testata e il piè di pagina', 'Heading 3')
    par2('La cornice istituzionale è della shell e vale uguale per ogni applicazione '
         'dell’ente: è la scelta che garantisce all’operatore di trovare le stesse '
         'informazioni nello stesso posto passando da un applicativo all’altro. Si compone '
         'di tre bande.')
    par2('In alto una **barra di servizio** sottile, con la preferenza di impaginazione a '
         'sinistra e a destra il nome dell’utente, le sue iniziali in un tondo e il menu del '
         'profilo. Sotto la **testata istituzionale**: il marchio ROMA, lo stemma e la '
         'dicitura «Roma Capitale», su fondo bianco. In fondo il **piè di pagina** su fondo '
         'scuro, con il marchio in negativo e due colonne — i contatti dell’ente, con sede, '
         'partita IVA, codice fiscale e i riferimenti del responsabile della protezione dei '
         'dati, e i canali di comunicazione, oggi sette. Chiude una riga più scura con i '
         'rimandi a privacy e cookie policy.')
    par2('⚠️ **Due osservazioni sulla versione aggiornata del piè di pagina.** La prima: '
         '**la colonna del menu non c’è più.** Le voci di navigazione che una versione '
         'precedente vi collocava sono state tolte, e il piè di pagina resta un elemento '
         'istituzionale e non di navigazione. La seconda: i canali sono passati da cinque a '
         'sette. Entrambe sono modifiche del design system condiviso e non del nostro '
         'perimetro: si recepiscono, non si discutono.')
    par2('⚠️ Resta aperto ciò che il documento sull’identità ha già registrato: **la cornice '
         'non prevede l’indicazione dell’ambiente né la versione di build.** Sono due '
         'elementi che questo disegno continua a raccomandare — il primo impedisce di '
         'operare in esercizio credendo di essere in collaudo, il secondo rende '
         'riconducibile una segnalazione al rilascio che l’ha prodotta — e che oggi non si '
         'disegnano, per scelta del committente, in attesa di parlarne con chi governa il '
         'design system.')
    fatti.append('sezione sulla testata e sul piè di pagina')

    # ───────────── 4. i certificati di postazione
    ancora3 = D.h(d, 2, 'Processi Back-End')

    def par3(t, stile='Normal'):
        return D.para(d, ancora3._p, t, stile=stile)

    par3('La gestione dei certificati di postazione', 'Heading 3')
    par3('È la quinta funzione da costruire e non compariva fra le quattro del paragrafo '
         'precedente, perché non appartiene al percorso di formazione dell’atto: '
         '**amministra gli strumenti con cui una postazione si presenta agli enti esterni.** '
         'Ogni postazione dispone di un certificato in formato PKCS#12, e il registro di '
         'quei certificati è ciò che consente al concentratore di firmare le richieste verso '
         'ANPR e verso ANSC.')
    par3('⚠️ Vale la pena dire che **in SIPO una gestione dei certificati di postazione già '
         'esiste**, con un registro su base dati e un’operazione di arruolamento: non si sta '
         'costruendo una funzione nuova, si sta dando un’interfaccia a una funzione che oggi '
         'è esposta come servizio e alimentata da un programma a riga di comando.')
    D.immagine(d, ancora3._p, os.path.join(IMG, 'cert_elenco.png'), 6.2,
               'La pagina di gestione: ricerca per nome, elenco dei certificati con la sede, '
               'menu di riga per le azioni.')
    par3('Il modulo di caricamento non è una pagina a sé: **sostituisce la riga di ricerca '
         'nella stessa pagina**, lasciando l’elenco visibile sotto. È una scelta economica '
         'in termini di navigazione e coerente con l’uso, perché chi carica un certificato '
         'ha quasi sempre bisogno di vedere che cosa c’è già.')
    D.immagine(d, ancora3._p, os.path.join(IMG, 'cert_inserimento.png'), 6.2,
               'Il modo inserimento: nome, password del contenitore e selezione del file, '
               'con l’elenco che resta visibile.')
    par3('Il flusso completo è in quattro passi, e quasi tutto è già nella libreria.')
    D.tabella(d, ancora3._p, CERT_FLUSSO, modello=d.tables[8], larghezze=[0.9, 2.7, 2.9])
    par3('⚠️ **Un rilievo sulla password.** Nelle schermate fornite il campo mostra il valore '
         'in chiaro mentre lo si digita. Per la parola d’ordine di un contenitore di chiavi '
         'private è una scelta da rivedere: va mascherata come qualunque credenziale, e '
         '**non va conservata dopo il caricamento** — serve soltanto ad aprire il file nel '
         'momento in cui lo si riceve. Il componente di campo riservato esiste già nella '
         'libreria.')
    par3('⚠️ **Una proposta sulle colonne.** L’elenco fornito mostra il nome del certificato '
         'e la sede. Mancano tre informazioni che un registro di certificati dovrebbe avere: '
         '**quando è stato caricato, da chi, e quando scade.** Senza la terza, in '
         'particolare, nessuno si accorge di una scadenza se non quando le chiamate verso '
         'l’ente esterno cominciano a fallire. La data di scadenza non va chiesta: sta '
         'dentro il certificato e si legge all’atto del caricamento.')
    fatti.append('sezione sui certificati di postazione, con le schermate reali')

    # ───────────── 5. requisiti e punti aperti
    t = D.trova_tabella(d, 'ID', 'Requisito', 'Nota')
    D.clona_riga(t, ('RF-FE-12', 'Gestione dei certificati di postazione',
                     'Il remote mfOperation deve offrire l’elenco dei certificati con '
                     'ricerca per nome, il caricamento di un contenitore PKCS#12 con la '
                     'relativa password, la conferma per notifica e l’eliminazione con '
                     'conferma esplicita. ⚠️ La password è mascherata in digitazione e non '
                     'conservata dopo il caricamento.'))
    fatti.append('RF-FE-12 aggiunto')

    t = D.trova_tabella(d, '#', 'Questione', 'Owner')
    n = len(t.rows) - 1
    for voce in (
        ('OP-FE-%d' % (n + 1),
         'Quando si rientra dal menu su ConfigMap al registro su base dati. ⚠️ La deroga è '
         'dichiarata per il primo rilascio: senza una data fissata diventa definitiva. Il '
         'rientro è a basso costo solo se il front-end legge il menu da un servizio e non '
         'dal file.', 'Architettura / Referenti SIPO'),
        ('OP-FE-%d' % (n + 2),
         'Se il registro dei certificati debba esporre data di caricamento, operatore e '
         'scadenza. La scadenza è dentro il certificato e si legge al caricamento: senza, '
         'nessuno se ne accorge finché le chiamate verso l’ente esterno non falliscono.',
         'Analisi / Referenti SIPO'),
        ('OP-FE-%d' % (n + 3),
         'Se l’indicazione dell’ambiente e la versione di build possano entrare nella '
         'cornice istituzionale. Oggi non sono previste dal design system e il committente '
         'non intende modificarlo.', 'Design system / Cliente'),
    ):
        D.clona_riga(t, voce)
    fatti.append('tre punti aperti nuovi')

    # ⚠️ riga vuota in coda alla tabella delle versioni di Angular: residuo
    # dell'aggiunta manuale di Angular 22 nella v0.4. La segnala verifica-conteggi.
    for tab in d.tables:
        if tab.rows[0].cells[0].text.strip() == 'Versione':
            for r in list(tab.rows)[1:]:
                if not any(c.text.strip() for c in r.cells):
                    r._tr.getparent().remove(r._tr)
                    fatti.append('rimossa una riga vuota dalla tabella delle versioni')
            break

    # ───────────── 6. testata e storia
    for tab in d.tables:
        if tab.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in tab.rows:
                v = {'Data consegna': '02/10/2026', 'Versione': '0.5',
                     'Documento': 'ANALISI_Front-End-Angular_v0.5'
                     }.get(r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    D.storia(d, '02/10/2026', '0.5',
             'Architettura · Design system · Processi Front-End · Requisiti · Punti aperti',
             'Aggiunta la gestione dei certificati di postazione, assegnata al remote '
             'mfOperation, con il flusso in quattro passi e le schermate fornite; segnalati '
             'la password mostrata in chiaro e le tre colonne mancanti nel registro. '
             'Descritte la testata e il piè di pagina secondo la versione aggiornata del '
             'design system, che toglie la colonna del menu e porta a sette i canali. Nuovo '
             'sottocapitolo sulla composizione del menu: per il primo rilascio si configura '
             'con un ConfigMap invece che dal registro su base dati, con il confronto fra le '
             'due vie, ciò a cui si rinuncia e la condizione che rende il rientro poco '
             'costoso. Un requisito e tre punti aperti nuovi.')
    fatti.append('testata e storia aggiornate')

    def grassetti(p):
        if '**' not in p.text or D.ha_commenti(p._p):
            return 0
        pezzi = D.segmenta('', p.text)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        for testo, gr in pezzi:
            run = p.add_run(testo)
            run.bold = gr
        return 1
    k = sum(grassetti(p) for p in d.paragraphs)
    for tab in d.tables:
        for r in tab.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    k += grassetti(p)
    fatti.append('%d paragrafi con grassetto applicato' % k)

    d.save(DST)
    print('\n'.join(' · ' + f for f in fatti))
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))
    print('scritto:', os.path.relpath(DST, BASE))


if __name__ == '__main__':
    main()
