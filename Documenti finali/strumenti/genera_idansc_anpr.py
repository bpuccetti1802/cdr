# -*- coding: utf-8 -*-
"""Genera «STUDIO_Identificativi-ANSC-verso-ANPR» — se e dove si possa passare ad ANPR
l'identificativo di ANSC, a partire dalla ricognizione dei processi che aggiornano ANPR.

La domanda del committente era se si possa passare ad ANPR il codice idSoggettoAnsc. La
risposta, verificata sullo schema ANPR generato nel client, è che per il SOGGETTO non si
può — ANPR ne ha un solo identificativo, il proprio — mentre per l'ATTO si può, perché lo
schema prevede `tipoAttoANSC` e il client lo implementa già in due punti.

Il documento riporta la risposta con le prove, mostra dove la catena si interrompe nei
cinque processi schedulati, e dice per ciascuno dove intervenire. Due risultati che vale
la pena non riscoprire: il ramo ANSC commuta **da sé** sul valore di `getIdAttoAnsc()`,
letto **per reflection**; e dei quindici generatori di scrittura solo **due** lo sanno
trattare.

Tutte le affermazioni portano il riferimento file:riga.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/genera_idansc_anpr.py
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = os.path.join(BASE, 'Template documentale', 'template_DAD_roma-capitale.docx')
OUT = os.path.join(BASE, 'Documenti finali',
                   'STUDIO_Identificativi-ANSC-verso-ANPR_v0.1.docx')

TAGLIA_DA = 29


# ------------------------------------------------------------------ primitive
def testo(par, s):
    if par.runs:
        par.runs[0].text = s
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(s)


def costruisci():
    shutil.copyfile(TEMPLATE, OUT)
    d = docx.Document(OUT)
    body = d.element.body
    for ch in list(body.iterchildren())[TAGLIA_DA:]:
        if ch.tag != qn('w:sectPr'):
            body.remove(ch)
    for p in d.paragraphs:
        if '“Lorem Ipsum”' in p.text:
            testo(p, '“Integrazione SIPO – ANSC”')
        elif p.style.name == 'Subtitle':
            testo(p, 'Gli identificativi di ANSC verso ANPR: che cosa è possibile '
                     'e dove si interviene')
    return d


def h(d, livello, t):
    p = d.add_paragraph(style='Heading %d' % livello)
    p.add_run(t)
    return p


def par(d, t, stile='Normal'):
    p = d.add_paragraph(style=stile)
    for pezzo, grassetto in D.segmenta('', t):
        r = p.add_run(pezzo)
        r.bold = grassetto
    return p


def voce(d, testa, corpo):
    p = d.add_paragraph(style='List Paragraph')
    r = p.add_run(testa + ' ')
    r.bold = True
    for pezzo, grassetto in D.segmenta('', corpo):
        rr = p.add_run(pezzo)
        rr.bold = grassetto
    return p


def codice(d, righe):
    for riga in righe:
        p = d.add_paragraph()
        r = p.add_run(riga)
        r.font.name = 'Courier New'
        r.font.size = Pt(8)
        p.paragraph_format.space_after = Pt(0)


def tabella(d, righe, larghezze=None):
    t = d.add_table(rows=len(righe), cols=len(righe[0]))
    t.style = 'Table Grid'
    for i, riga in enumerate(righe):
        for j, v in enumerate(riga):
            cel = t.cell(i, j)
            cel.text = ''
            p = cel.paragraphs[0]
            for pezzo, grassetto in D.segmenta('', str(v)):
                run = p.add_run(pezzo)
                run.bold = grassetto or i == 0
                run.font.size = Pt(9)
    if larghezze:
        for j, w in enumerate(larghezze):
            for riga in t.rows:
                riga.cells[j].width = Inches(w)
        griglia = t._tbl.find(qn('w:tblGrid'))
        if griglia is not None:
            for col, w in zip(griglia.findall(qn('w:gridCol')), larghezze):
                col.set(qn('w:w'), str(int(round(Inches(w).twips))))
    d.add_paragraph()
    return t


# ------------------------------------------------------------------ contenuti
RIFERIMENTI = [
    ['#', 'Riferimento', 'Natura'],
    ['[S1]', 'common/anpr-client — il client ANPR di SIPO: tipi generati dallo schema '
             'Sogei, generatori XML delle richieste, utilità comuni.', 'Sorgente applicativo'],
    ['[S2]', 'back-end/all-anpr-sipo — il concentratore verso ANPR per lo stato civile: '
             'endpoint di allineamento e conversione dei dati.', 'Sorgente applicativo'],
    ['[S3]', 'back-end/all-nas-be, all-dec-be, all-matr-be — i processi schedulati di '
             'allineamento per nascite, decessi e cittadinanza, matrimoni.',
     'Sorgente applicativo'],
    ['[S4]', 'back-end/prog-be, back-end/cri-on-anpr-be — i processi schedulati '
             'dell’elettorale e dei cambi di residenza online.', 'Sorgente applicativo'],
    ['[S5]', 'common/anagrafe-entities — le entità di ANAG_USR, fra cui ATTO con le '
             'colonne di ANSC, e la logica che le scrive.', 'Sorgente applicativo'],
    ['[S6]', 'back-end/varana-be — il back-end delle variazioni anagrafiche, da cui '
             'l’identificativo dell’atto ANSC entra oggi in SIPO.', 'Sorgente applicativo'],
    ['[S7]', 'ANALISI_Integrazione-ANSC (versione corrente) — il disegno della scrittura '
             'degli atti verso ANSC, da cui provengono le nozioni di atto e soggetto ANSC.',
     'Documento di progetto'],
]

SOGGETTO = [
    ['Che cosa cerchiamo', 'Che cosa offre lo schema ANPR', 'Esito'],
    ['Un identificativo del soggetto proveniente da ANSC',
     'Il tipo `TipoIdentificativi` espone **un solo campo**: `idANPR`. Nessun altro tipo '
     'dello schema porta un identificativo di soggetto di provenienza esterna.',
     '**Non esiste**'],
    ['Un identificativo dell’atto proveniente da ANSC',
     'Il tipo `TipoAttoANSC` porta `idANSC`, e il tipo `TipoAttoANPR` è il contenitore che '
     'ammette l’atto **in una delle due forme**: `atto` classico oppure `attoANSC`.',
     '**Esiste**'],
]

TIPO_ATTO_ANSC = [
    ['Elemento', 'Tipo nello schema', 'Obbligatorio', 'Che cos’è'],
    ['idANSC', 'tipoNumeroAttoANSC', 'sì',
     'L’identificativo nazionale dell’atto formato in ANSC.'],
    ['comuneRegistrazione', 'tipoComune', 'sì', 'Il comune che ha registrato l’atto.'],
    ['anno', 'numerico4', 'sì', 'L’anno dell’atto.'],
    ['ufficioMunicipio', 'tipoUfficioMunicipio', 'no',
     'Il municipio, dove l’organizzazione lo preveda.'],
    ['numeroComunale', 'tipoNumeroAtto', 'sì',
     'Il numero attribuito dal comune, distinto dall’identificativo nazionale.'],
    ['dataFormazioneAtto', 'tipoDataAnagrafica', 'no', 'La data di formazione.'],
    ['trascritto', 'numerico1', 'no',
     'Se l’atto sia iscritto o trascritto.'],
]

GENERATORI = [
    ['Operazione', 'Famiglia', 'Tratta l’atto in forma ANSC?'],
    ['1001, 1002, 1013, 1014', 'iscrizione', '**no**'],
    ['2001', 'cancellazione', '**sì**, attraverso l’utilità comune'],
    ['2003, 2011', 'cancellazione', '**no**'],
    ['5001, 5010, 5012, 5013', 'mutazione', '**no**'],
    ['5008', 'mutazione', '**sì**, in proprio e attraverso l’utilità comune'],
    ['A001, A002, A006', 'AIRE', '**no**'],
]

CATENA = [
    ['Tratta', 'Che cosa porta', 'L’identificativo ANSC c’è?'],
    ['Vista degli eventi del sistema legacy',
     'Codice individuale, codice evento, codice famiglia.',
     'no: la vista non lo espone'],
    ['Processo schedulato di allineamento',
     'Interroga il legacy per la scheda anagrafica della persona.',
     'no'],
    ['Scheda restituita dal legacy',
     'Gli estremi dell’atto nella forma classica: numero, parte, serie, anno, comune, '
     'esponente.',
     '**no: il tipo non ha il campo**'],
    ['Conversione nel concentratore',
     'Costruisce le entità del soggetto, della nascita e dell’atto a partire da quei dati.',
     '**no: è qui che si interrompe**'],
    ['Generatore della richiesta',
     'Legge l’identificativo dall’atto e, se c’è, usa la forma ANSC.',
     'legge un valore nullo e ricade sulla forma classica'],
]

INTERVENTI = [
    ['Processo', 'Operazioni', 'Atto coinvolto', 'Si può?', 'Dove si interviene'],
    ['Allineamento nascite, ogni giorno alle 05:00',
     '5008 per il soggetto esistente; 1001 e 1002 per il soggetto nuovo',
     'Atto di nascita',
     '**in parte**',
     'Per la mutazione: la conversione nel concentratore, che non valorizza le colonne di '
     'ANSC. Per l’iscrizione del soggetto nuovo **non basta**: i generatori 1001 e 1002 non '
     'trattano affatto la forma ANSC, e vanno estesi nel client.'],
    ['Allineamento decessi e cittadinanza, ogni giorno alle 05:00',
     '2001 per la cancellazione, 5008 per le mutazioni',
     'Atto di morte',
     '**sì**',
     'La sola conversione nel concentratore: entrambi i generatori trattano già la forma '
     'ANSC. È il caso più favorevole, e coincide con la famiglia del pilota.'],
    ['Allineamento matrimoni, ogni giorno alle 05:00',
     '5008, 5012, 5013',
     'Atto di matrimonio e atto di annullamento',
     '**in parte**',
     'La conversione per il ramo dell’annullamento, che passa dall’utilità comune. Per '
     '5012 e 5013 i generatori non trattano la forma ANSC: vanno estesi.'],
    ['Aggiornamento elettorale, ogni giorno alle 11:00',
     '5010',
     'Nessun atto di stato civile: iscrizioni e cancellazioni di elettori',
     '**non pertinente**',
     'Nulla. L’operazione non porta un atto di stato civile, e il generatore non prevede la '
     'forma ANSC.'],
    ['Recupero cambi di residenza online, alle 04:00, 11:00 e 17:00',
     'P001, P002, P003 e 3002',
     'Nessun atto',
     '**non pertinente**',
     'Nulla. Il processo è in prevalenza entrante e la sola scrittura riguarda lo stato '
     'della pratica.'],
]

APERTI = [
    ['#', 'Questione', 'A chi compete'],
    ['AN-01', 'Se le colonne di ANSC su ANAG_USR.ATTO siano popolate per gli atti che i '
              'processi schedulati trattano. Oggi le scrive l’operatore dalle maschere '
              'delle variazioni anagrafiche, per gli atti che passano da lì: non è detto '
              'che l’atto che genera l’evento notificato dal legacy sia fra quelli.',
     'Analisi / Esercizio'],
    ['AN-02', 'Con quale chiave il concentratore ritrovi l’atto in ANAG_USR a partire dal '
              'codice individuale del soggetto, dato che la vista degli eventi non porta un '
              'riferimento all’atto.', 'Analisi / Sviluppo'],
    ['AN-03', 'Se ANPR accetti la forma ANSC su tutte le operazioni che la dichiarano nello '
              'schema, o se la validazione lato Sogei la ammetta solo per alcune: lo schema '
              'la prevede, il comportamento in esercizio non è documentato fra le sorgenti.',
     'Presidio ANPR / Sogei'],
    ['AN-04', 'Se estendere i generatori che non trattano la forma ANSC — le iscrizioni e '
              'le mutazioni di matrimonio — e con quale priorità: è lavoro sul client '
              'condiviso, quindi con un raggio più ampio del singolo processo.',
     'Architettura / Referenti SIPO'],
    ['AN-05', 'Che cosa valorizzare nei due elementi facoltativi che nessuna delle due '
              'implementazioni popola oggi, ufficioMunicipio e trascritto.',
     'Analisi'],
    ['AN-06', 'Se l’atto di stato civile di MATR_USR debba a sua volta portare '
              'l’identificativo del proprio atto ANSC: oggi porta solo i riferimenti '
              'dell’atto trascritto perché formato altrove, e l’identificativo proprio vive '
              'sull’atto di anagrafe.', 'Analisi / Sviluppo'],
]


def scrivi(d):
    # =============================================================== 1
    h(d, 1, 'Scopo del documento')
    par(d, 'Questo studio risponde a una domanda precisa: se sia possibile passare ad ANPR '
           'l’identificativo con cui ANSC individua un soggetto, e in caso affermativo dove '
           'intervenire nei processi che aggiornano ANPR. La domanda nasce dal disegno '
           'dell’integrazione con ANSC [S7], dove il soggetto riceve un identificativo '
           'proprio, e dalla ricognizione dei processi schedulati che riversano su ANPR gli '
           'eventi di stato civile.')
    par(d, 'La risposta si è rivelata doppia, e per questo il documento non si limita a un '
           'sì o a un no: **per il soggetto non è possibile, per l’atto sì**, e il '
           'meccanismo per l’atto è in parte già realizzato. Le due cose vanno tenute '
           'distinte, perché confonderle porta a cercare un intervento dove non serve e a '
           'non farlo dove serve.')

    h(d, 2, 'Metodo e fonti')
    par(d, 'Ogni affermazione su che cosa ANPR accetti è verificata sui tipi generati dallo '
           'schema Sogei presenti nel client [S1]: sono la rappresentazione del contratto, '
           'e in assenza della documentazione ANPR nel perimetro di lavoro sono la fonte '
           'più vicina al contratto stesso. Ogni affermazione sul comportamento di SIPO è '
           'verificata sul codice, con il riferimento al file e alla riga. Dove una '
           'questione non si chiude sulle sorgenti, il documento la dichiara come punto '
           'aperto invece di colmarla con un’ipotesi.')

    h(d, 2, 'Riferimenti')
    tabella(d, RIFERIMENTI, [0.6, 4.4, 1.6])

    # =============================================================== 2
    h(d, 1, 'La domanda e la risposta')
    par(d, 'ANPR identifica una persona con un solo identificativo, il proprio. Lo schema '
           'non prevede, in nessun punto, un identificativo di soggetto di provenienza '
           'esterna: **l’identificativo con cui ANSC individua un soggetto non ha un posto '
           'dove essere scritto**.')
    par(d, 'Non è una lacuna: è la direzione del rapporto fra i due sistemi. ANPR è la fonte '
           'dell’identità della persona, e ANSC consuma proprio l’identificativo di ANPR per '
           'collegare l’atto alla persona. Passare ad ANPR l’identificativo che ANSC si è '
           'dato equivarrebbe a chiedere alla fonte di registrare il modo in cui un '
           'consumatore la chiama.')
    par(d, 'Lo schema prevede però **l’identificativo dell’atto**. Ed è quasi certamente ciò '
           'che serve: quando un processo comunica ad ANPR una nascita o un decesso, ciò che '
           'ANPR deve sapere non è come ANSC chiami la persona — quella la conosce già — ma '
           '**da quale atto discende la mutazione**. Il resto del documento tratta questa '
           'seconda domanda, che ha risposta affermativa.')
    tabella(d, SOGGETTO, [1.9, 3.3, 1.4])

    # =============================================================== 3
    h(d, 1, 'Che cosa ANPR accetta')
    par(d, 'Nello schema l’atto di un evento non è un tipo unico. Il contenitore '
           '`TipoAttoANPR` ha due proprietà e ne ammette una per volta: `atto`, la forma '
           'classica con numero, parte, serie e anno, oppure `attoANSC`, la forma nuova. '
           'La scelta fra le due non è una variante di comodo: è il modo in cui ANPR '
           'riconosce un atto formato nel registro nazionale invece che su un registro '
           'comunale.')
    par(d, 'La forma ANSC porta sette elementi.')
    tabella(d, TIPO_ATTO_ANSC, [1.5, 1.6, 0.9, 2.5])
    par(d, 'Si noti che la forma ANSC **non sostituisce il numero comunale**: lo richiede, '
           'accanto all’identificativo nazionale e obbligatorio come quello. È coerente con '
           'quanto il disegno dell’integrazione ha già stabilito, cioè che il numero '
           'comunale resta a carico del Comune e nessun servizio nazionale lo assegna.')

    # =============================================================== 4
    h(d, 1, 'Il meccanismo è già nel client, e commuta da sé')
    par(d, 'Il client ANPR implementa la forma ANSC in due punti, con lo stesso criterio.')
    voce(d, 'Nel generatore della mutazione.',
         'Nel metodo che costruisce l’atto di nascita, il generatore della 5008 legge '
         'l’identificativo dall’atto e, se lo trova, costruisce la forma ANSC; altrimenti '
         'costruisce quella classica (`Richiesta5008XmlGenerator.java:343-358`).')
    voce(d, 'In una utilità comune.',
         'Lo stesso criterio è in `AnprUtilities.generateTipoAttoANPR` '
         '(`AnprUtilities.java:427-449`), richiamata dal generatore della cancellazione '
         '(`Richiesta2001XmlGenerator.java:172`), dall’annullamento di matrimonio e '
         'dall’evento di morte della 5008 (`:705` e `:1036`).')
    par(d, 'Il punto da cui tutto dipende è una riga sola:')
    codice(d, [
        'String idAttoAnsc = (String) attoClass.getMethod("getIdAttoAnsc").invoke(atto);',
        'if (idAttoAnsc != null) { ... forma ANSC ... } else { ... forma classica ... }',
    ])
    par(d, 'Ne seguono tre conseguenze che conviene avere chiare prima di programmare un '
           'intervento.')
    voce(d, 'La commutazione è automatica.',
         'Nessun parametro la governa e nessuna configurazione la accende: **basta che '
         'l’atto che arriva al generatore porti l’identificativo**. Dove il valore c’è, la '
         'forma ANSC parte già oggi.')
    voce(d, 'L’accesso è per reflection.',
         'Il generatore non dipende da una classe: invoca `getIdAttoAnsc` per nome. Il '
         'meccanismo è quindi **invisibile a una ricerca per nome di campo**, ed è il motivo '
         'per cui è facile concludere che non esista. Ha un secondo effetto: un oggetto '
         'privo di quel metodo non produce un valore nullo, produce un errore — il metodo '
         'dichiara l’eccezione e non la cattura.')
    voce(d, 'Lo sanno trattare due generatori su quindici.',
         'È il limite vero del meccanismo, e si vede solo contandolo.')
    tabella(d, GENERATORI, [2.2, 1.7, 2.7])
    par(d, 'Le due operazioni che lo trattano sono la mutazione 5008 e la cancellazione '
           '2001. Tutte le altre — **comprese tutte le iscrizioni** — accettano soltanto la '
           'forma classica, perché il generatore non costruisce l’altra.')

    # =============================================================== 5
    h(d, 1, 'Da dove viene oggi l’identificativo dell’atto')
    par(d, 'L’atto di anagrafe di SIPO ha già le colonne che servono: l’identificativo ANSC, '
           'l’anno, il numero comunale e la data (`Atto.java:161`, `:164`, `:175`, `:180`). '
           'Non sono una previsione: sono popolate e lette.')
    par(d, 'A scriverle è una logica dedicata, che riceve l’identificativo come stringa e lo '
           'spezza nelle sue parti (`AttoAnscBL.creaAttoAnsc`). A chiamarla sono il '
           'back-end delle variazioni anagrafiche e quello dei matrimoni [S6], in sette '
           'punti, sempre in risposta a un dato **immesso dall’operatore in una maschera**; '
           'il formato è validato prima, da una utilità richiamata in tutte le aree — '
           'decessi, matrimoni, divorzi, cittadinanza, variazioni.')
    par(d, 'Questo è il fatto da cui dipende tutto il resto: **l’identificativo entra in '
           'SIPO a mano, nel percorso sincrono delle maschere**. I processi schedulati non '
           'lo scrivono e non lo leggono. Da qui la domanda AN-01, che il documento lascia '
           'aperta: se per gli atti che i processi trattano quelle colonne siano popolate.')
    par(d, 'Va ricordata anche l’asimmetria fra i due atti di SIPO, già rilevata altrove: '
           'l’identificativo ANSC dell’atto proprio vive sull’atto **di anagrafe**, mentre '
           'l’atto di stato civile porta soltanto i riferimenti dell’atto trascritto perché '
           'formato altrove. Chi cerca il dato partendo dallo stato civile non lo trova '
           '(AN-06).')

    # =============================================================== 6
    h(d, 1, 'Dove si interrompe la catena nei processi schedulati')
    par(d, 'I tre processi di allineamento seguono tutti lo stesso percorso. Conviene '
           'leggerlo tratta per tratta, perché il punto di rottura è uno solo e non è dove '
           'si immagina.')
    tabella(d, CATENA, [2.3, 3.0, 1.7])
    par(d, 'Il punto di rottura è **la conversione dentro il concentratore**. Il metodo che '
           'costruisce l’atto (`Convert.java:1033`) valorizza numero, parte, serie, anno, '
           'comune ed esponente a partire dagli estremi restituiti dal legacy, e **non '
           'tocca le colonne di ANSC**; lo stesso vale per l’altro convertitore di atto '
           '(`:177`). L’entità che arriva al generatore è quindi un atto completo nella '
           'forma classica e vuoto in quella nuova.')
    par(d, 'Due osservazioni rendono l’intervento più circoscritto di quanto sembri.')
    voce(d, 'Il dato non può venire dal legacy.',
         'Il tipo che il legacy restituisce porta soltanto comune, tipo di registrazione, '
         'anno, numero, parte, serie ed esponente: **non ha un campo per '
         'l’identificativo ANSC**. Non è una questione di valorizzazione mancata — la '
         'sorgente non lo trasporta. Il valore va letto da ANAG_USR, e il concentratore già '
         'interroga quella base dati per altri scopi.')
    voce(d, 'Il concentratore non sa nulla di ANSC.',
         'Nel modulo non c’è una sola occorrenza del termine. Non va quindi corretto un '
         'comportamento sbagliato: va aggiunto un comportamento che non c’è, e questo lo '
         'rende un intervento additivo.')
    par(d, 'Una trappola da evitare nella realizzazione: il concentratore ha un proprio '
           'oggetto per l’atto, privo del metodo che il generatore invoca per reflection. '
           'Passare quello invece dell’entità non produrrebbe il ripiego sulla forma '
           'classica, produrrebbe un errore in esecuzione. Oggi il percorso passa '
           'correttamente l’entità, e così deve restare.')

    # =============================================================== 7
    h(d, 1, 'Gli interventi, processo per processo')
    par(d, 'La ricognizione dei processi che aggiornano ANPR ne ha individuati cinque '
           'schedulati, più il processo che alimenta la coda degli eventi. La tabella dice, '
           'per ciascuno, se il passaggio della forma ANSC sia possibile e dove si '
           'intervenga.')
    tabella(d, INTERVENTI, [1.5, 1.3, 1.2, 0.8, 2.6])
    par(d, 'Ne risulta un ordine di convenienza che non era evidente all’inizio. **Il caso '
           'più favorevole è quello dei decessi**: entrambe le operazioni che il processo '
           'usa trattano già la forma ANSC, quindi l’intervento è circoscritto alla '
           'conversione, in un solo modulo, senza toccare il client condiviso. Ed è anche la '
           'famiglia su cui il progetto di integrazione ha scelto di partire.')
    par(d, 'Il caso delle **nascite** è a metà: la mutazione del soggetto già esistente si '
           'risolve come i decessi, ma l’iscrizione di un soggetto nuovo — che per una '
           'nascita è il caso normale — richiede di estendere due generatori nel client '
           'condiviso. È un lavoro di raggio diverso, da pianificare a parte (AN-04).')
    par(d, 'I due processi dell’elettorale e dei cambi di residenza online **non sono '
           'pertinenti**: non comunicano atti di stato civile, e la questione non li '
           'riguarda. Vale la pena dirlo esplicitamente, perché in una ricognizione dei '
           'processi che aggiornano ANPR compaiono accanto agli altri.')
    par(d, 'Infine il processo che alimenta la coda, in esecuzione ogni giorno alle 23:00: '
           'non chiama ANPR e non va modificato. Resta però il luogo in cui si potrebbe '
           'portare il riferimento all’atto, se si scegliesse di risolvere così la questione '
           'AN-02 invece di cercarlo a valle.')

    # =============================================================== 8
    h(d, 1, 'Che cosa non si può fare, e che cosa si fa invece')
    par(d, 'Il passaggio ad ANPR dell’identificativo con cui ANSC individua un soggetto '
           '**non è realizzabile e non va messo a piano**: lo schema non ha il campo, e '
           'nessuna estensione del client potrebbe crearlo, perché il contratto è di ANPR.')
    par(d, 'La corrispondenza fra i due mondi va dunque tenuta **dalla nostra parte**, e il '
           'disegno dell’integrazione lo prevede già: la tabella che raccoglie gli '
           'identificativi del soggetto associa, per ciascun soggetto di SIPO, '
           'l’identificativo di ANSC e quelli di ANPR. È il luogo giusto, per due ragioni. '
           'La prima è che i due codici di ANPR restano dove li scrivono i canali di ANPR, '
           'senza essere duplicati. La seconda è che la corrispondenza è una conoscenza del '
           'Comune, non un dato che i due sistemi nazionali si scambiano fra loro.')
    par(d, 'Quello che si passa ad ANPR è quindi **l’atto**, nella forma che lo schema '
           'prevede. E ciò che si guadagna è preciso: ANPR registra che la mutazione '
           'discende da un atto del registro nazionale, con il suo identificativo, invece di '
           'riceverne gli estremi comunali come se fosse un atto di registro locale.')

    # =============================================================== 9
    h(d, 1, 'Punti aperti')
    par(d, 'Sei questioni restano da chiudere. Le prime due decidono se l’intervento sia '
           'realizzabile oggi; la terza riguarda il comportamento di ANPR in esercizio, che '
           'le sorgenti non documentano.')
    tabella(d, APERTI, [0.8, 4.3, 1.5])


def intesta(d):
    """Compila il riquadro di testa, la testata di pagina e la Storia del Documento."""
    for r in d.tables[0].rows:
        if len(r.cells) < 2:
            continue
        etichetta = r.cells[0].text.strip().lower()
        valore = {'progetto': 'SIPO \u2014 Integrazione con ANSC',
                  'data consegna': '27/10/2026'}.get(etichetta)
        if valore:
            D.riscrivi_cella(r.cells[1], valore)
    for fatto in D.frontespizio(d, '0.1',
                                'STUDIO_Identificativi-ANSC-verso-ANPR_v0.1'):
        print('  ' + fatto)
    # ⚠️ Il template porta tre righe vuote nella Storia del Documento: restano
    # in pagina come righe bianche se non si rimuovono prima di aggiungere la propria.
    storia = D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')
    for r in list(storia.rows[1:]):
        if not any(c.text.strip() for c in r.cells):
            r._tr.getparent().remove(r._tr)
    D.storia(d, '27/10/2026', '0.1', 'Tutti',
             'Prima stesura. Verifica sullo schema ANPR generato nel client: per il '
             'soggetto non esiste un identificativo di provenienza esterna, per l\u2019atto '
             'lo schema prevede la forma ANSC e il client la implementa in due punti, '
             'commutando sul valore dell\u2019identificativo letto per reflection. Soltanto '
             'due generatori di scrittura su quindici la trattano. Individuato il punto in '
             'cui la catena si interrompe nei processi schedulati \u2014 la conversione nel '
             'concentratore \u2014 e indicati gli interventi per ciascuno dei cinque '
             'processi. Sei punti aperti.')


if __name__ == '__main__':
    d = costruisci()
    scrivi(d)
    intesta(d)
    d.save(OUT)
    print('scritto:', os.path.relpath(OUT, BASE))
    cap, tab, fig = D.riepilogo(OUT)
    print('capitoli %d | tabelle %d | figure %d' % (cap, tab, fig))
