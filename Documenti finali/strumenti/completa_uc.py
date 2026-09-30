# -*- coding: utf-8 -*-
"""Completa il foglio di mappatura di un UC: sezioni mancanti, obbligatorietà, proposte SIPO.

Il foglio di lavoro (`documenti elaborati intermedi/<UC>.xlsx`) è compilato a mano e porta
ciò che nessuna fonte contiene: la colonna SIPO e le regole di transcodifica. Qui NON si
tocca: si aggiunge.

Tre aggiunte, e ciascuna risponde a una domanda che il foglio lasciava aperta.

1. LE SEZIONI MANCANTI. Il mapping ufficiale dell'UC ha più righe del foglio: quelle in più
   si riportano in coda, evidenziate. Non sono errori del compilatore — «Luogo Redazione» è
   condizionata a un caso che SIPO non gestisce — ma vanno viste, non sottintese.

2. L'OBBLIGATORIETÀ. ⚠️ Il mapping la dichiara in DUE colonne che si contraddicono:
   «Obbligatorio» dice NO su `madre.cognome` mentre «Condizioni obbligatorieta'» dice
   «obbligatoria». Il cognome della madre non è facoltativo in un atto di nascita: la colonna
   che porta l'informazione utile è la seconda. Qui si calcola l'esito e si dichiara quale
   colonna ha deciso, così la scelta resta visibile e revocabile.

3. LE PROPOSTE SIPO. Dove la colonna a mano è vuota, il ponte propone tabella e colonna con
   la sua evidenza `file:riga`. ⚠️ Vanno in una colonna SEPARATA: sul banco di prova di
   Dic_Nasc_001 il ponte concorda con il compilatore 36 volte su 41, e le 5 divergenze gli
   danno torto tutte (l'interprete della nascita sta in ATTO_NASCITA_INTERPRETE, non in
   SOGGETTO_INTERPRETE). Una proposta non è un rilievo.

    /usr/bin/python3 "Documenti finali/strumenti/completa_uc.py" [UC] [famiglia] [area]
"""
import csv
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ponte_ansc_sipo as PT   # noqa: E402
import sipo_dizionario as SD   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CARTELLA = os.path.join(BASE, 'documenti elaborati intermedi')
BLU = PatternFill('solid', fgColor='1F3864')
GIALLO = PatternFill('solid', fgColor='FFF2CC')
VERDE = PatternFill('solid', fgColor='E6F2E6')
GRIGIO = PatternFill('solid', fgColor='F2F2F2')

NUOVE = ["Obbligatorio per l'UC", 'Come è stato deciso',
         'Proposta automatica SIPO', 'Evidenza della proposta']


def esito_obbligatorieta(obbligatorio, condizione):
    """L'obbligatorietà del campo per questo UC, e la colonna che l'ha decisa.

    ⚠️ Ordine di precedenza: la condizione governa; la colonna «Obbligatorio» interviene solo
    dove la condizione tace. È la lettura che regge nel merito — vedi il caso madre.cognome.
    """
    c = (condizione or '').strip()
    o = (obbligatorio or '').strip().upper()
    if c.lower() == 'obbligatoria':
        return 'SI', 'condizione = «obbligatoria»'
    if c.lower() == 'opzionale':
        return 'NO', 'condizione = «opzionale»'
    if c:
        return 'CONDIZIONATO', f'condizione: {c}'
    if o in ('SI', 'NO'):
        return o, 'colonna «Obbligatorio» (condizione assente)'
    return '', 'non dichiarato'


def righe_ufficiali(uc, famiglia):
    f = os.path.join(BASE, 'ansc', 'docs', 'Mapping_casi_uso', famiglia, uc + '.csv')
    with open(f, encoding='utf-8-sig') as fh:
        return [r for r in csv.DictReader(fh) if (r.get('Sezione') or '').strip() != 'Sezione']


def completa(uc='Dic_Nasc_001', famiglia='nascita', area='nascita', uscita=None):
    sorgente = os.path.join(CARTELLA, uc + '.xlsx')
    uscita = uscita or os.path.join(CARTELLA, uc + '_v0.2.xlsx')
    wb = openpyxl.load_workbook(sorgente)
    ws = wb.active
    intest = [(c.value or '') for c in ws[1]]
    n0 = len(intest)
    icol = {h: i for i, h in enumerate(intest)}
    i_bo = next(i for i, h in enumerate(intest) if h == 'Binding Object')
    i_bf = next(i for i, h in enumerate(intest) if h == 'Binding Field')
    i_sipo = next(i for i, h in enumerate(intest) if 'SIPO' in h)
    i_obb = next(i for i, h in enumerate(intest) if h.startswith('Dato obbligatorio'))
    i_cond = next(i for i, h in enumerate(intest) if h.startswith('Condizioni obbligatoriet'))
    i_sez = next(i for i, h in enumerate(intest) if h.startswith('Sezione'))

    ponte = PT.Ponte(area, ent=SD.entita(('common', 'back-end')))

    def proposta(bo, bf):
        perc = re.sub(r'^evento\.?', '', f'{bo}.{bf}').strip('.')
        if not perc:
            return '', ''
        e = ponte.risolvi(perc)
        d = ponte.dettaglio(e['percorso_sipo']) if e['percorso_sipo'] else None
        dest = (d or {}).get('destinazione') or {}
        if not dest.get('colonna'):
            return '', e['grado']
        return (f"{dest.get('tabella','')}.{dest['colonna']}",
                dest.get('evidenza', '') or e['grado'])

    # ------------------------------------------------------- intestazione delle nuove colonne
    for k, nome in enumerate(NUOVE):
        c = ws.cell(row=1, column=n0 + 1 + k, value=nome)
        c.font = Font(bold=True, color='FFFFFF', size=9)
        c.fill = BLU
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(n0 + 1 + k)].width = 30

    # ------------------------------------------------------- righe già presenti: si completa
    presenti, divergenze = set(), []
    for r in range(2, ws.max_row + 1):
        val = [ws.cell(row=r, column=j + 1).value for j in range(n0)]
        if not any(v not in (None, '') for v in val):
            continue
        bo, bf = str(val[i_bo] or ''), str(val[i_bf] or '')
        presenti.add((bo.strip(), bf.strip()))
        esito, come = esito_obbligatorieta(val[i_obb], val[i_cond])
        ws.cell(row=r, column=n0 + 1, value=esito).fill = (
            VERDE if esito == 'SI' else (GIALLO if esito == 'CONDIZIONATO' else GRIGIO))
        ws.cell(row=r, column=n0 + 2, value=come)
        prop, evid = proposta(bo, bf)
        a_mano = str(val[i_sipo] or '').strip()
        if not a_mano and prop:
            ws.cell(row=r, column=n0 + 3, value=prop)
            ws.cell(row=r, column=n0 + 4, value=evid)
        elif a_mano and prop:
            u = re.split(r'[ (\n]', a_mano)[0].strip().upper()
            if u != prop.upper():
                divergenze.append((f'{bo}.{bf}', a_mano.splitlines()[0], prop, evid))

    # ------------------------------------------------------- sezioni mancanti, in coda
    manca = [u for u in righe_ufficiali(uc, famiglia)
             if ((u.get('Binding Object') or '').strip(),
                 (u.get('Binding Field') or '').strip()) not in presenti]
    if manca:
        r = ws.max_row + 2
        c = ws.cell(row=r, column=1, value='RIGHE PRESENTI NEL MAPPING UFFICIALE E ASSENTI '
                                           'DAL FOGLIO — da valutare, non compilate')
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = BLU
        r += 1
        for u in manca:
            bo, bf = (u.get('Binding Object') or ''), (u.get('Binding Field') or '')
            esito, come = esito_obbligatorieta(u.get('Obbligatorio'),
                                               u.get("Condizioni obbligatorieta'"))
            valori = {i_sez: u.get('Sezione', ''),
                      icol.get('UseCase', 0): ws.cell(row=2, column=1).value,
                      i_bo: bo, i_bf: bf, i_obb: u.get('Obbligatorio', ''),
                      i_cond: u.get("Condizioni obbligatorieta'", '')}
            if 'Sezione FE ANSC' in icol:
                valori[icol['Sezione FE ANSC']] = u.get('Sezione', '')
            for j in range(n0):
                cell = ws.cell(row=r, column=j + 1, value=valori.get(j, ''))
                cell.fill = GIALLO
            # il nome del campo del mapping, dove il foglio non ha una colonna dedicata
            if not bf and u.get('Campo'):
                ws.cell(row=r, column=i_bf + 1, value=u['Campo']).fill = GIALLO
            ws.cell(row=r, column=n0 + 1, value=esito).fill = GIALLO
            ws.cell(row=r, column=n0 + 2, value=come).fill = GIALLO
            prop, evid = proposta(bo, bf)
            if prop:
                ws.cell(row=r, column=n0 + 3, value=prop).fill = GIALLO
                ws.cell(row=r, column=n0 + 4, value=evid).fill = GIALLO
            r += 1
    return wb, ws, manca, divergenze, uscita


def _foglio_obbligatorieta(wb, uc, famiglia, conta):
    """La spiegazione della regola, con i numeri su cui poggia: è la decisione da validare."""
    ws = wb.create_sheet('Obbligatorietà — come si legge')
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 104
    voci = [
        ('Il problema', None),
        ('Due colonne, non una',
         'Il mapping ufficiale dichiara l’obbligatorietà in «Obbligatorio» e in «Condizioni '
         'obbligatorieta’». Le due non coincidono: su tutto il dominio divergono in 23.111 '
         'righe su 67.300 (34 %) — 22.371 con «Obbligatorio = NO» e condizione '
         '«obbligatoria», 740 con «SI» e condizione «opzionale».'),
        ('Il caso che decide',
         'In questo UC, sezione Madre: tutte e 33 le righe hanno condizione «obbligatoria», '
         'ma «Obbligatorio» dice SI solo per 10 — e fra i NO c’è «cognome», mentre «nome» è '
         'SI. Il cognome della madre non è facoltativo in un atto di nascita: la colonna '
         '«Obbligatorio» non sta dichiarando l’obbligatorietà del campo.'),
        ('La regola applicata qui', None),
        ('  condizione = «obbligatoria»', '→ SI'),
        ('  condizione = «opzionale»', '→ NO'),
        ('  condizione = espressione', '→ CONDIZIONATO, e l’espressione si valuta a runtime: '
                                       'è la stessa grammatica campo,operatore,valore delle '
                                       'regole di determinazione e degli allegati, quindi la '
                                       'valuta lo stesso motore.'),
        ('  condizione assente', '→ vale la colonna «Obbligatorio».'),
        ('Che cosa resta aperto',
         'La semantica delle due colonne non è dichiarata in nessuna fonte ANSC: la regola '
         'qui sopra è la lettura che regge nel merito, non una definizione ufficiale. Va '
         'chiesta a Sogei — è un punto aperto nuovo, non una deduzione da dare per chiusa.'),
        ('L’esito su questo UC', None),
    ]
    for k, v in conta.items():
        voci.append((f'  {k}', f'{v} righe'))
    for a, b in voci:
        ws.append([a, b])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = Font(bold=True, color='FFFFFF' if b is None else '000000')
        if b is None:
            c.fill = BLU
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True,
                                                                vertical='top')


def _foglio_divergenze(wb, divergenze):
    ws = wb.create_sheet('Proposte divergenti')
    for j, (t, w) in enumerate((('Percorso ANSC', 46), ('Compilato a mano', 40),
                                ('Proposta automatica', 40), ('Evidenza', 52)), 1):
        c = ws.cell(row=1, column=j, value=t)
        c.font = Font(bold=True, color='FFFFFF', size=9)
        c.fill = BLU
        ws.column_dimensions[get_column_letter(j)].width = w
    for d in divergenze:
        ws.append(list(d))
    ws.freeze_panes = 'A2'
    return ws


if __name__ == '__main__':
    import collections
    uc = sys.argv[1] if len(sys.argv) > 1 else 'Dic_Nasc_001'
    fam = sys.argv[2] if len(sys.argv) > 2 else 'nascita'
    area = sys.argv[3] if len(sys.argv) > 3 else 'nascita'
    wb, ws, manca, divergenze, uscita = completa(uc, fam, area)
    n0 = ws.max_column - len(NUOVE)
    conta = collections.Counter()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=n0 + 1).value
        if v:
            conta[v] += 1
    _foglio_obbligatorieta(wb, uc, fam, dict(conta.most_common()))
    _foglio_divergenze(wb, divergenze)
    wb.save(uscita)
    print('scritto  :', os.path.relpath(uscita, BASE))
    print('righe aggiunte dal mapping ufficiale:', len(manca),
          '·', dict(collections.Counter(m['Sezione'] for m in manca)))
    print('proposte divergenti dalla compilazione a mano:', len(divergenze))
    print('esito obbligatorietà:', dict(conta.most_common()))
