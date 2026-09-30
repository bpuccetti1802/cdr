# -*- coding: utf-8 -*-
"""La mappatura a tre colonne di nascite e morti: maschera SIPO ↔ colonna Oracle ↔ ModelEvento.

Per ciascun caso d'uso ANSC delle due famiglie rilevate, ogni campo dichiarato dal mapping
ufficiale è riportato con: dove sta in ANSC (percorso del modello evento, tipo, decodifica,
descrizione dal contratto), dove sta nella maschera di SIPO (percorso del DTO, etichetta,
maschere che lo espongono) e dove finisce sul database (tabella e colonna, o nodo XML, o
riga di campo esteso), ciascuno con la propria evidenza `file:riga`.

    /usr/bin/python3 "Documenti finali/strumenti/workbook_nascite_morte.py" [file.xlsx]

⚠️ Non si compila a mano: si rigenera. Se ANSC pubblica una revisione del mapping o se il
codice di SIPO cambia, la correzione manuale andrebbe persa alla prima riesecuzione.

I tre stati che il documento chiede di distinguere, e che qui governano i colori:
  · COERENTE   il campo esiste in entrambi i mondi — si riporta la descrizione del contratto
  · MANCANTE   ANSC lo chiede e SIPO non ce l'ha (o non ha l'intero soggetto)
  · DEPRECATO  ANSC lo dichiara superato: non va configurato, nemmeno se SIPO ce l'ha
"""
import collections
import os
import sys

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mappatura_uc as MU        # noqa: E402
import ponte_ansc_sipo as PT     # noqa: E402
import sipo_dizionario as SD     # noqa: E402
import sipo_modello as SM        # noqa: E402
from workbook_mappatura import BLU, GIALLO, ROSSO, VERDE, foglio  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
USCITA = os.path.join(BASE, 'Documenti finali', 'MAPPATURA_Nascite-Morte_SIPO-ANSC_v0.1.xlsx')
VIOLA = PatternFill('solid', fgColor='EDE4F5')

FAMIGLIE = (('Morte', 'decessi'), ('Nascita', 'nascita'))

STATI = {
    'Deprecato in ANSC': ROSSO,
    'Mancante in SIPO': GIALLO,
    'Soggetto assente in SIPO': GIALLO,
    'Non risolto sul modello': VIOLA,
    'Coerente': VERDE,
    'Coerente (solo maschera)': None,
    'Blocco del modello': None,
}


def stato(riga, esito, dett):
    """Lo stato del campo. L'ordine è una scelta: il deprecato viene prima di tutto, perché
    un campo superato non va configurato nemmeno quando SIPO ce l'ha."""
    if riga['deprecato']:
        return 'Deprecato in ANSC'
    if not riga['risolto']:
        return 'Non risolto sul modello'
    if esito['grado'] == 'blocco':
        return 'Blocco del modello'
    if esito['percorso_sipo']:
        d = (dett or {}).get('destinazione')
        if d and d['genere'] != 'ignoto':
            return 'Coerente'
        return 'Coerente (solo maschera)'
    return ('Soggetto assente in SIPO' if esito['grado'] == 'soggetto assente in SIPO'
            else 'Mancante in SIPO')


def dettaglio_db(d):
    """Come si legge il dato sul database, quando non è una colonna semplice."""
    if not d:
        return ''
    if d['genere'] == 'xml':
        return f'nodo XML {d.get("percorso_xml", "")}'
    if d['genere'] == 'esteso':
        return f'campo esteso ID_EXT_FIELD = «{d.get("chiave_estesa", "")}»'
    if d['genere'] == 'ignoto':
        return 'non determinato: ' + d.get('motivo', '')
    return d.get('fusione') and f'fuso con: {d["fusione"]}' or ''


def costruisci(percorso=USCITA):
    modello = MU.Modello()
    ricognizione = MU.lato_sipo(modello)
    tutte = MU.mappatura(modello, ricognizione)
    ent = SD.entita(('common', 'back-end'))

    righe = []
    for famiglia, area in FAMIGLIE:
        ponte = PT.Ponte(area, ent=ent)
        for r in tutte:
            if r['famiglia'] != famiglia or not r['binding'] or r['sezione'] == 'Allegati':
                continue
            e = ponte.risolvi(r['percorso'])
            d = ponte.dettaglio(e['percorso_sipo']) if e['percorso_sipo'] else None
            dest = (d or {}).get('destinazione')
            # la ricognizione funzionale è un'evidenza indipendente dal codice: dove
            # concorda con il ponte la certezza è piena, dove diverge va guardata
            noti = ricognizione.get(r['percorso'], [])
            rilievo = next((s for s in noti if r['motore'] in s['motori']),
                           noti[0] if noti else None)
            st = stato(r, e, d)
            note = ' '.join(x for x in (e['nota_soggetto'], e['nota_campo'],
                                        (dest or {}).get('motivo', '')) if x)
            righe.append({
                'famiglia': famiglia, 'motore': r['motore'], 'uc': r['uc'],
                'descrizione_uc': r['descrizione_uc'], 'sezione': r['sezione'],
                'campo': r['campo'], 'obbligatorio': r['obbligatorio'],
                'condizione': r['condizione'],
                # ANSC
                'percorso': r['percorso'], 'tipo': r['tipo'], 'formato': r['formato'],
                'lista': 'Sì' if r['lista'] else '', 'decodifica': r['decodifica'],
                'deprecato': 'Sì' if r['deprecato'] else '',
                'descrizione_campo': (r['descrizione_campo'] or '')[:900],
                'classe': r['classe'],
                # SIPO — front-end
                'blocco': e['blocco'], 'percorso_sipo': e['percorso_sipo'],
                'etichetta': (d or {}).get('etichetta', ''),
                'maschere': ', '.join((d or {}).get('maschere', [])[:6]),
                'certezza_etichetta': (d or {}).get('certezza_etichetta', ''),
                'rilievo': rilievo['campo'] if rilievo else '',
                'rilievo_fonte': rilievo['fonte'] if rilievo else '',
                'evidenza_maschera': (d or {}).get('evidenza_maschera', ''),
                # SIPO — base dati
                'tabella': (dest or {}).get('tabella', ''),
                'colonna': (dest or {}).get('colonna', ''),
                'dettaglio_db': dettaglio_db(dest),
                'tipo_java': (dest or {}).get('tipo_java', ''),
                'evidenza_salvataggio': (dest or {}).get('evidenza', ''),
                'ambiguo': (d or {}).get('ambiguo', ''),
                # esito
                'stato': st, 'grado': e['grado'], 'note': note,
                'regola': r['regola'],
            })

    # ⚠️ write_only: le due famiglie fanno quasi un milione di celle e in modalità
    # ordinaria openpyxl le tiene tutte in memoria — il processo veniva ucciso prima di
    # arrivare al salvataggio. In streaming i fogli si scrivono nell'ordine di creazione,
    # quindi la legenda va creata per prima com'è giusto che sia.
    wb = openpyxl.Workbook(write_only=True)
    _legenda(wb, righe)

    col_uc = [
        ('UC (codice motore)', 15, 'motore'), ('UC (codice ANSC)', 13, 'uc'),
        ('Descrizione UC', 34, 'descrizione_uc'), ('Sezione', 20, 'sezione'),
        ('Campo (mapping ANSC)', 26, 'campo'), ('Obbl.', 7, 'obbligatorio'),
        ('Condizione di obbligatorietà', 26, 'condizione'),
        ('ANSC · percorso nel modello evento', 40, 'percorso'),
        ('ANSC · tipo', 10, 'tipo'), ('ANSC · formato', 11, 'formato'),
        ('ANSC · lista', 7, 'lista'), ('ANSC · decodifica', 12, 'decodifica'),
        ('ANSC · deprecato', 10, 'deprecato'),
        ('ANSC · descrizione dal contratto', 46, 'descrizione_campo'),
        ('SIPO FE · percorso nel DTO', 32, 'percorso_sipo'),
        ('SIPO FE · etichetta della maschera', 26, 'etichetta'),
        ('SIPO FE · maschere', 30, 'maschere'),
        ('SIPO FE · certezza etichetta', 15, 'certezza_etichetta'),
        ('SIPO FE · rilievo funzionale', 22, 'rilievo'),
        ('SIPO FE · evidenza', 40, 'evidenza_maschera'),
        ('SIPO DB · tabella', 24, 'tabella'), ('SIPO DB · colonna', 26, 'colonna'),
        ('SIPO DB · come si legge', 40, 'dettaglio_db'),
        ('SIPO DB · tipo Java', 13, 'tipo_java'),
        ('SIPO DB · evidenza', 44, 'evidenza_salvataggio'),
        ('SIPO DB · più destinazioni', 30, 'ambiguo'),
        ('Stato', 22, 'stato'), ('Grado del raccordo', 20, 'grado'),
        ('Note di raccordo e conversione', 60, 'note'),
        ('Classe di conversione', 20, 'classe'),
    ]
    for famiglia, _ in FAMIGLIE:
        foglio(wb, f'{famiglia} — campi per UC',
               col_uc, [r for r in righe if r['famiglia'] == famiglia],
               evidenzia=lambda r: STATI.get(r['stato']))

    _foglio_campi(wb, righe)
    _foglio_mancanti(wb, righe)
    _foglio_deprecati(wb, righe)
    _foglio_solo_sipo(wb, righe, ent)
    _foglio_soggetti(wb)

    wb.save(percorso)
    return percorso, righe


def _aggrega(righe):
    """Una riga per (famiglia, percorso ANSC): il campo si mappa una volta, non per UC."""
    agg = collections.OrderedDict()
    for r in righe:
        k = (r['famiglia'], r['percorso'])
        a = agg.get(k)
        if a is None:
            a = agg[k] = dict(r, n_uc=0, n_obbl=0, _uc=set(), _sez=set())
        a['n_uc'] += 1
        if (r['obbligatorio'] or '').upper().startswith('SI'):
            a['n_obbl'] += 1
        a['_uc'].add(r['motore'])
        a['_sez'].add(r['sezione'])
    for a in agg.values():
        a['uc_distinti'] = len(a['_uc'])
        a['elenco_uc'] = ', '.join(sorted(a['_uc'])[:40])
        a['sezioni'] = ', '.join(sorted(a['_sez'])[:6])
    return list(agg.values())


COL_CAMPO = [
    ('Famiglia', 11, 'famiglia'),
    ('ANSC · percorso nel modello evento', 42, 'percorso'),
    ('Sezioni del mapping', 26, 'sezioni'),
    ('UC che lo usano', 12, 'uc_distinti'),
    ('di cui obbligatorio', 12, 'n_obbl'),
    ('ANSC · tipo', 10, 'tipo'), ('ANSC · formato', 11, 'formato'),
    ('ANSC · decodifica', 12, 'decodifica'), ('ANSC · deprecato', 10, 'deprecato'),
    ('ANSC · descrizione dal contratto', 50, 'descrizione_campo'),
    ('SIPO FE · percorso nel DTO', 32, 'percorso_sipo'),
    ('SIPO FE · etichetta', 24, 'etichetta'),
    ('SIPO FE · maschere', 30, 'maschere'),
    ('SIPO DB · tabella', 24, 'tabella'), ('SIPO DB · colonna', 26, 'colonna'),
    ('SIPO DB · come si legge', 38, 'dettaglio_db'),
    ('SIPO DB · evidenza', 44, 'evidenza_salvataggio'),
    ('Stato', 22, 'stato'), ('Grado del raccordo', 20, 'grado'),
    ('Note di raccordo e conversione', 60, 'note'),
    ('Elenco degli UC', 50, 'elenco_uc'),
]


def _foglio_campi(wb, righe):
    foglio(wb, 'Campi (una riga per campo)', COL_CAMPO, _aggrega(righe),
           evidenzia=lambda r: STATI.get(r['stato']))


def _foglio_mancanti(wb, righe):
    """Il foglio di lavoro del divario: ciò che ANSC chiede e SIPO non ha, per peso."""
    m = [a for a in _aggrega(righe)
         if a['stato'] in ('Mancante in SIPO', 'Soggetto assente in SIPO')]
    m.sort(key=lambda a: (-a['n_obbl'], -a['uc_distinti']))
    foglio(wb, 'Mancanti in SIPO', COL_CAMPO, m,
           evidenzia=lambda r: ROSSO if r['n_obbl'] else GIALLO)


def _foglio_deprecati(wb, righe):
    d = [a for a in _aggrega(righe) if a['deprecato']]
    d.sort(key=lambda a: (-a['uc_distinti'],))
    foglio(wb, 'Deprecati in ANSC', COL_CAMPO, d, evidenzia=lambda r: ROSSO)


def _foglio_solo_sipo(wb, righe, ent):
    """Il verso opposto: i campi che le maschere raccolgono e nessun UC chiede.

    Serve a due cose: dice che cosa resta locale al Comune (e non va nel payload), e mette in
    guardia dai campi che si crederebbero da mappare e non lo sono.
    """
    usati = {r['percorso_sipo'] for r in righe if r['percorso_sipo']}
    fuori = []
    for famiglia, area in FAMIGLIE:
        alb = SM.albero(area, ent)
        for perc, c in sorted(alb.items()):
            if c['dto'] or not c['maschere'] or perc in usati:
                continue
            d = c['destinazione']
            fuori.append({
                'famiglia': famiglia, 'percorso_sipo': perc,
                'etichetta': c['etichetta'],
                'maschere': ', '.join(c['maschere'][:8]),
                'n_maschere': len(set(c['maschere'])),
                'certezza_etichetta': c['certezza_etichetta'],
                'tipo_java': c['tipo'],
                'tabella': (d or {}).get('tabella', ''),
                'colonna': (d or {}).get('colonna', ''),
                'dettaglio_db': dettaglio_db(d),
                'evidenza_maschera': c['evidenza_maschera'],
                'evidenza_salvataggio': (d or {}).get('evidenza', ''),
            })
    foglio(wb, 'Solo SIPO (non richiesti)', [
        ('Famiglia', 11, 'famiglia'),
        ('SIPO FE · percorso nel DTO', 34, 'percorso_sipo'),
        ('SIPO FE · etichetta', 26, 'etichetta'),
        ('SIPO FE · maschere', 34, 'maschere'),
        ('Numero di maschere', 12, 'n_maschere'),
        ('Certezza etichetta', 15, 'certezza_etichetta'),
        ('Tipo Java', 16, 'tipo_java'),
        ('SIPO DB · tabella', 24, 'tabella'), ('SIPO DB · colonna', 26, 'colonna'),
        ('SIPO DB · come si legge', 38, 'dettaglio_db'),
        ('SIPO FE · evidenza', 42, 'evidenza_maschera'),
        ('SIPO DB · evidenza', 44, 'evidenza_salvataggio'),
    ], fuori)


def _foglio_soggetti(wb):
    """La tabella di raccordo dei soggetti: è la parte del metodo che richiede giudizio, e
    per questo si pubblica insieme al risultato invece di restare dentro il programma."""
    righe = []
    for area, voci in PT.SOGGETTI.items():
        famiglia = 'Morte' if area == 'decessi' else 'Nascita'
        for blocco, prefisso, nota in voci:
            righe.append({
                'famiglia': famiglia,
                'blocco': blocco or '(radice dell’evento)',
                'prefisso': ('(radice del DTO)' if prefisso == ''
                             else prefisso if prefisso else '— nessun corrispondente —'),
                'esito': 'raccordato' if prefisso is not None else 'assente in SIPO',
                'nota': nota,
            })
    foglio(wb, 'Raccordo dei soggetti', [
        ('Famiglia', 11, 'famiglia'),
        ('ANSC · blocco del modello evento', 36, 'blocco'),
        ('SIPO · percorso nel DTO', 30, 'prefisso'),
        ('Esito', 20, 'esito'),
        ('Perché', 90, 'nota'),
    ], righe, evidenzia=lambda r: GIALLO if r['esito'] != 'raccordato' else None)


def _legenda(wb, righe):
    ws = wb.create_sheet('Legenda')
    c = collections.Counter(r['stato'] for r in righe)
    per_fam = {f: collections.Counter(r['stato'] for r in righe if r['famiglia'] == f)
               for f, _ in FAMIGLIE}
    campi = _aggrega(righe)
    voci = [
        ('Mappatura dei campi di nascita e morte', 'SIPO ↔ ANSC · Roma Capitale'),
        ('Che cos’è',
         'Per ciascun caso d’uso ANSC delle due famiglie rilevate, ogni campo dichiarato dal '
         'mapping ufficiale con: dove sta in ANSC, dove sta nella maschera di SIPO e in quale '
         'colonna Oracle finisce. Ogni riga porta le proprie evidenze file:riga.'),
        ('Come si aggiorna',
         'Si rigenera con «Documenti finali/strumenti/workbook_nascite_morte.py». Non si '
         'corregge a mano: alla prima riesecuzione la correzione andrebbe persa.'),
        ('Perimetro',
         'Morte e nascita, le due famiglie per cui esistono sia la ricognizione funzionale '
         'sia il codice dell’area. Le altre cinque famiglie non sono coperte.'),
        ('Le fonti', None),
        ('  · struttura ANSC', 'ansc/docs/openapi/model_evento.yaml — un solo albero, radice '
                               'ModelEvento.'),
        ('  · fabbisogno per UC', 'ansc/docs/Mapping_casi_uso/{morte,nascita}/*.csv.'),
        ('  · struttura SIPO', 'front-end/{decessi,nascita}-web — AttoDecessoModel e '
                               'AttoNascitaModel, l’albero che le maschere legano.'),
        ('  · maschere', 'templates/gestione{Decessi,Nascita}/*.html + messages.properties, '
                         'per l’etichetta che l’operatore vede.'),
        ('  · colonne', 'back-end/{decessi,nascita}-be/**/Salvataggio*.java + le entity JPA, '
                        'per la colonna su cui il dato si posa.'),
        ('  · rilievo funzionale', 'Sorgenti Documentali/Decessi_ANSC.xlsx e '
                                   'Nascite_ANSC_07.08.2026.xlsx.'),
        ('Come si raccorda', None),
        ('  1. il soggetto', 'evento.intestatari[] ↔ deceduto, madre ↔ madre, … La tabella è '
                             'nel foglio «Raccordo dei soggetti»: è la parte che richiede '
                             'giudizio e per questo si pubblica.'),
        ('  2. il campo', 'idComuneNascita ↔ comuneNascita, per regola di prefisso o per '
                          'sinonimo dichiarato. Il raccordo per solo nome di campo non '
                          'funziona: il mapping nomina «Cognome» senza dire di chi.'),
        ('Gli stati', None),
        ('  Coerente', 'Il campo esiste in entrambi i mondi e se ne conosce la colonna: si '
                       'riporta la descrizione del contratto ANSC.'),
        ('  Coerente (solo maschera)',
         'Il campo esiste nella maschera ma la colonna non è stata determinata leggendo il '
         'codice: va rilevata a mano prima di configurarlo.'),
        ('  Mancante in SIPO', 'ANSC lo chiede, il soggetto c’è, il campo no.'),
        ('  Soggetto assente in SIPO',
         'Manca l’intero blocco (i comparenti, l’interprete, l’ente dichiarante…): non è un '
         'campo da aggiungere ma una parte di maschera da progettare.'),
        ('  Deprecato in ANSC', 'ANSC lo dichiara superato: non va configurato nemmeno se '
                                'SIPO ce l’ha.'),
        ('  Non risolto sul modello',
         'Il mapping cita un percorso che il modello evento non contiene (OP-52).'),
        ('Il conto', None),
    ]
    for nome, cnt in (('  · complessivo', c), *((f'  · {f.lower()}', per_fam[f])
                                                for f, _ in FAMIGLIE)):
        voci.append((nome, ' · '.join(f'{k}: {v}' for k, v in cnt.most_common())))
    voci += [
        ('  · righe', f'{len(righe)} righe UC × campo, per {len({r["motore"] for r in righe})} '
                      f'casi d’uso'),
        ('  · campi distinti', f'{len(campi)} campi distinti per famiglia, dei quali '
                               f'{sum(1 for a in campi if a["stato"] == "Coerente")} coerenti '
                               f'con una colonna nota'),
        ('Che cosa resta da fare', None),
        ('  · rilevare le colonne',
         f'{sum(1 for a in campi if a["stato"] == "Coerente (solo maschera)")} campi sono '
         f'nella maschera ma senza colonna determinata dal codice.'),
        ('  · colmare le assenze',
         f'{sum(1 for a in campi if a["stato"] == "Mancante in SIPO")} campi mancano in SIPO '
         f'e {sum(1 for a in campi if a["stato"] == "Soggetto assente in SIPO")} appartengono '
         f'a blocchi interi che la maschera non prevede.'),
    ]
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 116
    for a, b in voci:
        titolo = WriteOnlyCell(ws, value=a)
        testo = WriteOnlyCell(ws, value=b)
        if b is None:
            titolo.font = Font(bold=True, color='FFFFFF')
            titolo.fill = BLU
        else:
            titolo.font = Font(bold=True)
        testo.alignment = Alignment(wrap_text=True, vertical='top')
        ws.append([titolo, testo])


if __name__ == '__main__':
    dove, righe = costruisci(sys.argv[1] if len(sys.argv) > 1 else USCITA)
    c = collections.Counter(r['stato'] for r in righe)
    print('scritto:', os.path.relpath(dove, BASE))
    print('righe  :', len(righe))
    for k, v in c.most_common():
        print(f'   {k:28} {v:6}')
