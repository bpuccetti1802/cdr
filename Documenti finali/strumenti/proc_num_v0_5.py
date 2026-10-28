# -*- coding: utf-8 -*-
"""v0.5 del processo di numerazione comunale, dalla v0.4 (28/10/2026).

Decisioni dell'utente recepite:
  - il servizio della numerazione si chiama be-numeratore-comunale;
  - le sue operazioni sono le quattro descritte nel capitolo del servizio (POST /richiesta,
    GET /ricerca, PUT /utilizza, PUT /scarta): le dieci di ansc-numerazione-v1.yaml non esistono
    più, e ciò che le quattro non coprono si dichiara mancante con rinvio a OP-79 dell'analisi;
  - colonne TIMESTAMP con il prefisso ts_ (NUMERATORE_COMUNALE.ts_inserimento/ts_aggiornamento).
Tolti anche i residui della serie a sequence (superata dal numeratore per anno).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/proc_num_v0_5.py
"""
import os
import sys

import docx
from docx.text.paragraph import Paragraph

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'PROCESSO_Numerazione-Comunale_SIPO-ANSC_v0.4.docx')
OUT = os.path.join(BASE, 'PROCESSO_Numerazione-Comunale_SIPO-ANSC_v0.5.docx')
IMG = os.path.join(QUI, 'img')
OP = 'OP-79 dell’analisi'


def _paragrafi(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs


def sost(doc, vecchio, nuovo, attese=1):
    """Sostituisce nel run che contiene il testo (conserva il grassetto d'avvio), altrimenti
    riscrive il paragrafo con testo_di, che protegge i commenti."""
    n, visti = 0, set()
    for p in _paragrafi(doc):
        if id(p._p) in visti or vecchio not in p.text:
            continue
        visti.add(id(p._p))
        for r in p.runs:
            if vecchio in r.text:
                r.text = r.text.replace(vecchio, nuovo)
                break
        else:
            D.testo_di(p, p.text.replace(vecchio, nuovo))
        n += 1
    if n != attese:
        raise SystemExit(f'«{vecchio[:70]}»: attese {attese}, trovate {n}')


def par(doc, inizio):
    trovati = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(trovati) != 1:
        raise SystemExit(f'«{inizio[:60]}»: {len(trovati)} paragrafi')
    return trovati[0]


def via(elemento):
    if D.ha_commenti(elemento):
        raise SystemExit('elemento con commento da eliminare')
    elemento.getparent().remove(elemento)


def tabella(doc, *intestazioni):
    for t in doc.tables:
        testa = [c.text.strip() for c in t.rows[0].cells]
        if all(i in testa for i in intestazioni):
            return t
    raise SystemExit('tabella non trovata: %s' % (intestazioni,))


def riga(t, primo):
    for r in t.rows:
        if r.cells[0].text.strip().startswith(primo):
            return r
    raise SystemExit('riga non trovata: ' + primo)


def cella(t, primo, col, testo):
    D.riscrivi_cella(riga(t, primo).cells[col], testo)


def riscrivi(doc, inizio, testo):
    """Riscrive un paragrafo conservando l'avvio in grassetto, se c'è."""
    p = par(doc, inizio)
    runs = p.runs
    if len(runs) > 1 and runs[0].bold and testo.startswith(runs[0].text):
        runs[1].text = testo[len(runs[0].text):]
        for r in runs[2:]:
            r.text = ''
    else:
        D.testo_di(p, testo)


# ───────────────────────────────────────────────────────────── testa e glossario
def testa(doc):
    t = tabella(doc, '#', 'Documento')
    cella(t, '[R2]', 1, 'ANALISI_Integrazione-ANSC_v3.54.docx, «La numerazione comunale» e '
          'Open Point')
    cella(t, '[R2]', 2, 'La serie per anno, le garanzie, le strutture ereditate, i percorsi, '
          'il servizio be-numeratore-comunale con le sue quattro operazioni, lo stacco passo per '
          'passo; OP-79, le operazioni che mancano al servizio.')
    cella(t, '[R4]', 1, 'ansc-numerazione-v1.yaml e SPEC_API_Integrazione-ANSC_v0.4')
    cella(t, '[R4]', 2, 'Le quattro operazioni del servizio della numerazione.')
    cella(t, '[R5]', 1, 'WORKFLOW_Interfacce-SIPO-ANSC_v0.5.docx')

    g = tabella(doc, 'Termine', 'Significato')
    cella(g, 'Serie', 1, 'I numeri di un anno: la riga dell’anno in NUMERATORE_COMUNALE, il cui '
          'num_atto è l’ultimo numero staccato.')
    cella(g, 'Buco', 1, 'Un numero non superiore all’ultimo staccato che non ha una riga in '
          'DATI_ATTO_ESTERNO. Lo stacco non ne produce: lo può lasciare solo una registrazione a '
          'posteriori che scavalca la serie.')
    r = riga(g, 'num-ansc-sipo')
    D.riscrivi_cella(r.cells[0], 'be-numeratore-comunale')
    D.riscrivi_cella(r.cells[1], 'Il servizio della numerazione: l’unico componente che stacca '
                     'numeri. Non parla con ANSC. Base path /api/v1/numerazione-comunale.')

    a = tabella(doc, 'Nome', 'Descrizione')
    cella(a, 'Ufficiale di stato civile', 1, 'Abilitato al back-office e ad ANSC. Chiede i numeri '
          'per gli atti da formare sulla web app e ne aggiorna lo stato. Ruolo USC.')
    cella(a, 'Amministratore del back-office', 1, 'Consulta i numeri staccati. La lettura della '
          'serie con i buchi e la sua creazione anticipata non hanno ancora un’operazione '
          '(' + OP + ').')
    cella(a, 'Concentratore operativo', 1, 'Stacca il numero di un atto formato in SIPO '
          'all’anteprima e lo conserva sull’atto (ANSC_STATO_ATTO.NUMERO_ATTO); lo porta a '
          'Utilizzato al deposito riuscito o a Scartato se l’atto è abbandonato. Agisce con il '
          'token dell’operatore.')
    r = riga(a, 'Servizio della numerazione')
    D.riscrivi_cella(r.cells[0], 'Servizio della numerazione (be-numeratore-comunale)')
    D.riscrivi_cella(r.cells[1], 'Governa la serie: stacchi serializzati e cambi di stato, con '
                     'le quattro operazioni POST /richiesta, GET /ricerca, PUT /utilizza e '
                     'PUT /scarta.')


# ───────────────────────────────────────────────────────────── processo e scenari
def processo(doc):
    sost(doc, 'Segno il numero come Utilizzato, indicando l’identificativo nazionale che ANSC ha '
              'assegnato: lo stato passa da 1 Disponibile a 2 Utilizzato.',
         'Segno il numero come Utilizzato (PUT /utilizza): lo stato passa da 1 Disponibile a 2 '
         'Utilizzato. L’identificativo nazionale non si registra sul numero, perché '
         'l’operazione non lo porta (' + OP + ').')
    sost(doc, 'Segno il numero come Scartato, dopo una conferma:',
         'Segno il numero come Scartato (PUT /scarta), dopo una conferma:')
    sost(doc, 'Leggo le richieste e i numeri in tabella, con le serie in testa.',
         'Leggo i numeri in tabella (GET /ricerca).')
    sost(doc, 'Sui numeri Disponibili posso cambiare lo stato; su un Utilizzato posso legare '
              'l’atto di SIPO.',
         'Sui numeri Disponibili posso cambiare lo stato.')
    sost(doc, 'All’anteprima il concentratore stacca un numero legato all’atto, senza che '
              'l’ufficiale apra il back-office.',
         'All’anteprima il concentratore stacca un numero (POST /richiesta) e lo conserva '
         'sull’atto, in ANSC_STATO_ATTO.NUMERO_ATTO, senza che l’ufficiale apra il back-office.')
    sost(doc, 'Se la validazione R009 riesce, il numero passa a Utilizzato con l’identificativo '
              'nazionale.',
         'Se la validazione R009 riesce, il numero passa a Utilizzato (PUT /utilizza), legato '
         'all’atto di SIPO.')
    sost(doc, 'Se la validazione è respinta, il numero resta all’atto e si riusa al tentativo '
              'successivo: la serie non ne consuma un altro.',
         'Se la validazione è respinta, il numero resta all’atto e si riusa al tentativo '
         'successivo: il riuso lo garantisce il concentratore, che legge il numero dall’atto e non '
         'ne chiede un altro, perché la richiesta al servizio non porta l’atto.')
    sost(doc, 'Se l’atto è abbandonato prima del deposito, il numero passa a Scartato.',
         'Se l’atto è abbandonato prima del deposito, il numero passa a Scartato (PUT /scarta).')
    sost(doc, 'Le righe sono visibili nella pagina di ricerca, con l’atto di SIPO valorizzato.',
         'Le righe sono visibili nella pagina di ricerca, con l’atto di SIPO valorizzato dal '
         'deposito riuscito.')
    sost(doc, 'È ciò che Side annuncia come «futuro allineamento».',
         'È ciò che Side annuncia come «futuro allineamento». Nessuna delle quattro operazioni '
         'del servizio imposta lo stato 4: lo scenario attende un’operazione (' + OP + ').')

    t = tabella(doc, 'Percorso', 'Chi chiama')
    cella(t, 'Atto formato in SIPO', 2, 'POST /richiesta all’anteprima, con il numero conservato '
          'sull’atto; PUT /utilizza con l’atto di SIPO al deposito riuscito; PUT /scarta se '
          'l’atto è abbandonato.')
    cella(t, 'Atto formato sulla web app', 2, 'POST /richiesta; PUT /utilizza o PUT /scarta. Lo '
          'stato 4 Evento acquisito all’aggancio non ha un’operazione (' + OP + ').')
    cella(t, 'Parto plurimo', 2, 'POST /richiesta con numeroElementi pari ai nati.')
    cella(t, 'SIPO non disponibile', 2, 'Nessuna: la registrazione a posteriori manca nel '
          'servizio (' + OP + ').')
    cella(t, 'Governo delle serie', 2, 'GET /ricerca. La lettura della serie con i buchi e la sua '
          'creazione anticipata non hanno operazioni (' + OP + ').')


# ───────────────────────────────────────────────────────────── le due pagine
def pagine(doc):
    sost(doc, 'Anno e numero di elementi sono obbligatori; l’atto di SIPO è facoltativo e ammesso '
              'solo con un numero, per legare subito il numero a un atto.',
         'Anno e numero di elementi sono obbligatori e sono i soli campi di POST /richiesta: '
         'l’atto di SIPO si lega dopo, con PUT /utilizza.')
    D.sostituisci_immagine(doc, 'Numerazione comunale — richiesta.',
                           os.path.join(IMG, 'pr_num_richiesta.png'))
    sost(doc, 'sotto, la registrazione a posteriori.',
         'sotto, la nota sulla registrazione a posteriori.')
    riscrivi(doc, 'RICHIEDI.', 'RICHIEDI. Chiama POST /richiesta: il servizio fa avanzare '
             'NUMERATORE_COMUNALE e inserisce una riga in PRENOTAZIONE_ATTI per la richiesta e '
             'una in DATI_ATTO_ESTERNO per ogni numero staccato, nello stato 1 Disponibile e '
             'senza identificativo nazionale.')
    riscrivi(doc, 'Spunta — Segna come utilizzato.', 'Spunta — Segna come utilizzato. Solo su un '
             'numero Disponibile. Apre una finestra che chiede l’atto di SIPO, facoltativo; '
             'chiama PUT /utilizza. L’identificativo nazionale non si chiede, perché '
             'l’operazione non lo porta (' + OP + ').')
    sost(doc, 'dopo una conferma; chiama aggiornaNumero.', 'dopo una conferma; chiama PUT /scarta.')
    riscrivi(doc, 'Registrazione a posteriori.', 'Registrazione a posteriori. Il servizio non ha '
             'un’operazione per registrare un numero già usato mentre SIPO non era disponibile '
             '(' + OP + '): la pagina la accoglierà quando l’operazione esisterà.')
    D.sostituisci_immagine(doc, 'La finestra «Segna come utilizzato».',
                           os.path.join(IMG, 'pr_num_utilizzato.png'))
    D.sostituisci_immagine(doc, 'Il menu del back-office: la mattonella',
                           os.path.join(IMG, 'pr_num_menu.png'))

    sost(doc, 'In testa, la tabella delle serie dice per ciascun anno il prossimo numero, i '
              'conteggi per stato e i buchi.',
         'I filtri sono quelli di GET /ricerca. Non c’è la tabella delle serie con il prossimo '
         'numero, i conteggi per stato e i buchi: nessuna operazione del servizio la alimenta '
         '(' + OP + ').')
    D.sostituisci_immagine(doc, 'Numerazione comunale — ricerca e consultazione.',
                           os.path.join(IMG, 'pr_num_ricerca.png'))
    t = tabella(doc, 'Elemento', 'Che cosa fa')
    for primo in ('Serie della numerazione', 'Occhio', 'Matita', 'CREA SERIE'):
        via(riga(t, primo)._tr)
    cella(t, 'CERCA', 1, 'I numeri secondo i filtri, paginati, ordinati per richiesta e numero.')
    cella(t, 'CERCA', 2, 'GET /ricerca')
    cella(t, 'Spunta e croce', 2, 'PUT /utilizza, PUT /scarta')


# ───────────────────────────────────────────────────────────── tabelle e DDL
def tabelle(doc):
    t = tabella(doc, 'Campo', 'Alimentazione')     # la prima: PRENOTAZIONE_ATTI
    dae = next(x for x in doc.tables if x.rows[0].cells[0].text.strip() == 'Campo'
               and any(r.cells[0].text.strip() == 'id_ansc_collegato' for r in x.rows))
    cella(dae, 'nr_atto_comunale', 3, 'Da NUMERATORE_COMUNALE (num_atto + 1)')
    cella(dae, 'id_ansc_collegato', 3, 'Nessuna operazione del servizio la scrive (' + OP + ')')
    cella(dae, 'id_atto_sipo', 3, 'PUT /utilizza')
    cella(dae, 'ts_allineamento_ansc_side', 3, 'Con il legame all’atto agganciato (stato 4), che '
          'non ha un’operazione (' + OP + ')')
    num = next(x for x in doc.tables if x.rows[0].cells[0].text.strip() == 'Campo'
               and any(r.cells[0].text.strip() == 'dt_inserimento' for r in x.rows))
    for vecchio, nuovo, nota in (('dt_inserimento', 'ts_inserimento',
                                  'Quando la riga dell’anno è stata creata'),
                                 ('dt_aggiornamento', 'ts_aggiornamento',
                                  'Ultimo stacco che ha fatto avanzare il numeratore')):
        r = riga(num, vecchio)
        D.riscrivi_cella(r.cells[0], nuovo)
        D.riscrivi_cella(r.cells[1], 'TIMESTAMP')
        D.riscrivi_cella(r.cells[2], nota)
    del t

    # DDL di NUMERATORE_COMUNALE: punto di troppo, prefisso ts_ e chiave primaria
    p = par(doc, 'num_atto              NUMBER(6,0).')
    D.testo_di(p, p.text.replace('NUMBER(6,0).       NOT NULL', 'NUMBER(6,0)        NOT NULL'))
    p = par(doc, 'dt_inserimento       TIMESTAMP,')
    D.testo_di(p, p.text.replace('dt_inserimento', 'ts_inserimento'))
    p = par(doc, 'dt_aggiornamento     TIMESTAMP')
    D.testo_di(p, p.text.replace('dt_aggiornamento     TIMESTAMP',
                                 'ts_aggiornamento     TIMESTAMP,'))
    D.riga_dopo(doc, p, '  CONSTRAINT numeratore_comunale_pk PRIMARY KEY (id_anno_documento)')


# ───────────────────────────────────────────────────────────── servizio, Side, punti aperti
def servizio(doc):
    h = par(doc, 'Il servizio be-numerazione-comunale')
    D.testo_di(h, 'Il servizio be-numeratore-comunale')
    p = par(doc, 'Il servizio espone le operazioni per staccare numeri comunali consecutivi')
    D.riga_dopo(doc, p, 'Sono le sole operazioni del servizio. Ciò che non coprono — la '
                'registrazione a posteriori di un numero già usato, lo stato 4 Evento acquisito, '
                'l’identificativo nazionale sul numero, la lettura della serie con i buchi e la '
                'sua creazione anticipata, l’operatore reale della richiesta (oggi sempre '
                '"admin") — manca ed è registrato come ' + OP + '.')


def side(doc):
    t = tabella(doc, 'Elemento', 'In Side', 'A Roma')
    cella(t, 'La serie', 2, 'Una riga per anno in NUMERATORE_COMUNALE, unica per tutti i '
          'registri, sei cifre al massimo.')
    cella(t, 'Segna come utilizzato', 2, 'Lega il numero all’atto di SIPO, se indicato; Scartato '
          'ed Evento acquisito sono finali. L’identificativo nazionale non si registra: '
          'PUT /utilizza non lo porta (' + OP + ').')
    c = riga(t, 'SIDe non disponibile').cells[2]
    D.riscrivi_cella(c, c.text.rstrip() + ' L’operazione non c’è ancora nel servizio (' + OP
                     + ').')
    cella(t, 'Numeri consecutivi', 2, 'Stacchi serializzati con il blocco della riga dell’anno '
          'in NUMERATORE_COMUNALE (SELECT … FOR UPDATE).')
    cella(t, 'Numeri consecutivi', 3, 'Il numeratore avanza nella stessa transazione che '
          'registra i numeri: i numeri di una richiesta sono consecutivi.')
    cella(t, 'Buchi della serie', 2, 'Lo stacco non ne produce; li può lasciare solo una '
          'registrazione a posteriori che scavalca la serie. La loro lettura non ha ancora '
          'un’operazione (' + OP + ').')
    cella(t, 'Buchi della serie', 3, 'Una transazione fallita annulla anche l’avanzamento del '
          'numeratore.')

    sost(doc, 'Lo stato Evento acquisito e i legami del numero si scrivono con aggiornaNumero '
              'all’aggancio dell’atto;',
         'Lo stato Evento acquisito e l’identificativo nazionale sul numero attendono '
         'un’operazione del servizio (' + OP + ');')

    q = tabella(doc, '#', 'Questione', 'Rimando')
    cella(q, 'NC-4', 1, 'L’utenza con cui il servizio della numerazione accede a '
          'NUMERATORE_COMUNALE, PRENOTAZIONE_ATTI e DATI_ATTO_ESTERNO.')
    cella(q, 'NC-6', 1, 'Chi può segnare come Scartato un numero staccato da un altro operatore: '
          'il servizio non dichiara ruoli e non distingue l’autore della richiesta, che registra '
          'oggi sempre come "admin".')
    c = riga(q, 'NC-8').cells[1]
    D.riscrivi_cella(c, c.text.rstrip() + ' L’operazione stessa manca nel servizio (OP-79).')
    D.clona_riga(q, ('NC-9', 'Le operazioni che le quattro del servizio non coprono: '
                     'registrazione a posteriori, stato 4 Evento acquisito, identificativo '
                     'nazionale sul numero, lettura della serie con i buchi e creazione '
                     'anticipata, operatore reale della richiesta.', 'OP-79'))


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    testa(doc)
    processo(doc)
    pagine(doc)
    tabelle(doc)
    servizio(doc)
    side(doc)
    D.storia(doc, '28/10/2026', '0.5',
             'Riferimenti · Glossario · Attori · Scenari · Le due pagine · Tabelle · Il servizio '
             'be-numeratore-comunale · Che cosa cambia rispetto a Side · Punti aperti',
             'Il servizio si chiama be-numeratore-comunale e le sue operazioni sono le quattro '
             'del capitolo del servizio: percorsi, scenari, pagine e schermate vi si allineano. '
             'Ciò che le quattro non coprono — registrazione a posteriori, Evento acquisito, '
             'identificativo nazionale sul numero, serie con i buchi — è dichiarato mancante '
             '(OP-79 dell’analisi, NC-9). Tolti i residui della serie a sequence; '
             'NUMERATORE_COMUNALE con ts_inserimento e ts_aggiornamento e la chiave primaria.')
    D.frontespizio(doc, '0.5', 'PROCESSO_Numerazione-Comunale_SIPO-ANSC_v0.5', '28/10/2026')
    doc.save(OUT)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
