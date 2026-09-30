# -*- coding: utf-8 -*-
"""DISEGNO_Back-Office_ANSC v0.3 → v0.4 (29/09/2026).

Quattro interventi chiesti dal committente:
  1. wireframe rifatti sul design system di Roma Capitale («Esempio header Footer.png»);
  2. cap. «Dizionari ANSC» allineato a ciò che R901 restituisce davvero;
  3. cap. «Riconciliazione» — colonne ANSC a sinistra, SIPO a destra, «da mappare» sul
     lato SIPO;
  4. la tabella SIPO di riferimento diventa un dato: schema, tabella e campo.

⚠️ Si modifica IL FILE REALE: la v0.3 porta sei commenti di Word che vanno preservati.
Il controllo in coda si interrompe se il conto dei commenti cambia.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v2.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_4_bo_designsystem.py
"""
import os
import shutil
import sys
import zipfile

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(BASE, 'Documenti finali', 'DISEGNO_Back-Office_ANSC_v0.3.docx')
DST = os.path.join(BASE, 'Documenti finali', 'DISEGNO_Back-Office_ANSC_v0.4.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')

FIGURE = [
    ('Figura 1 —', 'bo_mfe.png'), ('Figura 2 —', 'bo_chrome.png'),
    ('Figura 3 —', 'bo_home.png'), ('Figura 4 —', 'bo_atti.png'),
    ('Figura 5 —', 'bo_atto.png'), ('Figura 6 —', 'bo_allegati.png'),
    ('Figura 7 —', 'bo_notifiche.png'), ('Figura 8 —', 'bo_uc_elenco.png'),
    ('Figura 9 —', 'bo_uc_dettaglio.png'), ('Figura 10 —', 'bo_catalogo_uc.png'),
    ('Figura 11 —', 'bo_logiche.png'), ('Figura 12 —', 'bo_dizionari.png'),
    ('Figura 13 —', 'bo_riconciliazione.png'), ('Figura 14 —', 'bo_versioni.png'),
    ('Figura 15 —', 'bo_comandi.png'), ('Figura 16 —', 'bo_numerazione.png'),
    ('Figura 17 —', 'bo_amministrazione.png'),
]


def commenti(percorso):
    return zipfile.ZipFile(percorso).read('word/comments.xml').decode().count('<w:comment ')


def riga_di(tab, prima_cella):
    for r in tab.rows:
        if r.cells[0].text.strip().startswith(prima_cella):
            return r
    raise SystemExit('riga non trovata: ' + prima_cella)


def main():
    atteso = commenti(SRC)
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    fatti = []

    # ───────────────────────────────────── 1. i wireframe, tutti e diciassette
    for didascalia, png in FIGURE:
        D.sostituisci_immagine(d, didascalia, os.path.join(IMG, png))
    fatti.append('17 figure rifatte sul design system di Roma Capitale')

    # ───────────────────────────────── 2. la cornice: testo del capitolo 4
    D.sostituisci(
        d,
        'Header e footer appartengono, in questa ipotesi, alla shell e valgono uguali per '
        'ogni applicazione dell’ente',
        'Header e footer appartengono, in questa ipotesi, alla shell, seguono il design '
        'system di Roma Capitale e valgono uguali per ogni applicazione dell’ente',
        attese=1, etichetta='premessa della cornice', fatti=fatti)

    for p in d.paragraphs:
        if p.text.strip().startswith('Figura 2 —'):
            D.para(d, p._p,
                   'La cornice si compone di tre bande. La **barra di servizio** porta la '
                   'preferenza di layout a sinistra e l’utente a destra, con il menu del '
                   'profilo. La **testata istituzionale** porta il marchio ROMA, lo stemma e '
                   'la dicitura «Roma Capitale». Il **piè di pagina** raccoglie contatti, '
                   'menu, canali e riferimenti legali su fondo scuro. Nessuna delle tre è '
                   'disegnata dal micro-frontend.')
            D.para(d, p._p,
                   '⚠️ **Due elementi che questo documento raccomandava non compaiono nella '
                   'cornice, ed è una scelta.** L’indicazione dell’ambiente e la versione di '
                   'build non sono previste dal design system istituzionale, e il '
                   'committente non intende modificarlo. Restano qui come raccomandazione '
                   'motivata — le ragioni sono nelle tabelle che seguono — e non come '
                   'elemento disegnato, finché non siano concordate con chi governa il '
                   'design system.')
            fatti.append('premessa sulla cornice e sugli elementi non disegnati')
            break

    # ⚠️ nel .docx i marcatori di grassetto sono già stati consumati: si interviene sulle
    # celle, non sul testo.
    for tab in d.tables:
        intest = [c.text.strip() for c in tab.rows[0].cells]
        if intest[:2] == ['Elemento', 'Chi lo fornisce']:
            r = riga_di(tab, 'Ambiente')
            D.riscrivi_cella(r.cells[0], 'Ambiente (raccomandato)')
            D.riscrivi_cella(r.cells[-1],
                             '⚠️ DEV, COLLAUDO o ESERCIZIO. **Oggi non è previsto dal design '
                             'system e non è disegnato**, per scelta del committente. La '
                             'raccomandazione resta, perché è l’elemento che impedisce '
                             'l’errore più costoso: operare in esercizio credendo di essere '
                             'in collaudo. Negli ambienti diversi dall’esercizio dovrebbe '
                             'essere impossibile da non vedere, e il colore da solo non '
                             'basta — non è accessibile.')
            fatti.append('ambiente reso raccomandazione dichiarata')
        if intest == ['Elemento', 'Perché serve']:
            prime = [rr.cells[0].text.strip() for rr in tab.rows]
            if 'Versione e identificativo di build' in prime:
                r = riga_di(tab, 'Versione e identificativo di build')
                D.riscrivi_cella(r.cells[0], 'Versione e build (raccomandati)')
                D.riscrivi_cella(r.cells[1],
                                 '⚠️ **Oggi non previsti dal design system e non disegnati.** '
                                 'Durante un aggiornamento progressivo convivono due versioni '
                                 'dell’applicazione, e un difetto segnalato da un operatore è '
                                 'irriproducibile se non si sa quale stesse usando: è il '
                                 'primo dato che l’assistenza chiede.')
                D.clona_riga(tab, ('Contatti, menu e canali dell’ente',
                                   'Sono le tre colonne del piè di pagina istituzionale: '
                                   'recapiti e riferimenti dell’amministrazione, alcune voci '
                                   'di navigazione e i canali di comunicazione.'))
                fatti.append('versione resa raccomandazione e piè allineato al design system')

    D.sostituisci(
        d,
        'A destra del titolo, mai fra i filtri.',
        'A destra del titolo, mai fra i filtri. Nel design system sono bottoni pieni in '
        'rosso istituzionale.',
        attese=1, etichetta='stile delle azioni di pagina', fatti=fatti)

    # il timer OTP: dove sta oggi e dove andrà. ⚠️ Il paragrafo di chiusura su cui la v0.2
    # si appoggiava è stato rimosso dal committente: si àncora al capitolo successivo.
    for p in d.paragraphs:
        if p.style.name == 'Heading 1' and p.text.strip().startswith('Il registro delle '
                                                                     'pagine'):
            D.para(d, p._p,
                   '⚠️ **Dove sta il distintivo della sessione, e dove andrà.** Oggi è '
                   'accanto al titolo della pagina, cioè dentro la fascia che il '
                   'micro-frontend disegna: è il punto più vicino alle azioni che la '
                   'sessione abilita. È però in corso di predisposizione, a livello di '
                   'shell, **un’area dedicata sotto la testata**: quando sarà disponibile il '
                   'distintivo vi si trasferirà, guadagnando la stessa posizione in tutte le '
                   'applicazioni dell’ente. Il disegno non cambia, cambia chi lo ospita.')
            D.para(d, p._p,
                   'Alla scadenza la sessione **non si rinnova da sola**: le azioni verso '
                   'ANSC si disabilitano e l’ufficiale deve acquisire un nuovo codice sulla '
                   'web app di ANSC e consegnarlo di nuovo. Sotto la soglia di guardia il '
                   'distintivo cambia aspetto e invita a rinnovare **prima** di iniziare una '
                   'sequenza lunga, che è l’unico modo per non trovarsi interrotti a metà.')
            fatti.append('posizione e ciclo di vita del distintivo OTP')
            break

    # ──────────────────────────────── 3. capitolo Dizionari: allineamento a R901
    D.sostituisci(
        d,
        'I dizionari sono la replica locale delle decodifiche pubblicate da ANSC. La pagina '
        'li mostra in due elenchi affiancati — i domini a sinistra, i valori del dominio '
        'scelto a destra — perché è il modo in cui si consultano: si cerca un dominio e se '
        'ne leggono i valori.',
        'I dizionari sono la replica locale delle decodifiche pubblicate da ANSC. La pagina '
        'li mostra in due elenchi affiancati — i domini a sinistra, i valori del dominio '
        'scelto a destra — perché è il modo in cui si consultano: si cerca un dominio e se '
        'ne leggono i valori. ⚠️ **Le colonne mostrate sono esattamente quelle che R901 '
        'restituisce**, e nulla di più: la verifica condotta sul contratto ha fatto emergere '
        'che il modello ne prevedeva una senza fonte e ne ometteva una disponibile.',
        attese=1, etichetta='premessa dei dizionari', fatti=fatti)

    for p in d.paragraphs:
        if p.text.strip().startswith('La pagina serve il back-office e la diagnosi'):
            D.para(d, p._p,
                   '**Che cosa R901 restituisce davvero.** L’operazione di elenco '
                   '(«/config/decodifica/elenco») restituisce per ciascuna tabella **tre '
                   'soli dati**: l’identificativo, il nome della tabella e la **versione**. '
                   'L’operazione di dettaglio («/config/decodifica/dettaglio») non '
                   'restituisce una struttura ma **un CSV codificato in base64 e compresso '
                   'per impostazione predefinita**, con le colonne ID, DESCRIZIONE, '
                   'DATAINIZIOVALIDITA, DATAFINEVALIDITA e ORDINAMENTO. Due sole tabelle su '
                   'centoquarantacinque portano una colonna in più: quella dei casi d’uso e '
                   'quella dei consolati.')
            D.para(d, p._p,
                   '⚠️ **Da qui discendono tre correzioni al modello dati**, riportate '
                   'nella tabella delle strutture e da recepire nell’analisi '
                   'dell’integrazione: si elimina «cd_valore» da VALORE_DOMINIO, perché il '
                   'tracciato ha un solo identificativo e non due; si aggiunge la versione a '
                   'DOMINIO_DECODIFICA, perché è il dato che consente di accorgersi che una '
                   'tabella è cambiata **senza riscaricarla**; si aggiunge un attributo in '
                   'formato JSON a VALORE_DOMINIO per accogliere le due colonne aggiuntive, '
                   'che oggi non avrebbero dove andare.')
            fatti.append('capitolo dizionari allineato a R901')
            break

    # ──────────────────────────── 4. capitolo Riconciliazione: ordine e lato SIPO
    D.sostituisci(
        d,
        'Il raccordo si fa contro le tabelle di configurazione di SIPO — le CONF_* — e la '
        'pagina lo espone con le descrizioni di entrambi i lati, perché un confronto fra '
        'codici nudi non è verificabile da un funzionario.',
        'Il raccordo si fa contro le tabelle di configurazione di SIPO — le CONF_* — e la '
        'pagina lo espone con le descrizioni di entrambi i lati, perché un confronto fra '
        'codici nudi non è verificabile da un funzionario. ⚠️ **L’ordine delle colonne non è '
        'indifferente: prima ANSC, poi SIPO.** Il dominio di ANSC dichiara i propri valori e '
        'sono un dato certo; ciò che è incerto, e che costituisce il lavoro, è la '
        'corrispondenza dal lato di SIPO. Per la stessa ragione **il contrassegno «da '
        'mappare» compare sul valore SIPO**, non su quello di ANSC: è lì che manca qualcosa.',
        attese=1, etichetta='ordine delle colonne della riconciliazione', fatti=fatti)

    for p in d.paragraphs:
        if p.text.strip().startswith('Un valore SIPO senza corrispondenza blocca in '
                                     'preverifica'):
            D.para(d, p._p,
                   '**La tabella SIPO di riferimento diventa un dato.** Nella stesura '
                   'precedente il filtro «Tabella SIPO» non poggiava su alcuna colonna: era '
                   'un criterio che la pagina offriva e che il modello non sapeva '
                   'sostenere. RICONCILIAZ_DIZIONARI acquisisce quindi **SCHEMA_SIPO, '
                   'TABELLA_SIPO e CAMPO_SIPO**, con gli stessi nomi già usati da '
                   'ANSC_CFG_CAMPO. Lo schema serve perché le tabelle di configurazione '
                   'stanno in schemi diversi e in SIPO esistono tabelle omonime in schemi '
                   'distinti; il campo serve a sapere quale dato, in concreto, quella '
                   'corrispondenza traduce.')
            D.para(d, p._p,
                   '⚠️ Il prezzo della scelta va dichiarato: legando la riga anche al campo, '
                   '**la stessa decodifica usata da due campi diversi richiede due righe**. '
                   'È stato accettato perché rende la corrispondenza verificabile sul dato '
                   'reale, e perché è la stessa grana con cui la configurazione dei campi '
                   'già descrive il lato SIPO.')
            fatti.append('schema, tabella e campo sulla riconciliazione')
            break

    # ───────────────────────────────── 5. le tabelle delle due schede
    for tab in d.tables:
        intest = [c.text.strip() for c in tab.rows[0].cells]
        if intest[:2] == ['Tabella', 'Uso']:
            prime = [r.cells[0].text.strip() for r in tab.rows]
            if 'DOMINIO_DECODIFICA' in prime and 'VALORE_DOMINIO' in prime:
                D.riscrivi_cella(riga_di(tab, 'DOMINIO_DECODIFICA').cells[2],
                                 'id_dominio, nm_dominio e cd_versione: l’elenco di '
                                 'sinistra. Sono i tre dati che R901 restituisce '
                                 'nell’operazione di elenco.')
                D.riscrivi_cella(riga_di(tab, 'VALORE_DOMINIO').cells[2],
                                 'id_valore, ds_valore, nr_ordinamento, dt_inizio_validita '
                                 'e dt_fine_validita: le colonne del CSV. ⚠️ «cd_valore» è '
                                 'eliminata: il tracciato ha un solo identificativo.')
                fatti.append('tabelle sottese dei dizionari corrette')
            if 'RICONCILIAZ_DIZIONARI' in prime:
                D.riscrivi_cella(riga_di(tab, 'RICONCILIAZ_DIZIONARI').cells[2],
                                 'DECODIFICA, VALORE_ANSC, DESCRIZIONE_ANSC, VALORE_SIPO, '
                                 'DESCRIZIONE_SIPO, SCHEMA_SIPO, TABELLA_SIPO, CAMPO_SIPO, '
                                 'CONDIZIONE, DATA_INIZIO_VALIDITA e DATA_FINE_VALIDITA. '
                                 'Filtrata per ID_VERSIONE.')
                fatti.append('tabelle sottese della riconciliazione corrette')
        if intest[:2] == ['Elemento', 'Genere']:
            prime = [r.cells[0].text.strip() for r in tab.rows]
            if 'SCARICA DA ANSC (R901)' in prime:
                D.riscrivi_cella(riga_di(tab, 'SCARICA DA ANSC (R901)').cells[2],
                                 'Ricarica elenco e valori. È un comando manuale e non '
                                 'pianificato. ⚠️ Il dettaglio arriva come CSV in base64 e '
                                 'compresso: la decompressione e la lettura del tracciato '
                                 'sono a carico del componente dizionari.')
                D.clona_riga(tab, ('Versione', 'Colonna',
                                   'La versione della tabella dichiarata da ANSC. ⚠️ È il '
                                   'solo modo per accorgersi che una decodifica è cambiata '
                                   'senza riscaricarla e confrontarla riga per riga.'))
                fatti.append('azioni dei dizionari aggiornate')
            if 'Decodifica ANSC' in prime and 'Valore ANSC' in prime:
                D.riscrivi_cella(riga_di(tab, 'Tabella SIPO').cells[0],
                                 'Schema e tabella SIPO')
                D.riscrivi_cella(riga_di(tab, 'Schema e tabella SIPO').cells[2],
                                 'Poggia sulle colonne SCHEMA_SIPO e TABELLA_SIPO, '
                                 'introdotte in questa versione. Lo schema è necessario '
                                 'perché in SIPO esistono tabelle omonime in schemi diversi.')
                D.riscrivi_cella(riga_di(tab, 'Valore ANSC').cells[2],
                                 'Colonna di sola lettura: è il valore che il dominio '
                                 'dichiara. Non si sceglie e non si corregge da qui.')
                D.clona_riga(tab, ('Valore SIPO', 'Campo',
                                   'Tendina dei valori della tabella di configurazione '
                                   'indicata. ⚠️ Una riga senza valore SIPO porta il '
                                   'contrassegno «da mappare»: è il lato incerto della '
                                   'corrispondenza, ed è il lavoro della pagina.'))
                D.clona_riga(tab, ('Campo SIPO', 'Colonna e campo',
                                   'Il campo che in concreto contiene il valore da '
                                   'tradurre. ⚠️ Legando la riga al campo, la stessa '
                                   'decodifica usata da due campi diversi richiede due '
                                   'righe: è il prezzo dichiarato della precisione.'))
                fatti.append('campi e azioni della riconciliazione aggiornati')

    # ───────────────────────────────── 6. la copertura e i punti aperti
    for tab in d.tables:
        if [c.text.strip() for c in tab.rows[0].cells][:2] == ['Tabella',
                                                               'Pagine che la governano']:
            D.riscrivi_cella(riga_di(tab, 'VALORE_DOMINIO').cells[2],
                             'Sola lettura: si aggiorna da ANSC. ⚠️ In questa versione perde '
                             'cd_valore e acquisisce un attributo JSON per le due colonne '
                             'aggiuntive di consolati e casi d’uso.')
            D.riscrivi_cella(riga_di(tab, 'DOMINIO_DECODIFICA').cells[2],
                             'Sola lettura: si aggiorna da ANSC. ⚠️ In questa versione '
                             'acquisisce la versione della tabella, che R901 restituisce e '
                             'che il modello non conservava.')
            D.riscrivi_cella(riga_di(tab, 'RICONCILIAZ_DIZIONARI').cells[2],
                             'Completa. ⚠️ In questa versione acquisisce SCHEMA_SIPO, '
                             'TABELLA_SIPO e CAMPO_SIPO.')
            fatti.append('copertura aggiornata sulle tre tabelle modificate')
            break

    t = D.trova_tabella(d, '#', 'questione')
    for voce in (
        ('BO-13', 'Recepire nell’analisi le tre correzioni al modello dei dizionari',
         'L’eliminazione di «cd_valore», l’aggiunta della versione su DOMINIO_DECODIFICA e '
         'l’attributo JSON per le due colonne aggiuntive nascono da questo documento ma '
         'appartengono al modello dati dell’integrazione, dove le DDL sono definite. Vanno '
         'riportate là, altrimenti i due documenti divergono.'),
        ('BO-14', 'Recepire nell’analisi le tre colonne nuove della riconciliazione',
         'SCHEMA_SIPO, TABELLA_SIPO e CAMPO_SIPO su RICONCILIAZ_DIZIONARI valgono lo stesso '
         'rilievo: la struttura è dichiarata nell’analisi dell’integrazione.'),
        ('BO-15', 'Quando l’area dedicata della shell sarà disponibile',
         'Il distintivo della sessione sta oggi accanto al titolo della pagina. La shell sta '
         'predisponendo un’area dedicata sotto la testata: quando sarà pronta il distintivo '
         'vi si trasferirà, e con esso ogni altra informazione di contesto che oggi il '
         'micro-frontend disegna per conto proprio.'),
        ('BO-16', 'Se l’ambiente e la versione entreranno nella cornice istituzionale',
         '⚠️ Il committente non intende modificare l’interfaccia. Il disegno li mantiene come '
         'raccomandazione motivata: senza l’ambiente resta possibile operare in esercizio '
         'credendo di essere in collaudo, e senza la versione una segnalazione di difetto '
         'non è riconducibile al rilascio che l’ha prodotta.'),
    ):
        D.clona_riga(t, voce)
    fatti.append('quattro punti aperti nuovi (BO-13…BO-16)')

    # ───────────────────────────────────────────── 7. testata e storia
    for tab in d.tables:
        if tab.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in tab.rows:
                v = {'Data consegna': '29/09/2026', 'Versione': '0.4',
                     'Documento': 'DISEGNO_Back-Office_ANSC_v0.4'}.get(r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    D.storia(d, '29/09/2026', '0.4',
             'Cornice · Dizionari · Riconciliazione · Copertura · Punti aperti',
             'Wireframe rifatti sul design system di Roma Capitale: barra di servizio, '
             'testata istituzionale, tabelle con intestazione in rosso, impaginazione e piè '
             'di pagina a tre colonne. Ambiente e versione di build restano raccomandazioni '
             'e non sono disegnati. Il capitolo sui dizionari è allineato a ciò che R901 '
             'restituisce davvero — elenco con identificativo, nome e versione; dettaglio '
             'come CSV in base64 compresso — con tre correzioni al modello. Nel capitolo '
             'sulla riconciliazione le colonne di ANSC precedono quelle di SIPO e il '
             'contrassegno «da mappare» passa al lato SIPO; la tabella di riferimento '
             'diventa un dato con schema, tabella e campo.')
    fatti.append('testata e storia aggiornate')

    # ⚠️ `para` e `riscrivi_cella` non interpretano «**…**»: senza questo passaggio i
    # marcatori finirebbero visibili nel documento consegnato. I paragrafi che portano
    # un'ancora di commento non si toccano, perché ricostruire i run la distruggerebbe.
    def applica_grassetti(par):
        if '**' not in par.text or D.ha_commenti(par._p):
            return 0
        pezzi = D.segmenta('', par.text)
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        for testo_pezzo, gr in pezzi:
            run = par.add_run(testo_pezzo)
            run.bold = gr
        return 1
    n = sum(applica_grassetti(p) for p in d.paragraphs)
    for tab in d.tables:
        for r in tab.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    n += applica_grassetti(p)
    fatti.append(f'{n} paragrafi con grassetto applicato')

    d.save(DST)
    trovati = commenti(DST)
    assert trovati == atteso, f'commenti persi: {trovati} invece di {atteso}'
    print('\n'.join(' · ' + f for f in fatti))
    print('commenti preservati:', trovati)
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))
    print('scritto:', os.path.relpath(DST, BASE))


if __name__ == '__main__':
    main()
