# -*- coding: utf-8 -*-
"""Diagramma funzionale del capitolo «Il processo di gestione dei dati ANSC in SIPO», in
PowerPoint con forme e connettori NATIVI, così che l'utente lo possa modificare (richiesta del
15/10/2026). Dal .pptx PowerPoint stesso esporta il PNG che va nel .docx: la figura del documento
e il file modificabile sono la stessa cosa.

Fonti, tutte nel capitolo della v3.41:
  * la tabella del workflow (Passaggio | Immagine di riferimento | STATO ANSC | Stato SIPO);
  * la «Tabella degli stati dell'atto ANSC» (Momento | Stato ANSC / ANSC_11 | Generato da);
  * «Gli stati di SIPO» (1 In lavorazione, 2 Definito);
  * il testo di «Processo ANSC di validazione dell'atto» (KO di R009 → si torna alle pagine
    precedenti; senza firmatari la pagina delle firme si salta) e «Il passo Definisci».

Quattro colonne = quattro corsie: pagina di SIPO · servizio chiamato · stato in ANSC · stato in
SIPO. Le righe sono i passaggi, dall'alto in basso.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/diagramma_workflow_pptx.py
"""
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

QUI = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(QUI, 'img', 'workflow_ansc_sipo.pptx')

PAL = {   # bordo, fondo — la palette di diagrammi_comune.C
    'grigio': ('6b7280', 'e9eaec'), 'blu': ('2f6bb0', 'dbe6f5'), 'rosso': ('b03a48', 'f6dfe1'),
    'giallo': ('d9a441', 'fdf1d6'), 'verde': ('3f8f5f', 'dcefe1'), 'viola': ('7b5aa6', 'ece5f5'),
    'bianco': ('9ca3af', 'ffffff'),
}
INK, MUTED, ARROW = '1f2937', '6b7280', '4b5563'

W, H = 7.5, 10.2                       # pagina verticale, per entrare nella larghezza del .docx
COL = {'A': (0.25, 2.55), 'B': (2.92, 1.50), 'C': (4.54, 1.45), 'D': (6.23, 1.02)}
Y0, RH, GAP = 1.30, 0.56, 0.11         # prima riga, altezza riga, intervallo

# ─────────────────────────────────────────────────────────── le righe, dalla tabella del workflow
PASSI = [
    # (pagina di SIPO: titolo, sottotitolo) · servizio
    ('Accesso e scelte iniziali', 'Login · tipologia di gestione · tipologia di atto (invariati)', ''),
    ('Dati anagrafici', 'soggetti e intestatario, dall’UC (SOGGETTI_USECASE)', 'R005 · opzione intestatario'),
    ('Riepilogo dei soggetti', 'i tre codici per soggetto; non si prosegue se mancano dati obbligatori',
     'R005 · per ciascun soggetto'),
    ('Riepilogo', 'UC (ANSC_03), allegati, firmatari, dati mancanti', 'preverifica'),
    ('Anteprima · numero comunale', 'il numero definitivo sostituisce il provvisorio ed è mostrato',
     'num-ansc-sipo · creaPrenotazione'),
    ('Anteprima · Valida', 'validazione di metadati e allegati; si salvano gli idAnsc',
     'R009 /validazione/evento'),
    ('Anteprima · firme dei dichiaranti', 'firma digitale o scansione del cartaceo firmato',
     'R006 cartacea · R012 elettronica'),
    ('Anteprima · conformità', 'attestazione di conformità firmata dall’USC',
     'R007 /attestazione/firma'),
    ('Anteprima · firma dell’USC', 'utenza, PIN e OTP del servizio di firma remota',
     'R007 /firma_usc'),
    ('Definisci', 'pagina finale: salvataggio in SIPO, non si torna indietro', 'salvataggio in SIPO'),
    ('Chiudi', 'l’atto non si modifica più, nemmeno fino alle 24', ''),
]


def y(i):
    return Y0 + i * (RH + GAP)


def rgb(h):
    return RGBColor.from_string(h)


def riquadro(sl, x, y_, w, h, colore, righe, forma=MSO_SHAPE.ROUNDED_RECTANGLE, tratteggio=False,
             allinea=PP_ALIGN.CENTER):
    """righe = [(testo, pt, grassetto, colore)]"""
    s = sl.shapes.add_shape(forma, Inches(x), Inches(y_), Inches(w), Inches(h))
    bordo, fondo = PAL[colore]
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fondo)
    s.line.color.rgb = rgb(bordo)
    s.line.width = Pt(1.1)
    if tratteggio:
        s.line.dash_style = 4                # MSO_LINE.DASH
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(0.05))
    for m in ('margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(0.02))
    for k, (testo, pt, bold, col) in enumerate(righe):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = allinea
        r = p.add_run()
        r.text = testo
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.name = 'Arial'
        r.font.color.rgb = rgb(col or INK)
    return s


def testo(sl, x, y_, w, h, t, pt=8, bold=False, col=MUTED, allinea=PP_ALIGN.LEFT, corsivo=False):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y_), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Inches(0.01))
    p = tf.paragraphs[0]
    p.alignment = allinea
    r = p.add_run()
    r.text = t
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.italic = corsivo
    r.font.name = 'Arial'
    r.font.color.rgb = rgb(col)
    return tb


def verticale(sl, cx, cy, t, col, pt=6, lungh=1.8):
    """Etichetta ruotata di 270°, centrata in (cx, cy): la casella si crea orizzontale e si ruota
    attorno al proprio centro, quindi la si colloca a partire dal centro."""
    h = 0.11
    tb = testo(sl, cx - lungh / 2, cy - h / 2, lungh, h, t, pt, False, col, PP_ALIGN.CENTER)
    tb.rotation = 270
    return tb


def _punta(linea, inizio=False):
    ln = linea._element.spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(linea._element.spPr, qn('a:ln'))
    tag = 'a:headEnd' if inizio else 'a:tailEnd'
    e = etree.SubElement(ln, qn(tag))
    e.set('type', 'triangle')
    e.set('w', 'med')
    e.set('len', 'med')


def freccia(sl, a, ca, b, cb, colore=ARROW, tratteggio=False, spessore=1.25):
    """Connettore diritto agganciato alle forme: segue i riquadri se l'utente li sposta.
    ca/cb = punto di aggancio (0 alto, 1 sinistra, 2 basso, 3 destra)."""
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(a, ca)
    c.end_connect(b, cb)
    c.line.color.rgb = rgb(colore)
    c.line.width = Pt(spessore)
    if tratteggio:
        c.line.dash_style = 4
    _punta(c)
    return c


def gomito(sl, punti, colore=ARROW, tratteggio=False, spessore=1.1):
    """Spezzata libera (forma a mano libera) con la punta in fondo: per ritorni e salti."""
    fb = sl.shapes.build_freeform(Inches(punti[0][0]), Inches(punti[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in punti[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = rgb(colore)
    s.line.width = Pt(spessore)
    if tratteggio:
        s.line.dash_style = 4
    s.shadow.inherit = False
    _punta(s)
    return s


def linea(sl, x1, y1, x2, y2, colore, tratteggio=True, spessore=1.5):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(colore)
    c.line.width = Pt(spessore)
    if tratteggio:
        c.line.dash_style = 4
    return c


def genera(percorso=OUT):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb('ffffff')

    testo(sl, 0.25, 0.18, 7.0, 0.35, 'Il processo di gestione dei dati ANSC in SIPO', 15, True,
          '1f3a5f', PP_ALIGN.CENTER)
    testo(sl, 0.25, 0.52, 7.0, 0.3, 'Le pagine di SIPO, i servizi che chiamano e i cambi di stato '
          'dell’atto in ANSC (ANSC_11) e in SIPO (ID_STATO_PRATICA)', 8.5, False, MUTED,
          PP_ALIGN.CENTER, corsivo=True)

    intest = {'A': 'Pagina di SIPO', 'B': 'Servizio chiamato', 'C': 'Stato in ANSC',
              'D': 'Stato in SIPO'}
    for k, (x, w) in COL.items():
        testo(sl, x, 0.93, w, 0.25, intest[k], 9, True, INK, PP_ALIGN.CENTER)
        linea(sl, x, 1.19, x + w, 1.19, 'cbd5e1', tratteggio=False, spessore=1)

    # ── colonna A (pagine) e B (servizi)
    pag, srv = [], []
    for i, (titolo, sotto, servizio) in enumerate(PASSI):
        colore = 'grigio' if i == 0 else 'blu'
        pag.append(riquadro(sl, COL['A'][0], y(i), COL['A'][1], RH, colore,
                            [(titolo, 9, True, None), (sotto, 7, False, MUTED)]))
        if servizio:
            srv.append(riquadro(sl, COL['B'][0], y(i) + 0.09, COL['B'][1], RH - 0.18, 'bianco',
                                [(servizio, 7.5, False, None)], forma=MSO_SHAPE.RECTANGLE))
            freccia(sl, pag[i], 3, srv[-1], 1, colore='94a3b8', spessore=0.9)
        else:
            srv.append(None)
    for i in range(len(pag) - 1):
        freccia(sl, pag[i], 2, pag[i + 1], 0)

    # ── colonna C: gli stati dell'atto in ANSC
    xc, wc = COL['C']
    nessuno = riquadro(sl, xc, y(0), wc, y(4) + RH - y(0), 'grigio',
                       [('Nessun atto in ANSC', 9, True, None),
                        ('fino alla validazione', 7, False, MUTED)], tratteggio=True)
    stati = {}
    for i, (cod, nome) in zip((5, 6, 7, 8), (('2', 'CONFERMATO'), ('4', 'FIRMATO DAL DICHIARANTE'),
                                             ('11', 'FIRMATO PER CONFORMITÀ'), ('5', 'FIRMATO DA USC'))):
        h = RH if i < 8 else y(10) + RH - y(8)
        stati[cod] = riquadro(sl, xc, y(i), wc, h, 'giallo',
                              [(f'{cod} {nome}', 8.5, True, None)])
    freccia(sl, nessuno, 2, stati['2'], 0)
    freccia(sl, stati['2'], 2, stati['4'], 0)
    freccia(sl, stati['4'], 2, stati['11'], 0)
    freccia(sl, stati['11'], 2, stati['5'], 0)
    testo(sl, xc + wc * 0.53, y(4) + RH + 0.008, 0.8, 0.1, 'R009 esito positivo', 6, False, MUTED)
    # salto delle firme quando non ci sono firmatari: 2 → 11 sul fianco destro della colonna
    xr = xc + wc + 0.07
    gomito(sl, [(xc + wc, y(5) + RH / 2), (xr, y(5) + RH / 2), (xr, y(7) + RH / 2),
                (xc + wc, y(7) + RH / 2)], tratteggio=True)
    verticale(sl, xr + 0.085, y(6) + RH / 2, 'senza firmatari', MUTED)

    # ritorno sul KO di R009: dal passo Valida alle pagine dei dati (sul fianco sinistro)
    xl = COL['A'][0] - 0.12
    gomito(sl, [(COL['A'][0], y(5) + RH / 2), (xl, y(5) + RH / 2), (xl, y(1) + RH / 2),
                (COL['A'][0], y(1) + RH / 2)], colore='b03a48', tratteggio=True)
    verticale(sl, xl - 0.06, (y(1) + y(5)) / 2 + RH / 2, 'R009 KO: si corregge e si ripete',
              'b03a48')

    # ── colonna D: lo stato in SIPO
    xd, wd = COL['D']
    na = riquadro(sl, xd, y(0), wd, RH, 'grigio', [('N.A.', 8.5, True, MUTED)], tratteggio=True)
    lav = riquadro(sl, xd, y(1), wd, y(8) + RH - y(1), 'verde',
                   [('1', 11, True, None), ('IN LAVORAZIONE', 8, True, None),
                    ('nasce col numero provvisorio', 6.5, False, MUTED)])
    defi = riquadro(sl, xd, y(9), wd, y(10) + RH - y(9), 'verde',
                    [('2', 11, True, None), ('DEFINITO', 8, True, None)])
    freccia(sl, na, 2, lav, 0)
    freccia(sl, lav, 2, defi, 0)
    testo(sl, xd + wd * 0.55, y(8) + RH + 0.008, 0.6, 0.1, 'Definisci', 6, False, MUTED)

    # ── punto di non ritorno, dopo la firma dell'USC
    yn = y(9) - GAP / 2
    linea(sl, 0.2, yn, xc - 0.05, yn, 'b03a48')
    testo(sl, COL['B'][0], yn - 0.16, 1.6, 0.15, 'punto di non ritorno', 6.5, True, 'b03a48')

    # ── dopo la formazione: la tabella degli stati
    yb = y(10) + RH + 0.30
    dopo = riquadro(sl, 2.92, yb, 4.33, 0.86, 'viola',
                    [('Dopo la formazione', 8.5, True, None),
                     ('9 CANCELLATO — R011', 7.5, False, None),
                     ('3 ANNULLATO · 10 INEFFICACE — rettifica di tipo annullamento; annullamento '
                      'per inefficacia', 7.5, False, None)])
    freccia(sl, stati['5'], 2, dopo, 0, tratteggio=True)

    # ── legenda
    yl = yb
    voci = [('blu', 'pagina di SIPO'), ('bianco', 'servizio chiamato'), ('giallo', 'stato ANSC (ANSC_11)'),
            ('verde', 'stato SIPO'), ('viola', 'stati successivi')]
    for k, (col, t) in enumerate(voci):
        riquadro(sl, 0.25, yl + k * 0.17, 0.22, 0.12, col, [('', 6, False, None)],
                 forma=MSO_SHAPE.RECTANGLE)
        testo(sl, 0.52, yl + k * 0.17 - 0.015, 2.2, 0.15, t, 7)
    testo(sl, 0.25, yl + 0.88, 2.6, 0.3, 'Tratteggio: ritorno, salto o passaggio successivo alla '
          'formazione', 6.5, corsivo=True)

    os.makedirs(os.path.dirname(percorso), exist_ok=True)
    prs.save(percorso)
    return percorso


if __name__ == '__main__':
    print('scritto:', genera(sys.argv[1] if len(sys.argv) > 1 else OUT))
