# -*- coding: utf-8 -*-
"""v3.46 dell'analisi ANSC, dalla v3.45 (23/10/2026). Sette richieste dell'utente:

  1. ANSC_ANA_UC si unifica in ANSC_CFG_UC, che ne prende le colonne (codice motore,
     denominazione ANSC, famiglia, tipo contenuto). La riga è per UC e per Modello dentro una
     baseline: sezioni e formule richiamano l'UC per valore (COD_UC_ANSC + ID_VERSIONE), come
     già ALLEGATI_USECASE e SOGGETTI_USECASE. Il catalogo entra così nella baseline.
  2. ANSC_CFG_CAMPO riceve ID_VERSIONE e COD_UC_ANSC (commento di Word «Aggiungere
     uc+versione»), vincolati a quelli della sezione da una chiave esterna composta.
  3. ANSC_CFG_VERSIONE.COD_VERSIONE si elimina: nessun processo lo usa, e numero, revisione
     ANSC recepita e descrizione già dicono ciò che un nome libero ripeterebbe.
  4-6. comunicaz_anagrafica_usecase, forzatura_validazione e firmatari in sintassi Oracle.
  7. La tabella evento di Side («Ansc-evento») si toglie: resta ANSC_STATO_ATTO. Il confronto
     colonna per colonna è in CONFRONTO_evento-ANSC_STATO_ATTO_v0.1.xlsx
     (strumenti/confronto_evento_stato_atto.py).

I commenti di Word sulle righe toccate si spostano sulle righe che le sostituiscono.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_46_uc_ddl_oracle.py
"""
import os
import re
import subprocess
import sys

import docx

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.45.docx')
OUT = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.46.docx')
PY = sys.executable

fatti = []


# ------------------------------------------------------------------ primitive
def indice(doc, condizione, cosa, da=0):
    ps = doc.paragraphs
    trovati = [i for i in range(da, len(ps)) if condizione(ps[i])]
    if not trovati:
        raise SystemExit('non trovato: ' + cosa)
    return trovati[0]


def esatto(testo):
    return lambda p: p.text == testo


def inizia(testo):
    return lambda p: p.text.startswith(testo)


def titolo(livello, testo):
    return lambda p: p.style.name == f'Heading {livello}' and p.text.strip() == testo


def elimina(elementi):
    for e in elementi:
        if D.ha_commenti(e):
            raise SystemExit('elemento con commento, non lo elimino: ' +
                             ''.join(e.itertext())[:60])
    for e in elementi:
        e.getparent().remove(e)


def blocco(doc, i_da, i_a, righe, sposta=()):
    """Sostituisce i paragrafi [i_da, i_a) con le righe di DDL date.

    «sposta» = coppie (inizio della riga vecchia, inizio della riga nuova): le ancore dei
    commenti di Word passano dalla prima alla seconda prima che la vecchia sia eliminata.
    """
    ps = doc.paragraphs
    vecchi = [p for p in ps[i_da:i_a]]
    prima = vecchi[0]._p
    nuovi = [D.para(doc, prima, r, mono=True) for r in righe]
    for v_inizio, n_inizio in sposta:
        v = [p for p in vecchi if p.text.startswith(v_inizio)]
        n = [p for p in nuovi if p.text.startswith(n_inizio)]
        if len(v) != 1 or len(n) != 1:
            raise SystemExit(f'spostamento commenti: {v_inizio!r} {len(v)} / {n_inizio!r} {len(n)}')
        D.sposta_commenti(v[0]._p, n[0]._p)
    elimina([p._p for p in vecchi])
    return nuovi


def riga_ddl(doc, da, vecchia, nuove):
    """Sostituisce UNA riga di DDL (cercata dall'indice «da» in avanti) con una o più righe."""
    i = indice(doc, esatto(vecchia), vecchia, da)
    blocco(doc, i, i + 1, nuove if isinstance(nuove, list) else [nuove])
    return i


def cella(doc, prima_colonna, colonna_testo, nuovo, tabella=None):
    t = tabella or D.tabella_colonne(doc, prima_colonna)
    righe = [r for r in t.rows if r.cells[0].text.strip() == prima_colonna]
    if len(righe) != 1:
        raise SystemExit(f'riga {prima_colonna}: {len(righe)}')
    D.riscrivi_cella(righe[0].cells[colonna_testo], nuovo)
    return righe[0]


def S(vecchio, nuovo, attese=1):
    D.sostituisci(DOC, vecchio, nuovo, attese=attese, fatti=fatti)


def pezzi(p, testo):
    """Riscrive un paragrafo senza commenti conservando il carattere del primo run."""
    if D.ha_commenti(p._p):
        raise SystemExit('paragrafo con commento: ' + p.text[:60])
    D.testo_di(p, testo)


# ------------------------------------------------------------------ 1. ANSC_CFG_UC
CFG_UC_DESCR = [
    ('ID_UC_CFG', 'NUMBER (PK)', 'Chiave tecnica.'),
    ('COD_UC_ANSC', 'VARCHAR2(20)',
     'Codice numerico dell’UC pubblicato da ANSC. La lunghezza non è fissa: 11111000 per la '
     'dichiarazione di nascita, 2101 per la morte in abitazione — nel catalogo compaiono codici '
     'di quattro, cinque, sei e otto cifre. La notazione puntata (2.1.0.1) usata in altri '
     'documenti è redazionale e non va scritta qui. Insieme a ID_VERSIONE è il riferimento con '
     'cui sezioni, campi, formule, allegati e soggetti richiamano l’UC, per valore.'),
    ('COD_MOTORE', 'VARCHAR2(40)',
     'Codice parlante del motore (es. Dic_Nasc_001), dal catalogo di ANSC: attributo dell’UC, '
     'non sinonimo.'),
    ('DESC_UC_ANSC', 'VARCHAR2(400)',
     'Denominazione dell’UC pubblicata da ANSC (decodifica ANSC_03): è la descrizioneCasoUso '
     'del modello evento.'),
    ('COD_FAMIGLIA', 'VARCHAR2(30)', 'Famiglia di evento: nascita, morte, matrimonio, …'),
    ('ID_TIPO_EVENTO', 'VARCHAR2(10)',
     'Tipo evento secondo la decodifica ANSC_01 (1 nascita, 2 morte, 3 matrimonio, 4 unione '
     'civile, 5 cittadinanza), trasmesso nella testata del modello evento. Coincide con la prima '
     'cifra del codice dell’UC; si conserva perché il payload lo richiede e perché una regola '
     'implicita nel codice è una regola che nessuno verifica.'),
    ('ID_TIPO_CONTENUTO', 'VARCHAR2(10)',
     'Tipo di contenuto secondo la decodifica ANSC_02 (atto, annotazione, trascrizione…), '
     'pubblicato da ANSC per ciascun UC nella colonna IDTIPOCONTENUTO del catalogo dei casi '
     'd’uso: è la replica che la testata del modello evento richiede.'),
    ('DESCRIZIONE', 'VARCHAR2(200)',
     'Descrizione del Modello di atto come la nomina il Comune: «MORTE IN ABITAZIONE A ROMA». '
     'Serve a chi configura, che ragiona per Modelli e non per codici.'),
    ('ID_TIPO_DOCUMENTO', 'VARCHAR2(20)',
     'CREAZIONE, TRASCRIZIONE, … Nel catalogo che la tabella sostituisce la colonna nominava il '
     'tipo di documento del canale DMNM/TS (R022/R023), che è fuori perimetro: resta con il '
     'significato della configurazione.'),
    ('ID_MODELLO_ATTO', 'VARCHAR2(5)',
     'Modello di atto del Comune (CONF_TIPO_ATTI.ID_MODELLO_ATTO). Non è la chiave della riga: '
     'più UC condividono lo stesso Modello, ed è la ragione per cui la riga è per UC. Vuota per '
     'un UC del catalogo che nessun Modello raggiunge ancora.'),
    ('ID_CONF_TIPO_ATTO', 'NUMBER',
     'Tipo atto SIPO da cui discendono Modello e maschera (CONF_TIPO_ATTI). Vuota come il '
     'Modello.'),
    ('MASCHERA_UI', 'VARCHAR2(100)',
     'Maschera di immissione corrispondente. Riportata per leggibilità: la sorgente '
     'autoritativa resta CONF_TIPO_ATTI.'),
    ('STATO', 'VARCHAR2(1)',
     '[DA VERIFICARE: i valori ammessi. Un solo carattere; la sorgente dichiara la colonna ma '
     'non la valorizza. Da non confondere con lo stato dell’atto: qui si tratta della riga di '
     'configurazione.]'),
    ('NUM_PRIORITA', 'NUMBER',
     'Ordine di valutazione fra gli UC che condividono il Modello: il primo la cui logica è '
     'soddisfatta determina l’UC.'),
    ('ID_DOMINIO / COD_LOGICA_DI_SCELTA', 'NUMBER / VARCHAR2(30) (FK)',
     'La logica che stabilisce se l’UC si applica all’atto in lavorazione, presa dal catalogo '
     '(dominio 1). Vuota significa «si applica sempre»: è il caso dei Modelli che indirizzano '
     'un solo UC.'),
    ('SERIE', 'VARCHAR2(10)', 'Porzione di archivio in SIPO.'),
    ('ID_VERSIONE', 'NUMBER (FK)', 'Baseline di appartenenza (ANSC_CFG_VERSIONE).'),
    ('DATA_INIZIO_VALIDITA / DATA_FINE_VALIDITA', 'DATE',
     'Validità dell’UC come la dichiara ANSC. ANSC non cancella gli UC: ne chiude la validità — '
     'sono i tre casi ritirati, 374 validi su 377 righe del catalogo.'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)',
     'Campi tecnici previsti dallo standard di nomenclatura [R5].'),
]

CFG_UC_DDL = """CREATE TABLE ANSC_USR.ANSC_CFG_UC (
  ID_UC_CFG            NUMBER GENERATED ALWAYS AS IDENTITY,
  COD_UC_ANSC          VARCHAR2(20 CHAR)  NOT NULL,
  COD_MOTORE           VARCHAR2(40 CHAR)  NOT NULL,
  DESC_UC_ANSC         VARCHAR2(400 CHAR),
  COD_FAMIGLIA         VARCHAR2(30 CHAR)  NOT NULL,
  ID_TIPO_EVENTO       VARCHAR2(10 CHAR)  NOT NULL,
  ID_TIPO_CONTENUTO    VARCHAR2(10 CHAR),
  DESCRIZIONE          VARCHAR2(200 CHAR),
  ID_TIPO_DOCUMENTO    VARCHAR2(20 CHAR),
  ID_MODELLO_ATTO      VARCHAR2(5 CHAR),
  ID_CONF_TIPO_ATTO    NUMBER,
  STATO                VARCHAR2(1 CHAR),
  MASCHERA_UI          VARCHAR2(100 CHAR),
  SERIE                VARCHAR2(10 CHAR),
  NUM_PRIORITA         NUMBER DEFAULT 100 NOT NULL,
  ID_DOMINIO           NUMBER DEFAULT 1,
  COD_LOGICA_DI_SCELTA VARCHAR2(30 CHAR),
  ID_VERSIONE          NUMBER             NOT NULL,
  DATA_INIZIO_VALIDITA DATE DEFAULT SYSDATE            NOT NULL,
  DATA_FINE_VALIDITA   DATE DEFAULT DATE '9999-12-31'  NOT NULL,
  DATA_INS             TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,
  DATA_UPD             TIMESTAMP DEFAULT SYSTIMESTAMP NOT NULL,
  UTENTE_INS           VARCHAR2(40 CHAR),
  UTENTE_UPD           VARCHAR2(40 CHAR),
  CONSTRAINT PK_ANSC_CFG_UC PRIMARY KEY (ID_UC_CFG),
  CONSTRAINT UK_ANSC_CFG_UC
    UNIQUE (ID_VERSIONE, COD_UC_ANSC, ID_CONF_TIPO_ATTO, ID_MODELLO_ATTO),
  CONSTRAINT FK_ANSC_CFG_UC_VERSIONE
    FOREIGN KEY (ID_VERSIONE) REFERENCES ANSC_USR.ANSC_CFG_VERSIONE (ID_VERSIONE),
  CONSTRAINT FK_ANSC_CFG_UC_LOGICA
    FOREIGN KEY (ID_DOMINIO, COD_LOGICA_DI_SCELTA)
    REFERENCES ANSC_USR.LOGICHE_DI_SCELTA (ID_DOMINIO, COD_LOGICA_DI_SCELTA)
) TABLESPACE ANSC_USR;
CREATE INDEX ANSC_USR.IX_ANSC_CFG_UC_MODELLO
  ON ANSC_USR.ANSC_CFG_UC (ID_VERSIONE, ID_CONF_TIPO_ATTO, ID_MODELLO_ATTO,
                           NUM_PRIORITA) TABLESPACE ANSC_USR;
CREATE INDEX ANSC_USR.IX_ANSC_CFG_UC_UC
  ON ANSC_USR.ANSC_CFG_UC (ID_VERSIONE, COD_UC_ANSC) TABLESPACE ANSC_USR;
COMMENT ON TABLE ANSC_USR.ANSC_CFG_UC IS
  'Catalogo degli UC pubblicato da ANSC e loro adozione da parte del Comune: una riga per
   UC e per Modello che lo raggiunge, dentro una baseline. Senza Modello: UC del catalogo
   che nessun Modello raggiunge ancora.';
COMMENT ON COLUMN ANSC_USR.ANSC_CFG_UC.COD_MOTORE IS
  'Identificativo del mapping ufficiale (es. Dic_Nasc_001): attributo dell''UC, non sinonimo.';
COMMENT ON COLUMN ANSC_USR.ANSC_CFG_UC.COD_LOGICA_DI_SCELTA IS
  'Logica di scelta presa dal catalogo (dominio 1). Vuota significa: si applica sempre.';
COMMENT ON COLUMN ANSC_USR.ANSC_CFG_UC.NUM_PRIORITA IS
  'Ordine di valutazione fra gli UC dello stesso Modello: il primo la cui logica e
   soddisfatta determina l UC. Due righe con la stessa priorita sono una ambiguita.';""".split('\n')


def unifica_uc(doc):
    # --- descrizione: la riga commentata del catalogo passa sulla riga gemella di ANSC_CFG_UC
    t_ana = D.tabella_colonne(doc, 'COD_MOTORE')
    t_cfg = D.tabella_colonne(doc, 'ID_UC_CFG')
    D.riscrivi_sicura(t_cfg, CFG_UC_DESCR)
    src = [r for r in t_ana.rows if r.cells[0].text.strip() == 'ID_TIPO_DOCUMENTO'][0]
    dst = [r for r in t_cfg.rows if r.cells[0].text.strip() == 'ID_TIPO_DOCUMENTO'][0]
    for j, c in enumerate(src.cells):
        for p in c.paragraphs:
            if D.ha_commenti(p._p):
                D.sposta_commenti(p._p, dst.cells[j].paragraphs[0]._p)
    i = indice(doc, inizia('ANSC_ANA_UC — il catalogo degli UC'), 'titolo ANA_UC')
    h_ana = doc.paragraphs[i]._p
    vuoto = t_ana._tbl.getnext()
    elimina([h_ana, t_ana._tbl] + ([vuoto] if vuoto is not None and
                                    vuoto.tag.endswith('}p') and not ''.join(vuoto.itertext()).strip() else []))
    fatti.append('descrizione: ANSC_ANA_UC tolta, colonne riportate in ANSC_CFG_UC')

    h = doc.paragraphs[indice(doc, inizia('ANSC_CFG_UC — elenco delle regole'), 'titolo CFG_UC')]
    pezzi(h, 'ANSC_CFG_UC — il catalogo degli UC e le regole di identificazione.')
    p_eredita = doc.paragraphs[indice(doc, inizia('Questa tabella eredita dai file xlsx'), 'eredita')]
    for testo in reversed([
        'La tabella unisce il catalogo degli UC che ANSC pubblica e la configurazione con cui il '
        'Comune li adotta. Le colonne del catalogo — codice, codice motore, denominazione, '
        'famiglia, tipo evento, tipo contenuto, validità — si aggiornano dal mapping ufficiale e '
        'non si modificano a mano; le altre sono scelte del Comune.',
        'La riga è per UC e per Modello che lo raggiunge, dentro una baseline: più Modelli '
        'possono portare allo stesso UC, e ciascuno ha la propria priorità e la propria logica. '
        'Un UC che nessun Modello raggiunge ancora ha una riga sola, senza Modello: compare nel '
        'catalogo ma non partecipa alla determinazione. Ne segue che le colonne del catalogo si '
        'ripetono sulle righe dello stesso UC, e l’importazione le riscrive tutte insieme.',
        'Poiché lo stesso UC può stare su più righe, sezioni, campi, formule, allegati e soggetti '
        'non lo richiamano con una chiave esterna ma per valore, con la coppia COD_UC_ANSC e '
        'ID_VERSIONE: il codice pubblicato da ANSC è ciò che le righe dello stesso UC hanno in '
        'comune. L’esistenza dell’UC nella baseline la verificano l’importazione e i controlli di '
        'attivazione.',
        'Il catalogo entra così nella baseline: ogni baseline porta la fotografia del catalogo '
        'che ha recepito, e un UC ritirato da ANSC resta leggibile nelle baseline storiche. La '
        'revisione del mapping da cui il catalogo proviene non ha una colonna propria: è quella '
        'recepita dalla baseline (ANSC_CFG_VERSIONE.COD_VERSIONE_ANSC).',
    ]):
        nuovo = D.riga_dopo(doc, p_eredita, testo)
        nuovo.runs[0].bold = False
    fatti.append('ANSC_CFG_UC: quattro paragrafi sull’unificazione')

    # --- elenco delle tabelle del capitolo «Adeguamenti database»
    t_el = next(t for t in doc.tables
                if any(r.cells[0].text.strip() == 'ANSC_ANA_UC' for r in t.rows))
    r_ana = [r for r in t_el.rows if r.cells[0].text.strip() == 'ANSC_ANA_UC'][0]
    r_cfg = [r for r in t_el.rows if r.cells[0].text.strip().startswith('ANSC_CFG_UC —')][0]
    elimina([r_ana._tr])
    D.riscrivi_cella(r_cfg.cells[0], 'ANSC_CFG_UC — gli UC e la loro scelta')
    D.riscrivi_cella(r_cfg.cells[1], 'Catalogo degli UC dichiarati da ANSC e regole di '
                                     'determinazione, controllo e generazione (RF-11).')
    D.riscrivi_cella(r_cfg.cells[2],
                     'Una riga per UC e per Modello che lo raggiunge: codice, codice motore, '
                     'descrizione, famiglia e validità come li pubblica ANSC, e accanto le scelte '
                     'del Comune — il Modello, la priorità e la logica che decide, richiamata dal '
                     'catalogo delle logiche per codice (ID_DOMINIO, COD_LOGICA). Le colonne del '
                     'catalogo si aggiornano dal mapping ufficiale, non dall’operatore. Struttura '
                     'al capitolo «La determinazione dell’UC e le regole di controllo».')
    fatti.append('elenco delle tabelle: una riga sola per gli UC')

    # --- i tre livelli
    pezzi(doc.paragraphs[indice(doc, inizia('ANSC_ANA_UC: è il catalogo degli UC'), 'punto ANA')],
          'ANSC_CFG_UC: è il catalogo degli UC, con il codice, il codice motore e la descrizione '
          'pubblicati dal sistema nazionale, e insieme l’elenco delle condizioni utili '
          'all’identificazione dello UC;')
    pezzi(doc.paragraphs[indice(doc, inizia('ANSC_CFG_UC: è l’elenco delle condizioni'), 'punto CFG')],
          'ANSC_CFG_SEZIONE: le sezioni del modello evento che ciascun UC richiede e la condizione '
          'che le attiva, descritte sopra;')

    # --- sezione, formula: l'UC per valore
    t_sez = D.tabella_colonne(doc, 'SEZIONE_FE_ANSC')
    r = cella(doc, 'ID_UC', 2, 'UC di riferimento: il codice pubblicato da ANSC, che richiama '
              'ANSC_CFG_UC per valore insieme alla baseline, come in ALLEGATI_USECASE e '
              'SOGGETTI_USECASE. Le sezioni richieste differiscono fra UC dello stesso Modello.',
              tabella=t_sez)
    D.riscrivi_cella(r.cells[0], 'COD_UC_ANSC')
    D.riscrivi_cella(r.cells[1], 'VARCHAR2(20)')
    cella(doc, 'ID_VERSIONE', 2, 'Baseline di appartenenza (ANSC_CFG_VERSIONE). La sezione la '
          'dichiara e i suoi campi la ripetono, vincolati alla stessa.', tabella=t_sez)
    S('Vincolo di unicità su (ID_VERSIONE, ID_UC, OGGETTO_ANSC)',
      'Vincolo di unicità su (ID_VERSIONE, COD_UC_ANSC, OGGETTO_ANSC)')
    t_for = D.tabella_colonne(doc, 'COD_FORMULA')
    r = cella(doc, 'ID_UC', 2, 'UC di riferimento, per valore come per le sezioni e gli '
              'allegati. La formula pende dall’UC: lo stesso Modello di atto porta formule '
              'diverse secondo il caso d’uso determinato.', tabella=t_for)
    D.riscrivi_cella(r.cells[0], 'COD_UC_ANSC')
    D.riscrivi_cella(r.cells[1], 'VARCHAR2(20)')

    # --- gli altri richiami al catalogo
    S('Dal catalogo ANSC_ANA_UC, per codice.', 'Da ANSC_CFG_UC, per codice.')
    S('Si conserva sul catalogo (ANSC_ANA_UC.ID_TIPO_EVENTO)',
      'Si conserva sul catalogo (ANSC_CFG_UC.ID_TIPO_EVENTO)')
    S('Il legame è un riferimento per valore e non una chiave esterna verso il catalogo degli UC, '
      'per la stessa ragione per cui lo è in ANSC_CFG_UC: si nomina il codice pubblicato da ANSC, '
      'non la riga locale, così che la configurazione possa nominare un UC prima che il catalogo '
      'sia stato ricaricato.',
      'Il legame è un riferimento per valore — il codice dell’UC e la baseline — e non una '
      'chiave esterna verso ANSC_CFG_UC: nella stessa baseline un UC può stare su più righe, una '
      'per ciascun Modello che lo raggiunge, e il codice pubblicato da ANSC è ciò che le '
      'accomuna.')
    S('Va detto che il modello dati contiene già i semi di questo meccanismo, sparsi e non '
      'governati: ANSC_CFG_UC e ANSC_CFG_CAMPO portano una colonna di versione che entra nella '
      'chiave di unicità, ANSC_ANA_UC porta la versione del mapping, ANSC_STATO_ATTO porta la '
      'versione di configurazione dell’atto. Sono quattro stringhe libere: nessuna tabella ne '
      'governa i valori, nulla garantisce che siano coerenti fra loro e proprio ANSC_CFG_CAMPO — '
      'la tabella che cambia a ogni revisione — non ne ha alcuna.',
      'Il modello dati conteneva già i semi di questo meccanismo, sparsi e non governati: la '
      'configurazione degli UC, il catalogo degli UC e lo store di stato portavano ciascuno una '
      'versione come stringa libera, che nessuna tabella governava e nulla garantiva coerente '
      'con le altre, mentre proprio la tabella dei campi — quella che cambia a ogni revisione — '
      'non ne aveva alcuna.')
    S('Ne restano fuori il catalogo degli UC e i dizionari, che sono repliche di quanto ANSC '
      'dichiara e hanno una propria validità temporale: replicarli anche nella baseline '
      'significherebbe conservare copie di una verità che non ci appartiene.',
      'Il catalogo degli UC vi entra con ANSC_CFG_UC: ogni baseline porta la fotografia del '
      'catalogo che ha recepito. Ne restano fuori i dizionari, che sono repliche di quanto ANSC '
      'dichiara e hanno una propria validità temporale: replicarli anche nella baseline '
      'significherebbe conservare copie di una verità che non ci appartiene.')
    S('il codice dell’UC su ANSC_CFG_UC, su ALLEGATI_USECASE e su SOGGETTI_USECASE',
      'il codice dell’UC con cui ANSC_CFG_SEZIONE, ANSC_CFG_FORMULA, ALLEGATI_USECASE e '
      'SOGGETTI_USECASE richiamano ANSC_CFG_UC')
    # back-office
    for t in doc.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == 'Catalogo UC':
                D.riscrivi_cella(r.cells[1],
                                 'ANSC_CFG_UC letta come catalogo: gli UC dichiarati da ANSC con '
                                 'codice, codice motore, descrizione e famiglia, compresi quelli '
                                 'che nessun Modello raggiunge ancora. Le colonne del catalogo si '
                                 'aggiornano dal mapping ufficiale e qui si consultano soltanto: '
                                 'rappresentano ciò che il sistema nazionale dichiara, non una '
                                 'scelta del Comune.')
                fatti.append('back-office: pagina «Catalogo UC» su ANSC_CFG_UC')

    # --- DDL
    i = indice(doc, esatto('CREATE TABLE ANSC_USR.ANSC_CFG_UC ('), 'DDL CFG_UC')
    j = indice(doc, esatto("   soddisfatta determina l UC. Due righe con la stessa priorita "
                           "sono una ambiguita.';"), 'fine DDL CFG_UC', i)
    blocco(doc, i, j + 1, CFG_UC_DDL)
    i = indice(doc, titolo(3, 'ANSC_ANA_UC'), 'DDL ANA_UC')
    j = indice(doc, lambda p: p.style.name == 'Heading 3', 'dopo ANA_UC', i + 1)
    elimina([p._p for p in doc.paragraphs[i:j]])
    i = indice(doc, esatto('CREATE TABLE ANSC_USR.ANSC_CFG_SEZIONE ('), 'DDL sezione')
    riga_ddl(doc, i, '  ID_UC                NUMBER             NOT NULL,',
             '  COD_UC_ANSC          VARCHAR2(20 CHAR)  NOT NULL,')
    k = indice(doc, esatto('  CONSTRAINT FK_ANSC_CFG_SEZIONE_UC'), 'FK sezione', i)
    blocco(doc, k, k + 2, ['  CONSTRAINT UK_ANSC_CFG_SEZIONE_CAMPO',
                           '    UNIQUE (ID_SEZIONE, ID_VERSIONE, COD_UC_ANSC),'])
    riga_ddl(doc, i, '  CONSTRAINT UK_ANSC_CFG_SEZIONE UNIQUE (ID_VERSIONE, ID_UC, OGGETTO_ANSC),',
             '  CONSTRAINT UK_ANSC_CFG_SEZIONE UNIQUE (ID_VERSIONE, COD_UC_ANSC, OGGETTO_ANSC),')
    i = indice(doc, esatto('CREATE TABLE ANSC_USR.ANSC_CFG_FORMULA ('), 'DDL formula')
    riga_ddl(doc, i, '  ID_UC                NUMBER             NOT NULL,',
             '  COD_UC_ANSC          VARCHAR2(20 CHAR)  NOT NULL,')
    k = indice(doc, esatto('  CONSTRAINT FK_ANSC_CFG_FORMULA_UC'), 'FK formula', i)
    elimina([doc.paragraphs[k]._p, doc.paragraphs[k + 1]._p])
    riga_ddl(doc, i, '  CONSTRAINT UK_ANSC_CFG_FORMULA UNIQUE (ID_VERSIONE, ID_UC, COD_FORMULA),',
             '  CONSTRAINT UK_ANSC_CFG_FORMULA UNIQUE (ID_VERSIONE, COD_UC_ANSC, COD_FORMULA),')
    fatti.append('DDL: ANSC_ANA_UC tolta; CFG_UC con le colonne del catalogo; sezione e '
                 'formula con COD_UC_ANSC per valore')


# ------------------------------------------------------------------ 2. versione su CFG_CAMPO
def campo_versione(doc):
    t = D.tabella_colonne(doc, 'OPERATIVO')
    sez = [r for r in t.rows if r.cells[0].text.strip() == 'ID_SEZIONE'][0]
    nuove = [
        ('ID_VERSIONE', 'NUMBER (FK)',
         'Baseline di appartenenza (ANSC_CFG_VERSIONE), ripetuta dalla sezione. Con COD_UC_ANSC '
         'permette di leggere i campi di un UC in una baseline senza passare dalla sezione. La '
         'chiave esterna composta verso ANSC_CFG_SEZIONE (ID_SEZIONE, ID_VERSIONE, COD_UC_ANSC) '
         'impedisce che il campo dichiari una baseline o un UC diversi da quelli della sua '
         'sezione.'),
        ('COD_UC_ANSC', 'VARCHAR2(20) (FK)',
         'Codice dell’UC, ripetuto dalla sezione per la stessa ragione e con lo stesso vincolo.'),
    ]
    ancora = sez._tr
    for valori in nuove:
        r = D.clona_riga(t, valori)
        if D.ha_commenti(r._tr):
            raise SystemExit('riga clonata con commento')
        ancora.addnext(r._tr)
        ancora = r._tr
    D.riscrivi_cella(sez.cells[2], 'Sezione di appartenenza (ANSC_CFG_SEZIONE), che porta l’UC e '
                                   'la baseline: il campo le ripete.')

    i = indice(doc, esatto('CREATE TABLE ANSC_USR.ANSC_CFG_CAMPO ('), 'DDL campo')
    riga_ddl(doc, i, '  ID_SEZIONE           NUMBER             NOT NULL,',
             ['  ID_SEZIONE           NUMBER             NOT NULL,',
              '  ID_VERSIONE          NUMBER             NOT NULL,',
              '  COD_UC_ANSC          VARCHAR2(20 CHAR)  NOT NULL,'])
    riga_ddl(doc, i, '    FOREIGN KEY (ID_SEZIONE) REFERENCES ANSC_USR.ANSC_CFG_SEZIONE (ID_SEZIONE),',
             ['    FOREIGN KEY (ID_SEZIONE, ID_VERSIONE, COD_UC_ANSC)',
              '    REFERENCES ANSC_USR.ANSC_CFG_SEZIONE (ID_SEZIONE, ID_VERSIONE, COD_UC_ANSC),'])
    # commenti: il primo diceva che il campo non ripete UC e baseline; due COMMENT erano
    # doppi e nominavano BUSINESS_LOGIC, colonna che la CREATE non ha
    k = indice(doc, esatto('COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ID_SEZIONE IS'), 'commenti campo', i)
    fine = indice(doc, lambda p: not p.text.strip(), 'fine commenti campo', k)
    blocco(doc, k, fine, [
        'CREATE INDEX ANSC_USR.IX_ANSC_CFG_CAMPO_UC',
        '  ON ANSC_USR.ANSC_CFG_CAMPO (ID_VERSIONE, COD_UC_ANSC) TABLESPACE ANSC_USR;',
        'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ID_SEZIONE IS',
        "  'Sezione di appartenenza: porta l UC e la baseline, che il campo ripete; la chiave",
        "   esterna composta ne garantisce la coerenza.';",
        'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.OGGETTO_ANSC IS',
        "  'Binding Object del mapping ufficiale: l oggetto del modello evento che contiene il campo.';",
        'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ID_BUSINESS_LOGIC IS',
        "  'Trasformazione da applicare al valore letto da SIPO: codice del catalogo delle logiche",
        "   (dominio 3).';",
    ])
    fatti.append('ANSC_CFG_CAMPO: ID_VERSIONE e COD_UC_ANSC con FK composta verso la sezione; '
                 'COMMENT corretti (BUSINESS_LOGIC -> ID_BUSINESS_LOGIC, doppioni tolti)')


# ------------------------------------------------------------------ 3. COD_VERSIONE
def cod_versione(doc):
    t = D.tabella_colonne(doc, 'NUM_UC_TOCCATI')
    r = [r for r in t.rows if r.cells[0].text.strip() == 'COD_VERSIONE'][0]
    elimina([r._tr])
    cella(doc, 'ID_VERSIONE', 2,
          'Chiave tecnica. È la colonna che le tabelle di configurazione richiamano, e il numero '
          'con cui la baseline si nomina: un nome libero accanto ripeterebbe ciò che già dicono il '
          'numero, la revisione ANSC recepita e la descrizione.', tabella=t)
    i = indice(doc, esatto('CREATE TABLE ANSC_USR.ANSC_CFG_VERSIONE ('), 'DDL versione')
    k = indice(doc, esatto('  COD_VERSIONE         VARCHAR2(20 CHAR)  NOT NULL,'), 'COD_VERSIONE', i)
    D.sposta_commenti(doc.paragraphs[k]._p, doc.paragraphs[k - 1]._p)   # «Verificare che serva.»
    elimina([doc.paragraphs[k]._p])
    k = indice(doc, esatto('  CONSTRAINT UK_ANSC_CFG_VERSIONE UNIQUE (COD_VERSIONE),'), 'UK versione', i)
    elimina([doc.paragraphs[k]._p])
    fatti.append('ANSC_CFG_VERSIONE: COD_VERSIONE e il suo vincolo di unicità tolti '
                 '(commento di Word spostato su ID_VERSIONE)')


# ------------------------------------------------------------------ 4. comunicaz_anagrafica_usecase
COLONNE_COM = ('cd_usecase,ty_comunicazione_anagrafica,ty_sorgente_data_decorrenza,fg_operante,'
               'tx_note,cd_ope_ins_dt,ts_ins_dt,cd_ope_ult_agg_dt,ts_ult_agg_dt,ds_note_dt')


def _valore(v):
    if v == "''":
        return 'NULL'
    if re.fullmatch(r"'\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}(\.\d+)?'", v):
        return 'TIMESTAMP ' + v
    return v


def comunicaz(doc):
    i = indice(doc, esatto('-- DROP TABLE comunicaz_anagrafica_usecase;'), 'DDL comunicaz')
    h2 = indice(doc, titolo(2, 'Tabelle di controllo del flusso'), 'fine comunicaz', i)
    vecchi = [p.text for p in doc.paragraphs[i:h2]]
    blocchi, righe = [], None
    for t in vecchi:
        if t.startswith('INSERT INTO  comunicaz_anagrafica_usecase'):
            righe = []
            blocchi.append(righe)
        elif righe is not None and t.strip().startswith('('):
            dentro = t.strip()
            dentro = dentro[1:dentro.rindex(')')]
            valori = re.findall(r"'[^']*'|NULL", dentro)
            if len(valori) != 10:
                raise SystemExit('riga di carico non riconosciuta: ' + t)
            righe.append(', '.join(_valore(v) for v in valori))
    n = sum(len(b) for b in blocchi)
    vuote = sum(1 for b in blocchi for r in b if r.split(', ')[2] == 'NULL')
    nuove = """CREATE TABLE ANSC_USR.comunicaz_anagrafica_usecase (
  cd_usecase                   VARCHAR2(10 CHAR)  NOT NULL,
  ty_comunicazione_anagrafica  VARCHAR2(30 CHAR)  NOT NULL,
  ty_sorgente_data_decorrenza  VARCHAR2(30 CHAR),
  fg_operante                  CHAR(1 CHAR)       NOT NULL,
  tx_note                      CLOB,
  cd_ope_ins_dt                CHAR(6 CHAR)       NOT NULL,
  ts_ins_dt                    TIMESTAMP,
  cd_ope_ult_agg_dt            CHAR(6 CHAR),
  ts_ult_agg_dt                TIMESTAMP,
  ds_note_dt                   VARCHAR2(50 CHAR),
  CONSTRAINT comunicaz_anagraf_usecase_pk PRIMARY KEY (cd_usecase),
  CONSTRAINT comunicaz_anagraf_operante_ck CHECK (fg_operante IN ('0','1'))
) TABLESPACE ANSC_USR;
COMMENT ON TABLE ANSC_USR.comunicaz_anagrafica_usecase IS
  'Per ciascun UC di stato civile, il tipo di comunicazione anagrafica alla firma da
   utilizzare. Struttura ereditata dal sistema Side con i nomi di origine.';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.cd_usecase IS
  'Codice del caso d''uso di stato civile (codifica ANSC_3)';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.ty_comunicazione_anagrafica IS
  'Tipo di comunicazione anagrafica di default (codifica SIDe 107)';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.ty_sorgente_data_decorrenza IS
  'Tipo di campo sorgente da utilizzare come valore da impostare per la data decorrenza
   (codifica SIDe 108). Facoltativa: in Oracle la stringa vuota del carico di Side e NULL.';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.fg_operante IS
  'Flag che indica se i tipi di impostazioni previsti per la comunicazione anagrafica dello
   usecase sono attivi e quindi da applicare (1=Vero; 0=Falso)';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.tx_note IS 'Note eventuali';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.cd_ope_ins_dt IS
  'Utente che ha inserito il record';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.ts_ins_dt IS
  'Timestamp inserimento del record';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.cd_ope_ult_agg_dt IS
  'Utente che per ultimo ha aggiornato il record';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.ts_ult_agg_dt IS
  'Timestamp ultimo aggiornamento del record';
COMMENT ON COLUMN ANSC_USR.comunicaz_anagrafica_usecase.ds_note_dt IS
  'Descrizione contesto ultimo aggiornamento del record';
""".split('\n')
    nuove.append('-- Dati di caricamento: Oracle 12.2 non ammette più righe in una VALUES; '
                 'la stringa vuota')
    nuove.append('-- di Side diventa NULL, le date diventano letterali TIMESTAMP.')
    for b in blocchi:
        nuove.append(f'INSERT INTO ANSC_USR.comunicaz_anagrafica_usecase ({COLONNE_COM})')
        for k, r in enumerate(b):
            coda = ' FROM DUAL;' if k == len(b) - 1 else ' FROM DUAL UNION ALL'
            nuove.append(f'  SELECT {r}{coda}')
    nuove.append('')
    blocco(doc, i, h2, nuove)
    fatti.append(f'comunicaz_anagrafica_usecase in Oracle: {len(blocchi)} INSERT, {n} righe '
                 f'({vuote} con la data di decorrenza vuota → NULL)')
    return n, vuote


# ------------------------------------------------------------------ 5. forzatura_validazione
FORZATURA = """CREATE TABLE ANSC_USR.forzatura_validazione (
  id_forzatura_validaz   NUMBER GENERATED ALWAYS AS IDENTITY,
  id_evento              NUMBER             NOT NULL,
  cd_forzatura           VARCHAR2(30 CHAR)  NOT NULL,
  tx_anomalie            CLOB,
  cd_ope_ins_dt          VARCHAR2(40 CHAR)  NOT NULL,
  ts_ins_dt              TIMESTAMP          NOT NULL,
  cd_ope_ult_agg_dt      VARCHAR2(40 CHAR),
  ts_ult_agg_dt          TIMESTAMP,
  ds_note_dt             VARCHAR2(50 CHAR),
  CONSTRAINT forzatura_validazione_pk PRIMARY KEY (id_forzatura_validaz),
  CONSTRAINT forzatura_validaz_evento_fk
    FOREIGN KEY (id_evento) REFERENCES ANSC_USR.ANSC_STATO_ATTO (ID_STATO_ATTO),
  CONSTRAINT forzatura_validaz_anomalie_ck CHECK (tx_anomalie IS JSON)
) TABLESPACE ANSC_USR;
CREATE INDEX ANSC_USR.IX_FORZATURA_VALIDAZ_EVENTO
  ON ANSC_USR.forzatura_validazione (id_evento) TABLESPACE ANSC_USR;
COMMENT ON TABLE ANSC_USR.forzatura_validazione IS
  'Forzature della validazione di un evento di stato civile: il codice di forzatura
   registrato a operazione avvenuta, le anomalie superate e l''ufficiale dello stato civile
   che ha forzato';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.id_forzatura_validaz IS
  'Identificativo univoco del record di dati per possibile forzatura della validazione di un
   evento di stato civile';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.id_evento IS
  'L''atto in redazione (ANSC_STATO_ATTO.ID_STATO_ATTO): nome di origine, in Side punta alla
   tabella evento';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.cd_forzatura IS
  'Forcing code generato da ANSC e da utilizzare nel caso di forzatura della validazione di un
   evento di stato civile';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.tx_anomalie IS
  'Testo che rappresenta il json degli errori riscontrati da ANSC durante la richiesta di
   validazione evento e che sono considerati però anomalie superabili con l''applicazione del
   Forcing Code';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.cd_ope_ins_dt IS
  'Ufficiale dello stato civile che ha effettuato la forzatura: nome utente o altro
   identificativo, fino a 40 caratteri come le utenze di Roma';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.ts_ins_dt IS
  'Timestamp di inserimento del record';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.cd_ope_ult_agg_dt IS
  'Utente che per ultimo ha aggiornato il record';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.ts_ult_agg_dt IS
  'Timestamp di ultimo aggiornamento del record';
COMMENT ON COLUMN ANSC_USR.forzatura_validazione.ds_note_dt IS
  'Descrizione contesto ultimo aggiornamento del record';
""".split('\n')

FIRMATARI = """CREATE TABLE ANSC_USR.firmatari (
  id_atto_sipo             NUMBER             NOT NULL,
  id_soggetto              NUMBER             NOT NULL,
  ty_soggetto              VARCHAR2(30 CHAR)  NOT NULL,
  ty_firma                 VARCHAR2(30 CHAR),
  cd_stato_firma           VARCHAR2(30 CHAR),
  ds_email_notifica_firma  VARCHAR2(100 CHAR),
  id_presa_visione         NUMBER(10),
  cd_ope_ins_dt            VARCHAR2(40 CHAR)  NOT NULL,
  ts_ins_dt                TIMESTAMP          NOT NULL,
  cd_ope_ult_agg_dt        VARCHAR2(40 CHAR),
  ts_ult_agg_dt            TIMESTAMP,
  ds_note_dt               VARCHAR2(50 CHAR),
  CONSTRAINT firmatari_pk PRIMARY KEY (id_atto_sipo, id_soggetto)
) TABLESPACE ANSC_USR;
COMMENT ON TABLE ANSC_USR.firmatari IS
  'Firmatari di un atto di stato civile e stato della relativa firma. Struttura ereditata dal
   sistema Side; atto e soggetto sono riferimenti per valore a SIPO, perche nessuna chiave
   esterna attraversa gli schemi.';
COMMENT ON COLUMN ANSC_USR.firmatari.id_atto_sipo IS
  'Atto di SIPO a cui la firma appartiene (in Side id_evento, verso la tabella evento)';
COMMENT ON COLUMN ANSC_USR.firmatari.id_soggetto IS
  'Soggetto firmatario: MATR_USR.SOGGETTO.ID_SOGGETTO, per valore';
COMMENT ON COLUMN ANSC_USR.firmatari.ty_soggetto IS
  'Tipo di soggetto (intestatario, dichiarante...)';
COMMENT ON COLUMN ANSC_USR.firmatari.ty_firma IS
  'Codice SIDE per identificare il tipo di firma scelta per il firmatario (dominio 18 di
   tipo_dominio e valori_dominio) dovrebbe assumere i valori R006 / R012';
COMMENT ON COLUMN ANSC_USR.firmatari.cd_stato_firma IS
  'Codice SIDE per identificare lo stato della firma (dominio 19 di tipo_dominio e
   valori_dominio)';
COMMENT ON COLUMN ANSC_USR.firmatari.ds_email_notifica_firma IS
  'Email per inviare la notifica di firma digitale';
COMMENT ON COLUMN ANSC_USR.firmatari.id_presa_visione IS
  'Identificativo univoco della presa visione del firmatario tramite notifica via email del
   link all''esecuzione della firma digitale sull''atto di stato civile';
COMMENT ON COLUMN ANSC_USR.firmatari.cd_ope_ins_dt IS 'Utente che ha inserito la riga';
COMMENT ON COLUMN ANSC_USR.firmatari.ts_ins_dt IS 'Timestamp inserimento';
COMMENT ON COLUMN ANSC_USR.firmatari.cd_ope_ult_agg_dt IS
  'Utente che per ultimo ha aggiornato la riga';
COMMENT ON COLUMN ANSC_USR.firmatari.ts_ult_agg_dt IS 'Timestamp ultimo aggiornamento';
COMMENT ON COLUMN ANSC_USR.firmatari.ds_note_dt IS
  'Descrizione contesto ultimo aggiornamento';
CREATE INDEX ANSC_USR.IX_FIRMATARI_SOGGETTO
  ON ANSC_USR.firmatari (id_soggetto) TABLESPACE ANSC_USR;
""".split('\n')


def side_oracle(doc):
    i = indice(doc, titolo(3, 'firmatari'), 'titolo firmatari') + 1
    j = indice(doc, titolo(3, 'forzatura_validazione'), 'titolo forzatura', i)
    blocco(doc, i, j, FIRMATARI, sposta=[('  id_soggetto int8', '  id_soggetto ')])
    fatti.append('firmatari in Oracle: chiave (id_atto_sipo, id_soggetto), nessuna FK fuori '
                 'schema, commenti di Word spostati su id_soggetto')
    i = indice(doc, titolo(3, 'forzatura_validazione'), 'titolo forzatura') + 1
    j = indice(doc, lambda p: p.style.name == 'Heading 3', 'dopo forzatura', i)
    blocco(doc, i, j, FORZATURA)
    fatti.append('forzatura_validazione in Oracle: IDENTITY, CLOB IS JSON, FK su '
                 'ANSC_STATO_ATTO, indice su id_evento')
    S('(d) Il DDL di firmatari, ereditato da Side, dichiara una chiave esterna su id_evento ma la '
      'colonna è id_atto_sipo, e la chiave primaria è il solo id_soggetto, che consentirebbe a '
      'un soggetto di firmare un solo atto: va corretto prima dell’adozione.',
      '(d) Chiuso: la DDL di firmatari è in sintassi Oracle, con chiave primaria (id_atto_sipo, '
      'id_soggetto) e riferimenti per valore all’atto e al soggetto di SIPO.')


# ------------------------------------------------------------------ 7. la tabella evento di Side
def evento(doc):
    i = indice(doc, titolo(2, 'Tabelle di controllo del flusso'), 'evento')
    j = indice(doc, titolo(2, 'Tabelle dei dizionari'), 'dopo evento', i + 1)
    elimina([p._p for p in doc.paragraphs[i:j]])
    p = doc.paragraphs[indice(doc, inizia('Il modello evento è la copia di lavoro'), 'copia')]
    D.riga_dopo(doc, p,
                'La tabella evento del sistema Side svolge la stessa funzione e non si adotta '
                'accanto a questa: un atto ha una sola riga che lo descrive, ANSC_STATO_ATTO. Le '
                'due differiscono soprattutto per il corpo del modello evento: evento porta in '
                'colonna anche dichiarante, soggetti, consensi, atti collegati e composizione '
                'dell’atto, mentre qui quei dati restano nel modello evento conservato per intero '
                'e si promuove a colonna soltanto ciò su cui si cerca, si filtra, si decide o si '
                'riprende. Il codice di forzatura, che evento tiene in colonna, sta in '
                'forzatura_validazione insieme all’ufficiale che l’ha usato e alle anomalie '
                'superate.')
    fatti.append('Appendice A: tolta la tabella evento di Side («Tabelle di controllo del '
                 'flusso»); paragrafo in ANSC_STATO_ATTO')


# ------------------------------------------------------------------ main
def main():
    global DOC
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    DOC = doc
    unifica_uc(doc)
    campo_versione(doc)
    cod_versione(doc)
    comunicaz(doc)
    side_oracle(doc)
    evento(doc)

    subprocess.run([PY, os.path.join(QUI, 'diagrammi_erd.py')], check=True)
    D.sostituisci_immagine(doc, 'Schema ANSC_USR', os.path.join(QUI, 'img', 'erd_ansc_usr.png'))
    subprocess.run([PY, os.path.join(QUI, 'diagrammi_catena.py')], check=True)
    D.sostituisci_immagine(doc, 'Dalla maschera SIPO al payload ANSC',
                           os.path.join(QUI, 'img', 'catena_configurazione.png'))
    fatti.append('ERD e catena della configurazione rigenerati')

    D.storia(doc, '23/10/2026', '3.46',
             'Adeguamenti database · Costruzione del payload · Il versionamento della '
             'configurazione · Back-office · Modello dati (ERD) · Open Point · Appendice A',
             'Unificate ANSC_ANA_UC e ANSC_CFG_UC: la tabella della configurazione degli UC porta '
             'anche le colonne del catalogo (codice motore, denominazione, famiglia, tipo '
             'contenuto); sezioni e formule richiamano l’UC per valore con codice e baseline, e '
             'il catalogo entra nella baseline. ANSC_CFG_CAMPO riceve baseline e UC, vincolati a '
             'quelli della sezione. Tolto ANSC_CFG_VERSIONE.COD_VERSIONE, che nessun processo '
             'usava. comunicaz_anagrafica_usecase, forzatura_validazione e firmatari in sintassi '
             'Oracle (chiave di firmatari su atto e soggetto: OP-75 (d) chiuso). Tolta '
             'dall’Appendice A la tabella evento di Side: l’atto ha una sola tabella, '
             'ANSC_STATO_ATTO. Rigenerati lo schema ANSC_USR e la catena della configurazione.')
    fatti.extend(D.frontespizio(doc, '3.46', 'ANALISI_Integrazione-ANSC_v3.46'))
    doc.save(OUT)
    for f in fatti:
        print(' -', f)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
