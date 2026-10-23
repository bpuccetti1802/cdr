# -*- coding: utf-8 -*-
"""v3.45 dell'analisi ANSC, dalla v3.44 rivista a mano dall'utente (salvata il 23/10/2026, 08:17).

Tre interventi chiesti dall'utente:
  1. ripristino delle parole troncate: un «sostituisci tutto» di «atti.» le aveva tagliate. Il
     testo era rimasto nel run precedente, seguito da un run con il solo spazio: si aggiunge
     «atti.» in coda al run; per i riferimenti di DDL (atti.evento, atti.soggetto) si svuota
     anche il run dello spazio.
  2. numerazione: la serie è la tabella NUMERATORE_COMUNALE (struttura di Milano, una riga per
     anno), gestita dal servizio num-ansc-sipo. Via ogni riferimento alle sequence; stacco
     riscritto sul blocco della riga dell'anno; tabella descritta; DDL corretta (chiave
     sull'anno); ERD rigenerato da diagrammi_erd.py; OP-44 e OP-73 riallineati. I GRANT non
     servono alla documentazione (decisione dell'utente) e non si reintroducono.
  3. forzatura: la chiede l'ufficiale dello stato civile, unico abilitato, prima che ANSC
     risponda; il codice si registra in forzatura_validazione a operazione avvenuta con
     l'identificativo dell'ufficiale (cd_ope_ins_dt portato a 40 caratteri). Nota nel passo
     «VALIDA»; OP-21 chiuso.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_45_numeratore_forzatura.py
"""
import os
import sys

import docx

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.44.docx')
OUT = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.45.docx')
ERD = os.path.join(QUI, 'img', 'erd_ansc_usr.png')

fatti = []


def tutti_i_paragrafi(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


def unico(doc, condizione, cosa):
    trovati = [p for p in doc.paragraphs if condizione(p)]
    if len(trovati) != 1:
        raise SystemExit(f'{cosa}: attesi 1 paragrafo, trovati {len(trovati)}')
    return trovati[0]


def pezzi(p, parti):
    """Riscrive un paragrafo a pezzi (testo, grassetto). Si ferma se porta un commento."""
    if D.ha_commenti(p._p):
        raise SystemExit('paragrafo con commento, non lo riscrivo: ' + p.text[:60])
    modello = p.runs[0]._r.rPr if p.runs else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for testo, bold in parti:
        r = p.add_run(testo)
        if modello is not None:
            import copy
            r._r.insert(0, copy.deepcopy(modello))
        r.bold = bold
    return p


# ------------------------------------------------------------------ 1. «atti.»
TRONCATI = [
    # (fine del run troncato, svuotare il run di spazio che segue)
    ('non alla formazione di ', False),
    ('e non forma né firma ', False),
    ('che non mutano lo stato degli ', False),
    ('allinea lo stato, non forma ', False),
    ('non saper più leggere i propri ', False),
    ('fuori percorso, e resta agli ', False),
    ('come per lo store di stato degli ', False),
    ('il sistema colleghi i due ', False),
    ('dell’operatore e non forma ', False),
    ('in MATR_USR, insieme agli ', False),
    ("'Documenti allegati agli ", False),
    ('quando si configura una famiglia di ', False),
    ('si chiama id_evento e punta ad ', True),
    ('   referenzia ', True),
    ('firmatari_evento_fk FOREIGN KEY (id_evento) REFERENCES ', True),
    ('firmatari_soggetto_fk FOREIGN KEY (id_soggetto) REFERENCES ', True),
]


def atti(doc):
    for fine, attacca in TRONCATI:
        n = 0
        for p in tutti_i_paragrafi(doc):
            runs = p.runs
            for i, r in enumerate(runs):
                if not r.text.endswith(fine):
                    continue
                dopo = runs[i + 1].text if i + 1 < len(runs) else ''
                if dopo.strip():               # non è il punto troncato
                    continue
                r.text = r.text + 'atti.'
                if attacca and i + 1 < len(runs) and runs[i + 1].text == ' ':
                    runs[i + 1].text = ''
                n += 1
        if n != 1:
            raise SystemExit(f'«{fine}»: atteso 1 punto troncato, trovati {n}')
    fatti.append(f'«atti.» ripristinato in {len(TRONCATI)} punti')


# ------------------------------------------------------------------ 2. numerazione
def numerazione(doc):
    ps = doc.paragraphs

    def par(inizio):
        return unico(doc, lambda p: p.text.startswith(inizio), inizio[:40])

    pezzi(par('La gestione si eredita dal sistema Side del Comune di Milano'), [(
        'La gestione si eredita dal sistema Side del Comune di Milano [R12], con le sue tabelle, '
        'i suoi quattro stati e i suoi casi. La serie è tenuta dalla tabella NUMERATORE_COMUNALE, '
        'con una riga per anno, così che la numerazione riparta ogni anno: la struttura deriva da '
        'quella di Milano e soddisfa i requisiti di Roma. La tabella è gestita da un servizio '
        'specifico, num-ansc-sipo, che è l’unico a leggerla e a scriverla.', False)])
    p = par('Tale tabella prende il nome di NUMERATORE_COMUNALE')
    p._p.getparent().remove(p._p)

    pezzi(par('Riparte ogni anno per costruzione'), [
        ('Riparte ogni anno per costruzione. ', True),
        ('Ogni anno ha la propria riga nel numeratore, e la numerazione dell’anno nuovo comincia '
         'da uno.', False)])
    pezzi(par('La serie dell’anno prima resta disponibile'), [
        ('La serie dell’anno prima resta disponibile. ', True),
        ('A cavallo d’anno capita di dover staccare un numero della serie precedente: Side lo '
         'prevede esplicitamente, quando il gestionale torna disponibile dopo la fine dell’anno. '
         'Con una riga per anno basta indicare l’anno.', False)])
    pezzi(par('Sei cifre come limite del database'), [
        ('Sei cifre come limite del database. ', True),
        ('pg_atto è NUMBER(6): un numero di sette cifre fa fallire lo stacco invece di entrare '
         'nella serie.', False)])
    pezzi(par('Nessun numero perso per la cache'), [
        ('Nessun buco per una transazione fallita. ', True),
        ('Il numeratore avanza nella stessa transazione che registra i numeri: se la transazione '
         'fallisce, anche il numeratore torna al valore precedente, e nessun numero va perduto.',
         False)])
    pezzi(par('Si crea alla prima chiamata dell’anno'), [
        ('La riga dell’anno si crea alla prima richiesta. ', True),
        ('Se due richieste arrivano insieme, la chiave primaria sull’anno ne fa fallire una, che '
         'rilegge la riga creata dall’altra e prosegue.', False)])

    pezzi(par('Due tabelle, con i nomi di origine'), [(
        'Tre tabelle, con i nomi di origine. NUMERATORE_COMUNALE tiene la serie; '
        'PRENOTAZIONE_ATTI registra ogni richiesta; DATI_ATTO_ESTERNO registra ogni numero '
        'staccato e ciò che ne è stato. Di queste ultime due Side dichiara campi e significato ma '
        'non i tipi, che sono scelti qui.', False)])

    # descrizione di NUMERATORE_COMUNALE, prima di PRENOTAZIONE_ATTI
    ancora = par('PRENOTAZIONE_ATTI — la richiesta di numeri')
    modello = D.tabella_colonne(doc, 'nr_id_ansc_da_prenotare')
    D.para(doc, ancora._p, 'NUMERATORE_COMUNALE — la serie, una riga per anno.')
    D.tabella(doc, ancora._p, [
        ('Colonna', 'Tipo', 'Note'),
        ('id_anno_documento', 'NUMBER (PK)',
         'L’anno della serie. La riga si crea alla prima richiesta dell’anno; la chiave primaria '
         'impedisce che due richieste concorrenti ne creino due.'),
        ('pg_atto', 'NUMBER(6)',
         'L’ultimo numero staccato nell’anno, zero alla creazione della riga: il prossimo numero è '
         'pg_atto + 1. Sei cifre al massimo, secondo l’assunto del Comune (OP-73).'),
        ('dt_inserimento', 'TIMESTAMP', 'Quando la riga dell’anno è stata creata.'),
        ('dt_aggiornamento', 'TIMESTAMP',
         'Ultimo stacco, o ultima registrazione a posteriori che ha fatto avanzare il numeratore.'),
        ('nr_ordinamento', 'NUMBER',
         'Colonna della struttura di Milano, mantenuta; lo stacco non la usa.'),
    ], modello, larghezze=[1.6, 1.3, 3.6])
    fatti.append('NUMERATORE_COMUNALE descritta nel capitolo')

    D.sostituisci(doc, 'secondo l’assunto del Comune (OP-73): la sequence non può superarle.',
                  'secondo l’assunto del Comune (OP-73), come il contatore del numeratore.',
                  attese=1, fatti=fatti)
    D.sostituisci(doc, 'La sequence non sa che quel numero è stato consumato: la registrazione '
                  'la fa avanzare fino al numero, oppure colma un buco se il numero è inferiore al '
                  'prossimo,',
                  'Se il numero è superiore all’ultimo staccato, il numeratore è portato a quel '
                  'numero; se è inferiore e non registrato, colma un buco della serie;',
                  attese=1, fatti=fatti)
    D.sostituisci(doc, 'Nel percorso ordinario il numero si stacca al deposito, e l’anno è quello '
                  'del deposito;',
                  'Nel percorso ordinario il numero si stacca all’Anteprima, prima della '
                  'validazione, e l’anno è quello dello stacco;', attese=1, fatti=fatti)
    D.sostituisci(doc, 'creaNumero, che registra il numero già usato e tiene coerente la sequence.',
                  'creaNumero, che registra il numero già usato e allinea il numeratore.',
                  attese=1, fatti=fatti)
    D.sostituisci(doc, 'creaSerie all’installazione o se il job di dicembre non ha creato le serie;',
                  'creaSerie per creare in anticipo la riga dell’anno nel numeratore, facoltativa '
                  'perché la crea comunque la prima richiesta;', attese=1, fatti=fatti)

    # lo stacco, passo per passo
    pezzi(par('Lo stacco è la sola operazione del servizio'), [(
        'Lo stacco è la sola operazione del servizio che fa avanzare il numeratore, ed è quella '
        'da cui dipendono le due garanzie: numeri consecutivi per una stessa richiesta e nessun '
        'numero perduto. Si svolge in una sola transazione.', False)])
    pezzi(par('Serializzazione. Il servizio acquisisce il blocco'), [
        ('Serializzazione. ', True),
        ('Il servizio blocca la riga dell’anno in NUMERATORE_COMUNALE (SELECT … FOR UPDATE), '
         'creandola se è la prima richiesta dell’anno. Da qui alla conferma nessun altro stacco '
         'dello stesso anno procede; le letture continuano.', False)])
    pezzi(par('Riuso. Se la richiesta indica'), [
        ('Riuso. ', True),
        ('Se la richiesta indica un atto di SIPO che ha già un numero Disponibile nella stessa '
         'serie — un deposito respinto e ritentato — il servizio restituisce quel numero, conferma '
         'e si ferma: il numeratore non avanza.', False)])
    pezzi(par('Numeri. Per ciascuno dei numeri richiesti'), [
        ('Numeri. ', True),
        ('Per ciascuno dei numeri richiesti incrementa pg_atto e inserisce la riga di '
         'DATI_ATTO_ESTERNO con quel numero, nello stato Disponibile e legata all’atto se '
         'indicato: i numeri di una stessa richiesta sono quindi consecutivi. Se un numero '
         'risultasse già registrato, il vincolo di unicità della serie lo respinge: la '
         'transazione si annulla e l’errore segnala una serie incoerente, da esaminare. Non '
         'accade se la registrazione a posteriori segue la regola descritta sotto.', False)])
    pezzi(par('Conferma. La conferma rende visibili'), [
        ('Conferma. ', True),
        ('La conferma rende visibili richiesta, numeri e nuovo valore del numeratore, e rilascia '
         'il blocco. Se la transazione fallisce, si annulla tutto, compreso l’avanzamento del '
         'numeratore: nessun numero va perduto.', False)])
    pezzi(par('La registrazione a posteriori di un numero già usato'), [(
        'La registrazione a posteriori di un numero già usato fuori da SIPO — il percorso «SIPO '
        'non disponibile» — si svolge con lo stesso blocco sulla riga dell’anno e confronta il '
        'numero con l’ultimo staccato (pg_atto).', False)])
    pezzi(par('Numero inferiore al prossimo e non registrato'), [
        ('Numero non superiore all’ultimo staccato e non registrato. ', True),
        ('Colma un buco della serie e si registra nello stato Utilizzato; il numeratore non '
         'cambia.', False)])
    pezzi(par('Numero uguale o superiore al prossimo'), [
        ('Numero superiore all’ultimo staccato. ', True),
        ('Il numeratore è portato a quel numero, che si registra nello stato Utilizzato: così '
         'nessuno stacco successivo può restituirlo. I numeri eventualmente scavalcati restano '
         'senza registrazione e sono i soli buchi possibili della serie, visibili nella sua '
         'lettura.', False)])
    fatti.append('«Lo stacco, passo per passo» riscritto sul numeratore')

    # oggetti, ERD, Open Point
    D.sostituisci(doc, 'Una tabella che per per anno garantisce la numerazione; riparte ogni '
                  'anno.', 'Una tabella con una riga per anno, gestita dal servizio '
                  'num-ansc-sipo: la numerazione riparte ogni anno.', attese=1, fatti=fatti)
    D.sostituisci(doc, ', e con le sequence della numerazione comunale, che non sono tabelle.',
                  '.', attese=1, fatti=fatti)
    D.sostituisci_immagine(doc, 'Schema ANSC_USR, con ALLEGATO', ERD)
    fatti.append('ERD dello schema ANSC_USR sostituito')

    _op(doc, 'OP-44',
        'Chiuso. La serie è tenuta dalla tabella NUMERATORE_COMUNALE, una riga per anno, gestita '
        'dal servizio num-ansc-sipo: riparte ogni anno ed è unica per tutti i registri. Gli '
        'stacchi sono serializzati sulla riga dell’anno, così che i blocchi per il parto plurimo '
        'siano consecutivi; ogni numero è registrato nella tabella di Side DATI_ATTO_ESTERNO '
        'nella stessa transazione in cui è preso. Si veda il paragrafo «La numerazione '
        'comunale». Restano le conferme di OP-73.')
    _op(doc, 'OP-73',
        'Due conferme e un’utenza. (a) Che ANSC accetti il numero comunale di sei cifre e un '
        'progressivo unico fra i registri, come assume il Comune: il modello evento e il servizio '
        'di consultazione dichiarano per il numero dell’atto al massimo cinque caratteri, e il '
        'registro delle modifiche di ANSC riporta il controllo relativo. Se l’assunto non fosse '
        'confermato, la serie andrebbe tenuta per registro e anno, a cinque cifre, con il '
        'registro nella chiave del numeratore e delle tabelle di Side. (b) A quale anno '
        'appartenga il numero di un atto staccato in un anno e formato nel successivo. (c) '
        'L’utenza applicativa con cui il servizio num-ansc-sipo raggiunge ANSC_USR.')

    # Appendice A
    h = unico(doc, lambda p: p.style.name == 'Heading 3'
              and p.text.strip() == 'NUMERATORE_COMUNALE', 'titolo DDL numeratore')
    i = D.indice_di(doc, h)
    righe = doc.paragraphs[i + 1:i + 8]
    attese = ['CREATE TABLE ANSC_USR.NUMERATORE_COMUNALE (', '  id_anno_documento',
              '  pg_atto', '  dt_inserimento', '  dt_aggiornamento', '  nr_ordinamento',
              ') TABLESPACE ANSC_USR;']
    for p, a in zip(righe, attese):
        if not p.text.startswith(a):
            raise SystemExit(f'DDL del numeratore: atteso «{a}», trovato «{p.text}»')
    nuove = ['  id_anno_documento    NUMBER             NOT NULL,',
             '  pg_atto              NUMBER(6,0)        NOT NULL,',
             '  dt_inserimento       TIMESTAMP,',
             '  dt_aggiornamento     TIMESTAMP,',
             '  nr_ordinamento       NUMBER,',
             '  CONSTRAINT numeratore_comunale_pk PRIMARY KEY (id_anno_documento)']
    for p, t in zip(righe[1:6], nuove):
        D.testo_di(p, t)
    D.ddl(doc, righe[6]._p, [nuove[5]])
    dopo = doc.paragraphs[D.indice_di(doc, righe[6]) + 1]
    D.ddl(doc, dopo._p, [
        'COMMENT ON TABLE ANSC_USR.NUMERATORE_COMUNALE IS',
        "  'La serie dei numeri comunali: una riga per anno, creata alla prima richiesta.",
        "   Struttura di Milano, gestita dal servizio num-ansc-sipo.';",
        'COMMENT ON COLUMN ANSC_USR.NUMERATORE_COMUNALE.pg_atto IS',
        "  'Ultimo numero comunale staccato nell''anno; 0 alla creazione della riga.';",
    ])
    # la riga «nr_ordinamento» non chiude più la tabella: la virgola resta, il vincolo segue
    fatti.append('DDL di NUMERATORE_COMUNALE corretta, con chiave sull’anno e COMMENT')

    t = unico(doc, lambda p: p.style.name == 'Heading 3'
              and p.text.startswith('PRENOTAZIONE_ATTI, e le sequence'), 'titolo PRENOTAZIONE')
    D.testo_di(t, 'PRENOTAZIONE_ATTI')
    # la «/» isolata dopo l'ultimo COMMENT della numerazione
    i0 = D.indice_di(doc, unico(doc, lambda p: p.style.name.startswith('Heading')
                                and p.text.strip() == 'Numerazione Comunale', 'App. A'))
    barre = []
    for p in doc.paragraphs[i0 + 1:]:
        if p.style.name.startswith('Heading 1'):
            break
        if p.text.strip() == '/':
            barre.append(p)
    if len(barre) != 1:
        raise SystemExit(f'«/» isolata: attesa 1, trovate {len(barre)}')
    barre[0]._p.getparent().remove(barre[0]._p)
    fatti.append('App. A: titolo senza «sequence», tolta la «/» isolata')


def _op(doc, codice, testo, stato=None):
    for t in doc.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == codice and len(r.cells) >= 4:
                D.riscrivi_cella(r.cells[2], testo)
                if stato:
                    D.riscrivi_cella(r.cells[3], stato)
                fatti.append(f'{codice} riscritto' + (f', {stato}' if stato else ''))
                return
    raise SystemExit('open point non trovato: ' + codice)


# ------------------------------------------------------------------ 3. forzatura
def forzatura(doc):
    D.sostituisci(doc, 'Il suo uso è già registrato come OP-21, ma va ora letto insieme a un’altra '
                  'evidenza:', 'Il suo uso è regolato più sotto (OP-21) e va letto insieme a '
                  'un’altra evidenza:', attese=1, fatti=fatti)
    p = unico(doc, lambda p: p.text.startswith('A tal proposito è stata ereditata dal SIDE'),
              'paragrafo forzatura')
    pezzi(p, [(
        'La forzatura si chiede prima che ANSC risponda: è l’ufficiale dello stato civile che, '
        'alla validazione, decide di forzare, e la richiesta R009 porta il codice di forzatura. '
        'A operazione avvenuta il codice si registra nella tabella forzatura_validazione, '
        'ereditata da Side, con le anomalie superate e il nome o altro identificativo '
        'dell’ufficiale che ha forzato. Le forzature possono essere effettuate soltanto '
        'dall’ufficiale dello stato civile: la funzione è riservata al suo profilo. Sarà '
        'introdotta nel passo «VALIDA» in una prossima release del documento.', False)])

    # le stesse righe compaiono anche in comunicaz_anagrafica_usecase: si tocca solo la tabella
    # della forzatura, cercandola dopo il suo titolo
    h = unico(doc, lambda p: p.style.name == 'Heading 3'
              and p.text.strip() == 'forzatura_validazione', 'titolo DDL forzatura')
    i = D.indice_di(doc, h)
    blocco = doc.paragraphs[i + 1:i + 40]
    sostituzioni = {
        '  cd_ope_ins_dt bpchar(6) NOT NULL, -- Utente che ha inserito il record':
            '  cd_ope_ins_dt varchar(40) NOT NULL, -- Ufficiale dello stato civile che ha '
            'effettuato la forzatura (nome utente o altro identificativo)',
        '  cd_ope_ult_agg_dt bpchar(6) NULL, -- Utente che per ultimo ha aggiornato il record':
            '  cd_ope_ult_agg_dt varchar(40) NULL, -- Utente che per ultimo ha aggiornato il '
            'record',
        "COMMENT ON TABLE forzatura_validazione IS 'Tabella che contiene i dati che ANSC concede "
        "all''utente per eventualmente operare una forzatura della validazione di un evento di "
        "stato civile';":
            "COMMENT ON TABLE forzatura_validazione IS 'Forzature della validazione di un evento "
            "di stato civile: il codice di forzatura registrato a operazione avvenuta, le "
            "anomalie superate e l''ufficiale dello stato civile che ha forzato';",
        "COMMENT ON COLUMN forzatura_validazione.cd_ope_ins_dt IS 'Utente che ha inserito il "
        "record';":
            "COMMENT ON COLUMN forzatura_validazione.cd_ope_ins_dt IS 'Ufficiale dello stato "
            "civile che ha effettuato la forzatura: nome utente o altro identificativo, fino a "
            "40 caratteri come le utenze di Roma';",
    }
    fatte = 0
    for p in blocco:
        if p.text in sostituzioni:
            D.testo_di(p, sostituzioni[p.text])
            fatte += 1
    if fatte != len(sostituzioni):
        raise SystemExit(f'DDL forzatura: attese {len(sostituzioni)} righe, trovate {fatte}')
    fatti.append('forzatura_validazione: utente a 40 caratteri, COMMENT riallineati')

    _op(doc, 'OP-21',
        'Chiuso. La forzatura è riservata all’ufficiale dello stato civile: la chiede alla '
        'validazione, prima che ANSC risponda, e a operazione avvenuta il codice si registra in '
        'forzatura_validazione con il nome o altro identificativo dell’ufficiale. La funzione '
        'sarà introdotta nel passo «VALIDA» in una prossima release del documento. Resta aperto '
        'OP-41 sul flusso di conferma degli atti anomali.', stato='Chiuso')

    v = unico(doc, lambda p: p.style.name == 'Heading 2'
              and p.text.startswith('Il passo «VALIDA»'), 'titolo VALIDA')
    seguente = doc.paragraphs[D.indice_di(doc, v) + 2]
    D.para(doc, seguente._p, pezzi=[
        ('Nota. ', True),
        ('Nelle prossime release del documento sarà introdotta nel passo «VALIDA» la '
         'possibilità di effettuare una «Forzatura» della validazione. La chiede l’ufficiale '
         'dello stato civile, unico abilitato, prima che ANSC risponda; a operazione avvenuta il '
         'codice di forzatura si registra in forzatura_validazione con il nome o altro '
         'identificativo dell’ufficiale (si veda «Il codice di forzatura», OP-21).', False)])
    fatti.append('nota sulla forzatura nel passo «VALIDA»')


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    atti(doc)
    numerazione(doc)
    forzatura(doc)
    D.storia(doc, '23/10/2026', '3.45',
             'Tutto il documento · La numerazione comunale · Adeguamenti database · Modello dati '
             '(ERD) · Le funzioni offerte da ANSC · Il passo «VALIDA» · Open Point · Appendice A',
             'Ripristinate le parole troncate da una sostituzione di «atti.». Numerazione comunale: '
             'la serie è la tabella NUMERATORE_COMUNALE, una riga per anno, struttura di Milano '
             'gestita dal servizio num-ansc-sipo; tolti i riferimenti alle sequence, riscritto lo '
             'stacco sul blocco della riga dell’anno (una transazione fallita non lascia buchi), '
             'descritta la tabella, corretta la DDL con la chiave sull’anno, rigenerato lo schema '
             'ANSC_USR; OP-44 e OP-73 riallineati. Forzatura: riservata all’ufficiale dello stato '
             'civile, il codice si registra in forzatura_validazione a operazione avvenuta con il '
             'suo identificativo; nota nel passo «VALIDA»; OP-21 chiuso.')
    fatti.extend(D.frontespizio(doc, '3.45', 'ANALISI_Integrazione-ANSC_v3.45'))
    doc.save(OUT)
    for f in fatti:
        if f:
            print(' ·', f)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
