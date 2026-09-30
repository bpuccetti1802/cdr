# -*- coding: utf-8 -*-
"""v3.21 — allineamento fra le schede delle tabelle (capitoli) e il DDL (Appendice A).

Il confronto colonna per colonna fra le 13 schede «Colonna | Tipo | Note» e le 13 CREATE TABLE
ha trovato quattro tabelle divergenti. Criterio deciso con l'autore: **la scheda è la revisione
più recente, quindi vince e il DDL la segue**, con le eccezioni che seguono.

Decisioni dell'autore (15/09/2026):
  · ANSC_CFG_UC: TIPO_RITO e STATO si TOLGONO dalla scheda (erano incompleti: tipo vuoto l'uno,
    nome già occupato dallo stato dell'atto l'altro);
  · ANSC_XREF: ID_UC_ANSC e ID_OPERAZIONE_ANSC servono ENTRAMBI — non erano una rinomina ma due
    informazioni diverse (il caso d'uso e l'identificativo dell'operazione).

⚠️ Un'eccezione al criterio, dichiarata: TIPO_EVENTO/TIPO_OPERAZIONE della scheda di
ANSC_CFG_UC si riportano a COD_TIPO_EVENTO/COD_TIPO_OPERAZIONE, perché lo stesso concetto è già
così in ANSC_STATO_ATTO e in ANSC_XREF: seguire la scheda avrebbe dato tre nomi per una cosa
sola dentro lo stesso schema.

⚠️ Le tabelle si cercano per intestazione, mai per indice: `doc.tables[n]` slitta appena si
inserisce una tabella prima (è l'errore della v3.20).
"""
import copy
import os
import shutil
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')
SORG = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.20.docx')
DEST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.21.docx')


# ------------------------------------------------------------------ utilità

def scheda(doc, nome):
    """La tabella «Colonna | Tipo | Note» che segue il paragrafo che nomina la tabella."""
    import re
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    rx = re.compile(r'^([A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+|ALLEGATO)\b')
    ultimo = None
    for ch in doc.element.body.iterchildren():
        if ch.tag.endswith('}p'):
            m = rx.match(Paragraph(ch, doc).text.strip())
            if m:
                ultimo = m.group(1)
        elif ch.tag.endswith('}tbl'):
            t = Table(ch, doc)
            testa = [c.text.strip().lower() for c in t.rows[0].cells]
            if testa[:2] == ['colonna', 'tipo'] and ultimo == nome:
                return t
    raise SystemExit('scheda non trovata: ' + nome)


def riga_con(t, colonna):
    for r in t.rows[1:]:
        if r.cells[0].text.strip() == colonna:
            return r
    raise SystemExit(f'riga «{colonna}» non trovata')


def togli_riga(t, colonna):
    r = riga_con(t, colonna)
    if D.ha_commenti(r._tr):
        raise SystemExit(f'la riga «{colonna}» porta un commento: non si elimina')
    t._tbl.remove(r._tr)


def inserisci_dopo(t, colonna, valori):
    """Nuova riga subito dopo quella indicata, con la formattazione della riga precedente."""
    rif = riga_con(t, colonna)
    tr = copy.deepcopy(rif._tr)
    rif._tr.addnext(tr)
    from docx.table import _Row
    riga = _Row(tr, t)
    for cella, val in zip(riga.cells, valori):
        for p in cella.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        D.riscrivi_cella(cella, val)
    return riga


def riga_ddl(doc, contiene):
    trovati = [p for p in doc.paragraphs
               if contiene in p.text and any(r.font.name == 'Courier New' for r in p.runs)]
    if len(trovati) != 1:
        raise SystemExit(f'riga DDL «{contiene}»: trovate {len(trovati)} occorrenze')
    return trovati[0]


def ddl_dopo(doc, contiene, testo):
    rif = riga_ddl(doc, contiene)
    nuovo = copy.deepcopy(rif._p)
    rif._p.addnext(nuovo)
    from docx.text.paragraph import Paragraph
    p = Paragraph(nuovo, rif._parent)
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = testo
    return p


# ------------------------------------------------------------------ interventi

def main():
    shutil.copy2(SORG, DEST)
    doc = docx.Document(DEST)
    fatti = []

    # ── 1. ANSC_CFG_UC ────────────────────────────────────────────────────────
    t = scheda(doc, 'ANSC_CFG_UC')
    togli_riga(t, 'TIPO_RITO')
    togli_riga(t, 'STATO')
    D.riscrivi_cella(riga_con(t, 'TIPO_EVENTO').cells[0], 'COD_TIPO_EVENTO')
    D.riscrivi_cella(riga_con(t, 'TIPO_OPERAZIONE').cells[0], 'COD_TIPO_OPERAZIONE')
    D.riscrivi_cella(riga_con(t, 'SERIE').cells[1], 'VARCHAR2(10)')
    fatti.append('ANSC_CFG_UC: tolte TIPO_RITO e STATO; prefisso COD_ ai due tipi')

    # il DDL segue: LOGICA_DI_SCELTA e la nuova SERIE
    D.sostituisci(doc, 'REGOLA_DI_SCELTA', 'LOGICA_DI_SCELTA', attese=3,
                  etichetta='ANSC_CFG_UC.LOGICA_DI_SCELTA (DDL, commento, prospetto)', fatti=fatti)
    ddl_dopo(doc, '  MASCHERA_UI          VARCHAR2(100 CHAR),',
             '  SERIE                VARCHAR2(10 CHAR),')
    fatti.append('ANSC_CFG_UC: SERIE aggiunta al DDL')

    # ── 2. ANSC_CFG_CAMPO ─────────────────────────────────────────────────────
    t = scheda(doc, 'ANSC_CFG_CAMPO')
    inserisci_dopo(t, 'ID_UC', ['ID_VERSIONE', 'NUMBER (FK)',
                                'Baseline di appartenenza (ANSC_CFG_VERSIONE). ⚠️ È nella '
                                'chiave di unicità insieme a ID_UC e al campo: è ciò che fa '
                                'coesistere più versioni della stessa configurazione.'])
    D.riscrivi_cella(riga_con(t, 'SCHEMA_SIPO').cells[2],
                     'Schema di SIPO in cui risiede la tabella (ANAG_USR, MATR_USR, …).')
    D.riscrivi_cella(riga_con(t, 'TABELLA_SIPO').cells[2],
                     'Tabella di SIPO da cui il valore si prende.')
    D.riscrivi_cella(riga_con(t, 'OPERATIVO').cells[1], 'CHAR(1)')
    fatti.append('ANSC_CFG_CAMPO: reintrodotta ID_VERSIONE; note SCHEMA_SIPO/TABELLA_SIPO '
                 'rimesse al posto giusto')

    D.sostituisci(doc, 'DESC_COND_OBBLIGATORIETA   VARCHAR2(1000 CHAR)',
                  'DESC_COND_OBBLIGATORIETA_EVENTO   VARCHAR2(1000 CHAR)', attese=1,
                  etichetta='ANSC_CFG_CAMPO.DESC_COND_OBBLIGATORIETA_EVENTO', fatti=fatti)
    D.sostituisci(doc, 'REGOLA_COND_OBBLIGATORIETA VARCHAR2(1000 CHAR)',
                  'LOGICA_COND_OBBLIGATORIETA_EVENTO VARCHAR2(1000 CHAR)', attese=1,
                  etichetta='ANSC_CFG_CAMPO.LOGICA_COND_OBBLIGATORIETA_EVENTO', fatti=fatti)
    ddl_dopo(doc, '  FLG_OBBLIGATORIO     CHAR(1) DEFAULT \'N\' NOT NULL,',
             '  SCHEMA_SIPO          VARCHAR2(200 CHAR),')
    ddl_dopo(doc, '  ORDINAMENTO          NUMBER DEFAULT 1   NOT NULL,',
             "  OPERATIVO            CHAR(1) DEFAULT '1' NOT NULL,")
    ddl_dopo(doc, "  CONSTRAINT CK_ANSC_CFG_CAMPO_OBBL CHECK (FLG_OBBLIGATORIO IN ('S','N'))",
             "  ,CONSTRAINT CK_ANSC_CFG_CAMPO_OPERATIVO CHECK (OPERATIVO IN ('0','1'))")
    fatti.append('ANSC_CFG_CAMPO: SCHEMA_SIPO e OPERATIVO aggiunte al DDL')

    # ── 3. ANSC_ANA_UC ────────────────────────────────────────────────────────
    D.sostituisci(doc, '  COD_VERSIONE_MAPPING VARCHAR2(20 CHAR),',
                  '  COD_VERSIONE         VARCHAR2(20 CHAR),', attese=1,
                  etichetta='ANSC_ANA_UC.COD_VERSIONE', fatti=fatti)
    D.sostituisci(doc, '  COD_CATEGORIA        VARCHAR2(30 CHAR),',
                  '  ID_TIPO_DOCUMENTO    VARCHAR2(20 CHAR),', attese=1,
                  etichetta='ANSC_ANA_UC.ID_TIPO_DOCUMENTO', fatti=fatti)

    # ── 4. ANSC_XREF: le due informazioni convivono ───────────────────────────
    t = scheda(doc, 'ANSC_XREF')
    D.riscrivi_cella(riga_con(t, 'DATA_ACQUISIZIONE').cells[2], 'Momento della conferma.')
    inserisci_dopo(t, 'ID_UC_ANSC',
                   ['ID_OPERAZIONE_ANSC', 'VARCHAR2(50)',
                    'Identificativo dell’operazione comunicato ad ANSC '
                    '(idOperazioneComune). ⚠️ Non è la stessa cosa dell’UC: quello dice quale '
                    'caso d’uso è stato adottato, questo a quale chiamata l’atto si deve.'])
    ddl_dopo(doc, '  ID_OPERAZIONE_ANSC  VARCHAR2(50 CHAR)  NOT NULL,',
             '  ID_UC_ANSC          VARCHAR2(50 CHAR),')
    fatti.append('ANSC_XREF: ID_UC_ANSC e ID_OPERAZIONE_ANSC ora presenti in entrambe le parti')

    # ── 5. il vincolo spezzato di ANSC_STATO_ATTO ─────────────────────────────
    a = riga_ddl(doc, "CONSTRAINT CK_ANSC_STATO_ATTO_FLG_EMERGENZA")
    b = riga_ddl(doc, "CHECK (TIPO_ESITO IS NULL OR TIPO_ESITO IN")
    ta, tb = a.text, b.text
    D.testo_di(a, tb)
    D.testo_di(b, ta)
    fatti.append('ANSC_STATO_ATTO: rimesso in ordine CK_..._ESITO, che era spezzato in due')

    doc.save(DEST)
    print('\n'.join(' · ' + f for f in fatti))
    print('salvato:', os.path.relpath(DEST, BASE))


if __name__ == '__main__':
    main()
