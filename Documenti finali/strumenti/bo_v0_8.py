# -*- coding: utf-8 -*-
"""DISEGNO_Back-Office_ANSC v0.7 → v0.8 (10/10/2026).

Riallineamento alla v3.36 dell'analisi [R1]: la gestione dei documenti ADOTTA il processo
del sistema Side del Comune di Milano (indicazione del committente). Tre interventi:

  1. «Allegati dell'atto» rifatta: sei stati (0 locale + ANSC_08), INVIA e VERIFICA
     mutuamente esclusivi, forma a descrizione libera (998, D.LIB.M.01) invece dell'elenco
     separato degli «atti a testo libero», e la colonna dell'attestazione IN SOLA LETTURA —
     la scelta è dell'ufficiale atto per atto, non viene più dalla configurazione;
  2. due pagine nuove nell'area di configurazione, «Allegati per caso d'uso» e «Soggetti per
     caso d'uso», che governano per intero ALLEGATI_USECASE e SOGGETTI_USECASE;
  3. una pagina «Caricamento da file», comune alle due, nel formato in cui Side esporta le
     sue tabelle. ⚠️ Il solo file ricevuto è `Tabella SOGGETTI_USECASE.csv` (soggetti, non
     allegati): il formato si ricava da lì, e il verbale mostrato è quello vero del file.
     Per ALLEGATI_USECASE l'esportazione di Side non è fra le sorgenti → BO-29.

⚠️ Tre pagine nuove dopo la Figura 14: le figure 15…24 scalano di tre (al contrario).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v3.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/bo_v0_8.py
"""
import os
import re
import shutil
import subprocess
import sys

import docx
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402
import pagine_bo_v3 as P3   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.7.docx')
DST = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.8.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')

V = P3.verbale_soggetti()
M = P3.migl


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


# ════════════════════════════════════════════ 1. Allegati dell'atto (rifatta)
ALLEGATI_TESTO = [
    ('p', 'SIPO oggi non gestisce alcun documento allegato agli atti di stato civile, ma la '
          'funzione non si progetta: si adotta quella del sistema Side del Comune di Milano, '
          'con le sue tabelle, i suoi stati e le sue regole, come stabilito nell’analisi '
          '[R1, paragrafo «La gestione dei documenti: il processo di Side adottato in SIPO»]. '
          'Questa pagina ne è la versione di supervisione: mostra i documenti di un atto e '
          'consente le stesse azioni della maschera di stesura.'),
    ('p', 'L’elenco si legge dalla configurazione dell’UC determinato. Per ciascun tipo di '
          'documento la presenza dice se è obbligatorio, facoltativo o condizionato; la forma '
          'dice come si presenta la riga. Nella forma a descrizione fissa c’è un documento per '
          'tipo; nella forma a descrizione libera — il tipo 998 «Allegato generico» — '
          'l’operatore scrive la descrizione e può aggiungere più righe. È la stessa via per '
          'cui passano i documenti che l’UC non prevede: non serve un secondo elenco.'),
    ('p', 'Lo stato decide le azioni. Lo «0 — Caricato in SIPO» è locale: il file è nel '
          'database e non è ancora partito. Gli stati da 1 a 5 sono quelli di ANSC, e solo «3 '
          '— Inserito» consente di proseguire verso il deposito. INVIA trasmette i documenti '
          'nello stato 0; VERIFICA rilegge quelli negli stati 1 e 2; i due bottoni non sono mai '
          'presenti insieme, come in Side.'),
    ('h3', 'Tabelle sottese'),
    ('tab_uso', [
        ['Tabella', 'Uso', 'Che cosa se ne mostra'],
        ['ALLEGATI_USECASE', 'Lettura',
         'L’elenco previsto dall’UC: ty_allegato, ty_presenza, ty_configurazione, cd_logica '
         '(la condizione, quando la presenza è C…), nr_ordinamento. Filtrata per cd_usecase, '
         'id_versione e fg_operante.'],
        ['ALLEGATO', 'Lettura e scrittura',
         'I documenti dell’atto: ty_allegato, ds_tipo_allegato, nm_file, ty_file, cd_stato, '
         'id_ansc_allegato, cd_hash, fg_incluso_att_conformita. Il contenuto sta in '
         'oj_allegato.'],
        ['VALORE_DOMINIO', 'Lettura',
         'Le descrizioni dei tipi (ANSC_09), degli stati (ANSC_08) e dei formati (ANSC_10).'],
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('tab_el', [
        ['Elemento', 'Genere', 'Comportamento'],
        ['Riga a descrizione fissa', 'Elenco',
         'Viene dalla configurazione e non si aggiunge né si toglie da qui. Il segno più '
         'carica il file e crea la riga in ALLEGATO nello stato 0. L’asterisco segnala '
         'l’obbligatorietà; se la presenza è condizionata e la condizione non è soddisfatta, '
         'il documento è facoltativo.'],
        ['Descrizione + segno più', 'Campo e azione',
         'Solo nella forma a descrizione libera. La descrizione è obbligatoria e finisce in '
         'ds_tipo_allegato; il tipo è 998. «Rimuovi» toglie la riga.'],
        ['Stato', 'Colonna',
         'Il valore di cd_stato con la sua descrizione: 0 Caricato in SIPO, 1 Ricevuto, 2 In '
         'corso di verifica, 3 Inserito, 4 Rifiutato, 5 In elaborazione.'],
        ['Nell’attestazione', 'Colonna',
         'In sola lettura. Riporta fg_incluso_att_conformita, cioè la scelta che l’ufficiale '
         'fa nella pagina dell’attestazione di conformità del percorso di firma. ⚠️ Nella '
         'v0.7 il valore era proposto dalla configurazione: non lo è più [R1].'],
        ['Occhio', 'Azione di riga',
         'Anteprima. Si legge dal database negli stati 0, 1, 2 e 4; negli stati 3 e 5 si '
         'chiede ad ANSC con il servizio che restituisce il contenuto.'],
        ['Cestino', 'Azione di riga',
         'Cancella la riga in locale, in qualunque stato, e consente di caricare un altro '
         'file. ⚠️ Un documento già trasmesso resta in ANSC: il servizio non prevede la '
         'cancellazione, e che cosa ne sia è un punto aperto dell’analisi (OP-70).'],
        ['INVIA', 'Azione',
         'Visibile se almeno un documento è nello stato 0. Li trasmette tutti e registra '
         'identificativo, stato e impronta restituiti da ANSC. Richiede il codice di sessione.'],
        ['VERIFICA', 'Azione',
         'Visibile se nessun documento è nello stato 0 e almeno uno è negli stati 1 o 2. '
         'Aggiorna lo stato; sugli stati 3, 4 e 5 non fa nulla.'],
    ]),
    ('p', 'I vincoli sul file — dieci megabyte al massimo, nome senza spazi né caratteri '
          'speciali, PDF, JPG, TIF e P7M — sono dichiarati in testa alla pagina e non solo nel '
          'messaggio di errore: è la convenzione di S.I.De. e riduce i tentativi a vuoto. Due '
          'file con lo stesso nome non sono ammessi nello stesso atto.'),
]

# ════════════════════════════════════════════ 2. le tre pagine nuove
NUOVE = [
    ('h2', 'Allegati per caso d’uso'),
    ('img', ('bo_allegati_uc.png', 'Allegati per caso d’uso.')),
    ('p', 'La pagina governa per intero ALLEGATI_USECASE, cioè quali documenti ciascun caso '
          'd’uso richiede. La scheda «Allegati» del dettaglio di un UC ne è la vista filtrata '
          'su un solo caso d’uso; qui la stessa tabella si consulta e si modifica attraverso '
          'tutti, che è ciò che serve quando ANSC rivede un tipo di documento usato da molti UC '
          'o quando la configurazione si prepara per un’intera famiglia.'),
    ('p', 'La struttura e i valori sono quelli di Side [R1]: presenza O, F o C001…C004, forma '
          'BASE o D.LIB.M.01, ordine di presentazione, contrassegno di riga operante. Le '
          'aggiunte di Roma sono tre e si vedono nelle colonne: la versione della '
          'configurazione, che è un filtro obbligatorio; la descrizione del mapping, che '
          'consente di riconoscere la riga all’importazione; la condizione, presa dal catalogo '
          'delle logiche quando la presenza è condizionata.'),
    ('p', 'Si scrive soltanto nella versione in bozza. Una riga non si cancella: si disattiva, '
          'perché una versione storica deve continuare a dire che cosa si chiedeva quando era '
          'in esercizio.'),
    ('h3', 'Tabelle sottese'),
    ('tab_uso', [
        ['Tabella', 'Uso', 'Che cosa se ne mostra'],
        ['ALLEGATI_USECASE', 'Lettura e scrittura',
         'Tutte le colonne: cd_usecase, ty_allegato, ds_allegato, ty_presenza, '
         'ty_configurazione, id_dominio e cd_logica, nr_ordinamento, fg_operante, tx_note, e i '
         'campi tecnici in sola lettura. Filtrata per id_versione.'],
        ['ANSC_ANA_UC', 'Lettura per controllo',
         'La descrizione dell’UC, e la segnalazione di un UC ritirato da ANSC o assente dal '
         'catalogo: il legame è per valore e il controllo lo fa la pagina, non il database.'],
        ['VALORE_DOMINIO', 'Lettura', 'La descrizione del tipo di documento (ANSC_09).'],
        ['LOGICHE_DI_SCELTA', 'Lettura', 'Le condizioni selezionabili per la presenza C….'],
        ['ANSC_CFG_VERSIONE', 'Lettura', 'La versione in lavorazione e il suo stato.'],
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('tab_el', [
        ['Elemento', 'Genere', 'Comportamento'],
        ['Versione', 'Filtro',
         'Obbligatorio. Una versione diversa dalla bozza si consulta in sola lettura.'],
        ['Tipo allegato', 'Filtro',
         'Sul dizionario ANSC_09: trova tutti gli UC che chiedono quel documento.'],
        ['Presenza', 'Colonna e filtro',
         'O, F, C001…C004. ⚠️ Per C004 il valore del tipo di accertamento che la attiva non è '
         'più pubblicato da ANSC [R1]: la pagina segnala le righe che la usano.'],
        ['Forma', 'Colonna', 'BASE oppure D.LIB.M.01. La seconda ammette solo il tipo 998.'],
        ['Matita', 'Azione di riga',
         'Apre la riga in modifica: tipo, presenza, forma, condizione, ordine, note. La chiave '
         '— UC, tipo, versione — non si modifica: si disattiva la riga e se ne crea un’altra.'],
        ['Cestino', 'Azione di riga', 'Disattiva la riga (fg_operante a 0). Non cancella.'],
        ['NUOVA RIGA', 'Azione di pagina',
         'L’UC si sceglie dal catalogo e il tipo dal dizionario: niente a mano libera.'],
        ['CARICA DA FILE', 'Azione di pagina',
         'Apre il caricamento da file con la tabella già impostata. Si veda la pagina '
         'dedicata.'],
        ['ESPORTA CSV', 'Azione di pagina',
         'Scarica le righe filtrate nello stesso formato accettato dal caricamento, così che '
         'un file esportato si possa correggere fuori linea e ricaricare.'],
    ]),

    ('h2', 'Soggetti per caso d’uso'),
    ('img', ('bo_soggetti_uc.png', 'Soggetti per caso d’uso. Le righe sono quelle del file di '
                                   'Side per l’UC 2101.')),
    ('p', 'La pagina governa SOGGETTI_USECASE, la tabella con cui Side dichiara per ciascun '
          'caso d’uso quali persone l’atto coinvolge e con quale ruolo — dichiarante, '
          'comparente, firmatario. I ruoli non sono nel mapping di ANSC e si compilano qui '
          '[R1].'),
    ('p', f'Il contenuto non parte da zero. Il Comune di Milano ha fornito la tabella '
          f'compilata [R9]: {M(V["righe"])} righe su {V["uc"]} casi d’uso, nelle famiglie '
          f'nascita, morte, matrimonio, unione civile e cittadinanza. Caricata nella prima '
          f'bozza, dà alla configurazione di Roma il lavoro già fatto a Milano, da rivedere e '
          f'non da riscrivere.'),
    ('p', '⚠️ Il file rivela tre cose che la definizione della tabella non diceva. La prima: '
          'la colonna del tipo di soggetto riporta quasi sempre il nome del blocco nel modello '
          'evento — padre, madre, intestatari, interprete — ed è quindi il raccordo naturale '
          'con le sezioni della configurazione. La seconda: la presenza dei soggetti usa valori '
          'che per gli allegati non esistono — CSTCIV, FO, C02 — e che nessuna fonte spiega. '
          'La terza: i codici di configurazione sono '
          f'{V["configurazioni"]}, nomi di maschere di Side di cui non si ha la definizione. La '
          'pagina li accoglie e li mostra; non li valida, perché non c’è contro che cosa '
          'validarli (BO-27).'),
    ('h3', 'Tabelle sottese'),
    ('tab_uso', [
        ['Tabella', 'Uso', 'Che cosa se ne mostra'],
        ['SOGGETTI_USECASE', 'Lettura e scrittura',
         'Tutte le colonne: cd_usecase, cd_soggetto, ty_soggetto, nm_soggetto, ty_presenza, '
         'ty_configurazione, fg_dichiarante, fg_comparente, fg_firmatario, nr_ordinamento, '
         'fg_operante, tx_note, e i campi tecnici in sola lettura. Filtrata per id_versione.'],
        ['ANSC_ANA_UC', 'Lettura per controllo', 'Come nella pagina degli allegati.'],
        ['ANSC_CFG_SEZIONE', 'Lettura per controllo',
         'Segnala il soggetto la cui sezione nell’UC non risulta configurata, e viceversa: è '
         'la sovrapposizione fra le due tabelle che l’analisi chiede di governare (OP-69).'],
        ['ANSC_CFG_VERSIONE', 'Lettura', 'La versione in lavorazione e il suo stato.'],
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('tab_el', [
        ['Elemento', 'Genere', 'Comportamento'],
        ['Ruolo', 'Filtro',
         'Dichiaranti, comparenti, firmatari. È il filtro che serve per controllare il percorso '
         'di firma di un UC.'],
        ['Dich. · Comp. · Firma', 'Colonne', 'I tre contrassegni, modificabili in riga.'],
        ['Presenza', 'Colonna',
         'O, C, F e i valori non documentati, evidenziati finché la loro semantica non è '
         'chiarita.'],
        ['Matita · Cestino · NUOVA RIGA · CARICA DA FILE · ESPORTA CSV', 'Azioni',
         'Come nella pagina degli allegati per caso d’uso.'],
    ]),

    ('h2', 'Caricamento da file'),
    ('img', ('bo_carica_file.png', 'Caricamento da file. Il verbale è quello che produce il '
                                   'file dei soggetti fornito da Milano.')),
    ('p', 'Le due pagine precedenti condividono un caricamento da file, che è il modo in cui '
          'la configurazione ereditata da Side entra nel sistema di Roma e il modo in cui si '
          'applicano le correzioni in blocco. È un’operazione lunga e che scrive molte righe: '
          'per questo si prova prima di eseguirla, e il suo esito si conserva nel registro dei '
          'comandi come ogni altra operazione di servizio.'),
    ('h3', 'Il formato del file'),
    ('p', 'Il formato è quello in cui Side esporta le proprie tabelle, ricavato dal file '
          'ricevuto [R9]: non se ne inventa uno nuovo, così che un’esportazione di Milano si '
          'carichi senza essere riscritta.'),
    ('v', ('Intestazione. ', 'Prima riga con i nomi delle colonne, gli stessi della '
           'definizione della tabella. L’ordine non conta; una colonna sconosciuta è un '
           'errore, una colonna mancante lo è se la colonna è obbligatoria.')),
    ('v', ('Separatore e virgolette. ', 'Virgola. I valori possono essere fra doppi apici, e '
           'lo sono quando contengono virgole o ritorni a capo: nel file ricevuto una nota si '
           'estende su due righe.')),
    ('v', ('Valori. ', 'Contrassegni come 0 e 1; istanti nella forma AAAA-MM-GG hh:mm:ss.mmm; '
           'campo vuoto = valore nullo. I codici operatore sono di sei caratteri e possono '
           'arrivare completati da spazi.')),
    ('v', ('Colonne aggiunte da Roma. ', 'id_versione non sta nel file: la versione di '
           'destinazione si sceglie nella pagina. Le altre aggiunte (ds_allegato, id_dominio, '
           'cd_logica) sono facoltative, perché un file di Side non le ha.')),
    ('h3', 'I controlli'),
    ('p', 'Il caricamento distingue errori, che impediscono di applicarlo, e avvisi, che '
          'restano nel verbale ma non lo fermano.'),
    ('tab_uso', [
        ['Controllo', 'Specie', 'Che cosa verifica'],
        ['Intestazione', 'Errore', 'Colonne riconosciute e obbligatorie presenti.'],
        ['Chiave duplicata', 'Errore',
         'Due righe con la stessa chiave nel file: UC e tipo per gli allegati, UC e codice '
         'del soggetto per i soggetti.'],
        ['Valori dei contrassegni', 'Errore', 'Solo 0 e 1.'],
        ['Tipo di documento', 'Errore', 'Presente nel dizionario ANSC_09 (solo allegati).'],
        ['Forma e tipo', 'Errore', 'La forma D.LIB.M.01 ammette solo il tipo 998.'],
        ['UC assente dal catalogo', 'Avviso',
         'Il legame è per valore: la riga si carica e si segnala.'],
        ['Valori non documentati', 'Avviso',
         'Presenza o forma con valori fuori da quelli dichiarati da Side: si caricano e si '
         'segnalano.'],
        ['Righe non operanti', 'Informazione', 'Si caricano con il loro contrassegno.'],
    ]),
    ('h3', 'Le modalità e i campi tecnici'),
    ('v', ('Sostituisci le righe degli UC presenti nel file. ', 'È la modalità predefinita: '
           'per ciascun UC che compare nel file la versione di destinazione prende le righe del '
           'file, e quelle che non vi compaiono si disattivano. Gli UC assenti dal file restano '
           'intatti. È la forma che corrisponde al lavoro reale, che procede per caso d’uso.')),
    ('v', ('Aggiungi e aggiorna. ', 'Inserisce le righe nuove e aggiorna quelle esistenti, '
           'senza disattivare nulla. Serve per le correzioni puntuali.')),
    ('v', ('Solo simulazione. ', 'Predefinita a sì: produce il verbale degli effetti senza '
           'scrivere.')),
    ('v', ('Campi tecnici. ', 'I codici operatore e gli istanti del file non si riportano: '
           'identificano operatori di un altro ente e non dicono nulla su chi ha introdotto la '
           'riga a Roma. Si registrano chi carica e quando. Le note restano, perché documentano '
           'il perché delle scelte — nel file ricevuto, per esempio, i riallineamenti alle '
           'versioni di ANSC.')),
    ('h3', 'Il verbale del file ricevuto'),
    ('p', f'Caricato in simulazione su una bozza vuota, il file dei soggetti di Milano produce '
          f'il verbale della figura: {M(V["righe"])} righe e {V["uc"]} casi d’uso, nessuna '
          f'chiave duplicata, nessun errore. Restano tre avvisi. Un caso d’uso, '
          f'{", ".join(V["fuori_catalogo"])}, non compare nel catalogo pubblicato da ANSC. '
          f'{sum(V["non_documentati"].values())} righe usano valori di presenza non '
          f'documentati. I {V["configurazioni"]} codici di configurazione non hanno un catalogo '
          f'con cui confrontarli. Le {V["non_operanti"]} righe non operanti si caricano come '
          f'tali.'),
    ('p', '⚠️ Per ALLEGATI_USECASE non è stata fornita un’esportazione di Side: il formato '
          'qui descritto le si applica per analogia, con le colonne della sua definizione, e '
          'va confermato con Milano insieme alla richiesta del contenuto (BO-29).'),
    ('h3', 'Tabelle sottese'),
    ('tab_uso', [
        ['Tabella', 'Uso', 'Che cosa se ne mostra'],
        ['ALLEGATI_USECASE, SOGGETTI_USECASE', 'Scrittura',
         'Destinazione del caricamento, nella sola versione in bozza.'],
        ['ANSC_ANA_UC, VALORE_DOMINIO', 'Lettura per controllo',
         'Catalogo degli UC e dizionario dei tipi di documento.'],
        ['ANSC_LOG_AUDIT', 'Scrittura', 'Il verbale, come per gli altri comandi di servizio.'],
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('tab_el', [
        ['Elemento', 'Genere', 'Comportamento'],
        ['Tabella di destinazione', 'Campo',
         'ALLEGATI_USECASE o SOGGETTI_USECASE. Arrivando da una delle due pagine è già '
         'impostata.'],
        ['File', 'Campo', 'Un solo file CSV per volta.'],
        ['Versione di destinazione', 'Campo', 'Limitata alle versioni in stato BOZZA.'],
        ['Modalità', 'Campo', 'Sostituisci per UC (predefinita) oppure aggiungi e aggiorna.'],
        ['CONTROLLA', 'Azione',
         'Esegue i controlli e la simulazione e mostra il verbale. Non scrive.'],
        ['APPLICA', 'Azione',
         'Abilitata solo dopo un controllo senza errori sullo stesso file. Esegue in secondo '
         'piano; l’esito compare nel registro dei comandi.'],
    ]),
]

NUOVI_BO = [
    ('BO-27', 'Che cosa significano i valori dei soggetti che Side non documenta',
     'Nel file dei soggetti la presenza vale anche CSTCIV, FO e C02, e la forma assume '
     f'{V["configurazioni"]} codici di maschera (BASESD, NASC01, BASENOINT…). La '
     'documentazione di Side spiega soltanto i valori degli allegati. La pagina li accoglie '
     'senza validarli: per validarli, e per realizzare le maschere che quei codici '
     'designano, servono le definizioni. Si lega a OP-69 dell’analisi.'),
    ('BO-28', 'Chi può caricare da file',
     'Il caricamento sostituisce in un colpo le righe di molti casi d’uso. Va stabilito se '
     'basti l’abilitazione del funzionario configuratore o serva un profilo distinto, come '
     'per l’attivazione di una versione.'),
    ('BO-29', 'Il formato di esportazione di ALLEGATI_USECASE',
     'L’unico file ricevuto è quello dei soggetti. Il formato del caricamento si applica agli '
     'allegati per analogia; va confermato con Milano, a cui conviene chiedere insieme '
     'l’esportazione della tabella compilata, che popolerebbe la prima bozza come quella '
     'dei soggetti.'),
]

RIF = [
    ('[R9]', 'Sorgenti Documentali/Tabella SOGGETTI_USECASE.csv — esportazione della tabella '
             'dei soggetti per caso d’uso del sistema Side, compilata dal Comune di Milano',
     'Evidenza fornita'),
    ('[R10]', 'Documenti di processo «CdMI SIDe MI_ANSC»: gestione degli allegati di un atto '
              'v1.5 e gestione e firma dell’attestato di conformità v1.1',
     'Evidenza fornita'),
]

STORIA = (
    '10/10/2026', '0.8',
    'Le pagine (Allegati dell’atto; nuove Allegati per caso d’uso, Soggetti per caso d’uso, '
    'Caricamento da file) · Configurazione UC · Comandi · Mappa · Contesto tecnico · '
    'Copertura · Riferimenti · Punti aperti',
    'Riallineato il disegno alla v3.36 dell’analisi, che adotta il processo di Side per i '
    'documenti. La pagina degli allegati dell’atto passa a sei stati, con INVIA e VERIFICA '
    'mutuamente esclusivi; i documenti che l’UC non prevede seguono la forma a descrizione '
    'libera invece di un elenco separato; la colonna dell’attestazione di conformità diventa '
    'di sola lettura, perché la scelta è dell’ufficiale atto per atto. Nuove le pagine che '
    'governano per intero ALLEGATI_USECASE e SOGGETTI_USECASE e la pagina di caricamento da '
    'file, nel formato di esportazione di Side ricavato dal file dei soggetti fornito da '
    f'Milano ({M(V["righe"])} righe, {V["uc"]} casi d’uso), di cui si riporta il verbale. Le '
    'pagine passano da quindici a diciotto, le tabelle del modello coperte da sedici a '
    'diciotto. Figure dalla quindicesima in poi scalate di tre. Tre punti aperti nuovi.')


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    MOD_USO = D.trova_tabella(d, 'Tabella', 'Uso')
    MOD_EL = D.trova_tabella(d, 'Elemento', 'Genere', 'Comportamento')

    def blocchi(ancora, lista):
        for tipo, c in lista:
            if tipo == 'p':
                D.para(d, ancora, c)
            elif tipo in ('h2', 'h3'):
                D.para(d, ancora, c, stile='Heading ' + tipo[1])
            elif tipo == 'v':
                D.voce(d, ancora, *c)
            elif tipo == 'tab_uso':
                D.tabella(d, ancora, c, MOD_USO, [1.6, 1.2, 3.5])
                D.para(d, ancora, '')
            elif tipo == 'tab_el':
                D.tabella(d, ancora, c, MOD_EL, [1.6, 1.2, 3.5])
                D.para(d, ancora, '')
            elif tipo == 'img':
                png, did = c
                D.immagine(d, ancora, os.path.join(IMG, png), 6.2, 'Figura %d — ' % FIG[0] + did)
                FIG[0] += 1

    # ─────────────────────────── 0. rinumerazione: le figure da 15 in poi scalano di tre
    did = [p for p in d.paragraphs if re.match(r'^Figura \d+ —', p.text.strip())]
    assert len(did) == 24, 'attese 24 didascalie, trovate %d' % len(did)
    for p in reversed(did):
        n = int(re.match(r'^Figura (\d+) —', p.text.strip()).group(1))
        if n >= 15:
            D.testo_di(p, re.sub(r'^Figura \d+ —', 'Figura %d —' % (n + 3), p.text.strip()))
    FIG = [15]
    fatti.append('figure 15…24 rinumerate 18…27')

    # ─────────────────────────── 1. Allegati dell'atto
    elementi = D.sezione(d, 'Allegati dell’atto', 2)
    assert elementi, 'sezione Allegati dell’atto non trovata'
    for el in elementi:
        if D.ha_commenti(el):
            raise SystemExit('la sezione Allegati dell’atto porta un commento')
    # [0]=titolo [1]=figura [2]=didascalia: restano; il resto si riscrive
    did_all = Paragraph(elementi[2], d)
    assert did_all.text.startswith('Figura 11 —'), did_all.text
    D.testo_di(did_all, 'Figura 11 — Allegati dell’atto, sul processo di Side: caricamento '
                        'in SIPO, invio, verifica fino a «Inserito».')
    dopo = elementi[-1].getnext()
    for el in elementi[3:]:
        el.getparent().remove(el)
    blocchi(dopo, ALLEGATI_TESTO)
    D.sostituisci_immagine(d, 'Figura 11 —', os.path.join(IMG, 'bo_allegati.png'))
    fatti.append('«Allegati dell’atto» riscritta sul processo di Side, figura sostituita')

    # ─────────────────────────── 2. le tre pagine nuove, prima del catalogo UC
    ancora = D.h(d, 2, 'Catalogo dei casi d’uso di ANSC')._p
    blocchi(ancora, NUOVE)
    assert FIG[0] == 18
    fatti.append('nuove pagine: Allegati per caso d’uso, Soggetti per caso d’uso, '
                 'Caricamento da file (figure 15-17)')

    # ─────────────────────────── 3. configurazione UC: elenco e dettaglio
    D.sostituisci(d, 'Apre il dettaglio dell’UC con le sue quattro schede.',
                  'Apre il dettaglio dell’UC con le sue schede.', attese=1, fatti=fatti)
    D.sostituisci(d, 'Prepara nella bozza sezioni, campi, allegati e formule a partire dal '
                     'mapping ufficiale.',
                  'Prepara nella bozza sezioni, campi, allegati e formule a partire dal '
                  'mapping ufficiale. I soggetti no: il mapping non ne dichiara i ruoli, e si '
                  'caricano da file.', attese=1, fatti=fatti)
    D.sostituisci(d, 'è organizzata in quattro schede perché le quattro cose che un UC '
                     'dichiara hanno volumi molto diversi ovvero poche sezioni, decine di '
                     'campi, allegati, alcune formule.',
                  'è organizzata in schede perché le cose che un UC dichiara hanno volumi '
                  'molto diversi: poche sezioni, decine di campi, pochi allegati e soggetti, '
                  'alcune formule. Le schede «Allegati» e «Soggetti» sono le viste su un solo '
                  'UC delle pagine «Allegati per caso d’uso» e «Soggetti per caso d’uso», e ne '
                  'condividono i comportamenti.', attese=1, fatti=fatti)
    t = D.trova_tabella(d, 'Tabella', 'Uso', 'Che cosa deve essere mostrato')
    r = next(r for r in t.rows if r.cells[0].text.strip() == 'ALLEGATI_USECASE')
    D.riscrivi_cella(r.cells[2], 'Scheda «Allegati»: tipo, presenza, forma, condizione, '
                                 'ordine e contrassegno operante, per l’UC aperto.')
    D.clona_riga(t, ('SOGGETTI_USECASE', 'Lettura e scrittura',
                     'Scheda «Soggetti»: codice, tipo, nome, presenza, configurazione e i tre '
                     'contrassegni di ruolo, per l’UC aperto.'))
    D.sostituisci_immagine(d, 'Figura 14 —', os.path.join(IMG, 'bo_uc_dettaglio.png'))
    fatti.append('dettaglio UC: scheda «Soggetti», tabelle sottese, figura')

    # ─────────────────────────── 4. comandi di servizio
    D.sostituisci(d, 'lo scarico dei dizionari, l’importazione del mapping,',
                  'lo scarico dei dizionari, l’importazione del mapping, il caricamento da '
                  'file di allegati e soggetti per caso d’uso,', attese=1, fatti=fatti)
    t = next(t for t in d.tables if any(r.cells[0].text.strip().startswith(
        'ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO, ALLEGATI_USECASE') for r in t.rows))
    D.clona_riga(t, ('ALLEGATI_USECASE, SOGGETTI_USECASE', 'Scrittura',
                     'Destinazione del caricamento da file, nella sola versione in bozza.'))

    # ─────────────────────────── 5. mappa, contesto, conteggi
    t = D.trova_tabella(d, 'Area', 'Pagine', 'Chi la usa e quando')
    r = next(r for r in t.rows if r.cells[0].text.strip() == 'Configurazione')
    D.riscrivi_cella(r.cells[1], r.cells[1].text.strip() + ' · Allegati per caso d’uso · '
                                 'Soggetti per caso d’uso · Caricamento da file')
    D.sostituisci(d, 'L’applicazione è stata pensata con quindici pagine distribuite su '
                     'quattro aree funzionali, cui si aggiunge una sedicesima pagina',
                  'L’applicazione è stata pensata con diciotto pagine distribuite su quattro '
                  'aree funzionali, cui si aggiunge una diciannovesima pagina',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'Seguono le quindici pagine', 'Seguono le diciotto pagine',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'Un back-office di quindici pagine', 'Un back-office di diciotto pagine',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'ripeterle in quindici sezioni', 'ripeterle in diciotto sezioni',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'le sedici tabelle già progettate', 'le diciotto tabelle già progettate',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'da cui vengono le sedici tabelle', 'da cui vengono le diciotto tabelle',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'le sedici del modello dell’integrazione più le quattro del registro',
                  'le diciotto del modello dell’integrazione più le quattro del registro',
                  attese=1, fatti=fatti)
    D.sostituisci_immagine(d, 'Figura 1 —', os.path.join(IMG, 'bo_mfe.png'))

    # ─────────────────────────── 6. copertura
    t = D.trova_tabella(d, 'Tabella', 'Pagine che la governano', 'Copertura')
    r = next(r for r in t.rows if r.cells[0].text.strip() == 'ALLEGATI_USECASE')
    D.riscrivi_cella(r.cells[1], 'Allegati per caso d’uso · Configurazione UC (scheda '
                                 'Allegati) · Caricamento da file · Allegati dell’atto')
    D.riscrivi_cella(r.cells[2], 'Completa, compreso il caricamento da file: governata nelle '
                                 'prime tre, applicata nella quarta.')
    r = next(r for r in t.rows if r.cells[0].text.strip() == 'ALLEGATO')
    D.riscrivi_cella(r.cells[2], 'Completa. Il contrassegno dell’attestazione vi si legge e '
                                 'non vi si scrive: lo scrive il percorso di firma.')
    righe = list(t.rows)
    pos = next(i for i, r in enumerate(righe) if r.cells[0].text.strip() == 'ALLEGATI_USECASE')
    nuova = D.clona_riga(t, ('SOGGETTI_USECASE', 'Soggetti per caso d’uso · Configurazione UC '
                             '(scheda Soggetti) · Caricamento da file',
                             'Completa. Nuova in questa versione.'))
    righe[pos]._tr.addnext(nuova._tr)
    nuova = D.clona_riga(t, ('ANSC_CFG_DOMINIO_SIPO', 'Nessuna, per decisione',
                             'Si popola a mano, come stabilito nella v3.34 dell’analisi: è '
                             'l’unica tabella del modello senza una pagina, e lo è di proposito.'))
    riconc = next(r for r in t.rows if r.cells[0].text.strip() == 'RICONCILIAZ_DIZIONARI')
    riconc._tr.addnext(nuova._tr)
    fatti.append('copertura: ALLEGATI_USECASE, ALLEGATO, SOGGETTI_USECASE, '
                 'ANSC_CFG_DOMINIO_SIPO')

    # ─────────────────────────── 7. riferimenti, punti aperti, storia
    D.sostituisci(d, 'ANALISI_Integrazione-ANSC_v3.30.docx', 'ANALISI_Integrazione-ANSC_v3.36.docx',
                  attese=1, fatti=fatti, etichetta='[R1] alla v3.36')
    t = D.trova_tabella(d, '#', 'Riferimento', 'Natura')
    for r in RIF:
        D.clona_riga(t, r)
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == '#'
             and t.rows[-1].cells[0].text.strip() == 'BO-26')
    for r in NUOVI_BO:
        D.clona_riga(t, r)
    D.storia(d, *STORIA)
    fatti.append('[R1] alla v3.36, [R9] [R10], BO-27…29, storia')

    d.save(DST)
    dopo = commenti(DST)
    if dopo != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo} dopo')
    print('\n'.join(' · ' + f for f in fatti))
    print('commenti preservati:', dopo)
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))


if __name__ == '__main__':
    main()
