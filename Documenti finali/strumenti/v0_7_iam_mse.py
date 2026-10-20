# -*- coding: utf-8 -*-
"""ANALISI_Identita-Profilazione-IAM v0.6 → v0.7 (15/10/2026).

La strada è cambiata: SIPO non costruisce sipo-auth e sipo-profili, ma ADOTTA i servizi della
piattaforma a microservizi dell'ente (MSE, [F12] = Interventi_Innovazione_Progettazione_
Architettura_v1.1.docx): roma-capitale-auth (L2 nel diagramma del login) e
roma-capitale-management (L4). Il codice dei tre servizi che partecipano all'accesso è stato
letto ([F13]: auth 1.2.15, management 1.0.20, service mesh 1.0.7).

Il committente ha chiesto di valutare L4 rispondendo a tre domande, e di rivedere il cap. 6:
  a) servono interventi di manutenzione?          → cap. «Gli interventi di manutenzione»
  b) come renderlo più utile e più sicuro?        → cap. «Come rendere il servizio …»
  c) ci si integra la gestione AS-IS, con quali rischi? → cap. «L’integrazione con la gestione AS-IS di SIPO»
preceduti da un capitolo descrittivo («Il servizio di profilazione della piattaforma»).
La sicurezza ha un rilievo ma NON è il punto centrale (indicazione dell'utente).

⚠️ «token SIPO» resta il nome del token che la piattaforma emette per SIPO: rinominarlo
ovunque avrebbe reso false le frasi che descrivono che cosa deve contenere.
⚠️ «sipo-auth» compare anche nel NOME DI UN FILE in Riferimenti ([F9]) e nella Storia: lì
non si tocca.
⚠️ Le credenziali in chiaro trovate nella configurazione di management NON si riportano:
il documento dice dove sono (file e riga), non che cosa sono.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_7_iam_mse.py
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402
import diagrammi_iam      # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
IMG = os.path.join(FIN, 'strumenti', 'img')
SRC = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.6.docx')
DST = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.7.docx')

A = 'serv_roma-capitale-auth'          # prefissi brevi per le evidenze
M = 'roma-capitale-management'

# =============================================================== contenuti: capitolo 6

COMPONENTI = [
    ['Componente', 'Che cosa fa', 'Stato'],
    ['IAM di Roma Capitale (L3)',
     'Autentica la persona con OIDC. Resta l’unico luogo in cui si verificano le credenziali, '
     'SPID o CIE. Ogni applicazione della piattaforma vi è registrata come client proprio.',
     'Esistente'],
    ['roma-capitale-auth (L2)',
     'Il solo punto di autenticazione. Riceve ambito e applicazione, reindirizza a IAM con '
     'PKCE, scambia il codice, chiede i profili a management ed emette il token della '
     'piattaforma in un cookie HttpOnly.',
     'Esistente. ⚠️ Per SIPO va aggiunta l’applicazione: oggi ognuna ha la propria callback '
     'scritta nel codice (MI-10).'],
    ['roma-capitale-management (L4)',
     'Utenti, profili e abilitazioni per applicazione, in uno schema PostgreSQL proprio. È '
     'anche il registro delle applicazioni della piattaforma.',
     'Esistente. La valutazione è nel capitolo «Il servizio di profilazione della '
     'piattaforma».'],
    ['roma-capitale-service-mesh',
     'Instrada le chiamate fra i servizi risolvendo il contesto su management, e rinnova il '
     'token che sta per scadere.',
     'Esistente'],
    ['I 35 front-end Java',
     'Con la libreria ProfilazioneUtente aggiornata. Tengono in sessione LOGIN_USER — '
     'invariato — e accanto il token SIPO, mai nel browser.',
     'Si configurano, non si riscrivono'],
    ['I back-end',
     'Con RestSecurity aggiornata: verificano il token SIPO e concedono ROLE_<profiloBe>. '
     'Leggono chi opera dal token e non più dal corpo della richiesta.',
     'Una libreria sola da aggiornare'],
    ['Portale',
     'Resta il reverse proxy davanti a SIPO, ma non inietta più l’identità: passa soltanto '
     'x-real-ip e x-client-sn-sipo.',
     'Cambia configurazione'],
]

TOKEN = [
    ['Che cosa dice', 'Nel token della piattaforma, oggi', 'Che cosa serve a SIPO'],
    ['Chi è',
     'Nome utente, nome, cognome, codice fiscale e i claim di IAM.',
     'Adeguato. Il raccordo è il codice fiscale, come la piattaforma fa dal 4 marzo 2026 '
     '(decisione D3); la bonifica di ANAG_USR.CODICE_FISCALE resta il prerequisito.'],
    ['Dove opera',
     'Codice e descrizione di struttura, ufficio e tributo, scelti dall’utente dopo l’accesso.',
     '⚠️ L’ambito organizzativo deve venire dall’assegnazione, non dalla scelta di chi entra: '
     'organizzazione, sede di municipio, struttura convenzionata (MI-05).'],
    ['Che cosa può fare',
     'L’elenco dei codici di abilitazione dell’applicazione. Oltre una soglia di lunghezza, '
     'nel cookie va una versione senza abilitazioni e quella completa resta in Redis.',
     'Anche i ruoli e il profilo principale, perché i back-end decidono su un ruolo solo: '
     'nel solo StatoCivileBE le @PreAuthorize sono 393 (MI-06).'],
    ['Da dove',
     'Nessun claim.',
     'Postazione (x-client-sn-sipo) e indirizzo, perché All-Anpr dichiari ad ANPR la '
     'postazione reale.'],
    ['Per chi è valido',
     'Ambito e applicazione per cui è stato emesso.',
     'Adeguato, purché i back-end lo verifichino.'],
    ['Quanto dura',
     '480 minuti; il service mesh lo riemette quando mancano cento secondi.',
     '⚠️ La v0.6 chiedeva dieci-quindici minuti con rinnovo che rilegge i profili: oggi una '
     'revoca non ha effetto fino al nuovo accesso (MA-01, MA-02).'],
]

# =============================================================== capitolo L4: che cosa è

NUMERI = [
    ['Che cosa', 'Nel codice'],
    ['Operazioni esposte', '78, di cui 16 interne: riservate agli altri servizi e, come '
     'si vedrà, prive di autenticazione.'],
    ['Entità', '17, in uno schema PostgreSQL proprio.'],
    ['Pila', 'Java 21, Spring Boot 3.3.4, JPA con le interrogazioni in file separati.'],
    ['Test', 'Uno per servizio, che verifica soltanto l’avvio.'],
    ['Ultimo rilascio', 'Versione 1.0.20, del 20 marzo 2026.'],
]

ACCESSO = [
    ['Passo', 'Che cosa accade', 'Evidenza'],
    ['1 · Ingresso',
     'Il front-end chiama l’operazione di accesso con ambito e applicazione; '
     'roma-capitale-auth costruisce la richiesta a IAM con PKCE e risposta in form_post.',
     'LoginV1 /loginIAM'],
    ['2 · Ritorno',
     'Una callback per applicazione, scritta nel codice. Il codice di autorizzazione si '
     'scambia con i token di IAM usando il segreto dell’applicazione.',
     'LoginV1, LoginDomain.createAccessToken'],
    ['3 · Chi è',
     'Per i dipendenti l’utente si cerca in management per codice fiscale.',
     'LoginDomain.java:419'],
    ['4 · Che cosa può fare',
     'Management unisce le abilitazioni dei profili dell’applicazione, o delle applicazioni '
     'figlie se l’autorizzazione è di gruppo. Senza profili, le applicazioni di back-office '
     'negano l’accesso.',
     'UtenteDomain.java:304'],
    ['5 · Utente inattivo',
     'Oltre sei mesi dall’ultimo accesso l’utente viene disattivato al tentativo successivo.',
     'LoginDomain.java:147'],
    ['6 · Il token',
     'Firmato dalla piattaforma, dura 480 minuti, viaggia in un cookie HttpOnly sul dominio '
     'dell’ente. Oltre la soglia la copia completa va in Redis.',
     'LoginDomain.generateTokenAndSetProfile'],
    ['7 · Dove opera',
     'Struttura, ufficio e tributo si scelgono dopo l’accesso e finiscono nel token. Il '
     'server li confronta con le assegnazioni soltanto per lo sportello unico.',
     'LoginDomain.selezionaStruttura'],
    ['8 · Rinnovo',
     'L’operazione di rinnovo rilegge l’utente ma conserva le abilitazioni del token '
     'precedente; il service mesh riemette il token quando mancano cento secondi.',
     'LoginDomain.java:759 · AuthenticationFilter.java:112'],
]

GESTIONE = [
    ['Operazione', 'Come è realizzata', 'Che cosa ne segue'],
    ['Inserire un utente',
     'Controlla che codice fiscale e nome utente non esistano, richiede almeno un profilo, '
     'cifra la password se è indicata.',
     'Corretta. La password è cifrata in modo reversibile, non conservata come impronta.'],
    ['Modificare un utente e i suoi profili',
     'Aggiorna l’utente e, se arriva un elenco di profili, cancella tutte le assegnazioni e '
     'le reinserisce (UtenteDomain.java:240).',
     'Si perde chi ha concesso che cosa e quando: l’autore di ogni riga diventa l’ultima '
     'persona che ha salvato.'],
    ['Cancellare un utente',
     'Cancella le assegnazioni, poi l’utente (UtenteDomain.java:269).',
     'Cancellazione fisica: dell’utente non resta traccia.'],
    ['Creare e modificare un profilo',
     'Stessa tecnica: cancella le abilitazioni del profilo e le reinserisce.',
     'Nessun controllo che le abilitazioni appartengano all’applicazione del profilo.'],
    ['Cancellare un profilo',
     'Cancella le assegnazioni a tutti gli utenti, poi il profilo (ProfiloDomain.java:164).',
     'Toglie l’accesso a tutti, senza traccia.'],
    ['Cancellare un’abilitazione',
     'Cancella la riga (AbilitazioneDomain.java:178).',
     'Le associazioni ai profili restano senza oggetto [DA VERIFICARE: se un vincolo del '
     'database lo impedisce, l’operazione fallisce invece].'],
    ['Attribuire un nodo organizzativo',
     'Non esiste: l’assegnazione si crea con la sola chiave utente-profilo.',
     'Il nodo si può creare ma non attribuire a nessuno.'],
    ['Chi può fare tutto questo',
     'Nessuna operazione dichiara un’autorità richiesta: nei due servizi non compare alcuna '
     '@Secured né @PreAuthorize.',
     'Se il filtro del modulo comune non controlla le abilitazioni per percorso (PI-31), '
     'qualunque utente autenticato di qualunque applicazione gestisce utenti e profili.'],
]

# =============================================================== manutenzione (a)
# (#, intervento, evidenza, specie, priorità)
MANUTENZIONE = [
    ('MA-01', 'Il rinnovo del token non rilegge le abilitazioni: rilegge l’utente, poi rimette '
              'nel token quelle del token precedente. Una revoca non ha effetto finché '
              'l’utente non rifà l’accesso.',
     'LoginDomain.java:759', 'Difetto', 'Alta'),
    ('MA-02', 'Il service mesh riemette da sé il token in scadenza, con le stesse abilitazioni '
              'e senza controllare che l’utente sia ancora attivo: una sessione si estende '
              'senza limite.',
     'AuthenticationFilter.java:112 (service mesh)', 'Difetto', 'Alta'),
    ('MA-03', 'Il nodo organizzativo dell’assegnazione non è scritto da alcuna operazione, e le '
              'assegnazioni per struttura dello sportello unico non hanno percorso di '
              'scrittura.',
     'UtenteDomain.createUtentePerProfilo · UtentePerProfiloGet.xml', 'Difetto', 'Media'),
    ('MA-04', 'Il servizio che restituisce all’operatore le sue funzioni e la sua visibilità '
              'sui dati restituisce la visibilità sempre vuota.',
     'OperatoreDomain.java:161', 'Difetto', 'Media'),
    ('MA-05', 'Cancellare un profilo toglie l’accesso a tutti gli utenti senza traccia; '
              'cancellare un’abilitazione lascia le associazioni ai profili.',
     'ProfiloDomain.java:164 · AbilitazioneDomain.java:178', 'Difetto', 'Media'),
    ('MA-06', 'La modifica di un utente cancella e reinserisce tutte le sue assegnazioni: '
              'autore e data della concessione si perdono.',
     'UtenteDomain.java:240', 'Difetto', 'Media'),
    ('MA-07', 'Le abilitazioni arrivano nel token come codici senza applicazione: con '
              'l’autorizzazione di gruppo due applicazioni figlie con lo stesso codice non si '
              'distinguono.',
     'UtenteDomain.getAbilitazioniFromProfiloList', 'Difetto', 'Media'),
    ('MA-08', 'Logiche di una sola applicazione nel servizio comune: la struttura predefinita '
              'dell’amministratore dello sportello unico è scritta nel codice, come il ramo che '
              'la gestisce.',
     'LoginDomain.java:430', 'Difetto', 'Bassa'),
    ('MA-09', 'Nel login fra sistemi un’eccezione è ignorata e il risultato nullo usato subito '
              'dopo: un client inesistente produce un errore non gestito invece di «client non '
              'valido».',
     'LoginDomain.java:254', 'Difetto', 'Bassa'),
    ('MA-10', 'Dichiarazioni incoerenti nel modello: la descrizione dell’abilitazione è di '
              'dieci caratteri, e la password è obbligatoria anche per chi entra soltanto da '
              'IAM [DA VERIFICARE sulla DDL].',
     'AbilitazioneEntity.java:40 · UtenteEntity.java:36', 'Difetto', 'Bassa'),
    ('MA-11', 'L’operazione di accesso dichiarata superata è ancora pubblicata: emette un token '
              'per qualunque nome utente, senza password.',
     'LoginV1.java:80', 'Igiene', 'Alta'),
    ('MA-12', 'I percorsi interni di auth e di management sono pubblici, compreso l’accesso '
              'delle utenze tecniche senza credenziali; il service mesh, se una chiamata interna '
              'arriva senza token valido, vi appone il proprio.',
     'SecurityConfiguration.java:48 (auth), :27 (management) · ForwardDomain.java:109',
     'Igiene', 'Alta'),
    ('MA-13', 'Codici di autorizzazione, token di IAM e token rifiutati finiscono nei registri '
              'applicativi.',
     'LoginDomain.java:337, 364 · LoginV1.java:580', 'Igiene', 'Alta'),
    ('MA-14', 'La configurazione versionata di management contiene in chiaro la password del '
              'database, quella di Redis e la chiave privata con cui si firmano i token. Il '
              'documento non le riporta.',
     'application.properties di management, righe 11, 33, 74', 'Igiene', 'Alta'),
    ('MA-15', 'Controlli OIDC incompleti: il nonce si genera ma non si verifica, manca il '
              'parametro di stato, e la decodifica di un token di IAM non ne controlla il '
              'destinatario.',
     'LoginV1 /loginIAM · AbstractLoginApi.java:173', 'Igiene', 'Media'),
    ('MA-16', 'Le origini ammesse comprendono «null» con le credenziali, e l’operazione di '
              'decodifica restituisce il token nel corpo della risposta.',
     'SecurityConfiguration.java:83 (auth) · LoginV1.java:594', 'Igiene', 'Media'),
    ('MA-17', 'Il tracciamento delle operazioni nel service mesh è commentato, mentre il '
              'documento della piattaforma lo dichiara attivo [F12] §6.5.',
     'GeneralFilterServiceMesh.java:188', 'Strumento', 'Media'),
    ('MA-18', 'Un solo test per servizio, che verifica l’avvio; README di una o due righe in '
              'auth e nel service mesh; il controllo delle vulnerabilità delle librerie è '
              'commentato nei due pom, mentre il documento della piattaforma lo dichiara in uso '
              '[F12] §8.',
     'pom.xml di auth e service mesh · src/test', 'Strumento', 'Media'),
]

# =============================================================== miglioramenti (b)
MIGLIORAMENTI = [
    ['#', 'Capacità', 'Oggi', 'Che cosa proporre', 'Perché serve anche a SIPO'],
    ['MI-01', 'Storia delle assegnazioni', 'Cancellazione e reinserimento',
     'Validità dal-al e stato su ogni assegnazione; una revoca chiude la riga invece di '
     'cancellarla; interrogazione «chi aveva che cosa a una data».',
     'Per un sistema di stato civile la domanda «chi poteva farlo quel giorno» è un '
     'adempimento, non una curiosità.'],
    ['MI-02', 'Ciclo della concessione', 'Assente',
     'Richiesta, approvazione, attivazione, revoca con motivo; chi chiede non è chi concede.',
     'Oggi in SIPO l’abilitazione è un’azione diretta dell’amministratore: il processo, se '
     'c’è, sta fuori dal sistema.'],
    ['MI-03', 'Amministrazione per ambito', 'Nessun controllo',
     'Un amministratore d’ambito o d’applicazione concede solo i profili della propria '
     'applicazione; le operazioni di gestione richiedono un’abilitazione di amministrazione.',
     'Senza, l’amministratore di un’altra area può cambiare gli accessi allo stato civile.'],
    ['MI-04', 'Operazioni composte', 'Tre operazioni elementari per ogni cambiamento',
     'Copia dei profili da un collega; cambio di ufficio come operazione sola; cessazione '
     'che disattiva e revoca tutto; assegnazione a più utenti.',
     'Ingressi, spostamenti e uscite sono il lavoro ordinario di chi amministra.'],
    ['MI-05', 'Il nodo organizzativo', 'Previsto, non attribuibile',
     'Operazione che attribuisce il nodo; nodo e visibilità nel token o in una lettura '
     'dedicata; la struttura scelta dopo l’accesso verificata sulle assegnazioni.',
     'È l’ambito organizzativo di SIPO — struttura, sede di municipio, struttura '
     'convenzionata — che oggi nessuno applica lato server (RI-08).'],
    ['MI-06', 'Il profilo principale', 'Assente',
     'Un profilo principale per utente e applicazione, portato nel token insieme ai ruoli.',
     'I back-end di SIPO decidono su un ruolo solo: senza, il ruolo che conta torna a '
     'dipendere dall’ordinamento (RI-16).'],
    ['MI-07', 'Revisione periodica degli accessi', 'Assente',
     'Elenchi per ambito, profilo e nodo; profili mai usati; profili incompatibili fra loro.',
     'Rende sostenibile nel tempo un catalogo che cresce con ogni applicazione.'],
    ['MI-08', 'Catalogo dichiarato dall’applicazione', 'Inserimenti e cancellazioni a mano',
     'Le abilitazioni le pubblica l’applicazione al rilascio, da un elenco versionato con il '
     'codice; la cancellazione è impedita se l’abilitazione è in uso.',
     'Un’abilitazione esiste perché un controllo nel codice la cita: creata o cancellata a '
     'mano, catalogo e codice divergono.'],
    ['MI-09', 'Revoca che abbia effetto', 'Vale al prossimo accesso',
     'Il rinnovo rilegge le abilitazioni da management; invalidazione delle sessioni '
     'dell’utente; durata del token commisurata.',
     'Con 480 minuti e il rinnovo automatico, un operatore rimosso lavora per tutta la '
     'giornata.'],
    ['MI-10', 'Applicazioni configurabili', 'Una callback e una voce di elenco nel codice '
     'per ciascuna',
     'Una callback sola, parametrizzata da ambito e applicazione e letta dal registro delle '
     'applicazioni che management già tiene.',
     'SIPO entra con una configurazione invece che con un rilascio della piattaforma.'],
    ['MI-11', 'Eventi di gestione', 'Assenti',
     'Ogni concessione e revoca pubblica un evento sul gestore degli eventi della piattaforma '
     '[F12] §6.3.',
     'È il canale con cui la proiezione su ANAG_USR e il tracciamento restano allineati.'],
    ['MI-12', 'Logiche di una sola applicazione', 'Nel servizio comune',
     'Ricondotte a configurazione, o al servizio dell’applicazione che le richiede.',
     'Ogni applicazione che entra ne aggiungerebbe altre: SIPO compresa.'],
]

SICUREZZA = [
    ['Misura', 'Che cosa copre'],
    ['I percorsi interni non sono pubblici',
     'Autenticazione fra servizio e servizio e una regola di rete che li renda raggiungibili '
     'solo dall’interno. Copre MA-12 (PI-32).'],
    ['La chiave privata sta solo nell’emittente',
     'Gli altri servizi verificano con la chiave pubblica. Oggi la chiave di firma è nella '
     'configurazione di management e il service mesh conia token: chi ha la chiave può '
     'emettere qualunque abilitazione (MA-14, PI-31).'],
    ['Un solo servizio emette i token',
     'Il rinnovo nel service mesh diventa una chiamata all’emittente, che rilegge le '
     'abilitazioni (MA-02, MI-09).'],
    ['Il token non torna nelle risposte',
     'Altrimenti la libreria Angular lo conserva nel browser e il cookie HttpOnly perde il suo '
     'scopo (MA-16).'],
    ['Le password, dove servono, come impronta',
     'Oggi sono cifrate in modo reversibile con una chiave unica: chi la ha, le legge tutte.'],
]

# =============================================================== integrazione (c)
CORRISPONDENZE = [
    ['SIPO oggi (ANAG_USR)', 'roma-capitale-management', 'Come si porta'],
    ['UTENTI (NOME_UTENTE, CODICE_FISCALE)', 'utente (ds_username, cd_fiscale)',
     'Per codice fiscale, che management tratta come unico. In ANAG_USR la colonna è '
     'facoltativa e non univoca: la bonifica (PI-11) è prerequisito anche qui.'],
    ['CONF_RUOLI (RUOLO, DESCRIZIONE)', 'profilo (cd_ruolo_profilo, ds_profilo)',
     'Un profilo per ruolo e per applicazione SIPO, con il nome tecnico del ruolo nel codice '
     'di ruolo.'],
    ['CONF_APP_ROLE.PROFILO_BE', 'nessuna colonna',
     'Si ricava dal codice di ruolo, o diventa un parametro della configurazione '
     'dell’applicazione. È un dettaglio del vecchio mondo che il nuovo non deve conoscere.'],
    ['CONF_AREE_TEMATICHE', 'applicazione figlia',
     'SIPO applicazione padre con autorizzazione di gruppo, un’applicazione figlia per area. '
     'Il codice dell’applicazione ha al massimo dieci caratteri (PI-35).'],
    ['CONF_FUNZIONALITA', 'abilitazione',
     'Una per funzionalità e area. ⚠️ Il codice arriva nel token senza l’applicazione: va '
     'reso univoco su tutto SIPO, non solo nell’area (MA-07).'],
    ['R_UTENTI_RUOLI (utente × ruolo × area × funzionalità)',
     'utente_per_profilo + profilo_per_abilitazione',
     'Non c’è una corrispondenza diretta: vedi il sottocapitolo seguente.'],
    ['Organizzazione, sede di municipio, struttura convenzionata',
     'nodo_struttura e visibilita',
     'La forma c’è, la scrittura no (MA-03, MI-05).'],
    ['La schermata di gestione degli utenti di AnagrafeWeb',
     'Il front-end di amministrazione della piattaforma',
     'Non è fra le sorgenti: va visto prima di decidere (PI-36).'],
]

MODALITA = [
    ['Modalità', 'Come funziona', 'A favore', 'Contro'],
    ['A — SIPO resta la fonte, la piattaforma legge',
     'Gli accessi si amministrano in AnagrafeWeb come oggi; una sincronizzazione porta utenti '
     'e profili in management, che per l’applicazione SIPO è in sola lettura.',
     'Nessun cambiamento per chi amministra né per i trentacinque front-end. Si può fare '
     'subito, senza attendere i miglioramenti.',
     'La gestione resta quella di oggi, con i suoi difetti (RI-15). La quaterna va comunque '
     'tradotta in profili per poterla sincronizzare. Due copie da tenere allineate.'],
    ['B — la piattaforma è la fonte, SIPO legge la proiezione',
     'Si amministra in management; R_UTENTI_RUOLI diventa una proiezione alimentata dagli '
     'eventi di gestione (MI-11), che i trentacinque front-end continuano a leggere.',
     'Una sola gestione per l’ente. Storia, amministrazione per ambito e revisione valgono '
     'anche per SIPO.',
     'Richiede prima MI-01, MI-03, MI-05 e MI-06. ANAG_USR dipende da un servizio esterno per '
     'essere aggiornata.'],
    ['C — due gestioni parallele',
     'Si amministra in entrambi i posti.',
     '—',
     'Esclusa: due verità sugli stessi accessi divergono alla prima urgenza, e non c’è modo '
     'di sapere quale delle due abbia ragione.'],
]

RISCHI = [
    ['Rischio', 'Probabilità', 'Impatto', 'Come si contiene'],
    ['Le eccezioni della quaterna moltiplicano i profili', 'Alta', 'Medio',
     'Misurarle prima (PI-14); soglie del dieci e del novanta per cento; report di differenza '
     'letto prima del passaggio.'],
    ['Un operatore rimosso continua a operare', 'Alta', 'Alto',
     'MA-01, MA-02 e MI-09 prima che SIPO entri.'],
    ['L’ambito organizzativo non arriva al server: un operatore vede dati di un altro '
     'municipio', 'Media', 'Alto',
     'MI-05. Fino ad allora il filtro resta quello dei front-end di oggi, con il rischio già '
     'registrato (RI-08).'],
    ['Copia e fonte divergono', 'Media', 'Alto',
     'Un solo posto in cui si scrive, per applicazione; quadratura periodica fra le due.'],
    ['Il ruolo che conta cambia per chi ha più profili', 'Media', 'Alto',
     'Profilo principale (MI-06) e ruoli nuovi con identificativo alto, come già previsto nel '
     'capitolo dei ruoli.'],
    ['Codici fiscali mancanti o duplicati in ANAG_USR', 'Alta', 'Alto',
     'Bonifica (PI-11) prima della prima sincronizzazione: management rifiuta il duplicato '
     'all’inserimento.'],
    ['SIPO eredita le esposizioni della piattaforma', 'Media', 'Critico',
     'Percorsi interni chiusi e chiave di firma confinata prima dell’ingresso (MA-12, MA-14, '
     'PI-31, PI-32).'],
    ['Priorità e tempi di un’altra area', 'Alta', 'Medio',
     'Un accordo di governo con il titolare della piattaforma sulle richieste di SIPO (PI-33).'],
    ['Dimensionamento', 'Media', 'Medio',
     'Volumi di SIPO dichiarati al titolare della piattaforma: auth, management e service mesh '
     'sono sul percorso di ogni accesso e di ogni chiamata [DA VERIFICARE: volumi].'],
    ['Un database diverso sul percorso dell’accesso', 'Bassa', 'Medio',
     'PostgreSQL entra nel perimetro di esercizio di SIPO: alta affidabilità e salvataggi vanno '
     'verificati con chi lo gestisce.'],
]

PUNTI_NUOVI = [
    ('PI-31', 'Contenuto del modulo comune della piattaforma (roma-capitale-common): con quale '
              'algoritmo e quale chiave si firmano e si verificano i token, e se il filtro di '
              'autenticazione controlli le abilitazioni per percorso. Decide se la gestione dei '
              'profili sia aperta a qualunque utente autenticato e chi possa coniare token.',
     'Titolare della piattaforma'),
    ('PI-32', 'Se i percorsi interni di auth e di management, e il percorso interno del service '
              'mesh, siano raggiungibili dall’esterno del cluster o da altri spazi dei nomi. Da '
              'questa risposta dipende la gravità reale di MA-12.',
     'Sistemi / Titolare della piattaforma'),
    ('PI-33', 'Chi è il titolare della piattaforma, con quale contratto accoglie le richieste di '
              'manutenzione e di evoluzione di un’area diversa da quella per cui è nata, e con '
              'quale impegno di esercizio.',
     'Dipartimento'),
    ('PI-34', 'Come si popolano oggi, in esercizio, il nodo organizzativo delle assegnazioni e le '
              'assegnazioni per struttura dello sportello unico, che nessuna operazione scrive.',
     'Titolare della piattaforma'),
    ('PI-35', 'Codici di ambito e di applicazione per SIPO, e granularità: un’applicazione padre '
              'con le aree tematiche come figlie, oppure altro. I codici sono di due e di dieci '
              'caratteri.',
     'Architettura / Referenti SIPO'),
    ('PI-36', 'Il codice del front-end di amministrazione della piattaforma, che non è fra le '
              'sorgenti: può contenere regole di gestione che il back-end non ha, e va visto '
              'prima di scegliere la modalità di integrazione.',
     'Titolare della piattaforma'),
    ('PI-37', 'I valori di esercizio dei cookie e delle origini ammesse, e la durata del token in '
              'produzione. Il file versionato porta i valori di sviluppo; quelli di collaudo, '
              'indicati nei commenti, sono SameSite=None e Secure attivo.',
     'Titolare della piattaforma / Sistemi'),
]


# =============================================================== utilità

def sostituisci_nei_run(p, vecchio, nuovo):
    """Sostituisce dentro i run, così grassetti e corsivi restano; ripiega su testo_di."""
    n = 0
    for r in p.runs:
        if vecchio in r.text:
            n += r.text.count(vecchio)
            r.text = r.text.replace(vecchio, nuovo)
    if vecchio in p.text:                      # spezzato fra due run
        n += p.text.count(vecchio)
        D.testo_di(p, p.text.replace(vecchio, nuovo))
    return n


def paragrafo(d, inizio):
    trovati = [p for p in d.paragraphs if p.text.strip().startswith(inizio)]
    if len(trovati) != 1:
        raise SystemExit(f'attesi 1 paragrafo «{inizio[:50]}», trovati {len(trovati)}')
    return trovati[0]


def main():
    if os.path.exists(DST):
        os.remove(DST)                         # solo un'esecuzione precedente di questo script
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    modello = d.tables[3]
    fatti = []
    corpo = d.element.body
    fine = corpo[-1]
    assert fine.tag == qn('w:sectPr'), 'il corpo non termina con sectPr'

    png_mse = os.path.join(IMG, diagrammi_iam.mse(IMG))
    png_mgmt = os.path.join(IMG, diagrammi_iam.management(IMG))

    def inserisci(prima, blocchi):
        for tipo, contenuto in blocchi:
            if tipo == 'p':
                D.para(d, prima, contenuto)
            elif tipo == 'b':
                D.para(d, prima, pezzi=D.segmenta('', contenuto))
            elif tipo == 'v':
                D.voce(d, prima, *contenuto)
            elif tipo in ('h1', 'h2', 'h3'):
                D.para(d, prima, contenuto, stile='Heading ' + tipo[1])
            elif tipo == 'tab':
                righe, larghezze = contenuto
                D.tabella(d, prima, righe, modello, larghezze)
            elif tipo == 'img':
                png, pollici, didascalia = contenuto
                D.immagine(d, prima, png, pollici, didascalia)

    def ancora(titolo, livello=2):
        return D.h(d, livello, titolo)._p

    # ------------------------------------------------ Scopo, Glossario, Riferimenti
    inserisci(ancora('Perimetro e metodo'), [
        ('b', 'Dalla versione 0.7 il documento valuta anche il **servizio di profilazione della '
              'piattaforma a microservizi** che l’ente ha scelto di adottare [F12], e risponde '
              'alle tre domande poste dal committente: se richieda interventi di manutenzione, '
              'come renderlo più utile e più sicuro, e se la gestione dei profili che SIPO ha '
              'oggi vi si possa integrare, e con quali rischi.'),
    ])
    glo = D.trova_tabella(d, 'Termine', 'Significato')
    for riga in (
            ('MSE', 'La piattaforma a microservizi dell’area Entrate e Sportelli Unici di Roma '
                    'Capitale [F12], di cui SIPO adotta i servizi di autenticazione e di '
                    'profilazione.'),
            ('L2, L3, L4', 'Le sigle con cui il diagramma del login della piattaforma indica '
                           'roma-capitale-auth (L2), IAM (L3) e roma-capitale-management (L4).'),
            ('Abilitazione', 'Nella piattaforma, il diritto elementare di un’applicazione, citato '
                             'per codice nel token. Corrisponde alla funzionalità di SIPO.'),
            ('Profilo della piattaforma', 'Un insieme di abilitazioni di un’applicazione, con un '
                                          'codice di ruolo. Corrisponde al ruolo di SIPO.'),
            ('Nodo', 'Un elemento dell’albero organizzativo della piattaforma, che porta le '
                     'regole di visibilità sui dati.')):
        D.clona_riga(glo, riga)
    rif = D.trova_tabella(d, '#', 'Riferimento', 'Natura')
    D.clona_riga(rif, ('[F12]', 'Interventi_Innovazione_Progettazione_Architettura_v1.1.docx — '
                       '«Documento di Progettazione Architettura. Interventi di Innovazione», '
                       'v1.1 del 30/04/2025: la piattaforma MSE dell’area Entrate e Sportelli '
                       'Unici', 'Documento di piattaforma'))
    D.clona_riga(rif, ('[F13]', 'Sorgenti di roma-capitale-auth 1.2.15, '
                       'roma-capitale-management 1.0.20 e roma-capitale-service-mesh 1.0.7, '
                       'rilasciati fra il 16 e il 24 marzo 2026', 'Sorgente applicativo'))
    fatti.append('Scopo, 5 voci di glossario, riferimenti [F12] e [F13]')

    # sipo-auth / sipo-profili nel testo ESISTENTE, ma non nella Storia né nei Riferimenti.
    # ⚠️ Va fatto PRIMA di inserire il testo nuovo, che nomina sipo-auth di proposito.
    esclusi = {D.trova_tabella(d, 'versione', 'sintesi dei cambiamenti')._tbl, rif._tbl}
    n_auth = n_prof = 0
    for p in d.paragraphs:
        n_auth += sostituisci_nei_run(p, 'sipo-auth', 'roma-capitale-auth')
        n_prof += sostituisci_nei_run(p, 'sipo-profili', 'roma-capitale-management')
    for t in d.tables:
        if t._tbl in esclusi:
            continue
        visti = set()
        for r in t.rows:
            for c in r.cells:
                if c._tc in visti:
                    continue
                visti.add(c._tc)
                for p in c.paragraphs:
                    n_auth += sostituisci_nei_run(p, 'sipo-auth', 'roma-capitale-auth')
                    n_prof += sostituisci_nei_run(p, 'sipo-profili', 'roma-capitale-management')
    fatti.append(f'sipo-auth → roma-capitale-auth: {n_auth}; sipo-profili → management: {n_prof}')


    # ------------------------------------------------ Executive summary
    inserisci(ancora('Parte I — Lo stato di fatto', 1), [
        ('b', 'Dalla versione 0.7 la strada è cambiata in un punto: **l’autenticazione e i '
              'profili non si costruiscono, si adottano** i servizi della piattaforma a '
              'microservizi dell’ente [F12] — roma-capitale-auth per l’accesso, '
              'roma-capitale-management per i profili. Il codice di entrambi è stato letto '
              '[F13], e alle tre domande del committente sul servizio dei profili la risposta è '
              'questa.'),
        ('v', ('Servono interventi di manutenzione? ',
               'Sì: **%d**, di cui **%d** di priorità alta. Il più rilevante per l’ente è che '
               'una revoca non ha effetto finché l’utente non rifà l’accesso.'
               % (len(MANUTENZIONE), sum(1 for m in MANUTENZIONE if m[4] == 'Alta')))),
        ('v', ('Come renderlo più utile e più sicuro? ',
               'Facendo della gestione un processo invece che un insieme di operazioni sulle '
               'tabelle: storia delle assegnazioni, ciclo della concessione, amministrazione per '
               'ambito, nodo organizzativo, profilo principale. La sicurezza si copre con '
               'cinque misure.')),
        ('v', ('Ci si può integrare? ',
               'Sì, in due tempi: prima SIPO resta la fonte e la piattaforma legge, poi la '
               'piattaforma diventa la fonte. Il rischio principale è la forma del permesso di '
               'SIPO, che è per persona e va ricondotta a profili.')),
    ])
    fatti.append('Executive summary: le tre risposte')

    # ------------------------------------------------ Parte II, S4: la strada è questa
    inserisci(ancora('Il confronto in una tabella', 3), [
        ('b', '⚠️ **Dalla versione 0.7 questa è la strada intrapresa.** Il servizio è '
              'roma-capitale-auth della piattaforma MSE [F12], e le tre incognite della tabella '
              'sono state verificate sul codice [F13]: copre i dipendenti, cercati per codice '
              'fiscale; il token porta le abilitazioni dell’applicazione; è verificabile dai '
              'back-end con la chiave della piattaforma (PI-31). Il capitolo «Scelte effettuate» '
              'ne trae le conseguenze.'),
    ])

    # ------------------------------------------------ Capitolo 6
    p = paragrafo(d, 'La scelta è stata compiuta dal referente tecnico')
    D.testo_di(p, 'La scelta è stata compiuta dal referente tecnico del progetto ed è descritta in '
                  '[F9]: un solo punto di autenticazione, che parla con IAM ed emette un token '
                  'proprio. Dalla versione 0.7 quel punto ha un nome diverso: la strada '
                  'intrapresa è la piattaforma MSE [F12], che il punto di autenticazione e il '
                  'servizio dei profili li ha già, in esercizio per altre aree dell’ente. Questo '
                  'capitolo recepisce la scelta in termini operativi: che cosa si adotta, come si '
                  'arriva allo stato finale e — poiché il front-end di destinazione è Angular e il '
                  'back-end è Java — che cosa va verificato perché regga su quell’architettura. La '
                  'valutazione del servizio dei profili è nei capitoli che seguono quello dei '
                  'ruoli.')

    sez = D.sezione(d, 'La soluzione: un solo punto di autenticazione', 2)
    assert sez, 'sezione della soluzione non trovata'
    testa = sez[0]
    D.testo_di(Paragraph(testa, d), 'La soluzione: il punto di autenticazione della piattaforma')
    dopo = sez[-1].getnext()
    for el in sez[1:]:
        el.getparent().remove(el)
    inserisci(dopo, [
        ('b', 'Il punto di autenticazione di SIPO **non si costruisce**: è roma-capitale-auth, '
              'il componente della piattaforma MSE che le applicazioni dell’area Entrate e '
              'Sportelli Unici già usano [F12] [F13], e che nella sequenza del login della '
              'piattaforma compare come L2. I profili non li legge un servizio nuovo sopra '
              'ANAG_USR: li restituisce **roma-capitale-management** (L4), che tiene utenti, '
              'profili e abilitazioni in uno schema PostgreSQL proprio. Il ruolo che le versioni '
              'precedenti assegnavano a sipo-auth e a sipo-profili è dunque già coperto da due '
              'servizi esistenti: il lavoro cambia natura, da costruzione ad adozione e '
              'adeguamento.'),
        ('b', 'La libreria Angular dei nuovi front-end è il modulo comune della stessa '
              'piattaforma. Il percorso /msAuth/api/v1/autenticazione/loginIAM a cui rinvia '
              'coincide con l’operazione di accesso di roma-capitale-auth, che però richiede il '
              'codice di ambito e quello dell’applicazione — gli stessi che il commento della '
              'guardia dichiara mancanti. **Il servizio che le versioni precedenti trattavano '
              'come un’incognita è identificato** (S4, PI-17, PI-22).'),
        ('img', (png_mse, 6.4, 'Figura 6 — L’architettura adottata. I numeri sono i passi del '
                               'login; in verde ciò che esiste ed è adeguato, in giallo ciò che '
                               'esiste e va adeguato per SIPO, in rosso il lato SIPO.')),
        ('p', 'I componenti.'),
        ('tab', (COMPONENTI, [1.6, 3.4, 1.8])),
        ('b', 'Nel seguito si continua a chiamare **token SIPO** il token che la piattaforma '
              'emette per le applicazioni di SIPO. Che cosa contiene oggi, e che cosa deve '
              'contenere.'),
        ('tab', (TOKEN, [1.2, 2.6, 3.0])),
        ('b', 'Due proprietà che la versione 0.6 dava per acquisite **non sono garantite oggi** '
              'dalla piattaforma e restano requisiti per SIPO: la firma verificabile con la '
              'chiave pubblica da chi non può emettere — oggi la chiave privata è nella '
              'configurazione di più servizi (PI-31) — e il gettone che non raggiunge mai il '
              'browser — oggi viaggia in un cookie HttpOnly, ma l’operazione di decodifica lo '
              'restituisce nel corpo della risposta e la libreria Angular lo conserva. Sono '
              'trattate come interventi, nei capitoli che seguono, non come presupposti.'),
    ])
    fatti.append('cap. 6: sezione della soluzione riscritta (figura 6 nuova, 2 tabelle)')

    p = paragrafo(d, 'Si introduce pertanto un componente sul percorso del login')
    D.testo_di(p, 'Il componente sul percorso del login esiste già, ma resta il punto in cui un '
                  'disservizio ferma tutto, e con SIPO ferma anche lo stato civile. Per l’ingresso '
                  'di SIPO vanno verificate almeno due repliche di auth, di management e del '
                  'service mesh, le chiavi e le sessioni fuori dalla memoria locale, e una sonda '
                  'di prontezza che non dipenda dalla raggiungibilità di IAM — altrimenti un '
                  'disservizio esterno fa riavviare un componente sano.')
    p = paragrafo(d, 'L’operatore apre una maschera di SIPO e non ha una sessione.')
    D.testo_di(p, 'L’operatore apre una maschera di SIPO e non ha una sessione. Il front-end lo '
                  'manda a roma-capitale-auth con il codice di ambito e quello dell’applicazione, '
                  'conservando lato server la pagina che aveva chiesto. roma-capitale-auth lo '
                  'manda a IAM, che lo autentica con le credenziali, SPID o CIE — la scelta è di '
                  'IAM, non nostra. Al ritorno roma-capitale-auth scambia il codice con i token di '
                  'IAM, ne ricava il codice fiscale e chiede a roma-capitale-management chi sia '
                  'quella persona per l’applicazione SIPO e che cosa possa fare.')
    p = paragrafo(d, 'Se tutto procede correttamente roma-capitale-auth registra l’accesso')
    D.testo_di(p, 'Se tutto procede correttamente roma-capitale-auth emette il token SIPO e '
                  'riporta l’operatore al front-end, che costruisce l’oggetto utente come oggi e '
                  'lo mette in sessione sotto lo stesso nome di oggi. Accanto vi tiene il token, '
                  'solo lato server. L’identificativo di sessione viene rigenerato, e l’operatore '
                  'arriva alla pagina che aveva chiesto all’inizio. La registrazione dell’accesso, '
                  'che la v0.6 attribuiva al componente di autenticazione, oggi nella piattaforma '
                  'non avviene (MA-17).')

    # ⚠️ la piattaforma non tiene sessione: la continuità fra applicazioni è il cookie
    for vecchio, nuovo in (
            ('perché roma-capitale-auth ha già la sua sessione',
             'perché il cookie della piattaforma è già presente sul dominio dell’ente'),
            ('se la sua sessione è ancora valida il rientro è silenzioso',
             'se il cookie della piattaforma è ancora valido il rientro è silenzioso'),
            ('Il front-end chiude la propria sessione e quella di roma-capitale-auth.',
             'Il front-end chiude la propria sessione e la piattaforma cancella il proprio '
             'cookie.'),
            ('Una sola, per roma-capitale-auth: client, redirect e scope',
             'Una, per l’applicazione SIPO, a nome di roma-capitale-auth: client, redirect e '
             'scope')):
        n = 0
        for p in d.paragraphs:
            n += sostituisci_nei_run(p, vecchio, nuovo)
        for t in d.tables:
            for r in t.rows:
                for c in dict.fromkeys(r.cells):
                    for p in c.paragraphs:
                        n += sostituisci_nei_run(p, vecchio, nuovo)
        if n != 1:
            raise SystemExit(f'«{vecchio[:50]}»: attesa 1 occorrenza, trovate {n}')

    # tabelle del capitolo 6 che descrivevano un componente da costruire
    sost = D.trova_tabella(d, 'Punto', 'Come sta oggi', 'Che cosa serve fare')
    for r in sost.rows:
        if r.cells[0].text.strip() == 'msAuth':
            D.riscrivi_cella(r.cells[1], 'Identificato: è roma-capitale-auth della piattaforma '
                             'MSE [F13]. Il percorso della libreria coincide con l’operazione di '
                             'accesso del servizio, che richiede ambito e applicazione: sono i due '
                             'codici che il commento della guardia dichiara mancanti.')
            D.riscrivi_cella(r.cells[2], 'Configurare nella libreria i codici di ambito e di '
                             'applicazione di SIPO (PI-35) e aggiungere l’applicazione SIPO al '
                             'servizio (MI-10).')
        elif r.cells[0].text.strip().startswith('Il vocabolario dell’autorizzazione'):
            c = r.cells[1]
            D.riscrivi_cella(c, c.text + ' ⚠️ v0.7: il token della piattaforma porta già le '
                             'abilitazioni; mancano i ruoli e il profilo principale (MI-06).')
    cambia = D.trova_tabella(d, 'Aspetto', 'Oggi', 'A regime')
    for r in cambia.rows:
        if r.cells[0].text.strip() == 'Emittenti di gettoni':
            D.riscrivi_cella(r.cells[2], 'Solo roma-capitale-auth, con firma asimmetrica — oggi '
                             'anche il service mesh conia token e la chiave privata è in più '
                             'servizi (PI-31)')
    pre = D.trova_tabella(d, 'Prerequisito', 'Perché', 'A chi compete')
    for r in pre.rows:
        if r.cells[0].text.strip().startswith('La piattaforma di roma-capitale-auth'):
            D.riscrivi_cella(r.cells[0], 'L’ingresso di SIPO nella piattaforma')
            D.riscrivi_cella(r.cells[1], 'roma-capitale-auth va esteso con l’applicazione SIPO '
                             '— client presso IAM, callback, configurazione — e management va '
                             'popolato. È un rilascio della piattaforma, non di SIPO.')
            D.riscrivi_cella(r.cells[2], 'Titolare della piattaforma / Referenti SIPO')
    D.clona_riga(pre, ('Le manutenzioni di priorità alta della piattaforma',
                       'La revoca che non ha effetto e i percorsi interni aperti valgono per '
                       'chiunque entri: vanno chiusi prima che entri lo stato civile.',
                       'Titolare della piattaforma'))
    tab_seq = D.trova_tabella(d, 'Aspetto', 'Sequenza A', 'Sequenza B')
    D.para(d, tab_seq._tbl.getnext(),
           'Le stime della tabella vengono da [F10] e comprendono la costruzione del componente '
           'di autenticazione, che con la piattaforma non si costruisce più. Vanno rifatte sugli '
           'interventi dei capitoli che seguono [DA VERIFICARE].')
    fatti.append('cap. 6: tabelle di sostenibilità, cambiamenti, prerequisiti, nota sulle stime')

    # ------------------------------------------------ Capitolo dei ruoli
    p = paragrafo(d, 'Il committente chiede di introdurre un meccanismo di gestione dei ruoli.')
    D.para(d, p._p.getnext(), pezzi=D.segmenta('', (
        '⚠️ **Dalla versione 0.7 il livello della capacità esiste già, fuori da SIPO**: è la '
        'coppia profilo-abilitazioni di roma-capitale-management. L’analisi dello stato di fatto '
        'che segue resta valida — descrive ciò che va portato nella piattaforma — mentre la '
        'tabella CONF_RUOLO_FUNZIONALITA proposta più avanti diventa una delle modalità di '
        'integrazione, messa a confronto con le altre nel capitolo «L’integrazione con la '
        'gestione AS-IS di SIPO».')))
    inserisci(ancora('Due possibilità ancora aperte'), [
        ('p', 'Con la piattaforma MSE la gestione si sposta, alla fine del percorso, sul front-end '
              'di amministrazione della piattaforma; nella transizione resta qui. Le due fasi sono '
              'descritte in «L’integrazione con la gestione AS-IS di SIPO».'),
    ])
    fatti.append('cap. dei ruoli: due paragrafi di raccordo')

    # ------------------------------------------------ Punti aperti
    pa = D.trova_tabella(d, '#', 'Questione', 'A chi compete')
    aggiunte = {
        'PI-17': ('Chiuso. ', ' ⚠️ v0.7: risolto sul codice [F13]. Il servizio è '
                  'roma-capitale-auth della piattaforma MSE; copre i dipendenti, cercati per '
                  'codice fiscale in roma-capitale-management; la granularità è l’abilitazione '
                  'per applicazione; il token è verificabile dai back-end con la chiave della '
                  'piattaforma (PI-31).'),
        'PI-22': ('Chiuso. ', ' ⚠️ v0.7: è roma-capitale-auth, componente della piattaforma MSE '
                  'dell’area Entrate e Sportelli Unici [F12]. L’impegno di esercizio passa in '
                  'PI-33.'),
        'PI-20': ('', ' ⚠️ v0.7: la piattaforma porta le abilitazioni nel token e, oltre una '
                  'soglia, ne tiene la versione completa in Redis; per SIPO restano da aggiungere '
                  'i ruoli e il profilo principale (MI-06).'),
        'PI-21': ('', ' ⚠️ v0.7: la piattaforma non tiene sessione: token in un cookie e copia '
                  'completa in Redis.'),
        'PI-27': ('', ' ⚠️ v0.7: con la piattaforma si aggiunge una terza collocazione, la '
                  'matrice profilo-abilitazione di roma-capitale-management: vedi «L’integrazione '
                  'con la gestione AS-IS di SIPO».'),
    }
    for r in pa.rows:
        k = r.cells[0].text.strip()
        if k in aggiunte:
            prima_t, dopo_t = aggiunte.pop(k)
            D.riscrivi_cella(r.cells[1], prima_t + r.cells[1].text.strip() + dopo_t)
    assert not aggiunte, aggiunte
    for riga in PUNTI_NUOVI:
        D.clona_riga(pa, riga)
    fatti.append('punti aperti: 5 aggiornati, %d nuovi (PI-31…PI-37)' % len(PUNTI_NUOVI))

    # ------------------------------------------------ Capitolo: il servizio di profilazione
    inserisci(fine, [
        ('h1', 'Il servizio di profilazione della piattaforma'),
        ('b', 'Il committente ha chiesto di valutare il servizio dei profili della piattaforma, '
              'che nella sequenza del login compare come L4, rispondendo a tre domande: **se '
              'richieda interventi di manutenzione; come lo si possa migliorare perché sia più '
              'utile e più sicuro; se e con quali rischi la gestione dei profili che SIPO ha oggi '
              'vi si possa integrare.** Questo capitolo descrive che cosa il servizio è; i tre '
              'capitoli seguenti rispondono alle tre domande, uno per domanda.'),
        ('p', 'La valutazione è fatta sul codice dei tre servizi della piattaforma che partecipano '
              'all’accesso — autenticazione, profilazione, service mesh — nelle versioni '
              'rilasciate fra il 16 e il 24 marzo 2026 [F13]. Due parti non sono fra le sorgenti '
              'e limitano ciò che si può affermare: il modulo comune su cui i servizi poggiano, '
              'che contiene la verifica dei token e il filtro di autenticazione (PI-31), e il '
              'front-end di amministrazione, che potrebbe contenere regole di gestione che il '
              'back-end non ha (PI-36). Dove una conclusione dipende da quelle parti, è detto.'),
        ('h2', 'Che cosa è'),
        ('b', 'Il nome dice profilazione, ma il servizio fa di più: **è il registro della '
              'piattaforma**. Accanto a utenti, profili e abilitazioni tiene l’elenco delle '
              'applicazioni con il loro indirizzo, il contesto e il cloud su cui girano, la '
              'configurazione remota dei servizi, gli avvisi agli utenti, la manutenzione '
              'programmata e i menu. Il service mesh lo interroga a ogni instradamento per '
              'risolvere a quale servizio corrisponda un contesto: se management non risponde, '
              'non si raggiunge alcun servizio che passi dal mesh, non soltanto il login.'),
        ('tab', (NUMERI, [1.8, 4.8])),
        ('h2', 'Il modello dati'),
        ('img', (png_mgmt, 6.5, 'Figura 7 — Il modello dati di roma-capitale-management come '
                                'risulta dalle classi del servizio. In rosso ciò che il modello '
                                'prevede e nessuna operazione scrive.')),
        ('p', 'Il modello è più ricco di quello che il documento della piattaforma mostra [F12] '
              '§6.2, che ne riporta sette tabelle e precede il codice. Tre differenze contano per '
              'SIPO.'),
        ('v', ('Il profilo appartiene a un’applicazione e ha un codice di ruolo. ',
               'È il posto in cui il ruolo di SIPO — e da lì il profilo di back-end — trova '
               'collocazione senza inventare colonne.')),
        ('v', ('Le applicazioni hanno un padre. ',
               'Con l’autorizzazione di gruppo l’accesso a un’applicazione padre unisce le '
               'abilitazioni delle figlie: SIPO può essere un’applicazione padre, e le sue aree '
               'tematiche le figlie.')),
        ('v', ('L’assegnazione può nominare un nodo organizzativo. ',
               'Il nodo appartiene a un albero di strutture e porta regole di visibilità sui dati '
               '— tabella, colonna, valore —: è la forma che il municipio e la struttura di SIPO '
               'cercano. **È però la parte che il codice dichiara e non governa**: nessuna '
               'operazione scrive il nodo sull’assegnazione, e il servizio che restituisce la '
               'visibilità all’operatore la restituisce sempre vuota.')),
        ('p', 'Una sola tabella è specifica di un’applicazione: quella che lega le assegnazioni '
              'alle strutture del protocollo per lo sportello unico, e che non ha alcuna '
              'operazione di scrittura.'),
        ('h2', 'Come partecipa all’accesso'),
        ('tab', (ACCESSO, [1.3, 3.6, 1.8])),
        ('b', 'Sei degli otto passi sono quelli che le versioni precedenti di questo documento '
              'chiedevano: PKCE, ricerca per codice fiscale, disattivazione per inattività, token '
              'in un cookie non leggibile dagli script. **I due che non lo sono riguardano il '
              'tempo**: dove opera l’utente lo dice l’utente, e una revoca aspetta il prossimo '
              'accesso.'),
        ('h2', 'Come si gestisce'),
        ('p', 'Le funzioni di gestione sono quelle che il documento della piattaforma elenca, e il '
              'codice le realizza come operazioni elementari sulle tabelle.'),
        ('tab', (GESTIONE, [1.6, 2.6, 2.6])),
        ('b', 'È il quadro che la domanda del committente lasciava intuire: **la gestione è un '
              'insieme di operazioni di scrittura sulle tabelle, non un processo.** Non c’è nulla '
              'di sbagliato nelle singole operazioni; manca ciò che sta intorno — chi può '
              'concedere, a chi, per quanto tempo, con quale traccia — ed è quello che i capitoli '
              'seguenti trattano.'),
    ])

    # ------------------------------------------------ (a) manutenzione
    n_alta = sum(1 for m in MANUTENZIONE if m[4] == 'Alta')
    # ⚠️ i nomi di classe lunghi non vanno a capo e Word allarga la colonna delle evidenze:
    # uno spazio di larghezza nulla dopo ogni punto consente l'a capo, e il layout fisso
    # (impostato più sotto) fa rispettare le larghezze dichiarate.
    righe_ma = [['#', 'Intervento', 'Evidenza', 'Specie', 'Priorità']] + \
        [[m[0], m[1], m[2].replace('.', '.\u200b'), m[3], m[4]] for m in MANUTENZIONE]
    inserisci(fine, [
        ('h1', 'Gli interventi di manutenzione'),
        ('b', 'La risposta alla prima domanda è **sì**. Gli interventi sono %d, di tre specie: '
              '**difetti**, per cui il servizio si comporta diversamente da come dichiara; '
              'difetti di **igiene**, che espongono la piattaforma; lacune dello **strumento** '
              'di sviluppo e di esercizio. %d sono di priorità alta. La priorità è quella di chi '
              'scrive; la stima dell’impegno non è data, perché dipende dal modulo comune e dal '
              'front-end di amministrazione, che non sono fra le sorgenti.'
              % (len(MANUTENZIONE), n_alta)),
        ('tab', (righe_ma, [0.65, 3.05, 1.55, 0.75, 0.7])),
        ('h2', 'Che cosa non richiede interventi'),
        ('p', 'Va detto con la stessa chiarezza ciò che funziona, perché sono proprio le scelte '
              'che le versioni precedenti di questo documento chiedevano al componente da '
              'costruire: lo scambio del codice con PKCE e risposta in form_post; la ricerca '
              'dell’utente per codice fiscale; la disattivazione automatica dopo sei mesi di '
              'inattività; il token in un cookie non leggibile dagli script, con la copia ridotta '
              'quando supera la soglia; la pila tecnologica, Java 21 e Spring Boot 3.3.4, '
              'recente rispetto a quella di SIPO; un modello dati che prevede già applicazioni '
              'padre e figlie e un albero organizzativo.'),
    ])

    t_ma = D.trova_tabella(d, '#', 'Intervento', 'Evidenza', 'Specie')
    tblPr = t_ma._tbl.find(qn('w:tblPr'))
    for vecchio in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(vecchio)
    layout = tblPr.makeelement(qn('w:tblLayout'), {qn('w:type'): 'fixed'})
    tblPr.append(layout)

    # ------------------------------------------------ (b) miglioramenti
    inserisci(fine, [
        ('h1', 'Come rendere il servizio più utile e più sicuro'),
        ('b', 'La seconda domanda chiede che cosa manchi perché il servizio diventi **la gestione '
              'dei profili dell’ente**, e non soltanto la tabella da cui il login li legge. La '
              'risposta sta quasi tutta nella gestione; la sicurezza ne è il complemento, e se ne '
              'dà conto in fondo.'),
        ('h2', 'Dalla tabella al processo: la gestione'),
        ('tab', (MIGLIORAMENTI, [0.6, 1.2, 1.1, 2.1, 1.9])),
        ('b', 'Il filo comune è uno: **ogni concessione diventa un fatto con un autore, una data e '
              'una durata**, invece di una riga che c’è o non c’è. Da lì discendono la storia, '
              'la revoca, la revisione e l’amministrazione per ambito; e da lì discende anche '
              'l’utilità per le altre applicazioni della piattaforma, che hanno le stesse '
              'domande di SIPO senza averle ancora poste.'),
        ('h2', 'La sicurezza, in breve'),
        ('p', 'I rilievi di sicurezza non sono il tema di questo capitolo, ma un servizio che '
              'decide chi entra nello stato civile ne eredita il peso. Si riportano le cinque '
              'misure che li coprono; i difetti puntuali sono fra gli interventi di '
              'manutenzione.'),
        ('tab', (SICUREZZA, [2.2, 4.4])),
        ('h2', 'In che ordine'),
        ('b', 'Prima le manutenzioni di priorità alta, perché alcune di esse — la revoca, i '
              'percorsi interni — sono la condizione perché un utente in più entri nella '
              'piattaforma. Poi i miglioramenti **senza i quali SIPO non può entrare**: il nodo '
              'organizzativo (MI-05), il profilo principale (MI-06), la configurabilità delle '
              'applicazioni (MI-10), la storia delle assegnazioni (MI-01). Il resto può seguire, '
              'ma va chiesto adesso, perché cambia il modello dati e conviene cambiarlo una volta '
              'sola.'),
    ])

    # ------------------------------------------------ (c) integrazione
    inserisci(fine, [
        ('h1', 'L’integrazione con la gestione AS-IS di SIPO'),
        ('b', 'La risposta alla terza domanda è **sì, ma non per corrispondenza diretta**: i '
              'concetti ci sono quasi tutti, la forma del permesso no. Il capitolo mette a '
              'confronto le due gestioni, indica il punto che decide, propone come procedere e '
              'ne elenca i rischi. Lo stato di fatto è descritto nella prima parte e nel capitolo '
              'dei ruoli: qui lo si usa.'),
        ('h2', 'Le corrispondenze'),
        ('tab', (CORRISPONDENZE, [1.9, 1.8, 3.0])),
        ('h2', 'Il punto che decide: la quaterna'),
        ('b', 'In SIPO il permesso è **per persona**: l’utente, nel ruolo, ha una funzionalità in '
              'un’area. In management il permesso è **per profilo**: l’utente ha un profilo, il '
              'profilo ha le abilitazioni. Portare la prima forma nella seconda significa '
              'ricavare i profili dalle abilitazioni esistenti — è l’interrogazione del capitolo '
              'dei ruoli, con le stesse soglie — e decidere che cosa fare delle eccezioni.'),
        ('b', 'Le eccezioni sono il nodo. Una funzionalità che solo alcuni utenti di un ruolo '
              'hanno può diventare un profilo aggiuntivo — e se le eccezioni sono molte si '
              'finisce con un profilo per persona, cioè **con la quaterna di oggi sotto un altro '
              'nome** — oppure può essere tolta, e qualcuno perde una funzione che usa. Quante '
              'siano non si sa: dipende dalla configurazione di produzione, che non è fra le '
              'sorgenti (PI-14). La misura va fatta prima di scegliere la modalità, con '
              'l’interrogazione già pronta.'),
        ('p', 'Il capitolo dei ruoli non va quindi in archivio: la matrice ricavata dai dati, il '
              'report di differenza e il profilo principale diventano la procedura di migrazione '
              'verso la piattaforma. Cambia soltanto il luogo in cui la matrice vive.'),
        ('h2', 'Tre modalità, una da escludere'),
        ('tab', (MODALITA, [1.4, 2.0, 1.6, 1.7])),
        ('b', 'La proposta è **A come transizione e B come approdo**, con una regola che vale per '
              'entrambe: **per ciascuna applicazione un solo posto in cui si scrive**. In A è '
              'AnagrafeWeb, in B la piattaforma; il passaggio dall’una all’altra si fa per '
              'applicazione, non per tutto SIPO in un giorno.'),
        ('h2', 'I rischi'),
        ('tab', (RISCHI, [2.2, 0.8, 0.7, 3.0])),
    ])
    fatti.append('4 capitoli nuovi: servizio di profilazione · manutenzione (%d interventi) · '
                 'miglioramenti (%d) · integrazione (%d rischi)'
                 % (len(MANUTENZIONE), len(MIGLIORAMENTI) - 1, len(RISCHI) - 1))

    # ------------------------------------------------ Storia e frontespizio
    D.storia(d, '15/10/2026', '0.7',
             'Scopo · Executive summary · Parte II (S4) · Scelte effettuate · I ruoli · Punti '
             'aperti · quattro capitoli nuovi',
             'Recepita la strada intrapresa dall’ente: autenticazione e profili non si '
             'costruiscono, si adottano i servizi della piattaforma MSE [F12], letti sul codice '
             '[F13]. Riscritta la sezione della soluzione nel capitolo delle scelte, con la figura '
             'dell’architettura adottata e il confronto fra il token della piattaforma e ciò che '
             'serve a SIPO. Aggiunti quattro capitoli che rispondono alle domande del committente '
             'sul servizio dei profili: che cosa è, con lo schema dati ricavato dal codice; '
             'gli interventi di manutenzione (%d, %d di priorità alta); come renderlo più utile e '
             'più sicuro, con la sicurezza in forma sintetica; l’integrazione con la gestione di '
             'SIPO, con tre modalità, una esclusa, e dieci rischi. Chiusi PI-17 e PI-22, nuovi '
             'PI-31…PI-37.' % (len(MANUTENZIONE), n_alta))
    for f in D.frontespizio(d, '0.7', 'ANALISI_Identita-Profilazione-IAM_v0.7'):
        fatti.append(f)

    d.save(DST)
    print('\n'.join('  · ' + f for f in fatti))
    cap, tab, fig = D.riepilogo(DST)
    print(f'  capitoli {cap}, tabelle {tab}, immagini {fig}')
    print('  scritto', DST)


if __name__ == '__main__':
    main()
