# -*- coding: utf-8 -*-
"""DISEGNO_Back-Office_ANSC v0.9 → v0.10 (11/10/2026).

Riallineamento alla v3.40 dell'analisi, che introduce ANSC_SOGGETTO_IDENTIFICATIVI: la
consultazione R005 registra l'identificativo del soggetto in ANSC nella tabella nuova e i due
codici ANPR su ANAG_USR.SOGGETTO, dove già esistono, scrivendoli solo se le colonne sono vuote.
Una discordanza non sovrascrive: si registra e si segnala al back-office.
  * nuova pagina «Identificativi dei soggetti» (area operativa): le eccezioni della scrittura
    dei codici ANPR in testa, l'elenco, il confronto anagrafe ↔ risposta di R005;
  * dettaglio dell'atto: scheda «Soggetti», tabelle sottese e campi;
  * la ricerca R005 con la scelta fra più risultati NON sta nel back-office: sta nel riepilogo
    dei soggetti di SIPO, come dice l'analisi; il back-office legge e riporta all'atto;
  * mappa, unità di rilascio, conteggi (venti pagine, ventuno tabelle), copertura;
  * BO-31 (le operazioni del contratto che mancano), BO-32 (una discordanza vista e accettata
    non si distingue da una mai vista); [R1] = v3.40.

Le figure da 12 in poi scalano di uno (una pagina in più dopo la Figura 11).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v5.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v2.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/bo_v0_10.py
"""
import os
import re
import shutil
import subprocess
import sys

import docx
from docx.table import Table

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.9.docx')
DST = os.path.join(FIN, 'DISEGNO_Back-Office_ANSC_v0.10.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')
NUOVA_FIGURA = 12


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


USO = ('Tabella', 'Uso', 'Che cosa se ne mostra')
EL = ('Elemento', 'Genere', 'Comportamento')

PAGINA = [
    ('h2', 'Identificativi dei soggetti'),
    ('img', ('bo_identificativi.png', NUOVA_FIGURA,
             'Identificativi dei soggetti. In testa le eccezioni della scrittura dei codici '
             'ANPR, sotto l’elenco filtrato sulle discordanze e il confronto per il soggetto '
             'scelto.')),
    ('p', 'Per gli intestatari il modello evento chiede tre identificativi che senza la '
          'consultazione R005 non si conoscono: l’identificativo della scheda ANPR, '
          'l’identificativo unico nazionale e l’identificativo del soggetto in ANSC. L’analisi '
          '[R1] li registra in due luoghi, per la regola dell’atomicità del dato: i due codici '
          'ANPR su ANAG_USR.SOGGETTO, dove SIPO li conserva già, e l’identificativo ANSC nella '
          'tabella ANSC_SOGGETTO_IDENTIFICATIVI, una riga per soggetto di stato civile.'),
    ('p', 'La consultazione si fa nel riepilogo dei soggetti delle maschere di SIPO, dove '
          'l’operatore sceglie il soggetto giusto quando ANSC ne restituisce più d’uno. Il '
          'back-office non la ripete e non la sostituisce: mostra che cosa ha prodotto, e in '
          'particolare i casi in cui un codice ANPR non è stato scritto.'),
    ('p', 'I codici ANPR si scrivono solo se la colonna anagrafica è vuota. Se la trovano '
          'valorizzata con un valore diverso non la sovrascrivono: la riga resta DISCORDANTE ed '
          'è questa pagina a farla vedere. Accanto stanno gli altri due esiti che lasciano un '
          'codice senza posto — l’identificativo della scheda che non entra nella colonna '
          'numerica e il soggetto senza riga anagrafica, che l’analisi tiene aperto (OP-74).'),
    ('p', '⚠️ La pagina non scrive sull’anagrafe, e di proposito. Il codice presente in '
          'ANAG_USR.SOGGETTO lo hanno scritto i canali di allineamento con ANPR: correggerlo da '
          'qui significherebbe far prevalere la consultazione di ANSC sul dato anagrafico, '
          'cioè l’opposto della regola. La discordanza si risolve dove il dato nasce; una nuova '
          'consultazione dal riepilogo dei soggetti ne aggiorna poi l’esito.'),
    ('p', 'Il confronto mette accanto il valore registrato in anagrafe e quello restituito da '
          'R005. Il secondo non è conservato nella tabella: si legge dalla risposta della '
          'chiamata in ANSC_LOG_AUDIT, a cui la riga rimanda per valore. Quando l’audit è stato '
          'sfoltito il confronto mostra il solo valore anagrafico e lo dice.'),
    ('h3', 'Tabelle sottese'),
    ('uso', [
        USO,
        ('ANSC_SOGGETTO_IDENTIFICATIVI', 'Lettura',
         'È la tabella dell’elenco: ID_ANSC_SOGGETTO, ID_ATTO_SIPO, COD_ESITO_RICERCA, '
         'COD_ESITO_ANPR, DATA_RICERCA, UTENTE_RICERCA. I conteggi in testa raggruppano per '
         'COD_ESITO_ANPR.'),
        ('MATR_USR.SOGGETTO', 'Lettura per raccordo',
         'Cognome, nome e codice fiscale del soggetto, da ID_SOGGETTO_SIPO; CODICE_INDIVIDUALE '
         'per raggiungere il soggetto anagrafico.'),
        ('ANAG_USR.SOGGETTO', 'Lettura',
         'ID_SCHEDA_ANPR e ID_UNICO_NAZIONALE, nella colonna «In anagrafe» del confronto. '
         'Nessuna scrittura.'),
        ('ANSC_LOG_AUDIT', 'Lettura',
         'La risposta di R005 indicata da ID_AUDIT: i valori restituiti, nella colonna '
         '«Restituito da R005».'),
    ]),
    ('h3', 'Campi, colonne e azioni'),
    ('el', [
        EL,
        ('Da trattare', 'Pannello',
         'Un conteggio per ciascun esito che lascia un codice senza posto: Discordante, Non '
         'numerico, Assente in anagrafe. Il conteggio apre l’elenco già filtrato.'),
        ('Esito dei codici ANPR · Esito della ricerca · Atto SIPO · Codice fiscale · Ricerca da '
         '· Ricerca a', 'Filtri',
         'Predefinito l’esito Discordante. Gli esiti sono i valori ammessi dai vincoli della '
         'tabella; le due date agiscono su DATA_RICERCA.'),
        ('Occhio', 'Azione di riga',
         'Apre il confronto fra anagrafe e risposta di R005 e, da lì, la registrazione '
         'dell’audit con la risposta integrale.'),
        ('Matita', 'Azione di riga',
         'Riporta alla maschera di SIPO che ha in carico l’atto, dove si ripete la '
         'consultazione. Disabilitata se l’atto è firmato.'),
        ('ESPORTA', 'Azione di pagina',
         'Scarica l’elenco filtrato, per chi tratta le discordanze fuori dal back-office. Non '
         'tocca alcuna tabella.'),
    ]),
]

DETTAGLIO_P = (
    'La scheda «Soggetti» elenca i soggetti dell’atto con i tre identificativi che il modello '
    'evento chiede per gli intestatari e con gli esiti della consultazione R005: se il '
    'soggetto è stato trovato, scelto fra più risultati o non trovato, e che cosa è accaduto ai '
    'codici ANPR. Per l’intestatario la consultazione è obbligatoria anche quando non trova '
    'nulla, e la scheda lo segnala se manca. La consultazione non si fa qui ma nel riepilogo '
    'dei soggetti di SIPO; una discordanza rimanda alla pagina «Identificativi dei soggetti».')

DETTAGLIO_USO = [
    ('ANSC_SOGGETTO_IDENTIFICATIVI', 'Lettura',
     'Nella scheda «Soggetti»: per ciascun soggetto dell’atto l’identificativo ANSC e i due '
     'esiti, ricerca e codici ANPR.'),
    ('ANAG_USR.SOGGETTO', 'Lettura',
     'Nella scheda «Soggetti»: i due codici ANPR del soggetto, raggiunto per '
     'CODICE_INDIVIDUALE.'),
]
DETTAGLIO_EL = ('Soggetti', 'Scheda',
                'Una riga per soggetto dell’atto: ruolo, nome, identificativo ANSC, codici ANPR, '
                'esito della ricerca ed esito dei codici ANPR. Segnala l’intestatario non ancora '
                'consultato. Sola lettura.')

COPERTURA = [
    ('ANSC_SOGGETTO_IDENTIFICATIVI', 'Identificativi dei soggetti · Dettaglio atto (scheda '
                                     'Soggetti)',
     'In sola lettura, come deve essere: la scrive il concentratore con la consultazione R005 '
     'avviata dal riepilogo dei soggetti di SIPO. Nuova in questa versione.'),
    ('ANAG_USR.SOGGETTO (di SIPO)', 'Identificativi dei soggetti · Dettaglio atto (scheda '
                                    'Soggetti)',
     'Non appartiene al modello dell’integrazione: è il soggetto anagrafico di SIPO, di cui si '
     'leggono le due colonne dei codici ANPR. Nessuna pagina vi scrive, per la regola che '
     'vieta di sovrascrivere il dato dei canali ANPR.'),
]

BO_NUOVI = [
    ('BO-31', 'Le operazioni che servono alla pagina e alla consultazione dei soggetti',
     'Il contratto del concentratore [R6] non ha operazioni sugli identificativi dei soggetti: '
     'né la consultazione R005 che il riepilogo dei soggetti di SIPO deve avviare, né l’elenco '
     'e il confronto di questa pagina. Vanno aggiunte al contratto, e la specifica delle API '
     'rigenerata, prima di realizzare la pagina.'),
    ('BO-32', 'Una discordanza vista e accettata non si distingue da una mai vista',
     'La tabella registra l’esito della scrittura, non il suo trattamento: una riga resta '
     'DISCORDANTE finché una nuova consultazione non la trova allineata. Se chi tratta le '
     'discordanze — da stabilire, OP-74 dell’analisi — deve poter dichiarare un caso esaminato, '
     'serve una colonna in più, da recepire nell’analisi; altrimenti il pannello «Da trattare» '
     'cresce anche con i casi già esaminati.'),
]

STORIA = ('11/10/2026', '0.10',
          'Le pagine (Identificativi dei soggetti, nuova; Dettaglio dell’atto) · Mappa · '
          'Contesto tecnico · Copertura · Riferimenti · Punti aperti',
          'Riallineato il back-office alla v3.40 dell’analisi, che introduce '
          'ANSC_SOGGETTO_IDENTIFICATIVI e scrive i due codici ANPR sul soggetto anagrafico solo '
          'se le colonne sono vuote. Nuova pagina «Identificativi dei soggetti», nell’area '
          'operativa: le eccezioni della scrittura dei codici ANPR, l’elenco e il confronto fra '
          'anagrafe e risposta di R005, in sola lettura. Il dettaglio dell’atto acquisisce la '
          'scheda «Soggetti». La consultazione resta nel riepilogo dei soggetti di SIPO. La '
          'copertura sale a ventuno tabelle del modello, più il soggetto anagrafico di SIPO in '
          'lettura. BO-31 e BO-32 nuovi.')


def tabelle_di(elementi, d):
    return [Table(e, d) for e in elementi if e.tag.endswith('}tbl')]


def inserisci_dopo(t, dopo_tr, valori):
    nuova = D.clona_riga(t, valori)
    dopo_tr.addnext(nuova._tr)
    return nuova._tr


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    MOD_USO = D.trova_tabella(d, *USO)
    MOD_EL = D.trova_tabella(d, *EL)

    # ─────────────────────────── 0. le figure da 12 in poi scalano di uno
    for p in d.paragraphs:
        t = p.text.strip()
        m = re.match(r'^Figura (\d+) —', t)
        if m and int(m.group(1)) >= NUOVA_FIGURA:
            D.testo_di(p, re.sub(r'^Figura \d+ —', 'Figura %d —' % (int(m.group(1)) + 1), t))
        elif not m and re.search(r'Figura (1[2-9]|2\d)\b', t):
            raise SystemExit('rimando a una figura da rinumerare: ' + t[:80])
    fatti.append('figure 12-28 rinumerate 13-29')

    # ─────────────────────────── 1. la pagina nuova, dopo gli allegati dell'atto
    ancora = D.h(d, 2, 'Gestione delle notifiche')._p
    for tipo, c in PAGINA:
        if tipo == 'p':
            D.para(d, ancora, c)
        elif tipo in ('h2', 'h3'):
            D.para(d, ancora, c, stile='Heading ' + tipo[1])
        elif tipo == 'uso':
            D.tabella(d, ancora, c, MOD_USO, [1.6, 1.2, 3.5])
            D.para(d, ancora, '')
        elif tipo == 'el':
            D.tabella(d, ancora, c, MOD_EL, [1.6, 1.2, 3.5])
            D.para(d, ancora, '')
        elif tipo == 'img':
            png, n, did = c
            D.immagine(d, ancora, os.path.join(IMG, png), 6.2, 'Figura %d — %s' % (n, did))
    fatti.append('pagina «Identificativi dei soggetti»')

    # ─────────────────────────── 2. dettaglio dell'atto: scheda «Soggetti»
    D.sostituisci(d, 'lo stato di completezza sezione per sezione, i documenti',
                  'lo stato di completezza sezione per sezione, i soggetti con i loro '
                  'identificativi, i documenti', attese=1, fatti=fatti)
    sez = D.sezione(d, 'Dettaglio dell’atto', 2)
    assert sez, 'sezione del dettaglio non trovata'
    payload = [e for e in sez if e.tag.endswith('}p')
               and e.xpath('string(.)').startswith('La scheda «Payload inviato»')]
    assert len(payload) == 1
    D.para(d, payload[0].getnext(), DETTAGLIO_P)
    uso, el = tabelle_di(D.sezione(d, 'Dettaglio dell’atto', 2), d)
    assert uso.rows[0].cells[0].text.strip() == 'Tabella'
    assert el.rows[0].cells[0].text.strip() == 'Elemento'
    dopo = [r for r in uso.rows if r.cells[0].text.strip() == 'ANSC_LOG_AUDIT'][0]._tr
    for v in DETTAGLIO_USO:
        dopo = inserisci_dopo(uso, dopo, v)
    dopo = [r for r in el.rows if r.cells[0].text.strip() == 'Sezioni'][0]._tr
    inserisci_dopo(el, dopo, DETTAGLIO_EL)
    D.sostituisci_immagine(d, 'Figura 10 —', os.path.join(IMG, 'bo_atto.png'))
    fatti.append('dettaglio dell’atto: scheda Soggetti, tabelle e figura')

    # ─────────────────────────── 3. mappa, contesto tecnico, conteggi
    t = D.trova_tabella(d, 'Area', 'Pagine', 'Chi la usa e quando')
    r = [r for r in t.rows if r.cells[0].text.strip() == 'Operativa'][0]
    assert 'Allegati dell’atto · Numerazione' in r.cells[1].text
    D.riscrivi_cella(r.cells[1], r.cells[1].text.strip().replace(
        'Allegati dell’atto · Numerazione',
        'Allegati dell’atto · Identificativi dei soggetti · Numerazione'))
    t = D.trova_tabella(d, 'Unità di rilascio', 'Pagine', 'Che cosa la fa cambiare')
    rr = {r.cells[0].text.strip(): r for r in t.rows}
    c = rr['mfe-operativa'].cells[1]
    assert 'Allegati' in c.text and 'Identificativi' not in c.text, c.text
    testo = c.text.strip()
    i = testo.index('· Numerazione')
    D.riscrivi_cella(c, testo[:i] + '· Identificativi dei soggetti ' + testo[i:])
    D.sostituisci_immagine(d, 'Figura 1 —', os.path.join(IMG, 'bo_mfe.png'))
    fatti.append('mappa e unità di rilascio')
    for a, b in [
        ('L’applicazione è stata pensata con diciannove pagine distribuite su quattro aree '
         'funzionali, cui si aggiunge una ventesima pagina',
         'L’applicazione è stata pensata con venti pagine distribuite su quattro aree '
         'funzionali, cui si aggiunge una ventunesima pagina'),
        ('Seguono le diciannove pagine', 'Seguono le venti pagine'),
        ('Un back-office di diciannove pagine', 'Un back-office di venti pagine'),
        ('ripeterle in diciannove sezioni', 'ripeterle in venti sezioni'),
        ('le venti tabelle già progettate', 'le ventuno tabelle già progettate'),
        ('da cui vengono le venti tabelle', 'da cui vengono le ventuno tabelle'),
        ('le venti del modello dell’integrazione più le quattro del registro',
         'le ventuno del modello dell’integrazione più le quattro del registro'),
    ]:
        D.sostituisci(d, a, b, attese=1, fatti=fatti)

    # ─────────────────────────── 4. copertura
    t = D.trova_tabella(d, 'Tabella', 'Pagine che la governano', 'Copertura')
    rr = {r.cells[0].text.strip(): r for r in t.rows}
    dopo = rr['DATI_ATTO_ESTERNO']._tr
    for c in COPERTURA:
        dopo = inserisci_dopo(t, dopo, c)
    for k in ('PRENOTAZIONE_ATTI', 'DATI_ATTO_ESTERNO'):
        cella = rr[k].cells[2]
        D.riscrivi_cella(cella, cella.text.strip().replace(' Nuova in questa versione.', ''))
    fatti.append('copertura: ANSC_SOGGETTO_IDENTIFICATIVI e ANAG_USR.SOGGETTO')

    # ─────────────────────────── 5. riferimenti, punti aperti, scheda, storia
    D.sostituisci(d, 'ANALISI_Integrazione-ANSC_v3.39.docx', 'ANALISI_Integrazione-ANSC_v3.40.docx',
                  attese=1, fatti=fatti, etichetta='[R1]')
    t = D.trova_tabella(d, '#', 'Questione', 'Perché è aperta')
    assert t.rows[-1].cells[0].text.strip() == 'BO-30'
    for b in BO_NUOVI:
        D.clona_riga(t, b)
    scheda = d.tables[0]
    for r in scheda.rows:
        k = r.cells[0].text.strip()
        if k == 'Data consegna':
            D.riscrivi_cella(r.cells[1], '11/10/2026')
        elif k == 'Versione':
            D.riscrivi_cella(r.cells[1], '0.10')
        elif k == 'Documento':
            D.riscrivi_cella(r.cells[1], 'DISEGNO_Back-Office_ANSC_v0.10')
    D.storia(d, *STORIA)
    fatti.append('[R1] v3.40, BO-31 e BO-32, scheda, storia')

    d.save(DST)
    dopo_c = commenti(DST)
    if dopo_c != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo_c} dopo')
    print('\n'.join('· ' + f for f in fatti))
    print(f'commenti preservati: {dopo_c}')
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))


if __name__ == '__main__':
    main()
