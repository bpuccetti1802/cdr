# -*- coding: utf-8 -*-
"""DISEGNO_Back-Office_ANSC v0.4 → v0.5 (29/09/2026).

La scheda della riconciliazione descriveva l'elenco ma non il **dettaglio di inserimento e
modifica**: i bottoni «Nuova corrispondenza» e la matita non portavano ad alcuna maschera
disegnata. È la lacuna che questa versione colma, e conta perché è lì che l'inversione
decisa nella v3.31 dell'analisi si vede: lato ANSC in sola lettura, lato SIPO da completare.

Si correggono anche due frasi rimaste indietro: una persa in un'edizione manuale e una che
descriveva l'orientamento precedente.

⚠️ Si modifica IL FILE REALE: la v0.4 porta sei commenti di Word da preservare.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/pagine_bo_v2.py img
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_5_bo_dettaglio_riconciliazione.py
"""
import os
import re
import shutil
import sys
import zipfile

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(BASE, 'Documenti finali', 'DISEGNO_Back-Office_ANSC_v0.4.docx')
DST = os.path.join(BASE, 'Documenti finali', 'DISEGNO_Back-Office_ANSC_v0.5.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')

CAMPI = [
    ['Elemento', 'Genere', 'Comportamento'],
    ['Decodifica', 'Campo',
     'Il dominio su cui si sta lavorando. Arrivando dall’elenco è già valorizzato e in sola '
     'lettura; su una corrispondenza nuova si sceglie.'],
    ['Valore ANSC *', 'Campo',
     '**Obbligatorio.** Tendina dei valori che il dominio dichiara validi alla data '
     'indicata. ⚠️ Non è testo libero: un valore che ANSC non ammette non è riconciliabile, '
     'e lasciarlo scrivere a mano produrrebbe corrispondenze verso valori inesistenti.'],
    ['Descrizione ANSC', 'Campo', 'Sola lettura: viene dal dizionario insieme al valore.'],
    ['Validità in ANSC', 'Campo',
     'Sola lettura, informativa. ⚠️ Serve a rendere visibile il caso in cui si stia '
     'riconciliando un valore che ANSC ha già chiuso: è legittimo per gli atti vecchi, è un '
     'errore per quelli nuovi.'],
    ['Schema e Tabella', 'Campi',
     'Alimentano SCHEMA_SIPO e TABELLA_SIPO. La tabella si sceglie fra le CONF_* raggiungibili '
     'per sinonimo; lo schema è richiesto perché in SIPO esistono tabelle omonime in schemi '
     'diversi.'],
    ['Campo', 'Campo',
     'Alimenta CAMPO_SIPO. ⚠️ **Entra nella chiave di unicità**: cambiarlo non modifica '
     'questa corrispondenza, ne dichiara un’altra. La maschera lo segnala prima di salvare.'],
    ['Valore SIPO', 'Campo',
     '**Facoltativo**, ed è il punto dell’intera maschera. Tendina alimentata dalla tabella '
     'scelta. ⚠️ Salvare lasciandolo vuoto non è una bozza incompleta: è la **dichiarazione** '
     'che quel valore ammesso da ANSC non ha ancora un corrispondente locale, ed è la riga '
     'che l’elenco mostra come «da mappare».'],
    ['Descrizione SIPO', 'Campo',
     'Si compila da sola quando si sceglie il valore. Resta modificabile per le tabelle che '
     'non portano una descrizione leggibile.'],
    ['Stato', 'Campo calcolato',
     'Non si immette: vale «da mappare» finché il valore SIPO manca, «riconciliato» quando '
     'c’è. È lo stesso contrassegno mostrato nell’elenco.'],
    ['Condizione', 'Campo',
     'Facoltativa. Per i casi in cui la traduzione dipende dal contesto e non dal solo '
     'valore. Vuota nella grande maggioranza delle righe.'],
    ['Valida dal * / Valida fino al', 'Campi',
     '⚠️ È la validità **della corrispondenza**, da non confondere con quella del valore in '
     'ANSC mostrata sopra. La prima è una decisione del Comune, la seconda un dato di ANSC.'],
    ['SALVA', 'Azione',
     'Verifica la chiave di unicità — versione, decodifica, valore ANSC, campo, data di '
     'inizio — e, se una corrispondenza è già in vigore, **la chiude alla data indicata '
     'invece di sovrascriverla**, aprendo la nuova da lì in avanti.'],
    ['SALVA E PROSSIMA', 'Azione',
     'Salva e apre subito la corrispondenza non riconciliata successiva, secondo l’ordine '
     'dell’elenco da cui si è partiti. ⚠️ È l’azione che rende praticabile un lavoro fatto di '
     'molte righe brevi: senza, si torna all’elenco a ogni salvataggio.'],
    ['ANNULLA', 'Azione', 'Abbandona senza scrivere. Chiede conferma se qualcosa è cambiato.'],
]


def commenti(percorso):
    return zipfile.ZipFile(percorso).read('word/comments.xml').decode().count('<w:comment ')


def main():
    atteso = commenti(SRC)
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    fatti = []

    # ───────────── 1. le figure successive scalano di uno, dalla 14 in poi
    # ⚠️ si va a ritroso: rinominando 14→15 prima di 15→16 si sovrascriverebbe.
    for n in range(17, 13, -1):
        for p in d.paragraphs:
            if p.text.strip().startswith('Figura %d —' % n):
                D.testo_di(p, re.sub(r'^Figura %d —' % n, 'Figura %d —' % (n + 1),
                                     p.text.strip()))
                break
    fatti.append('figure da 14 a 17 rinumerate da 15 a 18')

    # ───────────── 2. due frasi rimaste indietro nella scheda
    D.sostituisci(
        d,
        'È la pagina che permette di sapere quale dizionario ANSC vale per un campo non dice '
        'come tradurre il valore che SIPO ha in tabella.',
        'Sapere quale dizionario di ANSC governa un campo non dice come tradurre il valore '
        'che SIPO ha in tabella.',
        attese=1, etichetta='frase di apertura ricomposta', fatti=fatti)
    D.sostituisci(
        d,
        'Un valore SIPO senza corrispondenza blocca in preverifica ogni atto che lo usi. Per '
        'questo il filtro «Solo non riconciliati» è predefinito a sì: la pagina si apre sul '
        'lavoro da fare, non sull’elenco completo.',
        'Un valore ammesso da ANSC e privo di corrispondente locale blocca in preverifica '
        'ogni atto che lo richieda. Per questo il filtro «Solo non riconciliati» è '
        'predefinito a sì: la pagina si apre sul lavoro da fare, non sull’elenco completo.',
        attese=1, etichetta='frase riorientata sul lato ANSC', fatti=fatti)

    # ───────────── 3. la sezione nuova, in coda alla scheda della riconciliazione
    ancora = D.h(d, 2, 'Versioni della configurazione')
    modello = D.trova_tabella(d, 'Elemento', 'Genere', 'Comportamento')

    D.para(d, ancora._p, 'Il dettaglio di inserimento e modifica', stile='Heading 3')
    D.para(d, ancora._p,
           'L’elenco dice che cosa manca; il dettaglio è dove si colma. Vi si arriva dal '
           'bottone «Nuova corrispondenza» o dalla matita di una riga, e la maschera è la '
           'stessa nei due casi: cambia soltanto se i campi arrivino già valorizzati.')
    D.immagine(d, ancora._p, os.path.join(IMG, 'bo_riconciliazione_dettaglio.png'), 6.4,
               'Figura 14 — Riconciliazione: inserimento e modifica di una corrispondenza.')
    D.para(d, ancora._p,
           '⚠️ **La disposizione dei due pannelli è la traduzione grafica della scelta di '
           'modello.** Sopra sta il lato ANSC, in sola lettura, perché è ciò che il dominio '
           'dichiara e non si discute; sotto il lato SIPO, da completare, perché è il lavoro. '
           'Una maschera che li mettesse alla pari, o che chiedesse prima il valore locale, '
           'suggerirebbe che i due lati abbiano lo stesso grado di certezza — e sono '
           'esattamente le due cose che la struttura distingue, con l’obbligatorietà su uno '
           'solo dei due.')
    D.para(d, ancora._p,
           '**Salvare senza valore SIPO è consentito, ed è un atto significativo.** La riga '
           'che ne risulta non è una bozza lasciata a metà: dichiara che quel valore ammesso '
           'da ANSC non ha, oggi, un corrispondente nelle tabelle del Comune. È la stessa '
           'riga che l’elenco mostra come «da mappare» e che il filtro predefinito porta in '
           'cima. ⚠️ Nella struttura precedente questo non era rappresentabile, perché il '
           'valore obbligatorio era quello di SIPO.')
    D.para(d, ancora._p,
           '**Una corrispondenza che cambia non si sovrascrive: si chiude e se ne apre una '
           'nuova.** Il motivo non è prudenza archivistica ma una necessità: gli atti già '
           'depositati sono stati costruiti con la corrispondenza in vigore allora, e '
           'devono restare leggibili così. Per questo la maschera, quando trova una '
           'corrispondenza già valida per la stessa chiave, lo dichiara prima di salvare e '
           'dice che cosa accadrà, invece di rifiutare o di sostituire in silenzio.')
    D.para(d, ancora._p, 'Campi e azioni del dettaglio', stile='Heading 3')
    D.tabella(d, ancora._p, CAMPI, modello, larghezze=[1.6, 1.1, 3.8])
    fatti.append('sezione «Il dettaglio di inserimento e modifica» con figura e tabella')

    # ───────────── 4. il punto aperto sulle corrispondenze molti-a-uno
    t = D.trova_tabella(d, '#', 'Questione')
    D.clona_riga(t, (
        'BO-17', 'Se due valori di SIPO possano convergere sullo stesso valore di ANSC',
        'La chiave di unicità adottata nell’analisi (versione, decodifica, valore ANSC, '
        'campo, data) ammette **un solo valore di SIPO per ciascun valore di ANSC e campo**. '
        '⚠️ Se in qualche decodifica esistessero due codici locali distinti da tradurre nello '
        'stesso valore — il caso tipico di un archivio stratificato, con un codice vecchio e '
        'uno nuovo — la maschera qui descritta non consentirebbe di dichiararlo. La verifica '
        'si fa sulle tabelle CONF_* raccordate; nell’analisi è il punto aperto OP-67.'))
    fatti.append('BO-17 aperto, collegato a OP-67 dell’analisi')

    # ───────────── 5. testata e storia
    for tab in d.tables:
        if tab.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in tab.rows:
                v = {'Data consegna': '29/09/2026', 'Versione': '0.5',
                     'Documento': 'DISEGNO_Back-Office_ANSC_v0.5'}.get(r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    D.storia(d, '29/09/2026', '0.5', 'Riconciliazione · Punti aperti',
             'Aggiunto il dettaglio di inserimento e modifica di una corrispondenza, che la '
             'scheda della riconciliazione non conteneva: i bottoni c’erano, la maschera no. '
             'La disposizione dei pannelli recepisce l’inversione decisa nella v3.31 '
             'dell’analisi — lato ANSC in sola lettura, lato SIPO da completare — e rende '
             'esplicito che salvare senza valore di SIPO dichiara un valore non ancora '
             'mappato, non una riga incompleta. Documentato che una corrispondenza in vigore '
             'si chiude e non si sovrascrive, perché gli atti già depositati devono restare '
             'leggibili. Corrette due frasi rimaste indietro. Le figure dalla quattordicesima '
             'in poi scalano di uno.')
    fatti.append('testata e storia aggiornate')

    # grassetti: `para` e `tabella` non interpretano «**…**»
    def grassetti(par):
        if '**' not in par.text or D.ha_commenti(par._p):
            return 0
        pezzi = D.segmenta('', par.text)
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        for testo, gr in pezzi:
            run = par.add_run(testo)
            run.bold = gr
        return 1
    n = sum(grassetti(p) for p in d.paragraphs)
    for tab in d.tables:
        for r in tab.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    n += grassetti(p)
    fatti.append('%d paragrafi con grassetto applicato' % n)

    d.save(DST)
    trovati = commenti(DST)
    assert trovati == atteso, 'commenti persi: %d invece di %d' % (trovati, atteso)
    print('\n'.join(' · ' + f for f in fatti))
    print('commenti preservati:', trovati)
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))
    print('scritto:', os.path.relpath(DST, BASE))


if __name__ == '__main__':
    main()
