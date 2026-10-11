# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.38 → v3.39 (11/10/2026).

La numerazione comunale diventa un MICROSERVIZIO DEDICATO, `num-ansc-sipo` (deciso con
l'utente l'11/10), e lo stacco sta nel microservizio, non in PL/SQL: la base dati tiene le
strutture, la logica sta nel componente (stesso principio delle regole, valutate nel
concentratore). Verifica che ha motivato la versione: la v3.38 aveva tabelle, sequence,
procedura e job, ma NESSUNA API, e lo stacco era descritto solo a parole.

Interventi:
  * paragrafo «La numerazione comunale»: chi stacca è il servizio, non il concentratore; due
    sottoparagrafi nuovi — il servizio e lo stacco passo per passo, con l'SQL;
  * ⚠️ correzione del percorso «SIPO non disponibile»: registrare a posteriori un numero usato
    sulla web app senza far avanzare la sequence faceva fallire lo stacco successivo sul
    vincolo di unicità della serie. Ora la registrazione fa avanzare la sequence fino al numero
    (o colma un buco, se il numero è inferiore al prossimo);
  * architettura: quattro unità di deployment nuove, riga nella tabella dei componenti;
  * Appendice A: CHECK su registro e anno, COMMENT sulle tabelle, indice per la ricerca, grant
    all'utenza del servizio (anche sulle sequence, dentro la procedura);
  * Appendice B: inventario del nuovo contratto `ansc-numerazione-v1.yaml` (10 operazioni) e
    riga nella corrispondenza con i servizi ANSC; [R7] e [R8] aggiornati (SPEC v0.2);
  * OP-73 esteso all'utenza applicativa del servizio.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_39_servizio_numerazione.py
"""
import copy
import os
import shutil
import subprocess
import sys

import docx
import yaml
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.38.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.39.docx')
CONTRATTO = os.path.join(FIN, 'ansc-numerazione-v1.yaml')


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


# ─────────────────────────────────────────────── testo nuovo nel paragrafo sulla numerazione
SERVIZIO = [
    ('h3', 'Il servizio della numerazione: num-ansc-sipo'),
    ('p', 'La serie si governa in un componente dedicato, num-ansc-sipo, che è la quarta unità di '
          'deployment nuova dell’integrazione. Le ragioni sono tre, e sono le stesse che hanno '
          'separato i dizionari dal concentratore operativo.'),
    ('v', ('La serie è una risorsa del Comune, non di ANSC. ', 'Il servizio non parla con ANSC e '
           'non detiene credenziali verso la piattaforma: non ha bisogno della sessione OTP, '
           'del certificato server né del client ANSC, che restano confinati in all-ansc-sipo.')),
    ('v', ('Ha consumatori propri. ', 'Lo chiamano il concentratore operativo, al deposito di un '
           'atto formato in SIPO; il back-office, dalla pagina della numerazione; le maschere di '
           'SIPO, quando l’ufficiale stacca un numero per un atto da formare sulla web app. Nessuno '
           'di questi usi passa da ANSC.')),
    ('v', ('È l’unico punto in cui gli stacchi si serializzano. ', 'Tenerlo in un solo componente '
           'rende la serializzazione verificabile in un luogo solo. Il blocco sta nella base dati e '
           'non nel processo, quindi il servizio può avere più repliche senza perdere la garanzia.')),
    ('p', '⚠️ Il prezzo va dichiarato: senza il servizio della numerazione nessun atto si deposita, '
          'perché il modello evento porta il numero. Il dominio di guasto del servizio coincide '
          'quindi con quello della formazione degli atti, e il servizio va dimensionato e '
          'sorvegliato come il concentratore. Il vantaggio è che il suo guasto è circoscritto: '
          'tocca la formazione, non la consultazione né i dizionari.'),
    ('p', 'Le operazioni sono dieci e stanno nel contratto ansc-numerazione-v1.yaml [R7]; '
          'l’inventario è nell’Appendice B. La tabella che segue dice quale operazione serve a '
          'ciascun percorso.'),
    ('tab', ([
        ('Percorso', 'Chi chiama', 'Operazioni'),
        ('Atto formato in SIPO', 'Il concentratore operativo, con il token dell’operatore',
         'creaPrenotazione con l’atto di SIPO, dopo il pre-filtro e prima del deposito; '
         'aggiornaNumero a Utilizzato con l’identificativo nazionale, a deposito riuscito; '
         'aggiornaNumero a Scartato, se l’atto è abbandonato.'),
        ('Atto formato sulla web app', 'L’ufficiale, dalla pagina della numerazione',
         'creaPrenotazione; poi aggiornaNumero a Utilizzato o a Scartato; all’aggancio dell’atto '
         'in SIPO, aggiornaNumero a Evento acquisito.'),
        ('Parto plurimo', 'L’ufficiale, prima della prenotazione R009',
         'creaPrenotazione con tanti numeri quanti i nati; ciascun numero segue poi la sorte del '
         'proprio atto.'),
        ('SIPO non disponibile', 'L’ufficiale, al ritorno di SIPO',
         'creaNumero, che registra il numero già usato e tiene coerente la sequence.'),
        ('Modalità di emergenza', 'Il concentratore operativo, al rientro',
         'Come per l’atto formato in SIPO.'),
        ('Governo delle serie', 'Amministratore, Supporto, Auditor',
         'elencaSerie e leggiSerie, con i buchi; creaSerie all’installazione o se il job di '
         'dicembre non ha creato le serie; elencaPrenotazioni, leggiPrenotazione, elencaNumeri, '
         'leggiNumero per la ricerca.'),
    ], [1.4, 1.8, 3.4])),

    ('h3', 'Lo stacco, passo per passo'),
    ('p', 'Lo stacco è la sola operazione del servizio che consuma la sequence, ed è quella da cui '
          'dipendono le due garanzie. Si svolge in una sola transazione.'),
    ('v', ('1. Controllo dei valori. ', 'Il registro deve appartenere alla decodifica ANSC_01; '
           'l’anno deve essere il corrente o il precedente. Il nome della sequence si compone solo '
           'da questi due valori controllati.')),
    ('v', ('2. Serializzazione. ', 'Il servizio acquisisce il blocco esclusivo su '
           'DATI_ATTO_ESTERNO. Da qui alla conferma nessun altro stacco procede; le letture '
           'continuano.')),
    ('v', ('3. Riuso. ', 'Se la richiesta indica un atto di SIPO che ha già un numero Disponibile '
           'nella stessa serie — un deposito respinto e ritentato — il servizio restituisce quel '
           'numero, conferma e si ferma: nessun valore della sequence è consumato.')),
    ('v', ('4. Richiesta. ', 'Inserisce la riga di PRENOTAZIONE_ATTI con registro, anno, quantità, '
           'operatore e istante.')),
    ('v', ('5. Numeri. ', 'Per ciascuno dei numeri richiesti prende il valore successivo della '
           'sequence e inserisce la riga di DATI_ATTO_ESTERNO, nello stato Disponibile e legata '
           'all’atto se indicato. Se un valore risultasse già registrato, il vincolo di unicità '
           'della serie lo respinge: la transazione si annulla e l’errore segnala una serie '
           'incoerente, da esaminare. Non accade se la registrazione a posteriori segue la regola '
           'del passo successivo.')),
    ('v', ('6. Conferma. ', 'La conferma rende visibili richiesta e numeri e rilascia il blocco. Se '
           'la transazione fallisce dopo aver preso valori dalla sequence, quei valori diventano '
           'buchi: è la finestra residua, che la lettura della serie rende visibile.')),
    ('mono', [
        '-- passo 2',
        'LOCK TABLE ANSC_USR.DATI_ATTO_ESTERNO IN EXCLUSIVE MODE;',
        '-- passo 3',
        'SELECT nr_atto_comunale FROM ANSC_USR.DATI_ATTO_ESTERNO',
        ' WHERE ty_evento = :registro AND nr_anno_atto = :anno',
        "   AND id_atto_sipo = :atto AND cd_stato_numero_prenotato = '1';",
        '-- passo 4',
        'INSERT INTO ANSC_USR.PRENOTAZIONE_ATTI (ty_evento, nr_anno_atto,',
        '       nr_id_ansc_da_prenotare, cd_operatore_prenotazione)',
        'VALUES (:registro, :anno, :quantita, :operatore)',
        'RETURNING id_prenotazione_atti INTO :id_prenotazione;',
        '-- passo 5, ripetuto per ciascun numero',
        'SELECT ANSC_USR.SEQ_NUM_COMUNALE_<registro>_<anno>.NEXTVAL FROM DUAL;',
        'INSERT INTO ANSC_USR.DATI_ATTO_ESTERNO (id_prenotazione_atti, nr_atto_comunale,',
        '       ty_evento, nr_anno_atto, id_atto_sipo)',
        'VALUES (:id_prenotazione, :numero, :registro, :anno, :atto);',
        '-- passo 6',
        'COMMIT;',
    ]),
    ('p', 'La registrazione a posteriori di un numero già usato fuori da SIPO — il percorso «SIPO '
          'non disponibile» — si svolge nello stesso blocco esclusivo e confronta il numero con '
          'il prossimo valore della sequence. Le sequence sono NOCACHE, quindi il prossimo valore '
          'è quello che il catalogo di Oracle riporta come ultimo numero.'),
    ('mono', [
        'SELECT LAST_NUMBER FROM ALL_SEQUENCES',
        " WHERE SEQUENCE_OWNER = 'ANSC_USR'",
        "   AND SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_<registro>_<anno>';",
    ]),
    ('v', ('Numero già registrato. ', 'La registrazione è respinta: due atti avrebbero lo stesso '
           'numero, ed è un caso da esaminare con l’ufficiale.')),
    ('v', ('Numero inferiore al prossimo e non registrato. ', 'Colma un buco della serie e si '
           'registra nello stato Utilizzato.')),
    ('v', ('Numero uguale o superiore al prossimo. ', 'Il servizio prende valori dalla sequence '
           'finché non raggiunge il numero, e registra quello. Così nessuno stacco successivo può '
           'restituirlo. I valori eventualmente scavalcati diventano buchi, visibili nella '
           'lettura della serie.')),
]

PERCORSO_4 = ('Si registra il numero già usato, con l’identificativo nazionale dell’atto. ⚠️ La '
              'sequence non sa che quel numero è stato consumato: la registrazione la fa avanzare '
              'fino al numero, oppure colma un buco se il numero è inferiore al prossimo, come '
              'descritto in «Lo stacco, passo per passo». Un numero già registrato è respinto.')

# ─────────────────────────────────────────────── Appendice A: completamenti della DDL
DDL_PRENOTAZIONE_CK = [
    "  CONSTRAINT prenotazione_atti_reg_ck CHECK (ty_evento IN ('1','2','3','4','5')),",
    "  CONSTRAINT prenotazione_atti_anno_ck CHECK (nr_anno_atto BETWEEN 2023 AND 9999)",
]
DDL_DATI_CK = [
    "  CONSTRAINT dati_atto_esterno_reg_ck CHECK (ty_evento IN ('1','2','3','4','5')),",
]
DDL_DOPO_INDICE = [
    'CREATE INDEX ANSC_USR.IX_PRENOTAZIONE_ATTI_SERIE',
    '  ON ANSC_USR.PRENOTAZIONE_ATTI (ty_evento, nr_anno_atto, ts_prenotazione)',
    '  TABLESPACE ANSC_USR;',
    'COMMENT ON TABLE ANSC_USR.PRENOTAZIONE_ATTI IS',
    "  'Richieste di numeri comunali, per registro (ANSC_01) e anno. Struttura di Side.';",
    'COMMENT ON TABLE ANSC_USR.DATI_ATTO_ESTERNO IS',
    "  'Ogni numero comunale staccato e la sua sorte. Struttura di Side.';",
]
DDL_GRANT_SEQ = [
    "      EXECUTE IMMEDIATE 'GRANT SELECT ON ANSC_USR.SEQ_NUM_COMUNALE_' || r || '_' ||",
    "        p_anno || ' TO <UTENZA_NUM>';",
]
DDL_GRANT = [
    '',
    '-- Utenza applicativa del servizio num-ansc-sipo (nome da assegnare, OP-73).',
    '-- Un privilegio sulla tabella basta anche per LOCK TABLE; le sequence ricevono',
    '-- il grant dalla procedura che le crea.',
    'GRANT SELECT, INSERT, UPDATE ON ANSC_USR.PRENOTAZIONE_ATTI TO <UTENZA_NUM>;',
    'GRANT SELECT, INSERT, UPDATE ON ANSC_USR.DATI_ATTO_ESTERNO TO <UTENZA_NUM>;',
    'GRANT EXECUTE ON ANSC_USR.PRC_CREA_SEQ_NUM_COMUNALE TO <UTENZA_NUM>;',
]

COMPONENTE = ('Servizio della numerazione num-ansc-sipo (nuovo)',
              'Stacca i numeri comunali dalla serie per registro e anno, serializzando gli stacchi, '
              'e ne registra la sorte (Disponibile, Utilizzato, Scartato, Evento acquisito) in '
              'PRENOTAZIONE_ATTI e DATI_ATTO_ESTERNO. Lo chiamano il concentratore al deposito, il '
              'back-office e le maschere SIPO. Non parla con ANSC e non detiene credenziali verso '
              'la piattaforma.',
              'Sì — pod dedicato')

STORIA = ('11/10/2026', '3.39',
          'Architettura (componenti) · SIPO e la web app («La numerazione comunale») · '
          'Appendice A · Appendice B · Riferimenti · Open Point',
          'La numerazione comunale ha un servizio dedicato, num-ansc-sipo: quarta unità di '
          'deployment nuova, senza credenziali verso ANSC, chiamata dal concentratore al '
          'deposito, dal back-office e dalle maschere SIPO. Lo stacco sta nel servizio ed è '
          'descritto passo per passo con l’SQL. Corretto il percorso «SIPO non disponibile»: '
          'registrare a posteriori un numero senza far avanzare la sequence avrebbe fatto '
          'fallire lo stacco successivo sul vincolo di unicità della serie. Nuovo contratto '
          'ansc-numerazione-v1.yaml con dieci operazioni, inventariate nell’Appendice B; '
          'specifica delle API alla v0.2. Completata la DDL con i vincoli su registro e anno, i '
          'commenti delle tabelle, l’indice per la ricerca e i grant all’utenza del servizio.')


def corpo(d):
    out = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith('}p'):
            out.append(Paragraph(el, d))
        elif el.tag.endswith('}tbl'):
            out.append(Table(el, d))
    return out


def para_ddl(d, testo):
    """Il paragrafo di Appendice A la cui riga di DDL coincide con il testo dato."""
    trovati = [p for p in d.paragraphs if p.text == testo]
    if len(trovati) != 1:
        raise SystemExit(f'riga di DDL attesa una volta, trovata {len(trovati)}: {testo}')
    return trovati[0]


def dopo(d, ancora, righe):
    """Inserisce righe di DDL dopo il paragrafo «ancora», ereditandone il carattere."""
    for r in reversed(righe):
        nuovo = copy.deepcopy(ancora._p)
        ancora._p.addnext(nuovo)
        np = Paragraph(nuovo, ancora._parent)
        for run in np.runs[1:]:
            run._r.getparent().remove(run._r)
        if np.runs:
            np.runs[0].text = r
        else:
            np.add_run(r)


def main():
    if os.path.exists(DST):
        raise SystemExit(f'{os.path.basename(DST)} esiste già: non si sovrascrive')
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)
    fatti = []
    modello = D.trova_tabella(d, 'Oggetto', 'Ruolo')

    # ─────────────────────────── 1. architettura
    D.sostituisci(d, 'Si propongono tre unità di deployment nuove',
                  'Si propongono quattro unità di deployment nuove', attese=1, fatti=fatti)
    D.sostituisci(d, 'e non forma né firma atti. A queste si potrebbe aggiungere',
                  'e non forma né firma atti. La quarta, «num-ansc-sipo», è il servizio della '
                  'numerazione comunale: governa la serie dei numeri che il Comune assegna agli '
                  'atti e non parla con ANSC. A queste si potrebbe aggiungere',
                  attese=1, fatti=fatti, etichetta='visione d’insieme: la quarta unità')
    t = D.trova_tabella(d, 'Componente', 'Responsabilità', 'Unità di deployment autonoma')
    auto = [r for r in t.rows if r.cells[0].text.startswith('Processo di automazione')]
    assert len(auto) == 1
    riga = D.clona_riga(t, COMPONENTE)
    auto[0]._tr.addnext(riga._tr)
    fatti.append('tabella dei componenti: num-ansc-sipo')

    # ─────────────────────────── 2. il paragrafo sulla numerazione
    D.sostituisci(d, 'si serializza: il concentratore acquisisce un blocco esclusivo',
                  'si serializza: il servizio della numerazione acquisisce un blocco esclusivo',
                  attese=1, fatti=fatti)
    D.sostituisci(d, 'Il nome della sequence lo compone il concentratore da due valori',
                  'Il nome della sequence lo compone il servizio della numerazione da due valori',
                  attese=1, fatti=fatti)
    tp = D.trova_tabella(d, 'Percorso', 'Quando si stacca')
    r4 = [r for r in tp.rows if r.cells[0].text.strip() == 'SIPO non disponibile']
    assert len(r4) == 1 and 'Coincide con quello della serie' in r4[0].cells[2].text
    D.riscrivi_cella(r4[0].cells[2], PERCORSO_4)
    fatti.append('percorso «SIPO non disponibile» corretto')

    # i due sottoparagrafi nuovi vanno in coda al paragrafo, prima dell'H2 successivo
    ancora = D.h(d, 2, 'La continuità operativa quando l’indisponibilità è nostra')._p
    for tipo, c in SERVIZIO:
        if tipo == 'p':
            D.para(d, ancora, c)
        elif tipo == 'h3':
            D.para(d, ancora, c, stile='Heading 3')
        elif tipo == 'v':
            D.voce(d, ancora, *c)
        elif tipo == 'tab':
            D.tabella(d, ancora, c[0], modello, c[1])
            D.para(d, ancora, '')
        elif tipo == 'mono':
            D.ddl(d, ancora, c)
    fatti.append('«Il servizio della numerazione» e «Lo stacco, passo per passo»')

    # ─────────────────────────── 3. Appendice A
    ck = para_ddl(d, '  CONSTRAINT prenotazione_atti_qta_ck CHECK (nr_id_ansc_da_prenotare >= 1)')
    D.testo_di(ck, ck.text + ',')
    dopo(d, ck, DDL_PRENOTAZIONE_CK)
    st = para_ddl(d, "    CHECK (cd_stato_numero_prenotato IN ('1','2','3','4'))")
    D.testo_di(st, st.text + ',')
    dopo(d, st, [r.rstrip(',') for r in DDL_DATI_CK])
    ix = para_ddl(d, '  ON ANSC_USR.DATI_ATTO_ESTERNO (id_atto_sipo) TABLESPACE ANSC_USR;')
    dopo(d, ix, DDL_DOPO_INDICE)
    seq = para_ddl(d, "        || ' NOCYCLE NOCACHE ORDER';")
    dopo(d, seq, DDL_GRANT_SEQ)
    inst = para_ddl(d, '-- All’installazione: EXEC ANSC_USR.PRC_CREA_SEQ_NUM_COMUNALE(<anno corrente>);')
    dopo(d, inst, DDL_GRANT)
    fatti.append('Appendice A: CHECK, COMMENT, indice, grant')

    # ─────────────────────────── 4. Appendice B
    spec = yaml.safe_load(open(CONTRATTO, encoding='utf-8'))
    ops = [(m.upper() + ' ' + p, o['operationId'], o['summary'], ', '.join(o.get('x-ruoli', [])))
           for p, v in spec['paths'].items() for m, o in v.items()
           if m in ('get', 'post', 'put', 'patch', 'delete')]
    assert len(ops) == 10
    inv = [t for t in d.tables
           if t.rows[0].cells[0].text.strip() == 'Metodo e percorso'
           and any('allineaDizionari' in r.cells[1].text for r in t.rows)]
    assert len(inv) == 1
    assert [c.text.strip() for c in inv[0].rows[0].cells] == [
        'Metodo e percorso', 'Identificativo', 'Titolo breve', 'Ruoli']
    nuova = copy.deepcopy(inv[0]._tbl)
    nt = Table(nuova, inv[0]._parent)
    for r in list(nt.rows)[2:]:
        nuova.remove(r._tr)
    prima_riga = nt.rows[1]
    for valori in ops:
        D.clona_riga(nt, valori)
    nuova.remove(prima_riga._tr)
    corr = D.h(d, 2, 'Corrispondenza con i servizi ANSC consumati')._p
    D.para(d, corr, 'Servizio della numerazione — num-ansc-sipo', stile='Heading 3')
    D.para(d, corr, spec['info']['summary'])
    corr.addprevious(nuova)
    D.para(d, corr, '')
    tc = D.trova_tabella(d, 'Operazione del componente', 'Servizio ANSC invocato')
    D.clona_riga(tc, ('Operazioni del servizio della numerazione',
                      'Nessuno. Il numero comunale è un dato del Comune e nessun servizio ANSC lo '
                      'assegna: il concentratore lo chiede al servizio prima di R009 e lo inserisce '
                      'nel modello evento.'))
    D.sostituisci(d, 'Sono tre e sono elencati al paragrafo seguente.',
                  'Sono quattro e sono elencati al paragrafo seguente.', attese=1, fatti=fatti)
    fatti.append('Appendice B: inventario del servizio (10 operazioni) e corrispondenza')

    # ─────────────────────────── 5. Riferimenti, Open Point, storia
    D.sostituisci(d, 'ansc-v1.yaml · ansc-dizionari-v1.yaml · ansc-automazione-v1.yaml',
                  'ansc-v1.yaml · ansc-dizionari-v1.yaml · ansc-automazione-v1.yaml · '
                  'ansc-numerazione-v1.yaml', attese=1, fatti=fatti, etichetta='[R7]')
    D.sostituisci(d, 'SPEC_API_Integrazione-ANSC_v0.1.docx', 'SPEC_API_Integrazione-ANSC_v0.2.docx',
                  attese=1, fatti=fatti, etichetta='[R8]')
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == '#'
             and any(r.cells[0].text.strip() == 'OP-73' for r in t.rows))
    op = [r for r in t.rows if r.cells[0].text.strip() == 'OP-73'][0]
    vecchio = op.cells[2].text
    assert vecchio.startswith('Due conferme e un permesso.')
    D.riscrivi_cella(op.cells[2], vecchio.replace(
        'Due conferme e un permesso.', 'Due conferme e due permessi.', 1)
        + ' (d) L’utenza applicativa del servizio num-ansc-sipo, con cui il servizio raggiunge '
          'ANSC_USR, e i suoi grant: nel DDL il nome è un segnaposto.')
    D.storia(d, *STORIA)
    fatti.append('[R7], [R8], OP-73, storia')

    d.save(DST)
    dopo_c = commenti(DST)
    if dopo_c != prima:
        raise SystemExit(f'commenti: {prima} prima, {dopo_c} dopo')
    print('\n'.join('· ' + f for f in fatti))
    print(f'commenti preservati: {dopo_c}')


if __name__ == '__main__':
    main()
