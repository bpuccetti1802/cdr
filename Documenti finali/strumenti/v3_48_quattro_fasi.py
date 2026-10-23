# -*- coding: utf-8 -*-
"""v3.48 dell'analisi ANSC, dalla v3.47 (23/10/2026): riscrittura di «Le quattro fasi».

Il capitolo «Il processo di gestione dei dati ANSC in SIPO» è il riferimento, e la sezione non
lo seguiva più. Le quattro fasi restano, ma raggruppano le pagine di quel capitolo:
  A  Tipo atto, Dati anagrafici, Dettaglio — R005 obbligatoria per l'intestatario (la ricerca,
     non il risultato); l'intestatario viene da SOGGETTI_USECASE;
  B  Riepilogo — UC, firmatari, documenti (R001), verifica locale (semplificabile nel primo
     rilascio);
  C  Anteprima e VALIDA — numero comunale, modello evento, R010, R009;
  D  Firme, «Definisci» e «Chiudi» — R006/R012, attestazione R007, firma USC R007, R014.
Nuova la regola del ritorno ai dati dopo la validazione (stato 0 e nuova validazione dello
stesso atto) al posto di «ridepositare o annullare»; tolta R013 dalla validazione (è la
rettifica); dichiarata la fine della correzione fino a mezzanotte. Riallineati la tabella delle
fasi, la tabella dei passi (13), la figura (diagrammi_flusso.py) e tre rimandi alla «fase A»
come luogo della determinazione dell'UC.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_48_quattro_fasi.py
"""
import copy
import os
import subprocess
import sys

import docx

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.47.docx')
OUT = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.48.docx')
fatti = []


def unico(doc, inizio):
    t = [p for p in doc.paragraphs if p.text.startswith(inizio)]
    if len(t) != 1:
        raise SystemExit(f'«{inizio[:50]}»: {len(t)} paragrafi')
    return t[0]


def pezzi(p, parti):
    if D.ha_commenti(p._p):
        raise SystemExit('paragrafo con commento: ' + p.text[:60])
    modello = p.runs[0]._r.rPr if p.runs else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for testo, bold in parti:
        r = p.add_run(testo)
        if modello is not None:
            r._r.insert(0, copy.deepcopy(modello))
        r.bold = bold


INTRO = ('Il percorso si legge in quattro fasi, e ciascuna ha una condizione di uscita che va '
         'soddisfatta prima di passare alla successiva. Le fasi raggruppano le pagine del processo '
         'descritto nel capitolo «Il processo di gestione dei dati ANSC in SIPO», che è il '
         'riferimento: le pagine dicono che cosa vede l’operatore, le fasi che cosa deve essere '
         'vero perché la lavorazione prosegua e in quale momento l’atto smette di essere '
         'modificabile.')

FASI = [
    ('Fase A — i dati anagrafici e il dettaglio. ',
     'L’operatore sceglie il tipo atto e compila l’atto nelle maschere che già conosce. Il tipo '
     'atto restringe gli UC possibili, e la tabella dei soggetti per UC (SOGGETTI_USECASE) dice '
     'quale soggetto sarà l’intestatario [DA VERIFICARE: quale UC si legga quando il Modello ne '
     'indirizza più d’uno con intestatari diversi, perché la determinazione vera avviene nel '
     'Riepilogo]. Nella pagina di dettaglio si cercano i soggetti in ANSC con R005, che richiede '
     'la sessione OTP: per l’intestatario la ricerca è obbligatoria, il risultato no. Un soggetto '
     'non trovato non ferma l’atto, ma senza il suo identificativo gli automatismi di ANSC non '
     'scattano e il deposito richiederebbe il collegamento manuale, che la nota di processo '
     'sconsiglia. Nella stessa pagina si mostrano i dati che ANSC richiede e SIPO non ha: finché '
     'mancano non si prosegue.'),
    ('Fase B — il Riepilogo. ',
     'La pagina determina l’UC: la determinazione è automatica, e la scelta da parte '
     'dell’operatore è ammessa soltanto quando la configurazione risulta ambigua, e resta '
     'registrata come tale. Determinato l’UC, la pagina ne mostra il codice e la descrizione, i '
     'firmatari con il contrassegno di firma, che l’operatore può modificare, i documenti che '
     'l’UC prevede e, se l’atto è già stato validato, il suo identificativo nazionale. Qui si '
     'caricano i documenti — firmati digitalmente o scansionati — e gli eventuali atti a testo '
     'libero, e si inviano ad ANSC con R001 (il punto esatto in cui farlo è aperto, OP-76). Qui '
     'infine si esegue la verifica locale contro la configurazione, che non chiama ANSC e '
     'intercetta le mancanze prima che un identificativo nazionale sia consumato. Nel primo '
     'rilascio la verifica può essere semplificata; in ogni caso non si prosegue finché i dati '
     'obbligatori dell’UC non sono soddisfatti.'),
    ('Fase C — l’Anteprima e «VALIDA». ',
     'Si stacca il numero comunale definitivo dal servizio num-ansc-sipo, si costruisce il '
     'modello evento e, se serve, se ne mostra il testo con R010. Il pulsante «VALIDA» chiama '
     'R009, che deposita l’atto in ANSC nello stato 2 CONFERMATO e ne restituisce '
     'l’identificativo nazionale. Se R009 respinge i dati, in ANSC non resta nulla: l’operatore '
     'torna alle pagine dei dati, corregge e valida di nuovo, e il numero comunale resta '
     'all’atto.'),
    ('Fase D — le firme e la definizione. ',
     'Seguono la firma dei dichiaranti e degli altri firmatari, quando prevista (R006 cartacea o '
     'R012 elettronica; senza firmatari la pagina si salta), l’attestazione di conformità (R007, '
     '11 FIRMATO PER CONFORMITÀ) e la firma dell’ufficiale con le proprie credenziali di firma '
     '(R007, 5 FIRMATO DA USC). Con «Definisci» l’atto passa in SIPO a 2 DEFINITO e, se l’UC la '
     'prevede, parte la comunicazione anagrafica (R014, secondo comunicaz_anagrafica_usecase); '
     '«Chiudi» chiude la lavorazione. La firma dell’ufficiale è il punto di non ritorno, e va '
     'resa evidente all’operatore prima che lo attraversi, non dopo: da lì ciò che si può ancora '
     'fare passa per gli istituti di ANSC — annotazione, rettifica, annullamento — e non per la '
     'maschera di compilazione. Viene meno, con ciò, la prassi di correggere l’atto fino alla '
     'mezzanotte del giorno in cui è stato redatto.'),
]

RITORNO = [
    ('Il ritorno ai dati dopo la validazione. ', True),
    ('Fino alla firma dell’ufficiale l’operatore può tornare alle pagine dei dati. Se i dati '
     'cambiano, il modello evento non è più quello validato: la sua impronta differisce da '
     'TXT_HASH_VALIDATO, l’atto passa allo stato 0 (BOZZA_APPLICATIVA) e le pagine delle firme '
     'non sono raggiungibili finché una nuova validazione dello stesso atto — con il suo '
     'identificativo nazionale, non con uno nuovo — non lo riporta a 2 CONFERMATO. Con quale '
     'campo, id o idAnsc, ANSC riconosca l’atto già validato è aperto (OP-75). Dopo le firme dei '
     'dichiaranti il ritorno ai dati è consentito soltanto se tutti i firmatari possono firmare di '
     'nuovo; altrimenti resta l’eliminazione dell’evento con R011 [DA VERIFICARE sul '
     'comportamento di ANSC]. Se la correzione cambia l’UC, l’atto validato non si rivalida come '
     'un altro caso d’uso: si elimina con R011 e se ne valida uno nuovo [DA VERIFICARE].', False),
]

TAB_FASI = [
    ('A — Dati anagrafici e dettaglio',
     'I soggetti sono stati cercati in ANSC, obbligatoriamente l’intestatario, e i dati che ANSC '
     'richiede e SIPO non ha sono stati completati.',
     'IN_LAVORAZIONE',
     'Nulla: l’atto resta pienamente modificabile.'),
    ('B — Riepilogo',
     'L’UC è determinato: una sola riga di configurazione applicabile, oppure la scelta '
     'dell’operatore fra le candidate in caso di ambiguità. I documenti obbligatori sono stati '
     'accettati da ANSC e la verifica locale non segnala mancanze.',
     'UC_DETERMINATO, DOCUMENTI_ACQUISITI, PREVALIDATO',
     'Nulla in SIPO; in ANSC un allegato trasmesso si sostituisce, non si modifica.'),
    ('C — Anteprima e «VALIDA»',
     'Il numero comunale è staccato e R009 ha validato l’atto restituendo l’identificativo '
     'nazionale (2 CONFERMATO).',
     'VALIDATO',
     'L’identificativo nazionale è consumato: una correzione porta l’atto allo stato 0 e a una '
     'nuova validazione dello stesso atto, non a un atto nuovo.'),
    ('D — Firme e definizione',
     'Le firme dei dichiaranti sono raccolte, l’attestazione di conformità è firmata '
     '(11 FIRMATO PER CONFORMITÀ), l’atto è firmato dall’ufficiale (5 FIRMATO DA USC) e definito '
     'in SIPO (2 DEFINITO): l’atto è formato.',
     'FIRMATO',
     'La modifica dell’atto in SIPO, anche nello stesso giorno: da qui si interviene con gli '
     'istituti previsti da ANSC.'),
]

TAB_PASSI = [
    ('1 — Tipo atto e dati anagrafici',
     'Le maschere di SIPO; SOGGETTI_USECASE per sapere quale soggetto è l’intestatario.',
     'L’atto nelle tabelle di SIPO, nello stato 1 In lavorazione: da qui l’identificativo '
     'ID_ATTO_SIPO, che correla tutto il resto.', '—'),
    ('2 — Dettaglio: la ricerca dei soggetti (R005)',
     'ANSC, per identificativo o per dati anagrafici. Richiede l’OTP di sessione.',
     'ANSC_SOGGETTO_IDENTIFICATIVI e i due codici ANPR su ANAG_USR.SOGGETTO; la riga dell’atto '
     'in ANSC_STATO_ATTO, in IN_LAVORAZIONE, con la pagina di ripresa. Per l’intestatario la '
     'ricerca è obbligatoria, il risultato no.', 'SOGGETTO'),
    ('3 — Riepilogo: la determinazione dell’UC',
     'ANSC_CFG_UC per la coppia tipo atto e Modello, nella baseline ATTIVA, in ordine di '
     'NUM_PRIORITA; per ciascuna riga la logica richiamata in LOGICHE_DI_SCELTA.',
     'ANSC_STATO_ATTO: ID_UC_ANSC, COD_LOGICA, COD_ORIGINE_UC, ID_VERSIONE e fase UC_DETERMINATO.',
     'DETERMINAZIONE'),
    ('4 — Riepilogo: i firmatari',
     'SOGGETTI_USECASE: quali soggetti dell’UC firmano (fg_firmatario).',
     'firmatari: chi firma, con il contrassegno che l’operatore può modificare.', '—'),
    ('5 — Riepilogo: i documenti (R001)',
     'ALLEGATI_USECASE: quali documenti l’UC richiede.',
     'ALLEGATO: il file, la sua impronta, lo stato restituito da ANSC e l’identificativo '
     'dell’allegato; fase DOCUMENTI_ACQUISITI. Solo lo stato «Inserito» consente di proseguire.',
     'ALLEGATI'),
    ('6 — Riepilogo: la verifica locale (RF-9)',
     'ANSC_CFG_SEZIONE per sapere quali blocchi servono, ANSC_CFG_CAMPO per i campi '
     'obbligatori, ALLEGATI_USECASE per i documenti, V_ANSC_DIZ_VALIDO per i valori ammessi.',
     'Nulla di definitivo: l’esito è l’elenco di ciò che manca, mostrato all’operatore; fase '
     'PREVALIDATO quando non manca nulla.', 'PREVERIFICA'),
    ('7 — Anteprima: il numero comunale',
     'Il servizio num-ansc-sipo.',
     'DATI_ATTO_ESTERNO: il numero e la sua sorte; ANSC_STATO_ATTO.NUM_COMUNALE.',
     '— (lo registra il servizio della numerazione)'),
    ('8 — Anteprima: il modello evento',
     'ANSC_CFG_SEZIONE, ANSC_CFG_CAMPO, ANSC_CFG_FORMULA e i dizionari per le transcodifiche.',
     'ANSC_STATO_ATTO.TXT_EVENTO: il modello evento in redazione, aggiornato a ogni salvataggio.',
     'PAYLOAD'),
    ('9 — Anteprima del testo (R010), facoltativa',
     'Il modello evento: R010 richiede l’intero modello, non un identificativo.',
     'Nulla: è una lettura. Serve a mostrare all’ufficiale il testo dell’atto prima che '
     'l’identificativo nazionale sia consumato.', 'ANTEPRIMA'),
    ('10 — «VALIDA»: la validazione (R009)',
     'Il modello evento, in ANSC_STATO_ATTO.TXT_EVENTO.',
     'ANSC_STATO_ATTO: ID_ANSC, ID_EVENTO_ANSC, ID_OPERAZIONE_ANSC, stato CONFERMATO, fase '
     'VALIDATO e l’impronta del modello validato. Da qui l’identificativo nazionale è consumato.',
     'DEPOSITO'),
    ('11 — Le firme dei dichiaranti (R006, R012)',
     'firmatari: chi firma e in quale modo.',
     'firmatari: lo stato della firma; ANSC_STATO_ATTO: stato FIRMATO_DICHIARANTE.',
     'FIRMA_DICH'),
    ('12 — L’attestazione di conformità (R007)',
     'ALLEGATO: i documenti inclusi nell’attestazione.',
     'ANSC_STATO_ATTO: stato FIRMATO_CONFORMITA.', 'FIRMA_CONFORMITA'),
    ('13 — La firma dell’USC (R007), «Definisci» e «Chiudi»',
     'Lo stato reale dell’atto, quando occorre riprendere una lavorazione interrotta; '
     'comunicaz_anagrafica_usecase per la comunicazione anagrafica (R014).',
     'ANSC_STATO_ATTO: stato FIRMATO_USC, che è l’atto formato, e fase FIRMATO; in SIPO l’atto '
     'passa a 2 DEFINITO.', 'FIRMA_USC'),
]


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)

    pezzi(unico(doc, 'Il percorso si legge in quattro fasi'), [(INTRO, False)])
    vecchi = ['Fase A — la lavorazione in SIPO.', 'Fase B — Gli allegati.',
              'Fase C — la validazione.', 'Fase D — le firme ANSC.']
    ultimo = None
    for inizio, (testa, corpo) in zip(vecchi, FASI):
        ultimo = unico(doc, inizio)
        pezzi(ultimo, [(testa, True), (corpo, False)])
    vuoto = ultimo._p.getnext()
    p = D.para(doc, vuoto, pezzi=RITORNO)
    fatti.append('6.1: introduzione, quattro fasi e il ritorno ai dati riscritti')

    t = D.trova_tabella(doc, 'fase', 'condizione di uscita', 'cod_fase raggiunto')
    D.riscrivi_sicura(t, TAB_FASI)
    t = D.trova_tabella(doc, 'passo', 'che cosa legge', 'che cosa scrive', 'traccia')
    D.riscrivi_sicura(t, TAB_PASSI)
    fatti.append(f'tabella delle fasi (4) e dei passi ({len(TAB_PASSI)}) riscritte')

    subprocess.run([sys.executable, os.path.join(QUI, 'diagrammi_flusso.py')], check=True)
    D.sostituisci_immagine(doc, 'Il flusso operativo nelle sue quattro fasi',
                           os.path.join(QUI, 'img', 'flusso_operativo.png'))
    fatti.append('figura del flusso operativo rigenerata (13 passi)')

    for vecchio, nuovo in [
        ('La fase A non si chiude senza UC: campi, documenti e formule ne dipendono.',
         'Il Riepilogo (fase B) non si chiude senza UC: campi, documenti e formule ne dipendono.'),
        ('È l’UC determinato nella fase A.', 'È l’UC determinato nel Riepilogo (fase B).'),
        ('finché la fase A non ha determinato l’UC',
         'finché il Riepilogo (fase B) non ha determinato l’UC'),
    ]:
        D.sostituisci(doc, vecchio, nuovo, attese=1, fatti=fatti)

    D.storia(doc, '23/10/2026', '3.48', 'Il flusso operativo (Le quattro fasi) · Premesse (RF-19) '
             '· Costruzione del payload',
             'Riscritta la sezione «Le quattro fasi» sul processo del capitolo «Il processo di '
             'gestione dei dati ANSC in SIPO»: le fasi raggruppano le pagine (dati anagrafici e '
             'dettaglio; riepilogo con UC, firmatari, documenti e verifica locale; anteprima e '
             '«VALIDA»; firme, «Definisci» e «Chiudi»). Nuova la regola del ritorno ai dati dopo '
             'la validazione (stato 0 e nuova validazione dello stesso atto); la ricerca R005 è '
             'obbligatoria per l’intestatario, il risultato no; con la firma dell’ufficiale viene '
             'meno la correzione fino a mezzanotte. Riallineate la tabella delle fasi, la tabella '
             'dei passi (13) e la figura; l’UC si determina nel Riepilogo.')
    D.frontespizio(doc, '3.48', 'ANALISI_Integrazione-ANSC_v3.48')
    doc.save(OUT)
    for f in fatti:
        print(' -', f)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
