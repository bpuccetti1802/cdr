# -*- coding: utf-8 -*-
"""v3.53 dell'analisi ANSC, dalla v3.52 dell'utente (27/10/2026).

Nella v3.52 l'utente ha riportato nell'Appendice A le DDL riviste con gli sviluppatori. Questo
script corregge gli errori di quelle DDL e adegua il resto del documento. Decisioni dell'utente:

  - TIPO_LOGICHE_DI_SCELTA e LOGICHE_DI_SCELTA sono deprecate e si tolgono dal documento:
    ANSC_CFG_UC.LOGICA_DI_SCELTA e ALLEGATI_USECASE.cd_logica restano codici liberi, senza
    chiave esterna (id_dominio si toglie);
  - ANSC_CFG_UC non ha più il Modello: il tipo atto di SIPO si lega all'UC con ID_TIPO_ATTO,
    UK (ID_VERSIONE, COD_UC_ANSC, ID_TIPO_ATTO);
  - ANSC_CFG_CAMPO: ALIAS_CAMPO_SIPO è popolato a mano; OBBLIGATORIETA_EVENTO si stabilirà in
    analisi; OPERATIVO resta 0/1; ID_DECODIFICA_ANSC si rimette; FK e UK tolte di proposito;
  - ANSC_STATO_ATTO: i meccanismi rimasti senza colonna si tolgono dal testo (impronta del
    modello validato, sospensione, esito qualificato, baseline sull'atto, ufficiale, campo id).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_53_ddl_sviluppatori.py
"""
import os
import sys

import docx
from docx.text.paragraph import Paragraph

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
import docx_comune as D  # noqa: E402

BASE = os.path.dirname(QUI)
SRC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.52.docx')
OUT = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.53.docx')
IMG = os.path.join(QUI, 'img')


# ───────────────────────────────────────────────────────────── primitive locali
def _paragrafi(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs


def sost(doc, vecchio, nuovo, attese=1):
    """Come D.sostituisci, ma se il testo sta in un solo run cambia solo quel run.

    Così i paragrafi con un avvio in grassetto («Il ritorno ai dati…») non perdono la
    formattazione; altrimenti si ripiega su testo_di, che protegge i commenti.
    """
    n, visti = 0, set()
    for p in _paragrafi(doc):
        if p._p in visti or vecchio not in p.text:
            continue
        visti.add(p._p)
        for r in p.runs:
            if vecchio in r.text:
                r.text = r.text.replace(vecchio, nuovo)
                break
        else:
            D.testo_di(p, p.text.replace(vecchio, nuovo))
        n += 1
    if n != attese:
        raise SystemExit(f'«{vecchio[:70]}»: attese {attese}, trovate {n}')


def par(doc, inizio):
    """Il paragrafo del corpo che comincia con il testo dato (uno solo)."""
    trovati = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(trovati) != 1:
        raise SystemExit(f'«{inizio[:60]}»: {len(trovati)} paragrafi')
    return trovati[0]


def via(elemento):
    if D.ha_commenti(elemento):
        raise SystemExit('elemento con commento da eliminare')
    elemento.getparent().remove(elemento)


def via_par(doc, inizio, quanti=1):
    trovati = [p for p in doc.paragraphs if p.text.strip().startswith(inizio)]
    if len(trovati) != quanti:
        raise SystemExit(f'«{inizio[:60]}»: {len(trovati)} paragrafi, attesi {quanti}')
    for p in trovati:
        via(p._p)


def riga_tabella(t, primo):
    for r in t.rows:
        if r.cells[0].text.strip().startswith(primo):
            return r
    raise SystemExit('riga non trovata: ' + primo)


def via_riga(t, primo):
    via(riga_tabella(t, primo)._tr)


def dopo(p, testo):
    """Un paragrafo nuovo dopo p, con lo stile e il carattere del primo run di p."""
    return D.riga_dopo(None, p, testo)


# ───────────────────────────────────────────────────────────── 1. DDL (Appendice A)
def ddl(doc):
    # ANSC_STATO_ATTO: virgola dopo l'ultimo vincolo e parentesi in più
    p = par(doc, "CONSTRAINT CK_ANSC_STATO_ATTO_CASA CHECK (FLG_CASA_COMUNALE IN ('S','N')),")
    D.testo_di(p, p.text.rstrip().rstrip(','))
    stray = Paragraph(p._p.getnext(), p._parent)
    assert stray.text.strip() == ')', 'parentesi isolata non trovata: ' + stray.text
    via(stray._p)
    sost(doc, 'ID_MOTIVO_RECUPERO      CHAR,', 'ID_MOTIVO_RECUPERO      VARCHAR2(3 CHAR),')

    # Le due tabelle del catalogo delle logiche, deprecate: via dall'Appendice
    for titolo in ('TIPO_LOGICHE_DI_SCELTA', 'LOGICHE_DI_SCELTA'):
        for el in D.sezione(doc, titolo, livello=3):
            via(el)
    for q in [q for q in doc.paragraphs if q.text.strip() == 'Deprecata']:
        via(q._p)

    # ANSC_CFG_UC
    sost(doc, 'UNIQUE (ID_VERSIONE, COD_UC_ANSC, ID_CONF_TIPO_ATTO, ID_MODELLO_ATTO),',
         'UNIQUE (ID_VERSIONE, COD_UC_ANSC, ID_TIPO_ATTO),')
    p = par(doc, 'CONSTRAINT FK_ANSC_CFG_UC_LOGICA')
    precedente = Paragraph(p._p.getprevious(), p._parent)
    D.testo_di(precedente, precedente.text.rstrip().rstrip(','))
    for _ in range(2):
        via(p._p.getnext())
    via(p._p)
    sost(doc, 'CREATE INDEX ANSC_USR.IX_ANSC_CFG_UC_MODELLO',
         'CREATE INDEX ANSC_USR.IX_ANSC_CFG_UC_TIPO_ATTO')
    p = par(doc, 'ON ANSC_USR.ANSC_CFG_UC (ID_VERSIONE, ID_CONF_TIPO_ATTO, ID_MODELLO_ATTO,')
    D.testo_di(p, '  ON ANSC_USR.ANSC_CFG_UC (ID_VERSIONE, ID_TIPO_ATTO, NUM_PRIORITA)')
    succ = Paragraph(p._p.getnext(), p._parent)
    assert succ.text.strip().startswith('NUM_PRIORITA) TABLESPACE'), succ.text
    D.testo_di(succ, '  TABLESPACE ANSC_USR;')
    sost(doc, 'UC e per Modello che lo raggiunge, dentro una baseline. Senza Modello: UC del catalogo',
         'UC e per tipo atto che lo raggiunge, dentro una baseline. Senza tipo atto: UC del catalogo')
    sost(doc, "che nessun Modello raggiunge ancora.';", "che nessun tipo atto raggiunge ancora.';")
    sost(doc, 'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_UC.COD_LOGICA_DI_SCELTA IS',
         'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_UC.LOGICA_DI_SCELTA IS')
    sost(doc, "'Logica di scelta presa dal catalogo (dominio 1). Vuota significa: si applica sempre.';",
         "'Codice libero della logica di scelta, valutato dal codice applicativo; vuota: sempre.';")
    sost(doc, "'Ordine di valutazione fra gli UC dello stesso Modello: il primo la cui logica e",
         "'Ordine di valutazione fra gli UC dello stesso tipo atto: il primo la cui logica e")
    sost(doc, 'e aggiunta dopo la creazione di ANSC_TEMPLATE. Distinta da COD_LOGICA_DI_SCELTA.',
         'e aggiunta dopo la creazione di ANSC_TEMPLATE. Distinta da LOGICA_DI_SCELTA.')

    # ANSC_CFG_CAMPO: ID_DECODIFICA_ANSC torna; via i commenti su colonne tolte
    p = par(doc, '--  ID_DECODIFICA_ANSC')
    D.testo_di(p, '  ID_DECODIFICA_ANSC\t\tNUMBER,')
    for inizio in ('COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.SEZIONE_FE_ANSC IS',
                   'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ID_BUSINESS_LOGIC IS'):
        p = par(doc, inizio)
        nxt = Paragraph(p._p.getnext(), p._parent)
        via(p._p)
        while True:                       # le righe di continuazione del commento
            fine = nxt.text.rstrip().endswith("';")
            succ = nxt._p.getnext()
            via(nxt._p)
            if fine:
                break
            nxt = Paragraph(succ, p._parent)
    p = par(doc, 'COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.OGGETTO_ANSC IS')
    ultimo = Paragraph(p._p.getnext(), p._parent)
    for riga in ["COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ALIAS_CAMPO_SIPO IS",
                 "  'Nome con cui il dato e noto in SIPO: compilato a mano da chi configura.';",
                 "COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.ID_DECODIFICA_ANSC IS",
                 "  'Dizionario ANSC del campo codificato: per valore, senza chiave esterna.';",
                 "COMMENT ON COLUMN ANSC_USR.ANSC_CFG_CAMPO.OPERATIVO IS",
                 "  '0/1: se il campo va inviato ad ANSC.';"]:
        ultimo = dopo(ultimo, riga)

    # punti dopo il nome di colonna
    sost(doc, 'TS_UPD.', 'TS_UPD', attese=2)

    # ALLEGATI_USECASE: niente più dominio né chiave esterna verso le logiche
    via_par(doc, 'id_dominio           NUMBER DEFAULT 2,')
    p = par(doc, 'CONSTRAINT allegati_usecase_logica_fk')
    for _ in range(2):
        via(p._p.getnext())
    via(p._p)
    sost(doc, "riferimento al catalogo delle logiche per la condizione.';",
         "codice della condizione (cd_logica), senza catalogo.';")

    # NUMERATORE_COMUNALE: la colonna si chiama num_atto
    sost(doc, 'COMMENT ON COLUMN ANSC_USR.NUMERATORE_COMUNALE.pg_atto IS',
         'COMMENT ON COLUMN ANSC_USR.NUMERATORE_COMUNALE.num_atto IS')


# ───────────────────────────────────────────────────────────── 2. ANSC_STATO_ATTO
STATO_ATTO = [
    ('ID_STATO_ATTO', 'NUMBER (PK)', 'Chiave tecnica (IDENTITY).'),
    ('ID_ATTO_SIPO', 'NUMBER', 'Chiave dell’atto in SIPO (ID_ATTO): l’elemento da cui ripartire '
     'per salvare i dati ANSC.'),
    ('ID_TIPO_EVENTO', 'VARCHAR2(10)', 'idTipoEvento, decodifica ANSC_01: 1 nascita, 2 morte, '
     '3 matrimonio, 4 unione civile, 5 cittadinanza. Coincide con la prima cifra del codice '
     'dell’UC determinato.'),
    ('ID_TIPO_CONTENUTO', 'VARCHAR2(10)', 'idtipocontenuto, decodifica ANSC_02: atto, '
     'annotazione, trascrizione, annotazione automatica, rifiuto. È quello dell’UC determinato '
     '(ANSC_CFG_UC.ID_TIPO_CONTENUTO).'),
    ('COD_TIPO_OPERAZIONE', 'VARCHAR2(20)', 'CREAZIONE, RETTIFICA, ANNOTAZIONE, ANNULLAMENTO.'),
    ('ID_MODELLO_ATTO', 'VARCHAR2(5)', 'Modello di atto di SIPO (CONF_TIPO_ATTI.ID_MODELLO_ATTO).'),
    ('ID_UC_ANSC', 'VARCHAR2(20)', 'L’UC determinato (decodifica ANSC_03).'),
    ('ID_ANSC', 'VARCHAR2(50)', 'Identificativo nazionale dell’evento (idAnsc) restituito da R009.'),
    ('COD_LOGICA', 'VARCHAR2(30)', 'Codice della logica che ha determinato l’UC '
     '(ANSC_CFG_UC.LOGICA_DI_SCELTA): senza di esso la determinazione non è spiegabile a '
     'posteriori, ed è ciò che la vista di supervisione mostra accanto all’UC.'),
    ('COD_ORIGINE_UC', 'VARCHAR2(20)', 'AUTOMATICA oppure OPERATORE: come l’UC è stato '
     'determinato. La seconda è ammessa solo in caso di ambiguità ed è il segnale di una '
     'configurazione da correggere.'),
    ('COD_FASE', 'VARCHAR2(30)', 'Fase della lavorazione locale: IN_LAVORAZIONE, '
     'UC_DETERMINATO, DOCUMENTI_ACQUISITI, PREVALIDATO, VALIDATO, FIRMATO, ANNULLATO. Non si '
     'confonde con lo stato, che replica ciò che ANSC ha registrato: la fase dice a che punto è '
     'il nostro percorso.'),
    ('NUMERO_ATTO', 'VARCHAR2(6)', 'numeroatto: il numero comunale dell’atto, staccato dal '
     'servizio della numerazione all’Anteprima. Sei cifre, come la serie del numeratore; il '
     'contratto ne dichiara cinque (OP-73). Sull’atto di stato civile di SIPO non esiste una '
     'colonna dedicata: se aggiungerla lì resta da decidere (OP-60).'),
    ('FLG_EMERGENZA', 'CHAR(1)', 'S/N: distingue un atto in preparazione ordinaria da uno '
     'posto in modalità di emergenza.'),
    ('ID_ATTO_CARTACEO', 'VARCHAR2(256)', 'idAttoCartaceo: l’identificativo dell’archivio '
     'temporaneo in emergenza, trasmesso ad ANSC al rientro e conservato dopo.'),
    ('ID_MOTIVO_RECUPERO', 'VARCHAR2(3)', 'motivoRecupero, decodifica ANSC_96 (1…5 e 101): '
     'il motivo del recupero, trasmesso al deposito.'),
    ('COD_MOTIVO_RECUPERO', 'VARCHAR2(256)', 'descrizioneMotivoRecupero: il testo libero che '
     'accompagna il motivo.'),
    ('TS_DECORRENZA / ORA_DECORRENZA / MINUTO_DECORRENZA', 'TIMESTAMP / NUMBER / NUMBER',
     'dataDecorrenza, oraDecorrenza e minutoDecorrenza dell’atto (R014).'),
    ('ID_VERSIONE_ANSC', 'VARCHAR2(10)', 'idVersion: la versione del modello evento (ANSC_100) '
     'usata per l’evento.'),
    ('ID_STATO / STATO', 'NUMBER / VARCHAR2(20)', 'Replica dello stato dichiarato da ANSC '
     '(decodifica ANSC_11): il codice e la sua descrizione. Accoglie anche gli stati che il '
     'percorso ordinario non produce, come INEFFICACE (si veda il capitolo su SIPO e la web app '
     'ANSC), e lo stato 0 ad uso applicativo. Prima della validazione in ANSC non esiste alcun '
     'atto e il codice è vuoto.'),
    ('FLG_SECRETATO', 'CHAR(1)', '0 non secretato, 1 secretato: i valori di ANSC, che lo gestisce.'),
    ('TS_FORMAZIONE', 'TIMESTAMP', 'dataformazione con ora e minuto; anno, mese e giorno se ne '
     'ricavano.'),
    ('COD_RUOLO_OPERATORE', 'VARCHAR2(10)', 'ruoloOperatore, decodifica ANSC_63.'),
    ('FLG_CASA_COMUNALE', 'CHAR(1)', 'redattoNellaCasaComunale, S/N.'),
    ('TXT_EVENTO', 'CLOB', 'Il modello evento intero, così come è in redazione: si riscrive a '
     'ogni salvataggio. Formato l’atto coincide con ciò che è stato depositato e firmato e non '
     'cambia più; è la copia che sopravvive alla purga dell’audit. Il vincolo IS JSON è '
     'sospeso nel DDL.'),
    ('TS_VALIDAZIONE', 'TIMESTAMP', 'Quando è avvenuta l’ultima validazione riuscita.'),
    ('COD_PAGINA_RIPRESA', 'VARCHAR2(30)', 'L’ultima pagina raggiunta del percorso: DETTAGLIO, '
     'RIEPILOGO, ANTEPRIMA, VALIDAZIONE, FIRME_DICHIARANTI, CONFORMITA, FIRMA_USC.'),
    ('NUM_REVISIONE', 'NUMBER', 'Cresce a ogni salvataggio. Un salvataggio che porta una '
     'revisione superata è respinto: due postazioni non si sovrascrivono.'),
    ('COD_MUNICIPIO', 'VARCHAR2(10)', 'Municipio di competenza, per l’instradamento della '
     'worklist di supervisione.'),
    ('ULTIMO_ERRORE', 'VARCHAR2(4000)', 'Ultimo errore restituito da ANSC (code e testo).'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)', 'Campi '
     'tecnici di audit previsti dallo standard di nomenclatura [R5]: istante e utenza di '
     'creazione e di ultima modifica.'),
]

# (riga del foglio, nuovo «dove si registra», nuova nota) — None lascia com'è
TESTATA = {
    '2': ('— nel JSON', 'Assegnato da ANSC «per uso interno»: resta nel modello evento.'),
    '8': ('ANSC_STATO_ATTO.ID_STATO / STATO', 'Replica di ANSC_11, più lo stato 0 ad uso '
          'applicativo.'),
    '10': ('ANSC_STATO_ATTO.ID_VERSIONE_ANSC', 'Versione del modello evento (ANSC_100). Non è '
           'la baseline della nostra configurazione.'),
    '11': ('ANSC_STATO_ATTO.TS_FORMAZIONE', 'Una colonna TIMESTAMP con ora e minuto.'),
    '12': (None, 'Da TS_FORMAZIONE.'), '13': (None, 'Da TS_FORMAZIONE.'),
    '14': (None, 'Da TS_FORMAZIONE.'), '15': (None, 'Da TS_FORMAZIONE.'),
    '16': (None, 'Da TS_FORMAZIONE.'),
    '17': (None, 'ANSC usa i dati dell’ufficiale autenticato; chi ha lavorato l’atto è in '
           'UTENTE_INS e UTENTE_UPD.'),
    '23': ('— nel JSON', 'Il comune di registrazione resta nel modello evento.'),
    '24': ('— nel JSON', 'Con idComuneRegistrazione.'),
    '25': ('— nel JSON', 'Con idComuneRegistrazione.'),
    '26': ('— nel JSON', 'Con idComuneRegistrazione.'),
    '27': ('ANSC_STATO_ATTO.TS_DECORRENZA', 'Una colonna TIMESTAMP (R014).'),
    '28': ('ANSC_STATO_ATTO.ORA_DECORRENZA', ''),
    '29': ('ANSC_STATO_ATTO.MINUTO_DECORRENZA', ''),
    '30': ('ANSC_STATO_ATTO.ID_MOTIVO_RECUPERO', 'Decodifica ANSC_96: modalità di emergenza e '
           'rientro.'),
    '31': ('ANSC_STATO_ATTO.COD_MOTIVO_RECUPERO', 'Testo libero che accompagna il motivo.'),
    '32': ('ANSC_STATO_ATTO.ID_ATTO_CARTACEO', 'L’identificativo dell’archivio temporaneo, '
           'fino a 256 caratteri.'),
    '33': ('ANSC_STATO_ATTO.NUMERO_ATTO', 'Staccato dal servizio della numerazione. Il '
           'contratto dichiara cinque caratteri, il Comune ne assume sei (OP-73).'),
}


def stato_atto(doc):
    t = D.tabella_colonne(doc, 'ID_STATO_ATTO')
    D.riscrivi_sicura(t, STATO_ATTO)
    sost(doc, 'Indici: chiave primaria ID_STATO_ATTO; indice di supervisione su (COD_MUNICIPIO, '
              'ID_UFFICIALE, STATO) per la vista degli atti incompleti per ufficiale.',
         'Indici: chiave primaria ID_STATO_ATTO; indice su ID_ATTO_SIPO, da cui si ritrova la '
         'riga dell’atto.')
    sost(doc, 'Il payload conservato, letto insieme alla baseline registrata sulla stessa riga, è '
              'quindi l’unico modo per sapere a posteriori che cosa il Comune abbia dichiarato e '
              'su quali regole.',
         'Il payload conservato è quindi l’unico modo per sapere a posteriori che cosa il Comune '
         'abbia dichiarato.')

    t = D.trova_tabella(doc, 'riga', 'campo', 'dove si registra')
    for r in t.rows[1:]:
        k = r.cells[0].text.strip()
        if k in TESTATA:
            dove, nota = TESTATA[k]
            if dove is not None:
                D.riscrivi_cella(r.cells[3], dove)
            D.riscrivi_cella(r.cells[4], nota)
    sost(doc, 'Dei trentadue campi 16 hanno una colonna in questa tabella, 11 si ricavano da altre '
              'colonne, 2 restano nel modello evento e 3 sono deprecati da ANSC.',
         'Dei trentadue campi 17 hanno una colonna in questa tabella, 6 si ricavano da altre '
         'colonne, 6 restano nel modello evento e 3 sono deprecati da ANSC.')

    # la ripresa: senza contrassegno di sospensione
    sost(doc, 'La ripresa di una redazione sospesa', 'La ripresa di una redazione interrotta',
         attese=2)
    sost(doc, 'I valori che ANSC restituisce (id, idAnsc, stato, contrassegno di secretazione) si '
              'scrivono nelle colonne e nel modello nella stessa transazione.',
         'I valori che ANSC restituisce (idAnsc, stato, contrassegno di secretazione) si '
         'scrivono nelle colonne, e nel modello, nella stessa transazione.')
    p = par(doc, 'La sospensione. Esplicita')
    for r in p.runs:
        if r.text.startswith('La sospensione'):
            r.text = r.text.replace('La sospensione', 'L’interruzione')
    sost(doc, 'Esplicita, con un comando «Sospendi» nelle pagine del percorso, o implicita, quando '
              'la sessione di SIPO si chiude o l’OTP scade a metà lavoro. In entrambi i casi la '
              'riga resta com’è all’ultimo salvataggio e porta il contrassegno di sospensione.',
         'Quando la sessione di SIPO si chiude, l’OTP scade a metà lavoro o l’operatore lascia '
         'l’atto, la riga resta com’è all’ultimo salvataggio: non serve un contrassegno, perché '
         'una redazione è interrotta finché la sua fase non è conclusiva.')
    sost(doc, 'Dalla ricerca degli atti di SIPO, che mostra le redazioni sospese.',
         'Dalla ricerca degli atti di SIPO, che mostra le redazioni non concluse.')
    sost(doc, 'Una redazione sospesa non porta con sé l’autorità di chi l’ha sospesa.',
         'Una redazione interrotta non porta con sé l’autorità di chi l’ha lasciata.')
    sost(doc, 'Se un’altra postazione possa riprendere la redazione sospesa da un collega',
         'Se un’altra postazione possa riprendere la redazione lasciata da un collega')
    sost(doc, 'è ciò che consente di sospendere la redazione e di riprenderla dal punto in cui era',
         'è ciò che consente di interrompere la redazione e di riprenderla dal punto in cui era')

    # tabella degli oggetti in testa al capitolo
    sost(doc, 'ANNULLATO (R011)), idAnsc, ufficiale, municipio, audit.',
         'ANNULLATO (R011)), idAnsc, municipio, audit.')
    sost(doc, 'il modello intero e la pagina da cui riprendere una redazione sospesa.',
         'il modello intero e la pagina da cui riprendere una redazione interrotta.')
    sost(doc, 'in modo presidiato (per ufficiale/municipio)', 'in modo presidiato (per municipio)')


# ───────────────────────────────────────────────────────────── 3. configurazione
CFG_UC = [
    ('ID_UC_CFG', 'NUMBER (PK)', 'Chiave tecnica.'),
    ('COD_UC_ANSC', 'VARCHAR2(20)', 'Codice numerico dell’UC pubblicato da ANSC. La lunghezza '
     'non è fissa: 11111000 per la dichiarazione di nascita, 2101 per la morte in abitazione — '
     'nel catalogo compaiono codici di quattro, cinque, sei e otto cifre. La notazione puntata '
     '(2.1.0.1) usata in altri documenti è redazionale e non va scritta qui. Insieme a '
     'ID_VERSIONE è il riferimento con cui campi, formule, allegati, soggetti e interrogazioni '
     'richiamano l’UC, per valore.'),
    ('COD_MOTORE', 'VARCHAR2(40)', 'Codice parlante del motore (es. Dic_Nasc_001), dal catalogo '
     'di ANSC: attributo dell’UC, non sinonimo.'),
    ('DESC_UC_ANSC', 'VARCHAR2(400)', 'Denominazione dell’UC pubblicata da ANSC (decodifica '
     'ANSC_03): è la descrizioneCasoUso del modello evento.'),
    ('COD_FAMIGLIA', 'VARCHAR2(30)', 'Famiglia di evento: nascita, morte, matrimonio, …'),
    ('ID_TIPO_CONTENUTO', 'VARCHAR2(10)', 'Tipo di contenuto secondo la decodifica ANSC_02 '
     '(atto, annotazione, trascrizione…), pubblicato da ANSC per ciascun UC nella colonna '
     'IDTIPOCONTENUTO del catalogo dei casi d’uso: è la replica che la testata del modello '
     'evento richiede.'),
    ('ID_CONF_TIPO_ATTO', 'NUMBER', 'Riferimento alla configurazione del tipo atto in SIPO '
     '(CONF_TIPO_ATTI). Non entra nella chiave di unicità.'),
    ('ID_TIPO_ATTO', 'NUMBER', 'Tipo atto di SIPO a cui la riga si applica: è la colonna con cui '
     'il tipo atto si lega all’UC, e con ID_VERSIONE e COD_UC_ANSC forma la chiave di unicità. '
     'Vuota per un UC del catalogo che nessun tipo atto raggiunge ancora.'),
    ('STATO', 'VARCHAR2(1)', '[DA VERIFICARE: i valori ammessi. Un solo carattere; la sorgente '
     'dichiara la colonna ma non la valorizza. Da non confondere con lo stato dell’atto: qui si '
     'tratta della riga di configurazione.]'),
    ('NUM_PRIORITA', 'NUMBER', 'Ordine di valutazione fra gli UC dello stesso tipo atto: il '
     'primo la cui logica è soddisfatta determina l’UC. Valore predefinito 100.'),
    ('LOGICA_DI_SCELTA', 'VARCHAR2(30)', 'Codice della logica che stabilisce se l’UC si applica '
     'all’atto in lavorazione. È un codice libero: nessuna tabella lo cataloga e lo valuta il '
     'codice applicativo. Vuota significa «si applica sempre»: è il caso dei tipi atto che '
     'indirizzano un solo UC.'),
    ('COD_QUERY_UC', 'VARCHAR2(20) (FK)', 'L’interrogazione da utilizzare per l’UC, in '
     'ANSC_TEMPLATE: chiave esterna composta con ID_VERSIONE, così che la riga richiami soltanto '
     'un’interrogazione della propria baseline. È un concetto distinto dalla logica di scelta: '
     'quella decide se l’UC si applica all’atto, l’interrogazione è quella che si usa per l’UC. '
     'Vuota se l’UC non ne prevede una.'),
    ('ID_VERSIONE', 'NUMBER (FK)', 'Baseline di appartenenza (ANSC_CFG_VERSIONE); valore '
     'predefinito 1.'),
    ('DATA_INIZIO_VALIDITA / DATA_FINE_VALIDITA', 'DATE', 'Validità dell’UC come la dichiara '
     'ANSC. ANSC non cancella gli UC: ne chiude la validità — sono i tre casi ritirati, 374 '
     'validi su 377 righe del catalogo.'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)', 'Campi '
     'tecnici previsti dallo standard di nomenclatura [R5].'),
]

CFG_CAMPO = [
    ('ID_CAMPO', 'NUMBER (PK)', 'Chiave tecnica.'),
    ('ID_VERSIONE', 'NUMBER', 'Baseline di appartenenza (ANSC_CFG_VERSIONE). Con COD_UC_ANSC '
     'dice a quale UC e a quale baseline il campo appartiene. Il riferimento è per valore: la '
     'tabella non ha chiavi esterne né vincolo di unicità, per scelta.'),
    ('COD_UC_ANSC', 'VARCHAR2(20)', 'Codice dell’UC pubblicato da ANSC: richiama ANSC_CFG_UC per '
     'valore insieme alla baseline, come ALLEGATI_USECASE e SOGGETTI_USECASE.'),
    ('OGGETTO_ANSC', 'VARCHAR2(1000)', 'Oggetto del modello evento che contiene il campo '
     '(Binding Object del mapping ufficiale), es. evento.datiDiMorte.'),
    ('CAMPO_ANSC', 'VARCHAR2(1000)', 'Campo dentro l’oggetto (Binding Field del mapping), '
     'es. dataMorte.'),
    ('ALIAS_CAMPO_SIPO', 'VARCHAR2(200)', 'Il nome con cui il dato è noto in SIPO. È popolato a '
     'mano da chi configura.'),
    ('FLG_OBBLIGATORIO', 'CHAR(1)', 'S/N: se il campo sia obbligatorio, come dichiarato dal '
     'mapping. Esprime l’obbligatorietà del campo, non della sezione che lo contiene.'),
    ('ID_DECODIFICA_ANSC', 'NUMBER', 'Dizionario ANSC con cui validare il valore, quando il '
     'campo è codificato. Il raccordo al catalogo è per valore e senza chiave esterna, per non '
     'riunificare i domini di guasto della parte operativa e della parte dizionari.'),
    ('OBBLIGATORIETA_EVENTO', 'VARCHAR2(30)', '[DA VERIFICARE: valori ammessi e rapporto con '
     'FLG_OBBLIGATORIO, da stabilire nel prosieguo dell’analisi.]'),
    ('ORDINAMENTO', 'NUMBER', 'Ordine di presentazione e di elaborazione dei campi (1..N).'),
    ('OPERATIVO', 'CHAR(1)', 'Valori 0/1: se il campo deve essere inviato ad ANSC oppure no.'),
    ('DATA_INS / DATA_UPD / UTENTE_INS / UTENTE_UPD', 'TIMESTAMP / VARCHAR2(40)', 'Campi '
     'tecnici di audit previsti dallo standard di nomenclatura [R5].'),
]


def configurazione(doc):
    # il catalogo delle logiche
    for el in D.sezione(doc, 'Il catalogo delle logiche', livello=3):
        via(el)
    t = D.trova_tabella(doc, 'oggetto', 'ruolo', 'note principali')
    via_riga(t, 'TIPO_LOGICHE_DI_SCELTA e LOGICHE_DI_SCELTA')
    sost(doc, 'Una riga per UC e per Modello che lo raggiunge: codice, codice motore, descrizione, '
              'famiglia e validità come li pubblica ANSC, e accanto le scelte del Comune — il '
              'Modello, la priorità e la logica che decide, richiamata dal catalogo delle logiche '
              'per codice (ID_DOMINIO, COD_LOGICA).',
         'Una riga per UC e per tipo atto che lo raggiunge: codice, codice motore, descrizione, '
         'famiglia e validità come li pubblica ANSC, e accanto le scelte del Comune — il tipo '
         'atto, la priorità e il codice della logica che decide (LOGICA_DI_SCELTA).')

    # ANSC_CFG_UC: il commento sulla riga ID_TIPO_DOCUMENTO passa all'intestazione della tabella
    t = D.tabella_colonne(doc, 'ID_UC_CFG')
    titolo = D.h(doc, 3, 'ANSC_CFG_UC — il catalogo degli UC')
    riga = riga_tabella(t, 'ID_TIPO_DOCUMENTO')
    for cella in riga.cells:
        for p in cella.paragraphs:
            if D.ha_commenti(p._p):
                D.sposta_commenti(p._p, titolo._p)
    assert not D.ha_commenti(riga._tr), 'commento rimasto sulla riga ID_TIPO_DOCUMENTO'
    D.riscrivi_sicura(t, CFG_UC)
    sost(doc, 'Le colonne del catalogo — codice, codice motore, denominazione, famiglia, tipo '
              'evento, tipo contenuto, validità —',
         'Le colonne del catalogo — codice, codice motore, denominazione, famiglia, tipo '
         'contenuto, validità —')
    p = par(doc, 'La riga è per UC e per Modello che lo raggiunge')
    D.testo_di(p, 'La riga è per UC e per tipo atto che lo raggiunge, dentro una baseline: più '
               'tipi atto possono portare allo stesso UC, e ciascuno ha la propria priorità e la '
               'propria logica. Un UC che nessun tipo atto raggiunge ancora ha una riga sola, '
               'senza tipo atto: compare nel catalogo ma non partecipa alla determinazione. Ne '
               'segue che le colonne del catalogo si ripetono sulle righe dello stesso UC, e '
               'l’importazione le riscrive tutte insieme.')
    dopo(p, 'La logica di scelta è nominata da un codice (LOGICA_DI_SCELTA) e non ha una tabella '
         'che la descriva: la valuta il codice applicativo. Lo stesso vale per la condizione dei '
         'documenti condizionati (ALLEGATI_USECASE.cd_logica).')

    # ANSC_CFG_CAMPO
    D.riscrivi_sicura(D.tabella_colonne(doc, 'ID_CAMPO'), CFG_CAMPO)

    # ANSC_TEMPLATE
    p = par(doc, 'COD_QUERY_UC e COD_LOGICA_DI_SCELTA convivono')
    D.testo_di(p, p.text.replace('COD_QUERY_UC e COD_LOGICA_DI_SCELTA convivono',
                                 'COD_QUERY_UC e LOGICA_DI_SCELTA convivono')
               .replace('La logica di scelta, presa dal catalogo delle logiche, stabilisce',
                        'La logica di scelta stabilisce'))
    sost(doc, 'una per ciascun Modello che lo raggiunge', 'una per ciascun tipo atto che lo '
         'raggiunge', attese=2)
    t = D.tabella_colonne(doc, 'ID_COD_QUERY')
    D.riscrivi_cella(riga_tabella(t, 'DATA_INS').cells[0],
                     'TS_INS / TS_UPD / UTENTE_INS / UTENTE_UPD')

    # ALLEGATI_USECASE e SOGGETTI_USECASE
    sost(doc, 'e il riferimento al catalogo delle logiche al posto della condizione scritta in '
              'chiaro.',
         'e il codice della condizione (cd_logica) al posto della condizione implicita nel '
         'codice di presenza.')
    t = D.tabella_colonne(doc, 'cd_usecase')
    r = riga_tabella(t, 'id_dominio / cd_logica')
    D.riscrivi_cella(r.cells[0], 'cd_logica')
    D.riscrivi_cella(r.cells[1], 'VARCHAR2(30)')
    D.riscrivi_cella(r.cells[2], 'Il codice della condizione che rende obbligatorio un '
                     'documento quando ty_presenza vale C…. È un codice libero, senza tabella '
                     'che lo cataloghi: la condizione la valuta il codice applicativo. '
                     'Adattamento rispetto a Side, dove la condizione è implicita nel codice di '
                     'presenza.')
    sost(doc, 'qui la condizione si scrive nel catalogo delle logiche e si richiama da cd_logica, '
              'e il codice resta quello di Side.',
         'qui la condizione si nomina con cd_logica, e il codice resta quello di Side.')
    sost(doc, 'Stesso codice; la condizione è scritta nel catalogo delle logiche e richiamata da '
              'cd_logica.',
         'Stesso codice; la condizione è nominata da cd_logica e valutata dal codice '
         'applicativo.')
    p = par(doc, 'Le scelte di adozione sono quelle già fatte per ALLEGATI_USECASE')
    testo = p.text
    i = testo.index('Non si aggiunge invece il richiamo al catalogo')
    D.testo_di(p, testo[:i] + 'Non si aggiunge invece un codice di condizione come cd_logica: '
               'che una persona sia prevista dall’UC lo dice già ty_presenza.')
    sost(doc, 'La configurazione dei campi non ha un livello di sezione con una propria condizione '
              'di attivazione: la sezione del mapping vi resta come dato descrittivo '
              '(SEZIONE_FE_ANSC).',
         'La configurazione dei campi non registra la sezione del mapping.')

    # riconciliazione: il lato SIPO del campo non è più in ANSC_CFG_CAMPO
    sost(doc, ', con gli stessi nomi già usati da ANSC_CFG_CAMPO', '')
    sost(doc, 'La configurazione dei campi contiene già l’indicazione: la colonna SIPO di un campo '
              'codificato dice da quale tabella proviene il valore, ed è da lì che il raccordo si '
              'ricava senza cercarlo a mano.',
         'La configurazione dei campi non porta la tabella di SIPO da cui proviene il valore '
         '(ha il solo alias, compilato a mano): il raccordo va individuato da chi configura.')

    # testata del modello evento: il tipo evento non sta più sul catalogo
    sost(doc, 'Si conserva sul catalogo (ANSC_CFG_UC.ID_TIPO_EVENTO) invece di ricalcolarlo.',
         'Si ricava dal codice dell’UC determinato e si registra sull’atto '
         '(ANSC_STATO_ATTO.ID_TIPO_EVENTO).')


# ───────────────────────────────────────────────────────────── 4. flusso, versioni, errori
def flusso(doc):
    sost(doc, 'Se i dati cambiano, il modello evento non è più quello validato: la sua impronta '
              'differisce da TXT_HASH_VALIDATO, l’atto passa allo stato 0 (BOZZA_APPLICATIVA)',
         'Se i dati cambiano, il modello evento non è più quello validato: l’atto passa allo '
         'stato 0 (BOZZA_APPLICATIVA)')
    t = D.trova_tabella(doc, 'passo', 'che cosa legge', 'che cosa scrive', 'traccia')
    r = riga_tabella(t, '3 — Riepilogo')
    D.riscrivi_cella(r.cells[1], 'ANSC_CFG_UC per il tipo atto, nella baseline ATTIVA, in ordine '
                     'di NUM_PRIORITA; per ciascuna riga la logica nominata da LOGICA_DI_SCELTA.')
    D.riscrivi_cella(r.cells[2], 'ANSC_STATO_ATTO: ID_UC_ANSC, COD_LOGICA, COD_ORIGINE_UC e '
                     'fase UC_DETERMINATO.')
    sost(doc, 'DATI_ATTO_ESTERNO: il numero e la sua sorte; ANSC_STATO_ATTO.NUM_COMUNALE.',
         'DATI_ATTO_ESTERNO: il numero e la sua sorte; ANSC_STATO_ATTO.NUMERO_ATTO.')
    r = riga_tabella(t, '10 — «VALIDA»')
    D.riscrivi_cella(r.cells[2], 'ANSC_STATO_ATTO: ID_ANSC, stato 2 CONFERMATO (ID_STATO e '
                     'STATO), TS_VALIDAZIONE e fase VALIDATO. Da qui l’identificativo nazionale '
                     'è consumato.')

    # la determinazione
    sost(doc, 'Il front-end chiede il completamento passando tre dati: l’identificativo dell’atto, '
              'il tipo atto e il Modello.',
         'Il front-end chiede il completamento passando due dati: l’identificativo dell’atto e '
         'il tipo atto.')
    sost(doc, 'e le righe senza logica si applicano sempre: sono i Modelli che indirizzano un solo '
              'UC.',
         'e le righe senza logica si applicano sempre: sono i tipi atto che indirizzano un solo '
         'UC.')
    sost(doc, 'per quel Modello nessun caso d’uso è applicabile', 'per quel tipo atto nessun caso '
         'd’uso è applicabile')

    # che cosa resta sull'atto: la baseline non è più registrata
    sost(doc, 'Conclusa la formazione, quattro informazioni devono restare sull’atto',
         'Conclusa la formazione, tre informazioni devono restare sull’atto')
    t = D.trova_tabella(doc, 'dato', 'dove risiede', 'perché deve restare')
    via_riga(t, 'Versione della configurazione')
    sost(doc, 'ANSC_STATO_ATTO.NUM_COMUNALE. Sull’atto di stato civile di SIPO',
         'ANSC_STATO_ATTO.NUMERO_ATTO. Sull’atto di stato civile di SIPO')

    # versionamento: l'atto non porta la baseline
    p = par(doc, 'Lo store di stato registra, per ciascun atto, la baseline in vigore')
    D.testo_di(p, 'Lo store di stato non registra la baseline con cui l’atto è stato formato. '
               'Dell’atto restano il modello evento depositato (TXT_EVENTO) e la versione del '
               'modello evento usata (ID_VERSIONE_ANSC); la baseline in vigore a una data si '
               'ricava dalle date di attivazione delle versioni.')
    sost(doc, 'Lo stato 0 si imposta quando l’impronta del modello evento salvato in '
              'ANSC_STATO_ATTO differisce da quella dell’ultima validazione riuscita '
              '(TXT_HASH_VALIDATO), e si chiude con una nuova validazione riuscita',
         'Lo stato 0 si imposta quando l’operatore modifica i dati di un atto già validato, e si '
         'chiude con una nuova validazione riuscita')

    # errori
    sost(doc, 'ANSC_STATO_ATTO.STATO è il campo che riflette lo stato reale (decodifica ANSC_11) e '
              'TIPO_ESITO distingue lo scarto in validazione (SCARTATO: l’atto resta '
              'IN_PREPARAZIONE, perché in ANSC non è stato creato nulla) dall’esito indeterminato '
              '(timeout tecnico);',
         'ANSC_STATO_ATTO.ID_STATO e STATO riflettono lo stato reale (decodifica ANSC_11); lo '
         'scarto in validazione e l’esito indeterminato si leggono dall’esito dell’ultima '
         'chiamata registrata in ANSC_LOG_AUDIT;')

    # emergenza
    sost(doc, 'un contrassegno nello store, che distingua un atto in preparazione ordinaria da uno '
              'in emergenza e ne conservi il motivo e la data di ingresso;',
         'un contrassegno nello store, che distingua un atto in preparazione ordinaria da uno '
         'in emergenza e ne conservi il motivo (FLG_EMERGENZA, ID_MOTIVO_RECUPERO);')

    # numerazione
    sost(doc, 'Il numero dell’atto depositato resta anche sullo store di stato '
              '(ANSC_STATO_ATTO.NUM_COMUNALE)',
         'Il numero dell’atto depositato resta anche sullo store di stato '
         '(ANSC_STATO_ATTO.NUMERO_ATTO)')
    sost(doc, 'pg_atto è NUMBER(6)', 'num_atto è NUMBER(6)')
    sost(doc, 'il prossimo numero è pg_atto + 1', 'il prossimo numero è num_atto + 1')
    sost(doc, 'e ANSC_STATO_ATTO.NUM_COMUNALE per l’atto', 'e ANSC_STATO_ATTO.NUMERO_ATTO per '
         'l’atto')


# ───────────────────────────────────────────────────────────── 5. back-office e open point
def back_office(doc):
    sost(doc, 'Worklist per ufficiale, ricerca e tracciabilità', 'Worklist per municipio, '
         'ricerca e tracciabilità')
    sost(doc, 'casi d’uso, campi, allegati, formule, logiche, riconciliazione dei valori',
         'casi d’uso, campi, allegati, formule, riconciliazione dei valori')
    sost(doc, 'supervisione atti per ufficiale (stato reale dell’atto in ANSC)',
         'supervisione atti per municipio (stato reale dell’atto in ANSC)')
    sost(doc, 'allegati e formule con versionamento, catalogo delle logiche, riconciliazione',
         'allegati e formule con versionamento, riconciliazione')
    sost(doc, 'filtrabile per municipio, ufficiale, categoria di eccezione',
         'filtrabile per municipio, categoria di eccezione')
    sost(doc, 'la descrizione del problema, l’ufficiale competente, l’ultimo errore',
         'la descrizione del problema, l’operatore che lo ha lavorato per ultimo, l’ultimo '
         'errore')
    sost(doc, 'Nessuna condizione di applicabilità ha risposto per il Modello dell’atto',
         'Nessuna condizione di applicabilità ha risposto per il tipo atto')
    sost(doc, 'Più UC dello stesso Modello hanno risposto', 'Più UC dello stesso tipo atto hanno '
         'risposto')
    sost(doc, 'e la logica che l’ha prodotto, e la baseline con cui il payload è stato costruito —',
         'e la logica che l’ha prodotto —')
    sost(doc, 'Scartato in validazione (IN_PREPARAZIONE, esito SCARTATO)',
         'Scartato in validazione (IN_PREPARAZIONE, ultimo deposito respinto)')
    sost(doc, 'Ogni atto in eccezione è instradato all’ufficiale competente per municipio: '
              'l’indice di worklist su (municipio, ufficiale, stato) serve proprio questo.',
         'Ogni atto in eccezione è instradato per municipio (COD_MUNICIPIO).')
    sost(doc, '(confermato, firmato dal dichiarante, firmato da USC, rifiutato, annullato), per '
              'ufficiale e municipio.',
         '(confermato, firmato dal dichiarante, firmato da USC, rifiutato, annullato), per '
         'municipio.')
    t = D.trova_tabella(doc, 'schermata', 'scopo', 'azioni principali', 'ruolo')
    r = riga_tabella(t, 'Configurazione — Casi d’uso')
    D.riscrivi_cella(r.cells[1], 'ANSC_CFG_UC: una riga per UC adottato, con il tipo atto a cui '
                     'si applica, la priorità fra gli UC dello stesso tipo atto e il codice della '
                     'logica che lo seleziona.')
    via_riga(t, 'Configurazione — Catalogo delle logiche')
    r = riga_tabella(t, 'Configurazione — Campi, allegati e formule per UC')
    D.riscrivi_cella(r.cells[1], 'ANSC_CFG_CAMPO, ALLEGATI_USECASE e ANSC_CFG_FORMULA per '
                     'singolo UC: come ogni dato è noto in SIPO e se va inviato, quali documenti '
                     'servono e quali diciture sono previste.')
    D.riscrivi_cella(r.cells[2], 'Precompila dal foglio (campi, allegati e formule insieme), '
                     'esamina e conferma riga per riga, compila l’alias SIPO, apre la '
                     'riconciliazione dei valori dove il campo è codificato.')
    sost(doc, 'Wireframe — Configurazione (Casi d’uso): una riga per UC adottato, con Modello, '
              'priorità, logica di scelta richiamata dal catalogo e origine della riga.',
         'Wireframe — Configurazione (Casi d’uso): una riga per UC adottato, con tipo atto, '
         'priorità e codice della logica di scelta.')
    sost(doc, 'tre schede per un solo UC, con la logica del campo richiamata dal catalogo e il '
              'rimando alla riconciliazione dei valori.',
         'tre schede per un solo UC, con l’alias SIPO di ciascun campo e il rimando alla '
         'riconciliazione dei valori.')
    sost(doc, 'il codice dell’UC con cui ANSC_CFG_FORMULA, ALLEGATI_USECASE e SOGGETTI_USECASE '
              'richiamano ANSC_CFG_UC, la decodifica su ANSC_CFG_CAMPO e la baseline sullo store '
              'di stato.',
         'il codice dell’UC con cui ANSC_CFG_CAMPO, ANSC_CFG_FORMULA, ALLEGATI_USECASE e '
         'SOGGETTI_USECASE richiamano ANSC_CFG_UC, la baseline e la decodifica su '
         'ANSC_CFG_CAMPO.')

    sost(doc, 'È il raccordo con OGGETTO_ANSC di ANSC_CFG_SEZIONE.',
         'È il raccordo con OGGETTO_ANSC di ANSC_CFG_CAMPO.')
    sost(doc, 'Finché lo schema non aveva altro, era un refuso senza conseguenze; da quando il '
              'Comune ha una tabella TIPO_LOGICHE_DI_SCELTA, quel nome allude a un legame che non '
              'c’è e manda fuori strada chi legge il DDL prima del modello.',
         'Il nome allude a un legame che non c’è e manda fuori strada chi legge il DDL prima '
         'del modello.')
    sost(doc, 'Mancano inoltre le operazioni per la configurazione delle sezioni, degli allegati, '
              'delle formule e del catalogo delle logiche, che il back-office prevede.',
         'Mancano inoltre le operazioni per la configurazione degli allegati e delle formule, '
         'che il back-office prevede.')
    sost(doc, '(c) Se una redazione sospesa possa essere ripresa da un altro operatore',
         '(c) Se una redazione interrotta possa essere ripresa da un altro operatore')
    t = next(t for t in doc.tables if t.rows[1].cells[0].text.startswith('OP-'))
    c = riga_tabella(t, 'OP-56').cells[2]
    D.riscrivi_cella(c, c.text.rstrip() + ' Le due tabelle sono state poi deprecate: la logica è '
                     'un codice libero, senza catalogo.')


FIGURE = [
    ('Schema ANSC_USR, con ALLEGATO', 'erd_ansc_usr.png'),
    ('Il flusso operativo nelle sue quattro fasi', 'flusso_operativo.png'),
    ('Dalla maschera SIPO al payload ANSC', 'catena_configurazione.png'),
    ('Wireframe — Configurazione (Casi d’uso)', 'bo_config_uc.png'),
    ('Wireframe — Configurazione (Campi, allegati e formule per UC)', 'bo_config_campi.png'),
]


def main():
    if os.path.exists(OUT):
        raise SystemExit('esiste già: ' + OUT)
    doc = docx.Document(SRC)
    ddl(doc)
    stato_atto(doc)
    configurazione(doc)
    flusso(doc)
    back_office(doc)
    for didascalia, png in FIGURE:
        D.sostituisci_immagine(doc, didascalia, os.path.join(IMG, png))
    D.storia(doc, '27/10/2026', '3.53',
             'Adeguamenti database · Il flusso operativo · Costruzione del payload · Gestione dei '
             'dizionari ANSC · Il versionamento della configurazione · Il processo di gestione '
             'dei dati ANSC in SIPO · La numerazione comunale · Il back-office · Open Point · '
             'Appendice A',
             'Allineato il documento alle DDL riviste con gli sviluppatori. Tolte le tabelle '
             'deprecate TIPO_LOGICHE_DI_SCELTA e LOGICHE_DI_SCELTA: la logica di scelta e la '
             'condizione dei documenti sono codici liberi. ANSC_CFG_UC si lega al tipo atto '
             '(ID_TIPO_ATTO) e non più al Modello; ANSC_CFG_CAMPO porta l’alias SIPO, compilato '
             'a mano, al posto di schema, tabella e colonna, e riprende ID_DECODIFICA_ANSC. '
             'ANSC_STATO_ATTO riallineata colonna per colonna: tolti dal testo impronta del '
             'modello validato, sospensione, esito qualificato, baseline e ufficiale sull’atto. '
             'Corretti gli errori di sintassi delle DDL; rigenerati ERD, flusso operativo, '
             'catena della configurazione e due wireframe.')
    D.frontespizio(doc, '3.53', 'ANALISI_Integrazione-ANSC_v3.53')
    doc.save(OUT)
    print('scritto:', OUT)


if __name__ == '__main__':
    main()
