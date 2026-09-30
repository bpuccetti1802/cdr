# -*- coding: utf-8 -*-
"""v3.14 — via ID_MAPPER, dentro TXT_TRANSCODIFICA: la trasformazione è configurazione.

⚠️ `ID_MAPPER` era un residuo della v2.x, quando si immaginava un adattatore di codice per
famiglia di atti. Il capitolo 9 stabilisce l'opposto — «il payload non si costruisce per caso
d'uso: si costruisce una volta sola sull'albero del modello e si riempie secondo la
configurazione» — e con esso una colonna che sceglie *quale* adattatore usare non ha oggetto.
Il documento diceva le due cose in capitoli diversi, e nel punto 6 della modalità operativa
arrivava a promettere «un mapper per tipo di atto», cioè ciò che il capitolo 9 dichiara di
voler evitare.

⚠️ Ma togliere la colonna senza guardare che cosa faceva sarebbe un errore: il mapper era la
sede delle TRASFORMAZIONI. Nel foglio di lavoro compilato a mano di Dic_Nasc_001, **53 righe
su 181** portano una regola di transcodifica — «estrarre YYYY-MM-DD dalla stringa», «usare
ID_DICHIARANTE come chiave in CONF_DICHIARANTE_NASCITA», «estrarre dal CLOB» — e
`ANSC_CFG_CAMPO` non aveva alcuna colonna dove metterle. `ID_MAPPER` non era tanto una
colonna inutile quanto il sintomo di una colonna mancante.

⚠️ Scelta dichiarata (decisa con l'utente): `TXT_TRANSCODIFICA` è **una specifica in lingua
corrente per chi sviluppa, non un'espressione eseguibile**. Non promette un automatismo che
non c'è: dove la provenienza è una ricerca su un altro schema — `dichiarante.idANPR` si
ricava cercando il soggetto in `ANAG_USR` — nessuna stringa la esegue, e dirlo è più onesto
che fingere una grammatica.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402
from v3_14_configurazione import riscrivi_tabella   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.14.docx')

TRANSCODIFICA = ('TXT_TRANSCODIFICA', 'VARCHAR2(600)',
                 'Come si ricava il valore dalla sorgente, quando non è un trasferimento '
                 'diretto: estrazione da una data o da un nodo XML, ricerca di una descrizione '
                 'in una tabella di decodifica, valore fisso. ⚠️ È una specifica in lingua '
                 'corrente per chi sviluppa, non un’espressione che il concentratore valuta: '
                 'dove la provenienza è una ricerca su un altro schema nessuna stringa la '
                 'esegue, e la configurazione lo dichiara invece di prometterlo.')

SPIEGAZIONE = [
    ('p', 'Una precisazione su che cosa la configurazione dei campi contenga e che cosa no. '
          'La colonna della sorgente dice da dove si prende il dato; quella della '
          'transcodifica dice come lo si ricava, perché in un numero non trascurabile di casi '
          'il trasferimento non è diretto. Sul caso d’uso di nascita usato come esempio, '
          'cinquantatré campi su centottantuno richiedono una trasformazione: estrarre la data '
          'o l’ora da una colonna che le contiene entrambe, leggere un nodo dentro un XML, '
          'risalire dalla chiave alla descrizione in una tabella di decodifica.'),
    ('p', '⚠️ Quella colonna è una specifica per chi sviluppa, non un’espressione eseguibile, e '
          'la distinzione va tenuta ferma. Alcune provenienze non sono trasformazioni ma '
          'ricerche su un altro schema — l’identificativo ANPR del dichiarante si ottiene '
          'cercando il soggetto in ANAG_USR — e nessuna stringa di configurazione le esegue. '
          'Dichiararlo è preferibile a costruire una grammatica che prometta un automatismo '
          'inesistente: il concentratore realizza quelle provenienze nel proprio codice, e la '
          'configurazione dice quali sono e che cosa devono produrre.'),
]

FRASI = [
    ('descrive, per ogni Modello di atto (tipo evento × tipo operazione × Modello), '
     'l’interfaccia ANSC da usare e il mapper, e — per ogni UC — i campi obbligatori',
     'descrive, per ogni caso d’uso adottato, l’interfaccia ANSC da usare e le condizioni che '
     'lo selezionano, e — sempre per UC — i campi obbligatori',
     'la configurazione come livello di metadati'),

    ('al deposito (R009) la configurazione della versione registrata seleziona il mapper e '
     'materializza il payload da SIPO',
     'al deposito (R009) l’adattatore costruisce il payload sull’albero del modello evento e lo '
     'riempie con la configurazione della versione registrata, applicando le transcodifiche '
     'dichiarate campo per campo',
     'deposito e firma'),

    ('6. Estensibilità: aggiungere un tipo di atto significa inserire righe di configurazione e '
     'un mapper; il concentratore e il pre-filtro restano invariati. È il fulcro della '
     'riduzione di effort.',
     '6. Estensibilità: aggiungere un tipo di atto significa configurare un caso d’uso — la sua '
     'condizione, i suoi campi con le rispettive sorgenti, gli allegati e le formule — senza '
     'scrivere codice: l’adattatore è uno solo e vale per tutto il dominio, perché il modello '
     'evento è un albero solo. Il concentratore e il pre-filtro restano invariati, ed è il '
     'fulcro della riduzione di effort.',
     'estensibilità senza un mapper per famiglia'),

    ('ANSC_CFG_UC e ANSC_CFG_CAMPO: i metadati che, per evento × operazione × Modello, guidano '
     'mapper e obbligatorietà.',
     'ANSC_CFG_UC e ANSC_CFG_CAMPO: i metadati che, per caso d’uso, guidano la selezione '
     'dell’UC, la provenienza dei dati e l’obbligatorietà.',
     'tabella degli oggetti dello schema'),

    ('La conseguenza progettuale è che il concentratore deve selezionare il mapper in base al '
     'tipo operazione risolto in configurazione: è il presupposto del requisito seguente.',
     'La conseguenza progettuale è che il concentratore deve trattare la trascrizione in base '
     'al tipo operazione risolto in configurazione — valorizzando i dati dell’atto di '
     'provenienza — e non dedurla dal codice dell’UC: è il presupposto del requisito seguente.',
     'trattamento della trascrizione'),

    ('il tipo atto, la maschera, l’interfaccia ANSC, il mapper, la priorità fra gli UC dello '
     'stesso Modello',
     'il tipo atto, la maschera, l’interfaccia ANSC, la priorità fra gli UC dello stesso Modello',
     'schermata del back-office'),
]


def applica(doc):
    fatti = []

    # ------------------------------------------- ANSC_CFG_UC: via ID_MAPPER
    t = D.tabella_colonne(doc, 'ID_UC_CFG')
    righe = [[c.text for c in r.cells] for r in t.rows[1:]]
    restano = [v for v in righe if not v[0].startswith('ID_MAPPER')]
    assert len(restano) == len(righe) - 1, 'ID_MAPPER non trovata in ANSC_CFG_UC'
    riscrivi_tabella(t, restano)
    fatti.append('ANSC_CFG_UC: rimossa ID_MAPPER')

    # ------------------------------------ ANSC_CFG_CAMPO: dentro TXT_TRANSCODIFICA
    tc = D.tabella_colonne(doc, 'ID_CAMPO')
    righe = [[c.text for c in r.cells] for r in tc.rows[1:]]
    if not any(v[0].startswith('TXT_TRANSCODIFICA') for v in righe):
        fuori = []
        for v in righe:
            fuori.append(v)
            if v[0].startswith('CAMPO_SIPO'):
                fuori.append(list(TRANSCODIFICA))
        riscrivi_tabella(tc, fuori)
        fatti.append('ANSC_CFG_CAMPO: aggiunta TXT_TRANSCODIFICA dopo CAMPO_SIPO')

    # ------------------------------------------------------------- il DDL
    D.sostituisci(doc, '  ID_MAPPER            VARCHAR2(40 CHAR)  NOT NULL,\n', '',
                  attese=None, etichetta='DDL: riga ID_MAPPER', fatti=fatti)
    riga = next((p for p in doc.paragraphs
                 if p.text.strip().startswith('ID_MAPPER') and 'VARCHAR2(40 CHAR)' in p.text),
                None)
    if riga is not None:
        riga._p.getparent().remove(riga._p)
        fatti.append('DDL di ANSC_CFG_UC: rimossa la riga ID_MAPPER')
    campo_sipo = next((p for p in doc.paragraphs
                       if p.text.strip().startswith('CAMPO_SIPO')
                       and 'VARCHAR2(120 CHAR)' in p.text), None)
    if campo_sipo is not None:
        D.riga_dopo(doc, campo_sipo, '  TXT_TRANSCODIFICA    VARCHAR2(600 CHAR),')
        fatti.append('DDL di ANSC_CFG_CAMPO: aggiunta TXT_TRANSCODIFICA')

    # ------------------------------------------------- INSERT ed esempio
    D.sostituisci(doc, 'SERVIZIO_ANSC, ID_MAPPER, NUM_PRIORITA,', 'SERVIZIO_ANSC, NUM_PRIORITA,',
                  attese=1, etichetta='INSERT del pilota (colonne)', fatti=fatti)
    D.sostituisci(doc, "   'R009','MorteDichiarazione',10,'N','S','S',:idVersione);",
                  "   'R009',10,'N','S','S',:idVersione);",
                  attese=1, etichetta='INSERT del pilota (valori)', fatti=fatti)
    D.sostituisci(doc, 'SERVIZIO_ANSC / ID_MAPPER', 'SERVIZIO_ANSC',
                  attese=1, etichetta='esempio completo (riga della tabella)', fatti=fatti)
    D.sostituisci(doc, 'R009 / NascitaDichiarazione', 'R009',
                  attese=1, etichetta='esempio completo (valore)', fatti=fatti)

    # ------------------------------------------------------- le frasi
    for vecchio, nuovo, etichetta in FRASI:
        D.sostituisci(doc, vecchio, nuovo, attese=1, etichetta=etichetta, fatti=fatti)

    # ------------------------------------ la spiegazione, dopo la tabella dei campi
    ancora = next(p for p in doc.paragraphs
                  if p.text.strip().startswith('Le tabelle della configurazione —'))
    for _, testo in SPIEGAZIONE:
        D.para(doc, ancora._p, testo)
    fatti.append('spiegazione della transcodifica (2 paragrafi)')
    return fatti


if __name__ == '__main__':
    percorso = sys.argv[1] if len(sys.argv) > 1 else DOC
    doc = docx.Document(percorso)
    for f in applica(doc):
        print('  ·', f)
    doc.save(percorso)
    print('salvato:', os.path.basename(percorso))
