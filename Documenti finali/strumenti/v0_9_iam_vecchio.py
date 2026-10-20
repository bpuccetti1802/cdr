# -*- coding: utf-8 -*-
"""ANALISI_Identita-Profilazione-IAM v0.8 → v0.9 (15/10/2026).

Due richieste dell'utente:
  1. togliere i «⚠️» anche dalle tabelle (nel corpo li aveva già tolti lui nella v0.7);
  2. dopo i rischi del capitolo dell'integrazione, dire in chiaro che implementare la libreria
     della piattaforma nel vecchio ambiente (AS-IS) e sostituirvi i meccanismi attuali può
     creare serie difficoltà e richiedere manutenzione evolutiva: **il suggerimento è lasciare il
     vecchio ambiente com'è** e adottare la piattaforma per il mondo nuovo.

Per coerenza si riallineano le due frasi che dicevano il contrario: la risposta (c)
dell'Executive summary e la proposta «A come transizione e B come approdo».
⚠️ Il capitolo «Scelte effettuate per il TO BE» descrive ancora il passaggio dei 35 front-end
alla piattaforma: NON si tocca senza un'indicazione dell'utente (glielo si segnala).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v0_9_iam_vecchio.py
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.8.docx')
DST = os.path.join(FIN, 'ANALISI_Identita-Profilazione-IAM_v0.9.docx')

AVVISO = '⚠️'

DUE_MONDI = [
    ['Che cosa', 'Vecchio ambiente: com’è oggi', 'Mondo nuovo: la piattaforma'],
    ['Accesso', 'Intestazioni del portale, lette dalla libreria di profilazione.',
     'roma-capitale-auth con IAM.'],
    ['Profili', 'ANAG_USR: ruoli, aree tematiche, funzionalità, R_UTENTI_RUOLI.',
     'roma-capitale-management: profili e abilitazioni per applicazione.'],
    ['Gestione degli accessi', 'La schermata di gestione degli utenti di AnagrafeWeb.',
     'L’interfaccia di amministrazione della piattaforma.'],
    ['Verso i back-end', 'Le utenze tecniche per profilo di back-end.',
     'Il token della piattaforma.'],
    ['Applicazioni', 'I trentacinque front-end Thymeleaf e i loro back-end.',
     'Le applicazioni Angular, a cominciare dal back-office ANSC.'],
]

RISPOSTA_C_VECCHIA = ('Sì, in due tempi: prima SIPO resta la fonte e la piattaforma legge, poi la '
                      'piattaforma diventa la fonte. Il rischio principale è la forma del permesso '
                      'di SIPO, che è per persona e va ricondotta a profili.')
RISPOSTA_C_NUOVA = ('Tecnicamente sì, ma nel vecchio ambiente non conviene: sostituire i meccanismi '
                    'attuali con quelli della piattaforma porterebbe serie difficoltà e manutenzione '
                    'evolutiva su entrambi i lati. Il suggerimento è lasciare il vecchio ambiente '
                    'com’è e adottare la piattaforma per il mondo nuovo.')


def togli_avvisi(p):
    """Toglie «⚠️» dentro i run, così grassetti e corsivi restano."""
    if AVVISO not in p.text:
        return 0
    n = p.text.count(AVVISO)
    for r in p.runs:
        if AVVISO in r.text:
            r.text = r.text.replace(AVVISO + ' ', '').replace(AVVISO, '')
    if AVVISO in p.text:                       # non dovrebbe accadere: spezzato fra run
        D.testo_di(p, p.text.replace(AVVISO + ' ', '').replace(AVVISO, ''))
    # lo spazio prima dell'avviso, se l'avviso stava in un run proprio, resta doppio
    for r in p.runs:
        while '  ' in r.text:
            r.text = r.text.replace('  ', ' ')
    if p.runs and p.runs[0].text.startswith(' '):
        p.runs[0].text = p.runs[0].text.lstrip()
    return n


def main():
    if os.path.exists(DST):
        os.remove(DST)                         # solo un'esecuzione precedente di questo script
    shutil.copy(SRC, DST)
    d = docx.Document(DST)
    modello = d.tables[3]
    fatti = []

    # ------------------------------------------------ 1. niente «⚠️» nelle tabelle (e nel corpo)
    n = sum(togli_avvisi(p) for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in dict.fromkeys(r.cells):
                for p in c.paragraphs:
                    n += togli_avvisi(p)
    assert AVVISO not in ''.join(p.text for p in d.paragraphs)
    fatti.append(f'tolti {n} «⚠️»')

    # ------------------------------------------------ 2. risposta (c) dell'Executive summary
    k = 0
    for p in d.paragraphs:
        for r in p.runs:
            if RISPOSTA_C_VECCHIA in r.text:
                r.text = r.text.replace(RISPOSTA_C_VECCHIA, RISPOSTA_C_NUOVA)
                k += 1
    if k != 1:
        raise SystemExit(f'risposta (c): attesa 1 occorrenza, trovate {k}')
    fatti.append('Executive summary: risposta (c) riallineata')

    # ------------------------------------------------ 3. la proposta fra le modalità
    vecchia = [p for p in d.paragraphs if p.text.startswith('La proposta è A come transizione')]
    assert len(vecchia) == 1
    D.para(d, vecchia[0]._p, pezzi=D.segmenta('', (
        'Delle tre, la C è esclusa. Fra A e B la scelta va letta insieme alla raccomandazione che '
        'chiude il capitolo: **nel vecchio ambiente nessuna delle due deve sostituire i '
        'meccanismi esistenti**, e B vale per le applicazioni nuove. Resta per tutte una regola: '
        '**per ciascuna applicazione un solo posto in cui si scrive**.')))
    assert not D.ha_commenti(vecchia[0]._p)
    vecchia[0]._p.getparent().remove(vecchia[0]._p)
    fatti.append('proposta fra le modalità riallineata')

    # ------------------------------------------------ 4. dopo i rischi: il vecchio ambiente
    fine = d.element.body[-1]
    assert fine.tag == qn('w:sectPr')
    for tipo, contenuto in [
        ('h2', 'Il vecchio ambiente: lasciarlo com’è'),
        ('b', 'I rischi della tabella precedente vanno letti per quello che dicono sul vecchio '
              'ambiente. **Implementare la libreria della piattaforma nell’AS-IS e sostituire i '
              'meccanismi attuali con quelli nuovi potrebbe creare serie difficoltà, e '
              'richiederebbe interventi di manutenzione evolutiva** su entrambi i lati.'),
        ('v', ('Sul lato SIPO. ',
               'Le librerie di profilazione e di accesso ai servizi, i trentacinque front-end da '
               'ricollaudare, i quarantaquattro back-end che dovrebbero verificare un token '
               'diverso, i 4.829 punti in cui conta il primo ruolo, e la traduzione della '
               'quaterna in profili, con le eccezioni che la accompagnano.')),
        ('v', ('Sul lato della piattaforma. ',
               'I miglioramenti senza i quali SIPO non può entrare — nodo organizzativo, profilo '
               'principale, storia delle assegnazioni, configurabilità delle applicazioni, revoca '
               'che abbia effetto (MI-01, MI-05, MI-06, MI-09, MI-10) — e le manutenzioni di '
               'priorità alta, da chiedere a un titolare che non è SIPO (PI-33).')),
        ('b', 'Il beneficio, per un ambiente che funziona e che non è destinato a crescere, non '
              'compensa questo costo. **Il suggerimento è quindi di lasciare, per il vecchio '
              'ambiente, tutto com’è**: accesso, profilazione e gestione degli accessi restano '
              'quelli di oggi. La piattaforma si adotta per il mondo nuovo, dove nasce senza un '
              'esistente da sostituire.'),
        ('tab', (DUE_MONDI, [1.4, 2.6, 2.6])),
        ('p', 'Ne discendono due conseguenze, da mettere in conto. Chi lavora in entrambi i mondi '
              'va abilitato in due posti, ciascuno per le proprie applicazioni: è il prezzo della '
              'scelta, e se i volumi lo giustificano si contiene con un’alimentazione in sola '
              'lettura dall’anagrafe degli utenti di SIPO verso la piattaforma, cioè con la '
              'modalità A, senza toccare i meccanismi esistenti. E la regola di un solo posto in '
              'cui si scrive resta rispettata, perché vale per applicazione: le applicazioni '
              'vecchie si amministrano in AnagrafeWeb, le nuove nella piattaforma.'),
        ('p', 'Resta fuori da questa raccomandazione, perché non riguarda la libreria, la clausola '
              'di chiusura dei servizi di back-end che oggi rispondono senza token (RI-01): è un '
              'intervento piccolo sul vecchio ambiente, che non ne cambia i meccanismi e che '
              'nessuna delle scelte di questo capitolo rende superfluo.'),
    ]:
        if tipo == 'p':
            D.para(d, fine, contenuto)
        elif tipo == 'b':
            D.para(d, fine, pezzi=D.segmenta('', contenuto))
        elif tipo == 'v':
            D.voce(d, fine, *contenuto)
        elif tipo == 'h2':
            D.para(d, fine, contenuto, stile='Heading 2')
        elif tipo == 'tab':
            righe, larghezze = contenuto
            D.tabella(d, fine, righe, modello, larghezze)
    fatti.append('nuovo sottocapitolo «Il vecchio ambiente: lasciarlo com’è»')

    D.storia(d, '15/10/2026', '0.9',
             'Executive summary · L’integrazione con la gestione AS-IS di SIPO · tutte le tabelle',
             'Aggiunto, dopo i rischi, il sottocapitolo sul vecchio ambiente: implementare la '
             'libreria della piattaforma nell’AS-IS e sostituirne i meccanismi potrebbe creare '
             'serie difficoltà e richiederebbe manutenzione evolutiva su entrambi i lati; il '
             'suggerimento è lasciare il vecchio ambiente com’è e adottare la piattaforma per il '
             'mondo nuovo. Riallineate di conseguenza la risposta sull’integrazione '
             'nell’Executive summary e la proposta fra le modalità. Tolti i simboli di avviso '
             'dalle tabelle.')
    for f in D.frontespizio(d, '0.9', 'ANALISI_Identita-Profilazione-IAM_v0.9'):
        fatti.append(f)

    d.save(DST)
    print('\n'.join('  · ' + f for f in fatti))
    cap, tab, fig = D.riepilogo(DST)
    print(f'  capitoli {cap}, tabelle {tab}, immagini {fig}')
    print('  scritto', DST)


if __name__ == '__main__':
    main()
