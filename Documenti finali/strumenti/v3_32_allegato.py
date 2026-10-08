# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.31 → v3.32 (07/10/2026).

Il committente fornisce il DDL aggiornato di ALLEGATO. La ragione dichiarata è il limite
di Oracle sulla lunghezza degli identificativi, ma il blocco porta anche uno spostamento
di schema, confermato: **la tabella passa da ANSC_USR a MATR_USR**.

⚠️ Il limite effettivo è **30 byte** (Oracle fino alla 12.1; 128 dalla 12.2), non 32: lo
dimostra il DDL stesso, che accorcia un nome di 31 caratteri.

⚠️ Il difetto non era isolato. Scansione di tutti i 75 identificativi del DDL: **cinque**
superano i 30 caratteri e il DDL fornito ne corregge uno. Gli altri quattro si correggono
qui, perché lasciarli significa scoprirli al primo CREATE:
    CK_ANSC_STATO_ATTO_FLG_EMERGENZA     32 → CK_STATO_ATTO_FLG_EMERGENZA
    allegati_usecase_fg_att_conf_check   34 → allegati_usecase_att_conf_ck
    allegati_usecase_fg_operante_check   34 → allegati_usecase_operante_ck
    valore_dominio_dominio_decodifica_fk 36 → valore_dominio_dominio_fk

⚠️ Lo spostamento di schema NON è una conseguenza del limite e tocca altri punti:
l'ERD (intitolato «Schema ANSC_USR» e contenente ALLEGATO), il capitolo degli adeguamenti
database, la premessa alle tabelle operative, e il regime dei permessi di scrittura — che
diventa un punto aperto nuovo.

⚠️ L'ERD portava ancora `fg_extra`, nome superato dalla v3.29 in favore di
`fg_testo_libero`: si corregge rigenerandolo.

⚠️ 22 commenti di Word da preservare: si contano prima e dopo.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_32_allegato.py
"""
import os
import shutil
import subprocess
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.31.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.32.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')

DDL = """CREATE TABLE MATR_USR.ALLEGATO (
  id_allegato          NUMBER GENERATED ALWAYS AS IDENTITY,
  id_atto_sipo         NUMBER NOT NULL,
  id_ansc_allegato     VARCHAR2(15 CHAR),
  ty_allegato          VARCHAR2(10 CHAR),
  ds_tipo_allegato     VARCHAR2(250 CHAR),
  fg_testo_libero      CHAR(1 CHAR) DEFAULT '0' NOT NULL,
  ty_file              VARCHAR2(10 CHAR),
  nm_file              VARCHAR2(250 CHAR),
  oj_allegato          BLOB,
  cd_hash              VARCHAR2(250 CHAR),
  cd_stato             VARCHAR2(10 CHAR),
  cd_ope_ins_dt        CHAR(6 CHAR) NOT NULL,
  ts_ins_dt            TIMESTAMP NOT NULL,
  cd_ope_ult_agg_dt    CHAR(6 CHAR),
  ts_ult_agg_dt        TIMESTAMP,
  ds_note_dt           VARCHAR2(50 CHAR),
  fg_attestazione      CHAR(1 CHAR),
  fg_incluso_att_conformita CHAR(1 CHAR),
  CONSTRAINT allegato_pk PRIMARY KEY (id_allegato),
  CONSTRAINT allegato_fg_testo_libero_ck CHECK (fg_testo_libero IN ('0','1')),
  CONSTRAINT allegato_fg_attestazione_ck CHECK (fg_attestazione IN ('0','1')),
  CONSTRAINT allegato_fg_incl_att_conf_ck CHECK (fg_incluso_att_conformita IN ('0','1'))
);

COMMENT ON TABLE MATR_USR.ALLEGATO IS
  'Documenti allegati agli atti. Struttura ereditata dal sistema Side del Comune di Milano,
   con i nomi di origine; i tipi sono tradotti da PostgreSQL a Oracle.';
COMMENT ON COLUMN MATR_USR.ALLEGATO.id_atto_sipo IS
  'Chiave dell atto in SIPO. Unico scostamento dall origine, dove la colonna e id_evento e
   referenzia atti.evento: quella tabella non esiste in SIPO.';
COMMENT ON COLUMN MATR_USR.ALLEGATO.oj_allegato IS
  'Contenuto del documento. Si veda la raccomandazione sull uso di un archivio a oggetti.';

CREATE INDEX MATR_USR.IX_ALLEGATO_ATTO ON MATR_USR.ALLEGATO(id_atto_sipo);""".split('\n')

# nome troppo lungo → nome conforme. Tutti verificati ≤ 30 caratteri.
RINOMINE = [
    ('CK_ANSC_STATO_ATTO_FLG_EMERGENZA', 'CK_STATO_ATTO_FLG_EMERGENZA'),
    ('allegati_usecase_fg_att_conf_check', 'allegati_usecase_att_conf_ck'),
    ('allegati_usecase_fg_operante_check', 'allegati_usecase_operante_ck'),
    ('valore_dominio_dominio_decodifica_fk', 'valore_dominio_dominio_fk'),
]


def main():
    if os.path.exists(DST):
        os.remove(DST)
    shutil.copy(SRC, DST)

    def commenti(p):
        return int(subprocess.run(
            ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                           "grep -o '<w:comment ' | wc -l" % repr(p)],
            capture_output=True, text=True).stdout.strip() or 0)
    prima = commenti(DST)

    d = docx.Document(DST)
    fatti = []

    # ------------------------------------------------ 1. il blocco DDL di ALLEGATO
    ps = d.paragraphs
    i = next(k for k, p in enumerate(ps)
             if p.text.strip() == 'CREATE TABLE ANSC_USR.ALLEGATO (')
    j = next(k for k, p in enumerate(ps)
             if k > i and p.text.strip().startswith('CREATE INDEX ANSC_USR.IX_ALLEGATO_ATTO'))
    vecchi = ps[i:j + 1]
    if any(D.ha_commenti(p._p) for p in vecchi):
        raise SystemExit('il blocco DDL porta commenti: non si riscrive in blocco')
    ancora = vecchi[0]._p
    D.ddl(d, ancora, DDL)
    for p in vecchi:
        p._p.getparent().remove(p._p)
    fatti.append('DDL di ALLEGATO sostituito con quello fornito (%d righe)' % len(DDL))

    # la prosa che precede il DDL deve dire perché lo schema è un altro
    h = D.h(d, 3, 'ALLEGATO', esatto=True)
    D.para(d, h._p.getnext(),
           '⚠️ **È l’unica tabella del disegno che non sta in ANSC_USR: sta in MATR_USR, '
           'insieme agli atti.** La ragione è che non è un dato dell’integrazione ma un '
           'dato dell’atto — i documenti allegati esistono per l’atto e con l’atto, e '
           '`id_atto_sipo` punta lì. Tenerla nello schema dedicato avrebbe prodotto un '
           'riferimento che attraversa due schemi su un legame che è invece interno a uno '
           'solo. **RNF-7 resta rispettato**: è una tabella nuova in uno schema esistente, '
           'non una modifica a tabelle esistenti. ⚠️ Ne discende però che il componente '
           'che la scrive ha bisogno di permessi di scrittura su MATR_USR, dove finora '
           'leggeva soltanto: è il punto aperto OP-68.')
    fatti.append('dichiarata la ragione dello schema e il suo prezzo')

    # ------------------------------------------------ 2. gli altri nomi troppo lunghi
    for vecchio, nuovo in RINOMINE:
        D.sostituisci(d, vecchio, nuovo, etichetta='%s → %s' % (vecchio, nuovo),
                      fatti=fatti)

    # ------------------------------------------------ 3. i punti che lo schema tocca
    D.sostituisci(
        d, 'Le tabelle seguono le convenzioni di SIPO (schema dedicato ANSC_USR, tipi '
           'Oracle 12.2).',
        'Le tabelle seguono le convenzioni di SIPO (schema dedicato ANSC_USR, tipi Oracle '
        '12.2), con una sola eccezione dichiarata: ALLEGATO, che sta in MATR_USR per la '
        'ragione detta nella sua sezione.',
        attese=1, etichetta='la premessa alle tabelle dichiara l’eccezione', fatti=fatti)
    D.sostituisci(
        d, 'La scelta mantiene la logica del “tutto additivo” e non impatta gli schemi '
           'esistenti (coerentemente con RNF-7): lo stato ANSC dell’atto vive in ANSC_USR, '
           'agganciato alla chiave dell’atto SIPO.',
        'La scelta mantiene la logica del “tutto additivo” e non modifica alcuna tabella '
        'esistente (coerentemente con RNF-7): lo stato ANSC dell’atto vive in ANSC_USR, '
        'agganciato alla chiave dell’atto SIPO. ⚠️ Una sola tabella nuova nasce invece in '
        'MATR_USR — ALLEGATO, i documenti dell’atto — perché è dato dell’atto e non '
        'dell’integrazione: resta additiva, ma richiede permessi di scrittura su uno '
        'schema in esercizio (OP-68).',
        attese=1, etichetta='«Adeguamenti database» allineato', fatti=fatti)
    D.sostituisci(d, 'ANSC_USR.ALLEGATO — colonne della firma',
                  'MATR_USR.ALLEGATO — colonne della firma',
                  attese=1, etichetta='voce di roadmap allineata', fatti=fatti)

    # ------------------------------------------------ 4. l'ERD
    subprocess.run([sys.executable, os.path.join(os.path.dirname(
        os.path.abspath(__file__)), 'diagrammi_erd.py')], check=True,
        cwd=os.path.dirname(os.path.abspath(__file__)))
    D.sostituisci_immagine(d, 'Schema ANSC_USR', os.path.join(IMG, 'erd_ansc_usr.png'))
    D.sostituisci(
        d, 'Schema ANSC_USR. Le frecce piene sono le chiavi esterne dichiarate nel DDL',
        'Schema ANSC_USR, con ALLEGATO che vi è adiacente perché sta in MATR_USR. Le '
        'frecce piene sono le chiavi esterne dichiarate nel DDL',
        attese=1, etichetta='didascalia dell’ERD allineata', fatti=fatti)
    fatti.append('ERD rigenerato: ALLEGATO marcata MATR_USR, «fg_extra» → «fg_testo_libero»')

    # ------------------------------------------------ 5. punto aperto
    to = D.trova_tabella(d, '#', 'Questione')
    esistenti = [r.cells[0].text.strip() for r in to.rows[1:]]
    n = max(int(x.split('-')[1]) for x in esistenti if x.startswith('OP-')) + 1
    D.clona_riga(to, (
        'OP-%02d' % n,
        'Con quali permessi il componente scrive ALLEGATO in MATR_USR. ⚠️ Il disegno '
        'prevedeva che il concentratore scrivesse solo nel proprio schema e leggesse SIPO; '
        'portando i documenti in MATR_USR si apre una via di scrittura su uno schema in '
        'esercizio. Da definire il permesso minimo (sulla sola tabella, non sullo schema) '
        'e chi lo concede. Vi si lega la crescita in volume: i documenti sono BLOB e '
        'pesano sul backup dello schema di Stato Civile, non più su quello dedicato.',
        'Base dati / Sistemi'))
    fatti.append('OP-%02d aperto (permessi di scrittura su MATR_USR)' % n)

    # ------------------------------------------------ 6. testata e storia
    for tb in d.tables:
        if tb.rows[0].cells[0].text.strip().lower().startswith('area organizzativa'):
            for r in tb.rows:
                v = {'Versione': '3.32',
                     'Documento': 'ANALISI_Integrazione-ANSC_v3.32'}.get(
                        r.cells[0].text.strip())
                if v:
                    D.riscrivi_cella(r.cells[1], v)
            break
    D.storia(d, '07/10/2026', '3.32',
             'Modello dati (ALLEGATO, ERD) · Adeguamenti database · Appendice DDL · '
             'Open Point',
             'Recepito il DDL aggiornato di ALLEGATO. Due le modifiche. La prima è la '
             'lunghezza dei nomi dei vincoli, che superavano il limite di Oracle: '
             'controllati tutti i settantacinque identificativi del DDL, cinque lo '
             'superavano e sono stati accorciati, non soltanto quello segnalato. La '
             'seconda è lo spostamento della tabella da ANSC_USR a MATR_USR, perché i '
             'documenti sono dato dell’atto e non dell’integrazione: la ragione è '
             'dichiarata nella sezione della tabella, l’ERD è stato rigenerato e i punti '
             'del documento che davano lo schema dedicato come unico sono stati allineati. '
             'Ne discende un punto aperto sui permessi di scrittura su uno schema in '
             'esercizio. Corretto nell’ERD anche «fg_extra», nome superato dalla v3.29.')
    fatti.append('testata e storia aggiornate')

    # ------------------------------------------------ 7. grassetti
    def grassetti(p):
        if '**' not in p.text or D.ha_commenti(p._p):
            return 0
        pezzi = D.segmenta('', p.text)
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        for testo, gr in pezzi:
            run = p.add_run(testo)
            run.bold = gr
        return 1
    k = sum(grassetti(p) for p in d.paragraphs)
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    k += grassetti(p)
    fatti.append('%d paragrafi con grassetto applicato' % k)

    d.save(DST)
    dopo = commenti(DST)
    if dopo != prima:
        raise SystemExit('COMMENTI PERSI: erano %d, sono %d' % (prima, dopo))
    fatti.append('commenti di Word conservati: %d → %d' % (prima, dopo))

    print('\n'.join(' · ' + f for f in fatti))
    print('capitoli/tabelle/immagini:', D.riepilogo(DST))
    print('scritto:', os.path.relpath(DST, BASE))


if __name__ == '__main__':
    main()
