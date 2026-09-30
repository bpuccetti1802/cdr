# -*- coding: utf-8 -*-
"""v3.14 — il back-office segue il modello dati: si configura per UC, non per Modello.

⚠️ Regola del progetto: se cambia la configurazione deve cambiare il back-office, perché è
l'unico luogo in cui i funzionari la toccano. Tre schermate su quattordici cambiano oggetto.
"""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOC = os.path.join(BASE, 'Documenti finali', 'ANALISI_Integrazione-ANSC_v3.14.docx')

RIGHE = {
    'Configurazione — Modelli': (
        'Configurazione — Casi d’uso',
        'ANSC_CFG_UC: una riga per UC adottato dal Comune, con il Modello di atto a cui si '
        'applica, il tipo atto, la maschera, l’interfaccia ANSC, il mapper, la priorità fra '
        'gli UC dello stesso Modello e lo stato di attivazione.',
        'Nuovo UC, attiva/disattiva, clona, ordina la priorità, apri condizioni, campi e '
        'allegati dell’UC; simula su un atto reale.',
        'Admin'),
    'Regole': (
        'Regole di controllo e di generazione',
        'ANSC_CFG_REGOLA e ANSC_CFG_REGOLA_CONDIZIONE: le regole di controllo e di '
        'generazione, che restano dichiarative. ⚠️ La determinazione dell’UC non è più qui: '
        'è una condizione dell’UC configurato (PC-9), e si governa nella sua schermata.',
        'Nuova regola, modifica delle condizioni, abilita/disabilita, verifica d’impatto.',
        'Amministratore'),
}

ORIGINE = (
    'In tutte le schermate della configurazione ogni riga mostra la propria origine — '
    'proposta, confermata, modificata, inserita — e il filtro predefinito è «non ancora '
    'esaminate». È così che si lavora una precompilazione di decine di migliaia di righe '
    'senza doverle scorrere tutte: ciò che nessuno ha guardato viene per primo, ciò che un '
    'funzionario ha già deciso non torna a chiedere attenzione a ogni revisione del mapping.')


def applica(doc):
    fatti = []
    t = D.trova_tabella(doc, 'Schermata', 'Scopo', 'Azioni principali', 'Ruolo')
    assert t is not None, 'tabella delle schermate non trovata'
    for riga in t.rows:
        nome = riga.cells[0].text.strip()
        if nome in RIGHE:
            for j, v in enumerate(RIGHE[nome]):
                D.testo_di(riga.cells[j].paragraphs[0], v)
                for p in riga.cells[j].paragraphs[1:]:
                    p._p.getparent().remove(p._p)
            fatti.append(f'schermata «{nome}» → «{RIGHE[nome][0]}»')

    # la corrispondenza dei campi resta per UC, ma cambia il verbo: non si importa soltanto
    for riga in t.rows:
        if riga.cells[0].text.strip().startswith('Configurazione — Campi'):
            D.testo_di(riga.cells[2].paragraphs[0],
                       'Precompila dal mapping ANSC (campi e allegati insieme), esamina e '
                       'conferma riga per riga, completa la corrispondenza con SIPO, dichiara '
                       'la maschera di caricamento degli allegati.')
            fatti.append('schermata «Campi e allegati per UC»: azioni riscritte')

    # la nota sull'origine, subito dopo la tabella
    i = list(doc.element.body).index(t._tbl)
    dopo = list(doc.element.body)[i + 1]
    D.para(doc, dopo, ORIGINE)
    fatti.append('nota sull’origine delle righe')
    return fatti


if __name__ == '__main__':
    doc = docx.Document(DOC)
    for f in applica(doc):
        print('  ·', f)
    doc.save(DOC)
    print('salvato')
