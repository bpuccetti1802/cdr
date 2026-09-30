# -*- coding: utf-8 -*-
"""La riga che separa i campi di ANSC dalle colonne aggiunte da Side va letta come un divisorio."""
import os
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from v3_20_notifiche import DOC  # noqa: E402

ETICHETTA = 'Colonne tecniche aggiuntive di SIDe'

doc = docx.Document(DOC)
fatti = 0
for t in doc.tables:
    for r in t.rows:
        if r.cells[0].text.strip() == ETICHETTA:
            for p in r.cells[0].paragraphs:
                for run in p.runs:
                    run.bold = True
            fatti += 1
assert fatti == 1, fatti
doc.save(DOC)
print('separatore evidenziato')
