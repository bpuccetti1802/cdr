# -*- coding: utf-8 -*-
"""Confronto fra la tabella evento di Side, ANSC_STATO_ATTO e il ModelEvento di ANSC (23/10/2026).

Richiesta dell'utente: le due tabelle assolvono la stessa funzione e ne deve restare una; la
lista delle differenze rispetto alle parti comuni del ModelEvento.

Le colonne si leggono dalla DDL della v3.45 (Appendice A: «Ansc-evento» e ANSC_STATO_ATTO), le
proprietà di primo livello dal contratto ansc/docs/openapi/model_evento.yaml. La corrispondenza
colonna → proprietà è dichiarata a mano in MAPPA_EVENTO e MAPPA_STATO: è l'unica parte non
ricavata, e il controllo in fondo verifica che ogni proprietà nominata esista nel contratto.

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/confronto_evento_stato_atto.py
"""
import os
import re
import sys

import docx
import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill

QUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(QUI)
RADICE = os.path.dirname(BASE)
DOC = os.path.join(BASE, 'ANALISI_Integrazione-ANSC_v3.45.docx')
YAML = os.path.join(RADICE, 'ansc', 'docs', 'openapi', 'model_evento.yaml')
OUT = os.path.join(BASE, 'CONFRONTO_evento-ANSC_STATO_ATTO_v0.1.xlsx')

# colonna di evento (Side) -> proprietà del ModelEvento ('' = nessuna)
MAPPA_EVENTO = {
    'id_evento': '', 'id_ansc': 'idAnsc', 'ty_evento': 'idTipoEvento',
    'ty_contenuto': 'idtipocontenuto', 'ty_use_case': 'idUsecase',
    'ds_use_case': 'descrizioneCasoUso', 'cd_stato_evento': 'stato',
    'fg_secretato': 'flagsecretato', 'id_version': 'idVersion',
    'ts_formazione_atto': 'dataformazione', 'cd_operatore': 'operatore',
    'ty_ruolo_operatore_usc': 'ruoloOperatore',
    'cd_fiscale_operatore': 'operatore', 'ds_nome_operatore': 'operatoreNome',
    'ds_cognome_operatore': 'operatoreCognome',
    'fg_redatto_nella_cc': 'redattoNellaCasaComunale', 'id_luogo_redazione': 'luogoRedazione',
    'id_comune_registrazione': 'idComuneRegistrazione',
    'nm_comune_registrazione': 'nomeComuneRegistrazione',
    'id_provincia_registraz': 'idProvinciaRegistrazione',
    'cd_sigla_prov_registraz': 'siglaProvinciaRegistrazione',
    'ts_decorrenza': 'dataDecorrenza', 'ty_motivo_recupero': 'motivoRecupero',
    'ds_motivo_recupero': 'descrizioneMotivoRecupero', 'nr_atto': 'numeroatto',
    'ty_dichiarante': 'tipoDichiarante', 'fg_ausilio_interprete': 'ausilioInterprete',
    'ds_note_ritardo': 'noteRitardo', 'ds_motivo_annullamento': 'motivoAnnullamento',
    'ty_consenso': 'tipoConsenso', 'ty_certificabilita': 'certificabile',
    'fg_sentenza_figlio_inc': 'sentenzaFiglioIncestuoso',
    'fg_proposta_annotazione': 'propostaAnnotazione',
    'fg_annotaz_non_certificabile': 'flagAnnotazioneNonCertificabile',
    'ty_dati_prodotti': 'datiProdotti', 'id_ente_dichiarante': 'enteDichiarante',
    'id_madre': 'madre', 'id_padre': 'padre', 'id_coniuge': 'coniuge',
    'id_dichiarante': 'dichiarante', 'id_interprete': 'interprete',
    'id_dati_evento_riconoscimento': 'datiEventoRiconoscimento',
    'id_riconoscimento': 'riconoscimento', 'id_riconoscimento_madre': 'riconoscimentoMadre',
    'id_riconoscimento_padre': 'riconoscimentoPadre', 'id_consenso_da_atto_coll': 'consenso',
    'id_consenso_madre': 'consensoMadre', 'id_consenso_padre': 'consensoPadre',
    'id_atto_not_cons_madre': 'attoNotarileConsensoMadre',
    'id_atto_not_cons_padre': 'attoNotarileConsensoPadre', 'id_tribunale': 'tribunale',
    'id_evento_collegato': 'eventoCollegato', 'tx_composiz_completa': 'composizioneCompleta',
    'ty_sentenza_riconoscim': 'idRiconoscimentoSentenza', 'cd_forzatura_validazione': '',
    'cd_ope_ins_dt': '', 'ts_ins_dt': '', 'cd_ope_ult_agg_dt': '', 'ts_ult_agg_dt': '',
    'ds_note_dt': '',
    'cd_fiscale_operat_firmatario': 'operatoreFirmatario',
    'ds_operat_firmatario_nome': 'operatoreFirmatarioNome',
    'ds_operat_firmatario_cognome': 'operatoreFirmatarioCognome',
    'id_madre_biologica': 'madreBiologica', 'id_padre_biologico': 'padreBiologico',
    'fg_certificato_semplice': 'flagCertificatoSemplice',
    'fg_estratto_per_riassunto': 'flagEstrattoPerRiassunto',
    'id_ufficiale_stato_civile': 'ufficialeStatoCivile', 'id_evento_primario': 'eventoPrimario',
    'id_atto_cartaceo': 'idAttoCartaceo', 'nr_ricezione': 'numeroRicezione',
    'id_dati_annotazione_sentenza': 'datiAnnotazioneSentenza',
    'id_evento_collegato_cittadinan': 'eventoCollegatoCittadinanza',
    'fg_omogenitoriale': 'flagOmogenitoriale',
    'tx_rett_annotaz_composiz_compl': 'rettificaAnnotazioneComposizioneCompleta',
}

# proprietà del ModelEvento -> colonna di ANSC_STATO_ATTO (dalla tabella «Testata del modello
# evento» della v3.45: 16 in colonna; le derivate si ricavano da una colonna)
MAPPA_STATO = {
    'id': 'ID_EVENTO_ANSC', 'idAnsc': 'ID_ANSC', 'idTipoEvento': 'ID_TIPO_EVENTO',
    'idtipocontenuto': 'ID_TIPO_CONTENUTO', 'idUsecase': 'ID_UC_ANSC', 'stato': 'STATO',
    'flagsecretato': 'FLG_SECRETATO', 'idVersion': 'COD_VERSIONE_ANSC',
    'dataformazione': 'DATA_FORMAZIONE', 'ruoloOperatore': 'COD_RUOLO_OPERATORE',
    'redattoNellaCasaComunale': 'FLG_CASA_COMUNALE',
    'idComuneRegistrazione': 'ID_COMUNE_REGISTRAZIONE', 'dataDecorrenza': 'DATA_DECORRENZA',
    'motivoRecupero': 'COD_MOTIVO_RECUPERO', 'idAttoCartaceo': 'ID_ATTO_EMERGENZA',
    'numeroatto': 'NUM_COMUNALE',
}
DERIVATE = {
    'descrizioneCasoUso': 'da ANSC_CFG_UC, per codice',
    'annoAtto': 'da DATA_FORMAZIONE', 'meseAtto': 'da DATA_FORMAZIONE',
    'giornoAtto': 'da DATA_FORMAZIONE', 'ora': 'da DATA_FORMAZIONE',
    'minuto': 'da DATA_FORMAZIONE', 'nomeComuneRegistrazione': 'dal comune',
    'idProvinciaRegistrazione': 'dal comune', 'siglaProvinciaRegistrazione': 'dal comune',
    'oraDecorrenza': 'da DATA_DECORRENZA', 'minutoDecorrenza': 'da DATA_DECORRENZA',
}

# Differenze puntuali sulle parti comuni (testata), verificate sulle due DDL
NOTE = {
    'idAnsc': 'Lunghezza: 60 in evento, 50 in ANSC_STATO_ATTO.',
    'idTipoEvento': 'Stessa lunghezza (10). Nella scheda di ANSC_STATO_ATTO il tipo è dichiarato '
                    'VARCHAR2(20) con valori MORTE, NASCITA…, mentre la DDL e ANSC_01 usano 1…5: '
                    'la scheda va allineata alla DDL.',
    'idUsecase': 'Lunghezza: 10 in evento, 20 in ANSC_STATO_ATTO (i codici UC arrivano a 8 cifre).',
    'descrizioneCasoUso': 'evento la ripete in colonna (text); ANSC_STATO_ATTO la ricava dal codice.',
    'stato': 'evento registra il codice ANSC_11 (varchar 10); ANSC_STATO_ATTO un nome (CONFERMATO, '
             'FIRMATO_USC…) più gli stati locali IN_PREPARAZIONE e BOZZA_APPLICATIVA, con CHECK.',
    'flagsecretato': 'Uguale: 0/1.',
    'idVersion': 'Uguale: 10 caratteri.',
    'dataformazione': 'Uguale: un timestamp; anno, mese, giorno, ora e minuto si ricavano.',
    'operatore': 'Deprecato da ANSC. evento ha cd_operatore CHAR(6) (codice utente di Milano); '
                 'ANSC_STATO_ATTO ha OPERATORE VARCHAR2(40) (utenze di Roma). Il commento di '
                 'evento nomina anche cd_fiscale_operatore, che la CREATE non dichiara.',
    'ruoloOperatore': 'Nella DDL di evento la colonna ty_ruolo_operatore_usc è scritta DENTRO il '
                      'commento «--» della riga di cd_operatore: la CREATE non la crea.',
    'operatoreNome': 'Deprecato. evento lo nomina solo nei COMMENT (ds_nome_operatore), non nella '
                     'CREATE: il COMMENT fallirebbe.',
    'operatoreCognome': 'Come operatoreNome (ds_cognome_operatore).',
    'redattoNellaCasaComunale': 'Dominio diverso: 0/1 in evento, S/N in ANSC_STATO_ATTO; il '
                                'contratto dichiara un booleano.',
    'luogoRedazione': 'evento lo normalizza in un’altra tabella (id_luogo_redazione); '
                      'ANSC_STATO_ATTO lo lascia nel modello evento (TXT_EVENTO).',
    'idComuneRegistrazione': 'Tipo: int4 in evento, VARCHAR2(10) in ANSC_STATO_ATTO; il contratto '
                             'dichiara una stringa.',
    'nomeComuneRegistrazione': 'evento lo ripete in colonna; ANSC_STATO_ATTO lo ricava.',
    'idProvinciaRegistrazione': 'evento lo ripete in colonna; ANSC_STATO_ATTO lo ricava.',
    'siglaProvinciaRegistrazione': 'evento lo ripete in colonna; ANSC_STATO_ATTO lo ricava.',
    'dataDecorrenza': 'Uguale: un timestamp.',
    'motivoRecupero': 'Lunghezza: 10 in evento, 20 in ANSC_STATO_ATTO.',
    'descrizioneMotivoRecupero': 'evento in colonna (text); ANSC_STATO_ATTO nel modello evento.',
    'idAttoCartaceo': 'Stessa lunghezza (256); nome diverso: id_atto_cartaceo / ID_ATTO_EMERGENZA.',
    'numeroatto': 'Tipo: int4 in evento, VARCHAR2(30) in ANSC_STATO_ATTO; il contratto dichiara '
                  'una stringa di 5 caratteri, il Comune ne assume 6 (OP-73).',
    'id': 'Il campo id di ANSC («per uso interno») non ha colonna in evento; in ANSC_STATO_ATTO è '
          'ID_EVENTO_ANSC. Serve a riconoscere l’evento alla rivalidazione (OP-46, OP-75).',
    'flagAnnotazioneNonCertificabile': 'In evento il COMMENT dice 1 = NON CERTIFICABILE, '
                                       '2 = CERTIFICABILE, ma il CHECK ammette 0/1; il contratto '
                                       'dichiara un booleano.',
    'attoNotarileConsensoMadre': 'In evento il COMMENT è identico a quello di id_consenso_madre: '
                                 'una delle due descrizioni è errata.',
    'attoNotarileConsensoPadre': 'Come attoNotarileConsensoMadre.',
}


def leggi_ddl():
    d = docx.Document(DOC)
    ps = [p.text for p in d.paragraphs]
    i = next(k for k, t in enumerate(ps) if t.startswith('CREATE TABLE  ansc-evento'))
    evento = []
    for t in ps[i + 1:]:
        if t.strip().startswith(('CONSTRAINT', ')')):
            break
        nome, tipo = t.split()[0], t.split()[1].rstrip(',')
        evento.append((nome, tipo, t))
        # la riga di cd_operatore nasconde ty_ruolo_operatore_usc dentro il commento
        m = re.search(r'\s{2}(ty_ruolo_operatore_usc) (varchar\(\d+\))', t)
        if m:
            evento.append((m.group(1), m.group(2) + ' (dentro il commento: non creata)', t))
    commenti = sorted({re.match(r'COMMENT ON COLUMN\s+evento\.(\w+)', t).group(1)
                       for t in ps if re.match(r'COMMENT ON COLUMN\s+evento\.', t)})
    j = next(k for k, t in enumerate(ps) if t.startswith('CREATE TABLE ANSC_USR.ANSC_STATO_ATTO'))
    stato = {}
    for t in ps[j + 1:]:
        if t.strip().startswith('CONSTRAINT'):
            break
        nome, tipo = t.split()[0], ' '.join(t.split()[1:]).rstrip(',')
        stato[nome] = tipo
    return evento, commenti, stato


def proprieta():
    s = yaml.safe_load(open(YAML))['components']['schemas']['ModelEvento']
    p = dict(s.get('properties', {}))
    for a in s.get('allOf', []):
        p.update(a.get('properties', {}))
    out = {}
    for k, v in p.items():
        t = v.get('type') or v.get('$ref', '').split('/')[-1] or 'oggetto'
        if t == 'array':
            t = '[' + (v['items'].get('$ref', '').split('/')[-1] or v['items'].get('type')) + ']'
        out[k] = (t, bool(v.get('deprecated')), v.get('maxLength'))
    return out


def main():
    evento, commenti, stato = leggi_ddl()
    prop = proprieta()
    for c, p in MAPPA_EVENTO.items():
        if p and p not in prop:
            raise SystemExit(f'proprietà inesistente nel contratto: {p} ({c})')
    for p, c in MAPPA_STATO.items():
        if p not in prop or c not in stato:
            raise SystemExit(f'corrispondenza errata: {p} -> {c}')
    nomi_evento = {n for n, _, _ in evento}
    solo_commento = [c for c in commenti if c not in nomi_evento]

    righe = []
    ordine = list(prop)
    testata = ordine[:ordine.index('numeroatto') + 1]
    inversa = {}
    for c, p in MAPPA_EVENTO.items():
        if p:
            inversa.setdefault(p, []).append(c)
    tipi_evento = {n: t for n, t, _ in evento}
    for k in ordine:
        t, depr, ml = prop[k]
        col_ev = ', '.join(inversa.get(k, []))
        tipo_ev = ', '.join(tipi_evento.get(c, 'solo nel COMMENT') for c in inversa.get(k, []))
        col_st = MAPPA_STATO.get(k, '')
        tipo_st = stato.get(col_st, '')
        if not col_st:
            col_st = ('— ' + DERIVATE[k]) if k in DERIVATE else '— nel modello evento (TXT_EVENTO)'
        parte = 'Testata' if k in testata else 'Corpo'
        if depr:
            esito = 'Deprecato da ANSC'
        elif inversa.get(k) and MAPPA_STATO.get(k):
            esito = 'In colonna in entrambe'
        elif inversa.get(k):
            esito = 'In colonna solo in evento'
        elif MAPPA_STATO.get(k):
            esito = 'In colonna solo in ANSC_STATO_ATTO'
        else:
            esito = 'In nessuna delle due (sta nel JSON o in altre tabelle)'
        righe.append((parte, k, t + (f' (max {ml})' if ml else ''), 'sì' if depr else '',
                      col_ev, tipo_ev, col_st, tipo_st, esito, NOTE.get(k, '')))

    solo_ev = [(n, t) for n, t, _ in evento if not MAPPA_EVENTO.get(n)]
    solo_st = [c for c in stato if c not in MAPPA_STATO.values()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ModelEvento'
    testa = ('Parte', 'Proprietà ModelEvento', 'Tipo nel contratto', 'Deprecata',
             'Colonna evento (Side)', 'Tipo evento', 'Colonna ANSC_STATO_ATTO',
             'Tipo ANSC_STATO_ATTO', 'Esito', 'Differenza / nota')
    _foglio(ws, testa, righe, (9, 30, 26, 9, 30, 22, 30, 24, 30, 70))

    ws2 = wb.create_sheet('Solo in evento')
    nota2 = {
        'id_evento': 'Chiave locale di Side. In ANSC_STATO_ATTO è ID_STATO_ATTO; l’aggancio a '
                     'SIPO è ID_ATTO_SIPO.',
        'cd_forzatura_validazione': 'Non è del modello evento: è un parametro della richiesta '
                                    'R009. Da noi sta in forzatura_validazione, con '
                                    'l’ufficiale e le anomalie.',
        'cd_ope_ins_dt': 'Campi tecnici di Side; in ANSC_STATO_ATTO DATA_INS/UTENTE_INS…',
    }
    _foglio(ws2, ('Colonna evento', 'Tipo', 'Nota'),
            [(n, t, nota2.get(n, 'Campo tecnico di Side; in ANSC_STATO_ATTO i campi tecnici '
                                 '[R5].' if n.startswith(('cd_ope', 'ts_', 'ds_note'))
                                 else '')) for n, t in solo_ev]
            + [(c, 'solo nel COMMENT', 'La CREATE non dichiara la colonna: il COMMENT fallirebbe.')
               for c in solo_commento],
            (32, 40, 90))

    ws3 = wb.create_sheet('Solo in ANSC_STATO_ATTO')
    nota3 = {
        'ID_STATO_ATTO': 'Chiave; richiamata da ANSC_LOG_AUDIT e forzatura_validazione.',
        'ID_ATTO_SIPO': 'L’atto di SIPO: in Side l’evento è l’atto stesso.',
        'COD_TIPO_OPERAZIONE': 'Creazione, rettifica, annotazione, annullamento.',
        'ID_MODELLO_ATTO': 'Il Modello del Comune da cui si è determinato l’UC.',
        'COD_LOGICA': 'La logica che ha determinato l’UC.',
        'COD_ORIGINE_UC': 'Automatica o scelta dall’operatore.',
        'COD_FASE': 'Fase della lavorazione locale.',
        'ID_VERSIONE': 'Baseline della configurazione usata.',
        'TIPO_ESITO': 'SCARTATO o INDETERMINATO.',
        'ID_OPERAZIONE_ANSC': 'Tracciamento della chiamata.',
        'FLG_EMERGENZA': 'Modalità di emergenza.',
        'DATA_INGRESSO_EMERGENZA': 'Modalità di emergenza (OP-48).',
        'TXT_EVENTO': 'Il modello evento intero: contiene tutto il corpo che evento porta in '
                      'colonna o in tabelle collegate.',
        'TXT_HASH_VALIDATO': 'Impronta all’ultima validazione: fa scattare lo stato 0.',
        'DATA_VALIDAZIONE': 'Ultima validazione riuscita.',
        'COD_PAGINA_RIPRESA': 'Ripresa della redazione.',
        'FLG_SOSPESO': 'Ripresa della redazione.', 'DATA_SOSPENSIONE': 'Ripresa della redazione.',
        'UTENTE_SOSPENSIONE': 'Ripresa della redazione.',
        'NUM_REVISIONE': 'Due postazioni non si sovrascrivono.',
        'OPERATORE': 'Chi ha lavorato l’atto (in evento cd_operatore, CHAR(6)).',
        'HOSTNAME': 'Postazione.', 'ID_UFFICIALE': 'Indice di supervisione.',
        'COD_MUNICIPIO': 'Indice di supervisione.', 'ULTIMO_ERRORE': 'Diagnosi.',
    }
    _foglio(ws3, ('Colonna ANSC_STATO_ATTO', 'Tipo', 'A che cosa serve'),
            [(c, stato[c], nota3.get(c, 'Campo tecnico [R5].' if c.startswith(('DATA_INS', 'DATA_UPD', 'UTENTE_')) else ''))
             for c in solo_st], (30, 44, 80))

    ws4 = wb.create_sheet('Legenda')
    conteggi = {}
    for r in righe:
        conteggi[(r[0], r[8])] = conteggi.get((r[0], r[8]), 0) + 1
    ws4.append(['Confronto fra la tabella evento di Side (DDL «Ansc-evento», App. A della v3.45), '
                'ANSC_STATO_ATTO e il ModelEvento (model_evento.yaml).'])
    ws4.append([])
    ws4.append(['Proprietà di primo livello del ModelEvento', len(prop)])
    ws4.append(['  di cui testata (righe 2-33 di Tracciato-Evento.xlsx)', len(testata)])
    ws4.append(['Colonne di evento (CREATE)', len(evento)])
    ws4.append(['Colonne nominate solo nei COMMENT di evento', len(solo_commento)])
    ws4.append(['Colonne di ANSC_STATO_ATTO', len(stato)])
    ws4.append([])
    for (parte, esito), n in sorted(conteggi.items()):
        ws4.append([f'{parte} — {esito}', n])
    ws4.column_dimensions['A'].width = 70
    wb.save(OUT)
    print('scritto:', OUT)
    print('proprietà', len(prop), 'testata', len(testata), 'evento', len(evento),
          'solo commento', solo_commento, 'stato', len(stato))
    for k, n in sorted(conteggi.items()):
        print(' ', k, n)


def _foglio(ws, testa, righe, larghezze):
    ws.append(testa)
    for r in righe:
        ws.append(list(r))
    blu = PatternFill('solid', fgColor='1F4E79')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = blu
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for riga in ws.iter_rows(min_row=2):
        for c in riga:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(larghezze):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


if __name__ == '__main__':
    main()
