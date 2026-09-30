# -*- coding: utf-8 -*-
"""v3.20 — la gestione delle notifiche ereditata dal sistema Side del Comune di Milano,
più il processo di allineamento delle decodifiche indicato dal fornitore.

Il documento aveva già un capitolo sul flusso in ingresso: che cosa sono le notifiche, il
cancello anagrafico, le tre situazioni, i limiti del contratto R008, la decisione (PC-10) di
lavorarle nel back-office. Quello dice **che cosa** va fatto. Il capitolo aggiunto qui dice
**come**, riportando il processo che a Milano è già in esercizio.

⚠️ Le tabelle si riportano as is, con i nomi di Side, e si costruiscono LEGGENDO LE SORGENTI:
non si ricopiano a mano, altrimenti alla prossima revisione divergono.

⚠️ Il rilievo che pesa di più: il batch di scarico è schedulabile, cioè funziona senza un
operatore autenticato. Presuppone quindi risolto OP-23.
"""
import os
import sys

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(BASE, 'Sorgenti Documentali')
DOC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.20.docx')

XLS = os.path.join(SRC, 'ANSC - Schema delle Notifiche.xlsx')
FUN = os.path.join(SRC, 'CdMI SIDe MI_ANSC _Gestione Notifiche ANSC (ID 14) v2.docx')


# ----------------------------------------------------------------- sorgenti

def _celle(ws):
    for r in ws.iter_rows(values_only=True):
        yield ['' if c is None else str(c).strip().replace('\n', ' — ') for c in r]


def righe_struttura():
    """Le 58 righe della notifica, as is: campo, descrizione (+ esempio), note (+ attività)."""
    ws = openpyxl.load_workbook(XLS, data_only=True)['Struttura del flusso']
    out = [['Campo', 'Descrizione e valori', 'Note']]
    for r in list(_celle(ws))[1:]:
        campo, descr, esempio, note, att = (r + [''] * 5)[:5]
        if not campo:
            continue
        d = ' '.join(x for x in (descr, esempio) if x)
        n = ' '.join(x for x in (note, att) if x)
        out.append([campo, d, n])
    return out


def righe_foglio(nome, colonne, intestazioni=None):
    ws = openpyxl.load_workbook(XLS, data_only=True)[nome]
    righe = [r for r in _celle(ws) if any(r)]
    testa = intestazioni or [righe[0][i] for i in colonne]
    out = [testa]
    for r in righe[1:]:
        r = r + [''] * 8
        if not any(r[i] for i in colonne):
            continue
        out.append([r[i] for i in colonne])
    return out


def righe_tabella_docx(indice, intestazioni=None):
    d = docx.Document(FUN)
    t = d.tables[indice]
    righe = [[c.text.strip().replace('\n', ' — ') for c in r.cells] for r in t.rows]
    righe = [r for r in righe if any(r)]
    if intestazioni:
        righe[0] = intestazioni
    return righe


# ----------------------------------------------------------------- contenuti

APERTURA = [
    'Il capitolo precedente stabilisce che cosa il Comune debba fare delle notifiche che ANSC '
    'gli invia, e perché non possa non farlo: finché le notifiche generate da un atto non sono '
    'confermate da tutti i comuni coinvolti la comunicazione all’ufficio anagrafe non parte, e '
    'un solo rifiuto la annulla. Il presente capitolo descrive come quel lavoro è organizzato '
    'nel sistema Side del Comune di Milano, che lo ha già in esercizio, e formula le ipotesi '
    'necessarie a portarlo in SIPO.',

    'La scelta di ereditare invece di progettare ha una ragione precisa. Il contratto R008 è '
    'povero: non prevede conferme massive, non ha un campo per la motivazione del rifiuto, e '
    'l’identità di chi conferma è quella della sessione. Attorno a quella povertà Milano ha '
    'costruito ciò che serve davvero all’ufficio — lo scarico periodico, l’incrocio con gli '
    'atti propri, la classificazione di che cosa fare di ciascuna notifica, e due elenchi su '
    'cui lavorare — ed è materiale che conviene riusare.',

    '⚠️ Una cosa va detta subito, perché cambia il dimensionamento di tutto il resto: **Side '
    'non si limita a ricevere le notifiche, le anticipa**. Al momento della firma dell’atto il '
    'sistema predispone già alcune delle azioni che le notifiche richiederanno — le '
    'comunicazioni per l’archivio e le PEC verso i comuni non aderenti — e quando la notifica '
    'arriva la si ritrova collegata a ciò che era stato preparato. È la ragione della colonna '
    'che lega la notifica alla comunicazione predisposta, ed è la parte del processo che in '
    'SIPO non ha alcun corrispondente.',

    'Le tabelle che seguono sono riportate nella forma in cui le sorgenti le esprimono, con i '
    'nomi di Side: la traduzione nella nomenclatura del Comune si fa una volta sola, quando le '
    'strutture vengono recepite, e non in sede di analisi.',
]

DUE = [
    'Le sorgenti trasmesse descrivono due realizzazioni distinte, prodotte a circa un anno di '
    'distanza, e conviene non confonderle perché rispondono a bisogni diversi.',
]
DUE_V = [
    ('La ricerca in linea.',
     'L’analisi tecnica del settembre 2024 aggiunge alla scrivania di stato civile un pulsante '
     '«Ricerca Notifiche ANSC» che interroga R008 al momento, con due modalità di ricerca — per '
     'parametri della notifica e per persona (codice fiscale, cognome e nome, identificativo '
     'ANPR o ANSC del soggetto) — e consente di approvare, rifiutare e vedere l’anteprima '
     'dell’atto. ⚠️ Il rifiuto non si esaurisce nella notifica: apre il provvedimento di '
     'rifiuto, che è un deposito vero e proprio (R016). La ricerca è richiamabile anche dai '
     'pulsanti di evento della scrivania, e in quel caso il registro arriva già impostato e '
     'protetto.'),
    ('Lo scarico periodico.',
     'L’analisi funzionale dell’agosto 2025 aggiunge un batch che porta le notifiche in una '
     'tabella locale, le incrocia con gli eventi del Comune, le classifica e le espone in due '
     'pagine. Non sostituisce la ricerca in linea: le si affianca, perché risponde a una '
     'domanda che la ricerca non sa porre — quali notifiche competono a un certo ufficiale, in '
     'un certo intervallo, e che cosa bisogna farne.'),
]
DUE_CODA = (
    'Il disegno del presente documento prevede entrambe le cose, e le ha già entrambe nominate: '
    'la worklist delle notifiche nel back-office è lo scarico periodico, la consultazione delle '
    'notifiche generate da un atto formato dal Comune è la ricerca in linea. Ciò che Side '
    'aggiunge è il dettaglio di come si fanno.'
)

ATTIVITA = [
    ('Scarico.', 'Un batch interroga R008 e registra in locale le notifiche ricevute.'),
    ('Incrocio.', 'Le notifiche si confrontano con gli eventi registrati localmente, per '
                  'recuperare l’operatore che ha steso l’atto.'),
    ('Arricchimento e classificazione.',
     'A ciascuna notifica si aggiungono l’identificativo dell’evento locale, l’operatore, '
     'l’ambito — cioè che cosa occorre farne — e l’identificativo della comunicazione '
     'anticipatamente predisposta alla firma.'),
    ('Pagina di consultazione.', 'L’elenco delle notifiche scaricate, con i filtri che '
                                 'l’ufficio usa davvero.'),
    ('Pagina delle notifiche senza riscontro.',
     'L’elenco speculare: le comunicazioni predisposte localmente che in ANSC non hanno una '
     'notifica corrispondente. È il controllo che rende visibili le divergenze.'),
]

BATCH = [
    'Il batch si chiama ACQUISIZIONE_NOTIFICHE_ANSC_DB, è schedulabile e registra in locale le '
    'informazioni ricevute da ANSC. In questa fase vengono individuati gli eventi presenti nel '
    'sistema locale, per i quali è possibile arricchire i dati con l’identificativo dell’evento, '
    'l’operatore che ha redatto l’atto, la classificazione per ambito e l’identificativo della '
    'comunicazione predisposta alla firma.',
]
BATCH_CODA = [
    'Accanto alla schedulazione esiste una via manuale a grana grossa: l’operatore scarica '
    'autonomamente le notifiche degli ultimi quindici giorni con un pulsante, mentre gli altri '
    'parametri restano a disposizione dei tecnici abilitati per gli allineamenti straordinari. '
    'È una ripartizione sensata — all’operatore si dà il caso ordinario senza parametri da '
    'sbagliare — e si propone di conservarla.',

    '⚠️ Il parametro numRecordPagina merita attenzione, perché la motivazione dichiarata è '
    'istruttiva: è impostato a 5.000 per ottenere una pagina sola, «perché c’è il rischio di '
    'perdersi qualche notifica». La paginazione di ANSC non offre garanzie di stabilità fra una '
    'pagina e la successiva, e chiedere tutto in una volta è il modo pratico di aggirarla. Alla '
    'scala di Roma il valore va verificato: se il volume di un intervallo superasse la pagina, '
    'la scelta va rifatta e l’accorgimento perde efficacia. È una delle ragioni per cui i '
    'volumi delle notifiche restano una questione aperta.',

    '⚠️ Il fatto che il batch sia schedulabile ha una conseguenza che va dichiarata: presuppone '
    'che la lettura delle notifiche sia eseguibile in modalità macchina-macchina, senza l’OTP '
    'del singolo ufficiale. È esattamente la domanda che il registro degli Open Point lascia '
    'aperta sul perimetro ammesso di quella modalità. Se la risposta fosse negativa il processo '
    'resterebbe valido nella sostanza ma cambierebbe forma: non un batch notturno, bensì un '
    'comando impartito da un amministratore dentro la propria sessione, con la stessa logica di '
    'confronto e di arricchimento. È una differenza di esercizio rilevante e conviene chiarirla '
    'prima di dimensionare il presidio.',
]

STRUTTURA = [
    'La notifica scaricata si conserva in una tabella — in Side notifica_stato_civile — che '
    'ricalca il messaggio di ANSC e vi aggiunge le colonne che servono al Comune. La struttura '
    'completa è riportata di seguito nella forma della sorgente.',

    'Tre osservazioni di lettura. **La prima**: le sezioni atto_collegato e atto_primario del '
    'messaggio ANSC sono riportate appiattite, una colonna per proprietà, perché la tabella è '
    'piana; l’atto collegato è l’atto che ha generato la notifica, l’atto primario è quello su '
    'cui l’annotazione va apposta e può essere cartaceo. **La seconda**: le ultime otto colonne '
    'non vengono da ANSC, le aggiunge Side, e sono quelle su cui si concentrano le ipotesi di '
    'raccordo più avanti. **La terza**: alcune caselle della sorgente contengono domande del '
    'redattore invece di specifiche — sulla struttura dell’identificativo, su quando '
    'l’identificativo del soggetto sia valorizzato — e si sono lasciate come stanno, perché '
    'sono domande che varranno anche per noi.',
]

DOMINI = [
    'I filtri e la classificazione si appoggiano a quattro decodifiche di ANSC e a una locale. '
    'Le prime quattro si ottengono con il servizio delle decodifiche e risiedono nelle tabelle '
    'di replica descritte al capitolo sui dizionari; la quinta è una decodifica del Comune, che '
    'ANSC non conosce.',
]
DOMINI_CODA = [
    '⚠️ L’ambito è il punto in cui il processo smette di essere una replica di ANSC e diventa '
    'organizzazione del Comune. I sette codici dicono che cosa fare della notifica, e tre di '
    'essi — archivio, PEC, gestione manuale — nominano attività che dipendono da come il Comune '
    'ha organizzato l’archivio cartaceo e l’invio della posta certificata. I codici si possono '
    'ereditare; la regola che li assegna va rifatta sull’organizzazione di Roma.',
]

SCENARI = [
    'La classificazione si ricava da tre attributi della notifica — il genere, il canale e il '
    'contrassegno che dice se l’evento è cartaceo — più una discriminante che non è nella '
    'notifica ma nel confronto: se il comune di formazione sia il Comune stesso oppure un '
    'altro. Negli scenari di Milano il confronto è con il codice 5646; per Roma va sostituito '
    'con il proprio codice ISTAT.',
]
SCENARI_CODA = [
    'Due righe della matrice vanno lette con attenzione perché non descrivono un automatismo ma '
    'un limite. **La prima** riguarda le annotazioni con atto primario cartaceo: la sorgente '
    'annota di aver trovato casi in cui l’atto primario risultava invece digitale, e si '
    'interroga se siano conseguenza di una compilazione incoerente o di un aggancio che ANSC '
    'compie per conto proprio. Finché la cosa non è chiarita esiste un ambito dedicato alla '
    'gestione manuale, che è il modo corretto di trattare ciò che l’automatismo non sa '
    'decidere. **La seconda** riguarda le richieste di trascrizione verso comuni non aderenti: '
    'ANSC non le trasmette affatto, quindi non esiste alcuna notifica da lavorare e la PEC va '
    'predisposta alla firma dell’atto. È il caso in cui l’elenco locale è l’unica traccia '
    'esistente, e spiega perché serva la seconda pagina.',
]

COMUNICAZIONI = [
    'Il collegamento fra la notifica e la comunicazione predisposta alla firma è il tratto più '
    'caratteristico del disegno di Side, e insieme quello che in SIPO non trova corrispondenza. '
    'Al momento della firma il sistema scrive già, in una propria tabella delle comunicazioni, '
    'il record della PEC da inviare o della comunicazione da mandare all’archivio; quando la '
    'notifica arriva, il batch la aggancia a quel record.',

    'La sorgente documenta il popolamento delle due comunicazioni tipiche — quella da inviare '
    'per posta certificata e quella destinata all’archivio — indicando per ciascun campo da '
    'dove il valore proviene. La si riporta perché mostra, meglio di qualunque descrizione, '
    'quanta parte della comunicazione sia ricavabile dalla notifica e quanta invece venga da un '
    'catalogo di modelli. Gli esempi di valore presenti nella sorgente, specifici di Milano, '
    'sono omessi.',
]
COMUNICAZIONI_CODA = [
    '⚠️ Il punto su cui decidere è netto: **SIPO non predispone nulla di simile**. Non esiste '
    'una tabella delle comunicazioni da inviare, né il processo che le genera alla firma, né il '
    'catalogo dei modelli con intestazione, oggetto e corpo. Ereditare la colonna senza '
    'ereditare il processo produrrebbe un riferimento sempre vuoto, cioè una struttura che '
    'sembra completa e non lo è. Le strade sono due: portare anche quel processo, che è lavoro '
    'eccedente le notifiche e tocca il protocollo del Comune, oppure conservare la colonna e '
    'dichiararla non valorizzata nel pilota, rimandando la scelta. La seconda è preferibile '
    'purché sia scritto, ed è la posizione che si assume qui.',
]

PAGINE = [
    'Le pagine sono due, entrambe raggiunte da un medesimo pulsante di menu.',
]
PAGINE_V = [
    ('Gestione notifiche per operatore.',
     'L’elenco delle notifiche scaricate, ordinato per identificativo discendente. I filtri '
     'sono: registro di destinazione, ambito, comune di formazione, proposte di annotazione, '
     'operatore, intervallo temporale — con le date obbligatorie — identificativo della '
     'notifica, identificativo dell’evento sorgente, stato, tipo contenuto, genere e '
     'descrizione del soggetto intestatario. La lista mostra intestatario, comune di '
     'formazione, evento sorgente, comune destinatario, registro e atto di destinazione, '
     'ambito, descrizione, data e ora di formazione, identificativo, stato, operatore e genere; '
     'le azioni disponibili sono approvazione, rifiuto e anteprima.'),
    ('Ricerca notifiche esclusivamente locali.',
     'Di sola consultazione, ordinata per identificativo della comunicazione discendente. '
     'Raccoglie almeno le PEC di richiesta di trascrizione e le comunicazioni predisposte alla '
     'firma che non trovano corrispondenza fra le notifiche di ANSC — perché non esistono, come '
     'nel caso delle trascrizioni, oppure perché il batch non le ha ancora scaricate. I filtri '
     'sono operatore, registro di destinazione, intervallo temporale, evento sorgente e genere.'),
]
PAGINE_CODA = [
    '⚠️ La sorgente dichiara il limite di questa impostazione, e va riportato perché riguarda '
    'ciò che l’operatore vede: poiché le pagine mostrano le notifiche a partire dai dati '
    'salvati in locale, non sono visibili quelle create dopo l’ultimo scarico, e gli stati non '
    'risentono di aggiornamenti compiuti altrove — per esempio sulla web app di ANSC. È la '
    'stessa natura di replica che il presente documento attribuisce allo store delle notifiche, '
    'e la conseguenza è che la lista locale va letta come una fotografia, non come lo stato '
    'corrente. Dove la certezza serve, la si ottiene rileggendo la singola notifica: ed è '
    'un’altra ragione per conservare accanto allo scarico la ricerca in linea.',
]

IPOTESI = [
    ('idEventoLocale → la chiave dell’atto in SIPO.',
     'In Side è l’identificativo dell’evento locale che ha generato la notifica, e la sorgente '
     'avverte che si riesce ad alimentarlo solo se l’atto è stato redatto o registrato nel '
     'sistema. In SIPO l’equivalente è la chiave dell’atto, la stessa che ANSC_STATO_ATTO '
     'registra in ID_ATTO_SIPO e che la tabella ALLEGATO usa in id_atto_sipo. ⚠️ L’incrocio '
     'non è però immediato: la notifica identifica l’atto con l’identificativo nazionale, e per '
     'risalire da quello all’atto di SIPO serve la tavola di corrispondenza ANSC_XREF. È il '
     'legame che rende possibile l’arricchimento, e senza il quale la notifica resta un '
     'documento senza padrone.'),

    ('cd_operatore → l’ufficiale che ha redatto l’atto.',
     'Side lo scrive in tabella perché non sempre riesce a dedurlo dalle comunicazioni, e serve '
     'a filtrare le notifiche per operatore: è il modo in cui l’ufficio se le distribuisce. In '
     'SIPO l’informazione esiste sull’atto e nello store di stato. ⚠️ Va però deciso se alla '
     'scala di Roma il filtro giusto sia l’operatore o il municipio: la seconda dimensione a '
     'Milano non esiste, qui probabilmente conta di più, e conviene prevedere entrambi i '
     'criteri fin dall’inizio.'),

    ('cd_azione_canale → la classificazione, che si eredita come regola e non come dato.',
     'I sette ambiti valgono anche per Roma, ma la loro attribuzione dipende da come il Comune '
     'organizza archivio e posta certificata. La matrice degli scenari è il punto di partenza; '
     'la sua validazione è un lavoro con l’organizzazione, non un lavoro di analisi.'),

    ('id_dati_comunicazione → una tabella che SIPO non ha.',
     'Come detto sopra: la colonna si conserva per fedeltà alla struttura e resta non '
     'valorizzata finché il Comune non decide se ereditare anche il processo che predispone le '
     'comunicazioni alla firma. La questione non riguarda solo le notifiche: è la stessa che il '
     'registro degli Open Point già annota a proposito dell’invio delle comunicazioni agli enti '
     'destinatari.'),

    ('id_comune_destinatario e le colonne di comune → l’instradamento.',
     'Servono a sapere a chi inviare PEC e comunicazioni, e a comporre la descrizione del '
     'destinatario nella forma «comune (provincia)». In SIPO i comuni risiedono in ANAG_USR e '
     'si raggiungono per sinonimo, come già avviene per le altre decodifiche territoriali: le '
     'colonne si valorizzano, il catalogo non si duplica.'),

    ('Il codice del Comune.',
     'Il confronto «la notifica riguarda un atto formato da noi oppure da altri» è la prima '
     'discriminante di tutto il processo e nella sorgente è scritto con il codice di Milano. '
     'Va parametrizzato, non sostituito con un’altra costante: il componente serve un solo '
     'comune, ma il codice appartiene alla configurazione e non al programma.'),

    ('Lo store delle notifiche già previsto.',
     'Il presente documento prevede già una tabella di replica delle notifiche. ⚠️ Le due '
     'strutture vanno riconciliate in un solo intervento, quando si deciderà se adottare i nomi '
     'di Side anche qui — come si è fatto per gli allegati e per i dizionari — o conservare la '
     'nomenclatura del Comune. La struttura di Side è più ricca e va presa come riferimento del '
     'contenuto; la decisione sui nomi è indipendente e va presa una volta sola per tutte e tre '
     'le eredità.'),
]

APPLICATIVO = [
    'Il committente osserva che questa parte potrebbe essere un applicativo distinto, '
    'richiamabile dal back-office o direttamente da SIPO. L’osservazione ha fondamento, e le '
    'ragioni sono le stesse che hanno portato a separare la parte dizionari dal concentratore.',

    'Il flusso in ingresso ha un ciclo di vita proprio: si alimenta a intervalli invece che su '
    'comando dell’operatore, lavora su un orizzonte di giorni e non di secondi, e la sua '
    'indisponibilità non impedisce di formare atti — impedisce di lavorare le notifiche, che è '
    'cosa diversa e meno urgente. Tenerlo dentro il concentratore legherebbe due domini di '
    'guasto che non hanno ragione di stare insieme.',

    'Ha inoltre consumatori diversi: le notifiche interessano l’ufficiale che ha redatto '
    'l’atto, l’archivista che tiene gli atti cartacei e chi spedisce le PEC — attori che nel '
    'percorso di formazione dell’atto non compaiono. La stessa Milano li colloca su console '
    'distinte, quella dell’archivista e quella dell’ufficiale.',

    '⚠️ Restano però due vincoli che un applicativo separato non può ignorare, e sono gli '
    'stessi della parte dizionari. **Il primo**: la chiave privata del certificato e la '
    'costruzione dei token restano in un solo luogo, quindi il nuovo componente raggiunge ANSC '
    'per il tramite del concentratore e non direttamente. **Il secondo**: la conferma di una '
    'notifica è un atto di volontà dell’ufficiale, non un’operazione di sistema, e va compiuta '
    'entro una sessione autenticata — il che vale anche se la lettura fosse automatizzabile. '
    'Separare il componente è quindi legittimo; separare la responsabilità no.',

    'La proposta è di trattarlo come **una terza unità di deployment a sé**, allo stesso titolo '
    'del componente dei dizionari: alimentata dal proprio batch, con le proprie schermate, '
    'richiamabile dal back-office e dai pulsanti di evento di SIPO, ma priva di credenziali '
    'verso ANSC. Il conto delle unità di deployment va aggiornato di conseguenza quando la '
    'decisione sarà confermata.',
]

APERTI = [
    'Il recepimento lascia aperte tre questioni, che si aggiungono al registro.',
]

OP_NUOVI = [
    ('OP-53', 'Comunicazioni predisposte alla firma',
     'Il disegno di Side lega ogni notifica alla comunicazione — PEC o comunicazione per '
     'l’archivio — predisposta al momento della firma dell’atto. SIPO non ha né la tabella né '
     'il processo che la popola, né il catalogo dei modelli di stampa. Va deciso se ereditare '
     'anche quel processo o dichiarare la colonna non valorizzata nel pilota. Si lega a OP-39.',
     'Aperto', 'Analisi / Cliente', 'Alta'),
    ('OP-54', 'Attribuzione dell’ambito alle notifiche',
     'I sette codici di ambito che classificano le notifiche presuppongono un’organizzazione '
     'dell’archivio cartaceo e dell’invio delle PEC. Chi è l’archivista alla scala di Roma, '
     'quale processo spedisce le PEC e su quali console vadano mostrate le notifiche sono '
     'scelte dell’organizzazione, non dell’analisi.',
     'Aperto', 'Cliente', 'Alta'),
    ('OP-55', 'Incongruenze rilevate sul contrassegno di evento cartaceo',
     'La sorgente segnala notifiche di proposta di annotazione con contrassegno di evento '
     'cartaceo il cui atto primario risulta invece digitale, e si interroga sulla struttura '
     'dell’identificativo delle notifiche automatiche, il cui numero comunale vale zero. '
     'Entrambe le cose vanno chiarite con il fornitore prima di costruire la classificazione '
     'automatica.',
     'Aperto', 'Analisi / Fornitore ANSC', 'Media'),
]


# ------------------------------------------------- decodifiche (cap. dizionari)

DEC = [
    'Il committente ha trasmesso, insieme ai documenti sulle notifiche, la descrizione del '
    'processo con cui il fornitore intende mantenere allineate le decodifiche. La si riporta '
    'qui perché conferma e precisa il comando descritto sopra.',

    'Le chiamate al servizio delle decodifiche sono due, in sequenza. La prima, sull’elenco, '
    'restituisce la lista dei domini pubblicati e aggiorna la tabella del catalogo. La seconda, '
    'sul dettaglio, si ripete per ciascun dominio elencato e restituisce i valori, che '
    'aggiornano la tabella dei valori.',

    'L’aggiornamento non sovrascrive: confronta con quanto già presente e distingue tre casi.',
]
DEC_V = [
    ('Record nuovi.', 'Si inseriscono.'),
    ('Record modificati.', 'Si aggiornano.'),
    ('Record non più presenti nella risposta.',
     'Non si cancellano: si valorizza la data di fine validità. ⚠️ La nota del fornitore '
     'osserva che è un caso «che non dovrebbe verificarsi», ed è corretto — ANSC chiude la '
     'validità invece di rimuovere — ma trattarlo è comunque necessario: un valore che sparisse '
     'senza essere chiuso resterebbe altrimenti valido per sempre.'),
]
DEC_CODA = [
    'Il processo coincide con il comando descritto in questo capitolo e ne conferma l’impianto: '
    'confronto e non ricarico, chiusura della validità e non cancellazione, sostituzione per '
    'singola decodifica. ⚠️ Una differenza va però segnalata: la nota descrive il processo come '
    'sequenza automatica, mentre qui si è scelto il comando manuale, per le ragioni dette sopra '
    'e perché la chiamata passa dal concentratore. Le due cose sono compatibili — è la stessa '
    'logica, avviata diversamente — ma la scelta va confermata.',
]

DEC_SERV = (
    'La nota del fornitore prevede infine un servizio di SIPO che interroga la tabella dei '
    'valori e restituisce i record di un dominio dato il suo identificativo. ⚠️ Il servizio si '
    'affianca alla vista, non la sostituisce: la vista serve alle maschere che leggono una '
    'decodifica per popolare una tendina — dove passare da una chiamata applicativa sarebbe un '
    'costo senza contropartita — mentre il servizio serve a chi ha bisogno dei valori senza '
    'accedere alla base dati, tipicamente il front-end e il back-office. La scelta di esporre '
    'la base dati come contratto verso SIPO resta quella motivata sopra; il servizio è un '
    'canale aggiuntivo e va aggiunto al contratto del componente dei dizionari.'
)


# ----------------------------------------------------------------- montaggio

def blocchi_capitolo():
    """La sequenza del capitolo: ('h',liv,testo) ('p',testo) ('v',testa,corpo) ('t',righe,larg)."""
    b = [('h', 1, 'La gestione delle notifiche: il processo ereditato dal sistema Side')]
    b += [('p', t) for t in APERTURA]

    b += [('h', 2, 'Due realizzazioni, non una')]
    b += [('p', t) for t in DUE]
    b += [('v', a, c) for a, c in DUE_V]
    b += [('p', DUE_CODA)]

    b += [('h', 2, 'Le attività del processo')]
    b += [('p', 'Il processo si articola in cinque attività.')]
    b += [('v', a, c) for a, c in ATTIVITA]

    b += [('h', 2, 'Il batch di scarico')]
    b += [('p', t) for t in BATCH]
    b += [('t', righe_tabella_docx(5, ['Parametro', 'Valori e descrizione']), [1.7, 4.6])]
    b += [('p', t) for t in BATCH_CODA]

    b += [('h', 2, 'La struttura della notifica')]
    b += [('p', t) for t in STRUTTURA]
    b += [('t', righe_struttura(), [1.6, 2.6, 2.1])]

    b += [('h', 2, 'I domini')]
    b += [('p', t) for t in DOMINI]
    b += [('p', 'Tipo evento (ANSC_1)')]
    b += [('t', righe_tabella_docx(6), [1.2, 5.1])]
    b += [('p', 'Tipo contenuto (ANSC_2)')]
    b += [('t', righe_tabella_docx(7), [1.2, 5.1])]
    b += [('p', 'Genere della notifica (ANSC_101)')]
    b += [('t', righe_tabella_docx(8), [1.2, 5.1])]
    b += [('p', 'Stato del flusso della notifica (ANSC_102)')]
    b += [('t', righe_tabella_docx(9), [1.2, 5.1])]
    b += [('p', 'Ambito (decodifica locale, in Side SIDe_30)')]
    b += [('t', righe_tabella_docx(10), [1.5, 4.8])]
    b += [('p', t) for t in DOMINI_CODA]

    b += [('h', 2, 'Gli scenari')]
    b += [('p', t) for t in SCENARI]
    b += [('t', righe_foglio('Scenari', [0, 1, 2, 3, 5],
                             ['Genere', 'Canale', 'Evento', 'Scenario', 'Ambito']),
           [1.0, 0.8, 0.8, 2.4, 1.3])]
    b += [('p', t) for t in SCENARI_CODA]

    b += [('h', 2, 'Le comunicazioni predisposte alla firma')]
    b += [('p', t) for t in COMUNICAZIONI]
    b += [('t', righe_foglio('Mapping PEC e ARCHIVIO', [0, 3, 4, 5],
                             ['Campo', 'Valore per la PEC', 'Valore per l’archivio', 'Note']),
           [1.5, 1.9, 1.9, 1.0])]
    b += [('p', t) for t in COMUNICAZIONI_CODA]

    b += [('h', 2, 'Le due pagine di consultazione')]
    b += [('p', t) for t in PAGINE]
    b += [('v', a, c) for a, c in PAGINE_V]
    b += [('p', t) for t in PAGINE_CODA]

    b += [('h', 2, 'Le ipotesi di raccordo con gli applicativi esistenti')]
    b += [('p', 'Il processo di Milano presuppone strutture e informazioni che in SIPO hanno '
                'un altro nome, o non esistono. Le ipotesi seguenti sono ciò che serve per '
                'colmare la distanza, e ciascuna dichiara quanto è verificato e quanto resta '
                'da decidere.')]
    b += [('v', a, c) for a, c in IPOTESI]

    b += [('h', 2, 'Un applicativo separato?')]
    b += [('p', t) for t in APPLICATIVO]

    b += [('h', 2, 'Che cosa resta aperto')]
    b += [('p', t) for t in APERTI]
    b += [('v', f'{c}.', t + ' ' + q.split('.')[0] + '.') for c, t, q, _, _, _ in OP_NUOVI]
    return b


def scrivi(doc, ancora, blocchi, modello):
    for blocco in blocchi:
        if blocco[0] == 'h':
            _, liv, testo = blocco
            D.para(doc, ancora, testo, stile=f'Heading {liv}')
        elif blocco[0] == 'p':
            testo = blocco[1]
            if '**' in testo:
                D.para(doc, ancora, pezzi=D.segmenta('', testo))
            else:
                D.para(doc, ancora, testo)
        elif blocco[0] == 'v':
            D.voce(doc, ancora, blocco[1] + ' ', blocco[2])
        elif blocco[0] == 't':
            D.tabella(doc, ancora, blocco[1], modello, blocco[2])


def main():
    doc = docx.Document(DOC)
    modello = doc.tables[97]

    # --- capitolo nuovo, prima di «SIPO vs la web app ANSC»
    ancora = None
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1' and p.text.strip().startswith('SIPO vs la web app'):
            ancora = p._p
            break
    assert ancora is not None, 'ancora del capitolo non trovata'
    scrivi(doc, ancora, blocchi_capitolo(), modello)

    # --- decodifiche: nuova sezione nel capitolo dizionari, prima di «Fruizione da parte di SIPO»
    anc2 = None
    for p in doc.paragraphs:
        if p.style.name == 'Heading 2' and p.text.strip().startswith('Fruizione da parte di SIPO'):
            anc2 = p._p
            break
    assert anc2 is not None, 'ancora della sezione dizionari non trovata'
    blocchi = [('h', 2, 'Il processo di allineamento indicato dal fornitore')]
    blocchi += [('p', t) for t in DEC]
    blocchi += [('v', a, c) for a, c in DEC_V]
    blocchi += [('p', t) for t in DEC_CODA]
    scrivi(doc, anc2, blocchi, modello)

    # il servizio di interrogazione: in coda alla sezione «Fruizione da parte di SIPO»
    coda = D.sezione(doc, 'Fruizione da parte di SIPO', livello=2)
    ultimo = coda[-1]
    p = D.para(doc, ultimo, DEC_SERV)
    ultimo.addnext(p._p)

    # --- open point
    t_op = doc.tables[97]
    for riga in OP_NUOVI:
        D.clona_riga(t_op, riga)

    D.storia(doc, '14/09/2026', '3.20',
             'Nuovo capitolo «La gestione delle notifiche»; capitolo «Gestione dei dizionari '
             'ANSC»; Registro degli Open Point',
             'Recepito il processo di gestione delle notifiche del sistema Side del Comune di '
             'Milano (batch di scarico, struttura della notifica, domini, scenari, '
             'comunicazioni predisposte alla firma, pagine di consultazione) con le ipotesi di '
             'raccordo con SIPO e la valutazione dell’ipotesi di un applicativo separato; '
             'recepito il processo di allineamento delle decodifiche indicato dal fornitore e '
             'il servizio di interrogazione dei valori di dominio; aggiunti OP-53, OP-54, OP-55.')

    doc.save(DOC)
    print('salvato:', DOC)
    print('capitoli / tabelle / immagini:', D.riepilogo(DOC))


if __name__ == '__main__':
    main()
