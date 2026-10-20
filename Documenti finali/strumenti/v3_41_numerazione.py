# -*- coding: utf-8 -*-
"""ANALISI_Integrazione-ANSC v3.40 → v3.41 (15/10/2026).

⚠️ Assunto dell'utente (15/10/2026): il numero comunale è di SEI cifre e NON si considera il
registro. Quindi:
  * una sola serie per anno, unica per tutti i registri: SEQ_NUM_COMUNALE_<anno>, MAXVALUE
    999999 (era una per registro e per anno, MAXVALUE 99999);
  * ty_evento esce da DATI_ATTO_ESTERNO (era nostra, serviva all'unicità della serie) e da
    PRENOTAZIONE_ATTI (era di Side: adattamento dichiarato, la colonna non avrebbe funzione);
  * vincolo di unicità della serie su (nr_anno_atto, nr_atto_comunale); nr_atto_comunale
    NUMBER(6);
  * la procedura del job crea UNA sequence; lo stacco compone il nome dal solo anno;
  * contratto ansc-numerazione-v1.yaml 1.1.0 (v3_41_contratto_numerazione.py) e SPEC_API v0.3
    → Appendice B e [R8] allineati.
⚠️ CONFLITTO CON LE FONTI, dichiarato e non nascosto: model_evento.yaml:149 e
R005_consultazione_ansc.yaml:543 hanno `numeroatto` maxLength 5; il Changelog ANSC riporta il
controllo «ammessi massimo 5 caratteri». La v3.40 (revisione dell'utente) scriveva che le fonti
dichiarano sei cifre: non è così → assunto del Comune + [DA VERIFICARE] + OP-73 (a).

Inoltre: corretto il refuso nella colonna virtuale DECODIFICA di ANSC_CFG_DOMINIO_SIPO
(`|| TABELLA_SIPO || 'ANSC'` generava «CONF_STATO_CIVILEANSC61»), ERD rigenerato (sequence
senza registro e ANSC_CFG_DOMINIO_SIPO senza CAMPO_SIPO, come nella CREATE rivista dall'utente).

    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_41_contratto_numerazione.py
    /Library/Developer/CommandLineTools/usr/bin/python3 genera-spec-api.py   (dalla radice)
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/diagrammi_erd.py
    /Library/Developer/CommandLineTools/usr/bin/python3 strumenti/v3_41_numerazione.py
"""
import os
import shutil
import subprocess
import sys

import docx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_comune as D   # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIN = os.path.join(BASE, 'Documenti finali')
SRC = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.40.docx')
DST = os.path.join(FIN, 'ANALISI_Integrazione-ANSC_v3.41.docx')
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'img')


def commenti(p):
    return int(subprocess.run(
        ['bash', '-c', "unzip -p %s word/comments.xml 2>/dev/null | "
                       "grep -o '<w:comment ' | wc -l" % repr(p)],
        capture_output=True, text=True).stdout.strip() or 0)


def par_esatto(d, testo):
    """L'unico paragrafo del corpo il cui testo è esattamente `testo` (spazi in coda esclusi)."""
    trovati = [p for p in d.paragraphs if p.text.rstrip() == testo.rstrip()]
    if len(trovati) != 1:
        raise SystemExit(f'«{testo[:60]}»: attesi 1 paragrafo, trovati {len(trovati)}')
    return trovati[0]


def riscrivi(d, vecchio, nuovo):
    D.testo_di(par_esatto(d, vecchio), nuovo)


def elimina(d, testo):
    p = par_esatto(d, testo)
    if D.ha_commenti(p._p):
        raise SystemExit(f'«{testo[:60]}» porta un commento: non si elimina')
    p._p.getparent().remove(p._p)


def riga_tabella(d, intestazioni, prima_cella):
    t = D.trova_tabella(d, *intestazioni)
    righe = [r for r in t.rows if r.cells[0].text.strip() == prima_cella]
    if len(righe) != 1:
        raise SystemExit(f'riga «{prima_cella}»: attesa 1, trovate {len(righe)}')
    return t, righe[0]


# ─────────────────────────────────────────────── prosa del paragrafo «La numerazione comunale»
PROSA = [
    ('ANSC fissa il numero comunale ad al massimo sei cifre: lo dichiarano il modello evento e il '
     'servizio di consultazione, e la nota ANSC sulla nuova numerazione degli atti digitali. La '
     'serie arriva quindi a 999.999 atti l’anno. Roma ne forma circa 140.000 secondo la '
     'ricognizione dei tipi atto.',
     'Il numero comunale si assume di sei cifre al massimo: la serie arriva quindi a 999.999 atti '
     'l’anno, e Roma ne forma circa 140.000 secondo la ricognizione dei tipi atto. Una serie sola '
     'basta con ampio margine. È un assunto del Comune che le fonti di ANSC non confermano: il '
     'modello evento e il servizio di consultazione dichiarano per il numero dell’atto una '
     'lunghezza massima di cinque caratteri, e il registro delle modifiche di ANSC riporta il '
     'controllo che ne ammette al massimo cinque. [DA VERIFICARE: che ANSC accetti un numero '
     'comunale di sei cifre (OP-73).]'),
    ('Ne discende una sequence per anno: cinque l’anno. Il nome segue lo standard di nomenclatura '
     '[R5]: SEQ_NUM_COMUNALE_<anno>, per esempio SEQ_NUM_COMUNALE_2027. [DA VERIFICARE: che ANSC si '
     'attenda il progressivo comunale per registro; le fonti lo implicano ma non lo dichiarano '
     '(OP-73).]',
     'Ne discende una sola sequence per anno, unica per tutti i registri: il progressivo non '
     'distingue nascite, morti, matrimoni, unioni civili e cittadinanze, e ogni anno ce n’è uno '
     'soltanto. Il nome segue lo standard di nomenclatura [R5]: SEQ_NUM_COMUNALE_<anno>, per '
     'esempio SEQ_NUM_COMUNALE_2027. È uno scostamento da Side, che organizza le richieste per '
     'registro e anno. [DA VERIFICARE: che ANSC accetti un progressivo comunale unico fra i '
     'registri; le fonti non lo dichiarano né lo escludono (OP-73).]'),
    ('sei cifre come limite del database. MAXVALUE 999999: superarlo fa fallire lo stacco invece di '
     'produrre un numero che ANSC rifiuterebbe.',
     'Sei cifre come limite del database. MAXVALUE 999999: superarlo fa fallire lo stacco invece '
     'di produrre un numero di sette cifre.'),
    ('Si creano in anticipo. Un job crea il 1° dicembre le sequence dell’anno successivo, e la '
     'procedura che le crea si esegue anche all’installazione per l’anno in corso. Lo stacco non '
     'crea mai una sequence: se manca, fallisce, ed è un guasto visibile invece di una struttura '
     'nata a metà giornata.',
     'Si crea in anticipo. Un job crea il 1° dicembre la sequence dell’anno successivo, e la '
     'procedura che la crea si esegue anche all’installazione per l’anno in corso. Lo stacco non '
     'crea mai una sequence: se manca, fallisce, ed è un guasto visibile invece di una struttura '
     'nata a metà giornata.'),
    ('Il nome della sequence lo compone il servizio della numerazione da due valori controllati — '
     'il registro, che deve appartenere alla decodifica ANSC_01, e l’anno, di quattro cifre — e mai '
     'da un testo ricevuto: è un nome che finisce in un’istruzione SQL.',
     'Il nome della sequence lo compone il servizio della numerazione da un solo valore '
     'controllato — l’anno, di quattro cifre — e mai da un testo ricevuto: è un nome che finisce '
     'in un’istruzione SQL.'),
    ('Controllo dei valori. Il registro deve appartenere alla decodifica ANSC_01; l’anno deve '
     'essere il corrente o il precedente. Il nome della sequence si compone solo da questi due '
     'valori controllati.',
     'Controllo dei valori. L’anno deve essere il corrente o il precedente. Il nome della sequence '
     'si compone solo da questo valore controllato.'),
    # SQL dello stacco
    (' WHERE ty_evento = :registro AND nr_anno_atto = :anno', ' WHERE nr_anno_atto = :anno'),
    ('INSERT INTO ANSC_USR.PRENOTAZIONE_ATTI (ty_evento, nr_anno_atto,',
     'INSERT INTO ANSC_USR.PRENOTAZIONE_ATTI (nr_anno_atto,'),
    ('VALUES (:registro, :anno, :quantita, :operatore)', 'VALUES (:anno, :quantita, :operatore)'),
    ('SELECT ANSC_USR.SEQ_NUM_COMUNALE_<registro>_<anno>.NEXTVAL FROM DUAL;',
     'SELECT ANSC_USR.SEQ_NUM_COMUNALE_<anno>.NEXTVAL FROM DUAL;'),
    ('       ty_evento, nr_anno_atto, id_atto_sipo)', '       nr_anno_atto, id_atto_sipo)'),
    ('VALUES (:id_prenotazione, :numero, :registro, :anno, :atto);',
     'VALUES (:id_prenotazione, :numero, :anno, :atto);'),
    ("   AND SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_<registro>_<anno>';",
     "   AND SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_<anno>';"),
    # ─────────────────────────────────────────────── DDL dell'Appendice A
    ('  nr_atto_comunale          NUMBER(5)          NOT NULL,',
     '  nr_atto_comunale          NUMBER(6)          NOT NULL,'),
    ('    UNIQUE (ty_evento, nr_anno_atto, nr_atto_comunale),',
     '    UNIQUE (nr_anno_atto, nr_atto_comunale),'),
    ("    CHECK (cd_stato_numero_prenotato IN ('1','2','3','4')),",
     "    CHECK (cd_stato_numero_prenotato IN ('1','2','3','4'))"),
    ('  ON ANSC_USR.PRENOTAZIONE_ATTI (ty_evento, nr_anno_atto, ts_prenotazione)',
     '  ON ANSC_USR.PRENOTAZIONE_ATTI (nr_anno_atto, ts_prenotazione)'),
    ("  'Richieste di numeri comunali, per registro (ANSC_01) e anno. Struttura di Side.';",
     "  'Richieste di numeri comunali, per anno. Struttura di Side, senza il registro.';"),
    ('-- Una sequence per registro (ANSC_01) e per anno. Modello, per le morti del 2027:',
     '-- Una sequence per anno, unica per tutti i registri. Modello, per il 2027:'),
    ('CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_2_2027', 'CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_2027'),
    ('  START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 99999 NOCYCLE NOCACHE ORDER;',
     '  START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 999999 NOCYCLE NOCACHE ORDER;'),
    ("     WHERE SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_' || r || '_' || p_anno;",
     "   WHERE SEQUENCE_NAME = 'SEQ_NUM_COMUNALE_' || p_anno;"),
    ("    SELECT COUNT(*) INTO n FROM USER_SEQUENCES",
     "  SELECT COUNT(*) INTO n FROM USER_SEQUENCES"),
    ("    IF n = 0 THEN", "  IF n = 0 THEN"),
    ("      EXECUTE IMMEDIATE 'CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_' || r || '_' ||",
     "    EXECUTE IMMEDIATE 'CREATE SEQUENCE ANSC_USR.SEQ_NUM_COMUNALE_' ||"),
    ("        p_anno || ' START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 99999'",
     "      p_anno || ' START WITH 1 INCREMENT BY 1 MINVALUE 1 MAXVALUE 999999'"),
    ("        || ' NOCYCLE NOCACHE ORDER';", "      || ' NOCYCLE NOCACHE ORDER';"),
    ("      EXECUTE IMMEDIATE 'GRANT SELECT ON ANSC_USR.SEQ_NUM_COMUNALE_' || r || '_' ||",
     "    EXECUTE IMMEDIATE 'GRANT SELECT ON ANSC_USR.SEQ_NUM_COMUNALE_' ||"),
    ("        p_anno || ' TO <UTENZA_NUM>';", "      p_anno || ' TO <UTENZA_NUM>';"),
    ("    END IF;", "  END IF;"),
    # refuso della v3.40 nella colonna virtuale di ANSC_CFG_DOMINIO_SIPO
    ("                         || TABELLA_SIPO || 'ANSC'",
     "                         || TABELLA_SIPO || '.ANSC_'"),
    # Appendice B: sommario del contratto della numerazione
    ('Stacco dei numeri comunali dalla serie per registro e anno e governo della loro sorte.',
     'Stacco dei numeri comunali dalla serie annuale e governo della loro sorte.'),
]

DA_ELIMINARE = [
    '  ty_evento                 VARCHAR2(10 CHAR)  NOT NULL,',          # due: le due tabelle
    "  CONSTRAINT prenotazione_atti_reg_ck CHECK (ty_evento IN ('1','2','3','4','5')),",
    "  CONSTRAINT dati_atto_esterno_reg_ck CHECK (ty_evento IN ('1','2','3','4','5'))",
    '  FOR r IN 1 .. 5 LOOP   -- i cinque registri di ANSC_01',
    '  END LOOP;',
]

# sostituzioni parziali, anche nelle celle, con il numero di paragrafi attesi
PARZIALI = [
    ('Ereditata da Side con i nomi di origine: registro, anno, quanti numeri, chi e quando.',
     'Ereditata da Side con i nomi di origine, senza il registro: anno, quanti numeri, chi e '
     'quando.', 1),
    ('Una sequence per anno, creata in anticipo da un job',
     'Una sequence per anno, unica per tutti i registri, creata in anticipo da un job', 1),
    ('Due richieste contemporanee sullo stesso registro si intercalano',
     'Due richieste contemporanee sulla stessa serie si intercalano', 1),
    ('Dopo, al ritorno di SIPO, indicando registro e anno:',
     'Dopo, al ritorno di SIPO, indicando l’anno:', 1),
    ('una pagina di richiesta, che propone registro e anno correnti e un numero',
     'una pagina di richiesta, che propone l’anno corrente e un numero', 1),
    ('Inserisce la riga di PRENOTAZIONE_ATTI con registro, anno, quantità, operatore e istante.',
     'Inserisce la riga di PRENOTAZIONE_ATTI con anno, quantità, operatore e istante.', 1),
    ('(a) Che ANSC si attenda il progressivo comunale per registro: le fonti lo implicano — cinque '
     'cifre al massimo, chiave di Side per registro e anno — ma non lo dichiarano.',
     '(a) Che ANSC accetti il numero comunale di sei cifre e un progressivo unico fra i registri, '
     'come assume il Comune: il modello evento e il servizio di consultazione dichiarano per il '
     'numero dell’atto al massimo cinque caratteri, e il registro delle modifiche di ANSC riporta '
     'il controllo relativo. Se l’assunto non fosse confermato, si torna alla serie per registro '
     'e anno a cinque cifre descritta nella v3.40.', 1),
    ('Crea le serie dei cinque registri per un anno', 'Crea la serie di un anno', 1),
    ('/ansc/v1/numerazione/serie/{codRegistro}/{numAnno}', '/ansc/v1/numerazione/serie/{numAnno}', 5),
    ('SPEC_API_Integrazione-ANSC_v0.2.docx', 'SPEC_API_Integrazione-ANSC_v0.3.docx', 1),
    ('Stacca i numeri comunali dalla serie per registro e anno, serializzando gli stacchi',
     'Stacca i numeri comunali dalla serie annuale, unica per tutti i registri, serializzando '
     'gli stacchi', 1),
    ('Chiuso nella v3.38. La serie è una sequence per registro e per anno, che riparte ogni anno;',
     'Chiuso nella v3.38. La serie è una sequence per anno, che riparte ogni anno — dalla v3.41 '
     'unica per tutti i registri;', 1),
]

ADATTAMENTO = ('Rispetto a Side manca una colonna, ty_evento, il registro della richiesta: in Side '
               'individua la serie, qui la serie è una sola per anno e la colonna non avrebbe più '
               'funzione. Il registro dell’atto resta dove già si trova, sull’atto di SIPO.')

STORIA = (
    '15/10/2026', '3.41',
    'SIPO e la web app («La numerazione comunale») · Appendice A (DDL della numerazione, '
    'ANSC_CFG_DOMINIO_SIPO) · Appendice B · Modello dati (ERD) · Riferimenti · Open Point',
    'La numerazione comunale si allinea all’assunto del Comune: numero di sei cifre e una sola '
    'serie per anno, unica per tutti i registri. Una sequence per anno invece di cinque, '
    'MAXVALUE 999999; il registro esce dalle tabelle ereditate da Side (adattamento dichiarato), '
    'dal vincolo di unicità della serie, dallo stacco e dal contratto del servizio, che passa '
    'alla 1.1.0 e rigenera la specifica delle API [R8] alla v0.3. L’assunto contrasta con le '
    'fonti di ANSC, che limitano il numero dell’atto a cinque caratteri: è dichiarato come da '
    'verificare e OP-73 (a) è riformulato. Corretto un refuso nella colonna virtuale di '
    'ANSC_CFG_DOMINIO_SIPO, che non separava tabella e dominio; ERD rigenerato.')


def main():
    if os.path.exists(DST):
        raise SystemExit('esiste già: ' + DST)
    shutil.copy(SRC, DST)
    prima = commenti(DST)
    d = docx.Document(DST)

    for vecchio, nuovo in PROSA:
        riscrivi(d, vecchio, nuovo)
    # i due ty_evento della DDL hanno testo identico: si eliminano uno alla volta
    for testo in DA_ELIMINARE:
        n = 2 if testo.startswith('  ty_evento') else 1
        for _ in range(n):
            if testo.startswith('  ty_evento'):
                p = [q for q in d.paragraphs if q.text.rstrip() == testo.rstrip()][0]
                p._p.getparent().remove(p._p)
            else:
                elimina(d, testo)
    if any(q.text.rstrip() == DA_ELIMINARE[0].rstrip() for q in d.paragraphs):
        raise SystemExit('ty_evento residuo nella DDL')

    fatti = []
    for vecchio, nuovo, attese in PARZIALI:
        D.sostituisci(d, vecchio, nuovo, attese=attese, fatti=fatti)

    # PRENOTAZIONE_ATTI: via la riga ty_evento
    for t in d.tables:
        intest = [c.text.strip() for c in t.rows[0].cells]
        if intest != ['Colonna', 'Tipo', 'Note']:
            continue
        prime = [row.cells[0].text.strip() for row in t.rows]
        if 'nr_id_ansc_da_prenotare' in prime and 'ty_evento' in prime:
            row = t.rows[prime.index('ty_evento')]
            if D.ha_commenti(row._tr):
                raise SystemExit('la riga ty_evento porta un commento')
            row._tr.getparent().remove(row._tr)
            fatti.append('PRENOTAZIONE_ATTI: tolta la riga ty_evento')
        if 'nr_atto_comunale' in prime and 'ty_evento / nr_anno_atto' in prime:
            row = t.rows[prime.index('nr_atto_comunale')]
            D.riscrivi_cella(row.cells[1], 'NUMBER(6) (PK)')
            D.riscrivi_cella(row.cells[2], 'Il numero. Sei cifre al massimo, secondo l’assunto '
                                           'del Comune (OP-73): la sequence non può superarle.')
            row = t.rows[prime.index('ty_evento / nr_anno_atto')]
            D.riscrivi_cella(row.cells[0], 'nr_anno_atto')
            D.riscrivi_cella(row.cells[1], 'NUMBER(4)')
            D.riscrivi_cella(row.cells[2], 'Aggiunta rispetto a Side: l’anno ripetuto dalla '
                                           'richiesta, perché il database possa garantire che '
                                           'un numero non compaia due volte nella stessa serie '
                                           '(vincolo di unicità su anno e numero). In Side '
                                           'l’unicità discende solo dal numeratore.')
            fatti.append('DATI_ATTO_ESTERNO: NUMBER(6), unicità su anno e numero')
    if len(fatti) < len(PARZIALI) + 2:
        raise SystemExit('tabelle della numerazione non trovate: ' + '; '.join(fatti))

    # adattamento dichiarato, dopo la frase che introduce le due tabelle
    ancora = par_esatto(d, 'Due tabelle, con i nomi di origine. PRENOTAZIONE_ATTI registra ogni '
                           'richiesta; DATI_ATTO_ESTERNO registra ogni numero staccato e ciò che ne '
                           'è stato. Side dichiara campi e significato ma non i tipi, che sono '
                           'scelti qui.')
    D.para(d, ancora._p.getnext(), ADATTAMENTO)

    D.sostituisci_immagine(d, 'Schema ANSC_USR, con ALLEGATO', os.path.join(IMG, 'erd_ansc_usr.png'))
    D.storia(d, *STORIA)
    for f in D.frontespizio(d, '3.41', data='15/10/2026'):
        fatti.append(f)
    # la testata scrive «Versione 3.40» in una forma che frontespizio() non riconosce
    for s in d.sections:
        for tt in s.header.tables:
            for rr in tt.rows:
                for c in rr.cells:
                    if c.text.strip() == 'Versione 3.40':
                        D.riscrivi_cella(c, 'Versione 3.41')
                        fatti.append('testata: Versione 3.41')
    d.save(DST)

    dopo = commenti(DST)
    for f in fatti:
        print('  ' + f)
    print(f'  commenti di Word conservati: {prima} → {dopo}')
    if dopo != prima:
        raise SystemExit('commenti persi')
    print('  scritto', DST, D.riepilogo(DST))


if __name__ == '__main__':
    main()
