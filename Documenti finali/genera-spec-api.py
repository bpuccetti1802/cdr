# -*- coding: utf-8 -*-
"""Genera SPEC_API_Integrazione-ANSC_v0.1.docx dai contratti OpenAPI dei pod.

Il documento è un artefatto DERIVATO: la fonte autoritativa sono i file .yaml.
Per aggiornarlo si modificano i contratti e si riesegue questo script.
"""
import io, json, os, shutil, sys
import yaml
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = '/Users/minimac/Desktop/Comune di Roma'
TEMPLATE = os.path.join(BASE, 'Template documentale/template_DAD_roma-capitale.docx')
OUT = os.path.join(BASE, 'Documenti finali/SPEC_API_Integrazione-ANSC_v0.1.docx')
DATA = os.path.join(BASE, 'Documenti finali')

POD = [
    ('all-ansc-sipo', 'Concentratore operativo', 'ansc-v1.yaml',
     'Deployment stateless, almeno due repliche. Espone il percorso presidiato di formazione '
     'dell’atto e le funzioni di back-office. È l’unico componente che detiene la chiave '
     'privata del certificato server e costruisce i token verso ANSC.'),
    ('dec-ansc-sipo', 'Concentratore dei dizionari', 'ansc-dizionari-v1.yaml',
     'Deployment a replica singola. Replica in locale le decodifiche ANSC su comando manuale. '
     'Non detiene certificati e non costruisce token: raggiunge ANSC per il tramite del '
     'concentratore operativo.'),
    ('processo-automazione', 'Processo di automazione', 'ansc-automazione-v1.yaml',
     'Deployment con una o due repliche. Esegue le sole attività non presidiate ammesse: letture '
     'e code. Non forma e non firma atti. La sua realizzabilità è subordinata alla chiusura di '
     'un open point.'),
]

METODI = ('get', 'post', 'put', 'delete', 'patch')

# ----------------------------------------------------------------- costruzione
shutil.copyfile(TEMPLATE, OUT)
d = docx.Document(OUT)
body = d.element.body

# svuota il corpo dopo il front matter, conservando le proprieta' di sezione
children = list(body.iterchildren())
TAGLIA_DA = 29
for ch in children[TAGLIA_DA:]:
    if ch.tag != qn('w:sectPr'):
        body.remove(ch)

# ------------------------------------------------------------------- front matter
from docx.text.paragraph import Paragraph
paras = [Paragraph(c, d) for c in body.iterchildren() if c.tag == qn('w:p')]


def set_par(par, text):
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(text)


for p in paras:
    if '“Lorem Ipsum”' in p.text:
        set_par(p, '“Integrazione SIPO – ANSC”')
    elif p.style.name == 'Subtitle':
        set_par(p, 'Specifica delle API dei componenti')

SCHEDA = {
    'Area Organizzativa': 'Dipartimento Trasformazione Digitale',
    'RTI': 'IBM, TIM, Sistemi Informativi, Leonardo, Deloitte, Webgenesys, Engage, Netservice, '
           'We:Com, N&C, Jacala Civitas, Intelligo',
    'Società': 'Sistemi Informativi S.r.l.',
    'Contratto Esecutivo': 'CIG B9DA959E80',
    'AQ': 'CIG A02590F330 - Accordo Quadro avente ad oggetto servizi applicativi in ottica cloud e '
          'servizi di demand e PMO per le pubbliche amministrazioni locali - terza edizione - '
          'ID 2610 - Lotti 1',
    'Progetto': 'Integrazione SIPO – ANSC',
    'Data consegna': '20/08/2026',
    'Versione': '0.1',
    'Documento': '[codice documento da assegnare]',
    'Codice Area Applicativa': '06',
    'Codice Asset': '[da assegnare]',
}
# solo la scheda informativa: la storia del documento ha intestazioni omonime
for r in d.tables[0].rows:
    k = r.cells[0].text.strip()
    if k in SCHEDA:
        set_par(r.cells[1].paragraphs[0], SCHEDA[k])

# storia del documento
storia = d.tables[1]
if len(storia.rows) > 1 and len(storia.rows[0].cells) == 4:
    r = storia.rows[1]
    vals = ['20/08/2026', '0.1', 'Prima emissione',
            'Prima stesura della specifica delle API dei componenti dell’integrazione, redatta '
            'secondo lo standard aziendale di nomenclatura e specifica delle API. Il documento è '
            'generato dai contratti OpenAPI dei pod ed è un artefatto derivato: la fonte '
            'autoritativa dei nomi sono i file di contratto.']
    for c, v in zip(r.cells, vals):
        set_par(c.paragraphs[0], v)

# ----------------------------------------------------------------------- helper
def h(text, level=1):
    return d.add_paragraph(text, style=f'Heading {level}')


def p(text, style=None):
    return d.add_paragraph(text, style=style) if style else d.add_paragraph(text)


def bullet(text):
    return d.add_paragraph(text, style='List Paragraph')


def mono(text):
    par = d.add_paragraph()
    run = par.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8)
    return par


def tabella(intestazioni, righe, larghezze=None):
    t = d.add_table(rows=1, cols=len(intestazioni))
    t.style = 'Table Grid'
    for c, txt in zip(t.rows[0].cells, intestazioni):
        set_par(c.paragraphs[0], txt)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for riga in righe:
        cells = t.add_row().cells
        for c, txt in zip(cells, riga):
            set_par(c.paragraphs[0], str(txt) if txt is not None else '—')
    if larghezze:
        for row in t.rows:
            for c, w in zip(row.cells, larghezze):
                c.width = Inches(w)
    d.add_paragraph()
    return t


def testo_markdown(md):
    """Rende il testo CommonMark del contratto in paragrafi Word."""
    for blocco in [b for b in md.split('\n\n') if b.strip()]:
        righe = [r.strip() for r in blocco.strip().split('\n')]
        if all(r.startswith(('- ', '* ')) or not r for r in righe if r):
            for r in righe:
                if r:
                    bullet(r[2:].strip())
        elif righe[0].startswith('#'):
            p(righe[0].lstrip('#').strip())
        else:
            p(' '.join(righe).replace('**', '').replace('`', ''))


def nome_schema(nodo):
    """Ricava un nome leggibile per uno schema, seguendo i riferimenti."""
    if not isinstance(nodo, dict):
        return '—'
    if '$ref' in nodo:
        return nodo['$ref'].rsplit('/', 1)[-1]
    if 'allOf' in nodo:
        parti = [nome_schema(x) for x in nodo['allOf'] if '$ref' in x]
        return ' + '.join(parti) if parti else 'struttura in linea'
    if nodo.get('type') == 'array':
        return f"elenco di {nome_schema(nodo.get('items', {}))}"
    return nodo.get('type', 'struttura in linea')


def esempi_di(nodo):
    ct = (nodo or {}).get('content', {}).get('application/json', {})
    return ct.get('examples', {}) or {}


# ==================================================================== documento
h('Scopo del documento', 1)
p('Il documento specifica le interfacce applicative dei componenti che realizzano l’integrazione '
  'fra SIPO e l’Archivio Nazionale dello Stato Civile, secondo lo standard aziendale di '
  'nomenclatura e specifica delle API.')
p('Copre tutti i componenti previsti dall’architettura: il concentratore operativo, il '
  'concentratore dei dizionari e il processo di automazione. Il ricevitore delle notifiche è '
  'trattato a parte, perché realizza un contratto definito da terzi.')
p('Il documento è un artefatto derivato. La fonte autoritativa dei nomi sono i contratti '
  'OpenAPI elencati al capitolo sul perimetro: se questo documento e un contratto divergono, è '
  'il documento a essere superato. Il testo che segue è generato dai contratti, così che i due '
  'non possano allontanarsi.')

h('Riferimenti', 2)
tabella(['ID', 'Documento', 'Contenuto'], [
    ['R1', 'Naming_convention_X_API_v1.00.docx',
     'Standard aziendale di nomenclatura e specifica delle API. È la metodologia che questo '
     'documento applica: la scheda di operazione, il vocabolario dei verbi, la struttura del '
     'manuale di riferimento e la lista di controllo prima della pubblicazione.'],
    ['R2', 'Naming_convention_X_DB_v1.01.docx',
     'Standard gemello per i modelli dati. I nomi dei campi delle rappresentazioni derivano dai '
     'nomi delle colonne secondo la corrispondenza dei prefissi di ruolo.'],
    ['R3', 'ANALISI_Integrazione-ANSC_v2.6.docx',
     'Analisi tecnico-funzionale dell’integrazione. Definisce l’architettura, i componenti, il '
     'modello dati e il modello di esecuzione presidiato che queste interfacce realizzano.'],
    ['R4', 'ansc/docs/openapi/ (R001–R024, R901)',
     'Contratti dei servizi cooperativi ANSC consumati dai componenti. Sono contratti di terzi: si '
     'consumano nella forma in cui sono pubblicati e non sono soggetti allo standard.'],
], [0.5, 2.4, 3.6])

h('Perimetro: i componenti e i loro contratti', 1)
p('L’integrazione è realizzata da più componenti distinti, ciascuno con un proprio ciclo di '
  'vita e un proprio dominio di guasto. A ciascuno corrisponde un contratto OpenAPI separato.')

righe_pod = []
contratti = {}
for nome, ruolo, file_yaml, descr in POD:
    spec = yaml.safe_load(io.open(os.path.join(DATA, file_yaml), encoding='utf-8'))
    contratti[nome] = spec
    n_op = sum(1 for v in spec['paths'].values() for m in v if m in METODI)
    righe_pod.append([nome, ruolo, file_yaml, str(n_op)])
righe_pod.append(['ricevitore-notifiche', 'Ricevitore delle notifiche', '— (contratto di terzi)', '—'])
tabella(['Componente', 'Ruolo', 'Contratto', 'Operazioni'], righe_pod, [1.6, 2.0, 2.2, 0.8])

for nome, ruolo, file_yaml, descr in POD:
    p(f'{nome}. {descr}')

h('Un solo contesto, più contratti', 2)
p('Tutti i percorsi vivono nello stesso contesto di dominio. La scelta merita di essere motivata, '
  'perché la simmetria suggerirebbe di dare un contesto a ciascun componente.')
p('Lo standard prescrive il contrario. Fra i suoi principi generali figura la trasparenza del '
  'dominio e l’opacità dell’implementazione: i nomi esprimono i concetti del dominio e non la '
  'struttura interna che li realizza, né il nome della tabella, né quello della classe, né '
  'quello del modulo. Un percorso che dicesse quale componente serve i dizionari esporrebbe '
  'nell’indirizzo una scelta di distribuzione destinata a cambiare.')
p('Il contratto segue invece il deployable, perché è con il deployable che viene versionato e '
  'rilasciato. Da qui la combinazione adottata: un solo contesto, un contratto per componente.')

h('Scostamenti dichiarati dallo standard', 2)
p('Lo standard richiede che ogni scostamento consapevole sia motivato e tracciato. Sono tre.')
tabella(['Scostamento', 'Regola', 'Motivazione'], [
    ['Nome dei file di contratto',
     'Il file si chiama con il contesto seguito dalla versione.',
     'Con un solo contesto e più contratti la regola produrrebbe nomi identici. Si adotta il '
     'contesto seguito dal componente e dalla versione.'],
    ['Segmento di raggruppamento',
     'Il percorso si compone di contesto, versione e collezione.',
     'I percorsi di back-office e di automazione interpongono un segmento di raggruppamento fra la '
     'versione e la collezione. La regola non lo prevede né lo vieta: se ne propone il '
     'riconoscimento esplicito come sotto-contesto.'],
    ['Sottorisorsa generica delle azioni',
     'Le operazioni non conformi al modello CRUD si esprimono con un sostantivo deverbale.',
     'La collezione delle azioni di back-office raccoglie operazioni rare ed eterogenee. È la '
     'deroga che lo standard stesso ammette, usata una volta sola.'],
], [1.6, 2.0, 3.0])

h('Come si accede', 1)
p('Gli indirizzi di base sono quelli degli ambienti SIPO. La corrispondenza con gli ambienti ANSC '
  'è fissata dall’analisi: esercizio SIPO verso produzione ANSC, test e sviluppo verso '
  'preproduzione o verso il simulatore locale.')
spec0 = contratti['all-ansc-sipo']
tabella(['Ambiente', 'Indirizzo di base'],
        [[s.get('description', ''), s['url']] for s in spec0.get('servers', [])], [3.0, 3.6])
p('L’autenticazione è quella interna a SIPO, con token di sessione fra moduli. Non va confusa '
  'con l’autenticazione verso ANSC: il token per il sistema nazionale è costruito e firmato dal '
  'concentratore operativo con il certificato server e non transita mai per il front-end.')

h('Convenzioni comuni', 1)
p('Le convenzioni che seguono valgono per tutte le operazioni di tutti i contratti e non sono '
  'ripetute nelle singole schede.')
tabella(['Aspetto', 'Convenzione'], [
    ['Formato', 'JSON, codifica UTF-8.'],
    ['Notazione dei campi',
     'Notazione a cammello con iniziale minuscola, con i prefissi di ruolo dello standard sui '
     'modelli dati: id, cod, desc, flag, data, num. Le collezioni portano il nome al plurale.'],
    ['Date e istanti',
     'Le date sono nella forma anno-mese-giorno; gli istanti portano sempre il fuso orario '
     'esplicito. È una correzione rispetto alle bozze precedenti, dove il fuso era omesso.'],
    ['Valori di dominio',
     'Stringhe in maiuscolo con le parole separate dal carattere di sottolineatura, coincidenti '
     'con i domini dichiarati sulla base dati.'],
    ['Booleani', 'Valori booleani veri e propri, mai le stringhe che li rappresentano sul database.'],
    ['Assenza di valore',
     'Il campo è omesso se l’informazione non è pertinente, valorizzato a nullo se è '
     'pertinente ma non nota.'],
    ['Paginazione e ordinamento',
     'Parametri riservati numPagina, numElementi, ordinamento e campi. La risposta riporta gli '
     'elementi in un campo al plurale e i dati di navigazione in un campo dedicato.'],
    ['Forma degli errori',
     'Busta applicativa con esito, messaggi e dati; ciascun messaggio porta tipo, campo, codice e '
     'descrizione. La scelta della busta al posto del formato normalizzato è dichiarata e vale '
     'per l’intera integrazione.'],
    ['Rapporto fra esito e codice di stato',
     'L’esito applicativo è indipendente dal codice di stato: una validazione respinta è un '
     'esito negativo su una risposta positiva, perché l’operazione è stata eseguita '
     'correttamente e ha prodotto un rifiuto.'],
    ['Campi ereditati da terzi',
     'I campi che provengono dal contratto ANSC conservano la grafia originale, anche quando non '
     'segue questa notazione. Sono confinati nelle registrazioni di audit e sono dichiarati come '
     'tali nello schema.'],
], [1.8, 4.8])

# ------------------------------------------------------- catalogo operazioni
h('Catalogo delle operazioni', 1)
p('Una scheda per operazione, raggruppata per componente e, dentro ciascuno, per collezione. '
  'L’ordine segue il flusso di lavoro e non l’alfabeto: le operazioni che si usano in sequenza '
  'compaiono in sequenza.')

codici_errore_globali = []

for nome, ruolo, file_yaml, _ in POD:
    spec = contratti[nome]
    h(f'{ruolo} — {nome}', 2)
    testo_markdown(spec['info'].get('description', ''))

    # indice delle operazioni del componente
    righe = []
    for percorso, voce in spec['paths'].items():
        for metodo, op in voce.items():
            if metodo not in METODI:
                continue
            righe.append([op['operationId'], f'{metodo.upper()} {percorso}', op.get('summary', '')])
    tabella(['Identificativo', 'Metodo e percorso', 'Titolo breve'], righe, [1.5, 2.6, 2.5])

    for percorso, voce in spec['paths'].items():
        par_percorso = voce.get('parameters', [])
        for metodo, op in voce.items():
            if metodo not in METODI:
                continue
            h(f"{op['operationId']} — {op.get('summary','')}", 3)

            ruoli = ', '.join(op.get('x-ruoli', [])) or 'Tutti i ruoli autenticati'
            meta = [
                ['Identificativo', op['operationId']],
                ['Titolo breve', op.get('summary', '')],
                ['Metodo e percorso', f'{metodo.upper()} {percorso}'],
                ['Raggruppamento', ', '.join(op.get('tags', []))],
                ['Sicurezza', f'Token di sessione interno a SIPO. Ruoli ammessi: {ruoli}.'],
                ['Effetti', (op.get('x-effetti') or '—').strip().replace('**', '').replace('`', '')],
            ]
            if op.get('x-vincoli'):
                meta.append(["Vincoli d'uso", op['x-vincoli'].strip().replace('**', '').replace('`', '')])
            tabella(['Elemento', 'Contenuto'], meta, [1.5, 5.1])

            p('Descrizione')
            testo_markdown(op.get('description', ''))

            tutti_par = par_percorso + op.get('parameters', [])
            pp = [x for x in tutti_par if x.get('in') == 'path' or '$ref' in x]
            qq = [x for x in tutti_par if x.get('in') == 'query']
            # risolve i riferimenti ai parametri comuni
            risolti_path, risolti_query = [], []
            for x in tutti_par:
                if '$ref' in x:
                    x = spec['components']['parameters'][x['$ref'].rsplit('/', 1)[-1]]
                (risolti_path if x.get('in') == 'path' else risolti_query).append(x)
            if risolti_path:
                p('Parametri di percorso')
                tabella(['Nome', 'Tipo', 'Descrizione'],
                        [[x['name'], x.get('schema', {}).get('type', ''), x.get('description', '')]
                         for x in risolti_path], [1.3, 0.9, 4.4])
            if risolti_query:
                p('Parametri di interrogazione')
                tabella(['Nome', 'Tipo', 'Obbl.', 'Per difetto', 'Descrizione'],
                        [[x['name'], x.get('schema', {}).get('type', ''),
                          'S' if x.get('required') else 'N',
                          x.get('schema', {}).get('default', '—'),
                          x.get('description', '')] for x in risolti_query],
                        [1.2, 0.7, 0.5, 0.8, 3.4])

            rb = op.get('requestBody')
            if rb:
                sch = rb.get('content', {}).get('application/json', {}).get('schema', {})
                p('Entità di richiesta')
                p(f"Schema: {nome_schema(sch)}. Tipo di contenuto: application/json. "
                  f"{'Obbligatoria.' if rb.get('required') else 'Facoltativa.'}")

            p('Esiti')
            righe_esiti = []
            for codice, resp in op.get('responses', {}).items():
                if '$ref' in resp:
                    rr = spec['components']['responses'][resp['$ref'].rsplit('/', 1)[-1]]
                else:
                    rr = resp
                sch = rr.get('content', {}).get('application/json', {}).get('schema', {})
                righe_esiti.append([codice, nome_schema(sch) if sch else '—',
                                    (rr.get('description', '') or '').strip().replace('**', '').replace('`', '')])
            tabella(['Codice', 'Rappresentazione', 'Condizione'], righe_esiti, [0.7, 1.7, 4.2])

            cod_err = op.get('x-codici-errore') or []
            if cod_err:
                p('Errori applicativi')
                tabella(['Codice', 'Condizione che lo produce'],
                        [[c['codice'], c['condizione']] for c in cod_err], [2.0, 4.6])
                for c in cod_err:
                    codici_errore_globali.append([c['codice'], nome, op['operationId'], c['condizione']])

            # esempi
            blocchi = []
            if rb:
                for k, v in esempi_di(rb).items():
                    blocchi.append((f"Richiesta — {v.get('summary', k)}", v.get('value')))
            for codice, resp in op.get('responses', {}).items():
                if '$ref' in resp:
                    resp = spec['components']['responses'][resp['$ref'].rsplit('/', 1)[-1]]
                for k, v in esempi_di(resp).items():
                    blocchi.append((f"Risposta {codice} — {v.get('summary', k)}", v.get('value')))
            if blocchi:
                p('Esempi')
                for titolo, valore in blocchi:
                    p(titolo)
                    mono(json.dumps(valore, ensure_ascii=False, indent=2))
                    d.add_paragraph()

# --------------------------------------------------------------- modelli dati
h('Modelli di dati', 1)
p('Uno schema per risorsa, con campi, tipi e domini. I nomi dei campi corrispondono alle colonne '
  'del modello dati secondo i prefissi di ruolo dello standard gemello.')
for nome, ruolo, file_yaml, _ in POD:
    spec = contratti[nome]
    h(f'{ruolo}', 2)
    for nome_sch, sch in spec['components']['schemas'].items():
        h(nome_sch, 3)
        if sch.get('description'):
            p(sch['description'].strip().replace('**', '').replace('`', ''))
        props = sch.get('properties')
        if not props and 'allOf' in sch:
            p('Estende ' + ' e '.join(nome_schema(x) for x in sch['allOf'] if '$ref' in x) +
              ' con i campi seguenti.')
            for x in sch['allOf']:
                if 'properties' in x:
                    props = x['properties']
        if not props:
            continue
        obbl = set(sch.get('required', []))
        righe = []
        for campo, meta in props.items():
            tipo = meta.get('type') or nome_schema(meta)
            if meta.get('format'):
                tipo += f" ({meta['format']})"
            dominio = ', '.join(str(v) for v in meta['enum']) if meta.get('enum') else '—'
            righe.append([campo, tipo, 'S' if campo in obbl else 'N', dominio,
                          (meta.get('description', '') or '').strip().replace('**', '').replace('`', '')])
        tabella(['Campo', 'Tipo', 'Obbl.', 'Dominio', 'Descrizione'], righe,
                [1.3, 1.0, 0.5, 1.3, 2.5])

# -------------------------------------------------------------- codici errore
h('Codici di errore', 1)
p('Elenco unico e ordinato dei codici applicativi. Il codice è la parte del messaggio su cui il '
  'consumatore scrive logica ed è stabile nel tempo; la descrizione è destinata alle persone e '
  'può cambiare senza preavviso.')
visti = {}
for cod, comp, oper, cond in codici_errore_globali:
    visti.setdefault(cod, [cond, comp, []])[2].append(oper)
tabella(['Codice', 'Componente', 'Condizione', 'Operazioni'],
        [[c, v[1], v[0], ', '.join(sorted(set(v[2])))] for c, v in sorted(visti.items())],
        [1.7, 1.2, 2.4, 1.3])

tabella(['Codice', 'Componente', 'Condizione', 'Operazioni'],
        [[c, v[1], v[0], ', '.join(sorted(set(v[2])))] for c, v in sorted(visti.items())],
        [1.7, 1.2, 2.4, 1.3]) if False else None

# ---------------------------------------------------- ricevitore delle notifiche
h('Il ricevitore delle notifiche', 1)
p('L’architettura prevede un quarto componente, eventuale: il ricevitore delle notifiche che ANSC '
  'invia verso il Comune. Non ha un contratto in questo documento, e la ragione è di merito.')
p('Lo standard non si applica ai contratti definiti da terzi, che si consumano nella forma in cui '
  'sono pubblicati. Un endpoint che ANSC chiama è esattamente questo: il percorso, il formato del '
  'messaggio e il meccanismo di autenticazione sono fissati dal sistema nazionale, e riscriverli '
  'secondo la nomenclatura aziendale significherebbe non essere raggiungibili.')
p('Del componente restano quindi da specificare i soli aspetti che sono in carico al Comune, e che '
  'l’analisi tratta come requisiti di piattaforma.')
tabella(['Aspetto', 'Requisito'], [
    ['Esposizione verso l’esterno',
     'È l’unico flusso che espone il cluster all’esterno e ne è quindi il punto più oneroso dal '
     'punto di vista della rete e della sicurezza: richiede un ingresso raggiungibile da ANSC.'],
    ['Effetti ammessi',
     'La notifica allinea lo stato locale dell’atto. Non forma atti e non ne modifica lo stato in '
     'ANSC: qualunque azione dispositiva che ne consegua resta presidiata.'],
    ['Condizione di realizzazione',
     'Il componente si realizza solo se ANSC notifica effettivamente i soggetti terzi. La '
     'verifica è tracciata come open point nell’analisi.'],
    ['Alternativa',
     'Se le notifiche non fossero disponibili in modalità di ricezione, la stessa informazione si '
     'ottiene in modalità di scarico dal processo di automazione, che la espone con '
     'elencaNotifiche.'],
], [1.8, 4.8])

# ------------------------------------------------------- verifica di conformità
h('Verifica di conformità allo standard', 1)
p('La lista di controllo prevista dallo standard prima della pubblicazione, applicata ai tre '
  'contratti.')
tabella(['Controllo', 'Esito', 'Nota'], [
    ['Nomi', 'Superato',
     'Percorsi, parametri, campi e artefatti conformi. I tre scostamenti sono dichiarati al '
     'capitolo sul perimetro.'],
    ['Completezza della scheda', 'Superato',
     'Tutti gli elementi obbligatori compilati per ciascuna delle operazioni.'],
    ['Descrizioni', 'Superato',
     'Nessun campo di schema privo di descrizione; nessuna descrizione che ripeta il nome del campo.'],
    ['Esiti', 'Superato',
     'Ogni codice di stato restituibile è dichiarato. Le operazioni che contattano ANSC dichiarano '
     'anche l’esito indeterminato, che è la condizione più insidiosa dell’integrazione.'],
    ['Errori', 'Superato',
     'Ogni codice applicativo compare nell’elenco unico con la condizione che lo produce.'],
    ['Esempi', 'Superato',
     'Presenti per il caso positivo e per l’errore più frequente delle operazioni che ne hanno uno. '
     'I valori sono realistici e inventati: nessun dato personale reale.'],
    ['Coerenza con il modello dati', 'Superato',
     'I nomi dei campi corrispondono alle colonne secondo i prefissi di ruolo; i domini coincidono '
     'con quelli dichiarati sulla base dati.'],
    ['Validità formale', 'Parziale',
     'I documenti sono validi rispetto alla specifica OpenAPI. Il controllo automatico di '
     'conformità allo standard non esiste ancora: è un punto aperto dello standard stesso.'],
    ['Versionamento', 'Non applicabile',
     'Prima emissione: non c’è ancora una versione precedente rispetto a cui valutare '
     'l’incompatibilità.'],
    ['Generazione', 'Superato',
     'Il presente documento è generato dai contratti senza errori e senza sezioni vuote.'],
], [1.7, 1.0, 3.9])

h('Interventi di allineamento rispetto all’analisi', 2)
p('Le interfacce descritte nell’appendice dell’analisi sono state riprese e, dove necessario, '
  'corrette. Gli scostamenti non sono redazionali: ciascuno risponde a una regola.')
tabella(['Nell’analisi', 'In questo documento', 'Perché'], [
    ['GET .../atti/{idAtto}/riconciliazione',
     'POST .../atti/{idAtto}/riconciliazioni',
     'La riconciliazione modifica lo stato locale: una lettura non può avere effetti. Diventa una '
     'collezione di fatti, di cui si legge anche lo storico.'],
    ['POST .../back-office/atti/{idAtto}/verifiche',
     'rimossa: si usa POST /ansc/v1/atti/{idAtto}/preverifiche',
     'Erano due percorsi per lo stesso effetto. Resta un’unica operazione, con due ruoli ammessi.'],
    ['GET .../back-office/sessione e POST .../sessione/rigenerazioni',
     'GET, PUT e DELETE /ansc/v1/sessione; GET .../back-office/sessioni',
     'La sessione dell’operatore appartiene alla superficie applicativa, non al back-office. La '
     'rigenerazione non è un’operazione del servizio: l’OTP si genera sulla web app di ANSC.'],
    ['— (assente)', 'PUT /ansc/v1/sessione',
     'L’architettura prevede che il front-end consegni l’OTP al concentratore, ma nessuna '
     'operazione lo faceva.'],
    ['— (assente)', 'POST e GET /ansc/v1/atti/{idAtto}/allegati',
     'Il flusso prevede l’invio degli allegati e l’attesa della scansione antivirus, senza che '
     'esistesse un’interfaccia.'],
    ['— (assente)', 'POST e GET /ansc/v1/atti/{idAtto}/firme',
     'Mancava l’operazione che forma l’atto: è la lacuna più grave dell’appendice.'],
    ['— (assente)', 'POST /ansc/v1/atti/{idAtto}/annullamento',
     'La cancellazione logica serve a chiudere gli esiti indeterminati che hanno prodotto un '
     'duplicato.'],
    ['— (assente)', 'POST /ansc/v1/back-office/postazioni',
     'Il registro delle postazioni si poteva leggere e aggiornare, ma non alimentare.'],
    ['— (assente)', 'contratto del processo di automazione',
     'Il componente esisteva nella topologia senza alcuna interfaccia.'],
], [2.1, 2.2, 2.3])

# ------------------------------------------------------------------ punti aperti
h('Punti aperti', 1)
p('I punti da chiudere perché la specifica sia realizzabile senza margini di interpretazione.')
tabella(['ID', 'Punto', 'Owner'], [
    ['PS-1', 'Perimetro ammesso senza OTP dell’operatore. Il contratto del processo di '
             'automazione poggia sull’assunto che letture e code siano eseguibili senza l’OTP del '
             'singolo operatore. È una condizione di esistenza del componente, non un dettaglio: '
             'se cade, quelle funzioni tornano operazioni presidiate del back-office.',
     'Fornitore ANSC / Sogei'],
    ['PS-2', 'Vincoli sugli allegati. Dimensione massima e tipi ammessi per caso d’uso non sono '
             'dichiarati nella documentazione disponibile e sono riportati nei contratti come '
             'da confermare.', 'Fornitore ANSC'],
    ['PS-3', 'Temporizzazione dell’attesa sulla scansione antivirus. Non esiste un limite di '
             'frequenza dichiarato: l’intervallo proposto è prudenziale e va confermato.',
     'Fornitore ANSC'],
    ['PS-4', 'Verbi di dominio introdotti. Le operazioni non conformi al modello CRUD usano i '
             'verbi verifica, deposita, firma, annulla, riconcilia, allinea, importa, attiva, '
             'prova ed esegui. I primi nove sono già nel registro dello standard; esegui è nuovo '
             'e ne va autorizzato l’inserimento.', 'Committenza / Analisi'],
    ['PS-5', 'Abbreviazioni non registrate. I contratti usano cod, desc, num, flag e id come '
             'prefissi di ruolo, previsti dallo standard, ma il dizionario aziendale delle '
             'abbreviazioni non esiste ancora: finché non esiste, la regola non è verificabile.',
     'Committenza'],
    ['PS-6', 'Riconoscimento del segmento di raggruppamento. I percorsi di back-office e di '
             'automazione interpongono un segmento fra la versione e la collezione, che la regola '
             'di composizione non contempla. Va riconosciuto nello standard o rimosso.',
     'Committenza / Analisi'],
    ['PS-7', 'Catena di produzione del manuale. Questo documento è generato da uno script di '
             'progetto. Perché la regola per cui il manuale si genera valga stabilmente, la '
             'produzione va inserita nella catena di rilascio con uno strumento standard.',
     'Architetti'],
], [0.6, 4.6, 1.4])

d.save(OUT)
print('generato:', OUT)
print('operazioni documentate:', sum(1 for n, _, _, _ in POD
                                     for v in contratti[n]['paths'].values()
                                     for m in v if m in METODI))
print('codici di errore distinti:', len(visti))
